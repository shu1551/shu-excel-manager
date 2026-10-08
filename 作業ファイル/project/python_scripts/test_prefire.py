# -*- coding: utf-8 -*-
"""先撃ちの登録簿（vbam_prefire）: Excel にも AI にも触らない部分。

2026-10-08: 鍛える回路（vbam_forge）を外したとき、test_forge.py から先撃ち（依頼の本体の入口）のテストだけを移した。
登録簿の読み（頭 2 行）／登録簿つきの先撃ちの計画／列の選び方／形の判定／不変条件の見分け。
"""
import json
import re
import os

import pytest

import vbam_prefire as vp


_MODULE_TEXT = """Attribute VB_Name = "表の整理"
' 表の整理 - …
Sub 表の書き方と罫線と列幅をそろえる()
    ' 依頼の語: これは既定なので読まない
    ' 扱う: x
End Sub
Sub 全列が同じ重複行を削除する()
End Sub

Sub 表の下に合計行を足す()
    ' 依頼の語: 合計|集計
    ' 扱う: 合計 集計
    Dim ws As Worksheet
End Sub

Private Sub 状態語をそろえる()
    ' 何をするか
    ' 依頼の語: 状態|ステータス
    ' 扱う: 判定
End Sub

Sub 頭の無いSub()
    Dim i As Long
End Sub

Sub 壊れた正規表現()
    ' 依頼の語: 合計(
    ' 扱う: 合計
End Sub
"""

_FURI = """Sub 決算統計区分に振り直す()
    ' 依頼の語: 振り直|決算統計|集計
    ' 扱う: 振り直し 集計 区分
    ' 見出し: 予算科目コード 支出額
End Sub
"""

_SHUYAKU = """Sub 課別シートを集約する()
    ' 依頼の語: 集約|まとめ|集計
    ' 扱う: 集約 まとめ
    ' 見出し: 課名 事業名 予算額 執行額
End Sub
"""

_TOTSUGO = """Sub 新システムと突合する()
    ' 依頼の語: 突合|新システム|突き合わせ
    ' 扱う: 突合 新システム突合 伝票番号突合
    ' 見出し: 伝票番号 科目名 支出額
End Sub
"""




def _log(tmp_path, prompt1, replies):
    p = tmp_path / '_last_agent_log.jsonl'
    lines = [json.dumps({'meta': {'book': 'b.xlsx', 'sheet': 's', 'mode': 'sheet', 'request': 'r'}}, ensure_ascii=False)]
    for i, r in enumerate(replies, 1):
        lines.append(json.dumps({'turn': i, 'prompt': prompt1 if i == 1 else '', 'reply': r}, ensure_ascii=False))
    p.write_text("\n".join(lines) + "\n", encoding='utf-8')
    return str(p)


def _snap(rows, r0=5, c0=1):
    return {'row': r0, 'col': c0, 'values': rows}


def _case():
    return {'name': 'n', 'time': 't', 'from_book': 'b.xlsx', 'sheet': 's', 'request': '合計を出して',
            'before': 'x.xlsx', 'expect': _snap([['a'], ['1'], ['1']]), 'before_values': _snap([['a'], ['1'], ['']]),
            'hands': ['write_cells cells={"A7": "=SUM(A6:A6)"}'], 'prefire': ['表の書き方と罫線と列幅をそろえる'], 'passed': False}


def _book(path, rows, sheet='明細', table=None, merges=(), other=None):
    """お題のブックを 1 冊作る（openpyxl）。other＝2 枚目のシートの行。"""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(list(r))
    for m in merges:
        ws.merge_cells(m)
    if table:
        from openpyxl.worksheet.table import Table
        ws.add_table(Table(displayName=table[0], ref=table[1]))
    if other is not None:
        ws2 = wb.create_sheet('別表')
        for r in other:
            ws2.append(list(r))
    wb.save(str(path))
    return str(path)


def _rich_book(tmp_path, name='rich.xlsx'):
    import datetime as dt
    rows = [['課名', '受付日', '金額'],
            ['総務課', dt.datetime(2026, 4, 1), 1000],
            ['財政課', dt.datetime(2026, 4, 2), 2000],
            ['税務課', dt.datetime(2026, 4, 3), 3000]]
    return _book(tmp_path / name, rows, table=('T明細', 'A1:C4'),
                 other=[['課名', '受付日', '金額'], ['総務課', dt.datetime(2026, 5, 1), 9]])


def _zip_book(path, sheet_rels):
    """xlsx の骨組みだけ（workbook.xml と関係ファイル）を書く。sheet_rels＝1 枚目のシートの関係（(種類, 行き先) の並び）。"""
    import zipfile
    R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
    with zipfile.ZipFile(path, 'w') as z:
        z.writestr('xl/workbook.xml', '<workbook xmlns:r="x"><sheets><sheet name="集計" sheetId="1" r:id="rId1"/>'
                                      '<sheet name="明細" sheetId="2" r:id="rId2"/></sheets></workbook>')
        z.writestr('xl/_rels/workbook.xml.rels', f'<Relationships><Relationship Id="rId1" Type="{R}worksheet" Target="worksheets/sheet1.xml"/>'
                                                 f'<Relationship Id="rId2" Type="{R}worksheet" Target="worksheets/sheet2.xml"/></Relationships>')
        rels = ''.join(f'<Relationship Id="rId{i}" Type="{R}{t}" Target="{g}"/>' for i, (t, g) in enumerate(sheet_rels, 1))
        z.writestr('xl/worksheets/_rels/sheet1.xml.rels', f'<Relationships>{rels}</Relationships>')
        z.writestr('xl/drawings/_rels/drawing1.xml.rels', f'<Relationships><Relationship Id="rId1" Type="{R}chart" '
                                                          'Target="../charts/chart1.xml"/></Relationships>')


def test_registry_from_text_reads_header_pairs_only():
    reg = vp.registry_from_text(_MODULE_TEXT, owner='秀コンボ.xlsm')
    assert [e['name'] for e in reg] == ['表の下に合計行を足す', '状態語をそろえる']
    assert reg[0] == {'name': '表の下に合計行を足す', 'owner': '秀コンボ.xlsm', 'ask': '合計|集計', 'handles': ['合計', '集計'],
                      'headers': [], 'select': None, 'shape': [], 'combo': None}
    assert reg[1]['ask'] == '状態|ステータス' and reg[1]['handles'] == ['判定']


def test_python_prefire_reads_the_word_combos_like_the_vba_entry_20260923():
    """弱点 7（2026-09-23）: 「依頼の組」を読んでいたのは VBA の先撃ちだけで、Python 側は素通りしていた。

    組＝+ で区切った組のどれにも 1 語ずつ当たって初めて点が付き、- の語が当たれば撃たない。
    """
    text = "\n".join([
        "Sub 重複行を消す2()",
        "' 依頼の語: 重複",
        "' 依頼の組: 重複,ダブ+消,削除+-一覧,洗い出し",
        "' 扱う: 重複",
        "End Sub",
        "Sub 選んだ列の重複する値を一覧にする()",
        "' 依頼の語: 重複キー",
        "' 依頼の組: 重複,ダブ+一覧,洗い出し",
        "' 扱う: 重複",
        "End Sub",
    ])
    reg = vp.registry_from_text(text)
    assert reg[0]['combo'] == '重複,ダブ+消,削除+-一覧,洗い出し'
    assert vp.combo_score('重複,ダブ+消,削除', '重複行を消して') == (3, False)   # 重複(2)+消(1)＝各組の最長の和
    assert vp.combo_score('重複,ダブ+消,削除', '重複を一覧に') == (0, False)     # 片方の組が当たらない＝0 点
    assert vp.combo_score('重複,ダブ+消,削除+-一覧', '重複を消して一覧に')[1] is True   # +- の組＝除外
    entries = vp.shelf_entries_from_text(text)
    assert vp.shelf_pick('重複行を消して', entries) == '重複行を消す2'
    assert vp.shelf_pick('重複キーを洗い出して一覧に', entries) == '選んだ列の重複する値を一覧にする'  # 「消」の方は除外の語で落ちる
    assert vp.shelf_pick('ピボットテーブルとは何ですか', entries) == ''          # 質問は棚に当てない
    assert vp.shelf_pick('重複を消すマクロを書いて', entries) == ''              # コードの依頼は AI の役目
    p = vp.plan_full('重複キーを洗い出して一覧に', reg)
    assert [e['name'] for e in p['extras']] == ['選んだ列の重複する値を一覧にする']          # 除外の語が効いて取り合いにならない


def test_plan_full_fires_registered_macro_and_removes_its_words_from_ai():
    reg = vp.registry_from_text(_MODULE_TEXT)
    p = vp.plan_full("表を整えて、下に合計を出して", reg)
    assert p['fire'] and p['tidy'] and [e['name'] for e in p['extras']] == ['表の下に合計行を足す']
    assert p['left'] == [] and p['other'] is False            # 合計は登録済みの Sub が片づける＝AI に回さない
    p = vp.plan_full("表を整えて、下に合計を出して、グラフも", reg)
    assert p['left'] == ['グラフ'] and p['other'] is True
    p = vp.plan_full("合計を出して", reg)                       # 整える語が無くても登録簿の語で撃つ
    assert p['fire'] and not p['tidy'] and [e['name'] for e in p['extras']] == ['表の下に合計行を足す']
    p = vp.plan_full("この表を見て", reg)
    assert not p['fire'] and p['extras'] == []


def test_plan_of_keeps_old_shape_without_registry():
    assert vp.plan_of("表を整えて、下に合計を出して") == (True, False, True)
    assert vp.plan_of("合計を出して") == (False, False, False)
    reg = vp.registry_from_text(_MODULE_TEXT)
    assert vp.plan_of("表を整えて、下に合計を出して", reg) == (True, False, False)
    assert vp.plan_of("合計を出して", reg) == (True, False, False)
    assert vp.plan_of("書き方をそろえて。重複行は削除してよい", reg) == (True, True, False)


def test_dedupe_registry_prefers_open_book_over_addin():
    reg = (vp.registry_from_text(_MODULE_TEXT, owner='秀コンボ.xlam')
           + vp.registry_from_text(_MODULE_TEXT, owner='秀コンボ.xlsm'))
    assert len(reg) == 4
    got = vp.dedupe_registry(reg, prefer=['秀コンボ.xlsm', 'Book1'])
    assert [(e['name'], e['owner']) for e in got] == [('表の下に合計行を足す', '秀コンボ.xlsm'), ('状態語をそろえる', '秀コンボ.xlsm')]
    got = vp.dedupe_registry(reg)                               # 開いているブックが無ければ先に読んだ方
    assert [e['owner'] for e in got] == ['秀コンボ.xlam', '秀コンボ.xlam']
    p = vp.plan_full("合計を出して", got)
    assert [e['name'] for e in p['extras']] == ['表の下に合計行を足す']  # 1 回だけ撃つ


def test_plan_full_skips_registered_macro_when_sheet_lacks_its_headers():
    reg = vp.registry_from_text(_FURI, owner='秀コンボ.xlsm')
    assert reg[0]['headers'] == ['予算科目コード', '支出額']
    staff = vp.sheet_words([['担当', '件数'], ['佐藤', 3]])
    p = vp.plan_full("担当ごとに件数を集計して", reg, staff)
    assert p['extras'] == [] and [e['name'] for e in p['skipped']] == ['決算統計区分に振り直す']
    assert '集計' in p['left'] and p['other'] is True              # 集計は AI に回る（片づいた扱いにしない）
    budget = vp.sheet_words([['令和7年度 支出明細', None], [None, None], ['予算科目コード', '支出額 ']])
    p = vp.plan_full("決算統計区分に振り直して集計して", reg, budget)
    assert [e['name'] for e in p['extras']] == ['決算統計区分に振り直す'] and p['other'] is False
    assert vp.plan_full("担当ごとに件数を集計して", reg)['extras']   # 見出しを渡さない（従来の呼び方）は絞らない
    # 実物の見出し（改行・全角の空白・単位の括弧）でも同じ見出しとみなす（2026-09-17 夜）
    real = vp.sheet_words([['予算科目\nコード', '科　目　名', '支出額\n（円）']])
    assert [e['name'] for e in vp.plan_full("決算統計区分に振り直して", reg, real)['extras']] == ['決算統計区分に振り直す']
    assert vp.head_key('支出額(千円)') == '支出額' and vp.head_key('（円）') == '(円)' and vp.head_key('人数（計）') == '人数'
    assert vp.head_key('支出額（税抜）') != '支出額'                     # 単位でない括弧は別の見出し


def test_release_instance_closes_only_that_excel(monkeypatch):
    import vbam_core

    class _WBs:
        Count = 0

    class _XL:
        def __init__(self):
            self.Workbooks = _WBs()
            self.quit = 0

        def Quit(self):
            self.quit += 1
    a, b = _XL(), _XL()
    monkeypatch.setattr(vbam_core, '_created_instances', [{'xl': a, 'pid': 999991}, {'xl': b, 'pid': 999992}])
    monkeypatch.setattr(vbam_core, '_created_xl', b)
    monkeypatch.setattr(vbam_core, '_created_xl_pid', 999992)
    monkeypatch.setattr(vbam_core, '_pid_is_excel', lambda pid: False)
    assert vbam_core.release_instance(999992) == 1
    assert b.quit == 1 and a.quit == 0
    assert [i['pid'] for i in vbam_core._created_instances] == [999991]
    assert vbam_core._created_xl is a and vbam_core._created_xl_pid == 999991
    assert vbam_core.release_instance(None) == 0 and vbam_core.release_instance(12345) == 0


def test_plan_full_longer_match_on_skipped_job_blocks_short_match():
    reg = vp.registry_from_text(_FURI + _SHUYAKU, owner='秀コンボ.xlam')
    shuyaku_sheet = {'課名', '事業名', '予算額', '執行額', '執行率'}
    furi_sheet = {'予算科目コード', '科目名', '支出額'}
    req = "対応表シートで予算科目コードを決算統計区分に振り直して、区分ごとの支出額を集計してください"
    p = vp.plan_full(req, reg, shuyaku_sheet)                     # 振り直しを頼んだのに集約のブック
    assert p['extras'] == [] and [e['name'] for e in p['shadowed']] == ['課別シートを集約する'] and p['other']
    p = vp.plan_full(req, reg, furi_sheet)                        # 本来の表
    assert [e['name'] for e in p['extras']] == ['決算統計区分に振り直す'] and p['shadowed'] == []
    p = vp.plan_full("各課のシートを集約シートにまとめて集計して", reg, shuyaku_sheet)
    assert [e['name'] for e in p['extras']] == ['課別シートを集約する']
    assert vp._ask_strength('集約|まとめ|集計', '全部まとめて集計') == 3 and vp._ask_strength('x(', 'x') == 0


def test_rehearse_drop_recomputes_plan_by_name():
    reg = vp.registry_from_text(_FURI, owner='o')
    p = vp.plan_full("決算統計区分に振り直して", reg, {'予算科目コード', '支出額'})
    dropped = p['extras']
    gone = {x['name'] for x in dropped}
    assert vp.plan_full("決算統計区分に振り直して", [e for e in reg if e['name'] not in gone], {'予算科目コード', '支出額'})['extras'] == []


def test_plan_full_clause_with_ask_word_is_covered():
    reg = vp.registry_from_text(_TOTSUGO, owner='o')
    sheet = {'伝票番号', '科目名', '支出額'}
    for req in ("新システムシートと突き合わせて、伝票番号ごとに突合結果を足し、その右に突合結果ごとの件数を出してください",
                "旧システムと新システムを照合して違いを出して", "新旧の伝票を突き合わせて一致・不一致を出して"):
        p = vp.plan_full(req, reg, sheet)
        assert [e['name'] for e in p['extras']] == ['新システムと突合する'] and p['left'] == [], (req, p['left'])
    p = vp.plan_full("新システムと突合して、結果をグラフにして", reg, sheet)
    assert p['left'] == ['グラフ']


def test_changed_cells_find_the_table_the_macro_wrote():
    # 2026-09-17 夜: 入口の tidy と検査が元の表（帳票・明細）に当たり、マクロが右に作った表を見ていなかった
    before = {'row': 1, 'col': 1, 'values': [['支出日', '課名', None, '課名'], ['2026/4/1', '総務課', None, '旧課']]}
    now = [['支出日', '課名', None, '課名', '4月'], ['2026/4/1', '総務課', None, '総務課', 100.0], [None, None, None, None, None]]
    assert vp._changed_cells(before, now, 1, 1) == [(1, 5), (2, 4), (2, 5)]
    assert vp._changed_cells(before, [['支出日', '課名', None, None]], 1, 1) == []          # 消しただけは数えない
    assert vp._changed_cells({'row': 3, 'col': 2, 'values': [['a']]}, [['a', 'b']], 3, 2) == [(3, 3)]


def test_old_findings_are_cells_the_macro_did_not_change():
    notes = ["「合計」行の C45（550,000円）が上の和と合いません（セル 16,519,000／上の和 5,063,000）",
             "D12: 対応なし の列に空欄があります", "見出しが 2 行あります"]
    assert vp.old_findings(notes, {'D12', 'F2'}) == [notes[0]]                 # 番地の無い指摘は古いと言わない
    assert vp.old_findings(notes, set()) == [notes[0], notes[1]]


def test_plan_full_generic_word_on_skipped_job_does_not_shadow():
    reg = vp.registry_from_text("Sub 表の下に合計行を足す()\n    ' 依頼の語: 合計|集計\n    ' 扱う: 合計 集計\n    ' 見出し: 品名 数量 金額\nEnd Sub\n"
                                "Sub 課別シートを集約する()\n    ' 依頼の語: 集約|各課\n    ' 扱う: 集約 合計\n    ' 見出し: 課名 事業名\nEnd Sub\n")
    p = vp.plan_full("各課シートの事業を集めて合計も出して", reg, {'課名', '事業名'})
    assert [e['name'] for e in p['extras']] == ['課別シートを集約する'] and p['other'] is False


def test_formula_copy_blockers_absolute_sum_per_column_is_fine():
    import vbam_view as vv
    cells = {(9, 5): ('=SUM($E$2:$E$8)', '=SUM(R2C5:R8C5)'), (9, 6): ('=SUM($F$2:$F$8)', '=SUM(R2C6:R8C6)'),
             (9, 7): ('=SUM($G$2:$G$8)', '=SUM(R2C7:R8C7)')}
    assert vv.formula_copy_blockers(cells) == []


def test_plan_full_graph_or_sort_in_same_clause_goes_to_ai():
    reg = vp.registry_from_text("Sub 決算統計区分を振り直して集計する()\n    ' 依頼の語: 決算統計|振り直\n    ' 扱う: 振り直し 区分 集計\n"
                                "    ' 見出し: 予算科目コード\nEnd Sub\n")
    p = vp.plan_full("決算統計区分に振り直して、区分ごとの円グラフも作って", reg, {'予算科目コード'})
    assert p['extras'] and p['other'] is True and 'グラフ' in p['left']
    p = vp.plan_full("決算統計区分に振り直して、区分ごとに集計して", reg, {'予算科目コード'})
    assert p['extras'] and p['other'] is False


def test_formula_copy_blockers_table_totals_row_is_fine():
    import vbam_view as vv
    cells = {(14, 4): ('=SUBTOTAL(109,[基本給])', '=SUBTOTAL(109,[基本給])'), (14, 5): ('=SUBTOTAL(109,[手当])', '=SUBTOTAL(109,[手当])'),
             (14, 6): ('=SUBTOTAL(109,[支給額])', '=SUBTOTAL(109,[支給額])')}
    assert vv.formula_copy_blockers(cells) == []


def test_select_line_and_column_pick():
    """' 選ぶ列: N ＝列を選ぶ仕事。入口は依頼の語と見出しが一致する列を、依頼に出てくる順に選ぶ（2026-09-18）。"""
    code = ("Sub ピボットを作る()\n    ' 依頼の語: ピボットにして\n    ' 扱う: ピボット\n    ' 見出し: なし\n"
            "    ' 選ぶ列: 3\nEnd Sub\n")
    reg = vp.registry_from_text(code)
    assert reg[0]['select'] == (3, 3)
    assert vp.registry_from_text(code.replace('選ぶ列: 3', '選ぶ列: 2-3'))[0]['select'] == (2, 3)
    cells = [(1, '支出日'), (2, '所属'), (3, '費目'), (4, '摘要'), (5, '金額')]
    assert vp.pick_columns('所属と費目で金額のピボットにして', cells, (3, 3)) == [2, 3, 5]
    assert vp.pick_columns('金額を所属ごとに', cells, (3, 3)) is None        # 2 列しか当たらない＝入口は撃たない
    assert vp.pick_columns('金額を所属ごとに', cells, (2, 3)) == [5, 2]      # 依頼に出てくる順
    assert vp.pick_columns('所属（部）と 金 額 で', [(1, '所属'), (2, '金額')]) == [1, 2]   # 空白・全角のならしも通す


def test_request_column_macros():
    """' 列は頼みから: のある棚は列を選ばずに撃つ（マクロが頼みの文の見出しで行・列・値を決める・2026-10-09）。"""
    code = ("Sub 選んだ3列でピボットを作る()\n    ' 依頼の語: ピボットで\n    ' 選ぶ列: 3\n    ' 列は頼みから: 行 列 値\nEnd Sub\n"
            "Sub 選んだ2列で構成比ピボットを作る()\n    ' 依頼の語: 構成比\n    ' 選ぶ列: 2\n    Dim x\n"
            "    ' 列は頼みから: 本文の中のコメントは数えない\nEnd Sub\n")
    assert vp.request_column_macros(code) == {'選んだ3列でピボットを作る'}
    assert vp.request_column_macros('') == set()


def test_pick_columns_prefers_left_when_header_repeats():
    """2026-09-18: 前の集計表が右にあると同じ見出しが 2 つ並び、入口が列を決められなかった＝左を採る。"""
    cells = [(1, '支出日'), (2, '事業名'), (5, '支出額'), (7, '事業名'), (8, '4月')]
    assert vp.pick_columns('事業名と支出日と支出額で作って', cells, (3, 3)) == [2, 1, 5]


def test_pick_columns_ignores_right_block_when_too_many_hit():
    """2026-09-18 朝: 前の集計表が右に残ると、依頼文の「4月」「3月」「合計」がその見出しに当たって列が決まらず
    （当たりすぎで None）、月次集計が入口から撃てなかった＝空の列で切れた左の塊だけで選び直す。"""
    cells = [(1, '伝票日付'), (2, '所属名'), (3, '節'), (4, '備考'), (5, '支払額'),
             (7, '所属名'), (8, '4月'), (19, '3月'), (20, '合計')]
    req = '所属名と伝票日付と支払額で 明細から月別の集計表を作る。4月から3月・合計も'
    assert vp.pick_columns(req, cells, (3, 3)) == [2, 1, 5]


def test_inv_ignores_filtered_hidden_rows_after_run():
    """2026-09-18 朝: 「0 の行は隠して」と頼んで絞り込んだのに「隠れていた行が変わった」で AI に回っていた。
    撃った後に絞り込みが掛かっていれば、隠れた行は見せ方＝咎めない（テーブルの絞り込みも見る）。"""
    import vbam_inv as vi
    b = {'sheets': {'S': {'used': 'A1:B9', 'values': [['a', 1]], 'formulas': {}, 'hidden': (), 'cf': 0,
                          'validation': '', 'filter': '', 'merged': False, 'tables': {}, 'charts': 0}},
         'order': ['S'], 'names': [], 'calc': None}
    import copy
    a = copy.deepcopy(b)
    a['sheets']['S']['hidden'] = (3, 5)
    a['sheets']['S']['filter'] = 'A1:B9'          # テーブルの絞り込みで隠れた
    assert not [x for x in vi._inv_compare(b, a, 'S', ()) if '隠れていた行' in x]
    a2 = copy.deepcopy(b)
    a2['sheets']['S']['hidden'] = (3, 5)          # 手で隠した（絞り込み無し）＝咎める
    assert [x for x in vi._inv_compare(b, a2, 'S', ()) if '隠れていた行' in x]


def test_inv_table_renamed_is_not_gone():
    """2026-09-18 朝: 「テーブル1」を T職員名簿 に直すのが仕事なのに「テーブルが消えた」で AI に回っていた。
    同じ左上に別の名前のテーブルがあれば、消えたのでなく名前を付け替えた。"""
    import copy
    import vbam_inv as vi
    b = {'sheets': {'S': {'used': 'A1:E11', 'values': [['a']], 'formulas': {}, 'hidden': (), 'cf': 0,
                          'validation': '', 'filter': '', 'merged': False, 'charts': 0,
                          'tables': {'テーブル1': 'A1:E11'}}},
         'order': ['S'], 'names': [], 'calc': None}
    a = copy.deepcopy(b)
    a['sheets']['S']['tables'] = {'T職員名簿': 'A1:E12'}          # 名前を付け替えて集計行を足した
    assert not [x for x in vi._inv_compare(b, a, 'S', ()) if 'テーブル' in x]
    a2 = copy.deepcopy(b)
    a2['sheets']['S']['tables'] = {}                              # 本当に消えた＝咎める
    assert [x for x in vi._inv_compare(b, a2, 'S', ()) if 'テーブル' in x and '消えた' in x]


def test_plan_full_skips_a_job_whose_shape_is_missing():
    """入口: 依頼の語が当たっても、表の形がそろわなければ撃たない（語のかすり当たりの誤爆を形で先に止める）。"""
    code = ("Sub テーブルを絞り込む()\n    ' 依頼の語: 絞り込\n    ' 扱う: 絞り込み\n    ' 見出し: なし\n"
            "    ' 形: テーブル\nEnd Sub\n")
    reg = vp.registry_from_text(code)
    assert reg[0]['shape'] == ['テーブル']
    p = vp.plan_full('この表を絞り込んで', reg, None, {'数の列'})
    assert not p['extras'] and [e['name'] for e in p['skipped']] == ['テーブルを絞り込む']
    assert p['skipped'][0]['why_shape'] == ['テーブル']
    assert [e['name'] for e in vp.plan_full('この表を絞り込んで', reg, None, {'数の列', 'テーブル'})['extras']] \
        == ['テーブルを絞り込む']
    # 形を読めなかったとき（shape=None）と、形の行が無い Sub は従来どおり撃つ
    assert vp.plan_full('この表を絞り込んで', reg, None, None)['extras']


def test_shape_signals_pivot_and_chart_words():
    import vbam_prefire as vp
    k = ['ss', 'sn', 'sn']
    v = [['課', '額'], ['a', '1'], ['b', '2']]
    assert {'ピボット', 'グラフ'} <= vp.shape_signals(v, k, has_pivot=True, has_chart=True)
    assert not ({'ピボット', 'グラフ'} & vp.shape_signals(v, k))
