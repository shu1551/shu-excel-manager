# -*- coding: utf-8 -*-
"""鍛える回路（vbam_forge・2026-09-17）と先撃ちの登録簿（vbam_prefire）: Excel にも AI にも触らない部分。

拾う（記録 → AI の手・先撃ちの名前）／前→後の差の要約／返事からコードを取り出す／規則の検査／
値の突き合わせ／.bas の書き出し（cp932）／登録簿の読み（頭 2 行）／登録簿つきの先撃ちの計画。
本物（写しで撃つ・登録）は agent --forge で。
"""
import json
import re
import os

import pytest

import vbam_forge as vf
import vbam_prefire as vp


# ----------------------------------------------------------------
# 拾う
# ----------------------------------------------------------------

def _log(tmp_path, prompt1, replies):
    p = tmp_path / '_last_agent_log.jsonl'
    lines = [json.dumps({'meta': {'book': 'b.xlsx', 'sheet': 's', 'mode': 'sheet', 'request': 'r'}}, ensure_ascii=False)]
    for i, r in enumerate(replies, 1):
        lines.append(json.dumps({'turn': i, 'prompt': prompt1 if i == 1 else '', 'reply': r}, ensure_ascii=False))
    p.write_text("\n".join(lines) + "\n", encoding='utf-8')
    return str(p)


def test_hands_of_log_skips_reads_and_reads_prefire_names(tmp_path):
    prompt = ("材料…\n--- 道具が先に撃ったマクロ（AI より先に・1.2 秒） ---\n"
              "表を整える → 重複行を消す を実行し、tidy で仕上げました。")
    reply = json.dumps({'say': 'x', 'actions': [
        {'op': 'read', 'range': 'A1:C3'},
        {'op': 'normalize', 'range': 'D6:D45', 'rules': ['phone'], 'overwrite': True},
        {'op': 'find_replace', 'range': 'H6:H45', 'search': 'Active', 'replace': '進行中', 'overwrite': True},
        {'op': 'write_cells', 'cells': {'G47': '=SUM(G6:G45)'}},
    ]}, ensure_ascii=False)
    hands, prefire = vf._hands_of_log(_log(tmp_path, prompt, [reply, '{"say":"done","actions":[],"done":true}']))
    assert prefire == ['表を整える', '重複行を消す']
    assert len(hands) == 3 and hands[0].startswith('normalize range=D6:D45') and 'overwrite' not in hands[0]
    assert 'Active' in hands[1] and 'SUM' in hands[2]


def test_hands_of_log_without_prefire_or_file(tmp_path):
    hands, prefire = vf._hands_of_log(str(tmp_path / 'none.jsonl'))
    assert hands == [] and prefire == []
    hands, prefire = vf._hands_of_log(_log(tmp_path, '材料', ['{"actions":[{"op":"eval","expr":"1"}]}']))
    assert hands == [] and prefire == []


# ----------------------------------------------------------------
# 前 → 後の差・表の頭
# ----------------------------------------------------------------

def _snap(rows, r0=5, c0=1):
    return {'row': r0, 'col': c0, 'values': rows}


def test_diff_summary_groups_by_column_and_counts_added_rows():
    before = _snap([['品名', '数量', '金額'], ['りんご', '2', '200'], ['みかん', '3', '150']])
    after = _snap([['品名', '数量', '金額'], ['りんご', '2', '200'], ['みかん', '3', '150'], ['合計', '5', '350']])
    text, total = vf._diff_summary(before, after)
    assert total == 3 and '行が増減' in text
    assert "列 A「品名」: 1 個  A8: ''→'合計'" in text
    assert "列 C「金額」: 1 個  C8: ''→'350'" in text


def test_diff_summary_counts_added_columns():
    before = _snap([['科目', '金額'], ['給料', '100']], r0=1)
    after = _snap([['科目', '金額', '区分', '', '区分'], ['給料', '100', '人件費', '', '人件費']], r0=1)
    text, total = vf._diff_summary(before, after)
    assert total == 4 and '列が増減' in text and '行が増減' not in text
    assert "列 C「区分」: 2 個" in text and "C2: ''→'人件費'" in text


def test_diff_summary_examples_are_capped():
    before = _snap([['電話']] + [[f'0{i}'] for i in range(20)])
    after = _snap([['電話']] + [[f'03-{i}'] for i in range(20)])
    text, total = vf._diff_summary(before, after)
    assert total == 20
    assert text.count('→') == vf._MAX_EXAMPLES


def test_table_text_shows_head_and_remaining_count():
    snap = _snap([['a', 'b']] + [[str(i), ''] for i in range(30)], r0=3, c0=2)
    t = vf._table_text(snap, max_rows=4)
    assert t.startswith('（左上 B3・31 行×2 列）') and '\n3\ta\tb' in t and '…（残り 27 行）' in t


# ----------------------------------------------------------------
# AI の返事 → コード → 規則
# ----------------------------------------------------------------

_GOOD = """Sub 合計行を足す()
    ' 依頼の語: 合計|集計
    ' 扱う: 合計 集計
    ' 見出し: a
    ' 本文の下に合計の行を足す
    Dim ws As Worksheet
    Set ws = ActiveSheet
    ws.Cells(10, 3).Formula = "=SUM(C2:C9)"
End Sub"""


def test_extract_code_strips_fences_and_prose():
    assert vf._extract_code("はい。\n```vba\n" + _GOOD + "\n```\n以上です") == _GOOD + "\n"
    assert vf._extract_code(_GOOD) == _GOOD + "\n"
    assert vf._extract_code("Sub x()\n何か") is None
    assert vf._extract_code("") is None


def test_validate_code_reads_header_and_rejects_rule_breaks():
    name, ask, handles = vf._validate_code(_GOOD + "\n")
    assert (name, ask, handles) == ('合計行を足す', '合計|集計', ['合計', '集計'])
    assert vf._validate_code(None)[0] is None
    assert 'Function' in vf._validate_code("Function f()\nEnd Function\n" + _GOOD)[1]
    assert '引数' in vf._validate_code(_GOOD.replace('合計行を足す()', '合計行を足す(n As Long)'))[1]
    assert 'Option Explicit' in vf._validate_code("Option Explicit\n" + _GOOD)[1]
    assert '頭の 2 行' in vf._validate_code(_GOOD.replace("    ' 扱う: 合計 集計\n", ''))[1]
    assert '既にある' in vf._validate_code(_GOOD.replace('合計行を足す', '表を整える'))[1]
    assert 'MsgBox' in vf._validate_code(_GOOD.replace('Set ws = ActiveSheet', 'MsgBox "x"'))[1]
    assert '正規表現' in vf._validate_code(_GOOD.replace('合計|集計', '合計('))[1]
    assert '1 本だけ' in vf._validate_code(_GOOD + "\n" + _GOOD.replace('合計行を足す', '別'))[1]


def test_write_bas_is_cp932_with_attribute_and_crlf(tmp_path):
    p = str(tmp_path / 'x.bas')
    assert vf._write_bas(p, '鍛冶_試し', "Sub a()\n    ' 依頼の語: x\nEnd Sub\n") is None
    raw = open(p, 'rb').read()
    assert raw.startswith('Attribute VB_Name = "鍛冶_試し"\r\n'.encode('cp932')) and b'\n' not in raw.replace(b'\r\n', b'')
    err = vf._write_bas(p, 'm', "Sub a()\n    ' 依頼の語: 𠮷\nEnd Sub\n")
    assert err and 'cp932' in err


# ----------------------------------------------------------------
# 突き合わせ
# ----------------------------------------------------------------

def test_compare_passes_on_equal_values_and_number_looks():
    e = _snap([['a', '1,320'], ['b', '1320.0']])
    g = _snap([['a', '1320'], ['b', '1320']])
    assert vf._compare(e, g) == []


def test_compare_reports_cells_and_size():
    e = _snap([['a', '1'], ['b', '2']])
    g = _snap([['a', '1'], ['b', '3']])
    out = vf._compare(e, g)
    assert out == ["B6: 期待 '2' ／ マクロ後 '3'"]
    out = vf._compare(e, _snap([['a', '1']]))
    assert out[0].startswith('表の大きさ／位置が違う') and any(x.startswith('A6') for x in out)


# ----------------------------------------------------------------
# 台帳・一覧・プロンプト
# ----------------------------------------------------------------

def _case():
    return {'name': 'n', 'time': 't', 'from_book': 'b.xlsx', 'sheet': 's', 'request': '合計を出して',
            'before': 'x.xlsx', 'expect': _snap([['a'], ['1'], ['1']]), 'before_values': _snap([['a'], ['1'], ['']]),
            'hands': ['write_cells cells={"A7": "=SUM(A6:A6)"}'], 'prefire': ['表を整える'], 'passed': False}


def test_prompt_contains_rules_request_hands_prefire_and_feedback():
    c = _case()
    p = vf._prompt(c)
    assert p.startswith(vf.FORGE_RULES) and '合計を出して' in p and 'write_cells' in p and '表を整える' in p
    assert '前回のマクロの不一致' not in p
    p2 = vf._prompt(c, "A7: 期待 '1' ／ マクロ後 ''", "Sub x()\nEnd Sub")
    assert '前回のマクロの不一致' in p2 and 'Sub x()' in p2


def test_forge_ledger_roundtrip_and_list(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(vf, '_AGENT_FORGE_FILE', str(tmp_path / '_agent_forge.json'))
    assert vf._forge_load() == {}
    assert vf.forged_list() is True
    assert '鍛えたマクロはありません' in capsys.readouterr().out
    c = _case()
    c.update({'passed': True, 'sub': '合計行を足す', 'ask': '合計', 'registered_to': 'b.xlsx'})
    assert vf._forge_save({'n': c})
    assert vf._forge_load()['n']['sub'] == '合計行を足す'
    assert vf.forged_list() is True
    out = capsys.readouterr().out
    assert '合格' in out and 'Sub 合計行を足す' in out and '登録: b.xlsx' in out


def test_forge_dry_run_prints_prompt_without_ai(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(vf, '_AGENT_FORGE_FILE', str(tmp_path / '_agent_forge.json'))
    c = _case()
    bed = tmp_path / 'x.xlsx'
    bed.write_bytes(b'x')
    c['before'] = str(bed)
    vf._forge_save({'n': c})
    monkeypatch.setattr(vf, '_ask_once', lambda *a, **k: (_ for _ in ()).throw(AssertionError('AI を呼んだ')))
    assert vf.forge('n', dry_run=True) is True
    out = capsys.readouterr().out
    assert '--dry-run' in out and vf.FORGE_RULES[:20] in out


def test_forge_loop_feeds_mismatch_back_and_records_pass(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(vf, '_AGENT_FORGE_FILE', str(tmp_path / '_agent_forge.json'))
    monkeypatch.setattr(vf, '_AGENT_FORGE_DIR', str(tmp_path))
    c = _case()
    bed = tmp_path / 'x.xlsx'
    bed.write_bytes(b'x')
    c['before'] = str(bed)
    vf._forge_save({'n': c})
    replies = iter(["文だけで Sub が無い返事", _GOOD, _GOOD])
    prompts = []

    def fake_ask(ai, model, prompt):
        prompts.append(prompt)
        return next(replies)
    tries = iter([(["A7: 期待 '1' ／ マクロ後 ''"], 0.1), ([], 0.1)])
    monkeypatch.setattr(vf, '_ask_once', fake_ask)
    monkeypatch.setattr(vf, '_try_macro', lambda case, bas, sub: next(tries))
    assert vf._check_bas.__name__ == '_check_bas'              # 本物の check-bas を通す（差し替えない）
    assert vf.forge('n', max_turns=3) is True
    rec = vf._forge_load()['n']
    assert rec['passed'] and rec['sub'] == '合計行を足す' and rec['turns'] == 3 and rec['mismatch'] == 0
    assert os.path.isfile(rec['bas'])
    assert '規則違反' in prompts[1] and '前回のマクロの不一致' in prompts[2]
    raw = open(rec['bas'], 'rb').read()
    assert raw.startswith('Attribute VB_Name = "表の整理"'.encode('cp932'))    # 合格後は登録用の名前
    out = capsys.readouterr().out
    assert '合格' in out and '--register' in out


def test_forge_gives_up_after_max_turns(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(vf, '_AGENT_FORGE_FILE', str(tmp_path / '_agent_forge.json'))
    monkeypatch.setattr(vf, '_AGENT_FORGE_DIR', str(tmp_path))
    c = _case()
    bed = tmp_path / 'x.xlsx'
    bed.write_bytes(b'x')
    c['before'] = str(bed)
    vf._forge_save({'n': c})
    monkeypatch.setattr(vf, '_ask_once', lambda *a: _GOOD)
    monkeypatch.setattr(vf, '_try_macro', lambda case, bas, sub: (["A7: 期待 '1' ／ マクロ後 ''"], 0.1))
    assert vf.forge('n', max_turns=2) is False
    rec = vf._forge_load()['n']
    assert not rec['passed'] and rec['turns'] == 2 and rec['mismatch'] == 1
    assert '不合格' in capsys.readouterr().out


# ----------------------------------------------------------------
# 登録簿（vbam_prefire）
# ----------------------------------------------------------------

_MODULE_TEXT = """Attribute VB_Name = "表の整理"
' 表の整理 - …
Sub 表を整える()
    ' 依頼の語: これは既定なので読まない
    ' 扱う: x
End Sub
Sub 重複行を消す()
End Sub

Sub 合計行を足す()
    ' 依頼の語: 合計|集計
    ' 扱う: 合計 集計
    Dim ws As Worksheet
End Sub

Private Sub 状態語をそろえる()
    ' 何をするか
    ' 依頼の語: 状態|ステータス
    ' 扱う: 判定
End Sub

Sub 頭の無いSub()
    Dim i As Long
End Sub

Sub 壊れた正規表現()
    ' 依頼の語: 合計(
    ' 扱う: 合計
End Sub
"""


def test_registry_from_text_reads_header_pairs_only():
    reg = vp.registry_from_text(_MODULE_TEXT, owner='秀コンボ.xlsm')
    assert [e['name'] for e in reg] == ['合計行を足す', '状態語をそろえる']
    assert reg[0] == {'name': '合計行を足す', 'owner': '秀コンボ.xlsm', 'ask': '合計|集計', 'handles': ['合計', '集計'],
                      'headers': [], 'select': None, 'shape': []}
    assert reg[1]['ask'] == '状態|ステータス' and reg[1]['handles'] == ['判定']


def test_plan_full_fires_registered_macro_and_removes_its_words_from_ai():
    reg = vp.registry_from_text(_MODULE_TEXT)
    p = vp.plan_full("表を整えて、下に合計を出して", reg)
    assert p['fire'] and p['tidy'] and [e['name'] for e in p['extras']] == ['合計行を足す']
    assert p['left'] == [] and p['other'] is False            # 合計は登録済みの Sub が片づける＝AI に回さない
    p = vp.plan_full("表を整えて、下に合計を出して、グラフも", reg)
    assert p['left'] == ['グラフ'] and p['other'] is True
    p = vp.plan_full("合計を出して", reg)                       # 整える語が無くても登録簿の語で撃つ
    assert p['fire'] and not p['tidy'] and [e['name'] for e in p['extras']] == ['合計行を足す']
    p = vp.plan_full("この表を見て", reg)
    assert not p['fire'] and p['extras'] == []


def test_plan_of_keeps_old_shape_without_registry():
    assert vp.plan_of("表を整えて、下に合計を出して") == (True, False, True)
    assert vp.plan_of("合計を出して") == (False, False, False)
    reg = vp.registry_from_text(_MODULE_TEXT)
    assert vp.plan_of("表を整えて、下に合計を出して", reg) == (True, False, False)
    assert vp.plan_of("合計を出して", reg) == (True, False, False)
    assert vp.plan_of("書き方をそろえて。重複行は削除してよい", reg) == (True, True, False)


# ----------------------------------------------------------------
# 2026-09-17 夕: 人の正解から鍛える・別の表で試す・登録先の名指し・登録簿の二重撃ち
# ----------------------------------------------------------------

def test_dedupe_registry_prefers_open_book_over_addin():
    reg = (vp.registry_from_text(_MODULE_TEXT, owner='秀コンボ.xlam')
           + vp.registry_from_text(_MODULE_TEXT, owner='秀コンボ.xlsm'))
    assert len(reg) == 4
    got = vp.dedupe_registry(reg, prefer=['秀コンボ.xlsm', 'Book1'])
    assert [(e['name'], e['owner']) for e in got] == [('合計行を足す', '秀コンボ.xlsm'), ('状態語をそろえる', '秀コンボ.xlsm')]
    got = vp.dedupe_registry(reg)                               # 開いているブックが無ければ先に読んだ方
    assert [e['owner'] for e in got] == ['秀コンボ.xlam', '秀コンボ.xlam']
    p = vp.plan_full("合計を出して", got)
    assert [e['name'] for e in p['extras']] == ['合計行を足す']  # 1 回だけ撃つ


def test_pick_register_target_names_or_single_owner():
    books = [('秀コンボ.xlsm', True), ('職場の表.xlsx', False)]
    assert vf._pick_register_target(books) == ('秀コンボ.xlsm', None)          # アクティブがどちらでも持ち主へ
    assert vf._pick_register_target(books, to='秀コンボ') == ('秀コンボ.xlsm', None)
    assert vf._pick_register_target(books, to='職場の表.xlsx') == ('職場の表.xlsx', None)
    name, why = vf._pick_register_target(books, to='無い.xlsm')
    assert name is None and '開いていません' in why
    name, why = vf._pick_register_target([('職場の表.xlsx', False)])
    assert name is None and '--to' in why
    name, why = vf._pick_register_target([('a.xlsm', True), ('b.xlsm', True)])
    assert name is None and '2 冊' in why


def test_prompt_without_hands_uses_truth_and_shows_other_sheets():
    c = _case()
    c.update({'hands': [], 'truth': 't.xlsx', 'tests': [{'label': 'x'}]})
    c['before_values'] = dict(c['before_values'], others=[{'name': '対応表', 'head': '（左上 A1）\n1\t科目\t区分'}])
    p = vf._prompt(c)
    assert '記録はありません' in p and '人が直した正解' in p and '別の表 1 枚' in p
    assert 'シート「対応表」' in p and '科目\t区分' in p


def test_tests_feedback_names_table_and_caps_lines():
    c = _case()
    c['tests'] = [{'label': '明細_2', 'expect': _snap([['区分', '金額'], ['人件費', '100']])}]
    mism = [f"A{i}: 期待 'x' ／ マクロ後 ''" for i in range(30)]
    text = vf._tests_feedback(c, [('明細_2', mism)])
    assert '別の表「明細_2」で不一致 30 件' in text and '…ほか' in text and '人件費' in text
    assert text.count('期待') == vf._MAX_TEST_MISMATCH


def test_add_tests_needs_pairs():
    assert vf.add_tests(_case(), ['a.xlsx']) is None


def test_forge_pair_retries_when_other_table_fails(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(vf, '_AGENT_FORGE_FILE', str(tmp_path / '_agent_forge.json'))
    monkeypatch.setattr(vf, '_AGENT_FORGE_DIR', str(tmp_path / 'forge'))
    for n in ('b.xlsx', 't.xlsx', 'b2.xlsx', 't2.xlsx'):
        (tmp_path / n).write_bytes(b'x')
    snaps = {'b': _snap([['a'], ['1']]), 't': _snap([['a'], ['1'], ['1']])}

    def fake_read(items, fallback_first=False):
        return [dict(snaps['t' if os.path.basename(p).split('_')[-1].startswith('truth') else 'b'], sheet='s', others=[])
                for p, _s in items]
    monkeypatch.setattr(vf, '_read_sheets', fake_read)
    replies = iter([_GOOD, _GOOD])
    prompts = []
    monkeypatch.setattr(vf, '_ask_once', lambda ai, model, prompt: (prompts.append(prompt), next(replies))[1])
    monkeypatch.setattr(vf, '_try_macro', lambda case, bas, sub: ([], 0.1))
    tries = iter([([('b2', ["A7: 期待 '1' ／ マクロ後 ''"])], 0.1), ([('b2', [])], 0.1)])
    monkeypatch.setattr(vf, '_try_tests', lambda case, bas, sub: next(tries))
    ok = vf.forge('n', before=str(tmp_path / 'b.xlsx'), truth=str(tmp_path / 't.xlsx'), request='合計を出して',
                  tests=[str(tmp_path / 'b2.xlsx'), str(tmp_path / 't2.xlsx')], max_turns=3)
    assert ok is True
    rec = vf._forge_load()['n']
    assert rec['passed'] and rec['tests_passed'] and rec['turns'] == 2 and rec['truth'] and rec['hands'] == []
    assert len(rec['tests']) == 1 and rec['tests'][0]['label'] == 'b2'
    assert '別の表「b2」で不一致' in prompts[1]
    # 合格済みに別の表を足すと、印を消して撃ち直す（AI は外れたときだけ）
    (tmp_path / 'b3.xlsx').write_bytes(b'x')
    (tmp_path / 't3.xlsx').write_bytes(b'x')
    retest = []
    monkeypatch.setattr(vf, '_try_tests', lambda case, bas, sub: (retest.append(len(case['tests'])), ([('b2', []), ('b3', [])], 0.1))[1])
    assert vf.forge('n', tests=[str(tmp_path / 'b3.xlsx'), str(tmp_path / 't3.xlsx')]) is True
    assert retest == [2] and vf._forge_load()['n']['tests_passed'] is True
    assert os.path.isfile(rec['before']) and rec['before'] != str(tmp_path / 'b.xlsx')     # 原本でなく写し
    assert '別の表 1 枚: 合格 0・外れ 1' in capsys.readouterr().out


def test_forge_before_without_truth_is_refused(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(vf, '_AGENT_FORGE_FILE', str(tmp_path / '_agent_forge.json'))
    assert vf.forge('n', before='x.xlsx') is False
    assert '--truth' in capsys.readouterr().out


def test_prefire_names_follow_request_words():
    assert vf._prefire_names('科目を振り直して区分ごとに集計して') == []
    assert vf._prefire_names('書き方をそろえて、合計を出して') == ['表を整える']
    assert vf._prefire_names('書き方をそろえて。重複行は削除してよい') == ['表を整える', '重複行を消す']


_FURI = """Sub 決算統計区分に振り直す()
    ' 依頼の語: 振り直|決算統計|集計
    ' 扱う: 振り直し 集計 区分
    ' 見出し: 予算科目コード 支出額
End Sub
"""


def test_plan_full_skips_registered_macro_when_sheet_lacks_its_headers():
    reg = vp.registry_from_text(_FURI, owner='秀コンボ.xlsm')
    assert reg[0]['headers'] == ['予算科目コード', '支出額']
    staff = vp.sheet_words([['担当', '件数'], ['佐藤', 3]])
    p = vp.plan_full("担当ごとに件数を集計して", reg, staff)
    assert p['extras'] == [] and [e['name'] for e in p['skipped']] == ['決算統計区分に振り直す']
    assert '集計' in p['left'] and p['other'] is True              # 集計は AI に回る（片づいた扱いにしない）
    budget = vp.sheet_words([['令和7年度 支出明細', None], [None, None], ['予算科目コード', '支出額 ']])
    p = vp.plan_full("決算統計区分に振り直して集計して", reg, budget)
    assert [e['name'] for e in p['extras']] == ['決算統計区分に振り直す'] and p['other'] is False
    assert vp.plan_full("担当ごとに件数を集計して", reg)['extras']   # 見出しを渡さない（従来の呼び方）は絞らない
    # 実物の見出し（改行・全角の空白・単位の括弧）でも同じ見出しとみなす（2026-09-17 夜）
    real = vp.sheet_words([['予算科目\nコード', '科　目　名', '支出額\n（円）']])
    assert [e['name'] for e in vp.plan_full("決算統計区分に振り直して", reg, real)['extras']] == ['決算統計区分に振り直す']
    assert vp.head_key('支出額(千円)') == '支出額' and vp.head_key('（円）') == '(円)' and vp.head_key('人数（計）') == '人数'
    assert vp.head_key('支出額（税抜）') != '支出額'                     # 単位でない括弧は別の見出し


def test_check_fit_needs_headers_present_and_request_hit():
    c = _case()
    c['before_values'] = _snap([['予算科目コード', '支出額'], ['02-01', '100']])
    c['request'] = '決算統計区分に振り直して'
    c['tests'] = [{'label': 'B', 'expect': _snap([['表題', ''], ['予算科目コード', '支出額']])}]
    assert vf._check_fit(c, _FURI, '振り直|決算統計') == (['予算科目コード', '支出額'], None)   # 集計は一般の依頼（課別の集計表を作って）に当たる
    heads, why = vf._check_fit(c, _FURI.replace('支出額', '金額'), '振り直')
    assert heads is None and '金額' in why and '撃つ前の表' in why
    c['tests'] = [{'label': 'B', 'expect': _snap([['予算科目コード', '執行額']])}]
    heads, why = vf._check_fit(c, _FURI, '振り直')
    assert heads is None and '別の表「B」' in why
    c['tests'] = []
    assert '当たりません' in vf._check_fit(c, _FURI, '合計行')[1]
    assert "見出し: …" in vf._check_fit(c, _FURI.replace("    ' 見出し: 予算科目コード 支出額\n", ''), '振り直')[1]


def test_forge_stops_when_same_cells_keep_failing(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(vf, '_AGENT_FORGE_FILE', str(tmp_path / '_agent_forge.json'))
    monkeypatch.setattr(vf, '_AGENT_FORGE_DIR', str(tmp_path))
    c = _case()
    bed = tmp_path / 'x.xlsx'
    bed.write_bytes(b'x')
    c['before'] = str(bed)
    vf._forge_save({'n': c})
    calls = []
    monkeypatch.setattr(vf, '_ask_once', lambda *a: (calls.append(1), _GOOD)[1])
    tries = iter([(["F2: 期待 '人件費' ／ マクロ後 '物件費'", "G2: 期待 '3' ／ マクロ後 '7'"], 0.1),
                  (["F2: 期待 '人件費' ／ マクロ後 '補助費等'", "G2: 期待 '3' ／ マクロ後 '5'"], 0.1)])
    monkeypatch.setattr(vf, '_try_macro', lambda case, bas, sub: next(tries))
    assert vf.forge('n', max_turns=4) is False
    assert len(calls) == 2                                     # 3・4 往復目の AI は呼ばない
    out = capsys.readouterr().out
    assert '同じセルが 2 往復続けて外れました（F2・G2）' in out and '依頼文にその決まり' in out
    assert vf._mismatch_key(["表の大きさ／位置が違う: …", "B3: 期待 'x'"]) == ('B3',)


def test_tests_feedback_shows_other_sheets_and_stall_on_tables(tmp_path, monkeypatch, capsys):
    c = _case()
    exp = _snap([['a', '支出額'], ['000123', '100']], r0=1)
    exp.update(sheet='旧システム出力', others=[{'name': '新システム出力', 'head': '（左上 A1）\n1\t新財務会計システム'}])
    c['tests'] = [{'label': 'G', 'expect': exp}]
    text = vf._tests_feedback(c, [('G', ["マクロは表に何も書かずに終わりました（…）", "D1: 期待 'x' ／ マクロ後 ''"])])
    assert 'シート「旧システム出力」の正解' in text and 'ほかのシート: 「新システム出力」' in text and '新財務会計システム' in text
    # 別の表が同じ所で 2 往復続けて外れたら止める・撃ち直しは前回のマクロと説明から
    monkeypatch.setattr(vf, '_AGENT_FORGE_FILE', str(tmp_path / '_agent_forge.json'))
    monkeypatch.setattr(vf, '_AGENT_FORGE_DIR', str(tmp_path))
    bed = tmp_path / 'x.xlsx'
    bed.write_bytes(b'x')
    c['before'] = str(bed)
    vf._forge_save({'n': c})
    calls = []
    monkeypatch.setattr(vf, '_ask_once', lambda ai, model, prompt: (calls.append(prompt), _GOOD)[1])
    monkeypatch.setattr(vf, '_try_macro', lambda case, bas, sub: ([], 0.1))
    monkeypatch.setattr(vf, '_try_tests', lambda case, bas, sub: ([('G', ["D1: 期待 'x' ／ マクロ後 ''"])], 0.1))
    assert vf.forge('n', max_turns=5) is False
    assert len(calls) == 2 and '同じ表が同じ所で 2 往復続けて外れました（G）' in capsys.readouterr().out
    assert vf.forge('n', max_turns=1) is False
    assert '前回のマクロの不一致' in calls[2] and 'Sub 合計行を足す' in calls[2]
    # 同じ規則違反が続いても止める
    calls.clear()
    monkeypatch.setattr(vf, '_ask_once', lambda ai, model, prompt: (calls.append(prompt), _GOOD.replace('ActiveSheet', 'ThisWorkbook.Sheets(1)'))[1])
    assert vf.forge('n', max_turns=5) is False
    assert len(calls) == 2 and '同じ規則違反が 2 往復続きました' in capsys.readouterr().out


def test_validate_code_rejects_thisworkbook_and_dotnet_objects():
    assert 'ThisWorkbook' in vf._validate_code(_GOOD.replace('ActiveSheet', 'ThisWorkbook.Sheets(1)'))[1]
    bad = _GOOD.replace('Set ws = ActiveSheet', 'Set ws = ActiveSheet\n    Set o = CreateObject("System.Collections.ArrayList")')
    assert 'Scripting.Dictionary' in vf._validate_code(bad)[1]
    ok = _GOOD.replace('Set ws = ActiveSheet', 'Set ws = ActiveSheet\n    Set d = CreateObject("Scripting.Dictionary")')
    assert vf._validate_code(ok)[0] == '合計行を足す'


def test_extract_code_unwraps_json_reply():
    import json as _json
    wrapped = _json.dumps({'code': _GOOD}, ensure_ascii=False)
    assert vf._extract_code(wrapped) == _GOOD + "\n"
    assert vf._extract_code('{"say": "x"}') is None


def test_drop_sub_removes_only_the_named_block():
    text = ("Attribute VB_Name = \"表の整理\"\r\nSub 表を整える()\r\n    x = 1\r\nEnd Sub\r\n"
            "Sub 合計行を足す()\r\n    ' 依頼の語: 合計\r\n    y = 2\r\nEnd Sub\r\nSub 重複行を消す()\r\nEnd Sub\r\n")
    cut = vf._drop_sub(text, '合計行を足す')
    assert 'Sub 合計行を足す' not in cut and 'y = 2' not in cut
    assert 'Sub 表を整える()' in cut and 'Sub 重複行を消す()' in cut and cut.count('End Sub') == 2
    assert vf._drop_sub(text, '無い') == text


def test_harness_calls_each_macro_and_returns_error_text():
    code = vf._harness_code(['表を整える', '合計行を足す'])
    assert code.startswith('Function 鍛冶_撃つ() As String') and 'On Error GoTo eh' in code
    assert code.index('    表を整える') < code.index('    合計行を足す')
    assert '"ERR|" & 段' in code and code.endswith('End Function\r\n')


def test_rule_violation_keeps_previous_mismatch_and_code(tmp_path, monkeypatch):
    monkeypatch.setattr(vf, '_AGENT_FORGE_FILE', str(tmp_path / '_agent_forge.json'))
    monkeypatch.setattr(vf, '_AGENT_FORGE_DIR', str(tmp_path))
    c = _case()
    bed = tmp_path / 'x.xlsx'
    bed.write_bytes(b'x')
    c['before'] = str(bed)
    vf._forge_save({'n': c})
    replies = iter([_GOOD, _GOOD.replace('ActiveSheet', 'ThisWorkbook.Sheets(1)'), _GOOD])
    prompts = []
    monkeypatch.setattr(vf, '_ask_once', lambda ai, model, prompt: (prompts.append(prompt), next(replies))[1])
    tries = iter([(["A7: 期待 '1' ／ マクロ後 ''"], 0.1), ([], 0.1)])
    monkeypatch.setattr(vf, '_try_macro', lambda case, bas, sub: next(tries))
    assert vf.forge('n', max_turns=3) is True
    assert 'ThisWorkbook' in prompts[2] and "A7: 期待 '1'" in prompts[2] and '前回のマクロ ---' in prompts[2]


def test_extract_code_keeps_trailing_functions_so_rules_reject_them():
    reply = _GOOD + "\n\nFunction NormalizeTkt(s As String) As String\n    NormalizeTkt = Trim(s)\nEnd Function\n"
    code = vf._extract_code("はい\n```vba\n" + reply + "```")
    assert 'Function NormalizeTkt' in code and code.rstrip().endswith('End Function')
    name, why, _h = vf._validate_code(code)
    assert name is None and 'GoSub' in why
    two = _GOOD + "\n\nSub 別の処理()\nEnd Sub\n以上です"
    assert vf._validate_code(vf._extract_code(two))[1].startswith('Sub は 1 本だけ')
    assert vf._extract_code(_GOOD + "\n\n以上です") == _GOOD + "\n"     # 後ろが文だけなら今までどおり


def test_context_rows_show_header_and_failed_rows():
    exp = _snap([['伝票番号', '科目名', '支出額', '突合'], ['000001', '旅費', '100', '一致'], ['', '合計', '100', '']], r0=1)
    text = vf._context_rows(exp, ["D3: 期待 '' ／ マクロ後 '新に無し'", "表の大きさ／位置が違う: …", "D3: 重複"])
    assert text.splitlines()[1] == "1\t伝票番号\t科目名\t支出額\t突合"
    assert text.splitlines()[2] == "3\t\t合計\t100\t" and len(text.splitlines()) == 3
    assert vf._context_rows(exp, ["Z99: 期待 'x'"]) == "" and vf._context_rows(exp, []) == ""
    c = _case()
    c['tests'] = [{'label': 'E', 'expect': exp}]
    assert '3\t\t合計\t100' in vf._tests_feedback(c, [('E', ["D3: 期待 '' ／ マクロ後 '新に無し'"])])


def test_hints_name_total_rows_and_row_count_shifts():
    mism = ["表の大きさ／位置が違う: 期待 30×5（左上 A1）・マクロ後 31×5（左上 A1）",
            "A18: 期待 '総務課' ／ マクロ後 '環境課'", "B18: 期待 '学校給食' ／ マクロ後 '合計'",
            "B19: 期待 '水道管更新' ／ マクロ後 '学校給食'", "C19: 期待 '2050000' ／ マクロ後 '680000'"]
    h = vf._hints(mism)
    assert '行が 1 行多い' in h and '「合計」「計」などの行を明細' in h and 'ずれています＝どこかで' not in h
    h = vf._hints(["表の大きさ／位置が違う: 期待 30×8（左上 A1）・マクロ後 30×3（左上 A1）"])
    assert '列が足りない' in h
    assert vf._hints(["D3: 期待 '一致' ／ マクロ後 '金額違い'"]) == ""
    c = _case()
    c['tests'] = [{'label': 'G', 'expect': _snap([['a']], r0=1)}]
    assert '原因の手がかり' in vf._tests_feedback(c, [('G', mism)])


class _FakeCM:
    def __init__(self, text):
        self.text = text
        self.CountOfLines = text.count('\n') + 1

    def Lines(self, a, n):
        return self.text


class _FakeBook:
    def __init__(self, name, text=None):
        self.Name = name
        self._text = text

    @property
    def VBProject(self):
        book = self

        class _P:
            def VBComponents(self, name):
                class _C:
                    CodeModule = _FakeCM(book._text)
                return _C()
        return _P()

    def SaveCopyAs(self, path):
        open(path, 'wb').write(b'x')


class _FakeXL:
    def __init__(self, books):
        self.books = {b.Name: b for b in books}

    def Workbooks(self, name):
        return self.books[name]


def test_rehearse_drops_extras_on_error_or_no_change(monkeypatch):
    import vbam_forge
    owner = _FakeBook('秀コンボ.xlam', _FURI)
    xl = _FakeXL([owner])
    wb = _FakeBook('支出.xlsx')
    reg = vp.registry_from_text(_FURI, owner='秀コンボ.xlam')
    plan = vp.plan_full("決算統計区分に振り直して。書き方もそろえて", reg, {'予算科目コード', '支出額'})
    seen = {}

    def fake(text, names, copy, sheet, module='表の整理', sel_cols=None):
        seen['names'] = names
        return steps
    monkeypatch.setattr(vbam_forge, 'rehearse_steps', fake)
    steps = [{'name': '表を整える', 'ok': True, 'changed': True, 'why': ''},
             {'name': '決算統計区分に振り直す', 'ok': True, 'changed': True, 'why': ''}]
    assert vp._rehearse(xl, wb, '支出明細', plan, '秀コンボ.xlam') == ([], '')
    assert seen['names'] == ['表を整える', '決算統計区分に振り直す']
    steps = [{'name': '表を整える', 'ok': True, 'changed': False, 'why': ''},
             {'name': '決算統計区分に振り直す', 'ok': False, 'changed': False, 'why': '実行時エラー 9 インデックスが有効範囲にありません。'}]
    dropped, why = vp._rehearse(xl, wb, '支出明細', plan, '秀コンボ.xlam')
    assert [e['name'] for e in dropped] == ['決算統計区分に振り直す'] and '実行時エラー 9' in why
    steps = [{'name': '表を整える', 'ok': True, 'changed': True, 'why': ''},
             {'name': '決算統計区分に振り直す', 'ok': True, 'changed': False, 'why': ''}]
    assert '何も変えずに' in vp._rehearse(xl, wb, '支出明細', plan, '秀コンボ.xlam')[1]
    steps = [{'name': '表を整える', 'ok': False, 'changed': False, 'why': 'コンパイルエラー'}]
    assert '表を整える がコンパイルエラー' in vp._rehearse(xl, wb, '支出明細', plan, '秀コンボ.xlam')[1]


def test_release_instance_closes_only_that_excel(monkeypatch):
    import vbam_core

    class _WBs:
        Count = 0

    class _XL:
        def __init__(self):
            self.Workbooks = _WBs()
            self.quit = 0

        def Quit(self):
            self.quit += 1
    a, b = _XL(), _XL()
    monkeypatch.setattr(vbam_core, '_created_instances', [{'xl': a, 'pid': 999991}, {'xl': b, 'pid': 999992}])
    monkeypatch.setattr(vbam_core, '_created_xl', b)
    monkeypatch.setattr(vbam_core, '_created_xl_pid', 999992)
    monkeypatch.setattr(vbam_core, '_pid_is_excel', lambda pid: False)
    assert vbam_core.release_instance(999992) == 1
    assert b.quit == 1 and a.quit == 0
    assert [i['pid'] for i in vbam_core._created_instances] == [999991]
    assert vbam_core._created_xl is a and vbam_core._created_xl_pid == 999991
    assert vbam_core.release_instance(None) == 0 and vbam_core.release_instance(12345) == 0


def test_phrases_make_ask_rule_and_reopen_passed_case(tmp_path, monkeypatch, capsys):
    c = _case()
    c['before_values'] = _snap([['予算科目コード', '支出額']], r0=1)
    c['request'] = '決算統計区分に振り直して'
    c['phrases'] = ['科目を決算統計の区分に分けて集計して', '統計の区分ごとにまとめて']
    heads, why = vf._check_fit(c, _FURI, '振り直|決算統計')
    assert heads is None and '「統計の区分ごとにまとめて」' in why and '科目を決算統計' not in why
    wide = _FURI.replace("' 依頼の語: 振り直|決算統計|集計", "' 依頼の語: 振り直|決算統計|統計の区分")
    assert vf._check_fit(c, wide, '振り直|決算統計|統計の区分')[1] is None
    assert '同じ仕事の言い換え' in vf._prompt(c) and '統計の区分ごとにまとめて' in vf._prompt(c)
    # 合格済みの弾に、依頼の語が当たらない言い換えを渡すと、鍛え直しに戻る
    monkeypatch.setattr(vf, '_AGENT_FORGE_FILE', str(tmp_path / '_agent_forge.json'))
    monkeypatch.setattr(vf, '_AGENT_FORGE_DIR', str(tmp_path))
    bed = tmp_path / 'x.xlsx'
    bed.write_bytes(b'x')
    bas = tmp_path / 'n.bas'
    bas.write_bytes(b'x')
    c.update({'before': str(bed), 'passed': True, 'bas': str(bas), 'sub': '決算統計区分に振り直す', 'code': _FURI,
              'ask': '振り直|決算統計', 'phrases': None})
    vf._forge_save({'n': c})
    ph = tmp_path / 'phrases.txt'
    ph.write_text('統計の区分ごとにまとめて\n', encoding='utf-8')
    prompts = []
    fixed = _FURI.replace("' 依頼の語: 振り直|決算統計|集計", "' 依頼の語: 振り直|決算統計|統計の区分")
    monkeypatch.setattr(vf, '_ask_once', lambda ai, model, prompt: (prompts.append(prompt), fixed)[1])
    monkeypatch.setattr(vf, '_try_macro', lambda case, bas, sub: ([], 0.1))
    monkeypatch.setattr(vf, '_check_bas', lambda p: True)
    assert vf.forge('n', phrases=str(ph), max_turns=2) is True
    rec = vf._forge_load()['n']
    # 2026-09-18: 値がそろっていて違反が依頼の語の行だけなら、道具が直す（AI を呼ばない・前の語も残す）
    assert rec['passed'] and '統計の区分' in rec['ask'].split('|') and '振り直' in rec['ask'].split('|') and len(prompts) == 0


_SHUYAKU = """Sub 課別シートを集約する()
    ' 依頼の語: 集約|まとめ|集計
    ' 扱う: 集約 まとめ
    ' 見出し: 課名 事業名 予算額 執行額
End Sub
"""


def test_plan_full_longer_match_on_skipped_job_blocks_short_match():
    reg = vp.registry_from_text(_FURI + _SHUYAKU, owner='秀コンボ.xlam')
    shuyaku_sheet = {'課名', '事業名', '予算額', '執行額', '執行率'}
    furi_sheet = {'予算科目コード', '科目名', '支出額'}
    req = "対応表シートで予算科目コードを決算統計区分に振り直して、区分ごとの支出額を集計してください"
    p = vp.plan_full(req, reg, shuyaku_sheet)                     # 振り直しを頼んだのに集約のブック
    assert p['extras'] == [] and [e['name'] for e in p['shadowed']] == ['課別シートを集約する'] and p['other']
    p = vp.plan_full(req, reg, furi_sheet)                        # 本来の表
    assert [e['name'] for e in p['extras']] == ['決算統計区分に振り直す'] and p['shadowed'] == []
    p = vp.plan_full("各課のシートを集約シートにまとめて集計して", reg, shuyaku_sheet)
    assert [e['name'] for e in p['extras']] == ['課別シートを集約する']
    assert vp._ask_strength('集約|まとめ|集計', '全部まとめて集計') == 3 and vp._ask_strength('x(', 'x') == 0


def test_rehearse_drop_recomputes_plan_by_name():
    reg = vp.registry_from_text(_FURI, owner='o')
    p = vp.plan_full("決算統計区分に振り直して", reg, {'予算科目コード', '支出額'})
    dropped = p['extras']
    gone = {x['name'] for x in dropped}
    assert vp.plan_full("決算統計区分に振り直して", [e for e in reg if e['name'] not in gone], {'予算科目コード', '支出額'})['extras'] == []


_TOTSUGO = """Sub 新システムと突合する()
    ' 依頼の語: 突合|新システム|突き合わせ
    ' 扱う: 突合 新システム突合 伝票番号突合
    ' 見出し: 伝票番号 科目名 支出額
End Sub
"""


def test_plan_full_clause_with_ask_word_is_covered():
    reg = vp.registry_from_text(_TOTSUGO, owner='o')
    sheet = {'伝票番号', '科目名', '支出額'}
    for req in ("新システムシートと突き合わせて、伝票番号ごとに突合結果を足し、その右に突合結果ごとの件数を出してください",
                "旧システムと新システムを照合して違いを出して", "新旧の伝票を突き合わせて一致・不一致を出して"):
        p = vp.plan_full(req, reg, sheet)
        assert [e['name'] for e in p['extras']] == ['新システムと突合する'] and p['left'] == [], (req, p['left'])
    p = vp.plan_full("新システムと突合して、結果をグラフにして", reg, sheet)
    assert p['left'] == ['グラフ']


def test_check_fit_requires_own_request_not_sent_to_ai():
    c = _case()
    c['before_values'] = _snap([['伝票番号', '科目名', '支出額']], r0=1)
    c['request'] = '新システムと突き合わせて突合して。件数も出して'
    heads, why = vf._check_fit(c, _TOTSUGO, '突合|新システム|突き合わせ')
    assert heads is None and '件数' in why and '扱う' in why
    fixed = _TOTSUGO.replace("' 扱う: 突合", "' 扱う: 件数 突合")
    assert vf._check_fit(c, fixed, '突合|新システム|突き合わせ')[1] is None


def test_ask_must_be_plain_words_and_ties_do_not_fire():
    bad = _GOOD.replace("' 依頼の語: 合計|集計", "' 依頼の語: 比較.*差異|合計")
    assert '語を | で並べるだけ' in vf._validate_code(bad)[1]
    assert vp._ask_strength('シート.*集|集計', '対応表シートで振り直して、支出額を集計して') == 4
    reg = vp.registry_from_text(_FURI + _SHUYAKU.replace('集約|まとめ|集計', '集約|まとめ|集計|区分'), owner='o')
    # 「決算統計」(4) vs 「区分」(2)／同じ長さの取り合いは撃たない
    p = vp.plan_full("決算統計区分に振り直して", reg, {'課名', '事業名', '予算額', '執行額'})
    assert p['extras'] == []
    tie = vp.registry_from_text(_FURI.replace('振り直|決算統計|集計', '振り直|集計') + _SHUYAKU, owner='o')
    p = vp.plan_full("振り直して集計して", tie, {'課名', '事業名', '予算額', '執行額'})
    assert p['extras'] == [] and [e['name'] for e in p['shadowed']] == ['課別シートを集約する']
    p = vp.plan_full("各課を集約して", tie, {'課名', '事業名', '予算額', '執行額'})
    assert [e['name'] for e in p['extras']] == ['課別シートを集約する']       # 撃てない仕事に当たらなければ撃つ


def test_ask_line_regex_is_split_into_words_and_phrase_words_are_shown():
    # 2026-09-17 夜: 「帳票.*一覧」「課別.{0,3}月別」を書き続けて規則違反だけで往復を使い切った＝道具が語に分ける
    code = _GOOD.replace("' 依頼の語: 合計|集計", "' 依頼の語: 帳票.*一覧|課別.{0,3}月別の支出|(受付簿)|帳票")
    fixed, line = vf._normalize_ask_line(code)
    assert line == '帳票|一覧|課別|月別の支出|受付簿'
    assert "' 依頼の語: 帳票|一覧|課別|月別の支出|受付簿" in fixed and vf._validate_code(fixed)[0]
    assert vf._normalize_ask_line(_GOOD) == (_GOOD, None)
    assert vf._phrase_words('課別・月別の支出額の表を作ってください') == ['課別', '月別', '支出額']


def test_hints_name_baked_dates_lost_zeros_and_column_shift():
    h = vf._hints(["G3: 期待 'R8-001' ／ マクロ後 '2026/1/1'", "G4: 期待 'R8-002' ／ マクロ後 '2026/2/1'"])
    assert '日付に化けて' in h and '"@"' in h
    assert '先頭の 0' in vf._hints(["A2: 期待 '0012' ／ マクロ後 '12'"])
    assert '日付に化けて' not in vf._hints(["A2: 期待 '2026/4/1' ／ マクロ後 '2026/4/2'"])
    shifted = [f"{c}1: 期待 '{e}' ／ マクロ後 '{g}'" for c, e, g in
               [('G', '課名', ''), ('H', '4月', ''), ('I', '5月', '課名'), ('J', '6月', '4月'), ('K', '7月', '5月'), ('L', '8月', '6月')]]
    h = vf._hints(shifted)
    assert '2 列右にずれて' in h and '行ごとにずれて' not in h


def test_compare_checks_number_date_text_kinds():
    # 2026-09-17 夜: 申請額・受付日を "@" の文字で書いた帳票のマクロが、値の見た目だけで合格していた
    exp = {'row': 1, 'col': 1, 'values': [['受付番号', '申請額', '受付日'], ['0001', '1000', '2026/4/1']], 'kinds': ['sss', 'snd']}
    got = {'row': 1, 'col': 1, 'values': [['受付番号', '申請額', '受付日'], ['0001', '1000', '2026/4/1']], 'kinds': ['sss', 'sss']}
    m = vf._compare(exp, got)
    assert len(m) == 2 and m[0].startswith("B2: 型が違う 期待 数") and '日付' in m[1]
    assert vf._compare(exp, dict(got, kinds=['sss', 'snd'])) == []
    assert vf._compare({k: v for k, v in exp.items() if k != 'kinds'}, got) == []      # 古い台帳の正解は型を比べない
    assert '文字で書いています' in vf._hints(m)
    assert vf._mismatch_key(m)
    assert vf._truth_of_test({'before': r'C:\x\月次_test_月次_試験A_前_before.xlsx'}) == r'C:\x\月次_test_月次_試験A_前_truth.xlsx'
    assert vf._kind_of(None) == '-' and vf._kind_of(3.0) == 'n' and vf._kind_of('a') == 's' and vf._kind_of(True) == 'b'


def test_changed_cells_find_the_table_the_macro_wrote():
    # 2026-09-17 夜: 入口の tidy と検査が元の表（帳票・明細）に当たり、マクロが右に作った表を見ていなかった
    before = {'row': 1, 'col': 1, 'values': [['支出日', '課名', None, '課名'], ['2026/4/1', '総務課', None, '旧課']]}
    now = [['支出日', '課名', None, '課名', '4月'], ['2026/4/1', '総務課', None, '総務課', 100.0], [None, None, None, None, None]]
    assert vp._changed_cells(before, now, 1, 1) == [(1, 5), (2, 4), (2, 5)]
    assert vp._changed_cells(before, [['支出日', '課名', None, None]], 1, 1) == []          # 消しただけは数えない
    assert vp._changed_cells({'row': 3, 'col': 2, 'values': [['a']]}, [['a', 'b']], 3, 2) == [(3, 3)]


def test_header_line_break_still_runs_macro_and_returns_values(tmp_path, monkeypatch, capsys):
    """頭の行だけの違反でも写しで撃ち、値の外れと一緒に返す（撃たずに往復を使い切らない）。"""
    monkeypatch.setattr(vf, '_AGENT_FORGE_FILE', str(tmp_path / '_agent_forge.json'))
    monkeypatch.setattr(vf, '_AGENT_FORGE_DIR', str(tmp_path))
    c = _case()
    bed = tmp_path / 'x.xlsx'
    bed.write_bytes(b'x')
    c.update({'before': str(bed), 'phrases': ['数量を足して', '合計を出して']})
    vf._forge_save({'n': c})
    replies = [_GOOD.replace("' 依頼の語: 合計|集計", "' 依頼の語: 合計"), _GOOD.replace("' 依頼の語: 合計|集計", "' 依頼の語: 合計|足して")]
    prompts, runs = [], []
    monkeypatch.setattr(vf, '_ask_once', lambda ai, model, prompt: (prompts.append(prompt), replies[len(prompts) - 1])[1])
    monkeypatch.setattr(vf, '_try_macro', lambda case, bas, sub: (runs.append(sub), ((['A1: 期待 1 ／ マクロ後 2'] if len(runs) == 1 else []), 0.1))[1])
    monkeypatch.setattr(vf, '_check_bas', lambda p: True)
    monkeypatch.setattr(vf, '_check_fit', lambda case, code, ask: ((['品名'], None) if '足して' in ask else (None, "' 依頼の語: が言い換えに当たりません")))
    assert vf.forge('n', max_turns=3) is True
    assert len(runs) == 2                                    # 1 往復目も撃った
    assert '言い換えに当たりません' in prompts[1] and 'A1: 期待 1' in prompts[1]


def test_passed_case_is_rechecked_against_tightened_rules(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(vf, '_AGENT_FORGE_FILE', str(tmp_path / '_agent_forge.json'))
    monkeypatch.setattr(vf, '_AGENT_FORGE_DIR', str(tmp_path))
    old = _FURI.replace('振り直|決算統計|集計', '振り直|決算統計.*区分')
    c = _case()
    bed = tmp_path / 'x.xlsx'
    bed.write_bytes(b'x')
    bas = tmp_path / 'n.bas'
    bas.write_bytes(b'x')
    c.update({'before': str(bed), 'passed': True, 'bas': str(bas), 'sub': '決算統計区分に振り直す', 'code': old,
              'ask': '振り直|決算統計.*区分', 'request': '決算統計区分に振り直して',
              'before_values': _snap([['予算科目コード', '支出額']], r0=1)})
    vf._forge_save({'n': c})
    prompts = []
    monkeypatch.setattr(vf, '_ask_once', lambda ai, model, prompt: (prompts.append(prompt), _FURI)[1])
    monkeypatch.setattr(vf, '_try_macro', lambda case, bas, sub: ([], 0.1))
    monkeypatch.setattr(vf, '_check_bas', lambda p: True)
    assert vf.forge('n', max_turns=2) is True
    out = capsys.readouterr().out
    assert '今の規則に合いません' in out and '語を | で並べるだけ' in prompts[0]
    assert vf._forge_load()['n']['ask'] in ('振り直|決算統計|集計', '振り直|決算統計')   # 2026-09-18: 集計は一般の依頼の語


def test_instrument_exits_marks_each_exit_sub_with_its_line():
    code = ("Attribute VB_Name = \"鍛冶_試し\"\r\nSub 振る()\r\n    Dim s As String\r\n    If s = \"\" Then Exit Sub\r\n    ' Exit Sub はコメント\r\n"
            "    s = \"Exit Sub\"\r\n    Exit Sub\r\nLbl:\r\n    Return\r\nEnd Sub\r\nSub 別()\r\n    Exit Sub\r\nEnd Sub\r\n")
    text, where = vf._instrument_exits(code, '振る')
    assert where == {3: 'If s = "" Then Exit Sub', 6: 'Exit Sub'}
    assert '    If s = "" Then 鍛冶_抜けた = "3": Exit Sub\r\n' in text
    assert "' Exit Sub はコメント" in text and 's = "Exit Sub"' in text      # コメント・文字の中は触らない
    assert text.count('鍛冶_抜けた') == 2                                   # 別の Sub は触らない
    assert 'Public 鍛冶_抜けた As String' in vf._harness_code(['振る'], track=True)
    assert '鍛冶_抜けた' not in vf._harness_code(['振る'])


def test_trap_hints_vbnarrow_with_katakana_literal():
    assert vf._trap_hints('ss = StrConv(ss, vbNarrow)\r\nIf ss = "予算科目コード" Then') 
    assert not vf._trap_hints('ss = StrConv(ss, vbNarrow)\r\nIf ss = "支出額" Then')


def test_trap_hints_sumif_over_text_numbers():
    code = 'ws.Cells(2, 7).Formula = "=SUMIF(D2:D9," & a & ",C2:C9)"'
    text_nums = {'values': [['コード', '支出額'], ['01', '12000'], ['02', '5,000円']], 'kinds': ['ss', 'ss', 'ss']}
    nums = {'values': [['コード', '支出額'], ['01', '12000']], 'kinds': ['ss', 'sn']}
    assert any('SUMIF' in h for h in vf._trap_hints(code, text_nums))
    assert not vf._trap_hints(code, nums) and not vf._trap_hints(code)


def test_trap_hints_yen_and_val_comma():
    t = {'values': [['支出額'], ['12,000円'], ['△1,200'], ['5,000']], 'kinds': ['s', 's', 's', 's']}
    h = vf._trap_hints('If IsNumeric(v) Then a = a + CDbl(v)', t)
    assert any('円・△' in x for x in h) and not any('SUMIF' in x for x in h)
    assert any('カンマで止まる' in x for x in vf._trap_hints('a = a + Val(v)', t))
    assert not vf._trap_hints('ws.Cells(9, 3).Formula = "=SUM(C2:C8)"', t)          # 合計行の =SUM は罠にしない


def test_trap_hints_nbsp():
    t = {'values': [['コード'], ['\u00a002-01']], 'kinds': ['s', 's']}
    assert any('NBSP' in x for x in vf._trap_hints('c = Trim(v)', t))
    assert not any('NBSP' in x for x in vf._trap_hints('c = Replace(Trim(v), ChrW(160), "")', t))
    assert any('ChrW' in x for x in vf._trap_hints('c = Replace(Trim(v), Chr(160), "")', t))   # Chr(160) は NBSP でない


def test_keep_best_reverts_worse_attempt():
    best = {}
    code, fb, rev = vf._keep_best(best, 2, 'A', 's', 'fbA')
    assert (code, fb, rev) == ('A', 'fbA', False) and best['score'] == 2
    code, fb, rev = vf._keep_best(best, 6, 'B', 's', 'fbB')
    assert rev and code == 'A' and '2 枚 → 6 枚' in fb and fb.endswith('fbA') and best['code'] == 'A'
    code, fb, rev = vf._keep_best(best, 1, 'C', 's', 'fbC')
    assert not rev and best['code'] == 'C'
    code, fb, rev = vf._keep_best(best, vf._MAIN_MISS_SCORE + 40, 'D', 's', 'fbD')
    assert rev and '弾の表で 40 セル' in fb


def test_hints_repeated_title_as_record():
    exp = {'row': 1, 'col': 1, 'values': [['令和8年度 地域づくり活動補助金 申請受付簿', '', ''], ['受付番号', '', '受付番号'], ['', '', 'R8-001']]}
    h = vf._hints(["G9: 期待 'R8-007' ／ マクロ後 '令和8年度 地域づくり活動補助金 申請受付簿'"], exp)
    assert '表題' in h and 'くり返した' in h
    assert '表題' not in vf._hints(["G9: 期待 'R8-007' ／ マクロ後 'R8-008'"], exp)


def test_similar_heads_contains_search():
    start = {'values': [['支出日', '課名', '支出額（税抜）', '支出額']], 'kinds': ['ssss']}
    code = 'If InStr(h, "支出額") > 0 Then amtCol = c'
    h = vf._trap_hints(code, start)
    assert any('「支出額」と「支出額（税抜）」' in x for x in h)
    exp = {'others': [{'name': '総務課', 'head': '（左上 A1・5 行×4 列）\n1\t事業名\t前年度予算額\t予算額\t執行額\n2\t除雪\t1\t2\t3'}]}
    assert any('「予算額」と「前年度予算額」' in x for x in vf._trap_hints('If InStr(v, "予算額") Then', {'values': []}, exp))
    assert not vf._trap_hints('If h = "支出額" Then', start)                      # 一致で探していれば言わない


def test_hints_total_in_header_row_points_to_list_sheet():
    exp = {'row': 1, 'col': 1, 'values': [['費目', '金額', '', '費目', '総務課', '合計']]}
    h = vf._hints(["F1: 期待 '合計' ／ マクロ後 '財政課'", "G1: 期待 '' ／ マクロ後 '合計'"], exp)
    assert '見出しの行に「合計」' in h
    assert '見出しの行に' in vf._hints(["L1: 期待 '合計' ／ マクロ後 '※ 4月1日現在の人数'"], exp)
    assert '※」で始まる注記' in vf._hints(["D30: 期待 '' ／ マクロ後 '※ 税込み'"], exp)
    h2 = vf._hints(["D9: 期待 '' ／ マクロ後 '合計'"], exp)
    assert '明細として扱っています' in h2 and '見出しの行' not in h2


def test_keep_best_names_tables_the_discarded_fix_broke():
    best = {}
    vf._keep_best(best, 1, 'A', 's', 'fbA', ['前年度'])
    code, fb, rev = vf._keep_best(best, 2, 'B', 's', 'fbB', ['前年度', '見出しの空白'])
    assert rev and code == 'A' and '「見出しの空白」を外しました' in fb
    assert not vf._trap_hints('If InStr(v, "（例") > 0 Then', {'values': [['事業名', '予算額'], ['（例）庁舎清掃', 1]], 'kinds': ['ss', 'sn']})


def test_hints_spaced_total_label():
    h = vf._hints(["B6: 期待 '公園整備' ／ マクロ後 '小　計'"])
    assert '空白の入った合計' in h and '明細として扱っています' in h


def test_old_findings_are_cells_the_macro_did_not_change():
    notes = ["「合計」行の C45（550,000円）が上の和と合いません（セル 16,519,000／上の和 5,063,000）",
             "D12: 対応なし の列に空欄があります", "見出しが 2 行あります"]
    assert vp.old_findings(notes, {'D12', 'F2'}) == [notes[0]]                 # 番地の無い指摘は古いと言わない
    assert vp.old_findings(notes, set()) == [notes[0], notes[1]]


def test_plan_full_generic_word_on_skipped_job_does_not_shadow():
    reg = vp.registry_from_text("Sub 合計行を足す()\n    ' 依頼の語: 合計|集計\n    ' 扱う: 合計 集計\n    ' 見出し: 品名 数量 金額\nEnd Sub\n"
                                "Sub 課別シートを集約する()\n    ' 依頼の語: 集約|各課\n    ' 扱う: 集約 合計\n    ' 見出し: 課名 事業名\nEnd Sub\n")
    p = vp.plan_full("各課シートの事業を集めて合計も出して", reg, {'課名', '事業名'})
    assert [e['name'] for e in p['extras']] == ['課別シートを集約する'] and p['other'] is False


def test_check_fit_rejects_ask_words_hitting_generic_requests():
    code = "Sub 課別シートを集約する()\n    ' 依頼の語: 集約|まとめ\n    ' 扱う: 集約\n    ' 見出し: 課名\nEnd Sub\n"
    case = {'request': '各課のシートを集約シートにまとめて', 'before_values': {'values': [['課名']]}, 'tests': []}
    heads, why = vf._check_fit(case, code, '集約|まとめ')
    assert heads is None and '「わかりやすくまとめてください」に「まとめ」' in why


def test_check_fit_rejects_long_memorized_ask_words():
    code = "Sub 新システムと突合する()\n    ' 依頼の語: 突合|旧システムと新システムの伝票を照らし合わせ\n    ' 扱う: 突合\n    ' 見出し: 伝票番号\nEnd Sub\n"
    case = {'request': '新システムと突合して', 'before_values': {'values': [['伝票番号']]}, 'tests': []}
    heads, why = vf._check_fit(case, code, '突合|旧システムと新システムの伝票を照らし合わせ')
    assert heads is None and '8 字まで' in why


def test_ask_candidates_avoid_generic_and_particles():
    c = vf._ask_candidates('課ごとの月ごとの支出額の表を作って', ['支出明細を課ごと・月ごとに集計', '月ごとの支出を課別に集計'])
    assert c and all(not w.startswith(('の', 'を')) and not w.endswith(('の', 'を')) for w in c)
    assert '支出額の表' not in c and any('月ごと' in w for w in c)


def test_check_fit_keeps_valid_words_of_passed_version():
    code = "Sub 課別シートを集約する()\n    ' 依頼の語: 集約シート|各課シート\n    ' 扱う: 集約\n    ' 見出し: 課名\nEnd Sub\n"
    case = {'request': '各課シートを集約シートに', 'before_values': {'values': [['課名']]}, 'tests': [],
            'keep_ask': '集約|一覧|まとめ|集約シート'}
    heads, why = vf._check_fit(case, code, '集約シート|各課シート')
    assert heads is None and '「集約」' in why and '「一覧」' not in why          # 一般語（一覧・まとめ）は外してよい


def test_check_fit_dropping_word_covered_by_shorter_word_is_fine():
    code = "Sub 帳票を一覧に直す()\n    ' 依頼の語: 受付簿|帳票\n    ' 扱う: 一覧\n    ' 見出し: 受付番号\nEnd Sub\n"
    case = {'request': '受付簿の帳票を一覧に直して', 'before_values': {'values': [['受付番号']]}, 'tests': [],
            'keep_ask': '受付簿|帳票|受付簿をリスト|帳票を一覧|一覧に直'}
    heads, why = vf._check_fit(case, code, '受付簿|帳票')
    assert '「一覧に直」' in why and '受付簿をリスト' not in why and '帳票を一覧' not in why


def test_check_fit_does_not_keep_word_gone_from_texts():
    """2026-09-18 第二期: 言い換えを替えて「中に小」が文に出なくなったのに「残せ」と言い、次の往復で「文に出ない語は外せ」と言った。"""
    code = "Sub スパークラインを足す()\n    ' 依頼の語: スパークライン|セル内\n    ' 扱う: スパークライン\n    ' 見出し: なし\nEnd Sub\n"
    case = {'request': 'スパークラインを入れて', 'phrases': ['セル内の小さなグラフで推移を'], 'before_values': {'values': [['a', 'b']]},
            'tests': [], 'keep_ask': 'スパークライン|中に小'}
    _heads, why = vf._check_fit(case, code, 'スパークライン|セル内')
    assert '中に小' not in (why or '')


def test_main_score_broken_version_is_worst():
    err = ["マクロ 前年度比較表を作る が実行時エラーで止まりました: 9 インデックスが有効範囲にありません。"]
    cells = [f"D{i}: 期待 '1' ／ マクロ後 '2'" for i in range(65)]
    assert vf._main_score(err) > vf._main_score(cells) > vf._main_score(cells[:3])
    best = {}
    vf._keep_best(best, vf._main_score(cells), 'A', 's', 'fbA', ['弾の表'])
    code, fb, rev = vf._keep_best(best, vf._main_score(err), 'B', 's', 'fbB', ['弾の表'])
    assert rev and code == 'A' and '撃てない版' in fb


def test_similar_heads_ignores_data_rows():
    start = {'values': [['科目', '予算額', '', '科目', '今年度'], ['旧科目1', 1, '', '旧科目1', 2]], 'kinds': ['sssss', 'sn-sn']}
    assert not any('旧科目1' in x for x in vf._trap_hints('If InStr(h, "科目") > 0 Then', start))


def test_check_fit_rejects_words_of_other_forged_jobs(monkeypatch):
    monkeypatch.setattr(vf, '_OTHER_JOB_TEXTS', {'突合': ['移行後のデータと突き合わせ']})
    code = "Sub 前年度比較表を作る()\n    ' 依頼の語: 前年度比較|突き合わせ\n    ' 扱う: 比較\n    ' 見出し: 科目\nEnd Sub\n"
    case = {'name': '比較', 'request': '前年度比較の表を', 'before_values': {'values': [['科目']]}, 'tests': []}
    heads, why = vf._check_fit(case, code, '前年度比較|突き合わせ')
    assert heads is None and '「突き合わせ」＝突合' in why


def test_keep_words_do_not_override_clash(monkeypatch):
    monkeypatch.setattr(vf, '_OTHER_JOB_TEXTS', {'突合': ['移行後のデータと突き合わせ']})
    code = "Sub 前年度比較表を作る()\n    ' 依頼の語: 前年度比較\n    ' 扱う: 比較\n    ' 見出し: 科目\nEnd Sub\n"
    case = {'name': '比較', 'request': '前年度比較の表を', 'before_values': {'values': [['科目']]}, 'tests': [],
            'keep_ask': '前年度比較|突き合わせ'}
    heads, why = vf._check_fit(case, code, '前年度比較')
    assert why is None or '突き合わせ' not in why


def test_formula_copy_blockers_absolute_sum_per_column_is_fine():
    import vbam_view as vv
    cells = {(9, 5): ('=SUM($E$2:$E$8)', '=SUM(R2C5:R8C5)'), (9, 6): ('=SUM($F$2:$F$8)', '=SUM(R2C6:R8C6)'),
             (9, 7): ('=SUM($G$2:$G$8)', '=SUM(R2C7:R8C7)')}
    assert vv.formula_copy_blockers(cells) == []


def test_compare_ignores_trailing_empty_rows_and_cols():
    e = _snap([['a', '1', ''], ['b', '2', ''], ['', '', '']])
    g = _snap([['a', '1'], ['b', '2']])
    assert vf._compare(e, g) == []
    assert vf._compare(e, _snap([['a', '1'], ['b', '3']])) == ["B6: 期待 '2' ／ マクロ後 '3'"]


def test_plan_full_graph_or_sort_in_same_clause_goes_to_ai():
    reg = vp.registry_from_text("Sub 決算統計区分を振り直して集計する()\n    ' 依頼の語: 決算統計|振り直\n    ' 扱う: 振り直し 区分 集計\n"
                                "    ' 見出し: 予算科目コード\nEnd Sub\n")
    p = vp.plan_full("決算統計区分に振り直して、区分ごとの円グラフも作って", reg, {'予算科目コード'})
    assert p['extras'] and p['other'] is True and 'グラフ' in p['left']
    p = vp.plan_full("決算統計区分に振り直して、区分ごとに集計して", reg, {'予算科目コード'})
    assert p['extras'] and p['other'] is False


def test_trap_hints_chart_gapwidth_on_series():
    assert any('ChartGroups(1).GapWidth' in x for x in vf._trap_hints('sr.GapWidth = 60'))
    assert not any('GapWidth' in x for x in vf._trap_hints('cht.ChartGroups(1).GapWidth = 60'))


def test_instrument_line_marks_skip_block_keywords():
    code = ("Sub 振る()\r\n    Dim a As Long\r\n    Select Case a\r\n        Case 1\r\n            a = 2\r\n        Case Else\r\n"
            "            a = 3\r\n    End Select\r\n    If a = 3 Then\r\n        a = 4\r\n    Else\r\n        a = 5\r\n    End If\r\n"
            "    GoSub Lbl\r\n    Exit Sub\r\nLbl:\r\n    a = a + _\r\n        1\r\n    Return\r\nEnd Sub\r\n")
    text, _where = vf._instrument_exits(code, '振る')
    lines = text.split('\r\n')
    marks = [i for i, ln in enumerate(lines) if '鍛冶_行 =' in ln]
    assert all(not re.match(r'\s*(Case|Else|End|Dim|Lbl:)', lines[i + 1]) for i in marks)
    assert not any('鍛冶_行' in lines[i] for i in range(len(lines)) if lines[i - 1].rstrip().endswith(' _'))
    assert 'Select Case a\r\n        Case 1' in text                     # Select Case と最初の Case のあいだに入れない


def test_formula_copy_blockers_table_totals_row_is_fine():
    import vbam_view as vv
    cells = {(14, 4): ('=SUBTOTAL(109,[基本給])', '=SUBTOTAL(109,[基本給])'), (14, 5): ('=SUBTOTAL(109,[手当])', '=SUBTOTAL(109,[手当])'),
             (14, 6): ('=SUBTOTAL(109,[支給額])', '=SUBTOTAL(109,[支給額])')}
    assert vv.formula_copy_blockers(cells) == []


def test_ask_cover_picks_word_hitting_all():
    miss = ['パワークエリで月ごとの支出を追加して課別に集計', 'T支出のテーブルをパワークエリで結合して課別集計', 'パワークエリで課別集計のクエリを作って']
    assert vf._ask_cover(miss, miss) == ['パワークエリ']


def test_check_fit_header_clash_only_when_distinguishable(monkeypatch):
    monkeypatch.setattr(vf, '_OTHER_JOB_WORDS', {'月次集計': {'支出日', '課名', '科目', '摘要', '支出額'}})
    monkeypatch.setattr(vf, '_OTHER_JOB_TEXTS', {})
    line = "Sub 推移を描く()\n    ' 依頼の語: 推移の折れ線\n    ' 扱う: グラフ\n    ' 見出し: 課名\nEnd Sub\n"
    case = {'name': '推移', 'request': '推移の折れ線を', 'tests': [],
            'before_values': {'values': [['課名', '4月', '5月', '合計'], ['総務課', 1, 2, 3]]}}
    heads, why = vf._check_fit(case, line, '推移の折れ線')
    assert heads is None and 'ほかの仕事「月次集計」' in why
    bar = line.replace('推移を描く', '棒を描く').replace('見出し: 課名', '見出し: 課名 支出額')
    case2 = dict(case, name='棒', before_values={'values': [['課名', '支出額'], ['総務課', 1]]})
    heads, why = vf._check_fit(case2, bar, '推移の折れ線')
    assert why is None or 'ほかの仕事' not in why


def test_auto_fix_ask_adds_cover_and_keeps_words(monkeypatch):
    monkeypatch.setattr(vf, '_OTHER_JOB_TEXTS', {'月次集計': ['支出明細を月別に集計']})
    monkeypatch.setattr(vf, '_OTHER_JOB_WORDS', {})
    code = ("Sub 課別科目別ピボットを作る()\n    ' 依頼の語: 課別科目別|支出明細ピボットのテーブルを作る\n    ' 扱う: ピボット\n"
            "    ' 見出し: 課名 科目\nEnd Sub\n")
    case = {'name': 'ピボット', 'request': '課別科目別のピボットを', 'tests': [], 'keep_ask': '課別科目別|ピボットテーブル',
            'phrases': ['クロス集計表をピボットで作成', '支出明細のピボットテーブルを別シートに'],
            'before_values': {'values': [['支出日', '課名', '科目', '摘要', '支出額']]}}
    new_code, ask = vf._auto_fix_ask(case, code, '課別科目別|支出明細ピボットのテーブルを作る')
    assert new_code and 'ピボットテーブル' in ask.split('|') and '課別科目別' in ask.split('|')
    assert not any(len(w) > 8 for w in ask.split('|'))
    assert vf._check_fit(case, new_code, ask)[1] is None


def test_trap_hints_chartarea_font_after_title_size():
    code = 'ch.ChartTitle.Characters.Font.Size = 14\nch.ChartArea.Format.TextFrame2.TextRange.Font.Name = "Meiryo UI"\n'
    assert any('タイトルの大きさと太字は ChartArea' in x for x in vf._trap_hints(code))


def test_ask_candidates_split_compound_only_when_nothing_else():
    """2026-09-18: 「課別月別のピボットをいつもの形で」に候補が無かった（課別月別はほかの仕事と衝突）→ 並みを 2 字以上ずつに切った語を出す。"""
    saved = dict(vf._OTHER_JOB_TEXTS)
    try:
        vf._OTHER_JOB_TEXTS.clear()
        vf._OTHER_JOB_TEXTS.update({'推移': ['課別月別支出の推移を折れ線に'], 'ピボット': ['課別・科目別のピボットテーブル']})
        assert vf._ask_candidates('課別月別のピボットをいつもの形で') == ['月別のピボット']
        assert '別のピボット' not in vf._ask_candidates('課別月別のピボットをいつもの形で', n=12)
        vf._OTHER_JOB_TEXTS.clear()
        assert '月別のピボット' not in vf._ask_candidates('課別月別のピボットをいつもの形で', n=12)   # 並みの頭からの語があれば使わない
    finally:
        vf._OTHER_JOB_TEXTS.clear()
        vf._OTHER_JOB_TEXTS.update(saved)

def test_auto_fix_ask_drops_header_missing_in_a_test_table(monkeypatch):
    """2026-09-18: 下半期だけの表を足したら「' 見出し: 課名 4月」の 4月 が無く、AI を 1 往復呼んだ → 道具が外す。"""
    monkeypatch.setattr(vf, '_OTHER_JOB_TEXTS', {})
    monkeypatch.setattr(vf, '_OTHER_JOB_WORDS', {})
    code = "Sub 縦持ちにする()\n    ' 依頼の語: 縦持ち\n    ' 扱う: 縦持ち\n    ' 見出し: 課名 4月\nEnd Sub\n"
    case = {'name': '縦持ち', 'request': '月別支出を縦持ちに', 'phrases': [], 'keep_ask': '',
            'before_values': {'values': [['課コード', '課名', '4月', '5月']]},
            'tests': [{'label': '下半期', 'expect': {'values': [['課名', '10月', '11月']]}}]}
    assert vf._check_fit(case, code, '縦持ち')[1]
    new_code, ask = vf._auto_fix_ask(case, code, '縦持ち')
    assert "' 見出し: 課名\n" in new_code and vf._check_fit(case, new_code, ask)[1] is None

def test_auto_fix_ask_adds_header_word_on_clash(monkeypatch):
    """2026-09-18: 課名を引くを足したら、突合の「' 見出し: 伝票番号 支出額」がその表にもそろい AI を呼んだ → 道具が語を足す。"""
    monkeypatch.setattr(vf, '_OTHER_JOB_TEXTS', {})
    monkeypatch.setattr(vf, '_OTHER_JOB_WORDS', {'課名を引く': {'伝票番号', '支出日', '課コード', '摘要', '支出額'}})
    code = "Sub 突合する()\n    ' 依頼の語: 突合\n    ' 扱う: 突合\n    ' 見出し: 伝票番号 支出額\nEnd Sub\n"
    case = {'name': '突合', 'request': '新旧を突合して', 'phrases': [], 'keep_ask': '',
            'before_values': {'values': [['伝票番号', '支出額', '移行後金額']]},
            'tests': [{'label': 'A', 'expect': {'values': [['伝票番号', '移行後金額', '支出額', '差異']]}}]}
    assert vf._check_fit(case, code, '突合')[1]
    new_code, ask = vf._auto_fix_ask(case, code, '突合')
    assert "' 見出し: 伝票番号 支出額 移行後金額\n" in new_code and vf._check_fit(case, new_code, ask)[1] is None

def test_clash_counts_any_other_job_text(monkeypatch):
    """2026-09-18 朝: 入口は当たった登録を全部撃つ＝弱く当たる語でも一緒に撃たれる（比較の依頼で突合が撃たれ
    「相手の金額」の列が付いた）。ほかの仕事の文に出る語は、勝ち負けに関係なくぶつかり。"""
    monkeypatch.setattr(vf, '_OTHER_JOB_TEXTS', {'月次集計': ['月別の支出を課ごとに集計'], '縦持ち': ['T月別支出をピボット解除して']})
    monkeypatch.setattr(vf, '_OTHER_JOB_ASKS', {'月次集計': '月別の支出を課', '縦持ち': 'ピボット解除'})
    assert vf._clash_jobs('月別の支出') == [('月次集計', '月別の支出を課ごとに集計')]
    assert vf._clash_jobs('支出をピボット') == [('縦持ち', 'T月別支出をピボット解除して')]
    assert '月別の支出' not in vf._ask_candidates('月別の支出をピボットでいつもの見せ方に', n=12)

def test_trap_hints_line_ends_with_ampersand_without_continuation():
    """2026-09-18: M の式を行をまたいで & で足し、行末の「 _」が無く構文エラー（ヒントは「呼んでいる Function が無い」だけだった）。"""
    bad = 'm = "let" & Chr(13) & Chr(10) &\n    "in"\n'
    good = 'm = "let" & Chr(13) & Chr(10) & _\n    "in"\n'
    assert any('行の継続' in x for x in vf._trap_hints(bad))
    assert not any('行の継続' in x for x in vf._trap_hints(good))

def test_header_alternatives_with_bar(monkeypatch):
    """2026-09-18: ピボットを値で貼り付けた集計表は「課名」でなく「行ラベル」＝「課名|行ラベル」で両方の表に撃てる。"""
    import vbam_prefire as vp
    assert vp.head_in('課名|行ラベル', {'行ラベル', '需用費'}) and not vp.head_in('課名|行ラベル', {'科目'})
    monkeypatch.setattr(vf, '_OTHER_JOB_TEXTS', {})
    monkeypatch.setattr(vf, '_OTHER_JOB_WORDS', {})
    code = "Sub 積み上げ()\n    ' 依頼の語: 積み上げ縦棒\n    ' 扱う: グラフ\n    ' 見出し: 課名|行ラベル\nEnd Sub\n"
    case = {'name': '積み上げ', 'request': '積み上げ縦棒にして', 'phrases': [], 'keep_ask': '',
            'before_values': {'values': [['課名', '需用費', '合計']]},
            'tests': [{'label': 'ピボット', 'expect': {'values': [['合計 / 支出額', '列ラベル'], ['行ラベル', '需用費', '総計']]}}]}
    assert vf._check_fit(case, code, '積み上げ縦棒')[1] is None
    reg = vp.registry_from_text(code)
    assert vp.plan_full('積み上げ縦棒にして', reg, {'行ラベル', '需用費'})['extras']

def test_trap_hints_object_without_set():
    assert any('Set が要ります' in x for x in vf._trap_hints('    wb = ActiveWorkbook\n'))
    assert not any('Set が要ります' in x for x in vf._trap_hints('    Set wb = ActiveWorkbook\n    n = wb.Worksheets(1).Name\n'))

# ----------------------------------------------------------------
# どの表でも動く形（2026-09-18: 鍛えた 27 本が課名・支出額・T支出明細を書き込み、職場の表で動かなかった）
# ----------------------------------------------------------------

def test_header_none_line_fires_on_any_table(monkeypatch):
    """' 見出し: なし ＝語に頼らないマクロ。登録簿の headers は []・見出しの検査もぶつかりの検査もしない。"""
    monkeypatch.setattr(vf, '_OTHER_JOB_TEXTS', {})
    monkeypatch.setattr(vf, '_OTHER_JOB_WORDS', {'月次集計': {'所属', '金額'}})
    code = "Sub 合計行を足す()\n    ' 依頼の語: 合計行\n    ' 扱う: 合計\n    ' 見出し: なし\nEnd Sub\n"
    reg = vp.registry_from_text(code)
    assert reg[0]['headers'] == []
    assert vp.plan_full('合計行を足して', reg, {'所属', '金額'})['extras']
    case = {'name': '合計行', 'request': '表の下に合計行を足して', 'phrases': [], 'keep_ask': '',
            'before_values': {'values': [['課名', '支出額'], ['総務課', '1']]},
            'tests': [{'label': 'A', 'expect': {'values': [['所属', '金額'], ['財政', '2'], ['合計', '2']]}}]}
    assert vf._check_fit(case, code, '合計行') == ([], None)
    new_code, _ask = vf._auto_fix_ask(case, code, '合計行')
    assert "' 見出し: なし\n" in new_code


def test_code_literals_skip_comments_and_unescape():
    code = ('    s = "課名" \' "支出額" はコメント\n'
            '    t = "a""b" & "所属"\n'
            "    ' 見出し: \"科目\"\n"
            '    Rem "費目"\n')
    assert vf._code_literals(code) == ['課名', 'a"b', '所属']


def test_literal_words_flag_words_that_differ_between_tables():
    before = {'sheet': '支出明細', 'values': [['課名', '支出額', '合計'], ['総務課', '100', '']], 'kinds': ['sss', 'sn-'],
              'others': [{'name': '課マスタ', 'head': '（左上 A1・2 行×2 列）\n1\t課コード\t課名'}],
              'objects': {'tables': [{'name': 'T支出明細'}], 'queries': []}}
    test = {'label': 'A', 'expect': {'sheet': 'データ', 'values': [['所属', '金額', '合計'], ['財政', '5', '']],
                                     'kinds': ['sss', 'sn-'], 'others': [], 'objects': {'tables': [], 'queries': []}}}
    case = {'request': '表の下に合計の行を足して（番号の列は足さない）', 'phrases': ['合計行をお願い'],
            'before_values': before, 'expect': before, 'tests': [test]}
    code = ('Sub x()\n    ws.Cells(1, 1) = "課名"\n    Set t = ws.ListObjects("T支出明細")\n    Set m = Sheets("課マスタ")\n'
            '    If v = "合計" Then Exit Sub\n    n = "番号"\n    k = "100"\n    Set s = Sheets("支出明細")\n'
            '    f = "#,##0"\n    a = "課名,所属"\nEnd Sub\n')
    words = [w for w, _has, _miss in vf._literal_words(case, code)]
    # 合計（一般の語）・番号（依頼の語）・数・書式は言わない。「課名,所属」の一片の「所属」（別の表にだけある語）は言う
    assert words == ['課名', 'T支出明細', '課マスタ', '支出明細', '所属']
    err = vf._literal_error(case, code)
    assert err and '表の形で探す' in err and '「課名」' in err
    # 表が 1 枚（別の表なし）なら違いが見えない＝言わない
    assert vf._literal_words(dict(case, tests=[]), code) == []
    # 全部の表に出る語は書いてよい
    both = dict(case, tests=[{'label': 'B', 'expect': dict(test['expect'], values=[['課名', '金額']])}])
    assert '課名' not in [w for w, _h, _m in vf._literal_words(both, 'Sub x()\n  a = "課名"\nEnd Sub\n')]


def test_forge_loop_does_not_pass_macro_with_table_words(tmp_path, monkeypatch, capsys):
    """値が全部そろっても、表の語を書き込んだ版は合格にしない（本体を直させる）。"""
    monkeypatch.setattr(vf, '_AGENT_FORGE_FILE', str(tmp_path / '_agent_forge.json'))
    monkeypatch.setattr(vf, '_AGENT_FORGE_DIR', str(tmp_path))
    monkeypatch.setattr(vf, '_OTHER_JOB_TEXTS', {})
    monkeypatch.setattr(vf, '_OTHER_JOB_WORDS', {})
    c = _case()
    bed = tmp_path / 'x.xlsx'
    bed.write_bytes(b'x')
    c.update({'before': str(bed), 'request': '表の下に合計行を足して', 'prefire': [],
              'before_values': {'row': 1, 'col': 1, 'values': [['課名', '支出額']], 'kinds': ['ss'], 'sheet': '明細'},
              'tests': [{'label': 'A', 'before': str(bed), 'expect': {'row': 1, 'col': 1, 'values': [['所属', '金額']],
                                                                     'kinds': ['ss'], 'sheet': 'データ'}}]})
    vf._forge_save({'n': c})
    bad = ("Sub 合計行を足す()\n    ' 依頼の語: 合計の行\n    ' 扱う: 合計\n    ' 見出し: なし\n"
           "    If ws.Cells(1, 1).Value = \"課名\" Then ws.Cells(3, 1).Value = \"合計\"\nEnd Sub\n")
    good = bad.replace('ws.Cells(1, 1).Value = "課名"', 'VarType(ws.Cells(2, 2).Value) = vbDouble')
    replies = iter([bad, good])
    prompts = []
    monkeypatch.setattr(vf, '_ask_once', lambda ai, model, prompt: (prompts.append(prompt), next(replies))[1])
    monkeypatch.setattr(vf, '_try_macro', lambda case, bas, sub: ([], 0.1))
    monkeypatch.setattr(vf, '_try_tests', lambda case, bas, sub: ([('A', [])], 0.1))
    monkeypatch.setattr(vf, '_check_bas', lambda p: True)
    assert vf.forge('n', max_turns=3) is True
    rec = vf._forge_load()['n']
    assert rec['passed'] and rec['turns'] == 2 and '課名' not in rec['code']
    assert '表の語をマクロに書き込んでいます' in prompts[1] and '「課名」' in prompts[1]


def test_prompt_shows_reference_code_only_on_first_turn():
    c = _case()
    c['ref_code'] = 'Sub 前の版()\nEnd Sub'
    assert '前の版（表の語を書き込んでいた' in vf._prompt(c)
    assert '前の版（表の語を書き込んでいた' not in vf._prompt(c, 'A7: 期待', 'Sub x()\nEnd Sub')


def test_select_line_and_column_pick():
    """' 選ぶ列: N ＝列を選ぶ仕事。入口は依頼の語と見出しが一致する列を、依頼に出てくる順に選ぶ（2026-09-18）。"""
    code = ("Sub ピボットを作る()\n    ' 依頼の語: ピボットにして\n    ' 扱う: ピボット\n    ' 見出し: なし\n"
            "    ' 選ぶ列: 3\nEnd Sub\n")
    reg = vp.registry_from_text(code)
    assert reg[0]['select'] == (3, 3)
    assert vp.registry_from_text(code.replace('選ぶ列: 3', '選ぶ列: 2-3'))[0]['select'] == (2, 3)
    cells = [(1, '支出日'), (2, '所属'), (3, '費目'), (4, '摘要'), (5, '金額')]
    assert vp.pick_columns('所属と費目で金額のピボットにして', cells, (3, 3)) == [2, 3, 5]
    assert vp.pick_columns('金額を所属ごとに', cells, (3, 3)) is None        # 2 列しか当たらない＝入口は撃たない
    assert vp.pick_columns('金額を所属ごとに', cells, (2, 3)) == [5, 2]      # 依頼に出てくる順
    assert vp.pick_columns('所属（部）と 金 額 で', [(1, '所属'), (2, '金額')]) == [1, 2]   # 空白・全角のならしも通す


def test_check_fit_requires_select_line_when_case_has_selection(monkeypatch):
    monkeypatch.setattr(vf, '_OTHER_JOB_TEXTS', {})
    monkeypatch.setattr(vf, '_OTHER_JOB_WORDS', {})
    base = "Sub ピボットを作る()\n    ' 依頼の語: ピボットにして\n    ' 扱う: ピボット\n    ' 見出し: なし\n{sel}End Sub\n"
    case = {'name': 'ピボット', 'request': '選んでいる列でピボットにして', 'phrases': [], 'keep_ask': '', 'sel': 'B1,C1,E1',
            'before_values': {'values': [['支出日', '課名', '科目', '摘要', '支出額']]},
            'tests': [{'label': 'A', 'expect': {'values': [['日付', '所属', '費目', '内容', '金額']]}}]}
    heads, why = vf._check_fit(case, base.format(sel=''), 'ピボットにして')
    assert heads is None and "' 選ぶ列: 3" in why
    assert vf._check_fit(case, base.format(sel="    ' 選ぶ列: 2\n"), 'ピボットにして')[1].startswith("' 選ぶ列: の数")
    assert vf._check_fit(case, base.format(sel="    ' 選ぶ列: 3\n"), 'ピボットにして') == ([], None)
    assert vf._check_fit(case, base.format(sel="    ' 選ぶ列: 2-3\n"), 'ピボットにして') == ([], None)
    assert vf._sel_of('x/お題_試験A_前.xlsx') is None


def test_validate_rejects_statements_vba_cannot_load():
    """2026-09-18: M の式を 27〜33 行「& _」でつないだ版が取り込めず、AI には理由が返らなかった（継続は 25 本まで）。"""
    body = "\n".join('    m = m & "x" & _' for _ in range(30)) + '\n    m = m & "end"\n'
    code = ("Sub 縦持ちにする()\n    ' 依頼の語: 縦持ち\n    ' 扱う: 縦持ち\n    ' 見出し: なし\n" + body + "End Sub\n")
    name, why, _h = vf._validate_code(code)
    assert name is None and '25 本まで' in why
    ok = ("Sub 縦持ちにする()\n    ' 依頼の語: 縦持ち\n    ' 扱う: 縦持ち\n    ' 見出し: なし\n"
          + "\n".join('    m = m & "x" & vbCrLf' for _ in range(30)) + "\nEnd Sub\n")
    assert vf._validate_code(ok)[0] == '縦持ちにする'
    long_line = ok.replace('    m = m & "x" & vbCrLf', '    m = "' + 'あ' * 1100 + '"', 1)
    assert '1023 字' in vf._validate_code(long_line)[1]


def test_hints_row_shift_and_duplicated_output():
    """2026-09-18: 表題の行を見出しにして 1 行上に書いた／2 回目に同じ列をもう 1 つ作った、を言葉にする。"""
    exp = {'row': 1, 'col': 1, 'values': [['', '', ''], ['', '', ''], ['コード', '区分', '金額']]}
    mism = ["A1: 期待 '' ／ マクロ後 'コード'", "B1: 期待 '' ／ マクロ後 '区分'", "C1: 期待 '' ／ マクロ後 '金額'",
            "A3: 期待 'コード' ／ マクロ後 ''"]
    h = vf._hints(mism, exp)
    assert '書いた行が 2 行上にずれています' in h
    dup = ["H1: 期待 '' ／ マクロ後 '統計区分'", "J1: 期待 '' ／ マクロ後 '統計区分'"]
    assert '同じ見出し「統計区分」を 2 か所' in vf._hints(dup)


def test_check_fit_rejects_words_not_in_request_or_phrases(monkeypatch):
    """2026-09-18: 依頼の語に「月次_試験」（お題の別の表の名前）が入った＝人の依頼には出ない語は外す。"""
    monkeypatch.setattr(vf, '_OTHER_JOB_TEXTS', {})
    monkeypatch.setattr(vf, '_OTHER_JOB_WORDS', {})
    code = "Sub 月次集計表を作る()\n    ' 依頼の語: 月次集計表|月次_試験\n    ' 扱う: 月次集計\n    ' 見出し: なし\nEnd Sub\n"
    case = {'name': '月次集計', 'request': '月次集計表を作って', 'phrases': ['月別の集計表を'], 'keep_ask': '',
            'before_values': {'values': [['課名', '支出日', '支出額']]}, 'tests': []}
    heads, why = vf._check_fit(case, code, '月次集計表|月次_試験')
    assert heads is None and '月次_試験' in why and '出ない語' in why
    new_code, ask = vf._auto_fix_ask(case, code, '月次集計表|月次_試験')
    assert new_code and '月次_試験' not in ask and '月次集計表' in ask


def test_trap_hints_strconv_drops_non_cp932_chars():
    """2026-09-18: 「補償補塡及び賠償金」が StrConv で「補償補?及び賠償金」になった（ANSI＝cp932 を通る）。"""
    start = {'values': [['科目', '予算額'], ['補償補塡及び賠償金', 1000]], 'kinds': ['ss', 'sn']}
    h = vf._trap_hints('    s = Trim$(StrConv(s, vbNarrow))\n', start)
    assert any('cp932（Shift-JIS）に無い字' in x for x in h)
    assert not vf._trap_hints('    s = Trim$(s)\n', start)


def test_pick_columns_prefers_left_when_header_repeats():
    """2026-09-18: 前の集計表が右にあると同じ見出しが 2 つ並び、入口が列を決められなかった＝左を採る。"""
    cells = [(1, '支出日'), (2, '事業名'), (5, '支出額'), (7, '事業名'), (8, '4月')]
    assert vp.pick_columns('事業名と支出日と支出額で作って', cells, (3, 3)) == [2, 1, 5]


def test_pick_columns_ignores_right_block_when_too_many_hit():
    """2026-09-18 朝: 前の集計表が右に残ると、依頼文の「4月」「3月」「合計」がその見出しに当たって列が決まらず
    （当たりすぎで None）、月次集計が入口から撃てなかった＝空の列で切れた左の塊だけで選び直す。"""
    cells = [(1, '伝票日付'), (2, '所属名'), (3, '節'), (4, '備考'), (5, '支払額'),
             (7, '所属名'), (8, '4月'), (19, '3月'), (20, '合計')]
    req = '所属名と伝票日付と支払額で 明細から月別の集計表を作る。4月から3月・合計も'
    assert vp.pick_columns(req, cells, (3, 3)) == [2, 1, 5]


def test_hints_constant_ratio_off():
    """2026-09-18: 按分で全部の値が同じ倍率（1/10000）でずれた＝割る基準の取り方が違う。"""
    mism = ["E2: 期待 '263521' ／ マクロ後 '26'", "F2: 期待 '503086' ／ マクロ後 '50'",
            "G2: 期待 '2204000' ／ マクロ後 '220'"]
    h = vf._hints(mism)
    assert '同じ倍率でずれています' in h and '基準の合計' in h


def test_trap_hints_chr_over_255_is_not_the_char_you_mean():
    """2026-09-18 実測: 日本語の Excel では Chr(12288)="0"（ChrW(12288) が全角スペース）。按分が金額から 0 を全部消した。"""
    h = vf._trap_hints('    If chn <> Chr(12288) Then tmp = tmp & chn\n')
    assert any('Chr(12288) は全角スペースではなく「0」' in x for x in h)
    assert not any('ANSI（cp932）で解釈' in x for x in vf._trap_hints('    s = Replace(s, ChrW(12288), "")\n'))


def test_trap_hints_select_job_loops_only_selection():
    """2026-09-18: 数値直しで For Each cell In Selection（見出しのセル 1 つ）だけを回し、本文に何も書かず 6 往復外れた。"""
    code = "    ' 選ぶ列: 1\n    Set sel = Selection\n    For Each area In sel.Areas\n        For Each cell In area.Cells\n"
    assert any('見出しのセル（1 列に 1 つ）だけ' in x for x in vf._trap_hints(code))
    ok = code + "    last = ws.Cells(ws.Rows.Count, c).End(xlUp).Row\n"
    assert not any('見出しのセル（1 列に 1 つ）だけ' in x for x in vf._trap_hints(ok))
    assert not any('見出しのセル（1 列に 1 つ）だけ' in x for x in vf._trap_hints(code.replace("' 選ぶ列: 1", "' 選ぶ列: なし")))


def test_trap_hints_black_triangle_codepoint():
    """2026-09-18: ▲ を ChrW(9652)（▴）と書き、▲ の数が負にならなかった。"""
    assert any('▲＝ChrW(9650)' in x for x in vf._trap_hints('    If fc = ChrW(9651) Or fc = ChrW(9652) Then\n'))
    assert not any('▲＝ChrW(9650)' in x for x in vf._trap_hints('    If fc = ChrW(9651) Or fc = ChrW(9650) Then\n'))


def test_trap_hints_sparkline_has_no_line():
    """2026-09-18 第二期: スパークラインの色を sg.Line.Color.RGB と書いて 438。"""
    bad = '    Set sg = ws.Range("P2").SparklineGroups.Add(1, "B2:M2")\n    sg.Line.Color.RGB = RGB(31, 78, 121)\n'
    assert any('SeriesColor.Color' in x for x in vf._trap_hints(bad))
    good = '    Set sg = ws.Range("P2").SparklineGroups.Add(1, "B2:M2")\n    sg.SeriesColor.Color = RGB(31, 78, 121)\n'
    assert not any('SeriesColor.Color' in x for x in vf._trap_hints(good))


def test_trap_hints_yen_not_said_when_macro_replaces_yen():
    """2026-09-18: 円を Replace で消してから IsNumeric しているマクロに「IsNumeric は円で False」と言い、4 往復迷わせた。"""
    start = {'values': [['金額'], ['5,900円']], 'kinds': ['s', 's']}
    bad = '    If IsNumeric(s) Then v = CDbl(s)\n'
    assert any('円・△が付いたもの' in x for x in vf._trap_hints(bad, start))
    good = '    s = Replace(s, "円", "")\n' + bad
    assert not any('円・△が付いたもの' in x for x in vf._trap_hints(good, start))

def test_inv_ignores_filtered_hidden_rows_after_run():
    """2026-09-18 朝: 「0 の行は隠して」と頼んで絞り込んだのに「隠れていた行が変わった」で AI に回っていた。
    撃った後に絞り込みが掛かっていれば、隠れた行は見せ方＝咎めない（テーブルの絞り込みも見る）。"""
    import vbam_inv as vi
    b = {'sheets': {'S': {'used': 'A1:B9', 'values': [['a', 1]], 'formulas': {}, 'hidden': (), 'cf': 0,
                          'validation': '', 'filter': '', 'merged': False, 'tables': {}, 'charts': 0}},
         'order': ['S'], 'names': [], 'calc': None}
    import copy
    a = copy.deepcopy(b)
    a['sheets']['S']['hidden'] = (3, 5)
    a['sheets']['S']['filter'] = 'A1:B9'          # テーブルの絞り込みで隠れた
    assert not [x for x in vi._inv_compare(b, a, 'S', ()) if '隠れていた行' in x]
    a2 = copy.deepcopy(b)
    a2['sheets']['S']['hidden'] = (3, 5)          # 手で隠した（絞り込み無し）＝咎める
    assert [x for x in vi._inv_compare(b, a2, 'S', ()) if '隠れていた行' in x]

def test_inv_table_renamed_is_not_gone():
    """2026-09-18 朝: 「テーブル1」を T職員名簿 に直すのが仕事なのに「テーブルが消えた」で AI に回っていた。
    同じ左上に別の名前のテーブルがあれば、消えたのでなく名前を付け替えた。"""
    import copy
    import vbam_inv as vi
    b = {'sheets': {'S': {'used': 'A1:E11', 'values': [['a']], 'formulas': {}, 'hidden': (), 'cf': 0,
                          'validation': '', 'filter': '', 'merged': False, 'charts': 0,
                          'tables': {'テーブル1': 'A1:E11'}}},
         'order': ['S'], 'names': [], 'calc': None}
    a = copy.deepcopy(b)
    a['sheets']['S']['tables'] = {'T職員名簿': 'A1:E12'}          # 名前を付け替えて集計行を足した
    assert not [x for x in vi._inv_compare(b, a, 'S', ()) if 'テーブル' in x]
    a2 = copy.deepcopy(b)
    a2['sheets']['S']['tables'] = {}                              # 本当に消えた＝咎める
    assert [x for x in vi._inv_compare(b, a2, 'S', ()) if 'テーブル' in x and '消えた' in x]


# ----------------------------------------------------------------
# 表の形（' 形:）＝入口が語の前に絞る必要条件（2026-09-18）
# ----------------------------------------------------------------

def _book(path, rows, sheet='明細', table=None, merges=(), other=None):
    """お題のブックを 1 冊作る（openpyxl）。other＝2 枚目のシートの行。"""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(list(r))
    for m in merges:
        ws.merge_cells(m)
    if table:
        from openpyxl.worksheet.table import Table
        ws.add_table(Table(displayName=table[0], ref=table[1]))
    if other is not None:
        ws2 = wb.create_sheet('別表')
        for r in other:
            ws2.append(list(r))
    wb.save(str(path))
    return str(path)


def _rich_book(tmp_path, name='rich.xlsx'):
    import datetime as dt
    rows = [['課名', '受付日', '金額'],
            ['総務課', dt.datetime(2026, 4, 1), 1000],
            ['財政課', dt.datetime(2026, 4, 2), 2000],
            ['税務課', dt.datetime(2026, 4, 3), 3000]]
    return _book(tmp_path / name, rows, table=('T明細', 'A1:C4'),
                 other=[['課名', '受付日', '金額'], ['総務課', dt.datetime(2026, 5, 1), 9]])


def test_shape_of_file_reads_kinds_table_and_other_sheets(tmp_path):
    """ブックのファイル → 表の形（openpyxl・Excel を起こさない）。入口の COM と同じ物差し。"""
    got = vf.shape_of_file(_rich_book(tmp_path))
    assert got == {'数の列', '日付の列', 'テーブル', '別のシート', '同じ見出しのシート', '共通の見出しのシート'}
    # 見出しがかすりもしない別のシートは「別のシート」だけ
    import datetime as dt
    p2 = _book(tmp_path / 'b.xlsx', [['課名', '受付日', '金額'], ['総務課', dt.datetime(2026, 4, 1), 1]],
               other=[['品名', '個数'], ['鉛筆', 3]])
    got2 = vf.shape_of_file(p2)
    assert '別のシート' in got2 and '同じ見出しのシート' not in got2 and '共通の見出しのシート' not in got2
    assert 'テーブル' not in got2
    assert vf.shape_of_file(str(tmp_path / '無い.xlsx')) == set()


def test_shape_of_file_merged_header_only_counts_the_header_rows(tmp_path):
    """結合の見出し＝見出しの行〜+2 行の結合。表題の行だけの結合は数えない（帳票と普通の表を見分ける）。"""
    head = _book(tmp_path / 'head.xlsx',
                 [['課名', '支出', None, '備考'], [None, '予算', '執行', None],
                  ['総務課', 1, 2, 'x'], ['財政課', 3, 4, 'y']], merges=('B1:C1',))
    assert '結合の見出し' in vf.shape_of_file(head)
    title = _book(tmp_path / 'title.xlsx',
                  [['支出一覧', None, None], ['課名', '予算', '執行'], ['総務課', 1, 2], ['財政課', 3, 4]],
                  merges=('A1:C1',))
    assert '結合の見出し' not in vf.shape_of_file(title)


def test_job_shape_is_the_intersection_of_all_tables(tmp_path):
    """' 形: は「合格した表の全部に共通する形」＝積集合。1 枚に無い形は書かない（証明した範囲でしか撃たない）。"""
    rich = _rich_book(tmp_path)
    plain = _book(tmp_path / 'plain.xlsx', [['所属', '金額'], ['財政', 5], ['税務', 6]])
    assert vf._job_shape({'before': rich, 'sheet': '明細', 'tests': []}) == \
        ['数の列', '日付の列', 'テーブル', '別のシート', '同じ見出しのシート', '共通の見出しのシート']
    case = {'before': rich, 'sheet': '明細', 'tests': [{'label': 'A', 'before': plain, 'sheet': '明細'}]}
    assert vf._job_shape(case) == ['数の列']
    # 読めない表は数えない（積集合を空にしない）
    case2 = dict(case, tests=[{'label': 'B', 'before': str(tmp_path / '無い.xlsx'), 'sheet': None}])
    assert vf._job_shape(case2) == ['数の列', '日付の列', 'テーブル', '別のシート', '同じ見出しのシート', '共通の見出しのシート']


def test_ensure_shape_line_is_written_by_the_tool(tmp_path):
    """' 形: は道具が入れる（AI には書かせない）。' 選ぶ列: があればその次・無ければ ' 見出し: の次。"""
    plain = _book(tmp_path / 'plain.xlsx', [['所属', '金額'], ['財政', 5]])
    case = {'before': plain, 'sheet': '明細', 'tests': []}
    code = "Sub 合計行を足す()\n    ' 依頼の語: 合計行\n    ' 扱う: 合計\n    ' 見出し: なし\n    Dim a\nEnd Sub\n"
    got, changed = vf.ensure_shape_line(case, code)
    assert changed and got.split('\n')[4] == "    ' 形: 数の列"
    assert vf.ensure_shape_line(case, got) == (got, False)          # 同じ形なら触らない
    sel = code.replace("    ' 見出し: なし\n", "    ' 見出し: なし\n    ' 選ぶ列: 2\n")
    got2, changed2 = vf.ensure_shape_line(case, sel)
    assert changed2 and got2.split('\n')[5] == "    ' 形: 数の列"   # 選ぶ列の次
    # 形が変わったら書き換える（前の行は残さない）
    got3, changed3 = vf.ensure_shape_line(case, got.replace("' 形: 数の列", "' 形: テーブル 日付の列"))
    assert changed3 and got3.count("' 形:") == 1 and "' 形: 数の列" in got3
    # 形が 1 つも共通しなければ「なし」＝形では絞らない
    other = _book(tmp_path / 'words.xlsx', [['所属', '備考'], ['財政', 'x']])
    got4, _c = vf.ensure_shape_line({'before': other, 'sheet': '明細', 'tests': []}, code)
    assert "' 形: なし" in got4
    assert vp.registry_from_text(got4)[0]['shape'] == []
    assert vp.registry_from_text(got)[0]['shape'] == ['数の列']


def test_check_fit_rejects_a_shape_missing_in_one_table(tmp_path, monkeypatch):
    """' 形: の語が 1 枚でも無ければ、その表では撃たれない＝規則違反。行が無いことは違反にしない（道具が足す）。"""
    monkeypatch.setattr(vf, '_OTHER_JOB_TEXTS', {})
    monkeypatch.setattr(vf, '_OTHER_JOB_WORDS', {})
    rich = _rich_book(tmp_path)
    plain = _book(tmp_path / 'plain.xlsx', [['所属', '金額'], ['財政', 5]])
    case = {'name': '合計行', 'request': '表の下に合計行を足して', 'phrases': [], 'keep_ask': '',
            'before': rich, 'sheet': '明細',
            'before_values': {'values': [['課名', '受付日', '金額'], ['総務課', '2026/4/1', '1000']]},
            'tests': [{'label': 'A', 'before': plain, 'sheet': '明細',
                       'expect': {'values': [['所属', '金額'], ['財政', '5'], ['合計', '5']]}}]}
    head = "Sub 合計行を足す()\n    ' 依頼の語: 合計行\n    ' 扱う: 合計\n    ' 見出し: なし\n"
    assert vf._check_fit(case, head + "End Sub\n", '合計行') == ([], None)          # 行が無い＝違反にしない
    assert vf._check_fit(case, head + "    ' 形: 数の列\nEnd Sub\n", '合計行') == ([], None)
    _h, why = vf._check_fit(case, head + "    ' 形: 数の列 テーブル\nEnd Sub\n", '合計行')
    assert why and 'テーブル' in why and '別の表「A」' in why
    _h2, why2 = vf._check_fit(case, head + "    ' 形: 太字の列\nEnd Sub\n", '合計行')
    assert why2 and '使えない語' in why2


def test_plan_full_skips_a_job_whose_shape_is_missing():
    """入口: 依頼の語が当たっても、表の形がそろわなければ撃たない（語のかすり当たりの誤爆を形で先に止める）。"""
    code = ("Sub テーブルを絞り込む()\n    ' 依頼の語: 絞り込\n    ' 扱う: 絞り込み\n    ' 見出し: なし\n"
            "    ' 形: テーブル\nEnd Sub\n")
    reg = vp.registry_from_text(code)
    assert reg[0]['shape'] == ['テーブル']
    p = vp.plan_full('この表を絞り込んで', reg, None, {'数の列'})
    assert not p['extras'] and [e['name'] for e in p['skipped']] == ['テーブルを絞り込む']
    assert p['skipped'][0]['why_shape'] == ['テーブル']
    assert [e['name'] for e in vp.plan_full('この表を絞り込んで', reg, None, {'数の列', 'テーブル'})['extras']] \
        == ['テーブルを絞り込む']
    # 形を読めなかったとき（shape=None）と、形の行が無い Sub は従来どおり撃つ
    assert vp.plan_full('この表を絞り込んで', reg, None, None)['extras']


def test_ask_len_allows_long_katakana_word_only():
    """2026-09-18 第二期: 「ウォーターフォール」（9 字）が 8 字の決まりで外され、「ウォーターフォール図を作って」で撃てなくなった。"""
    assert vf._ask_len_ok('ウォーターフォール') and vf._ask_len_ok('散布図')
    assert not vf._ask_len_ok('選んでいる列を散布図') and not vf._ask_len_ok('ア')
    assert 'ウォーターフォール' in vf._ask_candidates('ウォーターフォール図を作って', [])


def test_compare_book_names_and_other_sheet_values():
    """2026-09-18 第二期: 帳票の仕事はその仕事のシート以外に書く＝シートの並びとほかのシートの値も比べる。古い台帳（book 無し）は比べない。"""
    base = {'row': 1, 'col': 1, 'values': [['a']]}
    book = {'names': ['一覧', '101', '102'], 'sheets': {'101': {'row': 1, 'col': 1, 'values': [['氏名', '佐藤']]},
                                                       '102': {'row': 1, 'col': 1, 'values': [['氏名', '鈴木']]}}}
    e = dict(base, book=book)
    assert vf._compare(e, dict(base, book=book)) == []
    g = dict(base, book={'names': ['一覧', '101'], 'sheets': {'101': {'row': 1, 'col': 1, 'values': [['氏名', '田中']]}}})
    out = vf._compare(e, g)
    assert any('無いシート' in x and '102' in x for x in out) and any('シート「101」' in x and '田中' in x for x in out)
    assert vf._compare(base, g) == []


def test_compare_page_only_when_truth_has_page():
    """2026-09-18 第二期: 印刷設定の仕事。正解に印刷設定があるときだけ比べ、ずれには設定の書き方を添える。"""
    base = {'row': 1, 'col': 1, 'values': [['a']]}
    page = dict(vf._PAGE_DEFAULT, orientation='横', zoom=None, fit_wide=1, title_rows='3:3')
    out = vf._compare(dict(base, page=page), dict(base, page=dict(vf._PAGE_DEFAULT)))
    assert any('向き' in x and 'xlLandscape' in x for x in out) and any('見出しの繰り返し' in x for x in out)
    assert vf._compare(dict(base, page=page), dict(base, page=dict(page))) == []
    assert vf._compare(base, dict(base, page=page)) == []


def test_trap_hints_m_field_names_need_hash_quote():
    """2026-09-18 第二期: 見出し「金額（円）」を M の式で [金額（円）] と書いて式が読めなかった。"""
    bad = '    ActiveWorkbook.Queries.Add qn, "let S = ... List.Sum([" & amt & "])"\n'
    assert any('[#' in x and '金額（円）' in x for x in vf._trap_hints(bad))
    good = '    ActiveWorkbook.Queries.Add qn, "let S = ... List.Sum([#""" & amt & """])"\n'
    assert not any('金額（円）' in x for x in vf._trap_hints(good))


def test_compare_formula_cells_and_filter():
    """2026-09-18 第二期: 式かどうか・オートフィルタも比べる（古い台帳の正解には無い＝比べない）。"""
    base = {'row': 1, 'col': 1, 'values': [['a']]}
    e = dict(base, fx=['B2', 'B3'], filter='A1:C9')
    assert vf._compare(e, dict(base, fx=['B2', 'B3'], filter='A1:C9')) == []
    out = vf._compare(e, dict(base, fx=['B2', 'C5'], filter=''))
    assert any('式でない' in x and 'B3' in x for x in out) and any('式になっている' in x and 'C5' in x for x in out)
    assert any('オートフィルタ' in x for x in out)
    assert vf._compare(base, dict(base, fx=['Z9'], filter='A1:B2')) == []
    assert vf._addr_key('B12') == (12, 2) and vf._addr_key('AA3') == (3, 27)


def test_trap_hints_kabushiki_mark_codepoint():
    """2026-09-18 第二期: ㈱ を ChrW(&H339E)（㎞）と書いた。"""
    assert any('㈱＝ChrW(&H3231)' in x for x in vf._trap_hints('    s = Replace(s, ChrW(&H339E), "")\n'))
    assert not any('㈱＝ChrW(&H3231)' in x for x in vf._trap_hints('    s = Replace(s, ChrW(&H3231), "")\n'))


def test_ask_words_never_cross_punctuation():
    """2026-09-18 第二期: 語の候補に「出す。表」「式を、今」のような句読点をまたぐ切れ端が出た。"""
    c = vf._ask_candidates('累計を出す。表の右に足して', [])
    assert c and not any(ch in w for w in c for ch in '、。')
    code = "Sub 累計列を挿入する()\n    ' 依頼の語: 累計列|出す。表\n    ' 扱う: 累計\n    ' 見出し: なし\nEnd Sub\n"
    case = {'request': '累計列を足して。累計を出す。表の右に', 'phrases': [], 'before_values': {'values': [['a', 'b']]},
            'tests': [], 'keep_ask': ''}
    _h, why = vf._check_fit(case, code, '累計列|出す。表')
    assert '句読点をまたいだ語' in (why or '')


def _zip_book(path, sheet_rels):
    """xlsx の骨組みだけ（workbook.xml と関係ファイル）を書く。sheet_rels＝1 枚目のシートの関係（(種類, 行き先) の並び）。"""
    import zipfile
    R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
    with zipfile.ZipFile(path, 'w') as z:
        z.writestr('xl/workbook.xml', '<workbook xmlns:r="x"><sheets><sheet name="集計" sheetId="1" r:id="rId1"/>'
                                      '<sheet name="明細" sheetId="2" r:id="rId2"/></sheets></workbook>')
        z.writestr('xl/_rels/workbook.xml.rels', f'<Relationships><Relationship Id="rId1" Type="{R}worksheet" Target="worksheets/sheet1.xml"/>'
                                                 f'<Relationship Id="rId2" Type="{R}worksheet" Target="worksheets/sheet2.xml"/></Relationships>')
        rels = ''.join(f'<Relationship Id="rId{i}" Type="{R}{t}" Target="{g}"/>' for i, (t, g) in enumerate(sheet_rels, 1))
        z.writestr('xl/worksheets/_rels/sheet1.xml.rels', f'<Relationships>{rels}</Relationships>')
        z.writestr('xl/drawings/_rels/drawing1.xml.rels', f'<Relationships><Relationship Id="rId1" Type="{R}chart" '
                                                          'Target="../charts/chart1.xml"/></Relationships>')


def test_objects_of_file_reads_pivot_and_chart_per_sheet(tmp_path):
    """2026-09-18 第二期: ピボット・グラフの仕事は、そのシートにピボット・グラフがあるときだけ撃つ＝ファイルからも読む（openpyxl はグラフを読まない）。"""
    p = str(tmp_path / 'a.xlsx')
    _zip_book(p, [('pivotTable', '../pivotTables/pivotTable1.xml'), ('drawing', '../drawings/drawing1.xml')])
    assert vf._objects_of_file(p) == (True, True)
    assert vf._objects_of_file(p, '集計') == (True, True)
    assert vf._objects_of_file(p, '明細') == (False, False)
    q = str(tmp_path / 'b.xlsx')
    _zip_book(q, [])
    assert vf._objects_of_file(q) == (False, False)
    assert vf._objects_of_file(str(tmp_path / 'none.xlsx')) == (False, False)


def test_shape_signals_pivot_and_chart_words():
    import vbam_prefire as vp
    k = ['ss', 'sn', 'sn']
    v = [['課', '額'], ['a', '1'], ['b', '2']]
    assert {'ピボット', 'グラフ'} <= vp.shape_signals(v, k, has_pivot=True, has_chart=True)
    assert not ({'ピボット', 'グラフ'} & vp.shape_signals(v, k))


def test_job_shape_object_words_come_from_main_table_only(monkeypatch, tmp_path):
    """試験H（ピボットの無い表で何もしない）があっても「ピボット」を落とさない。ほかの語は今までどおり全部の表の積集合。"""
    main, h = tmp_path / 'm.xlsx', tmp_path / 'h.xlsx'
    main.write_bytes(b'x')
    h.write_bytes(b'x')
    shapes = {str(main): {'数の列', 'ピボット', '日付の列'}, str(h): {'数の列'}}
    monkeypatch.setattr(vf, 'shape_of_file', lambda p, s=None: shapes[str(p)])
    case = {'before': str(main), 'sheet': None, 'tests': [{'before': str(h), 'sheet': None}]}
    assert vf._job_shape(case) == ['数の列', 'ピボット']


def test_forge_writes_macros_with_opus_by_default(monkeypatch):
    """2026-09-19 shu「マクロを書く頭の既定は Opus に」。名指しが無く頭が Claude Code のときだけ Opus。名指しはそのまま。"""
    import vbam_forge as vf
    assert vf._forge_model(None, None) == 'opus'
    assert vf._forge_model('claude-code', None) == 'opus'
    assert vf._forge_model('claude-code', 'sonnet') == 'sonnet'       # --model の名指しが勝つ
    assert vf._forge_model('gemini', None) is None                    # ほかの頭は、その頭の既定に任せる
    assert vf._forge_model(None, 'gemini-3.7-flash') == 'gemini-3.7-flash'

    seen = {}
    import vbam_ai

    def fake_setup(ai, model):
        seen['setup'] = (ai, model)
        return 'claude-code', model, ''

    def fake_oneshot(prompt, model):
        seen['oneshot'] = model
        return ('Sub X()\nEnd Sub', {})

    monkeypatch.setattr(vbam_ai, '_ai_setup', fake_setup)
    monkeypatch.setattr(vbam_ai, '_cc_oneshot', fake_oneshot)
    assert vf._ask_once(None, None, 'p').startswith('Sub X')
    assert seen['setup'] == (None, 'opus') and seen['oneshot'] == 'opus'
