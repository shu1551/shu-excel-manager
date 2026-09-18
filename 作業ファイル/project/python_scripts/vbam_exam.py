# -*- coding: utf-8 -*-
"""vbam_exam.py — vba_manager 分割パート: 試験（正解の表を持つお題を種ごとに作り、本番と同じ道で撃ち、値で突き合わせる）

2026-09-11 夜。それまでの実射（--fire）は 6 行の練習台に細い答え合わせ、写し取り（--bed）は「壊していないか」だけ、
本番の台帳はお試し版と動画のお題（messy_customers）ばかりだった。動画のお題で初めて「重複 26 組中 3 組しか拾えない・
行ずれ・受注済みみ なのに合格」が出た＝**正解の表を持って、現実の汚れ方をした表で、値を突き合わせる**試験が無かった。

  agent --exam [--seed N] [--only お題,お題] [--dry-run] [--ai …] [--model …]
      正解の表（きれいな表）を 5 種類（顧客名簿・売上明細・会員名簿・経費精算・在庫表）作り、決まった汚れ
      （全角半角・前後の空白・（株）と株式会社・電話と郵便の書き方・文字の日付と和暦・文字の金額・「個」つきの数量・
      表記ゆれ込みの重複・文字の書式のばらつき）を種で散らしたブックを作る。お題の中身（行数・会社・値・汚れの場所）は
      種ごとに変わる＝1 つのお題に合わせて道具を育てても点は上がらない。
      各お題を本番と同じ run_agent（控え・関所・採点係・台帳）で撃ち、終わった表を正解の表とセルの値で突き合わせる。
      判定は道具の合格判定とも採点係とも別（道具が「合格」と言って試験が不合格なら、それは嘘の合格として数える）。
      --dry-run はお題のブックと依頼文を作って見せるだけ（Excel にも AI にも触らない）。

見るもの: 正しい行／足りない行（消しすぎ）／残った重複／余計な行／値の誤り／書き方の誤り（文字のままの日付・金額、
区切りの違う電話…）／列の中のばらつき／文字の書式のばらつき／秒・往復・道具の判定。
点数は _agent_exam.jsonl に 1 回 1 行で残る（種も残す＝同じ種で撃てば同じお題）。
状態ファイル（台帳・控え・覚書・記録）はこの回の置き場（_agent_exam/日時/_state）へ逃がす＝本番の台帳に混ぜない。
お題ごとの記録（.log.jsonl）はそこに控えるので、`agent --replay` でお題のブックに撃ち直せる。
"""
import os
import re
import sys
import json
import time
import random
import datetime
import unicodedata
import subprocess
from collections import Counter
from vbam_core import SCRIPT_DIR

_EXAM_DIR = os.path.join(SCRIPT_DIR, '_agent_exam')              # お題のブック・正解・結果の置き場
_EXAM_SCORE_FILE = os.path.join(SCRIPT_DIR, '_agent_exam.jsonl')  # 試験の点数（1 回 1 行・種つき）

_FONT = '游ゴシック'
_P_DIRT = 0.35               # 汚れを当てる割合（列ごと・セルごと）

# ----------------------------------------------------------------
# 材料（架空の値）
# ----------------------------------------------------------------

_SURNAMES = [('佐藤', 'サトウ'), ('鈴木', 'スズキ'), ('高橋', 'タカハシ'), ('田中', 'タナカ'), ('伊藤', 'イトウ'),
             ('渡辺', 'ワタナベ'), ('山本', 'ヤマモト'), ('中村', 'ナカムラ'), ('小林', 'コバヤシ'), ('加藤', 'カトウ'),
             ('吉田', 'ヨシダ'), ('山田', 'ヤマダ'), ('佐々木', 'ササキ'), ('山口', 'ヤマグチ'), ('松本', 'マツモト'),
             ('井上', 'イノウエ'), ('木村', 'キムラ'), ('斎藤', 'サイトウ'), ('清水', 'シミズ'), ('阿部', 'アベ'),
             ('工藤', 'クドウ'), ('石川', 'イシカワ'), ('小松', 'コマツ'), ('菅原', 'スガワラ'), ('三浦', 'ミウラ')]
_GIVEN = [('誠', 'マコト'), ('恵', 'メグミ'), ('学', 'マナブ'), ('香', 'カオリ'), ('亮', 'リョウ'), ('環', 'タマキ'),
          ('翔', 'ショウ'), ('葵', 'アオイ'), ('健一', 'ケンイチ'), ('美咲', 'ミサキ'), ('大輔', 'ダイスケ'),
          ('由美', 'ユミ'), ('拓也', 'タクヤ'), ('真理', 'マリ'), ('聡', 'サトシ'), ('陽子', 'ヨウコ'), ('隆', 'タカシ'),
          ('彩', 'アヤ'), ('直樹', 'ナオキ'), ('智子', 'トモコ')]
_COMPANY = ['アクアテック', '北斗精機', 'みらい物流', 'サンライズ食品', '東雲建設', 'ひかり電工', '大森商事',
            'グリーンファーム', 'はやて運輸', '秋桜印刷', 'ミナト化成', 'こだま製菓', '白樺ソフト', '鳥羽テクノ',
            '多摩川工業', 'ブルーオーシャン', '森の木工房', 'ユニバース通信', 'タカノ機械', '花園フーズ', 'ステラ設計',
            '稲穂農産', 'コスモ薬品', 'やまびこ観光', '銀河システム', 'しらかば住建', 'ノーザン電子', '朝日ビルメン',
            '若葉介護', '港町水産', '千歳エンジニアリング', 'オーロラ化学', '河口湖リゾート', '八ヶ岳ファーム',
            'かもしか工芸', '富士見電設', '太平洋マリン', '七夕企画', 'ひまわり製作所', '東海テック', 'わらびもち本舗',
            '双葉精密', '木葉メディカル', '紅葉ネット', '桜坂ホールディングス']
_FORMS = [('株式会社', 'pre'), ('株式会社', 'post'), ('株式会社', 'pre'), ('有限会社', 'pre')]
_CITIES = ['中央市', '北山市', '南川市', '東野市', '西原市', '港町', '川辺町', '青葉村']
_TOWNS = ['本町', '大手町', '旭町', '栄町', '宮下', '緑町', '桜台', '新町']
_AREAS = [('03', 4, 4), ('06', 4, 4), ('018', 3, 4), ('0185', 2, 4), ('022', 3, 4), ('011', 3, 4)]
_MOBILE = [('090', 4, 4), ('080', 4, 4), ('070', 4, 4)]
_PRODUCTS = [('A4コピー用紙', 450), ('B5ノート', 180), ('ボールペン黒', 120), ('ボールペン赤', 120),
             ('USBメモリ16GB', 980), ('トナーTN-27', 6800), ('LANケーブル3m', 640), ('単3電池', 90),
             ('クリアファイルA4', 30), ('付箋75mm', 210), ('ホッチキス針No.10', 150), ('電卓', 1980),
             ('封筒長3', 5), ('ファイルボックス', 580), ('マスキングテープ', 260), ('インクカートリッジBK', 1450)]
_SUBJECTS = ['旅費交通費', '会議費', '消耗品費', '通信費', '新聞図書費']
_PAYEES = ['JR東日本', '青葉中央交通', 'タクシー北斗', 'ホテル千歳', '文具のヤマダ', 'Amazon', 'NTTドコモ', '喫茶こだま']
_PLACES = ['倉庫A-1', '倉庫A-2', '倉庫B-1', '事務所棚3', '2F書庫']
_FONT_DIRT = [{'bold': True}, {'color': 'FFFF0000'}, {'size': 14}, {'italic': True}, {'name': 'ＭＳ 明朝'}]

# 汚れの種類（型ごと）
_DIRT = {
    'company': ['abbr', 'abbr', 'space', 'inner_space', 'linebreak'],
    'person': ['zspace', 'space', 'linebreak'],
    'kana': ['hankaku', 'space', 'hiragana'],
    'text': ['zenkaku', 'space', 'choon'],
    'code': ['zenkaku', 'space'],
    'phone': ['nohyphen', 'paren', 'zenkaku', 'spaces'],
    'postal': ['nohyphen', 'mark', 'zenkaku'],
    'date': ['fmt', 'text_slash', 'text_iso', 'text_kanji', 'serial', 'digits8'],
    'money': ['general', 'text_comma', 'text_yen', 'text_en', 'zenkaku'],
    'int': ['text', 'zenkaku'],
    'pct': ['text_pct', 'zen_pct', 'general', 'text_num'],
}
# 負の金額の汚れ（会計の △▲・かっこ・文字のマイナス。2026-09-12 に足した）
_DIRT_NEG = ['tri', 'blk', 'paren', 'minus_text', 'general']
_STR_TYPES = ('company', 'person', 'kana', 'text', 'code')

# ----------------------------------------------------------------
# 文字の道具
# ----------------------------------------------------------------

_KANA_F = 'アイウエオカキクケコサシスセソタチツテトナニヌネノハヒフヘホマミムメモヤユヨラリルレロワヲンァィゥェォッャュョー'
_KANA_H = 'ｱｲｳｴｵｶｷｸｹｺｻｼｽｾｿﾀﾁﾂﾃﾄﾅﾆﾇﾈﾉﾊﾋﾌﾍﾎﾏﾐﾑﾒﾓﾔﾕﾖﾗﾘﾙﾚﾛﾜｦﾝｧｨｩｪｫｯｬｭｮｰ'
_HALF_KANA_RE = re.compile(r'[ｦ-ﾟ]')
_FULL_ALNUM_RE = re.compile(r'[０-９Ａ-Ｚａ-ｚ]')
_HALF_ALNUM_RE = re.compile(r'[0-9A-Za-z]')


def _nfkc(s):
    return unicodedata.normalize('NFKC', str(s))


def _ws(s):
    return re.sub(r'\s+', ' ', s).strip()


def _zen(s):
    """英数字と - , . ( ) / を全角に。"""
    return ''.join(chr(ord(ch) + 0xFEE0) if (ch.isascii() and (ch.isalnum() or ch in '-,.()/')) else ch for ch in str(s))


def _hankaku_kana(s):
    out = []
    for ch in unicodedata.normalize('NFD', str(s)):
        if ch == '゙':
            out.append('ﾞ')
        elif ch == '゚':
            out.append('ﾟ')
        elif ch in _KANA_F:
            out.append(_KANA_H[_KANA_F.index(ch)])
        elif ch == '　':
            out.append(' ')
        else:
            out.append(ch)
    return ''.join(out)


def _pad(rng, s):
    return rng.choice([' ' + s, s + ' ', '　' + s, s + '　', ' ' + s + ' '])


def _hiragana(s):
    """カタカナ → ひらがな（長音・空白はそのまま）。"""
    return ''.join(chr(ord(ch) - 0x60) if 0x30A1 <= ord(ch) <= 0x30F6 else ch for ch in str(s))


def _digits(rng, n):
    return str(rng.randint(1, 9)) + ''.join(str(rng.randint(0, 9)) for _ in range(n - 1))


def _company_canon(v):
    s = re.sub(r'\s+', '', _nfkc(v))
    return s.replace('(株)', '株式会社').replace('(有)', '有限会社')


def _has_abbr(v):
    return bool(re.search(r'\((株|有)\)', _nfkc(v)))


def _parse_date(s):
    """文字の日付 → 'YYYY-MM-DD'（読めなければ None）。和暦（R8.9.1・令和8年9月1日）も読む。"""
    s = _nfkc(s).strip()
    y = mo = d = None
    m = re.fullmatch(r'(\d{4})[/\-.年](\d{1,2})[/\-.月](\d{1,2})日?', s)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
    m = m or re.fullmatch(r'(?:R|令和)(\d{1,2})[.年](\d{1,2})[.月](\d{1,2})日?', s)
    if m and y is None:
        y, mo, d = 2018 + int(m.group(1)), int(m.group(2)), int(m.group(3))
    m = m or re.fullmatch(r'(\d{1,2})/(\d{1,2})/(\d{4})', s)
    if m and y is None:
        y, mo, d = int(m.group(3)), int(m.group(1)), int(m.group(2))
    m = m or re.fullmatch(r'((?:19|20)\d{2})(\d{2})(\d{2})', s)          # 8 桁（20260901）
    if m and y is None:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if y is None:
        return None
    try:
        return datetime.date(y, mo, d).isoformat()
    except ValueError:
        return None


def _date(rng, y0, m0, y1, m1):
    a = datetime.date(y0, m0, 1).toordinal()
    b = datetime.date(y1, m1, 28).toordinal()
    return datetime.date.fromordinal(rng.randint(a, b))


def _person(rng):
    (s, sk), (g, gk) = rng.choice(_SURNAMES), rng.choice(_GIVEN)
    return f"{s} {g}", f"{sk} {gk}"


def _phone(rng, areas):
    """市外局番の後ろが、もっと長い別の局番と重ならない番号だけ作る（018-505-… は 0185 と読める＝
    本物の番号の決まりでも 018 の後ろに 5 は来ない。作ると正解の表のほうが間違う・2026-09-11 夜）。"""
    longer = [a for a, _n1, _n2 in _AREAS + _MOBILE]
    while True:
        a, n1, n2 = rng.choice(areas)
        s = f"{a}-{_digits(rng, n1)}-{_digits(rng, n2)}"
        flat = s.replace('-', '')
        if not any(len(x) > len(a) and flat.startswith(x) for x in longer):
            return s


def _postal(rng):
    return f"{rng.randint(0, 999):03d}-{rng.randint(0, 9999):04d}"


def _address(rng):
    return f"架空県{rng.choice(_CITIES)}{rng.choice(_TOWNS)}{rng.randint(1, 9)}-{rng.randint(1, 30)}-{rng.randint(1, 20)}"


# ----------------------------------------------------------------
# お題（正解の表）
# ----------------------------------------------------------------

def _ex_customers(rng):
    n = rng.randint(28, 38)
    rows = []
    for b in rng.sample(_COMPANY, n):
        form, pos = rng.choice(_FORMS)
        rows.append([form + b if pos == 'pre' else b + form, _person(rng)[0], _phone(rng, _AREAS), _postal(rng),
                     _address(rng), _date(rng, 2019, 1, 2026, 8), rng.randrange(50, 5000) * 1000])
    # 罠: 同じ会社の別の担当者（電話・登録日・取引額も違う）＝重複ではない
    for k in sorted(rng.sample(range(n), 2), reverse=True):
        r = list(rows[k])
        r[1] = _person(rng)[0]
        r[2] = _phone(rng, _AREAS)
        r[5] = _date(rng, 2019, 1, 2026, 8)
        r[6] = rng.randrange(50, 5000) * 1000
        rows.insert(k + rng.randint(3, 8), r)
    return {'name': '顧客名簿', 'sheet': '顧客名簿', 'title': None,
            'columns': [['会社名', 'company'], ['担当者', 'person'], ['電話', 'phone'], ['郵便番号', 'postal'],
                        ['住所', 'text'], ['登録日', 'date'], ['取引額', 'money']],
            'rows': rows, 'key': ['会社名', '担当者'], 'dups': True, 'font_dirt': True, 'blank_rows': True,
            'request': "この顧客名簿を修正してください。重複行は削除してよい。会社名の表記ゆれ（株式会社と（株））・"
                       "セルの中の改行・全角半角・前後の空白・電話番号・郵便番号・登録日・取引額の書き方と、"
                       "文字の書式のばらつきをそろえて。"}


def _ex_sales(rng):
    n = rng.randint(30, 42)
    staff = [_person(rng)[0] for _ in range(4)]
    seen, rows = set(), []
    while len(rows) < n:
        p, price = rng.choice(_PRODUCTS)
        q = rng.randint(1, 20)
        r = [_date(rng, 2026, 4, 2026, 8), rng.choice(staff), p, q, price, rng.choice([0, 0, 0.05, 0.1, 0.15]), price * q]
        k = (r[0], r[1], r[2])
        if k in seen:
            continue
        seen.add(k)
        rows.append(r)
    rows.sort(key=lambda r: r[0])
    # 罠: 同じ担当・商品・数量の別の日（重複ではない）
    for k in rng.sample(range(n), 2):
        r = list(rows[k])
        d = r[0] + datetime.timedelta(days=rng.randint(1, 5))
        if (d, r[1], r[2]) not in seen:
            seen.add((d, r[1], r[2]))
            r[0] = d
            rows.append(r)
    rows.sort(key=lambda r: r[0])
    return {'name': '売上明細', 'sheet': '売上明細', 'title': '2026年度 売上明細',
            'columns': [['日付', 'date'], ['担当', 'person'], ['商品', 'text'], ['数量', 'int'], ['単価', 'money'],
                        ['値引率', 'pct'], ['金額', 'money']],
            'rows': rows, 'key': ['日付', '担当', '商品'], 'dups': True, 'font_dirt': False,
            'request': "この売上明細の二重入力を消してください（同じ内容の行は削除してよい）。"
                       "日付・数量・単価・値引率（%）・金額を日付と数値にそろえて、商品名の全角半角と長音、"
                       "担当者名の空白もそろえて。"}


def _ex_members(rng):
    n = rng.randint(30, 40)
    start = rng.randint(1, 800)
    rows = []
    for i in range(n):
        name, kana = _person(rng)
        rows.append([f"M{start + i:04d}", name, kana, _postal(rng), _phone(rng, _MOBILE + _AREAS[:3]),
                     _date(rng, 2015, 4, 2026, 8)])
    # 罠: 同姓同名の別の会員（番号・電話・入会日が違う）＝重複ではない
    k = rng.randrange(n)
    twin = list(rows[k])
    twin[0] = f"M{start + n:04d}"
    twin[3], twin[4], twin[5] = _postal(rng), _phone(rng, _MOBILE), _date(rng, 2015, 4, 2026, 8)
    rows.append(twin)
    return {'name': '会員名簿', 'sheet': '会員名簿', 'title': None,
            'columns': [['会員番号', 'code'], ['氏名', 'person'], ['フリガナ', 'kana'], ['郵便番号', 'postal'],
                        ['電話', 'phone'], ['入会日', 'date']],
            'rows': rows, 'key': ['会員番号'], 'dups': True, 'font_dirt': False,
            'request': "会員名簿を整えてください。同じ会員が二重に入っている行は削除してよい。会員番号・氏名・"
                       "フリガナ（半角カナ・ひらがなは全角カタカナに）・郵便番号・電話・入会日の書き方と、"
                       "前後の空白・セルの中の改行をそろえて。"}


def _ex_expenses(rng):
    n = rng.randint(30, 40)
    staff = [_person(rng)[0] for _ in range(5)]
    rows = [[_date(rng, 2025, 4, 2026, 8), rng.choice(staff), rng.choice(_SUBJECTS), rng.choice(_PAYEES),
             rng.randrange(3, 480) * 100 * (-1 if rng.random() < 0.12 else 1)] for _ in range(n)]   # 1 割は返金（負の数）
    rows.sort(key=lambda r: r[0])
    # 罠: 同じ日・同じ人・同じ金額の本当に 2 回あった支払い（行は消さないで、と頼む）
    k = rng.randrange(n)
    rows.insert(k + 1, list(rows[k]))
    return {'name': '経費精算', 'sheet': '経費精算', 'title': None,
            'columns': [['日付', 'date'], ['氏名', 'person'], ['科目', 'text'], ['支払先', 'text'], ['金額', 'money']],
            'rows': rows, 'key': ['日付', '氏名', '金額'], 'dups': False, 'font_dirt': False, 'wareki': True,
            'blank_rows': True,
            'request': "経費精算の表の書き方をそろえてください。日付は和暦も混ざっているので西暦の日付に、"
                       "金額は数値に（△や（ ）の負の数も）、全角半角・前後の空白・セルの中の改行もそろえて。行は消さないでください。"}


def _ex_inventory(rng):
    n = rng.randint(28, 38)
    codes = rng.sample(range(1001, 1999), n)
    rows = []
    for c in codes:
        p, price = rng.choice(_PRODUCTS)
        rows.append([f"P-{c}", p, rng.choice(_PLACES), rng.randint(0, 300), price, _date(rng, 2025, 10, 2026, 8)])
    return {'name': '在庫表', 'sheet': '在庫表', 'title': '在庫表（2026年9月）',
            'columns': [['品番', 'code'], ['品名', 'text'], ['保管場所', 'text'], ['数量', 'int'], ['単価', 'money'],
                        ['仕入日', 'date']],
            'rows': rows, 'key': ['品番'], 'dups': True, 'font_dirt': True, 'ko': True, 'blank_rows': True,
            'request': "在庫表を整理してください。同じ品番の重複行は削除してよい。品番・品名・保管場所の全角半角と空白、"
                       "品名の長音、数量（「個」は取って数値に）、単価、仕入日をそろえて、文字の書式のばらつきも直して。"}


_EXAMS = [('顧客名簿', _ex_customers), ('売上明細', _ex_sales), ('会員名簿', _ex_members),
          ('経費精算', _ex_expenses), ('在庫表', _ex_inventory)]


# ----------------------------------------------------------------
# 汚す
# ----------------------------------------------------------------

def _clean_cell(typ, val):
    if typ == 'date':
        return {'v': val, 'fmt': 'yyyy/m/d'}
    if typ == 'money':
        return {'v': val, 'fmt': '#,##0'}
    if typ == 'pct':
        return {'v': val, 'fmt': '0%'}
    return {'v': val}


def _apply(typ, val, kind, rng):
    c = {'kind': kind}
    if typ == 'company':
        if kind == 'abbr':
            c['v'] = val.replace('株式会社', rng.choice(['（株）', '(株)', '㈱'])).replace('有限会社', rng.choice(['（有）', '(有)']))
        elif kind == 'inner_space':
            c['v'] = val.replace('株式会社', '株式会社 ' if val.startswith('株式会社') else ' 株式会社').replace(
                '有限会社', '有限会社　')
        elif kind == 'linebreak':
            # セルの中の改行（Alt+Enter）。会社の種類と名前のあいだで折り返した形
            c['v'] = (val.replace('株式会社', '株式会社\n', 1) if val.startswith('株式会社') else val.replace('株式会社', '\n株式会社', 1)
                      ).replace('有限会社', '有限会社\n', 1)
        else:
            c['v'] = _pad(rng, val)
    elif typ == 'person':
        c['v'] = {'zspace': val.replace(' ', '　'), 'linebreak': val.replace(' ', '\n')}.get(kind) or _pad(rng, val)
    elif typ == 'kana':
        c['v'] = {'hankaku': _hankaku_kana(val), 'hiragana': _hiragana(val)}.get(kind) or _pad(rng, val)
    elif typ in ('text', 'code'):
        if kind == 'zenkaku' and _HALF_ALNUM_RE.search(val):
            c['v'] = _zen(val)
        elif kind == 'choon' and re.search(r'[ァ-ヴ]ー', val):
            c['v'] = val.replace('ー', rng.choice(['-', '－', '―', '‐']))      # 長音を記号で打った（マスキングテ-プ）
        else:
            c['kind'], c['v'] = 'space', _pad(rng, val)
    elif typ == 'phone':
        c['v'] = {'nohyphen': val.replace('-', ''),
                  'paren': re.sub(r'^(\d+)-(\d+)-(\d+)$', r'\1(\2)\3', val),
                  'zenkaku': _zen(val),
                  'spaces': val.replace('-', ' ')}[kind]
    elif typ == 'postal':
        c['v'] = {'nohyphen': val.replace('-', ''), 'mark': '〒' + val, 'zenkaku': _zen(val)}[kind]
    elif typ == 'date':
        y, m, d = val.year, val.month, val.day
        if kind == 'fmt':
            c['v'], c['fmt'] = val, rng.choice(['yyyy年m月d日', 'yyyy-mm-dd', 'm/d/yyyy'])
        elif kind == 'serial':
            c['v'] = (val - datetime.date(1899, 12, 30)).days          # 日付が数値のまま（表示形式を標準に戻した）
        else:
            c['v'] = {'text_slash': f"{y}/{m}/{d}", 'text_iso': f"{y}-{m:02d}-{d:02d}",
                      'text_kanji': f"{y}年{m}月{d}日", 'wareki_short': f"R{y - 2018}.{m}.{d}",
                      'wareki_long': f"令和{y - 2018}年{m}月{d}日", 'digits8': f"{y}{m:02d}{d:02d}"}[kind]
    elif typ == 'money':
        a = abs(val)
        c['v'] = {'general': val, 'text_comma': f"{val:,}", 'text_yen': f"¥{val:,}",
                  'text_en': f"{val}円", 'zenkaku': _zen(str(val)),
                  'tri': f"△{a:,}", 'blk': f"▲{a:,}", 'paren': f"({a:,})", 'minus_text': f"-{a:,}"}[kind]
    elif typ == 'pct':
        p = round(val * 100)
        c['v'] = {'text_pct': f"{p}%", 'zen_pct': _zen(str(p)) + '％', 'general': val, 'text_num': str(val)}[kind]
    elif typ == 'int':
        c['v'] = {'text': str(val), 'zenkaku': _zen(str(val)), 'ko': f"{val}個"}[kind]
    else:
        c['v'] = val
    return c


def _dirty_cell(typ, val, rng, ex, p):
    if rng.random() >= p or typ not in _DIRT:
        return _clean_cell(typ, val)
    kinds = list(_DIRT[typ])
    if typ == 'money' and val < 0:
        kinds = list(_DIRT_NEG)
    if typ == 'date' and ex.get('wareki'):
        kinds += ['wareki_short', 'wareki_long', 'wareki_short']
    if typ == 'int' and ex.get('ko'):
        kinds += ['ko', 'ko']
    return _apply(typ, val, rng.choice(kinds), rng)


def _dirty_row(ex, row, rng, p):
    return [_dirty_cell(t, v, rng, ex, p) for (_n, t), v in zip(ex['columns'], row)]


def _soil(ex, rng):
    """正解の行 → 汚れた行（dirty: [{'truth': i, 'cells': [...], 'dup': bool}]）。"""
    dirty = [{'truth': i, 'cells': _dirty_row(ex, r, rng, _P_DIRT), 'dup': False} for i, r in enumerate(ex['rows'])]
    # 電話の市外局番の区切りは、同じ列のきれいな書き方が 1 つはあるときだけ決められる（道具にも人にも）＝局番ごとに 1 つ残す
    for j, (_n, t) in enumerate(ex['columns']):
        if t != 'phone':
            continue
        areas = {}
        for d in dirty:
            areas.setdefault(ex['rows'][d['truth']][j].split('-')[0], []).append(d)
        for lst in areas.values():
            if not any(d['cells'][j].get('kind') is None for d in lst):
                lst[0]['cells'][j] = _clean_cell(t, ex['rows'][lst[0]['truth']][j])
    if ex.get('dups'):
        n = len(ex['rows'])
        for i in rng.sample(range(n), max(3, round(n * 0.15))):
            pos = next(k for k, d in enumerate(dirty) if d['truth'] == i and not d['dup'])
            if rng.random() < 0.4:
                cells = [dict(c) for c in dirty[pos]['cells']]           # そのままの二重入力
            else:
                cells = _dirty_row(ex, ex['rows'][i], rng, 0.5)           # 表記ゆれ込みの重複
            dirty.insert(rng.randint(pos + 1, min(len(dirty), pos + 12)), {'truth': i, 'cells': cells, 'dup': True})
    if ex.get('font_dirt'):
        for d in dirty:
            for c in d['cells']:
                if rng.random() < 0.08:
                    c['font'] = rng.choice(_FONT_DIRT)
    if ex.get('blank_rows'):
        # 表の途中の空行（1 行ずつ・続けて 2 行は入れない＝続く 2 行の空行は「表の終わり」の決まり）
        n = len(dirty)
        for pos in sorted({rng.randint(2, n // 2), rng.randint(n // 2 + 2, n - 2)}, reverse=True):
            dirty.insert(pos, {'truth': None, 'cells': [{'v': None} for _ in ex['columns']], 'dup': False, 'blank': True})
    return dirty


def make_exams(seed, only=None):
    """種 → お題のリスト（正解の行・汚れた行・依頼文）。同じ種なら同じお題。"""
    want = {s.strip() for s in str(only or '').split(',') if s.strip()}
    out = []
    for i, (name, fn) in enumerate(_EXAMS):
        if want and name not in want:
            continue
        rng = random.Random(f"{seed}:{name}")
        ex = fn(rng)
        ex['dirty'] = _soil(ex, rng)
        out.append(ex)
    return out


# ----------------------------------------------------------------
# ブックに書く（openpyxl。文字は文字のまま・日付は日付のまま＝Excel の自動変換を通さない）
# ----------------------------------------------------------------

def _header_row(ex):
    return 3 if ex.get('title') else 1


def write_book(ex, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = ex['sheet']
    r0 = _header_row(ex)
    if ex.get('title'):
        ws.cell(1, 1, ex['title']).font = Font(name=_FONT, size=14, bold=True)
    for j, (name, _t) in enumerate(ex['columns'], 1):
        c = ws.cell(r0, j, name)
        c.font = Font(name=_FONT, size=11, bold=True)
        c.fill = PatternFill('solid', fgColor='DDEBF7')
    for i, d in enumerate(ex['dirty'], r0 + 1):
        for j, cell in enumerate(d['cells'], 1):
            c = ws.cell(i, j, cell['v'])
            if cell.get('fmt'):
                c.number_format = cell['fmt']
            f = cell.get('font') or {}
            c.font = Font(name=f.get('name', _FONT), size=f.get('size', 11), bold=f.get('bold', False),
                          italic=f.get('italic', False), color=f.get('color'))
    for j, (name, t) in enumerate(ex['columns'], 1):
        ws.column_dimensions[get_column_letter(j)].width = {'company': 26, 'text': 22, 'phone': 16, 'kana': 18,
                                                            'date': 13, 'money': 12}.get(t, 12)
    wb.save(path)


def _truth_json(ex):
    def enc(v):
        return f"D:{v.isoformat()}" if isinstance(v, datetime.date) else v
    return {'name': ex['name'], 'sheet': ex['sheet'], 'title': ex.get('title'), 'columns': ex['columns'],
            'key': ex['key'], 'font_dirt': bool(ex.get('font_dirt')), 'request': ex['request'],
            'truth': [[enc(v) for v in r] for r in ex['rows']],
            'dup_rows': sum(1 for d in ex['dirty'] if d['dup'])}


def _dirt_summary(ex):
    c = Counter()
    for d in ex['dirty']:
        for (name, _t), cell in zip(ex['columns'], d['cells']):
            if cell.get('kind'):
                c[name] += 1
            if cell.get('font'):
                c['文字の書式'] += 1
    return c


# ----------------------------------------------------------------
# 採点（正解の表と、終わった表のセルの値を突き合わせる。AI も道具の検査も使わない）
#   cells: [[{'v': 値, 't': 画面の文字, 'a': 番地}, …], …]（使用範囲の左上から）
#   値は 文字／数（float）／'D:YYYY-MM-DD'（日付）／None
# ----------------------------------------------------------------

def _empty(c):
    v = (c or {}).get('v')
    return v is None or (isinstance(v, str) and not v.strip())


def find_header(cells, names):
    """見出しの行（0 起点）と {見出し名: 列（0 起点）} → 見つからなければ None。"""
    for i, row in enumerate(cells[:15]):
        texts = [_ws(_nfkc(c.get('v'))) if isinstance(c.get('v'), str) else '' for c in row]
        if all(n in texts for n in names):
            return i, {n: texts.index(n) for n in names}
    return None


def _body_scan(cells, hdr, names):
    """見出しの下の本文 → [(cells の行 idx, 見出し順の値)]。途中の空行は飛ばし、空行が 2 行続いたら表の終わり。"""
    hi, cmap = hdr
    out, blanks = [], 0
    for i in range(hi + 1, len(cells)):
        row = cells[i]
        vals = [row[cmap[n]] if cmap[n] < len(row) else {'v': None, 't': ''} for n in names]
        if all(_empty(c) for c in vals):
            blanks += 1
            if blanks >= 2:
                break
            continue
        blanks = 0
        out.append((i, vals))
    return out


def body_rows(cells, hdr, names):
    return [vals for _i, vals in _body_scan(cells, hdr, names)]


def body_span(cells, hdr, names):
    """本文の最初と最後の行（cells の行 idx）。本文が無ければ None。"""
    scan = _body_scan(cells, hdr, names)
    return (scan[0][0], scan[-1][0]) if scan else None


def ident(typ, cell):
    """行を見分けるための緩い読み（書き方の違いは無視して値だけ）。"""
    v = (cell or {}).get('v')
    if v is None or (isinstance(v, str) and not v.strip()):
        return ''
    if typ in ('money', 'int'):
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return ('n', round(float(v), 2))
        s = re.sub(r'[¥,円個\s]', '', _nfkc(v))
        s = re.sub(r'^[△▲]', '-', s)
        s = re.sub(r'^\((.*)\)$', r'-\1', s)
        try:
            return ('n', round(float(s), 2))
        except ValueError:
            return ('s', _ws(_nfkc(v)))
    if typ == 'pct':
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return ('n', round(float(v), 4))
        s = _nfkc(v).strip()
        try:
            return ('n', round(float(s[:-1]) / 100, 4)) if s.endswith('%') else ('n', round(float(s), 4))
        except ValueError:
            return ('s', s)
    if typ == 'date':
        if isinstance(v, str) and v.startswith('D:'):
            return ('d', v[2:])
        if isinstance(v, (int, float)) and not isinstance(v, bool) and 20000 < v < 80000 and float(v).is_integer():
            return ('d', (datetime.date(1899, 12, 30) + datetime.timedelta(days=int(v))).isoformat())   # 数値のままの日付
        d = _parse_date(str(v))
        return ('d', d) if d else ('s', str(v))
    if typ in ('phone', 'postal'):
        if isinstance(v, (int, float)):
            return ('num', str(v))
        return ('p', re.sub(r'\D', '', _nfkc(v)))
    if typ == 'company':
        return ('s', _choon(_company_canon(v)))
    s = _choon(_ws(_nfkc(v)))
    if typ == 'kana':
        s = ''.join(chr(ord(ch) + 0x60) if 0x3041 <= ord(ch) <= 0x3096 else ch for ch in s)   # ひらがな → カタカナ
    return ('s', s)


_CHOON_RE = re.compile(r'(?<=[ァ-ヴ])[-－―‐−](?![0-9０-９])')   # コピ-用紙・テ-プ・トナ‐TN（後ろが数字なら区切り＝ルーム-2）


def _choon(s):
    """カタカナのあいだの記号（- － ― ‐ −）は長音「ー」と読む（行を見分けるためだけ・書き方の誤りは strict が咎める）。"""
    return _CHOON_RE.sub('ー', s)


def _show(cell):
    v = (cell or {}).get('v')
    if isinstance(v, str) and v.startswith('D:'):
        return (cell.get('t') or v[2:])
    return str(cell.get('t') if cell.get('t') not in (None, '') else v)


def strict(typ, cell, truth):
    """値は合っている前提で、書き方の誤り → 文（無ければ None）。"""
    v = (cell or {}).get('v')
    if typ in _STR_TYPES:
        if not isinstance(v, str):
            return "文字でなくなっている"
        if v != v.strip():
            return "前後の空白が残っている"
        if '\n' in v or '\r' in v:
            return "セルの中の改行が残っている"
        if typ in ('code', 'text') and _FULL_ALNUM_RE.search(v):
            return "英数字が全角のまま"
        if _HALF_KANA_RE.search(v):
            return "半角カナが残っている"
        if typ == 'kana' and re.search(r'[ぁ-ゖ]', v):
            return "ひらがなが残っている"
        if _CHOON_RE.search(v):
            return "長音が記号のまま（ー でない）"
        if typ == 'company' and re.search(r'(株式会社|有限会社)\s|\s(株式会社|有限会社)', _nfkc(v)) \
                and not re.search(r'\s', str(truth)):
            return "会社の種類の前後に空白が残っている"
        return None
    if typ == 'phone':
        if not isinstance(v, str):
            return "数になって先頭の 0 が消えた"
        if v != _nfkc(v):
            return "全角のまま"
        if re.fullmatch(r'\d+', v) or v == truth:
            return None
        return f"書き方が違う（正しくは {truth}）"
    if typ == 'postal':
        if not isinstance(v, str):
            return "数になって先頭の 0 が消えた"
        if re.fullmatch(r'\d{3}-\d{4}|\d{7}', v):
            return None
        return "書き方が違う（〒・全角・区切り）"
    if typ == 'date':
        return None if isinstance(v, str) and v.startswith('D:') else "日付になっていない（文字のまま）"
    if typ in ('money', 'int', 'pct'):
        return None if isinstance(v, (int, float)) and not isinstance(v, bool) else "数値になっていない（文字のまま）"
    return None


def _shape(text):
    return re.sub(r'\d+', 'N', str(text or ''))


def uneven(typ, cells):
    """列の中のばらつき → 文のリスト。"""
    vals = [c for c in cells if not _empty(c)]
    out = []
    if typ in _STR_TYPES:
        s = [c['v'] for c in vals if isinstance(c.get('v'), str)]
        if typ == 'company':
            ab = sum(1 for x in s if _has_abbr(x))
            if ab and ab < len(s):
                out.append(f"（株）と株式会社が混ざっている（略 {ab}・正式 {len(s) - ab}）")
        inner = [x.strip() for x in s]
        zs = sum(1 for x in inner if '　' in x)
        hs = sum(1 for x in inner if ' ' in x)
        if zs and hs:
            out.append(f"空白の全角半角が混ざっている（全角 {zs}・半角 {hs}）")
        fa = sum(1 for x in s if _FULL_ALNUM_RE.search(x))
        ha = sum(1 for x in s if _HALF_ALNUM_RE.search(x))
        if fa and ha:
            out.append(f"英数字の全角半角が混ざっている（全角 {fa}・半角 {ha}）")
    elif typ in ('phone', 'postal'):
        st = Counter('区切りあり' if '-' in str(c.get('v')) else '区切りなし' for c in vals)
        if len(st) > 1:
            out.append("区切りのある・ないが混ざっている（" + "・".join(f"{k} {v}" for k, v in st.items()) + "）")
    elif typ == 'date':
        sh = Counter(_shape(c.get('t') or c.get('v')) for c in vals)
        if len(sh) > 1:
            out.append("表示が混ざっている（" + "・".join(f"{k} {v}" for k, v in sh.most_common(4)) + "）")
    elif typ == 'money':
        big = [c for c in vals if isinstance(c.get('v'), (int, float)) and abs(c['v']) >= 1000]
        comma = Counter(',' in str(c.get('t') or '') for c in big)
        if len(comma) > 1:
            out.append(f"桁区切りのある・ないが混ざっている（ある {comma[True]}・ない {comma[False]}）")
        yen = Counter('¥' in str(c.get('t') or '') or '\\' in str(c.get('t') or '') for c in vals)
        if len(yen) > 1:
            out.append("¥ のある・ないが混ざっている")
    elif typ == 'pct':
        pc = Counter('%' in _nfkc(c.get('t') or '') for c in vals)
        if len(pc) > 1:
            out.append(f"% の表示のある・ないが混ざっている（ある {pc[True]}・ない {pc[False]}）")
    return out


def grade(exj, cells, fonts=None):
    """正解（_truth_json の形）と終わった表 → 結果の dict。pass は全部そろったときだけ。"""
    cols = exj['columns']
    names = [n for n, _t in cols]
    res = {'truth_rows': len(exj['truth']), 'got_rows': 0, 'rows_ok': 0, 'missing': 0, 'dup_left': 0,
           'extra': 0, 'value_errors': 0, 'format_errors': 0, 'uneven': [], 'problems': [], 'pass': False}
    hdr = find_header(cells, names)
    if hdr is None:
        res['problems'].append("見出し（" + "・".join(names) + "）の行が見つからない")
        return res
    body = body_rows(cells, hdr, names)
    res['got_rows'] = len(body)
    truth = [[{'v': v} for v in r] for r in exj['truth']]
    tid = [tuple(ident(t, c) for (_n, t), c in zip(cols, r)) for r in truth]
    gid = [tuple(ident(t, c) for (_n, t), c in zip(cols, r)) for r in body]
    pool = {}
    for j, k in enumerate(tid):
        pool.setdefault(k, []).append(j)
    pairs, un_g = [], []
    for i, k in enumerate(gid):
        if pool.get(k):
            pairs.append((i, pool[k].pop(0)))
        else:
            un_g.append(i)
    un_t = [j for lst in pool.values() for j in lst]
    kidx = [names.index(k) for k in exj['key']]
    val_notes, fmt_notes = [], Counter()
    fmt_example = {}
    wrong_pairs = []
    for i in list(un_g):
        kg = tuple(gid[i][c] for c in kidx)
        hit = next((j for j in un_t if tuple(tid[j][c] for c in kidx) == kg), None)
        if hit is None:
            continue
        un_g.remove(i)
        un_t.remove(hit)
        wrong_pairs.append((i, hit))
        for c, (n, _t) in enumerate(cols):
            if gid[i][c] != tid[hit][c]:
                res['value_errors'] += 1
                if len(val_notes) < 6:
                    a = body[i][c].get('a') or ''
                    val_notes.append(f"{n} {a}「{_show(body[i][c])}」→ 正しくは「{_show(truth[hit][c])}」")
    tset = set(tid)
    for i in un_g:
        if gid[i] in tset:
            res['dup_left'] += 1
        else:
            res['extra'] += 1
    res['missing'] = len(un_t)
    for i, j in pairs + wrong_pairs:
        row_ok = (i, j) in pairs
        for c, (n, t) in enumerate(cols):
            msg = strict(t, body[i][c], exj['truth'][j][c])
            if msg:
                res['format_errors'] += 1
                fmt_notes[(n, msg)] += 1
                fmt_example.setdefault((n, msg), f"{body[i][c].get('a') or ''}「{_show(body[i][c])}」")
                row_ok = False
        if row_ok:
            res['rows_ok'] += 1
    for c, (n, t) in enumerate(cols):
        for u in uneven(t, [r[c] for r in body]):
            res['uneven'].append(f"{n}: {u}")
    if exj.get('font_dirt') and fonts:
        for n in names:
            f = fonts.get(n) or {}
            mixed = [lab for key, lab in (('bold', '太字'), ('italic', '斜体'), ('size', '大きさ'), ('color', '色'),
                                          ('name', '書体')) if key in f and f[key] is None]
            if mixed:
                res['uneven'].append(f"{n}: 文字の書式が混ざっている（{'・'.join(mixed)}）")
    p = res['problems']
    if res['missing']:
        miss = [" / ".join(_show(truth[j][c]) for c in kidx) for j in un_t[:3]]
        p.append(f"足りない行 {res['missing']}（消しすぎ。例: {'、'.join(miss)}）")
    if res['dup_left']:
        p.append(f"残った重複 {res['dup_left']} 行")
    if res['extra']:
        ex_rows = [" / ".join(_show(body[i][c]) for c in kidx) for i in un_g if gid[i] not in tset][:3]
        p.append(f"余計な行 {res['extra']}（例: {'、'.join(ex_rows)}）")
    if res['value_errors']:
        p.append(f"値の誤り {res['value_errors']} セル: " + "／".join(val_notes))
    for (n, msg), k in fmt_notes.most_common(6):
        p.append(f"{n}: {msg} {k} セル（例 {fmt_example[(n, msg)]}）")
    p.extend(res['uneven'])
    res['pass'] = not (res['missing'] or res['dup_left'] or res['extra'] or res['value_errors']
                       or res['format_errors'] or res['uneven'])
    return res


# ----------------------------------------------------------------
# 撃つ（別プロセス。状態ファイルを逃がし、本番と同じ run_agent に撃つ）
# ----------------------------------------------------------------

def _probe_source():
    """撃つ子プロセスの台本。--by-macro のときだけ、使い捨ての新しいブックへ .bas を取り込む（注入経路。
    取り込む前に exam() が check-bas の検査＝文字コード・識別子を通す。人のブックには書かない）。"""
    return _PROBE


_PROBE = r'''# -*- coding: utf-8 -*-
import io, json, os, sys, time, shutil, contextlib, traceback
sys.path.insert(0, r"__SCRIPTS__")
import vbam_core
vbam_core.setup_encoding()
import vba_manager, vbam_agent as va, vbam_shake as vs, vbam_exam as ve   # _run_cmd が後から読む表も先に読む
from vbam_hands import _rows_of

plan_path = sys.argv[1]
with open(plan_path, encoding="utf-8") as f:
    plan = json.load(f)
vs._divert_state(plan["state"])
# 自分の台（見える窓）を gen_py キャッシュを通らずに起こす（DispatchEx はキャッシュが壊れていると
# CLSIDToClassMap で落ちる・2026-09-11 夜。人の Excel につなぐ道も late-binding＝同じ考え方）
import pythoncom, win32com.client.dynamic
pythoncom.CoInitialize()
xl = win32com.client.dynamic.Dispatch(pythoncom.CoCreateInstanceEx(
    "Excel.Application", None, pythoncom.CLSCTX_SERVER, None, (pythoncom.IID_IDispatch,))[0])
try:
    import win32process
    _pid = win32process.GetWindowThreadProcessId(xl.Hwnd)[1]
except Exception:
    _pid = None
vbam_core._created_instances.append({"xl": xl, "pid": _pid})   # 終わったら release_created_instances が畳む
xl.Visible = True
xl.DisplayAlerts = False
mwb = None
if plan.get("macro_bas"):
    # マクロで撃つ（AI を使わない）: 空のブックに .bas を読み込み、お題のシートをアクティブにして Run する
    mwb = xl.Workbooks.Add()
    mwb.VBProject.VBComponents.Import(plan["macro_bas"])


def enc(v):
    if v is None:
        return None
    if hasattr(v, "year") and hasattr(v, "month") and hasattr(v, "day"):
        return "D:%04d-%02d-%02d" % (v.year, v.month, v.day)
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return float(v)
    return str(v)


for item in plan["items"]:
    row = {"name": item["name"]}
    print("__START__" + item["name"], flush=True)
    buf = io.StringIO()
    t0 = time.time()
    wb = None
    try:
        wb = xl.Workbooks.Open(item["book"])
        vbam_core._wb_cache.clear()
        ws = wb.Worksheets(item["sheet"])
        ws.Activate()
        r, err = None, None
        if mwb is not None and plan.get("macro_only"):
            try:
                t1 = time.time()
                xl.Run("'%s'!表を整える" % mwb.Name)
                if item.get("dedupe"):
                    xl.Run("'%s'!重複行を消す" % mwb.Name)
                row["macro_sec"] = round(time.time() - t1, 2)
                buf.write("マクロ: 表を整える%s（%.2f 秒）\n" % ("・重複行を消す" if item.get("dedupe") else "", row["macro_sec"]))
            except Exception as ex:
                err = "%s: %s" % (type(ex).__name__, ex)
                buf.write(traceback.format_exc())
        else:
          with vbam_core.pinned_workbook(wb), contextlib.redirect_stdout(buf):
            try:
                kw = {}
                if plan.get("grade_always"):
                    kw["grade_always"] = True
                r = va.run_agent(item["request"], item["sheet"], wb, plan["ai"], plan.get("model"), **kw)
            except Exception as ex:
                err = "%s: %s" % (type(ex).__name__, ex)
                buf.write(traceback.format_exc())
        row["err"] = err
        row["sec"] = round(time.time() - t0, 1)
        if r is not None:
            row.update({"ok": bool(r.get("ok")), "done": bool(r.get("done")), "turns": r.get("turns"),
                        "gates": {k: v for k, v in (r.get("gates") or {}).items() if v},
                        "ai_sec": float((r.get("usage") or {}).get("sec") or 0)})
        ws = wb.Worksheets(item["sheet"])
        ur = ws.UsedRange
        r0, c0 = int(ur.Row), int(ur.Column)
        vals = _rows_of(ur.Value)
        cells = []
        for i, rr in enumerate(vals):
            line = []
            for j, v in enumerate(rr):
                line.append({"v": enc(v), "t": str(ws.Cells(r0 + i, c0 + j).Text),
                             "a": "%s%d" % (vbam_core._col_letter(c0 + j), r0 + i)})
            cells.append(line)
        row["cells"] = cells
        names = [n for n, _t in item["columns"]]
        hdr = ve.find_header(cells, names)
        fonts = {}
        if hdr is not None:
            span = ve.body_span(cells, hdr, names)
            if span:
                hi, cmap = hdr
                for n in names:
                    rg = ws.Range(ws.Cells(r0 + span[0], c0 + cmap[n]), ws.Cells(r0 + span[1], c0 + cmap[n]))
                    f = rg.Font
                    fonts[n] = {"bold": f.Bold, "italic": f.Italic, "size": f.Size, "color": f.Color, "name": f.Name}
        row["fonts"] = fonts
        try:
            wb.SaveCopyAs(item["after"])
        except Exception as ex:
            row["after_err"] = str(ex)
        try:
            if os.path.isfile(va._LAST_AGENT_LOG_FILE):
                shutil.copy2(va._LAST_AGENT_LOG_FILE, item["log"])
        except Exception:
            pass
    except Exception as ex:
        row["err"] = row.get("err") or "%s: %s" % (type(ex).__name__, ex)
        row["trace"] = traceback.format_exc()[-1500:]
    finally:
        ws = ur = None
        if wb is not None:
            try:
                wb.Close(SaveChanges=False)
            except Exception:
                pass
        wb = None
    with open(item["out"], "w", encoding="utf-8") as f:
        f.write(buf.getvalue())
    row.setdefault("sec", round(time.time() - t0, 1))
    print("__ROW__" + json.dumps(row, ensure_ascii=False, default=str), flush=True)

if mwb is not None:
    try:
        mwb.Close(SaveChanges=False)
    except Exception:
        pass
    mwb = None
try:
    vbam_core.release_created_instances()
except Exception:
    pass
xl = None
print("__END__", flush=True)
'''


def _w(s):
    return sum(2 if unicodedata.east_asian_width(ch) in ('W', 'F') else 1 for ch in str(s))


def _padw(s, n):
    s = str(s)
    return s + ' ' * max(0, n - _w(s))


def _verdict(row, g):
    if row.get('err'):
        return '落ちた'
    return '合格' if g['pass'] else '不合格'


def _line(row, g):
    tool = row.get('ok')
    lie = (tool and not g['pass'])
    return (f"  {_padw(row['name'], 10)} {_padw(_verdict(row, g), 6)} 正しい行 {g['rows_ok']:>2}/{g['truth_rows']:<2} "
            f"{row.get('sec', 0):6.1f} 秒 往復 {row.get('turns') if row.get('turns') is not None else '-'}  道具の判定 "
            + ('合格' if tool else ('不合格' if tool is False else '-'))
            + ('  ← 嘘の合格' if lie else ''))


def exam(seed=None, only=None, dry_run=False, ai=None, model=None, grade_always=False, timeout=3600, macro_bas=None,
         with_macro=None):
    """agent --exam の本体。戻り値は「全部のお題が合格したか」。
    macro_bas＝AI の代わりにマクロで撃つ（.bas の「表を整える」、依頼が消すことを承認していれば続けて「重複行を消す」）。
    with_macro＝その .bas を読み込んだ Excel で agent を撃つ（入口の先撃ちが効くか＝本番の秀コンボと同じ姿）。"""
    import vbam_agent as va
    macro_only = bool(macro_bas)
    macro_bas = macro_bas or with_macro
    if macro_bas:
        macro_bas = os.path.abspath(macro_bas)
        if not os.path.isfile(macro_bas):
            print(f"エラー: マクロの .bas がありません: {macro_bas}")
            return False
        from vbam_vba import _check_bas_one      # 取り込む前の検査（文字コード・識別子）＝注入経路のガード
        if not _check_bas_one(macro_bas):
            print("エラー: .bas の検査が通らないので撃ちません")
            return False
    seed = int(seed) if seed not in (None, '') else random.randrange(1, 100000)
    exams = make_exams(seed, only)
    if not exams:
        print("エラー: お題がありません（--only の名前: " + "・".join(n for n, _f in _EXAMS) + "）")
        return False
    stamp = time.strftime('%Y%m%d_%H%M%S')
    vdir = os.path.join(_EXAM_DIR, stamp)
    os.makedirs(vdir, exist_ok=True)
    items, truths = [], {}
    print(f"試験: 種 {seed}（同じお題で撃ち直すなら agent --exam --seed {seed}）　置き場 {vdir}")
    for ex in exams:
        book = os.path.join(vdir, f"{ex['name']}.xlsx")
        write_book(ex, book)
        tj = _truth_json(ex)
        truths[ex['name']] = tj
        dirt = _dirt_summary(ex)
        print(f"\n■ {ex['name']}: 正解 {len(ex['rows'])} 行・お題 {len(ex['dirty'])} 行（重複 {tj['dup_rows']} 行を混ぜた）"
              f"・見出し {_header_row(ex)} 行目")
        print("  汚れ: " + "・".join(f"{k} {v}" for k, v in dirt.items()))
        print(f"  依頼: {ex['request']}")
        items.append({'name': ex['name'], 'book': book, 'sheet': ex['sheet'], 'request': ex['request'],
                      'columns': ex['columns'],
                      'dedupe': bool(re.search(r'削除してよい|消してよい', ex['request'])
                                     and not re.search(r'消さないで', ex['request'])),
                      'after': os.path.join(vdir, f"{ex['name']}_after.xlsx"),
                      'log': os.path.join(vdir, f"{ex['name']}.log.jsonl"),
                      'out': os.path.join(vdir, f"{ex['name']}.out.txt")})
    with open(os.path.join(vdir, 'truth.json'), 'w', encoding='utf-8') as f:
        json.dump(truths, f, ensure_ascii=False, indent=1)
    if dry_run:
        print("\n（--dry-run: お題のブックと正解を作っただけ。Excel にも AI にも触っていません）")
        return True
    ai = (ai or va._CC_AI)
    state = os.path.join(vdir, '_state')
    os.makedirs(state, exist_ok=True)
    plan = {'state': state, 'ai': ai, 'model': model, 'grade_always': bool(grade_always), 'items': items,
            'macro_bas': macro_bas, 'macro_only': macro_only}
    plan_path = os.path.join(vdir, '_plan.json')
    with open(plan_path, 'w', encoding='utf-8') as f:
        json.dump(plan, f, ensure_ascii=False, indent=1)
    probe = os.path.join(vdir, '_probe.py')
    with open(probe, 'w', encoding='utf-8') as f:
        f.write(_probe_source().replace('__SCRIPTS__', SCRIPT_DIR))
    by = (f"マクロ {os.path.basename(macro_bas)}（AI なし）" if macro_only else
          f"頭 {ai}{(' ' + model) if model else ''}（本番と同じ run_agent"
          + (f"・{os.path.basename(macro_bas)} を読み込んだ Excel＝マクロの先撃ちあり" if macro_bas else "") + "）")
    print(f"\n撃ちます: {len(items)} 題・{by}。Excel は自分の台を起こす＝人の Excel には触らない。台帳・控え・覚書は {state} へ")
    env = dict(os.environ, PYTHONIOENCODING='utf-8')
    proc = subprocess.Popen([sys.executable, probe, plan_path], stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                            text=True, encoding='utf-8', errors='replace', env=env)
    rows, t0 = [], time.time()
    try:
        for line in proc.stdout:
            line = line.rstrip('\n')
            if line.startswith('__START__'):
                print(f"  {line[9:]} を撃っています…", flush=True)
            elif line.startswith('__ROW__'):
                row = json.loads(line[7:])
                g = grade(truths[row['name']], row.get('cells') or [], row.get('fonts'))
                row['grade'] = g
                rows.append(row)
                print(_line(row, g), flush=True)
                if row.get('err'):
                    print(f"      エラー: {row['err']}")
                for p in g['problems'][:8]:
                    print(f"      - {p}")
            elif line.startswith('__END__'):
                pass
            elif line.strip():
                print("    | " + line, flush=True)
            if time.time() - t0 > timeout:
                proc.kill()
                print(f"エラー: {timeout} 秒を超えたので止めました")
                break
    finally:
        try:
            proc.wait(timeout=60)
        except Exception:
            proc.kill()
    if not rows:
        print("エラー: 1 題も返ってきませんでした（出力は上の行）")
        return False
    ok = [r for r in rows if r['grade']['pass'] and not r.get('err')]
    lies = [r for r in rows if r.get('ok') and not r['grade']['pass']]
    right = sum(r['grade']['rows_ok'] for r in rows)
    total = sum(r['grade']['truth_rows'] for r in rows)
    secs = [r.get('sec', 0) for r in rows]
    print(f"\nまとめ（種 {seed}）: 合格 {len(ok)}/{len(rows)} 題　正しい行 {right}/{total}（{right * 100 // max(1, total)}%）"
          f"　嘘の合格 {len(lies)}　計 {sum(secs):.0f} 秒（1 題 平均 {sum(secs) / len(secs):.0f} 秒）")
    print(f"終わった表: {vdir}\\*_after.xlsx　撃ち直し: agent --replay {vdir}\\お題.log.jsonl（お題.xlsx を開いて）")
    report = {'seed': seed, 'stamp': stamp, 'ai': ai, 'model': model,
              'rows': [{k: v for k, v in r.items() if k != 'cells'} for r in rows]}
    with open(os.path.join(vdir, 'report.json'), 'w', encoding='utf-8') as f:
        json.dump(report, f, ensure_ascii=False, indent=1, default=str)
    rec = {'time': time.strftime('%Y-%m-%d %H:%M:%S'), 'seed': seed, 'ai': ai, 'model': model, 'dir': vdir,
           'by': (('macro:' if macro_only else 'agent+macro:') + os.path.basename(macro_bas)) if macro_bas else 'agent',
           'pass': len(ok), 'of': len(rows), 'rows_ok': right, 'rows': total, 'lies': len(lies),
           'sec': round(sum(secs), 1),
           'items': [{'name': r['name'], 'pass': r['grade']['pass'], 'rows_ok': r['grade']['rows_ok'],
                      'rows': r['grade']['truth_rows'], 'sec': r.get('sec'), 'turns': r.get('turns'),
                      'tool_ok': r.get('ok'), 'err': r.get('err'), 'why': r['grade']['problems'][:4]} for r in rows]}
    with open(_EXAM_SCORE_FILE, 'a', encoding='utf-8') as f:
        f.write(json.dumps(rec, ensure_ascii=False, default=str) + "\n")
    return len(ok) == len(rows)


__all__ = ['_EXAMS', '_EXAM_DIR', '_EXAM_SCORE_FILE', 'body_rows', 'exam', 'find_header', 'grade', 'ident',
           'make_exams', 'strict', 'uneven', 'write_book']
