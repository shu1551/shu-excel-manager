"""COM 不要の純粋ロジックの自動テスト（AI の機能: agent・試験・揺らし・手順書・キー・控えと戻す手）。

実行: このフォルダで `py -m pytest test_agent.py -v`
Excel には一切接続しない（COM オブジェクトはダミーで代用）。本物の AI も呼ばない（conftest の見張り）。
2026-09-12 に test_tools.py から分けた（中身は 1 行も変えていない）。本体のテストは test_tools.py。
"""
import os
import sys
import json
import inspect

import pytest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import vba_manager as vm
import form_layout as fl
import form_inspect as fi
import vbam_fire as vf                # 実射（練習台・弾・状態・写し取り・点数）は 2026-09-11 に別ファイルへ
import vbam_view_ai as vva               # vv は vbam_vba（1844 行目）で使っている
import vbam_grade as vg
import vbam_undo as vu
import vbam_ledger as vl
import vbam_clean as vc
import vbam_macro as vmac
import vbam_inv as vi
import vbam_hands as vh
import vbam_ai as vai
import vbam_keys as vk


# ================================================================
# open / close / rehearse（2026-07-15 追加）: COM に触る前の門番
# ================================================================

import argparse as _ap


# ================================================================
# vba_mcp_server: AI作業窓の関所（実行ログ・承認ゲート・差分確認）
# ================================================================
# 2026-08-26 追加。この層は「AI に任せる範囲が広がったこと」への備えなので、
# 壊れても普通の使い方では気づけない（AI_WIN_LOG が無いと全部素通しするため）。
# 素通しの向き——窓に聞けない異常時はフェイルオープン（方針A）、返事が無いときは
# フェイルクローズ（実行しない）——まで含めて機械で押さえる。
import re as _re                                                    # noqa: E402
import tempfile as _tf                                              # noqa: E402
import threading as _th                                             # noqa: E402
import time as _time                                                # noqa: E402

import pytest as _pytest                                            # noqa: E402

_cwd_before_srv_import = os.getcwd()
try:
    import vba_mcp_server as _srv
except Exception:                    # mcp SDK 未導入の環境ではこの節ごと飛ばす
    _srv = None
finally:
    # vba_mcp_server は import 時に os.chdir(SCRIPT_DIR) する。他のテストへ
    # 作業ディレクトリの変更を持ち込まないよう、ここで必ず戻す。
    os.chdir(_cwd_before_srv_import)


# ================================================================
# xlflow から輸入した診断4本と新コマンド5本（2026-08-28）
#   VBM008 Application の状態を戻さない / VBM009 呼び先が状態を変えたまま返る
#   VBM010 モジュール変数の書き手が散る / VBM011 Find/Replace の引数省略
# ================================================================

import vbam_vba as vv


def test_look_mixed_fonts_seen_by_materials_and_gate():
    """文字の色・大きさ・太字がばらばらの表を、材料も合格判定も見ていなかった（2026-09-11）。"""
    import types
    import vbam_view as vview

    class _Cols:
        Count = 2

        def __call__(self, k):
            return types.SimpleNamespace(HorizontalAlignment=1, Column=k)
    font = types.SimpleNamespace(Name=None, Size=11, Color=None, Bold=None, Italic=False,
                                 Underline=-4142, Strikethrough=False)
    rng = types.SimpleNamespace(Font=font, Interior=types.SimpleNamespace(ColorIndex=None),
                                Columns=_Cols(), Address="$A$2:$B$3")
    # 合格判定: 「書式のばらつき」だけで文字の書式も全部見る
    msg = " ".join(vg._look_misfits(rng, "書式のばらつきをそろえてください"))
    assert "フォント名" in msg and "文字色" in msg and "太字" in msg and "塗りつぶし" in msg
    assert vg._look_misfits(rng, "重複を消して") == []
    # 材料: 本文の混在を気づきに出す
    ws = types.SimpleNamespace(Cells=lambda r, c: (r, c), Range=lambda a, b: rng)
    grid = [["会社", "金額"], ["a", 1], ["b", 2]]
    addr, bad = vview._body_look_mixed(ws, grid, 1, 1, 0)
    assert addr == "A2:B3" and bad == ["フォント名", "文字色", "太字", "塗り"]
    text = "\n".join(vm.dirt_notes(grid, look_mixed=(addr, bad)))
    assert "文字の書式がセルごとに違う（本文 A2:B3）: フォント名・文字色・太字・塗り" in text
    assert "_body_look_mixed(" in inspect.getsource(vm.cmd_materials)


# ---- tidy の列の型（2026-09-02 深夜・「会員番号は左寄せ・金額はカンマが当然。道具でそうなるようにしろ」）----
import vbam_edit as ve                                              # noqa: E402


def test_build_sheet_ask_is_wired_and_prompt_is_one_shot(tmp_path, monkeypatch):
    # 手2: --ask / --ai / --model / --new-book が argparse に通る。プロンプトは規則＋例＋依頼の 1 通
    import vbam_build as vb
    import vbam_agent as va
    monkeypatch.setattr(vk, "_KEY_STORE", str(tmp_path / "keys.json"))   # 金庫は tmp に向ける＝本物のキーを見ない
    monkeypatch.setattr(vb, "_LAST_SHEET_ASK_FILE", str(tmp_path / "ask.txt"))   # 聞いた文も tmp へ（作業フォルダに残さない）
    p = vm.build_parser()
    ns, unknown = p.parse_known_args(["build-sheet", "--ask", "名簿を作って", "--ai", "gemini",
                                      "--model", "gemini-3.7-flash", "--new-book", "--dry-run"])
    assert ns.ask == "名簿を作って" and ns.ai == "gemini" and ns.new_book and ns.dry_run and not unknown
    prompt = vb._spec_prompt("会員名簿。氏名と会費。", "既存のシート: 目次, 台帳")
    assert prompt.startswith(vb.SPEC_RULES)
    assert '"sheet": "会員名簿"' in prompt            # 設計図の例が JSON で入っている
    assert "既存のシート: 目次, 台帳" in prompt and prompt.rstrip().endswith("設計図（JSON）:")
    assert prompt.index("依頼:") > prompt.index("手元のブックの様子:")
    # 返事の取り出し: 素の JSON / コードフェンス付き / 前後に文が混ざる
    assert vb._extract_json('{"sheet":"a","columns":["x"]}')["sheet"] == "a"
    assert vb._extract_json('```json\n{"sheet":"b","columns":["x"]}\n```')["sheet"] == "b"
    assert vb._extract_json('設計図です。\n{"sheet":"c","columns":["x"]}\n以上')["sheet"] == "c"
    # 末尾に } が 1 つ余分・後ろに別の JSON や文が続く（sonnet・effort low で実測）→ 最初の JSON を取る（2026-09-08）
    assert vb._extract_json('{"sheet":"d","columns":["x"]}}')["sheet"] == "d"
    assert vb._extract_json('{"sheet":"e","columns":["x"]}\n{"done":true}')["sheet"] == "e"
    # キーが無いときは止まる（ネットには出ない）。読み方は agent と同じ 1 か所＝環境変数 → 金庫
    # （2026-09-04: build だけ環境変数しか見ておらず、キーを金庫に移した後に止まっていた）
    import os, pytest
    saved = os.environ.pop("GEMINI_API_KEY", None)
    try:
        with pytest.raises(RuntimeError, match="APIキーの設定"):
            vb.ask_design("x", None, "gemini", None)
        # 金庫にキーがあれば環境変数が無くても通る（API 呼び出しの手前で止めて確かめる）
        monkeypatch.setattr(vai, "_key_load", lambda ai: "sk-test-from-store")   # _ai_setup は vbam_ai（2026-09-11）
        seen = {}

        def fake_call(prompt, model, api_key):
            seen['key'], seen['model'] = api_key, model
            raise KeyboardInterrupt        # ここまで来れば十分＝ネットには出ない

        monkeypatch.setattr(vb, "_ai_gemini", fake_call)
        with pytest.raises(KeyboardInterrupt):
            vb.ask_design("x", None, "gemini", None)
        assert seen == {'key': "sk-test-from-store", 'model': vb._AI_DEFAULT_MODEL['gemini']}
    finally:
        if saved is not None:
            os.environ["GEMINI_API_KEY"] = saved
    with pytest.raises(ValueError):
        vb.ask_design("x", None, "gpt", None)


# ---- agent（2026-09-03・既存シートの「読む→聞く→書く→整える」ループを Python が回す）----

def test_agent_is_wired():
    import vbam_agent as va
    table = vm._command_table()
    assert vm.raw_command(table["agent"]) is va.cmd_agent and table["エージェント"] is table["agent"]   # 台帳の包みつき（2026-09-16）
    p = vm.build_parser()
    ns, unknown = p.parse_known_args(["agent", "郵便番号を揃えて", "--sheet", "テスト用4", "--ai", "gemini",
                                      "--model", "gemini-3.7-flash", "--max-turns", "3", "--dry-run"])
    assert ns.command == "agent" and ns.posargs == ["郵便番号を揃えて"] and ns.sheet_opt == "テスト用4"
    assert ns.ai == "gemini" and ns.max_turns == "3" and ns.dry_run and not unknown
    ns = p.parse_args(["agent", "--fire", "金額の合計"])
    assert ns.fire and ns.posargs == ["金額の合計"]


def test_agent_prompt_and_reply_shape():
    import vbam_agent as va
    prompt = va._first_prompt("郵便番号を揃えて", "Book1", "名簿", "使用範囲: A1:E9\n（材料）")
    # 2026-09-06 夜: 重い道具の段は、依頼にも材料にもその語が無ければ 1 行に畳まれる（_rules_for）
    # 2026-09-07 未明: 規則文は依頼と材料しだいで段を畳む。前置き（先頭の段）はいつも同じ
    assert prompt.startswith(va.RULES[:400])
    assert va._HEAVY_STUB in prompt and "pivot_field" not in prompt.split("言葉の決め:")[0].split(va._HEAVY_STUB)[1]
    assert prompt.index("【依頼】") < prompt.index("【材料】") and "シート「名簿」" in prompt
    res = va._result_prompt([("write_cells 2 セル", True, "書き込み: 名簿!$C$7 ← 000-0002"),
                             ("tidy A3:E9", False, "エラー: x")], 2)
    assert res.startswith("【結果】") and "--- 2: tidy A3:E9 （失敗）---" in res and "残りの往復: 2 回" in res
    # 長い結果は切り詰める（AI に見せる字数の上限）。読む手（read / inspect）は 2 万字まで見せる（2026-09-04・大きい表）
    res = va._result_prompt([("write_cells 2 セル", True, "x" * (va._RESULT_LIMIT + 100))], 1)
    assert "字を省略" in res
    res = va._result_prompt([("read A1", True, "x" * (va._RESULT_LIMIT + 100))], 1)
    assert "字を省略" not in res
    res = va._result_prompt([("read A1", True, "x" * (va._RESULT_LIMIT_BIG + 100))], 1)
    assert "字を省略" in res
    # 返事: done は actions が空のときだけ有効
    say, actions, done, report, plan = va._parse_reply('{"say":"直す","actions":[{"op":"read","range":"A1"}],"done":true}')
    assert say == "直す" and len(actions) == 1 and done is False and plan == []
    say, actions, done, report, plan = va._parse_reply('```json\n{"say":"","actions":[],"done":true,"report":"C7 を直した"}\n```')
    assert done is True and report == "C7 を直した"
    import pytest
    with pytest.raises(ValueError):
        va._parse_reply("これは JSON ではありません")
    with pytest.raises(ValueError):
        va._parse_reply('{"actions": "A1"}')


def test_agent_actions_to_tokens_and_guards():
    import vbam_agent as va
    import pytest
    label, toks = va._action_to_tokens({"op": "read", "range": "a1:d10", "formula": True}, "名簿")
    assert toks == ["read-range", "--sheet", "名簿", "A1:D10", "--formula", "--width", "200"]
    label, toks = va._action_to_tokens({"op": "write_cells", "cells": {"C7": "000-0002", "K11": "=SUM(K6:K10)",
                                                                         "D3": 12000.0, "E3": None}}, "名簿")
    assert toks == ["write-cells", "--sheet", "名簿", "--show", "--", "C7", "000-0002", "K11", "=SUM(K6:K10)",
                    "D3", "12000", "E3", ""]
    label, toks = va._action_to_tokens({"op": "write_grid", "range": "F5", "rows": [["申込コース"], ["年間"]]}, "名簿")
    assert toks[:3] == ["write-range", "--sheet", "名簿"] and toks[-2:] == ["--show", "F5"] and "--tsv" in toks
    assert va._rows_to_tsv([["a", 1], ["b\tc", None]]) == "a\t1\nb c\t\n"
    label, toks = va._action_to_tokens({"op": "format", "range": "A5:E5", "bold": True, "bg": "#DDEBF7",
                                        "col_width": 12, "italic": False}, "名簿")
    assert toks == ["format-range", "--sheet", "名簿", "--show", "A5:E5", "--bold", "--bg=#DDEBF7",
                    "--col-width", "12"]
    label, toks = va._action_to_tokens({"op": "tidy", "ranges": ["A3:F9", "I5:K10"]}, "名簿")
    assert toks == ["tidy", "--sheet", "名簿", "A3:F9", "I5:K10"]
    # 番人: 他シート・列全体・許していない手・中身なし
    for bad in ({"op": "read", "range": "目次!A1"}, {"op": "tidy", "ranges": ["A:A"]},
                {"op": "clear", "range": "A1"}, {"op": "write_cells", "cells": {}},
                {"op": "format", "range": "A1"}, {"op": "write_grid", "range": "A1", "rows": "x"}):
        with pytest.raises(ValueError):
            va._action_to_tokens(bad, "名簿")
    # 規約外の手は実行せず、理由が結果になる（道具は呼ばない）
    results = va._execute([{"op": "row_delete", "row": 3}], "名簿")
    assert results[0][1] is False and "実行していません" in results[0][2]


def test_agent_table_ops_to_tokens():
    """表の足回り 14 本（2026-09-04）。「並べ替えて」を黙って落とした実測から足した手。"""
    import vbam_agent as va
    import pytest
    label, toks = va._action_to_tokens({"op": "sort", "range": "a7:f19", "key": "b", "desc": True}, "名簿")
    assert toks == ["sort", "--sheet", "名簿", "A7:F19", "--key", "B", "--desc", "--header"]
    label, toks = va._action_to_tokens({"op": "sort", "range": "A7:F19", "header": False}, "名簿")
    assert toks[-1] == "--no-header" and "--desc" not in toks
    # autofilter / cond-format / validation に --sheet は無い＝範囲に 'シート名'!A1 の形で渡す
    label, toks = va._action_to_tokens({"op": "autofilter", "range": "A7:F19"}, "名簿")
    assert toks == ["autofilter", "'名簿'!A7:F19"]
    assert va._action_to_tokens({"op": "autofilter", "off": True}, "名簿")[1] == ["autofilter", "'名簿'!A1", "--off"]
    label, toks = va._action_to_tokens({"op": "cond_format", "range": "A8:F19",
                                        "formula": '=$D8=""', "bg": "#FFC7CE"}, "名簿")
    assert toks == ["cond-format", "'名簿'!A8:F19", "--formula", '=$D8=""', "--bg", "#FFC7CE"]
    label, toks = va._action_to_tokens({"op": "cond_format", "range": "B2:B9", "between": [1, 10],
                                        "bold": True}, "名簿")
    assert toks == ["cond-format", "'名簿'!B2:B9", "--between", "1", "10", "--bold"]
    label, toks = va._action_to_tokens({"op": "validation", "range": "C2:C9", "list": ["佐藤", "鈴木"]}, "名簿")
    assert toks == ["validation", "'名簿'!C2:C9", "--list", "佐藤,鈴木"]
    label, toks = va._action_to_tokens({"op": "row", "action": "insert", "at": 7, "count": 2}, "名簿")
    assert toks == ["row", "insert", "7", "2", "--sheet", "名簿"]
    label, toks = va._action_to_tokens({"op": "col", "action": "insert", "at": "c"}, "名簿")
    assert toks == ["col", "insert", "C", "1", "--sheet", "名簿"]
    label, toks = va._action_to_tokens({"op": "copy_range", "src": "A7:F19", "dst": "H7", "values": True}, "名簿")
    assert toks[:3] == ["copy-range", "--sheet", "名簿"] and toks[4] == "'名簿'!H7" and toks[-1] == "--values"
    label, toks = va._action_to_tokens({"op": "find", "text": "貸出中", "whole": True}, "名簿")
    assert toks == ["find", "--sheet", "名簿", "--whole", "--", "貸出中"]
    label, toks = va._action_to_tokens({"op": "chart_config", "action": "legend", "chart": "グラフ 1",
                                        "args": ["bottom"]}, "名簿")
    assert toks == ["chart-config", "legend", "グラフ 1", "bottom"]
    label, toks = va._action_to_tokens({"op": "hyperlink", "cell": "A1", "url": "https://x.example",
                                        "text": "一覧へ"}, "名簿")
    assert toks == ["hyperlink", "'名簿'!A1", "https://x.example", "--text=一覧へ"]
    label, toks = va._action_to_tokens({"op": "comment", "cell": "D8", "text": "未返却"}, "名簿")
    assert toks == ["comment", "'名簿'!D8", "--", "未返却"]
    label, toks = va._action_to_tokens({"op": "sheet_op", "action": "add", "name": "集計"}, "名簿")
    assert toks == ["sheet", "add", "集計"]
    label, toks = va._action_to_tokens({"op": "sheet_op", "action": "rename", "name": "Sheet1",
                                        "to": "集計"}, "名簿")
    assert toks == ["sheet", "rename", "Sheet1", "集計"]
    label, toks = va._action_to_tokens({"op": "name_add", "name": "貸出簿", "range": "A7:F19"}, "名簿")
    assert toks == ["name", "add", "貸出簿", "'名簿'!A7:F19"]
    # 番人: 消す向きは 1 つも通さない（戻せない手は渡さない）
    for bad in ({"op": "row", "action": "delete", "at": 3},
                {"op": "col", "action": "delete", "at": "C"},
                {"op": "sheet_op", "action": "delete", "name": "集計"},
                {"op": "sheet_op", "action": "hide", "name": "集計"},
                {"op": "cond_format", "range": "A1:B2", "clear": True},
                {"op": "validation", "range": "A1:B2", "clear": True},
                {"op": "hyperlink", "cell": "A1", "remove": True},
                {"op": "comment", "cell": "A1", "remove": True}):
        with pytest.raises(ValueError):
            va._action_to_tokens(bad, "名簿")
    # 番人: 中身の欠け・条件の重なり・列文字でないキー・上書きになる PDF
    for bad in ({"op": "sort", "range": "A1:B2", "key": "貸出日"},
                {"op": "cond_format", "range": "A1:B2", "gt": 1, "lt": 9, "bg": "#fff"},
                {"op": "cond_format", "range": "A1:B2", "gt": 1},
                {"op": "cond_format", "range": "A1:B2", "formula": "D8>1", "bg": "#fff"},
                {"op": "validation", "range": "A1:B2", "list": []},
                {"op": "col", "action": "insert", "at": "名"},
                {"op": "row", "action": "insert", "at": 0},
                {"op": "chart_config", "action": "delete", "chart": "グラフ 1"},
                {"op": "export_pdf", "path": "C:/tmp/x.txt"},
                {"op": "find", "text": "  "}):
        with pytest.raises(ValueError):
            va._action_to_tokens(bad, "名簿")
    # 番人: 既にあるファイルには PDF を書かない
    import tempfile, os as _os
    p = _os.path.join(tempfile.gettempdir(), "_agent_pdf_番人.pdf")
    open(p, "w").close()
    try:
        with pytest.raises(ValueError):
            va._action_to_tokens({"op": "export_pdf", "path": p}, "名簿")
    finally:
        _os.remove(p)
    label, toks = va._action_to_tokens({"op": "export_pdf", "path": p, "range": "A1:F19"}, "名簿")
    assert toks == ["export-pdf", p, "--range", "'名簿'!A1:F19"]


def test_agent_tokens_pass_the_real_cli_parser():
    """agent が組み立てた引数列が、本物の CLI パーサを通るか（2026-09-04）。

    トークン列を突き合わせるだけのテストでは「そのコマンドに --sheet があるか」を確かめられず、
    autofilter / cond-format / validation の 3 本が実射で「引数エラー」で落ちた（AI は正しく手を出していた）。
    ここを通せば、手を足すたびに同じ食い違いが単体テストで捕まる。
    """
    import vba_manager
    import vbam_agent as va
    parser = vba_manager.build_parser()
    samples = [
        {"op": "read", "range": "A1:D10"},
        {"op": "write_cells", "cells": {"C7": "x"}},
        {"op": "write_grid", "range": "F5", "rows": [["a"]]},
        {"op": "format", "range": "A1:B2", "bold": True},
        {"op": "tidy", "ranges": ["A1:B2"]},
        {"op": "find_replace", "range": "A1:B2", "search": "a", "replace": "b", "overwrite": True},
        {"op": "sort", "range": "A1:F9", "key": "B", "desc": True},
        {"op": "sort", "range": "A1:F9", "header": False},
        {"op": "autofilter", "range": "A1:F9"},
        {"op": "autofilter", "off": True},
        {"op": "cond_format", "range": "A1:F9", "formula": '=$D1=""', "bg": "#FFC7CE"},
        {"op": "cond_format", "range": "B2:B9", "between": [1, 10], "bold": True},
        {"op": "cond_format", "range": "B2:B9", "gt": 100, "color": "#FF0000"},
        {"op": "validation", "range": "C2:C9", "list": ["佐藤", "鈴木"]},
        {"op": "row", "action": "insert", "at": 7, "count": 2},
        {"op": "col", "action": "insert", "at": "C"},
        {"op": "copy_range", "src": "A1:B2", "dst": "H1", "values": True},
        {"op": "find", "text": "貸出中", "whole": True},
        {"op": "export_pdf", "path": "C:/tmp/_agent_parser_試験.pdf", "range": "A1:B2"},
        {"op": "chart_config", "action": "legend", "chart": "グラフ 1", "args": ["bottom"]},
        {"op": "chart_config", "action": "axis-scale", "chart": "グラフ 1", "args": ["value"], "min": 0, "max": 10},
        {"op": "hyperlink", "cell": "A1", "url": "https://x.example", "text": "先"},
        {"op": "comment", "cell": "A1", "text": "メモ"},
        {"op": "sheet_op", "action": "add", "name": "集計"},
        {"op": "sheet_op", "action": "rename", "name": "Sheet1", "to": "集計"},
        {"op": "sheet_op", "action": "tab-color", "name": "集計", "to": "#FF0000"},
        {"op": "name_add", "name": "貸出簿", "range": "A1:F9"},
        {"op": "table", "action": "create", "range": "A1:E9", "name": "T名簿"},
        {"op": "pivot", "action": "list"},
    ]
    for act in samples:
        _, toks = va._action_to_tokens(act, "名簿")
        try:
            got = parser.parse_args(toks)          # 引数が食い違えば SystemExit
        except SystemExit:
            raise AssertionError(f"{act['op']} の引数が CLI パーサに通らない: {toks}")
        assert got.command == toks[0], f"{act['op']}: {toks}"
    # macro 側も同じく通す
    for mact, gate in (({"op": "get", "name": "集計"}, False), ({"op": "grep", "text": "x"}, False),
                       ({"op": "compile"}, False), ({"op": "list"}, False),
                       ({"op": "impact", "name": "集計"}, False), ({"op": "call_graph"}, False),
                       ({"op": "run", "name": "集計"}, True),
                       ({"op": "rehearse", "name": "集計"}, True)):
        _, toks, _w = vmac._macro_action_to_tokens(mact, allow_rehearse=gate)
        try:
            got = parser.parse_args(toks)
        except SystemExit:
            raise AssertionError(f"macro {mact['op']} の引数が CLI パーサに通らない: {toks}")
        assert got.command == toks[0]


def test_agent_macro_read_ops_and_run_gate():
    """macro の読む手 3 本と run の関所（2026-09-04）。run は --rehearse のときだけ。"""
    import vbam_agent as va
    import pytest
    assert vmac._macro_action_to_tokens({"op": "list"})[1] == ["list"]
    assert vmac._macro_action_to_tokens({"op": "impact", "name": "集計"})[1] == ["impact", "集計"]
    assert vmac._macro_action_to_tokens({"op": "call_graph"})[1] == ["call-graph"]
    assert vmac._macro_action_to_tokens({"op": "call_graph", "name": "集計"})[1] == ["call-graph", "--macro", "集計"]
    label, toks, w = vmac._macro_action_to_tokens({"op": "run", "name": "集計"}, allow_rehearse=True)
    assert toks[:2] == ["run-macro", "集計"] and w is False
    with pytest.raises(ValueError):
        vmac._macro_action_to_tokens({"op": "run", "name": "集計"})          # 関所が閉じている
    with pytest.raises(ValueError):
        vmac._macro_action_to_tokens({"op": "impact"})


def test_agent_lock_blocks_second_run(tmp_path, monkeypatch):
    """記録が混ざらないように、走っている agent がいる間は次を断る（PID が死んでいれば通す）。"""
    import vbam_agent as va
    import json as _json
    import pytest
    lock = tmp_path / "_agent_running.lock"
    monkeypatch.setattr(va, '_AGENT_LOCK_FILE', str(lock))
    with va._agent_lock():
        assert lock.exists()
        got = _json.loads(lock.read_text(encoding='utf-8'))
        assert got['pid'] == os.getpid()
        # 別プロセス（生きている PID）が持っていれば断る
        lock.write_text(_json.dumps({'pid': 4, 'at': 0, 'at_text': '00:00:00'}), encoding='utf-8')
        if va._pid_alive(4):
            with pytest.raises(RuntimeError):
                with va._agent_lock():
                    pass
        # 死んでいる PID の錠は無視して通る
        lock.write_text(_json.dumps({'pid': 999999, 'at': 0, 'at_text': '00:00:00'}), encoding='utf-8')
        with va._agent_lock():
            pass
    assert not lock.exists()


def test_agent_route_and_entry_options():
    # 薄い入口（手3）: --mode > --macro > --new-book > マクロの語 > 「新しいシート」の語か白紙 > sheet
    import vbam_agent as va
    import pytest
    assert va._route("郵便番号を揃えて") == 'sheet'
    assert va._route("合計の行を作って") == 'sheet'                     # 「作って」だけでは build にしない
    assert va._route("会員名簿を新しいシートに作って") == 'build'
    assert va._route("備品台帳", sheet_empty=True) == 'build'
    assert va._route("郵便番号を揃えて", new_book=True) == 'build'
    assert va._route("集計マクロが動かない") == 'macro'
    assert va._route("Sub 集計 のエラーを直して") == 'macro'
    assert va._route("何でも", macro="集計") == 'macro'
    assert va._route("マクロが動かない", forced="sheet") == 'sheet'        # --mode が最優先
    # マクロの語はあるが、ブックのどのマクロ名も依頼文に無い → sheet に落とす（2026-09-04。名前一覧を渡さなければ従来どおり）
    assert va._route("数式が動かない", macro_names=["集計", "印字"]) == 'sheet'
    assert va._route("集計マクロが動かない", macro_names=["集計", "印字"]) == 'macro'
    assert va._route("数式が動かない", macro_names=[]) == 'sheet'          # マクロが 1 本も無いブック
    assert va._route("数式が動かない") == 'macro'                          # 一覧なし＝語だけで決める（従来）
    assert va._route("VBA で作った表を新しいシートに写して", macro_names=["集計"]) == 'build'   # 落ちた先は普通の振り分け
    with pytest.raises(ValueError):
        va._route("x", forced="repair")
    p = vm.build_parser()
    ns, unknown = p.parse_known_args(["agent", "台帳を作って", "--mode", "build", "--new-book", "--macro", "集計",
                                      "--rehearse"])
    assert ns.mode == "build" and ns.new_book and ns.macro == "集計" and ns.rehearse and not unknown


def test_compound_notice_flags_a_request_spanning_macro_and_sheet():
    """8安: 「マクロを直してから表も整えて」は macro に振られ、シート側の仕事が黙って落ちていた（2026-09-05）。"""
    import vbam_agent as va
    notice = va._compound_notice('macro', 'マクロを直してから、表も見やすく整えて', None)
    assert notice and "シートの仕事も含まれています" in notice and "--mode sheet" in notice


def test_compound_notice_flags_an_embedded_macro_name_when_routed_to_sheet():
    """8安: sheet に振られたが依頼文に実在のマクロ名があるときの案内。"""
    import vbam_agent as va
    notice = va._compound_notice('sheet', '売上集計マクロと同じように色分けして', '売上集計マクロ')
    assert notice and "マクロ名「売上集計マクロ」" in notice and "--macro 売上集計マクロ" in notice


def test_compound_notice_is_silent_when_nothing_spans():
    """8安: またがっていない依頼・macro_name が無いとき・build のときは何も出さない（空文字）。"""
    import vbam_agent as va
    assert va._compound_notice('sheet', '見やすくして', None) == ''
    assert va._compound_notice('macro', 'コンパイルエラーを直して', None) == ''
    assert va._compound_notice('build', '新しいシートに集計して見やすくして', '何か') == ''
    assert 'agent' in {n for items in vv._CAPABILITIES.values() for n, _ in items}


def test_compound_notice_flags_a_build_that_also_asks_for_a_pivot():
    """2026-09-09: 「表を作って、そこからピボットも作って」を build に振ると、表だけ組んで黙って終わっていた。

    組む engine は設計図（列と型と仕上げ）しか持たず、ピボット・スライサー・クエリ・モデルを作れない。
    macro↔sheet と同じ形で案内を出す（黙って半分落とさない）。
    """
    import vbam_agent as va
    for word in ('ピボット', 'スライサー', 'パワークエリ', 'データモデル', 'メジャー'):
        notice = va._compound_notice('build', f'明細を作って、そこから{word}も作って', None)
        assert notice and word in notice and '--mode sheet' in notice, word
    # 組む engine で足りる依頼には出さない（テーブル・グラフ・集計は含めない）
    for plain in ('新しいシートに名簿の表を作って', '売上の表を作ってグラフも付けて', '集計表を組んで'):
        assert va._compound_notice('build', plain, None) == '', plain


def test_grade_materials_hands_the_pivot_inventory_to_the_grader(monkeypatch):
    """2026-09-09: 採点係は materials（そのシートのセル）しか渡されておらず、正しく作ったピボットを

    「値のみ貼り付けの可能性あり」と疑って done を突き返していた（実射で往復が 3 回増えた）。
    ピボットの現物はセルに写らない＝作業側と同じ一覧を渡す。0 本だけの回は足さない（送りを増やさない）。
    """
    import vbam_agent as va
    monkeypatch.setattr(va, '_run_cmd', lambda argv, wb: (True, f"使用範囲: A3:B7（{argv[1]}）"))
    monkeypatch.setattr(vg, '_grade_structure', lambda ws: [])
    monkeypatch.setattr(vg, '_grade_looks', lambda ws: [])
    monkeypatch.setattr(va, '_TIDY_RANGES', {})

    monkeypatch.setattr(vg, '_heavy_materials', lambda wb: [
        'テーブル（ブック全体）: 0 本', 'ピボット（ブック全体）: 1 本',
        '  [集計] P集計 出力=A3:E8 行=地域 列=商品 値=sum/売上', 'スライサー: 0 個'])
    mat = va._grade_materials('集計', object())
    assert '【ブックにある重い物' in mat and 'P集計' in mat and '行=地域' in mat
    assert '値の貼り付けではありません' in mat

    monkeypatch.setattr(vg, '_heavy_materials', lambda wb: [
        'テーブル（ブック全体）: 0 本', 'ピボット（ブック全体）: 0 本', 'スライサー: 0 個'])
    assert '【ブックにある重い物' not in va._grade_materials('名簿', object())


def test_pivot_covering_finds_the_pivot_under_a_range():
    """2026-09-09: ピボットは見出し行・ラベル行と本文で表示形式が違う。

    列ごとに読むと NumberFormat が None（＝まちまち）になり、採点係が「桁区切りがまちまち」と誤指摘する
    （実射で確認。桁区切りは全セルに付いていた）。当たり判定は番地の算術だけ＝COM の Intersect を使わない。
    """
    import vbam_agent as va

    class _R:
        def __init__(self, r, c, nr, nc):
            self.Row, self.Column = r, c
            self.Rows, self.Columns = type('X', (), {'Count': nr})(), type('X', (), {'Count': nc})()

    class _PT:
        def __init__(self, name, rng):
            self.Name, self.TableRange2 = name, rng

    class _WS:
        def __init__(self, pts):
            self._pts = pts

        def PivotTables(self):
            return self._pts

    ws = _WS([_PT('P集計', _R(3, 1, 6, 5))])                 # 集計!A3:E8
    assert va._pivot_covering(ws, _R(4, 1, 5, 5)).Name == 'P集計'    # 見出し推定 4 行目から広がる範囲
    assert va._pivot_covering(ws, _R(3, 1, 1, 1)).Name == 'P集計'    # 左上 1 セルでも重なれば当たり
    assert va._pivot_covering(ws, _R(20, 1, 3, 3)) is None           # 下に離れた表は当たらない
    assert va._pivot_covering(ws, _R(3, 9, 3, 3)) is None            # 右に離れた表も当たらない
    assert va._pivot_covering(_WS([]), _R(3, 1, 6, 5)) is None


def test_agent_macro_actions_and_guards():
    # 手2: マクロ修理の手 → コマンド。書き換えの手だけ「書き換え」印がつく（道具が compile を足す）
    import vbam_agent as va
    import pytest
    label, toks, w = vmac._macro_action_to_tokens({"op": "get", "name": "集計", "module": "shu003"})
    assert toks == ["get", "shu003", "集計"] and not w
    label, toks, w = vmac._macro_action_to_tokens({"op": "grep", "text": "ActiveSheet"})
    assert toks == ["grep", "ActiveSheet"] and not w
    label, toks, w = vmac._macro_action_to_tokens({"op": "code_replace", "search": "r.Valeu", "replace": "r.Value",
                                                 "module": "M1"})
    assert toks == ["code-replace", "-y", "r.Valeu", "r.Value", "--module", "M1"] and w
    code = "Sub 集計()\n    Range(\"A1\").Value = 1\nEnd Sub"
    label, toks, w = vmac._macro_action_to_tokens({"op": "replace", "name": "集計", "code": code})
    assert toks == ["replace-procedure", "-y"] and w
    label, toks, w = vmac._macro_action_to_tokens({"op": "compile"})
    assert toks == ["compile"] and not w
    label, toks, w = vmac._macro_action_to_tokens({"op": "rehearse", "name": "集計", "args": [1]}, allow_rehearse=True)
    assert toks == ["rehearse", "集計", "1", "--timeout", "120"] and not w
    for bad in ({"op": "delete", "name": "x"},                                  # 許していない手
                {"op": "code_replace", "search": "a\nb", "replace": "c"},        # 複数行
                {"op": "replace", "name": "集計", "code": "Sub 別名()\nEnd Sub"},  # 名前が違う
                {"op": "replace", "name": "集計", "code": code + "\nSub 二本目()\nEnd Sub"},  # 2 本
                {"op": "rehearse", "name": "集計"}):                             # --rehearse なし
        with pytest.raises(ValueError):
            vmac._macro_action_to_tokens(bad)
    assert vmac._find_macro_in_request("集計テスト が動かない", ["集計", "集計テスト", "印字"]) == "集計テスト"
    assert vmac._find_macro_in_request("何も無い", ["集計"]) is None
    # 実射の弾（macro）は修理 9 本＋新規作成 1 本。修理の弾は壊れたコード・依頼文・期待するセルを持つ
    # （2026-09-04 夜に 4 本 → 10 本。現場で実際に出る止まり方を足した＝母数不足の是正）
    assert [c["name"] for c in vf._FIRE_MACRO_CASES] == ["コンパイルエラーの修理", "打ち間違いの修理",
                                                          "存在しないマクロ呼びの修理", "新しいマクロを作る",
                                                          "存在しないシート", "型が一致しません", "Set の付け忘れ",
                                                          "引数付きで実行できない", "エラーは出ないが数が合わない",
                                                          "合計が0になる"]
    assert all(c["macro"] in c["code"] and c["request"] and c["expect"]
               for c in vf._FIRE_MACRO_CASES if not c.get('create'))


def test_agent_macro_create_hand_and_route():
    # 10（2026-09-04）: マクロを「作る」道。add は名前の重なり・宣言 1 本・入れ先を検査する
    import pytest
    import vbam_agent as va
    code = 'Sub 集計を書く()\n    Range("A1").Value = "合計"\nEnd Sub\n'
    label, toks, is_write = vmac._macro_action_to_tokens(
        {"op": "add", "module": "shu003", "name": "集計を書く", "code": code}, existing=["別のマクロ"])
    assert toks == ['add-procedure', '-y', 'shu003'] and is_write and label == "add 集計を書く → shu003"
    with pytest.raises(ValueError, match="既にあります"):
        vmac._macro_action_to_tokens({"op": "add", "module": "shu003", "name": "集計を書く", "code": code},
                                   existing=["集計を書く"])
    with pytest.raises(ValueError, match="module"):
        vmac._macro_action_to_tokens({"op": "add", "name": "集計を書く", "code": code})
    with pytest.raises(ValueError, match="名前"):
        vmac._macro_action_to_tokens({"op": "add", "module": "m", "name": "ちがう名前", "code": code})
    with pytest.raises(ValueError, match="1 本だけ"):
        vmac._macro_action_to_tokens({"op": "add", "module": "m", "name": "二本", "code": code + code})
    # 振り分け: マクロの語＋作る語 → macro（そのブックに無い名前でも通る）。直す語だけなら従来どおり
    assert va._route("A 列を消すマクロを作って", macro_names=["集計", "印字"]) == 'macro'
    assert va._route("マクロを追加して", macro_names=[]) == 'macro'
    assert va._route("数式が動かない", macro_names=["集計"]) == 'sheet'
    assert va._route("VBA で作った表を新しいシートに写して", macro_names=["集計"]) == 'build'
    # 実射の弾（新しく作る）は code を植えず、依頼文にシート名とモジュール名が入る
    case = [c for c in vf._FIRE_MACRO_CASES if c.get('create')][0]
    assert 'code' not in case and case['macro'] == "実射の集計を書く"
    req = case['request'].format(sheet="実射4", module="M実射4")
    assert "実射4" in req and "M実射4" in req and case['expect'] == {"A1": "合計", "B1": 60}


def test_agent_pins_target_book(tmp_path):
    # 対象ブックの固定: 保存済みならフルパスを先頭の位置引数に足す。未保存（Book1）は足さない（Activate で固定）
    import vbam_agent as va

    class _WB:
        def __init__(self, full):
            self.FullName = full
    saved = tmp_path / "台帳.xlsm"
    saved.write_bytes(b"")
    assert va._pin_target(_WB(str(saved))) == str(saved)
    assert va._pin_target(_WB("Book1")) is None and va._pin_target(None) is None
    # 解析後の posargs の先頭に足す（引数列に挟むと argparse が位置引数を分断して「不明な引数」になる）
    import vbam_core as vc
    ns, unknown = vm.build_parser().parse_known_args(["write-cells", "--sheet", "名簿", "--show", "C7", "x"])
    assert not unknown
    ns.posargs = [va._pin_target(_WB(str(saved)))] + ns.posargs
    target, rest = vc.parse_target_and_rest(ns.posargs)
    assert target == str(saved) and rest == ["C7", "x"]


def test_agent_big_table_normalize_rules():
    # 大きい表の手（2026-09-04）: AI に行を触らせない。規則は純 Python で固定する
    import datetime as dt
    import pytest
    import vbam_agent as va
    N = lambda v, rules, **kw: va._normalize_value(v, [(r, None) if isinstance(r, str) else r for r in rules], **kw)[0]
    assert N(" 青木 誠 ", ["trim"]) == "青木 誠" and N("　", ["trim"]) == "" and N(12000, ["trim"]) == 12000
    assert N("１２,０００", ["hankaku"]) == "12,000" and N("０００－０００１", ["hankaku"]) == "000-0001"
    assert N("000‐0001", ["hyphen"]) == "000-0001" and N("1ー2ー3", ["hyphen"]) == "1-2-3" and N("サーバー", ["hyphen"]) == "サーバー"
    assert N("12,000", ["number"]) == 12000 and N("１２,０００", ["number"]) == 12000 and N("9.5", ["number"]) == 9.5
    assert N("0001", ["number"]) == "0001"                                   # 先頭ゼロの番号は文字のまま
    assert N("0001", ["number"], keep_zero=False) == 1
    assert N("0", ["number"]) == 0 and N("abc", ["number"]) == "abc" and N("12,000円", ["number"]) == 12000
    d = N("2026/01/05", ["date"])
    assert d == dt.datetime(2026, 1, 5, tzinfo=dt.timezone.utc)             # 素の datetime は 9 時間ずれる＝tz 付き
    assert N("2026年1月5日", ["date"]).day == 5 and N("２０２６-２-１０", ["date"]).month == 2
    assert N("4月20日", ["date"]) == "4月20日" and N("2026/13/1", ["date"]) == "2026/13/1"
    rx = ('regex', (va.re.compile(r'^(\d{3})(\d{4})$'), r'\1-\2'))
    assert N("0000002", [rx]) == "000-0002" and N("000-0002", [rx]) == "000-0002"
    assert N(12000.0, ["as_text"]) == "12000" and N(d, ["as_text"]) == "2026/1/5" and N("0001", ["as_text"]) == "0001"
    # 規則は順に当たる: 全角の金額 → 半角 → 数値
    assert N("１２,０００", ["hankaku", "number"]) == 12000
    # 列単位: fill_down は空白を上の値で埋める。数式の位置（skip）は触らない
    out, counts, ex = va._normalize_column(["総務", None, "", "営業", " ", "=A1"], [("fill_down", None)],
                                           skip=[False, False, False, False, False, True])
    assert out == ["総務", "総務", "総務", "営業", "営業", "=A1"] and counts == {'fill_down': 3}
    out, counts, _ = va._normalize_column([" a ", "b", None, "=A1 "], [("trim", None)], skip=[False, False, False, True])
    assert out == ["a", "b", None, "=A1 "] and counts == {'trim': 1}
    # 手の検査（純 Python）: 元の列に書くには overwrite、規則名の綴り、regex の形、to は 1 セル
    rng, rules, to, header, ow, kz = va._normalize_spec({"op": "normalize", "range": "C2:C300", "rules": ["trim", "number"],
                                                         "to": "H2", "header": "電話（整形）"})
    assert (rng, to, header, ow, kz) == ("C2:C300", "H2", "電話（整形）", False, True) and [r[0] for r in rules] == ["trim", "number"]
    with pytest.raises(ValueError, match="overwrite"):
        va._normalize_spec({"op": "normalize", "range": "C2:C300", "rules": ["trim"]})
    with pytest.raises(ValueError, match="規則"):
        va._normalize_spec({"op": "normalize", "range": "C2:C3", "rules": ["upper"], "overwrite": True})
    with pytest.raises(ValueError, match="regex"):
        va._normalize_spec({"op": "normalize", "range": "C2:C3", "rules": [{"regex": {"search": "("}}], "overwrite": True})
    with pytest.raises(ValueError, match="1 セル"):
        va._normalize_spec({"op": "normalize", "range": "C2:C3", "rules": [], "to": "H2:H3"})
    _, rules, _, _, _, _ = va._normalize_spec({"op": "normalize", "range": "C2:C3", "overwrite": True,
                                               "rules": [{"regex": {"search": "^(\\d{3})(\\d{4})$", "replace": "\\1-\\2"}}]})
    assert rules[0][0] == 'regex' and rules[0][1][0].sub(rules[0][1][1], "0000002") == "000-0002"


def test_agent_big_table_profile_and_hands():
    import vbam_agent as va
    import pytest
    values = [["番号", "氏名", "郵便番号", "金額", "入会日"],
              ["0001", " 青木 誠", "000-0001", 12000, "2026/1/5"],
              ["0002", "石川 恵", "0000002", "１２,０００", "2026-02-10"],
              ["0003", "上田 学", "０００－０００３", 9500, None],
              ["0004", "遠藤 香", " 000-0004 ", "7000", "2026/3/1"],
              [None, "大西 亮", None, -2146826246, "4月20日"]]
    formulas = [[None] * 5 for _ in values]
    formulas[5][3] = "=1/0"
    lines = va._profile_columns(values, formulas, 1, 1)
    text = "\n".join(lines)
    assert text.count("\n") == 4                                              # 1 列 1 行
    assert "A 番号: 文字 4 / 空白 1" in text and "先頭ゼロ 4" in text and "形: 9999 ×4" in text
    assert "B 氏名:" in text and "前後空白 1（例 B2" in text
    assert "C 郵便番号: 文字 4 / 空白 1" in text and "全角 1（例 C4" in text and "数値に見える文字 1（例 C3" in text
    assert "999-9999 ×2" in text and "9999999 ×1" in text and "前後空白 1（例 C5" in text
    assert "D 金額: 数値 2 / 文字 2 / 数式 1" in text and "数値に見える文字 2" in text and "数値の範囲 9,500〜12,000" in text    # 数式は型に数えない
    assert "E 入会日: 文字 4 / 空白 1" in text and "日付に見える文字 4" in text
    assert "多い値" not in text                                                # 全部 1 件ずつの列（氏名）には出さない
    prof = va._profile_columns([["区分"]] + [["正会員"]] * 4 + [["準会員"]] * 2, None, 1, 1)
    assert "多い値: 正会員 ×4, 準会員 ×2" in prof[0]
    # 見出しが無い表は列文字だけ
    assert va._profile_columns([[1, 2], [3, 4]], None, 1, 1)[0].startswith("  A （見出しなし）: 数値 2")
    # 手の変換: find_replace は overwrite 必須。write_grid は 100 行まで
    label, toks = va._action_to_tokens({"op": "find_replace", "range": "C2:C300", "search": "－", "replace": "-",
                                        "overwrite": True, "whole": False}, "名簿")
    assert toks == ['find-replace', '--sheet', '名簿', '--', '－', '-', 'C2:C300'] and label.startswith("find_replace C2:C300")
    with pytest.raises(ValueError, match="overwrite"):
        va._action_to_tokens({"op": "find_replace", "range": "C2:C300", "search": "a", "replace": "b"}, "名簿")
    with pytest.raises(ValueError, match="100 行まで"):
        va._action_to_tokens({"op": "write_grid", "range": "A1", "rows": [["x"]] * 101}, "名簿")
    assert va._action_to_tokens({"op": "write_grid", "range": "A1", "rows": [["x"]] * 100}, "名簿")[0].endswith("100行")
    # 数式を読むときは切らない（既定の 40 字だと外部参照のパスが読めず、控えを報告に書けない）
    assert va._action_to_tokens({"op": "read", "range": "B2:B3", "formula": True}, "名簿")[1] == \
        ['read-range', '--sheet', '名簿', 'B2:B3', '--formula', '--width', '200']
    assert va._action_to_tokens({"op": "read", "range": "B2:B3"}, "名簿")[1] == \
        ['read-range', '--sheet', '名簿', 'B2:B3']
    for op in ("normalize", "fill", "find_replace"):
        assert op in va._ALLOWED_OPS
    assert "normalize" in va._DIRECT_OPS and "fill" in va._DIRECT_OPS and "find_replace" not in va._DIRECT_OPS
    for word in ("normalize", "fill", "find_replace", "列プロファイル", "100 行まで", "fill_down", "as_text"):
        assert word in va.RULES
    # 読む手の結果は 6,000 字で切らない（2 万字まで）
    big = "x" * 10000
    msg = va._result_prompt([("read A1:F300", True, big), ("write_cells 3 セル", True, big)], 2)
    assert "字を省略" in msg.split("--- 2:")[1] and "字を省略" not in msg.split("--- 2:")[0]
    # 弾: 300 行の練習台は決め打ちで 4 通りの乱れを含む
    rows = vf._big_dirty_rows()
    assert len(rows) == 300 and rows[0][0] == "0001"
    kinds = {va._shape_of(r[2].strip(va._WS_CHARS)) for r in rows}
    assert {"999-9999", "9999999", "９９９－９９９９"} <= kinds
    assert sum(1 for r in rows if isinstance(r[3], str)) == 200 and len(vf._big_match_right()) == 56


def test_agent_shows_the_image_to_the_ai(tmp_path, monkeypatch):
    # 11（2026-09-04）: 書いた直後の見た目を AI にも見せる（罫線・寄せ・### は文字では見えない）
    import base64
    import vbam_agent as va
    png = tmp_path / "view.png"
    png.write_bytes(b"\x89PNG\r\n\x1a\n" + b"x" * 20)
    b64 = base64.b64encode(png.read_bytes()).decode('ascii')
    hist = [('user', "1 通目", None), ('model', '{"actions":[]}'), ('user', "2 通目", str(png))]
    sent = {}

    class FakeResp:
        def read(self):
            return json.dumps({"candidates": [{"content": {"parts": [{"text": "{}"}]}}], "usageMetadata": {}}).encode()

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    def fake_urlopen(req, timeout=None):
        sent['body'] = json.loads(req.data.decode('utf-8'))
        return FakeResp()

    import json
    monkeypatch.setattr(va.urllib.request, "urlopen", fake_urlopen)
    monkeypatch.setattr(va.json, "load", lambda r: json.loads(r.read().decode()))
    va._chat_gemini(hist, "gemini-3.7-flash", "k")
    parts = sent['body']['contents']
    assert [len(c['parts']) for c in parts] == [1, 1, 2]                    # 画像が付くのは 3 通目だけ
    assert parts[2]['parts'][1]['inline_data'] == {"mime_type": "image/png", "data": b64}
    assert parts[1]['role'] == 'model' and parts[2]['role'] == 'user'
    va._chat_claude(hist, "claude-haiku-4-5-20251001", "k")
    msgs = sent['body']['messages']
    assert isinstance(msgs[0]['content'], str) and isinstance(msgs[2]['content'], list)
    assert msgs[2]['content'][1]['source'] == {"type": "base64", "media_type": "image/png", "data": b64}
    # 画像は「書く手があった往復」の後だけ（読むだけの往復では撮らない）＝_will_write と同じ判定
    assert va._will_write([{"op": "tidy", "ranges": ["A1:C3"]}]) and not va._will_write([{"op": "read", "range": "A1"}])
    p = vm.build_parser()
    ns, unknown = p.parse_known_args(["agent", "x", "--no-image"])
    assert ns.no_image and not unknown and not ns.image          # 既定は画像なし（2026-09-08 に反転。--no-image は互換）
    ns, unknown = p.parse_known_args(["agent", "x", "--image"])
    assert ns.image and not unknown
    # 見せる範囲は使用範囲の左上から 40 行 × 20 列まで＋上下左右 2 行 2 列の余白（行 1・列 A で止める）
    ws = type("WS", (), {"UsedRange": type("UR", (), {"Row": 3, "Column": 2,
                                                      "Rows": type("R", (), {"Count": 300})(),
                                                      "Columns": type("C", (), {"Count": 40})()})()})()
    assert va._view_range(ws) == ("A1:T40", 40, 20)
    ws2 = type("WS", (), {"UsedRange": type("UR", (), {"Row": 10, "Column": 5,
                                                       "Rows": type("R", (), {"Count": 3})(),
                                                       "Columns": type("C", (), {"Count": 2})()})()})()
    assert va._view_range(ws2) == ("C8:H14", 7, 6)              # 3 行 2 列の表に上下左右 2 の余白


def test_agent_undo_backup_and_diff(tmp_path, monkeypatch):
    # 取り消し（2026-09-04）: 最初の「書く手」の前だけ控えを取り、終わりに差分を数える
    import vbam_agent as va
    # 書く手があるかの判定（読むだけの手では控えを取らない）
    assert va._will_write([{"op": "read", "range": "A1"}]) is False
    assert va._will_write([{"op": "inspect", "sheet": "名簿"}]) is False
    assert va._will_write([{"op": "table", "action": "list"}, {"op": "pivot_calc", "action": "get_data"}]) is False
    assert va._will_write([{"op": "read", "range": "A1"}, {"op": "tidy", "ranges": ["A1:C3"]}]) is True
    assert va._will_write([{"op": "normalize", "range": "A2:A9", "overwrite": True}]) is True
    assert va._will_write([{"op": "table", "action": "create", "range": "A1:C3"}]) is True
    assert va._will_write(["こわれた手"]) is True
    # 差分: 値と数式の両方を見て、増えた行・消えた値も数える
    before = {'addr': 'A1:B2', 'row': 1, 'col': 1, 'values': [["a", 1], ["b", 2]],
              'formulas': [["a", 1], ["b", 2]], 'shapes': []}

    class FakeWS:
        def __init__(self, values, formulas=None, row=1, col=1):
            self._v, self._f, self._r, self._c = values, formulas or values, row, col

        @property
        def UsedRange(self):
            ws = self

            class UR:
                Row, Column = ws._r, ws._c
                Value = tuple(tuple(r) for r in ws._v)
                Formula = tuple(tuple(r) for r in ws._f)
                Address = "$A$1"

                class Cells:
                    CountLarge = 100
            return UR
        Shapes = ()

    n, ex = va._diff_sheet(before, FakeWS([["a", 1], ["b", 99]]))
    assert n == 1 and ex == ["B2 '2'→'99'"]
    n, ex = va._diff_sheet(before, FakeWS([["a", 1], ["b", 2], ["c", 3]]))
    assert n == 2 and ex == ["A3 ''→'c'", "B3 ''→'3'"]                 # 増えた行も差分
    n, _ = va._diff_sheet(before, FakeWS([["a", 1], ["b", 2]]))
    assert n == 0
    n, ex = va._diff_sheet(before, FakeWS([["a", 1], ["b", 2]], formulas=[["a", 1], ["b", "=1+1"]]))
    assert n == 1 and ex == ["B2 '2'→'=1+1'"]                          # 値が同じでも式に変わったら差分
    assert va._diff_sheet({'too_big': 999999}, FakeWS([["a"]])) == (None, [])
    assert va._diff_sheet(None, FakeWS([["a"]])) == (None, [])
    # 控えが無ければ --undo は止まる（Excel に触らない）
    monkeypatch.setattr(vu, "_LAST_AGENT_UNDO_FILE", str(tmp_path / "undo.json"))
    monkeypatch.setattr(vu, "get_workbook", lambda *a, **k: (_ for _ in ()).throw(AssertionError("Excel に触った")))
    assert va.undo_agent() is False
    (tmp_path / "undo.json").write_text('{"path": "C:/no/such.xlsx", "book": "x.xlsx", "sheet": "s"}', encoding='utf-8')
    assert va.undo_agent() is False                                    # 控えのファイルが無い＝ここでも Excel に触らない
    p = vm.build_parser()
    ns, unknown = p.parse_known_args(["agent", "--undo", "--dry-run"])
    assert ns.undo and ns.dry_run and not unknown
    # 控えは新しい 5 個だけ残す
    monkeypatch.setattr(vu, "BACKUP_DIR", str(tmp_path))
    for i in range(8):
        f = tmp_path / f"本{va._AGENT_BACKUP_MARK}2026090{i}_000000.xlsx"
        f.write_text("x", encoding='utf-8')
        os.utime(f, (1757000000 + i, 1757000000 + i))
    (tmp_path / "他のバックアップ.bas").write_text("x", encoding='utf-8')
    va._prune_agent_backups()
    left = sorted(f.name for f in tmp_path.iterdir() if va._AGENT_BACKUP_MARK in f.name)
    assert len(left) == 5 and left[0].endswith("20260903_000000.xlsx")
    assert (tmp_path / "他のバックアップ.bas").exists()                 # 他のバックアップは触らない


def test_build_hands_failed_verify_to_the_sheet_loop(monkeypatch):
    # 2026-09-04: 組み上げの検査に落ちたら、そのまま sheet のループに渡して直させる（前は一発勝負）
    import vbam_agent as va
    calls = {}
    monkeypatch.setattr(va, "ask_design", lambda *a, **k: {"sheet": "集計"})
    monkeypatch.setattr(va, "plan_sheet", lambda spec: {"sheet": "集計", "overwrite": False})
    monkeypatch.setattr(va, "_plan_summary", lambda plan: "（計画）")
    monkeypatch.setattr(va, "_apply", lambda xl, wb, plan, ov: type("WS", (), {"Name": "集計"})())
    verdicts = iter([False])
    monkeypatch.setattr(va, "_verify_build", lambda xl, ws, plan: next(verdicts, True))

    def fake_run_agent(request, sheet, wb, ai=None, model=None, max_turns=None, backup=None, **kw):
        calls.update(request=request, sheet=sheet, max_turns=max_turns, backup=backup)
        return {'verify': True, 'done': True}

    monkeypatch.setattr(va, "run_agent", fake_run_agent)
    assert va.run_build("会員名簿を作って", None, object(), "gemini", None) is True     # 直して合格になった
    assert calls['sheet'] == "集計" and calls['max_turns'] == va._BUILD_FIX_TURNS and calls['backup'] is False
    assert "###" in calls['request'] and "会員名簿を作って" in calls['request']
    # 検査に通ったときは直しのループを回さない
    calls.clear()
    monkeypatch.setattr(va, "_verify_build", lambda xl, ws, plan: True)
    assert va.run_build("x", None, object(), "gemini", None) is True and calls == {}
    # fix=False なら落ちてもそのまま返す（--dry-run と同じで余計な往復をしない）
    monkeypatch.setattr(va, "_verify_build", lambda xl, ws, plan: False)
    assert va.run_build("x", None, object(), "gemini", None, fix=False) is False and calls == {}


def test_agent_continue_resumes_the_conversation(tmp_path, monkeypatch):
    # 承認の続き（2026-09-04）: 記録から会話を読み戻し、追加の指示だけを足して続ける
    import json
    import vbam_agent as va
    log = tmp_path / "log.jsonl"
    log.write_text("\n".join(json.dumps(d, ensure_ascii=False) for d in [
        {"meta": {"book": "名簿.xlsx", "sheet": "会員", "mode": "sheet", "request": "手順書「データの掃除」", "model": "m"}},
        {"turn": 1, "prompt": "規則と材料…", "reply": '{"say":"見ます","actions":[{"op":"read","range":"A1:D9"}]}',
         "results": [{"label": "read A1:D9", "ok": True, "out": "格子…"}]},
        {"turn": 2, "prompt": "【結果】…", "reply": '{"say":"報告","actions":[],"done":true,"report":"C2 と C5 が全角"}'},
    ]) + "\n", encoding='utf-8')
    history, meta, last = va._load_resume(str(log))
    assert len(history) == 4 and history[0][0] == 'user' and history[1][0] == 'model'
    assert history[3][1].startswith('{"say":"報告"') and meta['sheet'] == "会員" and meta['book'] == "名簿.xlsx"
    assert last is None                                   # 最後の往復に results が無い（報告だけで終わった）
    _, _, last2 = va._load_resume(str(log))
    assert last2 is None
    # 途中で終わった記録なら、最後の結果も引き継ぐ（AI が自分の結果を見てから続けられる）
    log2 = tmp_path / "log2.jsonl"
    log2.write_text(json.dumps({"turn": 1, "prompt": "p", "reply": "r",
                                "results": [{"label": "tidy A1:D9", "ok": False, "out": "エラー: x"}]},
                               ensure_ascii=False) + "\n", encoding='utf-8')
    h2, m2, l2 = va._load_resume(str(log2))
    assert m2 is None and l2 == [("tidy A1:D9", False, "エラー: x")]
    assert va._load_resume(str(tmp_path / "no.jsonl")) == (None, None, None)
    (tmp_path / "empty.jsonl").write_text("", encoding='utf-8')
    assert va._load_resume(str(tmp_path / "empty.jsonl")) == (None, None, None)
    # 続きの 1 通目は「追加の指示」と「いまの材料」だけ（規則と前回の会話は history にある）
    p = va._continue_prompt("置き換えてよい", "使用範囲: A1:D9", l2)
    assert "【追加の指示】" in p and "置き換えてよい" in p and "【いまの材料】" in p and "使用範囲: A1:D9" in p
    assert "tidy A1:D9 （失敗）" in p and va.RULES not in p
    assert "【追加の指示】" in va._continue_prompt("x", "y")
    # CLI: --continue は補足を後ろに置ける（--undo・--fire と同じ入口）
    parser = vm.build_parser()
    ns, unknown = parser.parse_known_args(["agent", "--continue", "置き換えてよい"])
    assert ns.cont and ns.posargs == ["置き換えてよい"] and not unknown


def test_agent_verify_counts_only_new_errors(monkeypatch):
    # 終わりの検査は「始める前より増えたか」で見る（元から #N/A のあるシートで done でも要確認になっていた・2026-09-04）
    import vbam_agent as va
    assert va._count_err_hash("") == (0, 0)
    assert va._count_err_hash("エラーセル: 200+個  A1=#N/A …\n'###' で読めないセル: 3個（列幅不足）") == (200, 3)
    after = "使用範囲: (空)\nエラーセル: 2個  A1=#N/A A2=#N/A\n"
    monkeypatch.setattr(va, "_run_cmd", lambda toks, wb=None: (True, after))
    assert va._verify_sheet("s", None, base=(2, 0)) is True          # 元から 2 個＝増えていない
    assert va._verify_sheet("s", None, base=(3, 1)) is True          # 減ったのも合格
    assert va._verify_sheet("s", None, base=(1, 0)) is False         # 1 個増えた
    assert va._verify_sheet("s", None) is False                      # 既定＝元は 0


def test_agent_ai_setup_and_fire_cases(tmp_path, monkeypatch):
    import os, pytest
    import vbam_agent as va
    # 環境変数も金庫も無い状態＝キーが無いと言って止まる（金庫は tmp に向ける＝本物を見ない）
    monkeypatch.setattr(vk, "_KEY_STORE", str(tmp_path / "keys.json"))
    saved = os.environ.pop("GEMINI_API_KEY", None)
    try:
        with pytest.raises(RuntimeError, match="キーがありません"):
            va._ai_setup("gemini", None)
    finally:
        if saved is not None:
            os.environ["GEMINI_API_KEY"] = saved
    with pytest.raises(ValueError):
        va._ai_setup("gpt", None)
    # 実射の弾は sheet 3 本＋重い道具 4 本（2026-09-04）＋ピボットの難しい弾 6 本（同日深夜）、どれも依頼文と答え合わせ
    # （現物を見る関数）を持つ。練習台は plan_sheet を通る（難しい弾は build で明細 45 行）
    assert [c["name"] for c in vf.FIRE_CASES] == ["郵便番号の統一", "会員番号で突き合わせ", "金額の合計",
                                                  "テーブル化", "ピボット集計", "スライサー", "パワークエリで読み込み",
                                                  "月別×担当の構成比", "売上上位3地域", "前月比", "スライサー2枚連動",
                                                  "ピボットグラフ", "データモデルの平均単価",
                                                  # 2026-09-09: 既存のピボットを「直す」12 本。それまでの
                                                  # ピボットの弾 8 本は全部「白紙から作る」で、pivot_field /
                                                  # pivot_calc のほとんどの手は一度も撃たれていなかった
                                                  "ピボットに列を足す", "ピボットを平均に変える",
                                                  "ピボットを月別に組み替える", "ピボットの小計と総計を消す",
                                                  "ピボットを表形式にする", "ピボットに税込を足す",
                                                  "ピボットの元範囲を広げる", "ピボットを担当で絞る",
                                                  "ピボットを数量の区分でまとめる", "ピボットの空欄と見た目",
                                                  "ピボットのフィールドを外して並べ替える", "ピボットと言わずに集計",
                                                  "300行の掃除", "300行の突き合わせ",
                                                  "見やすくして", "未入力が分かるように", "合計を出して",
                                                  "多い順に並べて", "担当ごとの件数",
                                                  "部署の列を足す",          # 2026-09-06 夜: 手順書「列を足す」を畳んで弾だけ残した
                                                  "印刷したら切れる", "鈴木さんの分だけ",
                                                  "金額が足せない", "同じ人が二重",
                                                  # 2026-09-06: Claude for Excel の失敗型 4 本＋止まるべき弾 3 本
                                                  "3行目が見出し", "3行目が見出し・列を足す",
                                                  "行を足すと合計が漏れる", "文字の数字を数値に",
                                                  # 2026-09-06 夜: 失敗型の残り 2 本＋image の手の弾
                                                  "月別集計の式が横にコピーできる", "前提シートの参照を伸ばす", "写真を入れる",
                                                  "写真のファイルが無い", "人の表がある所へ書けと言われる",
                                                  "見るだけの依頼"]
    assert all(c["request"] and callable(c["check"]) for c in vf.FIRE_CASES)
    # 止まるべき弾は、現物が無事なだけでは合格にしない（report で言えたかも見る＝黙殺を不合格にする）
    weak = [c for c in vf.FIRE_CASES if c.get("kind") == "weak"]
    assert len(weak) == 10 and sum(1 for c in weak if callable(c.get("check_run"))) == 3
    assert "weak" in vf._FIRE_GROUPS
    import vbam_build as vb
    p = vb.plan_sheet(dict(vf._FIRE_BASE, sheet="実射1"))
    assert p['header_row'] == 3 and p['first'] == 4 and p['last'] == 9 and p['full_addr'] == 'A3:E9'


# ---- 手順書（2026-09-04・20 本を道具の手で回る形に書き直し、練習台と答え合わせを付けた。
#      同日夕、人の言葉の側から 5 本足して 25 本。2026-09-06 夜、Claude for Excel の査読で 20 本に切り直し。
#      2026-09-09、read_file（別のファイルを読む手）と一緒に 2 本足して 22 本）----

def test_recipes_are_22_and_well_formed():
    import vbam_recipes as vr
    import pytest
    assert len(vr.SHEET_RECIPES) == 22
    assert list(vr.SHEET_RECIPES) == ["値を直す", "表と数式の点検", "月次の集計表", "重複チェック", "説明書きを作る",
                                      "グラフ付き報告", "検索と抽出", "印刷とPDF", "図形と画像の整理", "見た目を直す",
                                      "リンクと外部参照の解消", "名前付き範囲の整理", "重いブックを軽くする",
                                      "住所と郵便番号の整形", "2つの表の突合", "他システムに渡すCSVを作る",
                                      "帳票を一覧表に直す", "並べ替えと絞り込み", "色分け", "入力ミスを防ぐ",
                                      "別ファイルとの突合", "コード対応表で振り直し"]
    for name, body in vr.SHEET_RECIPES.items():
        for key in ("いつ: ", "ゴール: ", "手順:", "確かめる: ", "承認: ", "報告: "):
            assert key in body, f"{name}: {key} が無い"
        assert "完了条件" in body, f"{name}: 完了条件が無い"
        assert vr.recipe_when(name)
        # 古い名前・無い手（inspect・row/col の action・「道具に無い」の嘘）が本文に残っていない
        for stale in ("inspect", '"op":"row"', "行の削除は道具に無い", "行や列を消す手は道具に無い",
                      "結合セルのレイアウト表を一覧表に直す", "データの掃除", "数式の監査", "PDFにして配る"):
            assert stale not in body, f"{name}: 古い記述「{stale}」が残っている"
    req = vr.compose_recipe_request("値を直す", "置き換えてよい")
    assert req.startswith("手順書「値を直す」") and req.rstrip().endswith("補足: 置き換えてよい")
    assert vr.compose_recipe_request("表と数式の点検").rstrip().endswith("補足: なし")
    with pytest.raises(ValueError):
        vr.compose_recipe_request("無い手順書")
    # 練習台と答え合わせ: 統合した手順書は練習台を複数持つ（variant）。名前の並びは手順書の並びと同じ
    names = [c["name"] for c in vr.RECIPE_CASES]
    assert list(dict.fromkeys(names)) == list(vr.SHEET_RECIPES)
    assert len(vr.RECIPE_CASES) == 26 and names.count("値を直す") == 3 and names.count("表と数式の点検") == 2
    assert names.count("印刷とPDF") == 2
    assert all(callable(c["build"]) and callable(c["check"]) and c["kind"] == "recipe" for c in vr.RECIPE_CASES)
    # 練習台と答え合わせは弾ごとに別物（2026-09-09: 後から足した手順書が _b25/_k25 という名前を
    # 先客（印刷とPDF の PDF 側）と取り合い、あとの def が前の def を黙って上書きした。
    # 全弾で「印刷とPDF（PDF）」の弾に突合の答え合わせが当たり、そこで初めて分かった）
    for key in ("build", "check"):
        fns = [c[key] for c in vr.RECIPE_CASES]
        assert len(set(fns)) == len(fns), (
            f"{key} を 2 つ以上の弾が共有しています（同じ名前で定義し直していないか）: "
            + " / ".join(sorted({f.__name__ for f in fns if fns.count(f) > 1})))
    assert [c["no"] for c in vr.RECIPE_CASES] == list(range(1, 27))
    assert vr.recipe_case_label(vr.RECIPE_CASES[0]) == "手順書「値を直す」（表記の乱れ）"
    assert vr.recipe_case_label(vr.RECIPE_CASES[5]) == "手順書「月次の集計表」"
    # 承認の言葉は補足だけで見る（本文の「置き換えてよい」は説明＝承認ではない）
    import vbam_agent as va
    assert va._has_approval(va._approval_scope(vr.compose_recipe_request("値を直す", ""))) is False
    assert va._has_approval(va._approval_scope(vr.compose_recipe_request("値を直す", "置き換えてよい"))) is True


def test_recipe_option_and_fire_groups():
    import vbam_agent as va
    p = vm.build_parser()
    ns, unknown = p.parse_known_args(["agent", "--recipe", "値を直す", "置き換えてよい"])
    assert ns.recipe == "値を直す" and ns.posargs == ["置き換えてよい"] and not unknown
    ns = p.parse_args(["agent", "--fire", "recipe", "表と数式の点検"])
    assert ns.fire and ns.posargs == ["recipe", "表と数式の点検"]
    cases = [{"name": "a", "kind": "sheet"}, {"name": "b", "kind": "macro"}, {"name": "c", "kind": "recipe"}]
    picked, unknown = vf._select_fire_cases(cases, None)
    assert [c["name"] for c in picked] == ["a", "b", "c"] and not unknown
    # 組の名前と弾の名前を混ぜたら弾の名前だけ（2026-09-04。組が効くと通っている弾まで撃ち直していた）
    picked, unknown = vf._select_fire_cases(cases, ["recipe", "a"])
    assert [c["name"] for c in picked] == ["a"] and not unknown
    picked, unknown = vf._select_fire_cases(cases, ["recipe"])
    assert [c["name"] for c in picked] == ["c"] and not unknown
    picked, unknown = vf._select_fire_cases(cases, ["x"])
    assert unknown == ["x"]
    assert va._RECIPE_MAX_TURNS >= va._DEFAULT_MAX_TURNS


def test_agent_new_hands_tokens_and_guards(tmp_path):
    import vbam_agent as va
    import pytest
    import datetime
    # read は他のシートを "sheet" で読める（書く手は対象シートだけ）
    label, toks = va._action_to_tokens({"op": "read", "range": "A1:C3", "sheet": "基準"}, "名簿")
    assert toks == ["read-range", "--sheet", "基準", "A1:C3"] and label == "read 基準!A1:C3"
    assert set(va._DIRECT_OPS) <= set(va._ALLOWED_OPS)
    for op in va._DIRECT_OPS:
        if op in va._ALIAS_OPS:              # 古い名前（inspect）は受けるだけで規則文には出さない（2026-09-06 夜）
            continue
        assert '"op":"%s"' % op in va.RULES
    # 直接の手はブックが無いと実行しない（理由が結果になる）。許していない手も同じ
    res = va._execute([{"op": "view", "freeze": "A2"}, {"op": "export_csv", "range": "A1", "path": "x.txt"},
                       {"op": "delete_shape", "name": "x"}], "名簿")
    assert all(r[1] is False and "実行していません" in r[2] for r in res)
    # CSV の部品（純 Python）
    assert va._date_fmt("yyyy-mm-dd") == "%Y-%m-%d" and va._date_fmt("yyyy/mm/dd") == "%Y/%m/%d"
    assert va._date_fmt("%d.%m.%Y") == "%d.%m.%Y"
    assert va._csv_text(datetime.datetime(2026, 1, 5), "%Y-%m-%d") == "2026-01-05"
    assert va._csv_text(12000.0, "%Y-%m-%d") == "12000" and va._csv_text("0001", "%Y") == "0001"
    assert va._csv_text(None, "%Y") == "" and va._csv_text(True, "%Y") == "TRUE" and va._csv_text(1.5, "%Y") == "1.5"
    with pytest.raises(ValueError):
        va._csv_check_path(str(tmp_path / "a.txt"))
    existing = tmp_path / "b.csv"
    existing.write_text("x")
    with pytest.raises(ValueError):
        va._csv_check_path(str(existing))
    assert va._csv_check_path(str(tmp_path / "c.csv")).endswith("c.csv")
    assert va._rows_of(((1, 2), (3, 4))) == [[1, 2], [3, 4]] and va._rows_of(5) == [[5]]


def test_agent_shape_delete_and_button_guard():
    # 2026-09-04: 図形を消す手。手順書「図形と画像の整理」が回るように足した。
    # 消せるのは図形だけ／マクロ付き（ボタン）は拒む／控えを取る手なので agent --undo で戻る。
    import pytest
    import vbam_agent as va

    class FakeShape:
        def __init__(self, name, onaction=""):
            self.Name, self.OnAction = name, onaction
            self.Left = self.Top = self.Width = self.Height = 10.0
            self.deleted = False

        def Delete(self):
            self.deleted = True

    class FakeShapes:
        def __init__(self, shapes):
            self._s = list(shapes)

        def __call__(self, key):
            for s in self._s:
                if s.Name == key and not s.deleted:
                    return s
            raise KeyError(key)

        def __iter__(self):
            return iter([s for s in self._s if not s.deleted])

        @property
        def Count(self):
            return len([s for s in self._s if not s.deleted])

    class FakeWS:
        def __init__(self, shapes):
            self.Shapes = FakeShapes(shapes)

    class FakeWB:
        def __init__(self, ws):
            self._ws = ws

        def Sheets(self, name):
            return self._ws

    shapes = [FakeShape("重なった要らない図"), FakeShape("はみ出した矢印"),
              FakeShape("ロゴ枠"), FakeShape("btn印刷", "印刷する")]
    wb = FakeWB(FakeWS(shapes))

    # 複数まとめて消せる（残りの数と戻し方を結果に出す）
    _, ok, out = va._direct_action(
        {"op": "shape", "names": ["重なった要らない図", "はみ出した矢印"], "delete": True}, "テスト用5", wb)
    assert ok and shapes[0].deleted and shapes[1].deleted and not shapes[2].deleted
    assert "2 個" in out and "--undo" in out and "残っている図形: 2 個" in out

    # マクロ付き＝ボタンは消さない・動かさない（消してから気づくのでは遅い＝当てる前に断る）
    with pytest.raises(ValueError, match="マクロ"):
        va._direct_action({"op": "shape", "name": "btn印刷", "delete": True}, "テスト用5", wb)
    with pytest.raises(ValueError, match="マクロ"):
        va._direct_action({"op": "shape", "name": "btn印刷", "left": 5}, "テスト用5", wb)
    assert not shapes[3].deleted
    # 1 つでもボタンが混じっていたら、その並びは 1 つも消さない
    with pytest.raises(ValueError, match="マクロ"):
        va._direct_action({"op": "shape", "names": ["ロゴ枠", "btn印刷"], "delete": True}, "テスト用5", wb)
    assert not shapes[2].deleted

    # 無い図形／複数を動かそうとした／当てる項目なし は実行せずに断る
    for bad in ({"op": "shape", "name": "無い図", "delete": True},
                {"op": "shape", "names": ["ロゴ枠", "無い図"], "delete": True},
                {"op": "shape", "names": ["ロゴ枠", "はみ出した矢印"], "left": 5},
                {"op": "shape", "names": "ロゴ枠", "delete": True},
                {"op": "shape", "delete": True},
                {"op": "shape", "name": "ロゴ枠"}):
        with pytest.raises(ValueError):
            va._direct_action(bad, "テスト用5", wb)

    # 位置と大きさは今までどおり
    _, ok, out = va._direct_action(
        {"op": "shape", "name": "ロゴ枠", "left": 50, "top": 60}, "テスト用5", wb)
    assert ok and shapes[2].Left == 50 and shapes[2].Top == 60 and "l=50" in out

    # 消す手は「シートを変える手」＝控えを取る（取らないと --undo で戻せない）
    assert va._will_write([{"op": "shape", "names": ["x"], "delete": True}]) is True
    assert '"delete":true' in va.RULES.replace(" ", "")


# ================================================================
# APIキーの預かり（2026-09-04）: set-key / clear-key
# ================================================================
import io                                                            # noqa: E402
import vbam_agent as va                                              # noqa: E402


def test_mask_key_shows_only_both_ends():
    assert va._mask_key("AIzaSyABCDEFGHIJKLMN1234x7Qd").startswith("AIza")
    assert va._mask_key("AIzaSyABCDEFGHIJKLMN1234x7Qd").endswith("x7Qd")
    assert "SyABCDEFGHIJKLMN" not in va._mask_key("AIzaSyABCDEFGHIJKLMN1234x7Qd")
    assert va._mask_key(None) == "（設定されていません）"
    assert va._mask_key("abc") == "***"                  # 短いものは全部伏せる


def test_key_problem_catches_paste_accidents():
    assert va._key_problem("") is not None                # 何も入っていない
    assert va._key_problem(" AIzaKey1234 ") is not None   # 前後の空白
    assert va._key_problem("AIza Key1234") is not None    # 途中で折れた貼り付け
    assert va._key_problem("AIzaKey\n1234") is not None   # 改行が混ざった
    assert va._key_problem("ＡＩｚａキー") is not None      # 全角で貼られた
    assert va._key_problem("x" * (va._KEY_MAX + 1)) is not None
    assert va._key_problem("AIzaSyABCDEFGHIJKLMN1234x7Qd") is None


def test_user_env_write_then_delete_removes_the_entry():
    """setx NAME "" は空の値を残す＝消えない。SetEnvironmentVariable(null) は項目ごと消える。"""
    name = "_VBAM_TEST_KEY"
    try:
        va._user_env_write(name, "dummy-value-1234")
        assert va._user_env_get(name) == "dummy-value-1234"
        va._user_env_write(name, None)
        assert va._user_env_get(name) is None             # 空文字ではなく、無いこと
    finally:
        try:
            va._user_env_write(name, None)
        except Exception:
            pass


def test_set_key_and_clear_key_are_wired():
    p = vm.build_parser()
    assert p.parse_args(["set-key", "gemini"]).posargs == ["gemini"]
    assert p.parse_args(["set-key", "--no-ping"]).no_ping is True
    assert p.parse_args(["clear-key", "claude", "-y"]).yes is True
    cmds = vm._command_table()
    assert vm.raw_command(cmds["set-key"]) is va.cmd_set_key
    assert vm.raw_command(cmds["clear-key"]) is va.cmd_clear_key


def test_set_key_refuses_a_broken_paste_without_writing(monkeypatch):
    """貼り付け事故は書く前に止める（環境変数に触らない）。"""
    wrote = []
    monkeypatch.setattr(vk, "_user_env_write", lambda n, v: wrote.append((n, v)))
    monkeypatch.setattr(sys, "stdin", io.StringIO("AIza Key1234\n"))
    args = _ap.Namespace(posargs=["gemini"], stdin=True, no_ping=True, model=None)
    assert va.cmd_set_key(args) is False
    assert wrote == []


def _tmp_store(monkeypatch, tmp_path):
    """金庫を tmp に向ける（本物の金庫を汚さない）。"""
    monkeypatch.setattr(vk, "_KEY_STORE", str(tmp_path / "keys.json"))


def test_set_key_puts_it_in_the_vault_not_in_the_environment(monkeypatch, tmp_path):
    """環境変数には書かない（平文・期限なしで残るため）。金庫に入れ、いまのプロセスにだけ渡す。"""
    _tmp_store(monkeypatch, tmp_path)
    wrote = []
    monkeypatch.setattr(vk, "_user_env_write", lambda n, v: wrote.append((n, v)))
    monkeypatch.setattr(vk, "_user_env_get", lambda n: None)
    monkeypatch.setattr(sys, "stdin", io.StringIO("AIzaSyABCDEFGHIJKLMN1234x7Qd\n"))
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    args = _ap.Namespace(posargs=["gemini"], stdin=True, no_ping=True, model=None, days=None)
    assert va.cmd_set_key(args) is True
    assert wrote == []                                                   # 環境変数は触らない
    assert va._key_load("gemini") == "AIzaSyABCDEFGHIJKLMN1234x7Qd"      # 金庫から戻る
    assert os.environ["GEMINI_API_KEY"] == "AIzaSyABCDEFGHIJKLMN1234x7Qd"  # このプロセスだけ
    body = (tmp_path / "keys.json").read_text(encoding="utf-8")
    assert "AIzaSyABCDEFGHIJKLMN1234x7Qd" not in body                    # 平文では置かない


def test_vault_forgets_by_itself_when_the_time_is_up(monkeypatch, tmp_path):
    """期限が過ぎたキーは読めない＝その場で消える（人が消し忘れても残らない）。"""
    _tmp_store(monkeypatch, tmp_path)
    va._key_save("gemini", "AIzaSyABCDEFGHIJKLMN1234x7Qd", days=30)
    assert va._key_load("gemini") == "AIzaSyABCDEFGHIJKLMN1234x7Qd"
    import datetime
    d = va._store_read()
    d["gemini"]["expires"] = ((datetime.datetime.now() - datetime.timedelta(minutes=1))
                              .strftime('%Y-%m-%d %H:%M'))
    va._store_write(d)
    assert va._key_info("gemini")[0] is False and "期限切れ" in va._key_info("gemini")[1]
    assert va._key_load("gemini") is None
    assert va._store_read() == {}                       # 読もうとした時点で消えている


def test_clear_key_removes_both_the_vault_and_the_old_environment(monkeypatch, tmp_path, capsys):
    _tmp_store(monkeypatch, tmp_path)
    va._key_save("gemini", "AIzaSyABCDEFGHIJKLMN1234x7Qd", days=30)
    monkeypatch.setattr(vk, "_user_env_get", lambda n: None)             # 環境変数には無い
    args = _ap.Namespace(posargs=["gemini"], yes=True)
    assert va.cmd_clear_key(args) is True
    assert va._store_read() == {}
    assert "金庫から消しました" in capsys.readouterr().out


def test_clear_key_fails_when_the_old_environment_value_survives(monkeypatch, tmp_path, capsys):
    """消したあと読み直して、残っていれば失敗として返す（消したつもりを許さない）。"""
    _tmp_store(monkeypatch, tmp_path)
    monkeypatch.setattr(vk, "_user_env_get", lambda n: "AIzaSyABCDEFGHIJKLMN1234x7Qd")
    monkeypatch.setattr(vk, "_user_env_write", lambda n, v: None)          # 消したふり
    args = _ap.Namespace(posargs=["gemini"], yes=True)
    assert va.cmd_clear_key(args) is False
    assert "消せていません" in capsys.readouterr().out


def test_agent_heavy_hands_tokens_and_guards():
    # 2026-09-04: 重い道具 7 本（table / pivot / pivot_field / pivot_calc / slicer / powerquery / datamodel）を AI の手に配線
    import vbam_agent as va
    import re
    import pytest
    assert set(va._HEAVY_OPS) <= set(va._ALLOWED_OPS) and not (set(va._HEAVY_OPS) & set(va._DIRECT_OPS))
    assert set(va._HEAVY_ACTIONS) == set(va._HEAVY_OPS)
    m = re.search(r"この (\d+) 個だけ", va.RULES)
    # 古い名前（inspect・row・col）は受けるだけで規則文には出さない＝数に入れない（2026-09-06 夜）
    assert m and int(m.group(1)) == len(va._ALLOWED_OPS) - len(va._ALIAS_OPS)
    assert set(va._ALIAS_OPS) <= set(va._ALLOWED_OPS)
    for op in va._HEAVY_OPS:
        assert '"op":"%s"' % op in va.RULES
    # 消す手はどの op にも無い
    for op, acts in va._HEAVY_ACTIONS.items():
        assert not any('delete' in a or a == 'column_remove' for a in acts), op
    parser = vm.build_parser()
    samples = [
        ({"op": "table", "action": "create", "range": "A3:E9", "name": "T名簿"},
         ["table", "create", "'名簿'!A3:E9", "T名簿"]),
        ({"op": "table", "action": "sort", "name": "T名簿", "column": "金額", "desc": True},
         ["table", "sort", "T名簿", "金額", "--desc"]),
        ({"op": "table", "action": "filter_values", "name": "T名簿", "column": "コース", "values": ["年間", "半年"]},
         ["table", "filter-values", "T名簿", "コース", "年間", "半年"]),
        ({"op": "table", "action": "column_add", "name": "T名簿", "column": "備考", "at": 3},
         ["table", "column", "add", "T名簿", "備考", "--at", "3"]),
        ({"op": "pivot", "action": "create", "range": "I5:K10", "rows": ["申込コース"], "values": "金額", "func": "sum",
          "at": "M5", "name": "P申込"},
         ["pivot", "create", "'名簿'!I5:K10", "--rows", "申込コース", "--values", "金額", "--func", "sum", "--at", "M5",
          "--name", "P申込"]),
        ({"op": "pivot", "action": "create", "range": "I5:K10", "cols": "申込コース", "values": "金額", "sheet": "集計"},
         ["pivot", "create", "'名簿'!I5:K10", "--cols", "申込コース", "--values", "金額", "--sheet", "集計"]),
        ({"op": "pivot_field", "action": "add_value", "pivot": "P申込", "field": "金額", "func": "count", "name": "件数"},
         ["pivot-field", "add-value", "P申込", "金額", "--func", "count", "--name", "件数"]),
        ({"op": "pivot_field", "action": "list", "pivot": "P申込"}, ["pivot-field", "list", "P申込"]),
        ({"op": "pivot_field", "action": "set_filter", "pivot": "P申込", "field": "申込コース", "values": ["年間"]},
         ["pivot-field", "set-filter", "P申込", "申込コース", "年間"]),
        ({"op": "pivot_field", "action": "group_numeric", "pivot": "P申込", "field": "金額", "start": 0, "end": 20000,
          "interval": 5000},
         ["pivot-field", "group-numeric", "P申込", "金額", "0", "20000", "5000"]),
        ({"op": "pivot_calc", "action": "get_data", "pivot": "P申込"}, ["pivot-calc", "get-data", "P申込"]),
        ({"op": "pivot_calc", "action": "calc_field_create", "pivot": "P申込", "name": "税込", "formula": "=金額*1.1"},
         ["pivot-calc", "calc-field", "create", "P申込", "税込", "=金額*1.1"]),
        ({"op": "pivot_calc", "action": "grand_totals", "pivot": "P申込", "which": "rows", "on": False},
         ["pivot-calc", "grand-totals", "P申込", "rows", "off"]),
        ({"op": "slicer", "action": "add", "source": "T申込", "field": "申込コース", "at": "M12"},
         ["slicer", "add", "T申込", "申込コース", "--at", "M12"]),
        ({"op": "powerquery", "action": "add", "name": "Q名簿", "m": 'let s = Excel.CurrentWorkbook(){[Name="T名簿"]}[Content] in s'},
         ["powerquery", "add", "Q名簿", "--m", 'let s = Excel.CurrentWorkbook(){[Name="T名簿"]}[Content] in s']),
        ({"op": "powerquery", "action": "load", "name": "Q名簿", "to": "sheet", "at": "M3"},
         ["powerquery", "load", "Q名簿", "--to", "sheet", "--sheet", "名簿", "--at", "M3"]),
        ({"op": "powerquery", "action": "refresh"}, ["powerquery", "refresh"]),
        ({"op": "datamodel", "action": "relation_add", "fk_table": "T売上", "fk_column": "商品", "pk_table": "T商品",
          "pk_column": "商品"},
         ["datamodel", "relation", "add", "T売上", "商品", "T商品", "商品"]),
        ({"op": "datamodel", "action": "measure_add", "table": "T売上", "name": "売上合計", "dax": "SUM(T売上[金額])",
          "format": "whole", "thousands": True},
         ["datamodel", "measure", "add", "T売上", "売上合計", "--dax", "SUM(T売上[金額])", "--format", "whole", "--thousands"]),
        # 2026-09-04 深夜: ピボットの難しい手（月別・構成比・前月比・上位 N・データモデル・ピボットグラフ）
        ({"op": "pivot", "action": "create", "range": "A1:F46", "rows": "日付", "cols": "担当", "filter": "地域",
          "values": "売上,売上", "sheet": "月別", "at": "B2"},
         ["pivot", "create", "'名簿'!A1:F46", "--rows", "日付", "--cols", "担当", "--filter", "地域", "--values", "売上,売上",
          "--sheet", "月別", "--at", "B2"]),
        ({"op": "pivot", "action": "create", "model": True, "rows": "T売上.地域", "measures": "平均単価", "sheet": "単価"},
         ["pivot", "create", "model", "--model", "--rows", "T売上.地域", "--measures", "平均単価", "--sheet", "単価"]),
        ({"op": "pivot_field", "action": "show_as", "pivot": "P", "field": "sum/売上2", "kind": "percent_total"},
         ["pivot-field", "show-as", "P", "sum/売上2", "percent_total"]),
        ({"op": "pivot_field", "action": "show_as", "pivot": "P", "field": "x", "kind": "diff-from", "base_field": "日付",
          "base_item": "(previous)"},
         ["pivot-field", "show-as", "P", "x", "diff_from", "--base-field", "日付", "--base-item", "(previous)"]),
        ({"op": "pivot_field", "action": "top_n", "pivot": "P", "field": "地域", "n": 3, "by": "sum/売上"},
         ["pivot-field", "top-n", "P", "地域", "3", "--by", "sum/売上"]),
        ({"op": "pivot_field", "action": "top_n", "pivot": "P", "field": "地域", "n": "2", "bottom": True},
         ["pivot-field", "top-n", "P", "地域", "2", "--bottom"]),
        ({"op": "pivot_field", "action": "filter_clear", "pivot": "P", "field": "地域"},
         ["pivot-field", "filter-clear", "P", "地域"]),
        ({"op": "pivot_field", "action": "sort", "pivot": "P", "field": "地域", "order": "desc", "by": "sum/売上"},
         ["pivot-field", "sort", "P", "地域", "desc", "--by", "sum/売上"]),
        ({"op": "pivot_field", "action": "position", "pivot": "P", "field": "商品", "n": 1},
         ["pivot-field", "position", "P", "商品", "1"]),
        ({"op": "pivot_calc", "action": "refresh", "pivot": "P"}, ["pivot-calc", "refresh", "P"]),
        ({"op": "pivot_calc", "action": "style", "pivot": "P", "style": "PivotStyleMedium9"},
         ["pivot-calc", "style", "P", "PivotStyleMedium9"]),
        ({"op": "pivot_calc", "action": "repeat_labels", "pivot": "P", "on": False},
         ["pivot-calc", "repeat-labels", "P", "off"]),
        ({"op": "pivot_calc", "action": "empty_as", "pivot": "P", "text": "0"}, ["pivot-calc", "empty-as", "P", "0"]),
        ({"op": "pivot_calc", "action": "set_source", "pivot": "P", "range": "A1:F60"},
         ["pivot-calc", "set-source", "P", "'名簿'!A1:F60"]),
    ]
    for act, want in samples:
        label, toks = va._action_to_tokens(act, "名簿")
        assert toks == want, (act, toks)
        assert label.startswith(act["op"] + " " + act["action"])
        # 写した引数列は CLI の parser がそのまま受ける（不明な引数ゼロ）
        ns, unknown = parser.parse_known_args(toks)
        assert not unknown and ns.command == toks[0], (toks, unknown)
    bads = [
        {"op": "table", "action": "delete", "name": "T名簿"},
        {"op": "pivot", "action": "delete", "name": "P申込"},
        {"op": "pivot", "action": "create", "range": "I5:K10", "values": "金額"},          # rows も cols も無い
        {"op": "pivot", "action": "create", "range": "I5:K10", "rows": "申込コース"},       # values が無い
        {"op": "pivot", "action": "create", "range": "申込!I5:K10", "rows": "a", "values": "b"},   # 番地にシート名
        {"op": "pivot_field", "action": "add_value", "pivot": "P申込", "field": "金額", "func": "median"},
        {"op": "pivot_field", "action": "group_date", "pivot": "P申込", "field": "日付", "by": "weeks"},
        {"op": "pivot_calc", "action": "layout", "pivot": "P申込", "layout": "flat"},
        {"op": "slicer", "action": "delete", "name": "x"},
        {"op": "powerquery", "action": "delete", "name": "Q"},
        {"op": "powerquery", "action": "add", "name": "Q"},                                   # M 式が無い
        {"op": "powerquery", "action": "load", "name": "Q", "to": "disk"},
        {"op": "datamodel", "action": "measure_delete", "name": "x"},
        {"op": "datamodel", "action": "measure_add", "table": "T", "name": "m", "dax": "1", "format": "money"},
        {"op": "pivot", "action": "create", "range": "A1:F46", "rows": "地域", "measures": "平均単価"},   # model 無しの measures
        {"op": "pivot", "action": "create", "model": True, "rows": "T売上.地域"},                      # 値もメジャーも無い
        {"op": "pivot_field", "action": "show_as", "pivot": "P", "field": "x", "kind": "percent_of_moon"},
        {"op": "pivot_field", "action": "top_n", "pivot": "P", "field": "地域", "n": 0},
        {"op": "pivot_field", "action": "position", "pivot": "P", "field": "地域"},
        {"op": "pivot_calc", "action": "style", "pivot": "P"},                                         # スタイル名が無い
        {"op": "pivot_calc", "action": "set_source", "pivot": "P", "range": "売上!A1:F60"},            # 番地にシート名
    ]
    for bad in bads:
        with pytest.raises(ValueError):
            va._action_to_tokens(bad, "名簿")
    # 「計算の種類」の名前は CLI 側（vbam_heavy._XL_SHOW_AS）と同じ
    import vbam_heavy as vh
    assert set(va._SHOW_AS_KINDS) == set(vh._XL_SHOW_AS)
    assert set(vh._SHOW_AS_BASE_ITEM) <= set(vh._SHOW_AS_BASE_FIELD) <= set(vh._XL_SHOW_AS)
    assert vh._cube_name("T売上.地域") == "[T売上].[地域]" and vh._cube_name("平均単価") == "[Measures].[平均単価]"
    assert vh._cube_name("[T].[c]") == "[T].[c]" and vh._cube_name("T売上.地域", measure=True) == "[Measures].[T売上.地域]"
    # 難しい弾 6 本が sheet の組にあり、練習台（build）と答え合わせを持つ
    for n in ("月別×担当の構成比", "売上上位3地域", "前月比", "スライサー2枚連動", "ピボットグラフ", "データモデルの平均単価"):
        c = next(c for c in vf.FIRE_CASES if c["name"] == n)
        assert callable(c["check"]) and callable(c["build"]), n
    rows = vf._fire_sales_rows()
    assert len(rows) == 45 and rows[0][0].tzinfo is not None and all(r[5] == r[4] * vf._FIRE_SALES_PRICE[r[3]] for r in rows)
    # RULES に新しい手が載っている
    # M の型付けの段は 2026-09-06 夜に規則文から外し、要る依頼のときだけ材料に足す（_DATAMODEL_NOTE）
    for word in ('show_as', 'top_n', 'group_date', '"model":true', '"pivot":"ピボット名"', 'set_source', 'TransformColumnTypes',
                 '同時に使えない'):
        assert word in va.RULES + va._DATAMODEL_NOTE, word
    # パワークエリの M に型付けを足す部品（モデルの列が全部「文字」になり SUM のメジャーが無言で消える・2026-09-04 実射）
    m = 'let ソース = Excel.CurrentWorkbook(){[Name="T売上"]}[Content] in ソース'
    assert vh._m_source_table(m) == "T売上"
    assert vh._m_source_table('let a = Excel.CurrentWorkbook() {[ Name = "名簿" ]} [Content] in a') == "名簿"
    assert vh._m_source_table('let a = Csv.Document(File.Contents("x.csv")) in a') is None
    typed = vh._m_with_types(m, [("日付", "type date"), ("数量", "Int64.Type"), ("売上", "type number"), ("地域", "type text")])
    assert typed.startswith("let __元 = (let ソース") and typed.endswith("in __型付け")
    assert '{"日付", type date}, {"数量", Int64.Type}, {"売上", type number}, {"地域", type text}' in typed
    assert vh._m_with_types(typed, [("x", "type text")]) == typed       # 既に型付きなら触らない
    assert vh._m_with_types(m, []) == m
    # 実射の弾（重い道具 4 本）が sheet の組に入っている
    names = [c["name"] for c in vf.FIRE_CASES]
    for n in ("テーブル化", "ピボット集計", "スライサー", "パワークエリで読み込み"):
        assert n in names and callable(next(c for c in vf.FIRE_CASES if c["name"] == n)["check"])
    assert callable(va._heavy_materials)


# ================================================================
# 2026-09-04 夕の欠陥直し（33 件）の見張り
# ================================================================

def test_recipe_entry_is_pinned_to_sheet():
    """1: 手順書は本文の言い回しで build／macro に振られていた。入口で sheet に固定する。"""
    import vbam_agent as va
    from vbam_recipes import SHEET_RECIPES, compose_recipe_request
    # 素の振り分けに通すと sheet 以外になる手順書が実際にある（だから固定が要る）
    routed = {n: va._route(compose_recipe_request(n, ""), macro_names=["図形を並べる"]) for n in SHEET_RECIPES}
    assert any(v != 'sheet' for v in routed.values())
    # 入口の判定（_cmd_agent_body と同じ式）: recipe があれば sheet
    for recipe in (None, "2つの表の突合"):
        for resume_hist in (None, ["h"]):
            forced = 'sheet' if (resume_hist or recipe) else None
            assert forced == ('sheet' if (recipe or resume_hist) else None)
    assert va._route("会員名簿を新しいシートに作って", forced='sheet') == 'sheet'


def test_format_choices_and_systemexit_are_caught():
    """2: align:"middle" で argparse が SystemExit(2)＝ループごと死んでいた。"""
    import vbam_agent as va
    import pytest
    # 道具の側で先に断る（AI は次の往復で直せる）
    for bad in ({"op": "format", "range": "A1", "align": "middle"},
                {"op": "format", "range": "A1", "valign": "middle"},
                {"op": "format", "range": "A1", "border": "all"}):
        with pytest.raises(ValueError):
            va._action_to_tokens(bad, "名簿")
    assert "--align" in va._action_to_tokens({"op": "format", "range": "A1", "align": "center"}, "名簿")[1]
    # 語彙は本物のパーサーの choices と同じ
    fmt = vm.build_parser()._subparsers._group_actions[0].choices['format-range']
    real = {a.dest: tuple(a.choices) for a in fmt._actions if a.choices}
    for key, want in va._FORMAT_CHOICES.items():
        assert tuple(real[key]) == tuple(want), key
    # それでも argparse が落ちるときは、失敗した手として AI に返す（ループは続く）
    ok, out = va._run_cmd(['format-range', '--sheet', '名簿', 'A1', '--align', 'middle'])
    assert ok is False and "引数の形が違います" in out


def test_image_is_attached_to_the_newest_turn_only(tmp_path, monkeypatch):
    """3: 古い user 発言にも画像が残り、往復 n で最新の 1 枚を n-1 枚送っていた。"""
    import vbam_agent as va
    png = tmp_path / "v.png"
    png.write_bytes(b"\x89PNG\r\n\x1a\n" + b"x" * 8)
    seen = []

    def fake_ask(ai, model, key, history):
        seen.append([len(h) for h in history])
        return '{"say":"","actions":[{"op":"read","range":"A1"}]}', {'in': 1, 'out': 1}, 0.0

    monkeypatch.setattr(va, "_ask", fake_ask)
    monkeypatch.setattr(va, "_ai_setup", lambda ai, model: ('gemini', 'm', 'k'))
    monkeypatch.setattr(va, "_execute", lambda a, s, wb=None: [("read A1", True, "x")])
    monkeypatch.setattr(va, "_shot_for_ai", lambda s, wb: (str(png), "A1:C5"))
    monkeypatch.setattr(va, "_will_write", lambda actions: True)
    monkeypatch.setattr(va, "_LAST_AGENT_LOG_FILE", str(tmp_path / "log.jsonl"))
    monkeypatch.setattr(va, "_LAST_AGENT_ASK_FILE", str(tmp_path / "ask.txt"))
    wb = type("WB", (), {"Name": "b.xlsx"})()
    va.run_agent("x", "名簿", wb, max_turns=3, materials="使用範囲: A1", verify=False, backup=False)
    # 3 往復目に送った history のうち、画像つき（3 つ組）の user は 1 本だけ
    assert len(seen) == 3 and seen[2].count(3) == 1 and seen[2][-1] == 3


def test_wants_first_image_detects_merged_cells_and_shapes():
    """3: 結合セル・図形の行がある materials だけ True。「未走査」（大きすぎて数えていない）は False。"""
    import vbam_agent as va
    assert va._wants_first_image("結合セル 1件: A1:D1") is True
    assert va._wants_first_image("図形・ボタン: 2個") is True
    assert va._wants_first_image("使用範囲: A1:E9") is False
    assert va._wants_first_image("結合セル: 未走査（500セル・大）") is False
    assert va._wants_first_image("") is False


def test_first_turn_gets_an_image_when_merged_cells_or_shapes_are_present(tmp_path, monkeypatch):
    """3: 結合セル・図形がある材料は、最初の返事の前にも画像を添える（文字だけでは崩れて見えるため・2026-09-05）。"""
    import vbam_agent as va
    png = tmp_path / "v.png"
    png.write_bytes(b"\x89PNG\r\n\x1a\n" + b"x" * 8)
    seen = []

    def fake_ask(ai, model, key, history):
        seen.append(list(history))
        return '{"say":"","actions":[],"done":true,"report":"ok"}', {'in': 1, 'out': 1}, 0.0

    monkeypatch.setattr(va, "_ask", fake_ask)
    monkeypatch.setattr(va, "_ai_setup", lambda ai, model: ('gemini', 'm', 'k'))
    monkeypatch.setattr(va, "_shot_for_ai", lambda s, wb: (str(png), "A1:D8"))
    monkeypatch.setattr(va, "_LAST_AGENT_LOG_FILE", str(tmp_path / "log.jsonl"))
    monkeypatch.setattr(va, "_LAST_AGENT_ASK_FILE", str(tmp_path / "ask.txt"))
    wb = type("WB", (), {"Name": "b.xlsx"})()
    va.run_agent("x", "名簿", wb, max_turns=3, materials="使用範囲: A1:D8\n結合セル 1件: A1:D1",
                 verify=False, backup=False, grade=False)
    assert len(seen) == 1                               # actions が空で done＝1 往復で終わる
    first_user_turn = seen[0][0]                        # 1 往復目に送った history の 1 本目（今回の user 発言）
    assert len(first_user_turn) == 3 and first_user_turn[2] == str(png)
    assert "結合セルの境目" in first_user_turn[1]


def test_continue_writes_the_old_turns_back_and_says_nothing_about_zero(tmp_path):
    """4・5: 2 回目の --continue で RULES が消えていた／1 通目に「残りの往復: 0 回」。"""
    import vbam_agent as va
    import json
    # 5: --continue の 1 通目には「残りの往復」を書かない
    p = va._continue_prompt("承認します", "使用範囲: A1", [("tidy A1", True, "ok")])
    assert "残りの往復" not in p and "【前回の結果】" in p
    assert "残りの往復: 1 回" in va._result_prompt([("tidy A1", True, "ok")], 1)
    # 4: 引き継いだ会話を新しい記録に書き戻す（次の --continue でも RULES が残る）
    log = tmp_path / "log.jsonl"
    resume = [('user', va.RULES + "【依頼】x"), ('model', '{"actions":[]}'),
              ('user', "【結果】…"), ('model', '{"done":true}')]
    with open(log, 'w', encoding='utf-8') as f:
        f.write(json.dumps({'meta': {'book': 'b', 'sheet': 's', 'mode': 'sheet'}}, ensure_ascii=False) + "\n")
        for i in range(0, len(resume) - 1, 2):
            f.write(json.dumps({'turn': i // 2 + 1, 'prompt': resume[i][1], 'reply': resume[i + 1][1],
                                'resumed': True}, ensure_ascii=False) + "\n")
    hist, meta, _last = va._load_resume(str(log))
    assert len(hist) == 4 and va.RULES in hist[0][1] and meta['mode'] == 'sheet'


def test_macro_log_has_meta_so_continue_refuses_it(tmp_path):
    """6: macro の記録に meta が無く、--continue が sheet の続きとして通っていた。"""
    import vbam_agent as va
    import json
    log = tmp_path / "log.jsonl"
    log.write_text("\n".join(json.dumps(d, ensure_ascii=False) for d in [
        {"meta": {"book": "b.xlsm", "sheet": None, "mode": "macro", "macro": "集計"}},
        {"turn": 1, "prompt": "p", "reply": '{"done":true}'}]) + "\n", encoding='utf-8')
    _h, meta, _l = va._load_resume(str(log))
    assert meta['mode'] == 'macro' and meta['mode'] != 'sheet'
    # meta の無い古い記録も続きにしない（None を sheet 扱いにしない）
    old = tmp_path / "old.jsonl"
    old.write_text(json.dumps({"turn": 1, "prompt": "p", "reply": "r"}, ensure_ascii=False) + "\n", encoding='utf-8')
    _h2, meta2, _l2 = va._load_resume(str(old))
    assert (meta2 or {}).get('mode') != 'sheet'


def test_leading_hyphen_values_pass_the_real_parser():
    """8: 「-要確認」のような値が「不明な引数」で落ちていた。"""
    import vbam_agent as va
    parser = vm.build_parser()
    _l, toks = va._action_to_tokens({"op": "find", "text": "-要確認"}, "名簿")
    ns, unknown = parser.parse_known_args(toks)
    assert ns.posargs == ["-要確認"] and not unknown and ns.sheet_opt == "名簿"
    _l, toks = va._action_to_tokens({"op": "comment", "cell": "D8", "text": "-要確認"}, "名簿")
    ns, unknown = parser.parse_known_args(toks)
    assert ns.posargs == ["'名簿'!D8", "-要確認"] and not unknown
    _l, toks = va._action_to_tokens({"op": "write_cells", "cells": {"A1": "-1 以下", "A2": "-x"}}, "名簿")
    ns, unknown = parser.parse_known_args(toks)
    assert ns.posargs == ["A1", "-1 以下", "A2", "-x"] and not unknown
    _l, toks = va._action_to_tokens({"op": "find_replace", "range": "C2:C9", "search": "－", "replace": "-",
                                     "overwrite": True}, "名簿")
    ns, unknown = parser.parse_known_args(toks)
    assert ns.posargs == ["－", "-", "C2:C9"] and not unknown
    # オプションの値は opt=value の形（「値が無い」で SystemExit にならない）
    _l, toks = va._action_to_tokens({"op": "hyperlink", "cell": "A1", "url": "https://x.example",
                                     "text": "-一覧"}, "名簿")
    ns, unknown = parser.parse_known_args(toks)
    assert ns.text == "-一覧" and not unknown
    _l, toks = va._action_to_tokens({"op": "format", "range": "A1", "number_format": "-#,##0"}, "名簿")
    ns, unknown = parser.parse_known_args(toks)
    assert ns.number_format == "-#,##0" and not unknown


def test_target_sheet_cannot_be_renamed_and_addresses_are_bounded():
    """9・25: 対象シートの改名で画像・差分・検査が落ちる／A0・XFE1 が素通しだった。"""
    import vbam_agent as va
    import pytest
    with pytest.raises(ValueError):
        va._action_to_tokens({"op": "sheet_op", "action": "rename", "name": "名簿", "to": "名簿2"}, "名簿")
    with pytest.raises(ValueError):
        va._action_to_tokens({"op": "sheet_op", "action": "copy", "name": "名簿", "to": "名簿"}, "名簿")
    assert va._action_to_tokens({"op": "sheet_op", "action": "rename", "name": "Sheet1", "to": "集計"},
                                "名簿")[1] == ["sheet", "rename", "Sheet1", "集計"]
    for bad in ("A0", "XFE1", "A1048577", "A1:XFE9"):
        with pytest.raises(ValueError):
            va._check_addr(bad)
    assert va._check_addr("xfd1048576") == "XFD1048576"


def test_find_looks_at_the_target_sheet():
    """10: --sheet が非アクティブのとき find がアクティブシートを探していた。"""
    import vbam_agent as va
    _l, toks = va._action_to_tokens({"op": "find", "text": "貸出中"}, "名簿")
    assert toks[:3] == ["find", "--sheet", "名簿"]
    _l, toks = va._action_to_tokens({"op": "find", "text": "貸出中", "book": True}, "名簿")
    assert "--sheet" not in toks and "--book" in toks
    ns, unknown = vm.build_parser().parse_known_args(["find", "--sheet", "名簿", "--", "x"])
    assert ns.sheet_opt == "名簿" and not unknown


def test_lock_sees_another_thread_of_the_same_process(tmp_path, monkeypatch):
    """15: MCP の世代交代（同じ PID の別スレッド）を素通ししていた。"""
    import vbam_agent as va
    import json as _json
    import threading
    import pytest
    lock = tmp_path / "lock.json"
    monkeypatch.setattr(va, '_AGENT_LOCK_FILE', str(lock))
    lock.write_text(_json.dumps({'pid': os.getpid(), 'thread': threading.get_ident() + 1,
                                 'at': 0, 'at_text': '00:00:00'}), encoding='utf-8')
    with pytest.raises(RuntimeError):
        with va._agent_lock():
            pass
    # 自分のスレッドの錠は素通し（入れ子でも止まらない）
    lock.write_text(_json.dumps({'pid': os.getpid(), 'thread': threading.get_ident(),
                                 'at': 0, 'at_text': '00:00:00'}), encoding='utf-8')
    with va._agent_lock():
        assert _json.loads(lock.read_text(encoding='utf-8'))['thread'] == threading.get_ident()


def test_code_replace_defaults_to_the_module_being_repaired():
    """17: module を省くとブック全体が置換の対象になっていた。"""
    import vbam_agent as va
    _l, toks, w = vmac._macro_action_to_tokens({"op": "code_replace", "search": "a", "replace": "b"},
                                             default_module="shu003")
    assert toks == ["code-replace", "-y", "a", "b", "--module", "shu003"] and w is True
    _l, toks, _w = vmac._macro_action_to_tokens({"op": "code_replace", "search": "a", "replace": "b",
                                               "module": "shu005"}, default_module="shu003")
    assert toks[-1] == "shu005"
    # 修理中のモジュールも AI の module も無い: search の行があるモジュールが 1 つならそこ・複数なら断る・
    # 0 なら CLI に任せる（一律必須は戻した・2026-09-04 夜）
    class _CM:
        def __init__(self, text): self.t = text
        @property
        def CountOfLines(self): return len(self.t.splitlines())
        def Lines(self, a, n): return self.t
    class _Comp:
        def __init__(self, name, text): self.Name, self.CodeModule = name, _CM(text)
    class _WB:
        def __init__(self, comps): self.VBProject = type("P", (), {"VBComponents": comps})()
    wb1 = _WB([_Comp("m1", "Sub x()" + chr(10) + " a = 1" + chr(10) + "End Sub"), _Comp("m2", "Sub y()" + chr(10) + "End Sub")])
    _l, toks, _w = vmac._macro_action_to_tokens({"op": "code_replace", "search": "a = 1", "replace": "b"}, wb=wb1)
    assert toks[-2:] == ["--module", "m1"]
    wb2 = _WB([_Comp("m1", "a = 1"), _Comp("m2", "a = 1")])
    with pytest.raises(ValueError, match="m1 / m2"):
        vmac._macro_action_to_tokens({"op": "code_replace", "search": "a = 1", "replace": "b"}, wb=wb2)
    _l, toks, _w = vmac._macro_action_to_tokens({"op": "code_replace", "search": "zzz", "replace": "b"}, wb=wb1)
    assert "--module" not in toks
    assert vmac._GET_MODULE_RE.search("マクロ名  : 集計\nモジュール  : shu003\n").group(1) == "shu003"


def test_readonly_hands_do_not_take_a_backup():
    """20: find／export_csv／export_pdf も書く手扱いで、控えと画像が余計に走っていた。"""
    import vbam_agent as va
    for act in ({"op": "find", "text": "x"}, {"op": "export_csv", "range": "A1:B2", "path": "x.csv"},
                {"op": "export_pdf", "path": "x.pdf"}, {"op": "read", "range": "A1"},
                {"op": "inspect", "sheet": "名簿"}):
        assert va._will_write([act]) is False, act
    for act in ({"op": "view", "freeze": "A2"}, {"op": "page_setup", "landscape": True},
                {"op": "tidy", "ranges": ["A1:C3"]}):
        assert va._will_write([act]) is True, act


def test_add_updates_the_existing_names(monkeypatch):
    """21: add の後も existing が更新されず、同じ名前を 2 回 add できていた。"""
    import vbam_agent as va
    import pytest
    code = 'Sub 新しい集計()\n    x = 1\nEnd Sub\n'
    names = ["集計"]
    monkeypatch.setattr(va, "_run_cmd", lambda toks, wb=None: (True, "追加しました"))
    monkeypatch.setattr(va, "LAST_PROC_FILE", str(_tmp_proc_file()))
    vmac._execute_macro([{"op": "add", "module": "shu003", "name": "新しい集計", "code": code}], existing=names)
    assert "新しい集計" in names
    with pytest.raises(ValueError):
        vmac._macro_action_to_tokens({"op": "add", "module": "shu003", "name": "新しい集計", "code": code},
                                   existing=names)


def _tmp_proc_file():
    import tempfile
    import os as _os
    return _os.path.join(tempfile.mkdtemp(), "_last_proc.vba")


def test_route_no_longer_swallows_cell_work():
    """22: 「合計を書いて」がマクロの新規作成に、「別シートの名簿と」が build に振られていた。"""
    import vbam_agent as va
    assert va._route("数式が動かない。合計を書いて", macro_names=["集計", "印字"]) == 'sheet'
    assert va._route("別シートの名簿と突き合わせて") == 'sheet'
    assert va._route("売上を別シートにまとめて") == 'build'
    assert va._route("A 列を消すマクロを作って", macro_names=["集計"]) == 'macro'
    assert va._route("会員名簿を新しいシートに作って") == 'build'
    # 「マクロを書いて」「マクロを 1 本足して」は作る（書いて／足してを外しすぎていた・2026-09-04 夜）
    assert va._route("合計を出すマクロを書いて", macro_names=["集計"]) == 'macro'
    assert va._route("印刷するマクロを 1 本足して", macro_names=["集計"]) == 'macro'


def test_the_key_is_kept_only_after_it_goes_through(monkeypatch, tmp_path):
    """23: 疎通に失敗しても壊れたキーが金庫に残っていた。"""
    _tmp_store(monkeypatch, tmp_path)
    monkeypatch.setattr(vk, "_user_env_get", lambda n: None)
    monkeypatch.setattr(sys, "stdin", io.StringIO("AIzaSyABCDEFGHIJKLMN1234x7Qd\n"))
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    monkeypatch.setattr(vk, "_ping_key", lambda ai, model, key: (False, "通りませんでした（HTTP 400）"))
    args = _ap.Namespace(posargs=["gemini"], stdin=True, no_ping=False, model=None, days=None)
    assert va.cmd_set_key(args) is False
    assert va._store_read() == {}                       # 金庫は空のまま
    monkeypatch.setattr(sys, "stdin", io.StringIO("AIzaSyABCDEFGHIJKLMN1234x7Qd\n"))
    monkeypatch.setattr(vk, "_ping_key", lambda ai, model, key: (True, "通りました"))
    assert va.cmd_set_key(args) is True
    assert va._key_load("gemini") == "AIzaSyABCDEFGHIJKLMN1234x7Qd"


def test_cell_text_keeps_the_line_break_but_tsv_does_not():
    """29: セル内改行を潰していた（TSV に並べるときだけ潰す）。"""
    import vbam_agent as va
    assert va._cell_text("上\n下") == "上\n下"
    assert va._tsv_text("上\n下") == "上 下"
    assert va._rows_to_tsv([["上\n下", 1]]) == "上 下\t1\n"


def test_created_heavy_objects_are_recorded_for_undo(tmp_path, monkeypatch):
    """12: --undo が対象シートのセルと図形しか戻さず、作ったピボット等が残っていた。"""
    import vbam_agent as va
    import json
    results = [
        ("table create", True, "テーブル作成: [名簿] T名簿  範囲=$A$3:$E$9"),
        ("pivot create", True, "ピボット作成: [集計] P申込  ソース=名簿!A3:E9"),
        ("slicer add", True, "スライサー追加: S申込コース  ソース=P申込(pivot)  フィールド=申込コース  シート=集計"),
        ("powerquery add", True, "クエリ作成: Q名簿（接続のみ。シート/モデルへの読み込みは別途）"),
        ("powerquery load", True, "シートに読み込みました: Q名簿 → 取込!A1（テーブル: Q名簿_1, 接続: x）"),
        ("sheet_op add", True, "シート追加: 集計"),
        ("name_add", True, "名前付き範囲を追加: 貸出簿 → =名簿!$A$7:$F$19"),
        ("datamodel", True, "メジャー作成: T売上[売上合計] = SUM('T売上'[金額])  (書式=whole)"),
        ("chart column", True, "グラフ作成: [名簿] グラフ 1  種別 column  データ F1:G4"),
        ("read", False, "テーブル作成: [x] 失敗しているので拾わない"),
    ]
    got = va._created_from_results(results)
    kinds = [(g['kind'], g['name']) for g in got]
    assert ('table', 'T名簿') in kinds and ('pivot', 'P申込') in kinds and ('slicer', 'S申込コース') in kinds
    assert ('query', 'Q名簿') in kinds and ('table', 'Q名簿_1') in kinds and ('sheet', '集計') in kinds
    assert ('name', '貸出簿') in kinds and ('measure', '売上合計') in kinds and ('chart', 'グラフ 1') in kinds
    assert ('table', 'x') not in kinds                       # 失敗した手からは拾わない
    # 控えの覚書に積む（--undo が逆順で消す）
    undo = tmp_path / "undo.json"
    undo.write_text(json.dumps({'book': 'b.xlsx', 'sheet': '名簿', 'path': 'x'}), encoding='utf-8')
    monkeypatch.setattr(vu, "_LAST_AGENT_UNDO_FILE", str(undo))
    va._record_created(got)
    va._record_created([{'kind': 'sheet', 'name': '後'}])
    saved = json.loads(undo.read_text(encoding='utf-8'))['created']
    assert len(saved) == len(got) + 1 and saved[-1]['name'] == '後'


def test_other_sheets_are_counted_too():
    """14: 別シートに出したとき「変わったセル 0 個」と報告していた。"""
    import vbam_agent as va
    acts = [{"op": "pivot", "action": "create", "range": "A1:E9", "sheet": "集計", "at": "A3"},
            {"op": "powerquery", "action": "load", "name": "Q", "to": "sheet", "sheet": "取込"},
            {"op": "sheet_op", "action": "add", "name": "控え"},
            {"op": "sheet_op", "action": "copy", "name": "名簿", "to": "名簿の写し"},
            {"op": "tidy", "ranges": ["A1:C3"]}]
    assert va._other_sheets_in(acts) == {"集計", "取込", "控え", "名簿の写し"}
    assert va._other_sheets_in([{"op": "pivot", "action": "create", "range": "A1:E9", "at": "H3"}]) == set()


def test_plan_checklist_blocks_a_silent_drop():
    """33: 「未」が残ったまま done を言わせない（落とし物の防止を構造で強制する）。"""
    import vbam_agent as va
    say, actions, done, report, plan = va._parse_reply(
        '{"say":"やります","plan":[{"item":"色分け","state":"未"},{"item":"並べ替え","state":"未"}],'
        '"actions":[],"done":true,"report":"色分けしました"}')
    assert done is True and [p['state'] for p in plan] == ['未', '未']
    assert va._plan_pending(plan) == ['色分け', '並べ替え']
    _s, _a, _d, _r, plan2 = va._parse_reply(
        '{"plan":[{"item":"色分け","state":"済"},{"item":"並べ替え","state":"不可"}],"actions":[],"done":true}')
    assert va._plan_pending(plan2) == []                     # 済と不可は通す
    assert "済 1 / 不可 1 / 要判断 0 / 未 0" in va._plan_line(plan2)
    assert va._parse_plan([{"item": "x", "state": "へんな値"}]) == [{"item": "x", "state": "未"}]
    assert va._parse_plan([{"state": "済"}, "文字列"]) == []
    for word in ('"plan"', '1 つでも残っている done は道具が受け付けない'):
        assert word in va.RULES, word


def test_both_mode_is_only_for_requests_that_really_span():
    """8高: またがる依頼だけ統合ループへ。マクロの中身に「列を」が出るだけの依頼は macro のまま（2026-09-05）。"""
    import vbam_agent as va
    names = ["集計", "印字"]
    # 2 つ並んでいる＝both
    assert va._route("集計マクロを直してから、表も整えて", macro_names=names) == 'both'
    assert va._route("集計マクロを直して、結果をシートに書いて", macro_names=names) == 'both'
    assert va._route("集計が動かない。直したあと見やすくして", macro_names=names) == 'both'
    # マクロの語が無ければ入口に届かない（既存の設計。マクロ名の一覧を読むのは語があるときだけ）
    assert va._route("集計を直してから、表も整えて", macro_names=names) == 'sheet'
    # 1 つの仕事＝今までどおり
    assert va._route("A 列を消すマクロを作って", macro_names=names) == 'macro'
    assert va._route("集計マクロが動かない", macro_names=names) == 'macro'
    assert va._route("行を挿入するマクロを作って", macro_names=names) == 'macro'
    assert va._route("郵便番号を揃えて", macro_names=names) == 'sheet'
    assert va._route("集計を直してから、表も整えて", forced='macro', macro_names=names) == 'macro'   # --mode が最優先
    assert 'both' in va._MODES and 'both' in va._MODE_LABEL


def test_both_rules_carry_both_hand_lists_and_one_reply_format():
    """8高: 統合の規約は sheet と macro の手をそのまま連結する＝どちらかを直せば自動で効く。"""
    import vbam_agent as va
    assert '"op":"write_grid"' in vmac.BOTH_RULES and '"op":"code_replace"' in vmac.BOTH_RULES
    assert vmac.BOTH_RULES.count('返事の形:') == 1        # 前置きは 1 回だけ
    assert '混ぜて並べてよい' in vmac.BOTH_RULES
    assert '動かして確かめてはいない' in vmac.BOTH_RULES  # 直しただけで「動く」と言わせない
    assert vmac._hands_part("前置き\nactions に並べられる手:\nA") == "actions に並べられる手:\nA"
    assert vmac._hands_part("目印なし") == "目印なし"


def test_mixed_execute_routes_each_hand_and_compiles_once():
    """8高: シートの手とマクロの手を書かれた順に振り分け、コンパイルは 1 往復に 1 回だけ。"""
    import vbam_agent as va
    calls = []

    def fake_run_cmd(toks, wb=None):
        calls.append(list(toks))
        return True, "OK"

    old = va._run_cmd
    va._run_cmd = fake_run_cmd
    try:
        res = vmac._execute_mixed([{"op": "code_replace", "search": "a", "replace": "b", "module": "shu003"},
                                 {"op": "write_cells", "cells": {"A1": "x"}},
                                 {"op": "get", "name": "集計"}], "名簿", None)
    finally:
        va._run_cmd = old
    heads = [c[0] for c in calls]
    assert heads == ['code-replace', 'write-cells', 'get', 'compile'], heads
    assert len(res) == 4 and res[-1][0].startswith("compile")
    # マクロを書き換えていなければコンパイルしない
    calls.clear()
    va._run_cmd = fake_run_cmd
    try:
        res = vmac._execute_mixed([{"op": "write_cells", "cells": {"A1": "x"}},
                                 {"op": "list"}], "名簿", None)
    finally:
        va._run_cmd = old
    assert [c[0] for c in calls] == ['write-cells', 'list'] and len(res) == 2


def test_normalize_cleaning_rules_and_charkind_facts():
    """掃除を Python に（2026-09-08）: 姓名の間の空白・半角カナの規則 3 本、材料の「列の中の文字種」、小さい表は画像なし。
    AI が 30 セルを書き直し（25 秒）・同じ見え方を read で読み直し（12.8 秒）・画像（2.6 秒）で 52 秒かかっていた分。"""
    import vbam_agent as va
    N = lambda v, rules: va._normalize_value(v, [(r, None) for r in rules])[0]
    assert N("佐藤 一郎", ["space_zenkaku"]) == "佐藤　一郎" and N("佐藤  一郎", ["space_zenkaku"]) == "佐藤　一郎"
    assert N("佐藤　　一郎", ["space_zenkaku"]) == "佐藤　一郎" and N(" 佐藤　一郎 ", ["space_zenkaku"]) == " 佐藤　一郎 "   # 前後は trim の仕事
    assert N("サトウ　イチロウ", ["space_hankaku"]) == "サトウ イチロウ" and N("山田太郎", ["space_hankaku"]) == "山田太郎"
    assert N("ｻﾄｳ ｷﾞﾝｼﾞ", ["kana_zenkaku"]) == "サトウ ギンジ" and N("ﾎﾞｰﾙ･ﾍﾟﾝ｡", ["kana_zenkaku"]) == "ボール・ペン。"
    assert N("サトウ イチロウ", ["kana_zenkaku"]) == "サトウ イチロウ" and N("１２", ["kana_zenkaku"]) == "１２"   # 全角英数は触らない
    assert N(12000, ["kana_zenkaku", "space_zenkaku"]) == 12000
    assert {'kana_zenkaku', 'space_zenkaku', 'space_hankaku'} <= set(va._NORMALIZE_RULES)
    # 綴りの近い規則名は読み替える（AI の 1 字違いで手が連鎖停止しない）。遠い名前は断る
    _, rules, _, _, _, _ = va._normalize_spec({"op": "normalize", "range": "C6:C20", "rules": ["kata_zenkaku", "Space_Zenkaku"],
                                               "overwrite": True})
    assert [r[0] for r in rules] == ["kana_zenkaku", "space_zenkaku"]
    # 列の中の文字種（見出しの下だけ数える。何も無い列は出さない）
    grid = [["会員番号", "氏名", "フリガナ", "種別"], ["1001", "佐藤　一郎", "ｻﾄｳ ｼﾞﾛｳ", "正会員"], ["１００２", "鈴木 花子", "", "学生"]]
    assert va._charkind_facts(grid, 5, 1, 0) == ["A列 全角英数 1", "B列 文字の間の空白=全角 1・半角 1",
                                                  "C列 文字の間の空白=全角 0・半角 1（この列へ書き足す値の空白も半角）・半角カナ 1"]
    assert va._charkind_facts(grid, 5, 1, None) == [] and va._charkind_facts([grid[0]], 5, 1, 0) == []


def test_clean_plan_decides_rules_without_ai():
    """clean-table の頭（2026-09-08）: 列ごとの規則・重複・空欄・日付列を純 Python で決める（AI 抜きで 3 秒）。"""
    import vbam_agent as va
    import datetime as dt
    d = lambda m, day: dt.datetime(2026, m, day)
    grid = [
        ["名簿", None, None, None, None],
        ["会員番号", "氏名", "フリガナ", "種別", "入会日"],
        ["1001", "佐藤　一郎", "ｻﾄｳ ｲﾁﾛｳ", "正会員", d(4, 1)],
        ["1002", "鈴木 花子", "スズキ ハナコ", "正会員", d(5, 12)],
        ["１００３", "田中　　次郎", "タナカ ジロウ", "準会員", "2026/6/3"],
        ["1004", "高橋　三郎", "", "学生", d(7, 20)],
        ["1004", "高橋　三郎", "", "学生", d(7, 20)],      # 完全な重複
    ]
    p = vc.clean_plan(grid, 1)
    assert p['rules'][0] == ['hankaku']                                  # 会員番号の全角数字
    assert p['rules'][1] == ['space_zenkaku']                            # 氏名: 全角 3・半角 1 → 多数派の全角 1 つに
    assert p['rules'][2] == ['kana_zenkaku']                             # フリガナの半角カナ（空白は半角で揃っている）
    assert 3 not in p['rules'] and p['rules'][4] == ['date']             # 種別は乱れなし／入会日は文字の日付あり
    assert p['dups'] == [(6, 5)] and p['date_cols'] == [4]
    assert p['blanks'] == [(5, 2)]                                       # 高橋のフリガナ（重複行の分は数えない）
    # 見出しが無い・空の格子は何も返さない（落ちない）
    assert vc.clean_plan([], 0)['rules'] == {} and vc.clean_plan(grid, None)['dups'] == []
    # 空欄は「同じ列の他の行が埋まっている」ときだけ拾う
    g2 = [["氏名", "フリガナ"], ["佐藤", "サトウ"], ["鈴木", "スズキ"], ["田中", "タナカ"], ["山田", None]]
    assert vc.clean_plan(g2, 0)['blanks'] == [(4, 1)]


def test_rows_and_columns_can_be_deleted_only_with_approval():
    """10: 行・列を消す手を開けた（2026-09-05）。控えで戻せる（⑥⑦）が、承認の言葉が要る。"""
    import vbam_agent as va
    import pytest
    label, toks = va._action_to_tokens({"op": "row", "action": "delete", "at": 7, "count": 2,
                                        "overwrite": True}, "名簿")
    assert toks == ["row", "delete", "7", "2", "--sheet", "名簿"] and label == "row delete 7 x2"
    label, toks = va._action_to_tokens({"op": "col", "action": "delete", "at": "c",
                                        "overwrite": True}, "名簿")
    assert toks == ["col", "delete", "C", "1", "--sheet", "名簿"]
    # 承認が無ければ実行しない（消す候補は report に書かせる）
    for bad in ({"op": "row", "action": "delete", "at": 7},
                {"op": "col", "action": "delete", "at": "C"}):
        with pytest.raises(ValueError, match="overwrite"):
            va._action_to_tokens(bad, "名簿")
    # insert は今までどおり承認なしで通る／知らない action は断る
    label, toks = va._action_to_tokens({"op": "row", "action": "insert", "at": 7}, "名簿")
    assert toks[:2] == ["row", "insert"]
    with pytest.raises(ValueError, match="insert か delete"):
        va._action_to_tokens({"op": "row", "action": "clear", "at": 7}, "名簿")
    # 行・列が動くと番地も行番号もずれる＝不変条件の hidden／formula は外す（delete でも同じ）
    assert va._inv_skip_for([{"op": "row", "action": "delete", "at": 7, "overwrite": True}]) == ('formula', 'hidden', 'tables')
    # 消せるものと、消す前に見るものを RULES に書いた（2026-09-06 夜に手の名前を row_delete／col_delete に）
    assert "row_delete／col_delete" in va.RULES and "#REF! に化ける" in va.RULES


def test_progress_file_says_what_it_is_doing(tmp_path, monkeypatch):
    """5: MCP 越しだと終わるまで無言だった。最新の 1 状態をファイルに置いて待っている側が読む（2026-09-05）。"""
    import vbam_agent as va
    f = tmp_path / "progress.json"
    monkeypatch.setattr(vai, "_AGENT_PROGRESS_FILE", str(f))
    assert va.progress_read() is None                     # まだ無い＝黙って None
    rec = va.progress_write('ask', book="b.xlsm", sheet="名簿", turn=2, max_turns=4, sec=12.34)
    assert rec['state'] == 'ask' and rec['pid'] == os.getpid()
    got = va.progress_read()
    assert got['turn'] == 2 and got['book'] == "b.xlsm"
    assert got['line'] == "AI に聞いています / 往復 2/4 / b.xlsm!名簿 / 12.3 秒"
    # 手を実行しているところ・終わったところ
    va.progress_write('run', book="b.xlsm", sheet="名簿", turn=2, max_turns=4,
                      say="整えます", hands="手 2 本: tidy / format")
    assert "手 2 本: tidy / format" in va.progress_read()['line']
    va.progress_write('done', book="b.xlsm", sheet="名簿", note="整えました", sec=30.0)
    assert va.progress_read()['line'].startswith("終わりました")
    # 空の値は書かない（行が「/ / /」で埋まらない）
    assert 'say' not in va.progress_write('ask', say="", note=None)
    assert va.progress_line(None) == "" and va.progress_line({}) == ""


def test_tests_never_write_the_real_progress_file():
    """pytest が本物の _agent_progress.json を上書きしていた（2026-09-05 に実測で発覚）。

    偽のブックで run_agent を回すテストが本物のファイルへ「Book1!名簿 / 使い捨て /
    何もしていない」と書き、agent_status がそれを最新の仕事として返していた。
    conftest.py の autouse で全テストをテンポラリへ逃がす＝この検査がその見張り番。
    """
    import vbam_agent as va
    real = os.path.join(va.SCRIPT_DIR, '_agent_progress.json')
    assert va._AGENT_PROGRESS_FILE != real
    assert not va._AGENT_PROGRESS_FILE.startswith(va.SCRIPT_DIR)
    # 走行台帳も同じ（2026-09-06）。偽の走行が 1 行混ざると「実運用の量」の数字が嘘になる
    assert not va._AGENT_RUNS_FILE.startswith(va.SCRIPT_DIR)
    assert not va._AGENT_LOGS_DIR.startswith(va.SCRIPT_DIR)


def _build_stubs(monkeypatch, va, plan):
    """run_build を Excel 抜きで回すための最小の差し替え。"""
    monkeypatch.setattr(va, "ask_design", lambda *a, **k: {})
    monkeypatch.setattr(va, "plan_sheet", lambda spec: plan)
    monkeypatch.setattr(va, "_plan_summary", lambda p: "")
    monkeypatch.setattr(va, "_book_materials", lambda wb: "使用範囲: A1")


def test_run_build_reports_progress_and_pushes_its_own_clock(tmp_path, monkeypatch):
    """build は進み具合を書かず、仕事の時計も押していなかった（2026-09-05 の実機確認）。

    その結果 agent_status には前の仕事の行が出て、報告の「経過: materials から N 秒」は
    ディスクに残った前の仕事の時計＝1030.2 秒という嘘の数字になっていた。
    """
    import vbam_agent as va
    f = tmp_path / "progress.json"
    monkeypatch.setattr(vai, "_AGENT_PROGRESS_FILE", str(f))
    pushed = []
    monkeypatch.setattr(va, "job_clock_start", lambda label="": pushed.append(label))
    _build_stubs(monkeypatch, va, {"sheet": "名簿", "overwrite": False})
    wb = type("WB", (), {"Name": "b.xlsx"})()
    assert va.run_build("表を作って", None, wb, dry_run=True) is True
    assert pushed and "b.xlsx" in pushed[0]               # 自分の時計を押した
    got = va.progress_read(str(f))
    assert got['state'] == 'done' and got['book'] == "b.xlsx"
    assert got['sec'] < 60                               # 前の仕事の秒を引き継がない


def test_run_build_does_not_add_a_second_book_when_given_a_fresh_one(tmp_path, monkeypatch):
    """--new-book で白紙を足してから run_build に渡すと、もう 1 冊足していた（2026-09-05）。"""
    import vbam_agent as va
    monkeypatch.setattr(vai, "_AGENT_PROGRESS_FILE", str(tmp_path / "p.json"))
    monkeypatch.setattr(va, "job_clock_start", lambda label="": None)
    _build_stubs(monkeypatch, va, {"sheet": "名簿", "overwrite": False})
    monkeypatch.setattr(va, "_mark_undo_stale", lambda *a, **k: None)
    monkeypatch.setattr(va, "_apply", lambda xl, wb, plan, ow: type("WS", (), {"Name": "名簿"})())
    monkeypatch.setattr(va, "_verify_build", lambda xl, ws, plan: True)
    added = []

    class _XL:
        class Workbooks:
            @staticmethod
            def Add():
                added.append(1)
                return type("WB", (), {"Name": "Book2"})()

    wb = type("WB", (), {"Name": "Book1"})()
    assert va.run_build("表を作って", _XL(), wb, new_book=True, book_is_new=True) is True
    assert added == []                                   # 渡された白紙をそのまま使う
    assert va.run_build("表を作って", _XL(), wb, new_book=True) is True
    assert added == [1]                                  # 白紙を渡されなければ今までどおり足す


def test_agent_workbook_adds_a_blank_book_only_for_new_book(monkeypatch, capsys):
    """--new-book なのに「開いているブックが見つかりません」で断られていた（2026-09-05 実機）。"""
    import vbam_agent as va

    def _no_book(*a, **k):
        raise Exception("起動中の Excel に開いているブックが見つかりません。")

    monkeypatch.setattr(va, "get_workbook", _no_book)

    class _XL:
        Visible = False

        class Workbooks:
            @staticmethod
            def Add():
                return type("WB", (), {"Name": "Book1"})()

    monkeypatch.setattr(va, "_get_active_excel", lambda: _XL())
    xl, wb, is_new = va._agent_workbook(None, True)
    assert wb is not None and wb.Name == "Book1" and is_new is True
    # --new-book でなければ今までどおり断る（勝手に白紙を足さない）
    assert va._agent_workbook(None, False) == (None, None, False)
    assert "開いているブックが見つかりません" in capsys.readouterr().out
    # ブック名を名指ししているときも足さない（その 1 冊を開くのが筋）
    assert va._agent_workbook("台.xlsm", True) == (None, None, False)


def test_mcp_agent_can_run_without_waiting(monkeypatch):
    """5: agent(wait=False) は待たずに id を返し、agent_status がキューを通さずに様子を答える。"""
    import queue as _queue
    import threading as _threading
    import vba_mcp_server as ms
    import vbam_agent as va
    jobs = _queue.Queue()
    monkeypatch.setattr(ms, "_jobs", jobs)
    monkeypatch.setattr(ms, "_live_jobs", {})
    monkeypatch.setattr(ms, "_log_call", lambda line: None)
    monkeypatch.setattr(ms, "_gate_check", lambda line, wait=None: None)
    monkeypatch.setattr(va, "progress_read", lambda path=None: {'time': '12:00:00', 'line': '往復 1/4 / tidy'})

    out = ms.agent(request="整えて", wait=False)
    jid = out.split(": ")[1].split("\n")[0]
    assert jid.startswith("job") and "agent_status" in out
    line, box, done = jobs.get_nowait()                   # キューには入っている（ワーカーが拾う）
    assert line.startswith("agent ") and "整えて" in line

    # 終わる前: 進み具合を返す（ワーカーは塞がっているが、こちらは答えられる）
    assert "まだ走っています" in ms.agent_status(jid) and "往復 1/4" in ms.agent_status(jid)
    # 終わった後: 全文を返し、覚えは捨てる（2 回目は「覚えがありません」）
    box.update({'ok': True, 'out': "できました", 'err': ""})
    done.set()
    got = ms.agent_status(jid)
    assert got.startswith("終わりました") and "できました" in got
    assert "覚えがありません" in ms.agent_status(jid)
    # id を省くといちばん新しい仕事／仕事が無ければ進み具合だけ返す
    assert "走らせた仕事はありません" not in ms.agent_status()
    assert isinstance(_threading.Event(), _threading.Event)


def test_undo_meta_grows_to_every_written_sheet(tmp_path, monkeypatch):
    """7（後半）: 別シートに書けるようにした以上、--undo は書いた全部のシートを戻せないといけない。

    覚書は書き換える前に足す（図形の位置もそのときに写す）。まだ無いシートは控えに入っていないので足さない。
    """
    import json
    import vbam_agent as va
    f = tmp_path / "undo.json"
    monkeypatch.setattr(vu, "_LAST_AGENT_UNDO_FILE", str(f))
    monkeypatch.setattr(vu, "_shapes_geometry", lambda ws: [{"name": "図1", "left": 1, "top": 2,
                                                             "width": 3, "height": 4}])
    f.write_text(json.dumps({"book": "b.xlsm", "sheet": "名簿", "sheets": ["名簿"], "path": "p",
                             "shapes": [{"name": "元の図", "left": 0, "top": 0, "width": 1, "height": 1}]}),
                 encoding="utf-8")

    class WB:
        def Sheets(self, name):
            if name == "まだ無い":
                raise Exception("no such sheet")
            return object()

    va._undo_add_sheet(WB(), "集計")
    va._undo_add_sheet(WB(), "集計")          # 2 回呼んでも増えない
    va._undo_add_sheet(WB(), "まだ無い")      # これから作るシートは戻す対象にしない（created の側が消す）
    meta = json.loads(f.read_text(encoding="utf-8"))
    assert va._undo_sheets(meta) == ["名簿", "集計"]
    assert meta["shapes_by_sheet"]["集計"][0]["name"] == "図1"
    # 図形は「シートごと」に引ける。古い覚書（shapes だけ）は 1 枚目のものとして読む
    by = va._undo_shapes(meta, "名簿")
    assert by["名簿"][0]["name"] == "元の図" and by["集計"][0]["name"] == "図1"
    assert va._undo_shapes({"shapes": [{"name": "旧"}]}, "名簿")["名簿"][0]["name"] == "旧"
    assert va._undo_sheets({"sheet": "名簿"}) == ["名簿"]        # 'sheets' が無い古い覚書
    assert va._undo_sheets({}) == [] and va._undo_sheets(None) == []


def test_undo_restores_every_recorded_sheet():
    """7（後半）: undo_agent が覚書のシートを 1 枚ずつ戻す（1 枚目だけで終わらない）。"""
    import inspect as _inspect
    import vbam_agent as va
    src = _inspect.getsource(va.undo_agent)
    assert "sheets = _undo_sheets(meta)" in src
    assert "for name in sheets:" in src, "1 枚ずつ回して戻す"
    # 戻す先のシートは「作った物」として消さない（消したら貼る先が無い）
    assert "c.get('name') in sheets" in src
    # 図形はシートごとに引く
    assert "shapes_by = _undo_shapes(meta, sheet)" in src and "shapes_by.get(name)" in src


def test_write_hands_can_name_another_sheet():
    """7（前半）: 書く手に "sheet" を書けばそのシートへ。重い道具の "sheet" は置き場なので読み替えない。"""
    import vbam_agent as va
    # 番地の修飾も --sheet も、指した側のシートになる
    label, toks = va._action_to_tokens({"op": "write_cells", "cells": {"A1": "x"}, "sheet": "集計"}, "名簿")
    assert toks[:3] == ["write-cells", "--sheet", "集計"]
    label, toks = va._action_to_tokens({"op": "tidy", "ranges": ["A1:C9"], "sheet": "集計"}, "名簿")
    assert toks == ["tidy", "--sheet", "集計", "A1:C9"]
    # 表の足回り（範囲に 'シート'!A1 の形で渡す手）も同じ
    label, toks = va._action_to_tokens({"op": "autofilter", "range": "A1:F9", "sheet": "集計"}, "名簿")
    assert any("'集計'!A1:F9" == t for t in toks), toks
    # "sheet" が無ければ今までどおり対象シート
    label, toks = va._action_to_tokens({"op": "write_cells", "cells": {"A1": "x"}}, "名簿")
    assert toks[:3] == ["write-cells", "--sheet", "名簿"]
    # 重い道具の "sheet" は「ピボットの置き場」＝書き先の読み替えに使わない
    assert va._act_sheet({"op": "pivot", "action": "create", "sheet": "集計"}, "名簿") == "名簿"
    assert va._act_sheet({"op": "write_grid", "range": "A1", "sheet": "集計"}, "名簿") == "集計"
    assert va._act_sheet({"op": "tidy"}, "名簿") == "名簿" and va._act_sheet("文字列", "名簿") == "名簿"


def test_other_sheets_and_written_sheets_include_the_named_ones():
    """7（前半）: 別シートに書いたら、控え・差分・不変条件がそのシートも見る。読むだけの手は入れない。"""
    import vbam_agent as va
    acts = [{"op": "write_cells", "cells": {"A1": "x"}, "sheet": "集計"},
            {"op": "read", "range": "A1:C3", "sheet": "基準"},          # 読むだけ＝控えは要らない
            {"op": "tidy", "ranges": ["A1:C9"]}]
    assert va._other_sheets_in(acts) == {"集計"}
    assert va._act_writes(acts[0]) is True and va._act_writes(acts[1]) is False
    # 意図して書いたシートは「頼んでいない変化」に数えない（⑥との噛み合わせ）
    before = {'sheets': {'名簿': {'used': 'A1', 'values': [['a']], 'formulas': {}, 'hidden': (), 'cf': 0,
                                  'validation': '', 'filter': ''},
                         '集計': {'used': 'A1', 'values': [['x']], 'formulas': {}, 'hidden': (), 'cf': 0,
                                  'validation': '', 'filter': ''}},
              'names': {}}
    now = {'sheets': {'名簿': before['sheets']['名簿'],
                      '集計': dict(before['sheets']['集計'], values=[['y']])}, 'names': {}}
    assert va._inv_compare(before, now, '名簿') == ["頼んでいないシート「集計」が変わった（A1: 'x' → 'y'）"]
    assert va._inv_compare(before, now, {'名簿', '集計'}) == []      # 意図して書いた＝違反ではない


def test_writing_to_a_sheet_that_does_not_exist_is_refused():
    """7（前半）: 無いシートを名指ししたら実行せずに断り、後ろの手も止める（分かりにくい COM エラーにしない）。"""
    import vbam_agent as va

    class FakeSheets:
        def __iter__(self):
            return iter([type("S", (), {"Name": n})() for n in ("名簿", "目次")])

    wb = type("WB", (), {"Sheets": FakeSheets()})()
    res = va._execute([{"op": "write_cells", "cells": {"A1": "x"}, "sheet": "無い表"},
                       {"op": "tidy", "ranges": ["A1:C9"]}], "名簿", wb)
    assert res[0][1] is False and "シート '無い表' がありません" in res[0][2] and "名簿, 目次" in res[0][2]
    # 2026-09-06 夜: 実行前に全部の手を見るので、後ろの手は「止めました」でなく「1 手も実行していません」
    assert res[1][1] is False and "1 手も実行していません" in res[1][2]


def test_inv_skip_is_decided_from_the_hands():
    """1: 本番には弾ごとの inv_skip が無い。依頼どおりに動かせば当然変わるものを手から決める（2026-09-05）。"""
    import vbam_agent as va
    assert va._inv_skip_for([{"op": "row", "action": "insert", "at": 7}]) == ('formula', 'hidden', 'tables')
    assert va._inv_skip_for([{"op": "col", "action": "insert", "at": "C"}]) == ('formula', 'hidden', 'tables')
    assert va._inv_skip_for([{"op": "sort", "range": "A1:F9", "key": "B"}]) == ('formula', 'hidden', 'tables')
    assert va._inv_skip_for([{"op": "autofilter", "range": "A1:F9", "equals": "x"}]) == ('hidden',)
    assert va._inv_skip_for([{"op": "clear_range", "range": "A1:F9", "overwrite": True}]) == ('formula', 'used')
    assert va._inv_skip_for([{"op": "clear_range", "range": "A1:F9"}]) == ()      # 承認が無いなら外さない
    assert va._inv_skip_for([{"op": "sheet_op", "action": "rename", "to": "新"}]) == ('order', 'other_sheet')
    # 普通の手では 1 つも外さない（外しすぎると検査が効かない）
    assert va._inv_skip_for([{"op": "tidy", "ranges": ["A1:F9"]}, {"op": "write_cells", "cells": {"A1": "x"}}]) == ()
    assert va._inv_skip_for([]) == () and va._inv_skip_for(["文字列", None]) == ()
    # 複数の手はまとめて外す
    assert va._inv_skip_for([{"op": "sort", "range": "A1:F9"},
                             {"op": "sheet_op", "action": "rename", "to": "新"}]) == (
        'formula', 'hidden', 'order', 'other_sheet', 'tables')
    # 2026-09-06 に足した分: 結合を外すのが依頼なら merged を外す（帳票→一覧）
    assert va._inv_skip_for([{"op": "format", "range": "A1:F9", "unmerge": True}]) == ('merged',)


def test_the_loop_refuses_done_before_undone_changes_are_explained(tmp_path, monkeypatch):
    """1: 実射でしか見ていなかった不変条件を本番の done の関所にした（2026-09-05）。

    頼んでいない変化があるうちは done を受け付けず、残ったまま終われば ok にしない。
    """
    import vbam_agent as va
    replies = ['{"say":"やります","actions":[{"op":"write_cells","cells":{"A1":"x"}}],"done":false}',
               '{"say":"終わりました","actions":[],"done":true,"report":"書きました"}',
               '{"say":"直せません","actions":[],"done":true,"report":"別のシートはピボットの置き場です"}']
    sent = []

    def fake_ask(ai, model, key, history):
        sent.append(history[-1][1])
        return replies[min(len(sent) - 1, len(replies) - 1)], {'in': 1, 'out': 1}, 0.0

    monkeypatch.setattr(va, "_ask", fake_ask)
    monkeypatch.setattr(va, "_ai_setup", lambda ai, model: ('gemini', 'm', 'k'))
    monkeypatch.setattr(va, "_execute", lambda a, s, wb=None: [("write_cells 1 セル", True, "ok")])
    monkeypatch.setattr(va, "_sheet_snapshot", lambda ws: {})
    monkeypatch.setattr(va, "_agent_backup", lambda *a, **k: None)   # 2026-09-06: 走行の名札 run_id が増えた
    monkeypatch.setattr(va, "_book_cell_count", lambda wb: 1000)
    monkeypatch.setattr(va, "_snapshot_book", lambda wb, own=None: {'sheets': {}, 'names': {}})
    monkeypatch.setattr(va, "_inv_violations",
                        lambda wb, before, own, skip=(): ["頼んでいないシート「目次」が変わった（A1: '' → 'x'）"])
    monkeypatch.setattr(va, "_audit_now", lambda wb, sheet, targets: [])
    monkeypatch.setattr(va, "_changes_of", lambda *a, **k: [])
    monkeypatch.setattr(va, "_format_changes", lambda *a, **k: [])
    monkeypatch.setattr(va, "_LAST_AGENT_LOG_FILE", str(tmp_path / "log.jsonl"))
    monkeypatch.setattr(va, "_LAST_AGENT_ASK_FILE", str(tmp_path / "ask.txt"))
    wb = type("WB", (), {"Name": "b.xlsx", "Path": "", "Sheets": staticmethod(lambda name: object())})()
    r = va.run_agent("A1 に x と書いて", "名簿", wb, max_turns=4, materials="使用範囲: A1",
                     verify=False, grade=False, show_image=False)
    # 差し戻しは 1 回だけ（2 通目の done が「頼んでいない変化」で押し返される）
    inv_msgs = [m for m in sent if "依頼に無い変化があります" in m]
    assert len(inv_msgs) == 1 and "目次" in inv_msgs[0]
    # 直らないまま終わった＝報告には残り、合格にはしない
    assert r['inv'] and r['ok'] is False and r['done'] is True


def test_the_loop_refuses_done_while_something_is_pending(tmp_path, monkeypatch):
    """33（ループ側）: 未が残る done は受け付けず、往復を続ける（上限に当たったら終わる）。"""
    import vbam_agent as va
    replies = ['{"say":"やります","plan":[{"item":"色分け","state":"未"},{"item":"並べ替え","state":"未"}],'
               '"actions":[],"done":true,"report":"色分けだけ"}',
               '{"plan":[{"item":"色分け","state":"済"},{"item":"並べ替え","state":"不可"}],'
               '"actions":[],"done":true,"report":"並べ替えはできない"}']
    sent = []

    def fake_ask(ai, model, key, history):
        sent.append(history[-1][1])
        return replies[min(len(sent) - 1, len(replies) - 1)], {'in': 1, 'out': 1}, 0.0

    monkeypatch.setattr(va, "_ask", fake_ask)
    monkeypatch.setattr(va, "_ai_setup", lambda ai, model: ('gemini', 'm', 'k'))
    monkeypatch.setattr(va, "_LAST_AGENT_LOG_FILE", str(tmp_path / "log.jsonl"))
    monkeypatch.setattr(va, "_LAST_AGENT_ASK_FILE", str(tmp_path / "ask.txt"))
    wb = type("WB", (), {"Name": "b.xlsx"})()
    r = va.run_agent("色分けして並べ替えて", "名簿", wb, max_turns=3, materials="使用範囲: A1",
                     verify=False, backup=False)
    assert r['done'] is True and r['turns'] == 2
    assert "未の項目が残っています" in sent[1] and "並べ替え" in sent[1]
    assert [p['state'] for p in r['plan']] == ['済', '不可']


def test_replay_ask_returns_recorded_replies_then_raises(tmp_path):
    """6: --replay の下地。記録の reply を実行順に返す。'resumed' の行（--continue の再掲）は数えない。
    記録より長く回ったら「記録が尽きました」で止める（想定外の枝分かれ＝道具側の直しが記録と食い違った合図）。
    """
    import json
    import vbam_agent as va
    log = tmp_path / "log.jsonl"
    rows = [
        {"meta": {"book": "b.xlsx", "sheet": "名簿", "mode": "sheet"}},
        {"turn": 1, "prompt": "旧の質問", "reply": "旧の返事", "resumed": True},
        {"turn": 1, "prompt": "p1", "reply": "r1"},
        {"turn": 2, "prompt": "p2", "reply": "r2"},
        {"turn": 3, "prompt": "p3", "reply": "r3"},
        {"turn": 3, "grade": True, "prompt": "採点", "reply": "r4"},
    ]
    with open(log, "w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")
    ask = va._replay_ask(str(log))
    got = [ask(None, None, None, None)[0] for _ in range(4)]
    assert got == ["r1", "r2", "r3", "r4"]           # resumed の行は数えられていない
    with pytest.raises(RuntimeError, match="記録が尽きました"):
        ask(None, None, None, None)


def test_replay_reproduces_the_recorded_loop_without_calling_the_api(tmp_path, monkeypatch):
    """6: run_agent(ask=_replay_ask(...)) が記録どおりの plan・done を、API を呼ばず（課金 0）再現する。"""
    import json
    import vbam_agent as va
    src = tmp_path / "src.jsonl"
    rows = [
        {"meta": {"book": "b.xlsx", "sheet": "名簿", "mode": "sheet", "request": "色分けして並べ替えて"}},
        {"turn": 1, "prompt": "p1",
         "reply": '{"say":"やります","plan":[{"item":"色分け","state":"未"},{"item":"並べ替え","state":"未"}],'
                  '"actions":[],"done":true,"report":"色分けだけ"}'},
        {"turn": 2, "prompt": "p2",
         "reply": '{"plan":[{"item":"色分け","state":"済"},{"item":"並べ替え","state":"不可"}],'
                  '"actions":[],"done":true,"report":"並べ替えはできない"}'},
    ]
    with open(src, "w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")
    monkeypatch.setattr(va, "_LAST_AGENT_LOG_FILE", str(tmp_path / "log.jsonl"))    # run_agent 自身の記録（src とは別）
    monkeypatch.setattr(va, "_LAST_AGENT_ASK_FILE", str(tmp_path / "ask.txt"))
    wb = type("WB", (), {"Name": "b.xlsx"})()
    r = va.run_agent("色分けして並べ替えて", "名簿", wb, max_turns=3, materials="使用範囲: A1",
                     verify=False, backup=False, grade=False, ask=va._replay_ask(str(src)))
    assert r['done'] is True and r['turns'] == 2
    assert [p['state'] for p in r['plan']] == ['済', '不可']
    assert r['in_tokens'] == 0 and r['out_tokens'] == 0     # API を呼んでいない＝課金 0


def test_max_turns_is_carried_over_and_bounded():
    """28: --continue が上限 4 に戻る／--max-turns 0 や負数が通っていた。"""
    import vbam_agent as va
    from vbam_recipes import _RECIPE_MAX_TURNS
    for resume_meta, recipe, given, want in (({'max_turns': 6}, None, None, 6),
                                             (None, "手順書", None, _RECIPE_MAX_TURNS),
                                             ({'max_turns': 6}, None, 2, 2),
                                             (None, None, None, va._DEFAULT_MAX_TURNS)):
        default_turns = (resume_meta or {}).get('max_turns') if resume_meta else None
        default_turns = default_turns or (_RECIPE_MAX_TURNS if recipe else va._DEFAULT_MAX_TURNS)
        assert int(given or default_turns) == want
    args = _ap.Namespace(posargs=[], max_turns=0, recipes=False, undo=False, fire=False, recipe=None,
                         cont=False, request_file=None, ai=None, model=None)
    assert va._cmd_agent_body(args) is False


def test_inspect_does_not_reset_the_job_clock():
    """27: 往復の途中の inspect が仕事の時計を押し直していた。"""
    import vbam_core as vc
    import vbam_agent as va
    saved = vc.job_clock_get()
    try:
        vc.job_clock_start("材料")
        first = vc.job_clock_get()
        vc.job_clock_start("inspect が押し直した")
        vc.job_clock_set(first)
        assert vc.job_clock_get()['label'] == "材料"
        assert callable(va.job_clock_set) and callable(va.job_clock_get)
    finally:
        vc.job_clock_set(saved)


# ---- 2026-09-04 夕: 監査 10 件の直し（build の控え・錠・落とし物・結合と clear・控えの間引き） ----

def test_agent_lock_ignores_broken_lock_file(tmp_path, monkeypatch):
    """錠のファイルが辞書でない JSON だと agent が全部起動できなくなっていた（AttributeError）。"""
    import vbam_agent as va
    lock = tmp_path / "lock.json"
    monkeypatch.setattr(va, "_AGENT_LOCK_FILE", str(lock))
    lock.write_text("[1,2,3]", encoding="utf-8")
    with va._agent_lock():
        pass                                   # 例外が出ないこと＝直っている
    assert not lock.exists()                   # 抜けるときに外す
    lock.write_text("これは JSON ですらない", encoding="utf-8")
    with va._agent_lock():
        pass


def test_agent_clear_range_and_unmerge_hands():
    """結合の解除（format --unmerge）と範囲クリア（clear_range）。帳票を元の場所で一覧に直す手。"""
    import vbam_agent as va
    import pytest
    assert 'clear_range' in va._ALLOWED_OPS and '"op":"clear_range"' in va.RULES
    label, toks = va._action_to_tokens({"op": "clear_range", "range": "A1:F30", "overwrite": True}, "帳票")
    assert toks == ["clear-range", "--sheet", "帳票", "--contents", "A1:F30"] and "contents" in label
    _l, toks = va._action_to_tokens({"op": "clear_range", "range": "A1:F30", "what": "all",
                                     "overwrite": True}, "帳票")
    assert toks[-2] == "--all"
    with pytest.raises(ValueError):            # 承認の言葉が無ければ消さない
        va._action_to_tokens({"op": "clear_range", "range": "A1:F30"}, "帳票")
    with pytest.raises(ValueError):            # what は 3 つだけ
        va._action_to_tokens({"op": "clear_range", "range": "A1", "what": "everything",
                              "overwrite": True}, "帳票")
    _l, toks = va._action_to_tokens({"op": "format", "range": "A1:F9", "unmerge": True}, "帳票")
    assert "--unmerge" in toks and "--sheet" in toks
    _l, toks = va._action_to_tokens({"op": "format", "range": "A1:B1", "merge": True}, "帳票")
    assert "--merge" in toks
    assert va._will_write([{"op": "clear_range", "range": "A1", "overwrite": True}]) is True


def test_agent_plan_pending_is_not_waived_on_the_last_turn():
    """最後の往復だけ「未」チェックが素通りしていた（落とし物が合格で出ていく）。"""
    import vbam_agent as va
    src = open(va.__file__, encoding='utf-8').read()
    assert "if pending:" in src and "dropped = pending" in src
    assert "落とし物: 未の項目が残ったまま往復が尽きました" in src
    plan = [{"item": "色分け", "state": "済"}, {"item": "並べ替え", "state": "未"}]
    assert va._plan_pending(plan) == ["並べ替え"]


def test_agent_undo_stops_when_the_backup_is_stale(tmp_path, monkeypatch, capsys):
    """控えを取った後に build／macro を回すと、--undo がその前の仕事の控えで上書きしていた。"""
    import vbam_agent as va
    import json
    memo = tmp_path / "undo.json"
    monkeypatch.setattr(vu, "_LAST_AGENT_UNDO_FILE", str(memo))
    memo.write_text(json.dumps({"book": "B.xlsx", "sheet": "S", "path": str(tmp_path / "no.xlsx"),
                                "request": "元の仕事", "time": "2026-09-04 10:00:00"},
                               ensure_ascii=False), encoding="utf-8")
    va._mark_undo_stale("build（新しいシートを組んだ）")
    meta = json.loads(memo.read_text(encoding="utf-8"))
    assert meta["stale"].startswith("build") and meta.get("stale_time")
    assert va.undo_agent(None, dry_run=True) is False          # 既定では止まる（COM に触る前に返る）
    out = capsys.readouterr().out
    assert "いまの仕事のものではありません" in out and "--force" in out
    # 控えが取り直されれば印は消える（_agent_backup が覚書を書き直す）
    memo.write_text(json.dumps({"book": "B.xlsx", "sheet": "S"}, ensure_ascii=False), encoding="utf-8")
    assert json.loads(memo.read_text(encoding="utf-8")).get("stale") is None


def test_agent_prune_backups_is_per_book(tmp_path, monkeypatch):
    """控えの間引きがブック横断だった＝別のブックを 5 回回すと前のブックの控えが消えていた。"""
    import vbam_agent as va
    import time as _t
    monkeypatch.setattr(vu, "BACKUP_DIR", str(tmp_path))
    for i in range(6):
        p = tmp_path / f"甲{va._AGENT_BACKUP_MARK}2026090{i}_000000.xlsx"
        p.write_text("x", encoding="utf-8")
        os.utime(p, (_t.time() + i, _t.time() + i))
    keep = tmp_path / f"乙{va._AGENT_BACKUP_MARK}20260901_000000.xlsx"
    keep.write_text("x", encoding="utf-8")
    va._prune_agent_backups(keep=5, stem="甲")
    left = sorted(p.name for p in tmp_path.iterdir())
    assert len([n for n in left if n.startswith("甲")]) == 5
    assert keep.name in left                    # 別のブックの控えは消さない


def test_agent_materials_warns_about_calc_mode_and_protection():
    """手動計算・保護シートは「書いたのに直らない」の正体。材料に出す。"""
    import vbam_agent as va
    src = open(vh.__file__, encoding='utf-8').read()
    assert "計算モード" in src and "-4135" in src and "手動" in src
    assert "ProtectContents" in src and "書き込みは全部失敗します" in src


def test_agent_build_takes_a_backup_and_asks_before_wiping(tmp_path):
    """build は控えを取らず、AI の overwrite だけで既存シートを全消しできていた。"""
    import vbam_agent as va
    import inspect
    src = inspect.getsource(va.run_build)
    assert "max_turns" in inspect.signature(va.run_build).parameters
    assert "_agent_backup" in src and "_mark_undo_stale" in src
    assert "中身のあるシートを潰す判断は人がします" in src


def test_mcp_agent_passes_changes_and_prune_days(monkeypatch):
    """CLI にだけあった --changes（人が検分する口）と --prune-days を MCP から渡せるようにした（2026-09-10）。"""
    import vba_mcp_server as ms
    import vbam_agent as va
    seen = []
    monkeypatch.setattr(ms, "_submit", lambda line, timeout=600: seen.append(line) or "ok")
    call = ms.agent.fn if hasattr(ms.agent, 'fn') else ms.agent
    call(changes=True)
    call(prune_days=30)
    call(prune_days=30, force=True)
    assert seen == ["agent --changes", "agent --prune-days 30", "agent --prune-days 30 --force"]
    parser = vm.build_parser()
    for line in seen:
        ns, unknown = parser.parse_known_args(line.split())
        assert not unknown and ns.command == "agent", line
    # --mode のエラー文は 4 つ全部を言う（both が抜けていた）
    import pytest
    with pytest.raises(ValueError, match=r'sheet \| build \| macro \| both'):
        va._route("x", forced="all")


def test_the_counts_in_the_help_match_the_code():
    """説明文の本数がコードとずれていた（手順書 20→22 本・手 42→43）。数を書く所はコードから照らす（2026-09-10）。"""
    import vbam_agent as va
    import vbam_recipes as vr
    here = os.path.dirname(va.__file__)
    n, hands = len(vr.SHEET_RECIPES), len(va._ALLOWED_OPS) - len(va._ALIAS_OPS)
    srv = open(os.path.join(here, 'vba_mcp_server.py'), encoding='utf-8').read()
    cli = open(os.path.join(here, 'vba_manager.py'), encoding='utf-8').read()
    assert f"recipes=True で {n} 本の一覧" in srv and f"手 {hands} 本" in srv
    assert f"手順書 {n} 本は --recipes" in cli and f"sheet の定型依頼・{n} 本" in cli
    assert f"この {hands} 個だけ" in va.RULES
    assert all(m in cli for m in va._MODES)                                       # --mode の説明に 4 つとも


# ----------------------------------------------------------------
# 2026-09-04 夜: 監査で「効果のあるもの」だけ直した 4 件
# ----------------------------------------------------------------

def test_undo_restores_formulas_so_they_do_not_become_external_links():
    """1: 控えは別ブック＝貼り付けで ='明細'!B2 が ='[控え.xlsx]明細'!B2 に化ける（実機で再現済み）。

    控えから読んだ .Formula の字面で入れ直せば消える。控えは 5 個で間引かれるので、
    放っておくと #REF! になる（貼り戻しが数式を壊す＝戻すための道具が壊す側だった）。
    """
    import vbam_agent as va

    class FakeRange:
        def __init__(self, nr, nc):
            self.Rows = type("R", (), {"Count": nr})()
            self.Columns = type("C", (), {"Count": nc})()
            self.Formula = None

    class FakeWs:
        def __init__(self, nr, nc):
            self.rng = FakeRange(nr, nc)

        def Range(self, _addr):
            return self.rng

    ws = FakeWs(2, 2)
    snap = {'formulas': [["=明細!B2", 1], ["=SUM(明細!B2:B5)", 2]]}
    assert va._restore_formulas(ws, "A1:B2", snap) is True
    assert ws.rng.Formula == (("=明細!B2", 1), ("=SUM(明細!B2:B5)", 2))
    # 1 セルだけの控えは 2 次元で渡さない（COM が受け取れない）
    ws1 = FakeWs(1, 1)
    assert va._restore_formulas(ws1, "A1", {'formulas': [["=B1"]]}) is True
    assert ws1.rng.Formula == "=B1"
    # 控えが大きすぎて写せなかったときは None（呼び側が警告を出して外部リンクを外しに行く）
    assert va._restore_formulas(ws, "A1:B2", {'formulas': None}) is None
    src = inspect.getsource(va.undo_agent)
    assert "_restore_formulas(ws, addr, src_snap)" in src and "_drop_backup_link(wb, path)" in src


def test_build_fix_loop_requires_zero_errors_not_merely_no_increase():
    """2: 直しのループが材料を読み直す＝壊れた組み上げ後の数が基準になり、何もしなくても合格していた。"""
    import vbam_agent as va
    assert "base=None" in inspect.signature(va.run_agent).parameters['base'].__str__() or \
        va.run_agent.__defaults__ is not None
    src = inspect.getsource(va.run_agent)
    assert "if base is None:" in src
    assert "base=(0, 0)" in inspect.getsource(va.run_build)


def test_macro_mode_returns_ok_not_the_dict():
    """3: dict をそのまま返す＝`ok is not False` が真で、失敗しても終了コード 0 だった。"""
    import vbam_agent as va
    src = inspect.getsource(va._cmd_agent_body)
    assert "return rm['ok']" in src
    assert "return run_macro_agent(" not in src
    assert ({'ok': False} is not False) is True      # dict は必ず「成功」に見える＝これが元の穴


def test_code_replace_zero_match_is_reported_as_a_failure():
    """4: 1 行も当たらなくても成功で返り、AI は「直した」と思って done にしていた。"""
    import vbam_agent as va
    assert vmac._NO_MATCH_RE.search("'x' にマッチする行はありません（置換なし）")
    assert vmac._NO_MATCH_RE.search("'x' は 3 行に一致しましたが、置換後も同じ内容です（変更なし）")
    assert not vmac._NO_MATCH_RE.search("3 行を書き換えました")
    calls = []

    def fake_run_cmd(toks, wb=None):
        calls.append(list(toks))
        if toks[0] == 'code-replace':
            return True, "'a' にマッチする行はありません（置換なし）"
        return True, "コンパイル: 通過"

    old = va._run_cmd
    va._run_cmd = fake_run_cmd
    try:
        res = vmac._execute_macro([{"op": "code_replace", "search": "a", "replace": "b", "module": "shu003"}])
    finally:
        va._run_cmd = old
    label, ok, out = res[0]
    assert ok is False and "1 行も書き換わっていません" in out
    # 書き換わっていないので、道具の自動コンパイルも走らない（＝「直った」の誤解を重ねない）
    assert [c[0] for c in calls] == ['code-replace']


def test_execute_stops_after_a_failed_hand():
    """2: 手は上から順に実行し、1 つ失敗したら残りは実行しない（2026-09-05）。"""
    import vbam_agent as va
    calls = []

    def fake_run_cmd(toks, wb=None):
        calls.append(list(toks))
        if len(calls) == 2:
            return False, "エラー: だめでした"
        return True, "OK"

    old = va._run_cmd
    va._run_cmd = fake_run_cmd
    try:
        res = va._execute([{"op": "write_cells", "cells": {"A1": "x"}},
                           {"op": "write_cells", "cells": {"A2": "y"}},
                           {"op": "write_cells", "cells": {"A3": "z"}}], "名簿")
    finally:
        va._run_cmd = old
    assert len(res) == 3
    assert res[0][1] is True
    assert res[1][1] is False and "だめでした" in res[1][2]
    assert res[2][1] is False and "止めました" in res[2][2]
    assert len(calls) == 2          # 3 本目は実行していない（_run_cmd が呼ばれていない）


def test_execute_runs_every_hand_when_none_fail():
    """2: 全部成功すれば従来どおり最後まで実行される（止める条件に引っかからない）。"""
    import vbam_agent as va
    calls = []

    def fake_run_cmd(toks, wb=None):
        calls.append(list(toks))
        return True, "OK"

    old = va._run_cmd
    va._run_cmd = fake_run_cmd
    try:
        res = va._execute([{"op": "write_cells", "cells": {"A1": "x"}},
                           {"op": "write_cells", "cells": {"A2": "y"}},
                           {"op": "write_cells", "cells": {"A3": "z"}}], "名簿")
    finally:
        va._run_cmd = old
    assert len(res) == 3 and all(r[1] is True for r in res) and len(calls) == 3


def test_execute_macro_stops_after_failed_code_replace_and_skips_compile():
    """2: _execute_macro も同じ型。code_replace が失敗したら後ろの get は実行せず、コンパイルも足さない。"""
    import vbam_agent as va
    calls = []

    def fake_run_cmd(toks, wb=None):
        calls.append(list(toks))
        return False, "エラー: だめでした"

    old = va._run_cmd
    va._run_cmd = fake_run_cmd
    try:
        res = vmac._execute_macro([{"op": "code_replace", "search": "a", "replace": "b", "module": "shu003"},
                                 {"op": "get", "name": "何か"}])
    finally:
        va._run_cmd = old
    assert len(res) == 2
    assert res[0][1] is False and "だめでした" in res[0][2]
    assert res[1][1] is False and "止めました" in res[1][2]
    assert [c[0] for c in calls] == ['code-replace']       # get も compile も呼ばれていない


# ----------------------------------------------------------------
# 2026-09-04 夜（続き）: 残りの実害のあるもの
# ----------------------------------------------------------------

def test_normalize_number_keeps_long_ids_as_text():
    """5: Excel の数値は 15 桁まで。16 桁超を数値にすると下位の桁が黙って変わる（実測）。"""
    import vbam_agent as va
    assert va._normalize_value("12345678901234567890", [('number', None)])[0] == "12345678901234567890"
    assert va._normalize_value("1234-5678-9012-3456-7890", [('number', None)])[0] == "1234-5678-9012-3456-7890"
    # 15 桁までは今までどおり数値にする
    assert va._normalize_value("123456789012345", [('number', None)])[0] == 123456789012345
    assert va._normalize_value("1000", [('number', None)])[0] == 1000
    assert va._normalize_value("0001", [('number', None)])[0] == "0001"


def test_csv_writes_error_cells_as_error_text():
    """6: #N/A が -2146826246 として書き出され、開いた人には元が分からなかった。"""
    import vbam_agent as va
    assert va._csv_text(-2146826246, "%Y-%m-%d") == "#N/A"
    assert va._csv_text(-2146826281, "%Y-%m-%d") == "#DIV/0!"
    assert va._csv_text(1000, "%Y-%m-%d") == "1000"


def test_text_of_keeps_the_time():
    """7: as_text と差分の照合で 9:30 の予定が 0:00 に見えていた。"""
    import datetime
    import vbam_agent as va
    assert va._text_of(datetime.datetime(2024, 4, 1, 9, 30)) == "2024/4/1 9:30"
    assert va._text_of(datetime.datetime(2024, 4, 1, 9, 30, 15)) == "2024/4/1 9:30:15"
    assert va._text_of(datetime.datetime(2024, 4, 1)) == "2024/4/1"


def test_fill_has_a_ceiling_and_guards_the_first_cell():
    """8: 列全体を渡すと 100 万行に膨らみ、先頭セルは検査の外で無条件に潰していた。"""
    import vbam_agent as va
    assert va._FILL_MAX_CELLS == 200000
    src = inspect.getsource(va._do_fill)
    assert "_FILL_MAX_CELLS" in src
    assert "head_used" in src and "n_exist or head_used" in src
    # 先頭に「これから書くのと同じ数式」があるときは上書き扱いにしない（人の数式を伸ばす頼み方）
    assert "str(cur) != str(v)" in src


def test_plan_items_cannot_be_dropped_by_the_next_reply():
    """9: 2 往復目に項目を書き落とすと「未」ごと消えて、落とし物の検査をすり抜けた。"""
    import vbam_agent as va
    old = [{'item': '色分け', 'state': '未'}, {'item': '並べ替え', 'state': '未'}]
    merged = va._merge_plan(old, [{'item': '色分け', 'state': '済'}])      # 並べ替えを書き落とした
    assert merged == [{'item': '色分け', 'state': '済'}, {'item': '並べ替え', 'state': '未'}]
    assert va._plan_pending(merged) == ['並べ替え']
    # 新しく分けた項目は足す
    merged2 = va._merge_plan(merged, [{'item': '並べ替え', 'state': '済'}, {'item': '列幅', 'state': '未'}])
    assert [p['item'] for p in merged2] == ['色分け', '並べ替え', '列幅']
    assert va._plan_pending(merged2) == ['列幅']
    assert "_merge_plan(plan, plan_now)" in inspect.getsource(va.run_agent)


def test_done_with_actions_on_the_last_turn_is_not_thrown_away():
    """10: 最後の往復で手と done を一緒に返すと、手は実行済みなのに失敗扱いだった。"""
    import vbam_agent as va
    assert va._wanted_done('{"done":true,"actions":[{"op":"read","range":"A1"}]}') is True
    assert va._wanted_done('{"done":false,"actions":[]}') is False
    assert va._wanted_done('これは JSON ではありません') is False
    src = inspect.getsource(va.run_agent)
    assert "wanted_done and last_actions" in src


def test_macro_material_says_when_the_code_was_cut():
    """11: 本文が 6000 字で無言で切られ、AI は見えていない後半ごと replace で消していた。"""
    import vbam_agent as va
    assert va._RESULT_LIMIT_BIG > va._RESULT_LIMIT
    assert 'get ' in va._RESULT_BIG_LABELS
    src = inspect.getsource(vmac._macro_materials)
    assert "_RESULT_LIMIT_BIG if label == 'get' or label.startswith('repair')" in src   # 2026-09-17: 材料は repair 1 手
    assert "replace で全文を書き直さないでください" in src


def test_audit_targets_picks_the_written_ranges_only():
    """仕上げ検査を当てる先は、書いた手の番地だけ（read や pivot は拾わない）。"""
    import vbam_agent as va
    acts = [{"op": "write_grid", "range": "F5"}, {"op": "tidy", "ranges": ["A5:F13", "I5:K10"]},
            {"op": "write_cells", "cells": {"C7": "x", "C11": "y"}},
            {"op": "read", "range": "A1:D10"}, {"op": "pivot", "action": "create", "range": "I5:K10"},
            {"op": "copy_range", "src": "A1:B2", "dst": "H7"},
            {"op": "normalize", "range": "C2:C300", "to": "H2"},
            {"op": "write_grid", "range": "ZZZ9999999"}]
    got = va._audit_targets(acts)
    assert got == ["F5", "A5", "I5", "C7", "H7", "H2"]


def test_done_is_refused_while_the_look_is_broken(tmp_path, monkeypatch):
    """仕上げ検査に引っかかるうちは done を受け付けない（1 回だけ差し戻す）。"""
    import vbam_agent as va
    replies = ['{"say":"書く","plan":[{"item":"表を書く","state":"未"}],'
               '"actions":[{"op":"write_grid","range":"A1"}],"done":false}',
               '{"say":"終わり","plan":[{"item":"表を書く","state":"済"}],"actions":[],"done":true,"report":"書いた"}',
               '{"say":"整える","plan":[{"item":"表を書く","state":"済"}],'
               '"actions":[{"op":"tidy","ranges":["A1:C3"]}],"done":false}',
               '{"say":"終わり","plan":[{"item":"表を書く","state":"済"}],"actions":[],"done":true,"report":"整えた"}',
               '{"unmet":[]}']
    sent = []

    def fake_ask(ai, model, key, history):
        sent.append(history[-1][1])
        return replies[min(len(sent) - 1, len(replies) - 1)], {'in': 1, 'out': 1}, 0.0

    calls = {'n': 0}

    def fake_audit(wb, sheet, targets):
        calls['n'] += 1
        return ["罫線がありません（A1:C3）→ tidy を当てる"] if calls['n'] == 1 else []

    monkeypatch.setattr(va, "_ask", fake_ask)
    monkeypatch.setattr(va, "_ai_setup", lambda ai, model: ('gemini', 'm', 'k'))
    monkeypatch.setattr(va, "_execute", lambda a, s, wb=None: [("write_grid A1", True, "ok")])
    monkeypatch.setattr(va, "_audit_now", fake_audit)
    monkeypatch.setattr(va, "_shot_for_ai", lambda s, wb: (None, None))
    monkeypatch.setattr(va, "_LAST_AGENT_LOG_FILE", str(tmp_path / "log.jsonl"))
    monkeypatch.setattr(va, "_LAST_AGENT_ASK_FILE", str(tmp_path / "ask.txt"))
    wb = type("WB", (), {"Name": "b.xlsx", "Path": ""})()
    r = va.run_agent("表を書いて", "名簿", wb, max_turns=6, materials="使用範囲: A1",
                     verify=False, backup=False, show_image=False)
    assert r['done'] is True
    assert any("仕上げ検査" in m and "罫線" in m for m in sent), sent
    # 検査は 2 回（done の関所で 1 回＋終わりの報告で 1 回）。差し戻しは 1 回だけで、
    # 報告に出るのは直した後の姿（関所で拾った古い指摘をそのまま出さない）
    assert calls['n'] == 2 and r['audit'] == []


def test_grade_turn_blocks_done_until_the_unmet_points_are_fixed(tmp_path, monkeypatch):
    """自己採点: 満たしていない点を挙げたら、その分を直すまで done にしない。"""
    import vbam_agent as va
    replies = ['{"say":"色分け","plan":[{"item":"色分け","state":"未"}],'
               '"actions":[{"op":"cond_format","range":"A1:C3"}],"done":false}',
               '{"say":"終わり","plan":[{"item":"色分け","state":"済"}],"actions":[],"done":true,"report":"色分けした"}',
               '{"unmet":["金額の多い順に並べ替えていない"],"note":"並べ替えが残っている"}',
               '{"say":"並べ替え","plan":[{"item":"色分け","state":"済"},'
               '{"item":"金額の多い順に並べ替えていない","state":"未"}],'
               '"actions":[{"op":"sort","range":"A1:C3","key":"C","desc":true}],"done":false}',
               '{"say":"終わり","plan":[{"item":"色分け","state":"済"},'
               '{"item":"金額の多い順に並べ替えていない","state":"済"}],'
               '"actions":[],"done":true,"report":"色分けと並べ替え"}']
    sent = []

    def fake_ask(ai, model, key, history):
        sent.append(history[-1][1])
        return replies[min(len(sent) - 1, len(replies) - 1)], {'in': 1, 'out': 1}, 0.0

    monkeypatch.setattr(va, "_ask", fake_ask)
    monkeypatch.setattr(va, "_ai_setup", lambda ai, model: ('gemini', 'm', 'k'))
    monkeypatch.setattr(va, "_execute", lambda a, s, wb=None: [("cond_format", True, "ok")])
    monkeypatch.setattr(va, "_audit_now", lambda wb, sheet, t: [])
    monkeypatch.setattr(va, "_shot_for_ai", lambda s, wb: (None, None))
    monkeypatch.setattr(va, "_LAST_AGENT_LOG_FILE", str(tmp_path / "log.jsonl"))
    monkeypatch.setattr(va, "_LAST_AGENT_ASK_FILE", str(tmp_path / "ask.txt"))
    wb = type("WB", (), {"Name": "b.xlsx", "Path": ""})()
    r = va.run_agent("色分けして多い順に並べて", "名簿", wb, max_turns=6, materials="使用範囲: A1",
                     verify=False, backup=False, show_image=False)
    assert r['done'] is True and r['graded'] is True
    assert any("【採点】" in m for m in sent)
    assert any("満たしていない点が挙がりました" in m for m in sent)
    assert [p['item'] for p in r['plan']] == ["色分け", "金額の多い順に並べ替えていない"]
    sent.clear()
    r2 = va.run_agent("色分けして", "名簿", wb, max_turns=6, materials="使用範囲: A1",
                      verify=False, backup=False, show_image=False, grade=False)
    assert r2['graded'] is False and not any("【採点】" in m for m in sent)


def test_code_violations_catch_the_collateral_rewrite():
    """マクロ側の「頼んでいない変化」（2026-09-06）。

    code_replace は**そのモジュールの中で search に当たる行を全部**書き換える。
    On Error Resume Next のような行を search にすると隣のプロシージャまで変わり、
    コンパイルは通るので誰も気づかない。撃つ前の全コードと照らして道具が数える。
    """
    import vbam_agent as va
    code = ("Option Explicit\n"
            "Sub 直す対象()\n  Dim n As String\n  n = 1\nEnd Sub\n"
            "Sub 別のマクロ()\n  On Error Resume Next\n  MsgBox 1\nEnd Sub\n")
    procs = vmac._split_procs(code)
    assert set(procs) == {"(宣言部)", "直す対象", "別のマクロ"}
    assert "Dim n As String" in procs["直す対象"] and "MsgBox 1" in procs["別のマクロ"]
    before = {("M1", k): v for k, v in procs.items()}
    # 直す対象だけを直した＝違反なし
    after = dict(before)
    after[("M1", "直す対象")] = "Sub 直す対象()\n  Dim n As Long\n  n = 1\nEnd Sub"
    assert vmac._code_violations(before, after, allowed=["直す対象"]) == []
    # 隣のプロシージャまで変わった／消えた／勝手に増えた
    after2 = dict(after)
    after2[("M1", "別のマクロ")] = "Sub 別のマクロ()\n  MsgBox 1\nEnd Sub"
    bad = vmac._code_violations(before, after2, allowed=["直す対象"])
    assert bad == ["頼んでいないマクロが変わった: M1.別のマクロ"]
    after3 = {k: v for k, v in after.items() if k != ("M1", "別のマクロ")}
    assert "消えた: M1.別のマクロ" in vmac._code_violations(after, after3, allowed=["直す対象"])[0]
    after4 = dict(after)
    after4[("M1", "勝手に足した")] = "Sub 勝手に足した()\nEnd Sub"
    assert "増えた: M1.勝手に足した" in vmac._code_violations(after, after4, allowed=["直す対象"])[0]
    # 宣言部（Option Explicit を足す等）も咎める＝家の決まりが機械的に効く
    after5 = dict(after)
    after5[("M1", "(宣言部)")] = "Option Explicit\nPublic 余計 As Long"
    assert vmac._code_violations(before, after5, allowed=["直す対象"]) == [
        "頼んでいないマクロが変わった: M1.(宣言部)"]
    # 控えが取れなかった（読めないブック）ときは黙って通す＝関所を落とし穴にしない
    assert vmac._code_violations(None, after, allowed=[]) == []


def test_macro_loop_has_the_same_two_gates_as_the_sheet_loop():
    """macro のループにも plan の関所とコードの不変条件がある（2026-09-06・配線の見張り番）。"""
    import inspect as _inspect
    import vbam_agent as va
    src = _inspect.getsource(vmac.run_macro_agent)
    assert "va._merge_plan(plan, va._parse_plan(d.get('plan')))" in src, "plan を読んでいない"   # macro は vbam_macro（va. 経由・2026-09-11）
    assert "未の項目が残っているので done を受け付けません" in src, "落とし物の関所が無い"
    assert "code_before = _code_snapshot(wb)" in src, "撃つ前の控えを取っていない"
    assert src.index("code_before = _code_snapshot(wb)") < src.index("results = _execute_macro("), \
        "控えは最初の書き換えの前に取る"
    assert "ok = ok_c and not code_notes" in src, "頼んでいない変化が残ったまま合格にしている"
    # AI 側の規約にも書く（道具だけ厳しくして、AI に知らせないのは往復の無駄）
    assert "plan" in vmac.MACRO_RULES and "「未」が 1 つでも残っている done" in vmac.MACRO_RULES
    assert "直すと言った 1 本以外のコードを変えない" in vmac.MACRO_RULES


def test_grading_is_a_separate_conversation(tmp_path, monkeypatch):
    """採点は別の会話で聞く（2026-09-06）。

    前は同じ会話の続きに質問を足していた＝自分の手順も言い分も全部見た相手が採点していて、
    追認しやすい作りだった。渡すのは依頼文・道具が読み直した現物・画像だけ。
    """
    import vbam_agent as va
    replies = ['{"say":"終わり","plan":[{"item":"整える","state":"済"}],"actions":[],"done":true,"report":"整えた"}',
               '{"unmet":[]}']
    seen = []

    def fake_ask(ai, model, key, history):
        seen.append({'ai': ai, 'model': model, 'n': len(history), 'last': history[-1][1]})
        return replies[min(len(seen) - 1, len(replies) - 1)], {'in': 1, 'out': 1}, 0.0

    monkeypatch.setattr(va, "_ask", fake_ask)
    monkeypatch.setattr(va, "_ai_setup", lambda ai, model: (ai or 'gemini', model or 'm', 'k'))
    monkeypatch.setattr(va, "_audit_now", lambda wb, sheet, t: [])
    monkeypatch.setattr(va, "_shot_for_ai", lambda s, wb: (None, None))
    monkeypatch.setattr(va, "_grade_materials", lambda s, wb, extra=(): "使用範囲: A1:C3（読み直した現物）")
    monkeypatch.setattr(va, "_LAST_AGENT_LOG_FILE", str(tmp_path / "log.jsonl"))
    monkeypatch.setattr(va, "_LAST_AGENT_ASK_FILE", str(tmp_path / "ask.txt"))
    wb = type("WB", (), {"Name": "b.xlsx", "Path": ""})()
    r = va.run_agent("整えて", "名簿", wb, max_turns=4, materials="使用範囲: A1",
                     verify=False, backup=False, show_image=False, grade_ai="claude")
    assert r['graded'] is True and r['done'] is True
    work, grading = seen[0], seen[1]
    assert grading['n'] == 1, "採点が作業の会話を引き継いでいる（別の会話で聞く）"
    assert "【採点】" in grading['last'] and "読み直した現物" in grading['last']
    # 作業側の言い分（report・plan・手順）は採点係に渡さない
    assert "整えた" not in grading['last'] and work['last'] not in grading['last']
    # --grade-ai で採点だけ別の AI に振れる
    # 2026-09-09: 作業側の既定は claude-code（実射も対話も。金のかかる頭は名指ししたときだけ）
    assert work['ai'] == 'claude-code' and grading['ai'] == 'claude'


def test_grade_materials_does_not_reset_the_job_clock(monkeypatch):
    """採点のために materials を読み直しても、仕事の時計は押し直さない（経過が嘘になる）。"""
    import inspect as _inspect
    import vbam_agent as va
    src = _inspect.getsource(va._grade_materials)
    assert "job_clock_get()" in src and "job_clock_set(clock)" in src
    assert src.index("job_clock_get()") < src.index("_run_cmd(['materials'")


def test_gates_count_every_pushback_and_split_the_pass(tmp_path, monkeypatch):
    """関所の通信簿（2026-09-06）: どの関所で何回差し戻したかを数え、合格を 3 段に割る。

    「59/59 合格」だけでは、AI が一発で正しかったのか・関所が毎回直したのかが分からない。
    """
    import vbam_agent as va
    # 1 往復目に規約外の返事 → 2 往復目に手 → 3 往復目に done（採点が 1 件挙げる）→ 4 往復目で直して done
    replies = ['これは JSON ではない',
               '{"say":"書く","plan":[{"item":"色分け","state":"未"}],'
               '"actions":[{"op":"cond_format","range":"A1:C3"}],"done":false}',
               '{"say":"終わり","plan":[{"item":"色分け","state":"済"}],"actions":[],"done":true,"report":"色分け"}',
               '{"unmet":["並べ替えていない"]}',
               '{"say":"並べる","plan":[{"item":"色分け","state":"済"},'
               '{"item":"並べ替えていない","state":"済"}],"actions":[],"done":true,"report":"両方"}']
    sent = []

    def fake_ask(ai, model, key, history):
        sent.append(history[-1][1])
        return replies[min(len(sent) - 1, len(replies) - 1)], {'in': 1, 'out': 1}, 0.0

    monkeypatch.setattr(va, "_ask", fake_ask)
    monkeypatch.setattr(va, "_ai_setup", lambda ai, model: ('gemini', 'm', 'k'))
    monkeypatch.setattr(va, "_execute", lambda a, s, wb=None: [("cond_format", False, "失敗")])
    monkeypatch.setattr(va, "_audit_now", lambda wb, sheet, t: [])
    monkeypatch.setattr(va, "_shot_for_ai", lambda s, wb: (None, None))
    monkeypatch.setattr(va, "_LAST_AGENT_LOG_FILE", str(tmp_path / "log.jsonl"))
    monkeypatch.setattr(va, "_LAST_AGENT_ASK_FILE", str(tmp_path / "ask.txt"))
    wb = type("WB", (), {"Name": "b.xlsx", "Path": ""})()
    r = va.run_agent("色分けして並べて", "名簿", wb, max_turns=6, materials="使用範囲: A1",
                     verify=False, backup=False, show_image=False)
    g = r['gates']
    assert g['format'] == 1, "規約外の返事を数えていない"
    assert g['hand'] == 1, "手の失敗を数えていない"
    assert g['grade'] == 1, "自己採点の差し戻しを数えていない"
    assert va._gate_total(g) == 3
    # 1 行に畳める・3 段に割れる
    assert "返事の形 1" in va._gate_line(g) and "自己採点 1" in va._gate_line(g)
    assert va._gate_line({}) == "関所の差し戻し: なし（一発）"
    assert va._gate_grade(True, {}) == '一発合格'
    assert va._gate_grade(True, g) == '関所で直して合格'
    assert va._gate_grade(False, {}) == '不合格'


def test_first_shot_line_reads_the_inside_of_the_pass():
    """点数の中身: 一発合格 / 関所で直して合格 / 不合格 を数える（古い 3 つ組も受ける）。"""
    import vbam_agent as va
    rows = [("A", True, "ok", {'turns': 1, 'gates': {'audit': 0}}),
            ("B", True, "ok", {'turns': 3, 'gates': {'audit': 1, 'grade': 1}}),
            ("C", False, "だめ", {'turns': 4, 'gates': {'inv': 2}})]
    line = vf._first_shot_line(rows)
    assert "一発合格 1" in line and "関所で直して合格 1" in line and "不合格 1" in line
    assert "仕上げ検査 1" in line and "頼んでいない変化 2" in line
    # 古い 3 つ組（関所の記録が無い）でも落ちない
    assert vf._fire_rows([("A", True, "ok")]) == [("A", True, "ok", {})]
    assert "数えていません" in vf._first_shot_line([("A", True, "ok")])
    assert "弾なし" in vf._first_shot_line([])


def test_runs_ledger_records_the_real_work_and_the_undo(tmp_path, monkeypatch):
    """本番の走行台帳（2026-09-06）: 1 走行 1 行・記録も控える・undo で印が付く。

    実射の点数は「私が作った問題の点数」。実運用は本番の走行でしか測れず、いちばん正直な
    失敗率は「人が戻した割合（undo 率）」になる。
    """
    import vbam_agent as va
    replies = ['{"say":"書く","plan":[{"item":"整える","state":"未"}],'
               '"actions":[{"op":"tidy","ranges":["A1:C3"]}],"done":false}',
               '{"say":"終わり","plan":[{"item":"整える","state":"済"}],"actions":[],"done":true,"report":"整えた"}']
    sent = []

    def fake_ask(ai, model, key, history):
        sent.append(history[-1][1])
        return replies[min(len(sent) - 1, len(replies) - 1)], {'in': 7, 'out': 3}, 0.0

    monkeypatch.setattr(va, "_ask", fake_ask)
    monkeypatch.setattr(va, "_ai_setup", lambda ai, model: ('gemini', 'm', 'k'))
    monkeypatch.setattr(va, "_execute", lambda a, s, wb=None: [("tidy", True, "ok")])
    monkeypatch.setattr(va, "_audit_now", lambda wb, sheet, t: [])
    monkeypatch.setattr(va, "_shot_for_ai", lambda s, wb: (None, None))
    monkeypatch.setattr(va, "_LAST_AGENT_LOG_FILE", str(tmp_path / "log.jsonl"))
    monkeypatch.setattr(va, "_LAST_AGENT_ASK_FILE", str(tmp_path / "ask.txt"))
    wb = type("WB", (), {"Name": "台帳.xlsx", "Path": "", "Sheets": staticmethod(lambda n: None)})()
    # 実射（backup=False）は台帳に残さない＝練習台の走行を実戦の記録に混ぜない
    va.run_agent("整えて", "名簿", wb, max_turns=4, materials="使用範囲: A1",
                 verify=False, backup=False, show_image=False, grade=False)
    assert os.path.exists(va._AGENT_RUNS_FILE) is False
    # 本番（backup=True）は 1 行残す。控えは偽物にしてブックを触らせない
    monkeypatch.setattr(va, "_agent_backup", lambda *a, **k: None)
    monkeypatch.setattr(va, "_sheet_snapshot", lambda ws: {})
    monkeypatch.setattr(va, "_book_cell_count", lambda wb_: 10)
    monkeypatch.setattr(va, "_snapshot_book", lambda wb_, own=None: {'sheets': {}, 'names': []})
    monkeypatch.setattr(va, "_inv_violations", lambda *a, **k: [])
    monkeypatch.setattr(va, "_changes_of", lambda *a, **k: [])
    sent.clear()
    r = va.run_agent("整えて", "名簿", wb, max_turns=4, materials="使用範囲: A1",
                     verify=False, backup=True, show_image=False, grade=False)
    rows = va._runs_load()
    assert len(rows) == 1
    rec = rows[0]
    assert rec['book'] == "台帳.xlsx" and rec['sheet'] == "名簿" and rec['mode'] == 'sheet'
    # 最後が tidy の往復は done の返事を待たずに検査へ進む（2026-09-11）＝AI に聞くのは 1 往復（7 トークン）だけ
    assert rec['done'] is True and rec['run_id'] == r['run_id'] and rec['in'] == 7
    assert rec["log"] and os.path.exists(rec["log"]), "走行ごとの記録が控えられていない"
    assert 'undone' not in rec
    # --undo が「人が戻した」の印を付ける（台帳は追記だけ＝読むときに畳む）
    va.runs_append({'undo': r['run_id'], 'time': '2026-09-06 10:00'})
    rows = va._runs_load()
    assert len(rows) == 1 and rows[0]['undone'] == '2026-09-06 10:00'
    line = va._runs_summary(rows)
    assert "1 走行" in line and "人が戻した（undo）1/1" in line and "一発 1/1" in line


def test_keep_case_turns_a_real_run_into_a_bullet(tmp_path, monkeypatch, capsys):
    """本番の 1 走行を弾にする（2026-09-06）。依頼文は記録から、練習台はそのブックの写し取り。

    練習台は全部こちらの想像で組んだもので、撃つ依頼まで想像だった。ここが実戦から弾を供給する口。
    """
    import vbam_agent as va
    log = tmp_path / "log.jsonl"
    log.write_text(json.dumps({'meta': {'book': "件名簿.xlsm", 'sheet': "2024", 'mode': 'sheet',
                                        'request': "観光関係一般の分だけ抜き出して。"}},
                              ensure_ascii=False) + "\n"
                   + json.dumps({'turn': 1, 'prompt': 'p', 'reply': 'r'}, ensure_ascii=False) + "\n",
                   encoding='utf-8')
    monkeypatch.setattr(va, "_LAST_AGENT_LOG_FILE", str(log))
    monkeypatch.setattr(vf, "_AGENT_CASES_DIR", str(tmp_path / "cases"))
    src = tmp_path / "件名簿.xlsm"
    src.write_text("dummy", encoding='utf-8')
    wb = type("WB", (), {"Name": "件名簿.xlsm", "Path": str(tmp_path),
                         "FullName": str(src), "Saved": True})()
    monkeypatch.setattr(va, "get_workbook", lambda t=None: (None, wb))
    harvested = []
    monkeypatch.setattr(vf, "harvest_book", lambda p, o=None, k=1, seed=None: harvested.append((p, o)) or o)
    assert vf.keep_case("観光だけ抜き出す") is True
    # 元のブックは触らない（写し取りの元は元ファイル・出力は弾の置き場）
    assert harvested and harvested[0][0] == str(src)
    assert harvested[0][1].startswith(str(tmp_path / "cases"))
    saved = vf._cases_load()
    assert list(saved) == ["観光だけ抜き出す"]
    c = saved["観光だけ抜き出す"]
    assert c['sheet'] == "2024" and "観光関係一般" in c['request'] and c['from_book'] == "件名簿.xlsm"
    # 一覧に出る
    vf.cases_list()
    assert "観光だけ抜き出す" in capsys.readouterr().out
    # 別のブックを開いたまま登録しようとしたら断る（違うブックを写し取らない）
    wb2 = type("WB", (), {"Name": "別.xlsm", "Path": str(tmp_path), "FullName": str(src), "Saved": True})()
    monkeypatch.setattr(va, "get_workbook", lambda t=None: (None, wb2))
    assert vf.keep_case("だめな弾") is False
    assert "だめな弾" not in vf._cases_load()


def test_keep_case_refuses_what_it_cannot_replay(tmp_path, monkeypatch):
    """弾にできるのはシートの走行だけ・保存していないブックは断る。"""
    import vbam_agent as va
    log = tmp_path / "log.jsonl"
    monkeypatch.setattr(va, "_LAST_AGENT_LOG_FILE", str(log))
    monkeypatch.setattr(vf, "_AGENT_CASES_DIR", str(tmp_path / "cases"))
    log.write_text(json.dumps({'meta': {'book': "b.xlsm", 'sheet': None, 'mode': 'macro',
                                        'request': "マクロを直して"}}, ensure_ascii=False) + "\n",
                   encoding='utf-8')
    assert vf.keep_case("マクロの弾") is False          # macro の走行は弾にしない
    log.write_text(json.dumps({'meta': {'book': "b.xlsm", 'sheet': "名簿", 'mode': 'sheet',
                                        'request': "整えて"}}, ensure_ascii=False) + "\n",
                   encoding='utf-8')
    wb = type("WB", (), {"Name": "b.xlsm", "Path": "", "FullName": "b.xlsm", "Saved": False})()
    monkeypatch.setattr(va, "get_workbook", lambda t=None: (None, wb))
    assert vf.keep_case("未保存の弾") is False          # 保存していない＝写し取る元が無い
    assert vf._cases_load() == {}


def test_fire_mine_shoots_the_saved_bullets_with_the_bed_judgement(tmp_path, monkeypatch, capsys):
    """--fire mine: 登録した弾を、写し取りと同じ判定（正解表なし）で撃つ。"""
    import vbam_agent as va
    bed = tmp_path / "bed.xlsm"
    bed.write_text("x", encoding='utf-8')
    vf._cases_save({"抜き出す": {'name': "抜き出す", 'bed': str(bed), 'sheet': "2024",
                                 'request': "観光関係一般の分だけ抜き出して。", 'mode': 'sheet',
                                 'from_book': "件名簿.xlsm", 'time': '2026-09-06 09:00'}})
    closed = []
    wb = type("WB", (), {"Name": "bed.xlsm", "Close": staticmethod(lambda SaveChanges=True: closed.append(1))})()
    monkeypatch.setattr(va, "get_workbook", lambda p=None: (None, wb))
    seen = {}

    def fake_case(book, sheet, request, ai, model, twice=True, max_turns=4):
        seen.update(sheet=sheet, request=request, twice=twice)
        return True, "壊さず・指摘は増えず", {'turns': 2, 'gates': {'audit': 1}}

    monkeypatch.setattr(vf, "_fire_bed_case", fake_case)
    rows = vf._fire_mine(None, 'gemini', 'm')
    assert seen['sheet'] == "2024" and "観光関係一般" in seen['request']
    assert closed, "弾の練習台を閉じていない（次に撃つときは同じ姿から）"
    assert len(rows) == 1 and rows[0][1] is True and rows[0][3]['gates'] == {'audit': 1}
    assert "本番の弾" in rows[0][0]
    # 練習台のファイルが消えていたら、黙って合格にしない
    vf._cases_save({"消えた": {'name': "消えた", 'bed': str(tmp_path / "no.xlsm"), 'sheet': "S",
                               'request': "整えて", 'mode': 'sheet', 'time': '2026-09-06 09:00'}})
    rows = vf._fire_mine(None, 'gemini', 'm')
    assert rows[0][1] is False and "ありません" in rows[0][2]


def test_fire_mine_is_opt_in_and_named():
    """mine は組の名前として通り、名前を書かない全部撃ちには黙って混ざらない。"""
    import inspect as _inspect
    import vbam_agent as va
    assert 'mine' in vf._FIRE_GROUPS
    src = _inspect.getsource(vf.fire_agent) + _inspect.getsource(vf._fire_agent_body)
    assert "want_mine = 'mine' in names" in src
    assert "if want_mine or mine_pick:" in src
    ns = vm.build_parser().parse_args(["agent", "--keep-case", "観光だけ"])
    assert ns.keep_case == "観光だけ"
    assert vm.build_parser().parse_args(["agent", "--cases"]).cases is True


def test_undo_marks_the_run_it_rolled_back():
    """--undo は控えの覚書にある run_id を台帳へ書き戻す（配線の見張り番）。"""
    import inspect as _inspect
    import vbam_agent as va
    assert "'run_id': run_id" in _inspect.getsource(va._agent_backup), "控えに走行の名札が無い"
    assert "run_id" in _inspect.signature(va._agent_backup).parameters
    src = _inspect.getsource(va.undo_agent)
    assert "runs_append({'undo': meta['run_id']" in src, "戻したことを台帳に残していない"
    # 実射・build の直しループは台帳に残さない（実戦の記録を練習で薄めない）
    assert "if backup and not dry_run:" in _inspect.getsource(va.run_agent)
    assert "ledger=False" in _inspect.getsource(vf._fire_macro_cases)


def test_grade_reply_is_read_loosely():
    """採点の返事が読めない・空でも落ちない（採点は関所であって落とし穴にしない）。"""
    import vbam_agent as va
    assert va._parse_unmet('{"unmet":["A","B"]}') == ["A", "B"]
    assert va._parse_unmet('{"unmet":[{"item":"C"}]}') == ["C"]
    assert va._parse_unmet('{"unmet":[]}') == []
    assert va._parse_unmet('これは JSON ではない') == []
    assert va._parse_unmet('{"note":"満たしている"}') == []
    assert len(va._parse_unmet('{"unmet":[' + ",".join('"x"' for _ in range(20)) + ']}')) == 8


def test_score_history_records_and_finds_the_regression(tmp_path, monkeypatch):
    """点数: 撃つたびに記録し、前回と比べて落ちた弾（退行）を出す。"""
    import vbam_agent as va
    monkeypatch.setattr(vf, "_AGENT_SCORE_FILE", str(tmp_path / "score.jsonl"))
    vf._score_save({'time': '2026-09-04 10:00', 'ai': 'gemini', 'model': 'm', 'n': 3, 'ok': 3,
                    'sec': 10, 'cases': {'A': True, 'B': True, 'C': True}})
    rec = {'time': '2026-09-04 12:00', 'ai': 'gemini', 'model': 'm', 'n': 3, 'ok': 2, 'sec': 9,
           'cases': {'A': True, 'B': False, 'C': True}}
    prev, worse, better = vf._score_compare(rec, vf._score_load())
    assert prev['ok'] == 3 and worse == ['B'] and better == []
    vf._score_save(rec)
    rec2 = dict(rec, time='2026-09-04 13:00', ok=3, cases={'A': True, 'B': True, 'C': True})
    prev2, worse2, better2 = vf._score_compare(rec2, vf._score_load())
    assert worse2 == [] and better2 == ['B']
    assert vf.score_history() is True
    other = {'time': 't', 'n': 1, 'ok': 1, 'cases': {'Z': True}}
    assert vf._score_compare(other, vf._score_load())[0] is None


def test_notes_remember_the_last_jobs_per_sheet(tmp_path, monkeypatch):
    """覚書: ブック・シートごとに直近 3 件だけ残し、材料に足す文にする。"""
    import vbam_agent as va
    monkeypatch.setattr(vg, "_AGENT_NOTES_FILE", str(tmp_path / "notes.json"))
    assert va._notes_recall("b.xlsx", "名簿") == ""
    for i in range(5):
        va._notes_remember("b.xlsx", "名簿", f"依頼{i}", f"報告{i}")
    va._notes_remember("b.xlsx", "別表", "こっちの依頼", "こっちの報告")
    text = va._notes_recall("b.xlsx", "名簿")
    assert "依頼4" in text and "依頼2" in text and "依頼1" not in text
    assert "こっちの依頼" not in text
    assert "材料を優先" in text
    assert va._notes_recall("b.xlsx", "別表").count("こっちの依頼") == 1


def test_vague_cases_are_on_the_bench_with_machine_checks():
    """曖昧な依頼の弾 5 本が台に載っていて、答え合わせが全部「実物を見る関数」であること。"""
    import vbam_agent as va
    vague = [c for c in vf.FIRE_CASES if c.get('kind') == 'vague']
    assert [c['name'] for c in vague] == ["見やすくして", "未入力が分かるように", "合計を出して",
                                          "多い順に並べて", "担当ごとの件数", "部署の列を足す",
                                          "印刷したら切れる", "鈴木さんの分だけ",
                                          "金額が足せない", "同じ人が二重"]
    for c in vague:
        assert callable(c['check']) and callable(c['build'])
        assert len(c['request']) < 40
    assert 'vague' in vf._FIRE_GROUPS
    picked, unknown = vf._select_fire_cases(vf.FIRE_CASES, ['vague'])
    assert not unknown and [c['name'] for c in picked] == [c['name'] for c in vague]


def test_notes_go_into_the_materials_and_are_written_back(tmp_path, monkeypatch):
    """覚書の配線: 材料に前回の依頼が足され、終わった仕事が（保存済みブックのときだけ）覚えられる。"""
    import vbam_agent as va
    monkeypatch.setattr(vg, "_AGENT_NOTES_FILE", str(tmp_path / "notes.json"))
    va._notes_remember("b.xlsx", "名簿", "郵便番号を統一して", "C 列を 000-0000 に直した")

    sent = []

    def fake_ask(ai, model, key, history):
        sent.append(history[-1][1])
        return '{"say":"終わり","plan":[],"actions":[],"done":true,"report":"何もしていない"}', {'in': 1, 'out': 1}, 0.0

    monkeypatch.setattr(va, "_ask", fake_ask)
    monkeypatch.setattr(va, "_ai_setup", lambda ai, model: ('gemini', 'm', 'k'))
    monkeypatch.setattr(va, "_run_cmd", lambda tokens, wb=None: (True, "使用範囲: A1:E9"))
    monkeypatch.setattr(va, "_extra_materials", lambda wb, ws: "")
    monkeypatch.setattr(va, "_LAST_AGENT_LOG_FILE", str(tmp_path / "log.jsonl"))
    monkeypatch.setattr(va, "_LAST_AGENT_ASK_FILE", str(tmp_path / "ask.txt"))

    class _WB:
        Name = "b.xlsx"
        Path = r"C:\somewhere"

        def Sheets(self, name):
            return object()

    r = va.run_agent("続きをやって", "名簿", _WB(), max_turns=2, verify=False, backup=False,
                     show_image=False, grade=False)
    assert r['done'] is True
    assert "このシートで前にやったこと" in sent[0]
    assert "郵便番号を統一して" in sent[0]          # 前回の依頼が材料に入っている
    # 終わった仕事が覚えられている（保存済みブック＝Path があるとき）
    text = va._notes_recall("b.xlsx", "名簿")
    assert "続きをやって" in text and "郵便番号を統一して" in text
    # 未保存のブック（Path が空）は覚えない＝実射の使い捨てブックで覚書が汚れない
    class _New(_WB):
        Name = "Book1"
        Path = ""
    va.run_agent("使い捨て", "名簿", _New(), max_turns=2, verify=False, backup=False,
                 show_image=False, grade=False)
    assert va._notes_recall("Book1", "名簿") == ""


def test_agent_changes_ledger(tmp_path, monkeypatch, capsys):
    """変わったセルの明細（2026-09-05）。

    それまで報告に出るのは「変わったセル: N 個」＋例 8 個だけで、**何が何になったかを
    戻さずに確かめる道が無かった**。番地・前・後を 1 セル 1 行で出し、全部をファイルに残す。
    """
    import vbam_agent as va

    class FakeWS:
        def __init__(self, values, formulas=None, row=1, col=1, name="名簿"):
            self._v, self._f, self._r, self._c, self.Name = values, formulas or values, row, col, name

        @property
        def UsedRange(self):
            ws = self

            class UR:
                Row, Column = ws._r, ws._c
                Value = tuple(tuple(r) for r in ws._v)
                Formula = tuple(tuple(r) for r in ws._f)
                Address = "$A$1"

                class Cells:
                    CountLarge = 100
            return UR
        Shapes = ()

    before = {'addr': 'A1:B2', 'row': 1, 'col': 1, 'values': [["a", 1], ["b", 2]],
              'formulas': [["a", 1], ["b", 2]], 'shapes': []}
    rows = va._changes_of(before, FakeWS([["a", 1], ["b", 99]]))
    assert rows == [{'sheet': "名簿", 'addr': "B2", 'before': "2", 'after': "99",
                     'before_formula': "2", 'after_formula': "99"}]
    # 数式に変わったセルは、後ろに数式の字面を出す（値だけでは検分にならない）
    rows = va._changes_of(before, FakeWS([["a", 1], ["b", 2]], formulas=[["a", 1], ["b", "=1+1"]]))
    assert len(rows) == 1 and va._change_side(rows[0], 'after') == "=1+1"
    assert va._change_side(rows[0], 'before') == "2"
    # 数えられないとき（大きすぎる・控えが無い）は None。数と例を返す旧い口はそのまま
    assert va._changes_of({'too_big': 999999}, FakeWS([["a"]])) is None
    assert va._changes_of(None, FakeWS([["a"]])) is None
    assert va._diff_sheet(before, FakeWS([["a", 1], ["b", 99]])) == (1, ["B2 '2'→'99'"])
    # 明細はファイルに残る（1 セル 1 行・見出しつき・Excel で開ける）
    path = str(tmp_path / "changes.tsv")
    many = [{'sheet': "名簿", 'addr': f"A{i}", 'before': "", 'after': str(i),
             'before_formula': "", 'after_formula': ""} for i in range(1, 31)]
    fmt = [{'sheet': "名簿", 'addr': "C列", 'what': "表示形式", 'before': "G/標準", 'after': "#,##0"},
           {'sheet': "名簿", 'addr': "A1:E9", 'what': "罫線・内・横", 'before': "なし", 'after': "あり"}]
    assert va._write_changes_file(many, path, fmt_rows=fmt) == path
    body = open(path, encoding='utf-8-sig').read().splitlines()
    assert body[0].split("\t") == ["種類", "シート", "番地", "前", "後", "前の数式", "後の数式"]
    assert len(body) == 33 and body[1].split("\t")[:5] == ["値", "名簿", "A1", "", "1"]
    assert body[-1].split("\t")[:5] == ["書式・罫線・内・横", "名簿", "A1:E9", "なし", "あり"]
    # 書式の変化: 列ごと（幅・表示形式・寄せ・太字）と範囲の罫線を、控えの姿と突き合わせる
    before_fmt = {'addr': "A1:B2", 'cols': {'A': {'width': "8.00", 'numfmt': "G/標準", 'halign': "標準", 'bold': "並"},
                                            'B': {'width': "8.00", 'numfmt': "G/標準", 'halign': "標準", 'bold': "並"}},
                  'borders': {'内・横': "なし", '下': "あり"}}

    class FmtWS:
        Name = "名簿"

        class _Col:
            def __init__(self, numfmt, halign, bold):
                self.NumberFormatLocal, self.HorizontalAlignment = numfmt, halign
                self.Font = type("F", (), {'Bold': bold})

        @property
        def UsedRange(self):
            ws = self

            class UR:
                Row, Column, Address = 1, 1, "$A$1:$B$2"

                class Columns:
                    Count = 2

                    @staticmethod
                    def __call__(j):
                        return ws._cols[j - 1]

                @staticmethod
                def Borders(idx):
                    return type("B", (), {'LineStyle': 1 if idx == 12 else -4142})
            UR.Columns = type("C", (), {'Count': 2, '__call__': lambda s, j: ws._cols[j - 1]})()
            return UR

        def __init__(self):
            self._cols = [self._Col("#,##0", -4152, True), self._Col("G/標準", 1, False)]
            self.Columns = lambda i: type("W", (), {'ColumnWidth': 12.5})

    changes = va._format_changes(before_fmt, FmtWS())
    got = {(c['addr'], c['what']): (c['before'], c['after']) for c in changes}
    assert got[("A列", "表示形式")] == ("G/標準", "#,##0")
    assert got[("A列", "寄せ")] == ("標準", "右") and got[("A列", "太字")] == ("並", "太字")
    assert got[("A列", "列幅")] == ("8.00", "12.50")
    assert got[("A1:B2", "罫線・内・横")] == ("なし", "あり")
    assert ("A1:B2", "罫線・下") in got and got[("A1:B2", "罫線・下")] == ("あり", "なし")
    assert va._format_changes(None, FmtWS()) == []
    # 画面には頭だけ出して、残りは行数で言う（往復の報告を明細で埋めない）
    table = va._changes_table(many, show=12)
    assert len(table) == 14 and "…ほか 18 行" in table[-1]
    assert va._changes_table(many, show=12, multi_sheet=True)[0].startswith("  シート")
    # 後から見る口（Excel に触らない）
    monkeypatch.setattr(vu, "_LAST_AGENT_CHANGES_FILE", str(tmp_path / "none.tsv"))
    assert va.show_changes(str(tmp_path / "none.tsv")) is False
    assert va.show_changes(path) is True
    out = capsys.readouterr().out
    assert "変わったセル: 30 個" in out and "A30" in out          # --changes は全部出す
    assert "書式の変化: 2 件" in out and "罫線・内・横" in out     # 書式も後から読める
    ns, unknown = vm.build_parser().parse_known_args(["agent", "--changes"])
    assert ns.changes and not unknown


# ----------------------------------------------------------------
# 不変条件の検査（2026-09-05）: 弾の check は「頼んだことが当たったか」しか見ない。
# 「頼んでいないところを壊していないか」を全弾に自動で当てる目のテスト。
# ----------------------------------------------------------------

def _snap(values, formulas=None, used="A1", hidden=(), cf=0, validation="", filt=""):
    return {"used": used, "values": values, "formulas": dict(formulas or {}),
            "hidden": tuple(hidden), "cf": cf, "validation": validation, "filter": filt}


def test_inv_cell_name_maps_grid_to_address():
    import vbam_agent as va
    assert va._cell_name(0, 0, "A1") == "A1" and va._cell_name(2, 3, "A1") == "D3"
    # 使用範囲が A1 起点とは限らない（右の表 I5:K10 の控えは I5 が原点）
    assert va._cell_name(0, 0, "I5:K10") == "I5" and va._cell_name(1, 2, "I5:K10") == "K6"
    # 2 文字の列（AA 以降）でも桁上がりする
    assert va._cell_name(0, 1, "Z1") == "AA1"


def test_inv_catches_changes_outside_the_request():
    import vbam_agent as va
    before = {"sheets": {"実射1": _snap([["会員番号", "氏名"], ["2001", "青木 誠"]]),
                         "他の表": _snap([["触るな", "1"]])},
              "names": {"名簿": "=実射1!$A$1:$B$2"}}

    # 何も変わっていなければ違反なし
    assert va._inv_compare(before, before, "実射1") == []

    # (1) 頼んでいないシートが変わった＝本物のブックに書いた事故（2026-09-04 に実際に起きた形）
    now = {"sheets": {"実射1": before["sheets"]["実射1"], "他の表": _snap([["書いた", "1"]])},
           "names": dict(before["names"])}
    bad = va._inv_compare(before, now, "実射1")
    assert len(bad) == 1 and "他の表" in bad[0] and "'触るな' → '書いた'" in bad[0]
    assert va._inv_compare(before, now, "実射1", skip=("other_sheet",)) == []

    # 練習台そのものが変わるのは正当（それが依頼）
    now = {"sheets": {"実射1": _snap([["会員番号", "氏名"], ["2001", "青木誠"]]),
                      "他の表": before["sheets"]["他の表"]}, "names": dict(before["names"])}
    assert va._inv_compare(before, now, "実射1") == []

    # (2) 名前定義が消えた（材料に出た名前を消さない＝手の型の完了条件）
    now = {"sheets": before["sheets"], "names": {}}
    bad = va._inv_compare(before, now, "実射1")
    assert len(bad) == 1 and "名前定義が消えた: 名簿" in bad[0]

    # シートごと消された
    now = {"sheets": {"実射1": before["sheets"]["実射1"]}, "names": dict(before["names"])}
    assert "頼んでいないシート「他の表」が消えた" in va._inv_compare(before, now, "実射1")


def test_inv_catches_hidden_rows_and_crushed_formulas():
    import vbam_agent as va
    vals = [["番号", "金額"], ["1", "100"], ["2", "200"]]
    forms = {"B3": "=SUM(B2:B2)"}
    before = {"sheets": {"実射1": _snap(vals, forms, hidden=(2,))}, "names": {}}

    # (3) 絞り込み中の行を解いたまま返した（本物のブックで出た欠陥 1 の形。道具は解いて書き、掛け直す）
    now = {"sheets": {"実射1": _snap(vals, forms, hidden=())}, "names": {}}
    bad = va._inv_compare(before, now, "実射1")
    assert len(bad) == 1 and "隠れていた行が変わった" in bad[0] and "[2] → なし 行目" in bad[0]
    assert va._inv_compare(before, now, "実射1", skip=("hidden",)) == []

    # 絞り込み（オートフィルタ）で隠れていた場合は、道具が「解いてから書く」のが仕様なので咎めない。
    # ここを咎めると、正しく直した仕事まで落ちる（2026-09-05・実射で踏んだ。「郵便番号の統一 ×
    # 絞り込み中」が、道具が正しく全行を直したのに不合格になった）
    b_f = {"sheets": {"実射1": _snap(vals, forms, hidden=(2,), filt="A1:B3")}, "names": {}}
    n_f = {"sheets": {"実射1": _snap(vals, forms, hidden=(), filt="A1:B3")}, "names": {}}
    assert va._inv_compare(b_f, n_f, "実射1") == []
    # 手で隠した行（絞り込みではない）が開いたら、それは頼んでいない変化
    assert va._inv_compare(before, now, "実射1") != []

    # (4) 元からあった数式が値に潰された
    now = {"sheets": {"実射1": _snap(vals, {}, hidden=(2,))}, "names": {}}
    bad = va._inv_compare(before, now, "実射1")
    assert len(bad) == 1 and "数式が値に潰された" in bad[0] and "B3: =SUM(B2:B2)" in bad[0]
    # 数式が増えるのは正当（突き合わせの列を足すのが依頼）
    now = {"sheets": {"実射1": _snap(vals, {"B2": "=A2", "B3": "=SUM(B2:B2)"}, hidden=(2,))}, "names": {}}
    assert va._inv_compare(before, now, "実射1") == []


def test_uncovered_formula_cells_catches_unread_ranges():
    """見ていない範囲を「なし」と報告させない（2026-09-05・本物の 161 列の表で出た欠陥）。

    AI は A1:P13 など 5 ブロックだけ read して、残り約 90 列を一度も見ずに
    「不整合なし・取りこぼしなし」と断言した。道具が数えて差し戻す。
    """
    import vbam_agent as va
    assert va._a1_to_rc("A1") == (1, 1)
    assert va._a1_to_rc("$BR$13") == (13, 70)
    assert va._a1_to_rc("FE2") == (2, 161)
    assert va._a1_to_rc("") is None and va._a1_to_rc("A") is None
    assert va._range_box("A1:P13") == (1, 1, 13, 16)
    assert va._range_box("計算表!$A$1:$P$13") == (1, 1, 13, 16)
    assert va._range_box("C7") == (7, 3, 7, 3)
    assert va._range_box("P13:A1") == (1, 1, 13, 16)      # 逆順でも同じ箱
    assert va._range_box("なにこれ") is None

    # 実際に踏んだ形: A1:P13 だけ読んで、BO2・FE2 を見ていない
    formulas = ["C4", "H4", "P13", "BO2", "BR2", "FE2"]
    read = ["A1:P13"]
    assert va._uncovered_formula_cells(formulas, read) == ["BO2", "BR2", "FE2"]
    # 全部覆えば空
    assert va._uncovered_formula_cells(formulas, ["A1:FE13"]) == []
    # 範囲を足していけば減る
    assert va._uncovered_formula_cells(formulas, ["A1:P13", "BO1:BR13"]) == ["FE2"]
    # 読んだ範囲が無ければ全部が未読／数式が無ければ空
    assert va._uncovered_formula_cells(formulas, []) == formulas
    assert va._uncovered_formula_cells([], ["A1:P13"]) == []


def test_inv_allows_external_links_to_become_values():
    """外部参照を値にするのは人が普通に頼む仕事＝咎めない（2026-09-05・実射で踏んだ）。

    手順書「リンクと外部参照の解消」は、依頼どおり動くと必ず (4) に引っかかって落ちていた
    （58/59 の 1 本）。(3) の絞り込みと同じ型の誤検知。
    """
    import vbam_agent as va
    ext1 = r"='C:\Users\user\AppData\Local\Temp\[外部.xlsx]元'!A1"
    ext2 = "=[外部.xlsx]元!A2"
    ext3 = "=[1]Sheet1!A1"                      # リンクが名前で保存された形
    for f in (ext1, ext2, ext3):
        assert va._is_external_formula(f), f
    # 構造化参照・自ブックの他シート参照・ふつうの式は「外部」ではない
    for f in ("=SUM(表1[金額])", "=元!A1", "=SUM(B2:B3)", "=IFERROR(VLOOKUP(A2,名簿,2,FALSE),\"なし\")"):
        assert not va._is_external_formula(f), f

    vals = [["項目", "値"], ["外部1", "100"], ["外部2", "200"], ["合計", "300"]]
    before = {"sheets": {"実射1": _snap(vals, {"B2": ext1, "B3": ext2, "B4": "=SUM(B2:B3)"})}, "names": {}}
    # 外部参照 2 本だけが値になった＝依頼どおり。咎めない
    now = {"sheets": {"実射1": _snap(vals, {"B4": "=SUM(B2:B3)"})}, "names": {}}
    assert va._inv_compare(before, now, "実射1") == []
    # 同じ回に自前の数式（合計）まで潰したら、それは頼んでいない変化として残る
    now2 = {"sheets": {"実射1": _snap(vals, {})}, "names": {}}
    bad = va._inv_compare(before, now2, "実射1")
    assert len(bad) == 1 and "B4: =SUM(B2:B3)" in bad[0] and "外部.xlsx" not in bad[0]


def test_inv_does_not_fire_when_a_column_is_added():
    """誤検知の回し車を止める（2026-09-05・入れた直後に見つけた穴）。

    隠れ行を「可視範囲の文字」で持つと、F 列を 1 本足しただけで A1:E9 → A1:F9 になり、
    突き合わせの弾（F 列を足すのが依頼）が毎回落ちる。行番号の集合で持てば動かない。
    """
    import vbam_agent as va
    assert va._addr_rows("A1:B2,A4:B5") == {1, 2, 4, 5}
    assert va._addr_rows("$A$3") == {3} and va._addr_rows("") == set()
    # 使用範囲 A1:E9（3 行目が隠れ）に F 列を足して A1:F9 になっても、隠れ行は 3 のまま
    used_before, used_after = "A1:E9", "A1:F9"
    vis_before, vis_after = "A1:E2,A4:E9", "A1:F2,A4:F9"
    assert (va._addr_rows(used_before) - va._addr_rows(vis_before)
            == va._addr_rows(used_after) - va._addr_rows(vis_after) == {3})
    before = {"sheets": {"実射1": _snap([["a"], ["b"]], {}, used=used_before, hidden=(3,))}, "names": {}}
    now = {"sheets": {"実射1": _snap([["a", "x"], ["b", "y"]], {}, used=used_after, hidden=(3,))}, "names": {}}
    assert va._inv_compare(before, now, "実射1") == []


def test_inv_catches_lost_format_and_validation():
    import vbam_agent as va
    vals = [["番号"], ["1"]]
    before = {"sheets": {"実射1": _snap(vals, cf=2, validation="A2:A9")}, "names": {}}
    # 条件付き書式・入力規則は足すのは正当、消すのは頼んでいない
    now = {"sheets": {"実射1": _snap(vals, cf=1, validation="A2:A9")}, "names": {}}
    assert "条件付き書式が減った（2 → 1）" in va._inv_compare(before, now, "実射1")
    now = {"sheets": {"実射1": _snap(vals, cf=3, validation="")}, "names": {}}
    bad = va._inv_compare(before, now, "実射1")
    assert len(bad) == 1 and "入力規則が消えた（A2:A9）" in bad[0]


def test_inv_reports_row_count_and_limits_noise():
    import vbam_agent as va
    b = [["x"] for _ in range(9)]
    a = [["x"] for _ in range(5)]
    d = va._grid_diff(b, a, "A1", "値")
    assert d and "値の行数が 9 → 5" in d[0]
    # 明細で埋めない（頭 3 つまで）
    b = [[str(i)] for i in range(20)]
    a = [["ちがう"] for _ in range(20)]
    assert len(va._grid_diff(b, a, "A1", "値")) == 3


def test_inv_is_wired_into_every_fire_case():
    """弾を 1 本も書き足さずに、全弾へ自動で当たっていること（そこが①の値打ち）。"""
    import inspect as _inspect
    import vbam_agent as va
    import vbam_recipes as vr
    for fn in (vf._fire_sheet_cases, vr._fire_recipe_cases):
        src = _inspect.getsource(fn)
        assert "_snapshot_book(wb" in src, f"{fn.__name__} が撃つ前の控えを取っていない"
        assert "_inv_violations(" in src and "inv_skip" in src, f"{fn.__name__} が照らしていない"
        assert "ok = False" in src, f"{fn.__name__} が違反を不合格にしていない"


# ----------------------------------------------------------------
# 練習台の状態（2026-09-05）: 弾はそのままに、撃つ前のシートの姿を本物のブック寄りにする。
# 59 本が同じきれいな表から始まる限り、本数を増やしても踏む地雷は増えない、への手当て。
# ----------------------------------------------------------------

def test_fire_states_are_defined_and_reproducible():
    import vbam_agent as va
    # 本物のブックにある姿が一通りそろっている（きれいな表だけで撃たない）
    for want in ("なし", "絞り込み中", "隠れ行", "前のフィルタ", "枠固定", "名前定義",
                 "条件付き書式", "他シート参照", "結合タイトル", "書式ばらばら"):
        assert want in vf._FIRE_STATES, f"状態「{want}」が無い"
    cases = [{"name": f"弾{i}"} for i in range(12)]

    # seed が同じなら同じ姿＝落ちた組み合わせをそのまま撃ち直せる（--seed）
    a = vf._pick_states(cases, None, 12345)
    assert a == vf._pick_states(cases, None, 12345)
    assert a != vf._pick_states(cases, None, 999), "seed を変えても同じ姿では掛け算にならない"
    # 弾ごとにばらけている（全部同じ状態を引いていたら掛け算の意味がない）
    assert len(set(a.values())) >= 3
    # 「なし」は自動では引かない（きれいな表に戻らない）
    assert "なし" not in set(a.values())

    # --state で全弾に同じ状態を固定できる
    b = vf._pick_states(cases, "隠れ行", 1)
    assert set(b.values()) == {"隠れ行"} and len(b) == len(cases)


def test_fire_states_are_wired_and_validated():
    import inspect as _inspect
    import vbam_agent as va
    import vbam_recipes as vr
    for fn in (vf._fire_sheet_cases, vr._fire_recipe_cases):
        src = _inspect.getsource(fn)
        assert "_apply_state(ws" in src, f"{fn.__name__} が状態を当てていない"
        assert "states" in _inspect.signature(fn).parameters, f"{fn.__name__} が状態を受け取らない"
    src = _inspect.getsource(vf.fire_agent) + _inspect.getsource(vf._fire_agent_body)
    assert "_pick_states(" in src and "--seed" in src, "fire_agent が seed を出していない"
    # 状態は控えを取る前に当てる（「元からそうだった」姿にする＝後から当てたら不変条件が誤検知する）
    body = _inspect.getsource(vf._fire_sheet_cases)
    assert body.index("_apply_state(ws") < body.index("_snapshot_book(wb")
    # 知らない状態名は撃つ前に弾く
    ns = vm.build_parser().parse_args(["agent", "--fire", "--state", "隠れ行", "--seed", "7"])
    assert ns.state == "隠れ行" and ns.seed == "7"


def test_fire_state_functions_touch_the_right_things():
    """状態の中身（COM を呼ぶ前に、何を触るつもりかを型で押さえる）。"""
    import inspect as _inspect
    import vbam_agent as va
    src = {n: _inspect.getsource(f) for n, f in vf._FIRE_STATES.items()}
    assert "Hidden = True" in src["隠れ行"]
    # 「絞り込み中」は AutoFilter でなければならない。手で隠した行（Rows.Hidden）では書き込みは
    # 壊れず、本物のブックの欠陥 1 を再現できない（2026-09-05 実測で判明。最初これを取り違えた）
    assert "AutoFilter(" in src["絞り込み中"] and "Hidden = True" not in src["絞り込み中"]
    assert "AutoFilter" in src["前のフィルタ"] and "AutoFilterMode = False" in src["前のフィルタ"]
    assert "FreezePanes = True" in src["枠固定"]
    assert "Names.Add" in src["名前定義"]
    assert "FormatConditions.Add" in src["条件付き書式"]
    assert "Worksheets.Add" in src["他シート参照"] and "COUNTA" in src["他シート参照"]
    assert "Merge()" in src["結合タイトル"]
    assert "Interior.Color" in src["書式ばらばら"]
    # 状態は「答えを変えない」のが約束。行の挿入は番地で書かれた弾（名簿は A3:E9）を全部ずらす
    # ＝入れた直後に実機で踏んだ（2026-09-05）。どの状態も行・列を挿入しない
    for n, s in src.items():
        assert ".Insert(" not in s and "Rows(1).Insert" not in s, f"状態「{n}」が行を挿入している"
        # 消す手を持たせない。Worksheet.Delete は削除の確認ダイアログを出し、実射が丸ごと固まる
        # （2026-09-05・vague 9 本の 2 本目で 12 分止まった。Excel は「応答あり」のままだった）
        assert ".Delete()" not in s, f"状態「{n}」が消す手を持っている（確認ダイアログで実射が固まる）"
    # 実射は 1 冊のブックに 実射1・実射2 … と並ぶ＝状態が置く名前は弾ごとに一意でなければ 2 本目で衝突する
    assert "{ws.Name}" in src["名前定義"] and "{ws.Name}" in src["他シート参照"]
    # 当てられない練習台（小さすぎる等）でも実射を止めない
    class Boom:
        def __getattr__(self, k):
            raise RuntimeError("この練習台には当てられない")
    assert vf._apply_state(Boom(), "隠れ行") == ""
    assert vf._apply_state(Boom(), "なし") == "" and vf._apply_state(Boom(), "無い状態") == ""


# ----------------------------------------------------------------
# 冪等性の検査（2026-09-05）: 正解の表を持たずに判定できる関係。
# 「合計行がもう 1 本足された」は 1 回目の答え合わせでは PASS のまま通る。
# ----------------------------------------------------------------

def test_idempotence_catches_double_application():
    import vbam_agent as va
    done = {"sheets": {"実射1": _snap([["氏名", "金額"], ["青木", "100"], ["合計", "100"]],
                                      {"B3": "=SUM(B2:B2)"})}, "names": {}}
    # 2 回目で何も起きないのが正しい
    assert vf._idempotence_violations(done, done, "実射1") == []

    # 合計行がもう 1 本足された（1 回目の check は「合計が合っている」で PASS のまま通る形）
    twice = {"sheets": {"実射1": _snap([["氏名", "金額"], ["青木", "100"], ["合計", "100"], ["合計", "100"]],
                                       {"B3": "=SUM(B2:B2)", "B4": "=SUM(B2:B3)"})}, "names": {}}
    bad = vf._idempotence_violations(done, twice, "実射1")
    assert any("値が変わった" in b and "行数が 3 → 4" in b for b in bad)
    assert any("数式が変わった" in b and "増えた 1 本" in b and "B4" in b for b in bad)

    # ピボットをもう 1 枚作った
    more = {"sheets": {"実射1": done["sheets"]["実射1"], "Sheet4": _snap([["集計"]])}, "names": {}}
    assert vf._idempotence_violations(done, more, "実射1") == ["シートが増えた: Sheet4"]


def test_idempotence_is_wired_behind_a_flag():
    import inspect as _inspect
    import vbam_agent as va
    src = _inspect.getsource(vf._fire_sheet_cases)
    assert "twice" in _inspect.signature(vf._fire_sheet_cases).parameters
    assert "_idempotence_violations(" in src and "if twice and ok:" in src, "2 回目が配線されていない"
    # 合格した弾にだけ撃つ（落ちた弾にもう 1 回課金しない）
    assert src.index("if twice and ok:") > src.index("ok = ok and (r['done']")
    ns = vm.build_parser().parse_args(["agent", "--fire", "--twice"])
    assert ns.twice
    # 既定では撃たない（往復が 2 倍になるので、明示したときだけ）
    assert vm.build_parser().parse_args(["agent", "--fire"]).twice is False


# ----------------------------------------------------------------
# 他人のブックの構造を写し取る（2026-09-05）: 自分の想像で組んだ練習台の外から的を供給する口。
# ----------------------------------------------------------------

def test_harvest_dummy_keeps_the_shape_not_the_content():
    import datetime as _dt
    import random as _rnd
    import vbam_agent as va
    memo, rnd = {}, _rnd.Random(1)

    # 同じ元値は同じダミーへ（左の会員番号と右の会員番号の関係が壊れない＝突き合わせが成り立つ）
    a1 = vf._dummy_value("青木 誠", memo, rnd)
    a2 = vf._dummy_value("青木 誠", memo, rnd)
    b1 = vf._dummy_value("石川 恵", memo, rnd)
    assert a1 == a2 and a1 != b1
    # 一目で偽物と分かる形（実在しそうな氏名・住所を作らない）
    assert a1.startswith("ダミー") and b1.startswith("ダミー")

    # 数値は桁数を保つ（列幅の ### も桁区切りの検査も、桁が変わると意味を失う）
    for src in (5, 1234, 1875000):
        got = vf._dummy_value(src, memo, rnd)
        assert len(str(abs(int(got)))) == len(str(abs(src))), f"{src} → {got} で桁が変わった"
    # 同じ数値は同じダミー
    assert vf._dummy_value(1234, memo, rnd) == vf._dummy_value(1234, memo, rnd)

    # 日付は同じ年月のまま（月別集計・前月比の弾が成り立つ）
    d = vf._dummy_value(_dt.datetime(2026, 4, 17, tzinfo=_dt.timezone.utc), memo, rnd)
    assert d.year == 2026 and d.month == 4

    # 数式・空・真偽は触らない（None＝そのまま）
    assert vf._dummy_value("=SUM(B2:B4)", memo, rnd) is None
    assert vf._dummy_value(None, memo, rnd) is None and vf._dummy_value(True, memo, rnd) is None
    assert vf._dummy_value("   ", memo, rnd) is None

    # 引いたダミーが元と同じ値になったら引き直す。1 桁の数値は 1/10、同じ月の日付は 1/27 でぶつかり、
    # そのぶんだけ中身が元のまま残る（2026-09-05・読み戻しの関所が実物で 22 個捕まえた）
    for s in range(40):
        for src in (0, 1, 5, 9, 42):
            got = vf._dummy_value(src, {}, _rnd.Random(s))
            assert got != src, f"seed {s}: 数値 {src} がそのまま残った"
        d0 = _dt.datetime(2026, 4, 17, tzinfo=_dt.timezone.utc)
        got = vf._dummy_value(d0, {}, _rnd.Random(s))
        assert got.day != d0.day and got.month == 4, f"seed {s}: 日付 {d0} がそのまま残った"


def test_harvest_counts_what_is_left_after_writing():
    """書きっぱなしにせず、読み戻して「元の中身が残ったセル」を数えること。

    2026-09-05: 範囲でまとめて書いたら、隠れた行のある 2023 シートで 70 セルが元の氏名のまま残った
    （9/4 に本物のブックで見つけた欠陥 1 と同じ形を、写し取りの側でもう一度踏んだ）。
    人に渡せるかが懸かっているので、作った後に数える。
    """
    import inspect as _inspect
    import vbam_agent as va
    src = _inspect.getsource(vf.harvest_book)
    # 1 行ずつ書く（隠れ行を飛ばさない）。範囲へのまとめ書きに戻したらここで落ちる
    assert "for i, line in enumerate(new):" in src and "ur.Value = tuple(" not in src
    # 書いた後に読み戻して数え、残っていたら「作りました」で終わらせない
    assert "back = va._rows_of(sh.UsedRange.Value)" in src
    assert "中身が元のまま残ったセル" in src and "人に渡せません" in src


def test_fire_bed_judges_only_what_the_run_changed():
    """写し取った本物の構造に撃つとき、元からの歪みを咎めないこと。

    本物のブックは元から歪んでいる（見出しの無い列・継ぎはぎの罫線・旧範囲の _FilterDatabase）。
    2026-09-05 の初回実射は「見出しが空の列（O,P）」で 3 本とも落ちた＝元からそうだった。
    直せと頼んでいないものを不合格にすると、どんな本物のブックでも 0 点になり、テストにならない。
    """
    import inspect as _inspect
    import vbam_agent as va
    # 判定の中身は _fire_bed_case（2026-09-06・本番から拾った弾 --fire mine も同じ判定を使うので部品化した）
    src = _inspect.getsource(vf._fire_bed_case)
    # 撃つ前の指摘を控え、増えた分だけを見る
    assert "before_audit = set(" in src and "if str(x) not in before_audit" in src
    assert src.index("before_audit = set(") < src.index("r = va.run_agent("), "控えは撃つ前に取る"
    # 正解の表を持たない＝どんな構造にも撃てる（判定は 不変条件・仕上げ検査の差分・冪等性の 3 つだけ）
    assert "_inv_violations(" in src and "_idempotence_violations(" in src
    assert "expect" not in src and "check" not in src.replace("仕上げ検査", "").replace("audit_table", "")
    # audit_table は ws.UsedRange でなく A1 の CurrentRegion で見る（2026-09-05・本物のブックの実射で発見。
    # 使用範囲が実データより大きく膨張した台帳で「G1001 まで罫線が無い」と誤検知していた）
    assert "audit_table(ws, ws.UsedRange)" not in src
    assert 'ws.Range("A1").CurrentRegion' in src
    # 練習台は使い捨て（保存しない）
    assert "wb.Close(SaveChanges=False)" in _inspect.getsource(vf._fire_bed)
    assert "_fire_bed_case(" in _inspect.getsource(vf._fire_bed), "写し取りは共通の判定を通す"
    for c in vf._BED_REQUESTS:
        assert set(c) == {"name", "request"}, "写し取りの依頼に正解表を持たせない"
    ns = vm.build_parser().parse_args(["agent", "--fire", "--bed", r"C:\x\台.xlsm", "2024"])
    assert ns.bed.endswith("台.xlsm") and ns.posargs == ["2024"]


def test_harvest_never_touches_the_original(tmp_path, capsys):
    import inspect as _inspect
    import vbam_agent as va
    src = _inspect.getsource(vf.harvest_book)
    # 元は複製してから複製の側だけを開く（元のブックを開いて書き換えない）
    assert "shutil.copy2(path, out)" in src and "get_workbook(out)" in src
    assert "Workbooks.Open(path" not in src
    # 出力先が元と同じなら断る
    f = tmp_path / "本物.xlsx"
    f.write_bytes(b"dummy")
    assert vf.harvest_book(str(f), str(f)) is None
    assert "元のブックと同じ" in capsys.readouterr().out
    # 無いファイルは静かに断る
    assert vf.harvest_book(str(tmp_path / "無い.xlsx")) is None

    # 配線
    ns = vm.build_parser().parse_args(["agent", "--harvest", r"C:\x\本物.xlsx", "--keep-rows", "2"])
    assert ns.harvest.endswith("本物.xlsx") and ns.keep_rows == "2"
    body = _inspect.getsource(va._cmd_agent_body)
    assert "harvest_book(args.harvest" in body
    assert body.index("harvest") < body.index("if getattr(args, 'fire'"), "--harvest が --fire より後ろだと届かない"


def test_harvest_falls_back_to_saveas_when_save_fails():
    """9-2: 修復モードで開いたブック（_open_maybe_repair）は素の Save も同じ理由で弾かれることがある
    （2026-09-05・本物のブックの実射で発見。Google スプレッドシート書き出しの xlsx）。
    パスと形式を明示した SaveAs に切り替える。
    """
    import inspect as _inspect
    import vbam_agent as va
    src = _inspect.getsource(vf.harvest_book)
    assert "wb.Save()" in src
    assert "wb.SaveAs(out, FileFormat=out_fmt)" in src
    assert src.index("wb.Save()") < src.index("wb.SaveAs(out, FileFormat=out_fmt)"), "素の Save をまず試す"


def test_harvest_output_path_is_recovered_when_swallowed_by_target_file():
    """9-0: 'agent --harvest 元.xlsm 出力.xlsm' で出力名が既定のテンプレへ落ちていた（2026-09-05）。

    .xlsm に見える先頭の位置引数は parse_target_and_rest が target_file 側に取る＝rest は空になり、
    harvest_book への出力パスが None のまま渡っていた（harvest には「触る対象ブック」という概念が無い）。
    """
    import inspect as _inspect
    import vbam_agent as va
    target_file, rest = va.parse_target_and_rest(["出力.xlsm"])
    assert target_file == "出力.xlsm" and rest == []          # 出力名が rest でなく target_file 側に落ちる
    assert (rest[0] if rest else target_file) == "出力.xlsm"   # 直した式はここから拾い直す
    body = _inspect.getsource(va._cmd_agent_body)
    assert "out = rest[0] if rest else target_file" in body


# ----------------------------------------------------------------
# 2026-09-06 午後: 90 点 → 100 点の回収 4 件
#   ①画像に行番号・列番号 ②--undo 番号（5 個のどれにでも戻せる） ③--drop-case ④台帳の「人が保存した」
# ----------------------------------------------------------------

def _fake_ws_for_headers(r0=3, c0=2, nr=4, nc=5, col_w=20.0, row_h=15.0):
    """_draw_headers に渡す偽シート: Range(addr) が行・列の数と幅・高さ（ポイント）を答える。"""
    class Line:
        def __init__(self, size):
            self.Width = size
            self.Height = size

    class Axis:
        def __init__(self, n, size):
            self.Count = n
            self._s = size

        def __call__(self, i):
            return Line(self._s)

    class Rng:
        Row, Column = r0, c0
        Rows, Columns = Axis(nr, row_h), Axis(nc, col_w)

    class WS:
        def Range(self, addr):
            return Rng()
    return WS()


def test_view_image_gets_row_and_column_headers(tmp_path, monkeypatch):
    """①: AI に見せる画像に行番号・列番号のヘッダーを描き足す（Claude for Excel が注文していた形）。

    それまでは添え文で「左上のセルが B3 です。ここから数えてください」と AI に数えさせていた。
    """
    from PIL import Image
    import vbam_agent as va
    png = tmp_path / "view.png"
    Image.new('RGB', (100, 60), (255, 255, 255)).save(png)       # 5 列 × 20pt・4 行 × 15pt を 1pt=1px で
    assert va._draw_headers(str(png), _fake_ws_for_headers(), "B3:F6") is True
    im = Image.open(png)
    assert im.size[0] > 100 and im.size[1] > 60                  # 上と左に帯が足された
    assert im.getpixel((im.size[0] - 1, im.size[1] - 1)) == (255, 255, 255)   # 元の絵は右下にそのまま
    assert im.getpixel((2, 2)) == va._VIEW_HEADER_BG              # 左上の角は帯の色
    # 描けない（PIL が読めない・幅が 0）ときは False＝画像はそのまま・添え文は「左上から数えて」に戻る
    assert va._draw_headers(str(tmp_path / "no.png"), _fake_ws_for_headers(), "B3:F6") is False
    assert va._draw_headers(str(png), _fake_ws_for_headers(col_w=0.0), "B3:F6") is False
    monkeypatch.setattr(vva, "_VIEW_HEADED", True)
    note = va._view_note("B3:F6", "末尾")
    assert "行番号のヘッダーが写っています" in note and "末尾）" in note and "ここから数えて" not in note
    monkeypatch.setattr(vva, "_VIEW_HEADED", False)
    note = va._view_note("B3:F6")
    assert "左上のセルが B3 です" in note and "ここから数えて" in note
    # 配線: 撮った直後に描く・人が見る最後の画像にも描く・添え文は全部 _view_note を通る
    import inspect as _inspect
    src = _inspect.getsource(va._shot_for_ai)
    assert "_VIEW_HEADED = _draw_headers(_LAST_AGENT_VIEW_PNG" in src
    assert "_draw_headers(_LAST_AGENT_PNG" in _inspect.getsource(va._verify_sheet)
    loop = _inspect.getsource(va.run_agent)
    assert loop.count("_view_note(") == 4 and "ここから数えてください" not in loop


def _plant_backup(va, tmp_path, name, when, run_id, book="台帳.xlsx", sheet="名簿", request="依頼"):
    """控えとその隣の覚書を 1 組置く（mtime を when にそろえる）。"""
    p = tmp_path / f"{os.path.splitext(book)[0]}{va._AGENT_BACKUP_MARK}{name}.xlsx"
    p.write_text("x", encoding="utf-8")
    meta = {'book': book, 'sheet': sheet, 'sheets': [sheet], 'path': str(p), 'request': request,
            'time': _time.strftime('%Y-%m-%d %H:%M:%S', _time.localtime(when)), 'run_id': run_id, 'shapes': []}
    assert va._save_undo_meta(meta) is True
    os.utime(p, (when, when))
    return p, meta


def test_backups_keep_their_own_memo_and_undo_can_pick_one(tmp_path, monkeypatch, capsys):
    """②: 控えはブックごとに 5 個残るのに、戻せるのは直前の 1 つだけだった。

    控えごとに覚書を隣に置き、--backups の番号（か走行の名札）で --undo 番号。その控えより後に
    同じブックで走行があれば、後からやったぶんが消えるので既定で止まる（--force）。
    """
    import vbam_agent as va
    monkeypatch.setattr(vu, "BACKUP_DIR", str(tmp_path))
    monkeypatch.setattr(vu, "_LAST_AGENT_UNDO_FILE", str(tmp_path / "_last_agent_undo.json"))
    monkeypatch.setattr(vu, "get_workbook",
                        lambda *a, **k: (_ for _ in ()).throw(AssertionError("Excel に触った")))
    base = _time.mktime(_time.strptime('2026-09-06 10:00:00', '%Y-%m-%d %H:%M:%S'))   # 10:00／11:00／12:00 の控え
    p1, m1 = _plant_backup(va, tmp_path, "20260906_100000", base, "20260906_100000", request="古い仕事")
    p2, m2 = _plant_backup(va, tmp_path, "20260906_110000", base + 3600, "20260906_110000", request="中の仕事")
    p3, m3 = _plant_backup(va, tmp_path, "20260906_120000", base + 7200, "20260906_120000", request="直前の仕事")
    # 隣の覚書ができている（直前用の覚書も最後の 1 つ）
    assert os.path.isfile(va._backup_meta_path(p1)) and os.path.isfile(va._backup_meta_path(p3))
    assert va._load_undo_meta()['run_id'] == "20260906_120000"
    items = va._list_backups()
    assert [d['n'] for d in items] == [1, 2, 3] and items[0]['path'] == str(p3)   # 新しい順に番号
    assert va._pick_backup("2")['meta']['request'] == "中の仕事"
    assert va._pick_backup("20260906_100000")['meta']['request'] == "古い仕事"     # 走行の名札でも
    assert va._pick_backup(p1.name)['meta']['request'] == "古い仕事"               # ファイル名でも
    assert va._pick_backup("9") is None and va._pick_backup("") is None
    # 一覧: 番号・直前の印・依頼
    va.backups_list()
    out = capsys.readouterr().out
    assert "[1]" in out and "[3]" in out and "← 直前" in out and "古い仕事" in out
    # 覚書の無い控え（この作りより前のもの）は番号を持たず、一覧では数だけ畳んで出る。名前で指せば「戻せない」と言う
    old = tmp_path / f"台帳{va._AGENT_BACKUP_MARK}20260901_000000.xlsx"
    old.write_text("x", encoding="utf-8")
    os.utime(old, (base - 9999, base - 9999))
    tail = va._list_backups()[-1]
    assert tail['meta'] is None and tail['n'] is None
    assert [d['n'] for d in va._list_backups()] == [1, 2, 3, None]
    va.backups_list()
    out = capsys.readouterr().out
    assert "ほか 1 個は覚書なし" in out and old.name in out and "[4]" not in out
    assert va._pick_backup("4") is None
    assert va.undo_agent(which=old.name) is False
    assert "覚書がありません" in capsys.readouterr().out
    # 無い番号は止まる（Excel に触らない）
    assert va.undo_agent(which="99") is False
    # その控えより後に同じブックで走行がある → 既定で止まる（--force で越える。越えた先で Excel に触る）
    va.runs_append({'time': '2026-09-06 11:30:00', 'run_id': 'x1', 'mode': 'sheet', 'book': "台帳.xlsx",
                    'request': "後の仕事"})
    assert va.undo_agent(which="2") is False
    out = capsys.readouterr().out
    assert "後に、同じブックで走行" in out and "後の仕事" in out and "--force" in out
    assert va._runs_after(m3) == []                             # 直前の控えより後には何も無い
    with _pytest.raises(AssertionError, match="Excel に触った"):
        va.undo_agent(which="2", force=True)
    # 間引き: 隣の覚書は数に入れず、控えと一緒に消える
    va._prune_agent_backups(keep=2, stem="台帳")
    left = sorted(q.name for q in tmp_path.iterdir())
    assert p1.name not in left and va._backup_meta_path(p1.name) not in left
    assert p2.name in left and p3.name in left and os.path.basename(va._backup_meta_path(p3)) in left
    # 覚書を書き足す道（別シート・作った物・stale）も隣の覚書へ通る
    va._record_created([{'kind': 'table', 'name': 'T1'}])
    assert va._load_undo_meta(va._backup_meta_path(p3))['created'] == [{'kind': 'table', 'name': 'T1'}]
    va._mark_undo_stale("build")
    assert va._load_undo_meta(va._backup_meta_path(p3))['stale'] == "build"
    # 引数: --undo だけ＝直前、--undo 2＝番号、--backups
    p = vm.build_parser()
    ns, unknown = p.parse_known_args(["agent", "--undo", "--dry-run"])
    assert ns.undo == "latest" and ns.dry_run and not unknown
    ns, _ = p.parse_known_args(["agent", "--undo", "2", "--force"])
    assert ns.undo == "2" and ns.force
    ns, _ = p.parse_known_args(["agent", "--backups"])
    assert ns.backups and ns.undo is None


def test_drop_case_removes_the_bullet_and_its_bed(tmp_path, monkeypatch, capsys):
    """③: 弾を消す口。台帳の項目と、弾の置き場にある練習台だけを消す（人のブックを指していたら残す）。"""
    import vbam_agent as va
    cases_dir = tmp_path / "cases"
    cases_dir.mkdir()
    monkeypatch.setattr(vf, "_AGENT_CASES_DIR", str(cases_dir))
    bed = cases_dir / "観光.xlsx"
    bed.write_text("x", encoding="utf-8")
    outside = tmp_path / "人のブック.xlsx"
    outside.write_text("x", encoding="utf-8")
    vf._cases_save({"観光": {'name': "観光", 'bed': str(bed), 'sheet': "S", 'request': "r", 'time': "t"},
                    "外": {'name': "外", 'bed': str(outside), 'sheet': "S", 'request': "r", 'time': "t"}})
    assert vf.drop_case("無い弾") is False
    assert "その弾はありません" in capsys.readouterr().out
    assert vf.drop_case("観光") is True
    assert not bed.exists() and "観光" not in vf._cases_load()
    assert vf.drop_case("外") is True
    assert outside.exists() and vf._cases_load() == {}          # 置き場の外のファイルは残す
    assert "置き場の外" in capsys.readouterr().out
    p = vm.build_parser()
    ns, _ = p.parse_known_args(["agent", "--drop-case", "観光"])
    assert ns.drop_case == "観光"


def test_runs_ledger_sees_whether_the_book_was_saved_afterwards(tmp_path, monkeypatch):
    """④: undo 率は失敗の側の人の判断。受け入れの側＝「人が保存した」も台帳に出す。

    道具は保存しないので、ブックの更新日時が走行より新しければ人が保存した。macro は道具が保存するので数えない。
    """
    import vbam_agent as va
    book = tmp_path / "台帳.xlsx"
    book.write_text("x", encoding="utf-8")
    t_run = _time.time() - 600
    stamp = _time.strftime('%Y-%m-%d %H:%M:%S', _time.localtime(t_run))
    va.runs_append({'time': stamp, 'run_id': 'a', 'mode': 'sheet', 'book': "台帳.xlsx", 'sheet': "S",
                    'request': "整えて", 'done': True, 'gates': {}, 'path': str(book)})
    va.runs_append({'time': stamp, 'run_id': 'b', 'mode': 'sheet', 'book': "台帳.xlsx", 'sheet': "S",
                    'request': "戻された", 'done': True, 'gates': {}, 'path': str(book)})
    va.runs_append({'undo': 'b', 'time': stamp})
    va.runs_append({'time': stamp, 'run_id': 'c', 'mode': 'macro', 'book': "台帳.xlsx", 'sheet': None,
                    'request': "直して", 'done': True, 'gates': {}, 'path': None})
    os.utime(book, (t_run - 100, t_run - 100))                  # まだ保存していない
    rows = va._runs_load()
    assert all('saved' not in r for r in rows)
    assert "人が保存した（受け入れ）0/2" in va._runs_summary(rows)  # path のある 2 行が分母（macro は入らない）
    os.utime(book, (t_run + 100, t_run + 100))                  # 走行の後に保存した
    rows = va._runs_load()
    assert rows[0].get('saved') and rows[1].get('saved') and 'saved' not in rows[2]
    assert "人が保存した（受け入れ）1/2" in va._runs_summary(rows)  # 戻した走行は受け入れに数えない
    assert va.runs_history() is True
    # 配線: sheet と build は path を渡す・macro は渡さない
    import inspect as _inspect
    assert "path=_book_path(wb)" in _inspect.getsource(va.run_agent)
    assert "path=_book_path(wb)" in _inspect.getsource(va.run_build)
    assert "path=_book_path(wb)" not in _inspect.getsource(vmac.run_macro_agent)
    assert va._book_path(type("WB", (), {"Path": "", "FullName": "x"})()) is None
    assert va._book_path(type("WB", (), {"Path": "C:/d", "FullName": "C:/d/x.xlsx"})()) == "C:/d/x.xlsx"


# ================================================================
# 2026-09-06: Claude for Excel への 24 問の聞き取りから輸入した分
#   ①事前の門（非空セルへの書き込み） ②plan の「要判断」 ③不変条件の追加
#   ④報告の 4 段と道具が足す分 ⑤関所の差し戻しを往復に数えない ⑥派生値の数式率
#   ⑦同じ依頼の 2 回目 ⑩見出し行の推定・和暦
# ================================================================

class _GateWS:
    """書き先の中身だけ答える作り物のシート（Range(a).Formula）。"""

    def __init__(self, filled=None, grid=None):
        self._filled = filled or {}          # {'A1': '既にある値'}
        self._grid = grid                    # write_grid 用の 2 次元

    def Range(self, a):
        if self._grid is not None and ':' in str(a):
            outer = self
            return type("R", (), {"Formula": outer._grid})()
        return type("R", (), {"Formula": self._filled.get(str(a).replace('$', ''), '')})()


class _GateWB:
    def __init__(self, ws):
        self._ws = ws
        self.Sheets = lambda _n: ws


def test_writing_onto_someone_elses_value_is_refused_before_it_runs():
    """①: write_cells / write_grid は、書き先に人の値があると実行前に手ごと止める（2026-09-06）。

    不変条件は「書いた後」に咎める事後検査で、戻すのは人の手番になる。事前の門を一枚足す。
    同じ値の書き直しは通す（仕上げ検査の差し戻しで同じ表を書き直す流れを殺さないため）。
    """
    import vbam_agent as va
    ws = _GateWS({'C7': '前からある', 'C11': ''})
    wb = _GateWB(ws)
    act = {"op": "write_cells", "cells": {"C7": "新しい", "C11": "空だから通る"}}
    assert va._overwrite_hits(act, "名簿", wb) == ['C7']
    # 承認があれば通す
    assert va._overwrite_hits(dict(act, overwrite=True), "名簿", wb) == []
    # 同じ値の書き直しは破壊ではない
    assert va._overwrite_hits({"op": "write_cells", "cells": {"C7": "前からある"}}, "名簿", wb) == []
    assert va._cell_taken('0001', "'0001") is False and va._cell_taken('', 'x') is False
    # 文面は次の往復でそのまま直せる形
    msg = va._overwrite_error('write_cells', ['C7'])
    assert '"overwrite": true' in msg and '承認の言葉' in msg and 'C7' in msg
    # 手ごと止まり、後ろの手も実行されない
    res = va._execute([act, {"op": "tidy", "ranges": ["A1:C9"]}], "名簿", wb)
    assert res[0][1] is False and 'もとからある値が 1 セル' in res[0][2]
    assert res[1][1] is False and '1 手も実行していません' in res[1][2]     # 実行前の全手検査（2026-09-06 夜）


def test_the_gate_counts_only_the_cells_write_grid_will_actually_write():
    """①: write_grid は左上だけの番地でも rows の大きさへ広げて見る。行の短い分は書かない＝数えない。"""
    import vbam_agent as va
    assert va._grid_box('F5', 3, 2) == 'F5:G7' and va._grid_box('A1:C5', 1, 1) == 'A1:C5'
    grid = (('見出し', '金額', '備考'), ('人の値', '', ''))      # F5:H6 の中身
    wb = _GateWB(_GateWS(grid=grid))
    act = {"op": "write_grid", "range": "F5", "rows": [["新見出し", "新金額"], ["a", "=B6*2"]]}
    assert va._overwrite_hits(act, "名簿", wb) == ['F5', 'G5', 'F6']   # H 列は書かない＝数えない
    assert va._overwrite_hits(dict(act, overwrite=True), "名簿", wb) == []


def test_plan_can_hold_a_question_for_the_human_without_blocking_done():
    """②: 「要判断」＝人に聞かないと決められない項目。安全側で仮処理して done は塞がない（2026-09-06）。

    向こう（Claude for Excel）は ask_user_question で**止まる**が、こちらの往復は無人。
    止めると必ず往復を使い切るので、「仮に処理して report に質問を残す」に読み替えた。
    """
    import vbam_agent as va
    _s, _a, done, _r, plan = va._parse_reply(
        '{"plan":[{"item":"色分け","state":"済"},'
        '{"item":"重複行を消す","state":"要判断","ask":"0002 と 0005 が二重です。消してよいですか"}],'
        '"actions":[],"done":true,"report":"色分けしました"}')
    assert done is True and va._plan_pending(plan) == []          # 要判断は done を塞がない
    assert va._plan_asks(plan) == ['重複行を消す　→ 聞きたいこと: 0002 と 0005 が二重です。消してよいですか']
    assert '要判断 1' in va._plan_line(plan) and '消してよいですか' in va._plan_line(plan)
    # 未が残っていれば今までどおり塞ぐ
    _s, _a, _d, _r, plan2 = va._parse_reply('{"plan":[{"item":"並べ替え","state":"未"}],"actions":[],"done":true}')
    assert va._plan_pending(plan2) == ['並べ替え']
    # 引き継ぎ（項目は減らさない・質問は残す）
    merged = va._merge_plan(plan, [{"item": "色分け", "state": "済"}])
    assert len(merged) == 2 and merged[1].get('ask')


def _inv_snap(**kw):
    base = {'used': 'A1:C3', 'values': [['a']], 'formulas': {}, 'hidden': (), 'cf': 0,
            'validation': '', 'filter': '', 'merged': None, 'tables': {}, 'charts': 0}
    base.update(kw)
    return base


def test_the_invariants_now_watch_merges_tables_charts_calc_and_the_sheet_order():
    """③: 頼んでいない変化に、結合・テーブル・グラフ・使用範囲・シートの並び・計算モードを足した。

    どれも「減る・消える・入れ替わる」方だけ咎める（作る・足すのは正当な仕事）。
    2026-09-06 に Claude for Excel が「検査する価値がある」と挙げた項目から、COM が安く読める分。
    """
    import vbam_agent as va
    before = {'sheets': {'名簿': _inv_snap(merged=None, tables={'T名簿': 'A1:E9'}, charts=1)},
              'names': {}, 'order': ['名簿', '目次'], 'calc': -4105}
    same = {'sheets': {'名簿': _inv_snap(merged=None, tables={'T名簿': 'A1:E9'}, charts=1)},
            'names': {}, 'order': ['名簿', '目次'], 'calc': -4105}
    assert va._inv_compare(before, same, '名簿') == []

    def bad(**kw):
        now = {'sheets': {'名簿': _inv_snap(**kw.pop('sheet', {'tables': {'T名簿': 'A1:E9'}, 'charts': 1}))},
               'names': {}, 'order': kw.pop('order', ['名簿', '目次']), 'calc': kw.pop('calc', -4105)}
        return " / ".join(va._inv_compare(before, now, '名簿'))

    assert '結合セルが全部外れた' in bad(sheet={'merged': False, 'tables': {'T名簿': 'A1:E9'}, 'charts': 1})
    assert 'テーブル「T名簿」が消えた' in bad(sheet={'tables': {}, 'charts': 1})
    assert '範囲が狭くなった' in bad(sheet={'tables': {'T名簿': 'A1:E5'}, 'charts': 1})
    assert 'グラフが減った' in bad(sheet={'tables': {'T名簿': 'A1:E9'}, 'charts': 0})
    assert '使用範囲が' in bad(sheet={'used': 'A1:C99999', 'tables': {'T名簿': 'A1:E9'}, 'charts': 1})
    assert 'シートの並びが変わった' in bad(order=['目次', '名簿'])
    assert '計算モードが変わった' in bad(calc=-4135)
    # 足す方は違反ではない（テーブルを作る・グラフを作る・シートを増やす）
    grew = {'sheets': {'名簿': _inv_snap(tables={'T名簿': 'A1:F9', 'T集計': 'H1:I5'}, charts=2)},
            'names': {}, 'order': ['名簿', '新', '目次'], 'calc': -4105}
    assert va._inv_compare(before, grew, '名簿') == []
    # 依頼で結合を外した回は咎めない
    unmerged = {'sheets': {'名簿': _inv_snap(merged=False, tables={'T名簿': 'A1:E9'}, charts=1)},
                'names': {}, 'order': ['名簿', '目次'], 'calc': -4105}
    assert va._inv_compare(before, unmerged, '名簿', skip=('merged',)) == []
    assert va._box_shrank('A1:E10', 'A1:E5') is True and va._box_shrank('A1:E10', 'A1:F12') is False


def test_the_tool_adds_the_missing_report_sections_itself():
    """④: 要判断・落とし物・頼んでいない変化は、AI の作文でなく道具が report の末尾に足す。

    向こうの報告は 4 段（やったこと／できなかったこと／決めたこと／確認していないこと）。
    3 段目・4 段目を作文に任せると忘れられ、忘れたことは報告に出ない＝読む側が気づけない。
    """
    import vbam_agent as va
    assert '【やったこと】' in va.RULES and '【確認していないこと】' in va.RULES
    src = open('vbam_agent.py', encoding='utf-8').read()
    body = src[src.index('def run_agent('):src.index('# 実射（練習台・弾・答え合わせ・状態・写し取り・点数）は vbam_fire.py へ')]
    for head in ('【人に判断してほしいこと】', '【できなかったこと（往復が尽きました）】',
                 '【頼んでいない変化', '【報告に出た番地のうち、いま空のもの】'):
        assert head in body
    # 報告に出た番地を道具が読み直す素（問い番号を拾う誤検知は承知のうえで注意書きにする）
    # 範囲は左上だけ見る・色コードは番地でない（2026-09-06 実測。報告の #FFC7CE から 'FFC7' を、
    # 「A1:H3 のタイトル行」から H3 を拾って「いま空です」と言っていた＝報告に嘘が混じる）
    assert va._report_addrs('A2:E301 を読み戻し、B14 と D22 を直しました。') == ['A2', 'B14', 'D22']
    assert va._report_addrs('背景を薄い赤（#FFC7CE）にしました。') == []
    assert va._report_addrs('300 行の表です。F 列は触っていません。') == []


def test_the_gates_do_not_eat_the_turns_meant_for_the_work():
    """⑤: 仕上げ検査・不変条件・読み残し・採点の差し戻しは、依頼をこなす往復とは別枠にする。

    前は 4 回の枠を食い合い、最短でも「書く→検査の直し→採点の直し→done」で使い切っていた
    （2026-09-06・規則文を Claude for Excel に査読させて出た指摘）。
    """
    import vbam_agent as va
    assert va._GATE_EXTRA_TURNS == 4
    src = open('vbam_agent.py', encoding='utf-8').read()
    body = src[src.index('def run_agent('):src.index('# 実射（練習台・弾・答え合わせ・状態・写し取り・点数）は vbam_fire.py へ')]
    assert 'while turn < limit:' in body and 'limit = max_turns' in body
    # 4 つの関所＋事前の門＋承認の言葉（2026-09-06 夜）で枠が 1 回ずつ伸びる
    assert body.count('limit = min(limit + 1, max_turns + _GATE_EXTRA_TURNS)') == 6
    assert 'max_turns - turn' not in body                 # 残りの往復も伸びた枠で数える


def test_the_run_counts_how_much_of_what_it_wrote_is_a_formula():
    """⑥: 派生値の数式率。合計・割合を値で置くと、人が元の行を直しても黙って古いままになる。"""
    import vbam_agent as va
    rows = [{'after': '1250000', 'after_formula': '=SUM(D2:D9)'},   # 数式で書いた数
            {'after': '980000', 'after_formula': ''},               # 値で置いた数
            {'after': '青木 誠', 'after_formula': ''},               # 文字は数えない
            {'after': '', 'after_formula': ''}]                     # 消したセルは数えない
    assert va._formula_rate(rows) == (2, 1)
    assert va._formula_rate([{'after': '氏名', 'after_formula': ''}]) is None
    assert va._formula_rate([]) is None
    assert '導ける値' in va.RULES and '必ず数式で書く' in va.RULES


def test_the_same_request_on_an_unchanged_sheet_is_flagged_before_it_runs(tmp_path, monkeypatch):
    """⑦: 同じ依頼を、前の走行の直後と同じ姿のシートに撃とうとしたら材料で告げる（二重に足す事故）。"""
    import vbam_agent as va
    monkeypatch.setattr(vg, '_AGENT_NOTES_FILE', str(tmp_path / 'notes.json'))
    fp = {'used': 'A1:E9', 'head': '会員番号 氏名 金額', 'formulas': 3}
    va._notes_remember('名簿.xlsx', '名簿', '担当の列を足して', '足しました', fp)
    # 同じ依頼＋同じ姿＝二重になる
    warn = va._notes_recall('名簿.xlsx', '名簿', '担当の列を足して', fp)
    assert '同じ依頼を、この姿のシートに、前にも撃っています' in warn and '二重になります' in warn
    # 姿が変われば告げない（人が後から直した・別の仕事が入った）
    assert '二重になります' not in va._notes_recall('名簿.xlsx', '名簿', '担当の列を足して',
                                                    dict(fp, used='A1:F9'))
    # 依頼が違えば告げない。空白のゆれは同じ依頼と見る
    assert '二重になります' not in va._notes_recall('名簿.xlsx', '名簿', '金額を集計して', fp)
    assert va._request_key('担当の列を足して') == va._request_key(' 担当の列を 足して ')


def test_the_header_row_is_guessed_and_wareki_dates_are_understood():
    """⑩: 日本の表は 1 行目がタイトルで見出しが 3 行目のことがある。和暦も日付に直す。

    どちらも 2026-09-06 に Claude for Excel が「自分の失敗型」「DATEVALUE が通らない形」として挙げた分。
    """
    import vbam_agent as va
    grid = [['売上一覧', None, None], [None, None, None],
            ['日付', '担当', '金額'], ['2026/1/1', '佐藤', 100]]
    assert va._pick_header_row(grid) == (2, '日付 担当 金額')          # 0 起点＝3 行目
    assert va._pick_header_row([['日付', '担当'], ['2026/1/1', '佐藤']]) == (0, '日付 担当')
    assert va._pick_header_row([[None, None], [1, 2]]) is None        # 見出しらしい行が無い
    d = va._parse_wareki_text('令和6年3月1日')
    assert (d.year, d.month, d.day) == (2024, 3, 1)
    assert va._parse_wareki_text('R6.3.1').year == 2024
    assert va._parse_wareki_text('平成元年1月8日').year == 1989
    assert va._parse_wareki_text('昭和64年1月7日').year == 1989
    assert va._parse_wareki_text('ただの文字') is None
    assert va._parse_date_text('令和6年3月1日').year == 2024          # date の規則から通る
    assert va._parse_date_text('2026/1/5').year == 2026               # 西暦は今までどおり


def test_the_rules_text_says_what_the_reviewer_found_missing():
    """⑤: 規則文の直し（Claude for Excel の査読 25 件のうち、文で直す分）。"""
    import vbam_agent as va
    r = va.RULES
    assert '必ず守る 5 つ' in r and r.index('必ず守る 5 つ') < r.index('返事の形')   # 重い規則を先頭へ
    assert '承認の言葉' in r and '元の値が変わってよい' in r                        # 語の定義
    assert '並べ替える・書式を整える・列を足す' in r                                # sort と承認の衝突を解く
    assert '見出し行の推定' in r and '1 行目とは限らない' in r
    assert '数式の中身はこの制限を受けない' in r                                    # A:A 禁止は range の話
    assert '前提は、ラベルを付けたセルに置いて参照' in r
    assert 'read と安全な手を先に' in r                                             # 失敗で止まる話と並べる
    assert 'done は actions が空のときだけ有効' not in r                            # 道具が弾く＝規則文から外す


def test_a_formula_that_moved_down_a_row_is_not_called_crushed():
    """③の直し: 表の形が変わった回は、番地でなく数の増減で数式を見る（2026-09-06・実射で踏んだ）。

    合計の行がある表に 1 行足すと、=SUM が 1 行下がる。前は「C5 の数式が値に潰された」と言っていた
    （row の手なら inv_skip で外れるが、承認をもらって write_grid で書き直すと同じ形になる）。
    """
    import vbam_agent as va
    before = {'sheets': {'名簿': _inv_snap(values=[['日付', '金額'], ['1/5', 120000], ['合計', 120000]],
                                           formulas={'B3': '=SUM(B2:B2)'})},
              'names': {}, 'order': ['名簿'], 'calc': -4105}
    # 行が 1 つ増えて、合計が 1 行下がった（数式の数は同じ）＝咎めない
    grew = {'sheets': {'名簿': _inv_snap(values=[['日付', '金額'], ['1/5', 120000],
                                                 ['1/8', 50000], ['合計', 170000]],
                                         formulas={'B4': '=SUM(B2:B3)'})},
            'names': {}, 'order': ['名簿'], 'calc': -4105}
    assert va._inv_compare(before, grew, '名簿') == []
    # 形が変わっても、数式そのものが減ったら咎める
    crushed = {'sheets': {'名簿': _inv_snap(values=[['日付', '金額'], ['1/5', 120000],
                                                    ['1/8', 50000], ['合計', 170000]],
                                            formulas={})},
               'names': {}, 'order': ['名簿'], 'calc': -4105}
    assert '数式の数が減った' in " / ".join(va._inv_compare(before, crushed, '名簿'))
    # 形が同じなら今までどおり番地で見る（潰したセルを名指しできる）
    same_shape = {'sheets': {'名簿': _inv_snap(values=[['日付', '金額'], ['1/5', 120000], ['合計', 120000]],
                                               formulas={})},
                  'names': {}, 'order': ['名簿'], 'calc': -4105}
    assert 'B3: =SUM(B2:B2)' in " / ".join(va._inv_compare(before, same_shape, '名簿'))
    assert va._shape_changed({'values': [[1, 2]]}, {'values': [[1, 2, 3]]}) is True
    assert va._shape_changed({'values': [[1, 2]]}, {'values': [[1, 2]]}) is False


def test_other_sheet_references_may_shift_when_rows_are_inserted():
    """③の直し 2: 行を足すと Excel が他シートの参照を自分で直す。それは「壊れた」ではない。

    2026-09-06 の実射で踏んだ。練習台の状態「他シート参照」＝ =COUNTA(実射1!$A$1:$C$5) が
    行の挿入で $C$6 になっただけで不合格になっていた。#REF! に化けたときと、消えたときだけ咎める。
    """
    import vbam_agent as va
    before = {'sheets': {
        '名簿': _inv_snap(values=[['日付', '金額'], ['1/5', 120000], ['合計', 120000]],
                          formulas={'B3': '=SUM(B2:B2)'}),
        '集計': _inv_snap(values=[['件数', 3]], formulas={'B1': '=COUNTA(名簿!$A$1:$C$5)'})},
        'names': {}, 'order': ['名簿', '集計'], 'calc': -4105}
    # 名簿に 1 行足した → 集計の参照が $C$6 に伸びた（Excel が自分で直した）＝咎めない
    after = {'sheets': {
        '名簿': _inv_snap(values=[['日付', '金額'], ['1/5', 120000], ['1/8', 50000], ['合計', 170000]],
                          formulas={'B4': '=SUM(B2:B3)'}),
        '集計': _inv_snap(values=[['件数', 4]], formulas={'B1': '=COUNTA(名簿!$A$1:$C$6)'})},
        'names': {}, 'order': ['名簿', '集計'], 'calc': -4105}
    assert va._inv_compare(before, after, '名簿') == []
    # #REF! に化けたら咎める（行を消して参照が壊れた）
    broken = {'sheets': {
        '名簿': _inv_snap(values=[['日付', '金額'], ['合計', 120000]], formulas={'B2': '=SUM(B2:B2)'}),
        '集計': _inv_snap(values=[['件数', 0]], formulas={'B1': '=COUNTA(#REF!)'})},
        'names': {}, 'order': ['名簿', '集計'], 'calc': -4105}
    assert '数式が壊れた' in " / ".join(va._inv_compare(before, broken, '名簿'))
    # 数式が消えたのも咎める
    gone = {'sheets': {
        '名簿': _inv_snap(values=[['日付', '金額'], ['1/5', 120000], ['1/8', 50000], ['合計', 170000]],
                          formulas={'B4': '=SUM(B2:B3)'}),
        '集計': _inv_snap(values=[['件数', 4]], formulas={})},
        'names': {}, 'order': ['名簿', '集計'], 'calc': -4105}
    assert '数式が壊れた' in " / ".join(va._inv_compare(before, gone, '名簿'))
    # 形が変わっていない回に他シートの数式が変わったら、今までどおり咎める
    meddled = {'sheets': {
        '名簿': before['sheets']['名簿'],
        '集計': _inv_snap(values=[['件数', 9]], formulas={'B1': '=COUNTA(名簿!$A$1:$Z$99)'})},
        'names': {}, 'order': ['名簿', '集計'], 'calc': -4105}
    assert '数式が壊れた' in " / ".join(va._inv_compare(before, meddled, '名簿'))


# ---- 2026-09-06 夜: 残っていた輸入候補（承認の言葉・実行前の全手検査・不変条件 3 つ・指標 4 つ・手の追加・手順書 20 本）----

def test_approval_words_are_checked_by_the_tool():
    """査読 (h): 承認の言葉があるかを AI に解釈させず、道具が依頼文で確かめる。"""
    import vbam_agent as va
    ok = va._has_approval
    assert ok("重複した行は消してよい")
    assert ok("元の列を置き換えてよい")
    assert ok("空行を消して")                       # 強い動詞は命令形でも承認（元の値が消えると人が言っている）
    assert ok("リンクを切って値にして")
    assert ok("値を直接書いて")
    assert ok("空欄は上の値で埋めてよい")
    assert ok("表記を直してもいいです")
    assert not ok("表記を直して")                    # 弱い動詞の命令形は承認ではない
    assert not ok("見やすくして。並べ替えて。集計して。")
    assert not ok("空行は消してはいけない")
    assert not ok("元の値は消さないで")
    assert not ok("上書きしないで")
    assert not ok("")
    # --cont の答えは「はい」「お願いします」も承認（道具が聞いた問いへの返事）
    assert va._has_approval("はい、消してください", answer_mode=True)
    assert va._has_approval("お願いします", answer_mode=True) and not va._has_approval("お願いします")
    # 手順書は補足だけを見る（本文の説明に「置き換えてよい」が出ても承認ではない）
    body = "手順書「値を直す」\n承認: 補足に「置き換えてよい」があるときだけ\n補足: なし"
    assert not ok(va._approval_scope(body))
    assert ok(va._approval_scope(body.replace("補足: なし", "補足: 置き換えてよい")))
    assert va._approval_scope("空行を消して") == "空行を消して"
    # 承認が要る手
    assert va._needs_approval({"op": "clear_range", "range": "A1:B2", "overwrite": True})
    assert va._needs_approval({"op": "row_delete", "at": 3, "count": 1, "overwrite": True})
    assert va._needs_approval({"op": "row", "action": "delete", "at": 3, "overwrite": True})
    assert va._needs_approval({"op": "shape", "names": ["a"], "delete": True})
    assert va._needs_approval({"op": "normalize", "range": "A2:A9", "rules": ["trim"], "overwrite": True})
    assert not va._needs_approval({"op": "normalize", "range": "A2:A9", "rules": ["trim"], "to": "H2"})
    assert not va._needs_approval({"op": "row_insert", "at": 3})
    assert not va._needs_approval({"op": "sort", "range": "A1:C9", "key": "B"})


def test_cells_the_tool_wrote_in_this_run_are_not_someone_elses(tmp_path):
    """承認の門を入れた直後の実射で踏んだ: 自分が同じ走行で書いた式を直せなかった（人の値と見なした）。"""
    import vbam_agent as va

    class _WS:                                        # 番地でも範囲でも中身を答える作り物のシート
        def __init__(self, cells):
            self.c = cells

        def Range(self, a):
            a = str(a).replace('$', '')
            if ':' in a:
                r1, c1, r2, c2 = va._range_box(a)
                grid = tuple(tuple(self.c.get(f"{va._col_letter(c)}{r}", '') for c in range(c1, c2 + 1))
                             for r in range(r1, r2 + 1))
                return type("R", (), {"Formula": grid})()
            return type("R", (), {"Formula": self.c.get(a, '')})()

    va._own_reset()
    va._approval_set("")                              # 承認の言葉なし
    ws = _WS({'F4': '=IFERROR(VLOOKUP(A4,$I$6:$K$10,2,FALSE),"申込なし")', 'F5': '前任者のメモ'})
    wb = _GateWB(ws)
    act = {"op": "write_cells", "cells": {"F4": "=VLOOKUP(--A4,$I$6:$K$10,2,FALSE)", "F5": "x"}}
    assert va._overwrite_hits(act, "名簿", wb) == ['F4', 'F5']
    # F4 は自分が書いた（fill で）と覚えたら、人の値は F5 だけ
    va._own_record({"op": "fill", "range": "F4:F4", "formula": "=1"}, "名簿")
    assert va._own_has("名簿", "F4") and not va._own_has("名簿", "F5")
    assert va._overwrite_hits(act, "名簿", wb) == ['F5']
    # fill の overwrite も、書き先が自分のセルだけなら承認なしで通る（人の値があれば止まる）
    va._own_record({"op": "write_grid", "range": "G1", "rows": [["a", "b"], ["c", "d"]]}, "名簿")
    assert va._own_has("名簿", "H2")
    assert va._prevalidate([{"op": "fill", "range": "F4:F4", "formula": "=2", "overwrite": True}], "名簿", wb) == {}
    got = va._prevalidate([{"op": "fill", "range": "F4:F5", "formula": "=2", "overwrite": True}], "名簿", wb)
    assert list(got) == [0] and 'F5' in got[0] and va._APPROVAL_MARK in got[0]
    # 行・列が動いたら番地がずれる＝忘れる（以後は人の値として扱う）
    va._own_record({"op": "row_insert", "at": 2}, "名簿")
    assert not va._own_has("名簿", "F4")
    # normalize の to（右の整形列）と header、copy_range の貼り先も覚える
    va._own_record({"op": "normalize", "range": "C2:C4", "rules": ["trim"], "to": "H2", "header": "整形"}, "名簿")
    assert va._own_has("名簿", "H4") and va._own_has("名簿", "H1") and not va._own_has("名簿", "H5")
    va._own_record({"op": "copy_range", "src": "A1:B2", "dst": "J5"}, "名簿")
    assert va._own_has("名簿", "K6") and not va._own_has("名簿", "L6")
    # --cont は前回の明細から読み戻す
    p = tmp_path / "changes.tsv"
    p.write_text("種類\tシート\t番地\t前\t後\t前の数式\t後の数式\n値\t名簿\tF4\t\t1\t\t\n書式・列幅\t名簿\tA\t8\t12\t\t\n",
                 encoding="utf-8-sig")
    va._own_reset()
    assert va._own_load_from_changes(str(p)) == 1 and va._own_has("名簿", "F4")
    va._own_reset()
    va._approval_set("")


def test_prevalidate_refuses_the_whole_reply_when_one_hand_is_bad():
    """輸入候補 12: 実行前に全部の手を見て、1 つでも通らなければ 1 手も実行しない（並べ替えはしない）。"""
    import vbam_agent as va
    va._approval_set("この表を見やすくして")           # 承認の言葉なし
    acts = [{"op": "read", "range": "A1:C9"},
            {"op": "row_delete", "at": 3, "count": 1, "overwrite": True},
            {"op": "tidy", "ranges": ["A1:C9"]}]
    problems = va._prevalidate(acts, "名簿", None)
    assert list(problems) == [1] and va._APPROVAL_MARK in problems[1]
    results = va._execute(acts, "名簿", None)
    assert [ok for _l, ok, _t in results] == [False, False, False]
    assert va._APPROVAL_MARK in results[1][2] and "1 手も実行していません" in results[0][2]
    # 承認があれば通る（wb が無いので門は見ない）
    va._approval_set("3 行目は消してよい")
    assert va._prevalidate(acts, "名簿", None) == {}
    # 形の違う手・許していない手・オブジェクトでない要素も一緒に返る
    va._approval_set("")
    bad = [{"op": "write_grid", "range": "A1"}, {"op": "hide_rows"}, "x"]
    assert set(va._prevalidate(bad, "名簿", None)) == {0, 1, 2}
    va._approval_set("")


def test_row_col_split_names_and_new_hands_are_wired():
    import pytest
    import vbam_agent as va
    label, toks = va._action_to_tokens({"op": "row_insert", "at": 7, "count": 2}, "名簿")
    assert toks == ["row", "insert", "7", "2", "--sheet", "名簿"] and label == "row insert 7 x2"
    label, toks = va._action_to_tokens({"op": "col_delete", "at": "C", "overwrite": True}, "名簿")
    assert toks == ["col", "delete", "C", "1", "--sheet", "名簿"]
    with pytest.raises(ValueError):
        va._action_to_tokens({"op": "row_delete", "at": 7}, "名簿")          # overwrite が要る
    for op in ("row_insert", "row_delete", "col_insert", "col_delete", "row_group", "col_group",
               "image", "eval", "read_sheet"):
        assert op in va._ALLOWED_OPS
    for op in ("row_group", "col_group", "image", "eval", "read_sheet"):
        assert op in va._DIRECT_OPS
    assert not va._act_writes({"op": "eval", "formula": "=SUM(A1:A3)"})
    assert not va._act_writes({"op": "read_sheet", "sheet": "名簿"})
    assert va._act_writes({"op": "image", "path": "x.png", "at": "A1"})
    assert va._inv_skip_for([{"op": "row_delete", "at": 3, "overwrite": True}]) == ('formula', 'hidden', 'tables')
    assert va._inv_skip_for([{"op": "view", "freeze": "A2"}]) == ('frozen',)
    assert va._inv_skip_for([{"op": "chart_config", "action": "set-source", "chart": "g", "args": ["A1:B3"]}]) == ('chart_src',)
    r = va.RULES
    assert '"op":"row_insert"' in r and '"op":"col_delete"' in r and '"op":"image"' in r and '"op":"eval"' in r
    assert '"op":"read_sheet"' in r and '"op":"row_group"' in r and '"op":"inspect"' not in r
    assert '新しく作る表の型' in r and '行や列を隠さない' in r and '承認の言葉があるかは道具も依頼文で確かめる' in r
    assert 'Table.TransformColumnTypes' not in r and 'Table.TransformColumnTypes' in va._DATAMODEL_NOTE
    assert va._wants_datamodel_note("売上をデータモデルに載せてメジャーを作って", "") is True
    assert va._wants_datamodel_note("この表を見やすくして", "") is False
    # 材料の「0 本」には反応しない（最初の実射で毎回足していた）／1 本でも載っていれば足す
    zero = "パワークエリ: 0 本 \nデータモデル: テーブル 0 本   メジャー 0 本\n"
    assert va._wants_datamodel_note("この表を見やすくして", zero) is False
    assert va._wants_datamodel_note("この表を見やすくして", zero.replace("メジャー 0 本", "メジャー 2 本")) is True
    assert va._wants_datamodel_note("この表を見やすくして", "パワークエリ: 1 本 \n") is True
    assert '元からそうだったもの' in va._GRADE_PROMPT and '元から隠れている行・列' in va.RULES
    assert va._eval_text(-2146826281) == '#DIV/0!' and va._eval_text(12.0) == '12'


def test_new_invariants_frozen_chart_source_circular():
    import copy
    import vbam_agent as va
    base = {'sheets': {'名簿': _inv_snap(values=[['a', 1]], formulas={})}, 'names': {}, 'order': ['名簿'],
            'calc': -4105, 'frozen': {'名簿': (1, 0)}}
    base['sheets']['名簿']['chart_src'] = {'グラフ 1': ['=SERIES(,名簿!$A$2:$A$5,名簿!$B$2:$B$5,1)']}
    base['sheets']['名簿']['circular'] = ''
    now = copy.deepcopy(base)
    assert va._inv_compare(base, now, '名簿') == []
    now['frozen'] = {'名簿': (0, 0)}
    assert '枠固定が外れた' in " / ".join(va._inv_compare(base, now, '名簿'))
    assert va._inv_compare(base, now, '名簿', skip=('frozen',)) == []
    now = copy.deepcopy(base)
    now['sheets']['名簿']['chart_src'] = {'グラフ 1': ['=SERIES(,#REF!,#REF!,1)']}
    assert '#REF! に化けた' in " / ".join(va._inv_compare(base, now, '名簿'))
    now['sheets']['名簿']['chart_src'] = {'グラフ 1': []}
    assert '系列が減った' in " / ".join(va._inv_compare(base, now, '名簿'))
    now = copy.deepcopy(base)
    now['sheets']['名簿']['circular'] = 'B3'
    assert '循環参照ができた' in " / ".join(va._inv_compare(base, now, '名簿'))
    # 古い控え（新しい鍵が無い）でも落ちない
    old = {'sheets': {'名簿': _inv_snap(values=[['a', 1]], formulas={})}, 'names': {}, 'order': ['名簿'], 'calc': -4105}
    assert va._inv_compare(old, copy.deepcopy(old), '名簿') == []


# --- 控えを取る側（COM を読む）の偽物。_snap_sheet / _snapshot_book / _pivot_sheets_in 用 ---
class _InvObj:
    def __init__(self, **kw):
        self.__dict__.update(kw)


class _InvList(list):
    @property
    def Count(self):
        return len(self)


class _InvRange:
    def __init__(self, address, value, formula, visible=None, validation=None, cf=0, merged=False):
        self.Address, self.Value, self.Formula = address, value, formula
        self.Rows, self.Columns = _InvObj(Count=len(value)), _InvObj(Count=len(value[0]))
        self.FormatConditions, self.MergeCells = _InvObj(Count=cf), merged
        self._special = {12: visible, -4174: validation}

    def SpecialCells(self, kind):
        got = self._special.get(kind)
        if got is None:
            raise RuntimeError("該当するセルが見つかりません")      # COM と同じく「無い」は例外
        return _InvObj(Address=got)


class _InvSheet:
    def __init__(self, name, used, filter_range=None, tables=None, charts=None, circular=None, pivots=()):
        self.Name, self.UsedRange = name, used
        self.AutoFilterMode = filter_range is not None
        self.AutoFilter = _InvObj(Range=_InvObj(Address=filter_range))
        self.ListObjects = _InvList(_InvObj(Name=n, Range=_InvObj(Address=a)) for n, a in (tables or {}).items())
        self.CircularReference = None if circular is None else _InvObj(Address=circular)
        self._charts, self._pivots = charts or {}, pivots

    def ChartObjects(self):
        out = _InvList()
        for name, forms in self._charts.items():
            series = _InvList(_InvObj(Formula=f) for f in forms)
            out.append(_InvObj(Name=name, Chart=_InvObj(
                SeriesCollection=lambda k=None, s=series: s if k is None else s[k - 1])))
        return out

    def PivotTables(self):
        return _InvList(_InvObj(Name=n) for n in self._pivots)


def _inv_book(sheets, names=None, active=None, freeze=None, calc=-4105):
    wb = _InvObj(Name='台.xlsx', Worksheets=sheets, Sheets=sheets,
                 Names=_InvList(_InvObj(Name=n, RefersTo=r) for n, r in (names or {}).items()))
    fr = freeze or (0, 0)
    wb.Application = _InvObj(ActiveWorkbook=_InvObj(Name=active or wb.Name),
                             ActiveWindow=_InvObj(FreezePanes=bool(freeze), SplitRow=fr[0], SplitColumn=fr[1]),
                             ActiveSheet=_InvObj(Name=sheets[0].Name), Calculation=calc)
    return wb


def test_the_snapshot_reads_what_the_invariants_compare(monkeypatch):
    """不変条件の控えを取る側（_snap_sheet / _snapshot_book）を偽のブックで確かめる（2026-09-10）。

    照らす側（_inv_compare）の 14 項目にはテストがあったが、控えが正しく取れているかは誰も見ていなかった。
    控えを読み違えれば、照らす側が正しくても「何も変わっていない」で素通りする。
    """
    import vbam_agent as va
    used = _InvRange('$A$1:$C$3',
                     value=(('名前', '点', '合計'), ('佐藤', '5', None), ('鈴木', '7', '12')),
                     formula=(('名前', '点', '合計'), ('佐藤', '5', ''), ('鈴木', '7', '=SUM(B2:B3)')),
                     visible='$A$1:$C$1,$A$3:$C$3', validation='$B$2:$B$3', cf=2, merged=False)
    sh = _InvSheet('名簿', used, filter_range='$A$1:$C$3', tables={'T名簿': '$A$1:$C$3'},
                   charts={'グラフ 1': ['=SERIES(,名簿!$A$2:$A$3,名簿!$B$2:$B$3,1)']})
    s = va._snap_sheet(sh)
    assert s['used'] == 'A1:C3'
    assert s['values'] == [['名前', '点', '合計'], ['佐藤', '5', ''], ['鈴木', '7', '12']]
    assert s['formulas'] == {'C3': '=SUM(B2:B3)'}            # 数式だけを番地で持つ
    assert s['hidden'] == (2,)                                # 使用範囲の行 − 見えている行
    assert (s['cf'], s['validation'], s['merged'], s['filter']) == (2, 'B2:B3', False, 'A1:C3')
    assert s['tables'] == {'T名簿': 'A1:C3'} and s['charts'] == 1
    assert s['chart_src'] == {'グラフ 1': ['=SERIES(,名簿!$A$2:$A$3,名簿!$B$2:$B$3,1)']}
    assert s['circular'] == ''
    # 入力規則が無い・隠れ行が無い（SpecialCells が例外）→ 空のまま
    bare = _InvRange('$A$1:$B$2', value=(('a', 'b'), ('c', 'd')), formula=(('a', 'b'), ('c', 'd')))
    s2 = va._snap_sheet(_InvSheet('素', bare, circular='$B$2'))
    assert (s2['validation'], s2['hidden'], s2['filter'], s2['circular']) == ('', (), '', 'B2')
    # 上限を超えるシートは値だけ（数式は見ない）
    monkeypatch.setattr(vi, '_INV_MAX_CELLS', 4)   # _snap_sheet は vbam_inv（2026-09-11）
    s3 = va._snap_sheet(sh)
    assert s3['formulas'] == {} and len(s3['values']) == 3
    monkeypatch.undo()
    # 何も読めないシートでも落ちない
    s4 = va._snap_sheet(_InvObj(Name='壊'))
    assert s4['used'] == '' and s4['formulas'] == {} and s4['tables'] == {}

    # ブック: 名前定義・並び・計算モード・枠固定（アクティブなシートだけ）
    wb = _inv_book([sh, _InvObj(Name='壊')], names={'単価': '=名簿!$B$2'}, freeze=(1, 0))
    b = va._snapshot_book(wb)
    assert list(b['sheets']) == ['名簿', '壊'] and b['names'] == {'単価': '=名簿!$B$2'}
    assert b['order'] == ['名簿', '壊'] and b['calc'] == -4105 and b['frozen'] == {'名簿': (1, 0)}
    assert 'frozen' not in va._snapshot_book(_inv_book([sh], active='別.xlsx'))   # 別のブックが前にいる

    # 控え → 現物を変える → 控え、を照らすと、読んだ側の変化がそのまま違反に出る
    sh.ListObjects = _InvList()
    used.Formula = (('名前', '点', '合計'), ('佐藤', '5', ''), ('鈴木', '7', '12'))
    wb.Application.ActiveWindow.FreezePanes = False
    bad = " / ".join(va._inv_compare(b, va._snapshot_book(wb), '名簿'))
    assert 'テーブル「T名簿」が消えた' in bad and 'C3: =SUM(B2:B3)' in bad and '枠固定が外れた' in bad


def test_pivot_sheets_in_adds_only_the_sheet_of_the_named_pivot():
    """ピボットを直す手が名指ししたピボットの出力シートだけを「書いたシート」に足す（2026-09-10 にテストを足した）。"""
    import vbam_agent as va
    wb = _InvObj(Worksheets=[_InvSheet('元', None), _InvSheet('集計', None, pivots=('P売上',)),
                             _InvSheet('別集計', None, pivots=('P在庫',))])
    assert va._pivot_sheets_in([{"op": "pivot_field", "pivot": "P売上", "action": "add_row"}], wb) == {'集計'}
    assert va._pivot_sheets_in([{"op": "slicer", "source": " P在庫 "}], wb) == {'別集計'}
    assert va._pivot_sheets_in([{"op": "pivot_calc", "pivot": "P売上"},
                                {"op": "pivot_calc", "pivot": "P在庫"}], wb) == {'集計', '別集計'}
    # 直す手でない・見つからない・形が違う・読めないブックは空
    assert va._pivot_sheets_in([{"op": "pivot", "pivot": "P売上"}, {"op": "write_cells"}], wb) == set()
    assert va._pivot_sheets_in([{"op": "pivot_field", "pivot": "無い"}], wb) == set()
    assert va._pivot_sheets_in(["文字列", None], wb) == set()
    assert va._pivot_sheets_in([{"op": "pivot_field", "pivot": "P売上"}], _InvObj()) == set()


def test_the_invariants_catch_a_table_that_shrank_but_not_one_that_grew():
    import copy
    import vbam_agent as va
    base = {'sheets': {'名簿': _inv_snap(tables={'T名簿': 'A1:E9'})}, 'names': {}, 'order': ['名簿'], 'calc': -4105}
    now = copy.deepcopy(base)
    now['sheets']['名簿']['tables'] = {'T名簿': 'A1:E5'}
    assert 'テーブル「T名簿」の範囲が狭くなった（A1:E9 → A1:E5）' in " / ".join(va._inv_compare(base, now, '名簿'))
    assert va._inv_compare(base, now, '名簿', skip=('tables',)) == []
    now['sheets']['名簿']['tables'] = {'T名簿': 'A1:D9'}                 # 列が減っても狭くなった
    assert '範囲が狭くなった' in " / ".join(va._inv_compare(base, now, '名簿'))
    now['sheets']['名簿']['tables'] = {'T名簿': 'A1:F12'}                # 広がるのは正当
    assert va._inv_compare(base, now, '名簿') == []


# --- read_file（別のファイルを読む手）。2026-09-09 に足したが pytest が 0 本だった（2026-09-10） ---
class _RfSheet:
    def __init__(self, name):
        self.Name, self.cols, self.last = name, {}, None

    def Columns(self, c):
        return self.cols.setdefault(c, _InvObj(NumberFormat='General'))

    def Cells(self, r, c):
        return (r, c)

    def Range(self, a, b):
        self.last = _InvObj(a=a, b=b, Value=None)
        return self.last


class _RfSheets(list):
    @property
    def Count(self):
        return len(self)

    def __call__(self, key):
        if isinstance(key, int):
            return self[key - 1]
        return next(s for s in self if s.Name == key)

    def Add(self, After=None):
        ws = _RfSheet(f"Sheet{len(self) + 1}")
        self.append(ws)
        return ws


def _rf_book(*names):
    sheets = _RfSheets(_RfSheet(n) for n in names)
    return _InvObj(Sheets=sheets, Worksheets=sheets, Application=_InvObj(Workbooks=_InvObj(Count=0)))


def test_read_file_checks_the_path_and_guesses_the_encoding(tmp_path):
    import pytest
    import vbam_agent as va
    with pytest.raises(ValueError, match='path にファイルの場所'):
        va._read_file_path('  ')
    with pytest.raises(ValueError, match='読めるのは'):
        va._read_file_path(str(tmp_path / 'a.doc'))
    with pytest.raises(ValueError, match='ファイルがありません'):
        va._read_file_path(str(tmp_path / '無い.csv'))
    f = tmp_path / '旧.CSV'
    f.write_bytes(b'a,b\n')
    assert va._read_file_path(f'"{f}"') == (str(f), '.csv')          # 引用符は外す・拡張子は小文字で見る
    assert va._read_text_decode(b'\xef\xbb\xbfabc') == ('abc', 'utf-8-sig')
    assert va._read_text_decode('日本'.encode('utf-8')) == ('日本', 'utf-8')
    assert va._read_text_decode('日本'.encode('cp932')) == ('日本', 'cp932')   # 古い CSV


def test_read_file_reads_csv_tsv_and_books(tmp_path):
    import openpyxl
    import pytest
    import vbam_agent as va
    c = tmp_path / '名簿.csv'
    c.write_bytes('番号,氏名\n0001,佐藤\n0002,鈴木\n'.encode('cp932'))
    assert va._read_csv_rows(str(c), 100) == ([['番号', '氏名'], ['0001', '佐藤'], ['0002', '鈴木']], 'cp932')
    assert len(va._read_csv_rows(str(c), 2)[0]) == 2                          # limit で切る
    t = tmp_path / '名簿.tsv'
    t.write_text('番号\t氏名,敬称\n1\t佐藤,様\n', encoding='utf-8')
    assert va._read_csv_rows(str(t), 100)[0] == [['番号', '氏名,敬称'], ['1', '佐藤,様']]
    x = tmp_path / '台帳.xlsx'
    bk = openpyxl.Workbook()
    bk.active.title = '一覧'
    bk.active.append(['番号', '氏名'])
    bk.active.append([1, '佐藤'])
    bk.create_sheet('別').append(['x'])
    bk.save(x)
    assert va._read_book_rows(str(x), None, 100) == ([['番号', '氏名'], [1, '佐藤']], '一覧', ['一覧', '別'])
    assert va._read_book_rows(str(x), '別', 100)[0] == [['x']]
    assert va._read_book_rows(str(x), None, 1)[0] == [['番号', '氏名']]
    with pytest.raises(ValueError, match='そのシートがありません'):
        va._read_book_rows(str(x), '無い', 100)


def test_read_file_names_the_new_sheet_and_keeps_leading_zeros():
    import vbam_agent as va
    wb = _rf_book('名簿', '取込_旧', '取込_旧_2')
    assert va._new_sheet_name(wb, None, 'C:/受領/旧.csv') == '取込_旧_3'     # 同名があれば番号
    assert va._new_sheet_name(wb, 'a:b/c[1]', 'x.csv') == 'a_b_c_1_'        # シート名に使えない字
    assert len(va._new_sheet_name(wb, 'あ' * 40, 'x.csv')) == 31
    ws = _RfSheet('新')
    assert va._read_file_write(ws, [['0001', '名'], ['0002', '佐藤', '様'], ['12', '鈴木']]) == (3, 3, ['A'])
    assert ws.cols[1].NumberFormat == '@' and 2 not in ws.cols                  # 先頭ゼロの列だけ文字に
    assert ws.last.Value == (('0001', '名', None), ('0002', '佐藤', '様'), ('12', '鈴木', None))
    assert va._read_file_write(_RfSheet('空'), []) == (0, 0, [])


def test_read_file_adds_a_sheet_and_never_writes_the_source(tmp_path):
    import pytest
    import vbam_agent as va
    src = tmp_path / '旧システム.csv'
    raw = '番号,氏名\n0001,佐藤\n0002,鈴木\n'.encode('cp932')
    src.write_bytes(raw)
    wb = _rf_book('名簿')
    name, ok, out = va._direct_action({'op': 'read_file', 'path': str(src)}, '名簿', wb)
    assert (name, ok) == ('read_file', True)
    assert [s.Name for s in wb.Sheets] == ['名簿', '取込_旧システム']
    assert '3 行 x 2 列' in out and '先頭ゼロを守った列: A' in out and '元のファイルには書いていません' in out
    assert wb.Sheets('取込_旧システム').last.Value[1] == ('0001', '佐藤')
    assert src.read_bytes() == raw                                               # 元のファイルは 1 バイトも変えない
    empty = tmp_path / '空.csv'
    empty.write_bytes(b'')
    with pytest.raises(ValueError, match='中身が空'):
        va._direct_action({'op': 'read_file', 'path': str(empty)}, '名簿', wb)


def test_the_request_paths_are_peeked_before_the_first_turn(tmp_path):
    import vbam_agent as va
    assert va._paths_in_request('「C:/受領/旧 システム.csv」と C:\\data\\a.xlsx と D:/b.tsv を突き合わせて') == \
        ['C:/受領/旧 システム.csv', 'C:\\data\\a.xlsx']                         # 空白を含む道を切らない・2 本まで
    assert va._paths_in_request('この表を見やすくして') == []
    f = tmp_path / '会員.csv'
    f.write_text('会員番号,氏名\nA01,佐藤\nA02,鈴木\n', encoding='utf-8')
    note = va._peek_files_note(f'「{f}」と突き合わせて')
    assert '全 3 行（これで全部）' in note and 'A02 | 鈴木' in note and '"op":"read_file"' in note
    assert '\\' not in note.split('取り込む手:')[1]                             # 手の見本は「/」で書く
    assert '**ありません**' in va._peek_files_note(f'「{tmp_path / "無い.csv"}」を取り込んで')


def test_the_score_ledger_keeps_why_a_case_failed(tmp_path, monkeypatch, capsys):
    """点数台帳に落ちた理由を残す（2026-09-10）。83 本全滅の行に理由が無かった。"""
    import json
    import vbam_agent as va
    f = tmp_path / 'score.jsonl'
    monkeypatch.setattr(vf, '_AGENT_SCORE_FILE', str(f))
    vf._score_save({'time': 't', 'ai': 'claude-code', 'model': 'sonnet', 'n': 2, 'ok': 1, 'sec': 3.0,
                    'cases': {'郵便番号の統一': True, '鈴木さんの分だけ': False},
                    'why': {'鈴木さんの分だけ': 'HTTP 429 残高なし'}})
    assert json.loads(f.read_text(encoding='utf-8'))['why'] == {'鈴木さんの分だけ': 'HTTP 429 残高なし'}
    vf.score_history()
    assert '└ 鈴木さんの分だけ: HTTP 429 残高なし' in capsys.readouterr().out
    src = open(vf.__file__, encoding='utf-8').read()
    assert "**({'why': why} if why else {})" in src                             # 実射の記録に載せている


def test_every_hand_translates_to_what_the_cli_accepts():
    """pytest にも実射にも出ていなかった手の翻訳（JSON → CLI の引数列）を 1 本ずつ当てる（2026-09-10）。

    重い道具の action・chart_config の 11 本・sheet_op・format の全項目。写した引数列を CLI の parser が
    そのまま受けること、CLI 側にその action の受け口が本当にあることまで見る（名前のずれで空振りしない）。
    """
    import pytest
    import vbam_agent as va
    P = {"pivot": "P"}
    samples = [
        # 重い道具
        ({"op": "table", "action": "list"}, ["table", "list"]),
        ({"op": "table", "action": "read", "name": "T"}, ["table", "read", "T"]),
        ({"op": "table", "action": "filter", "name": "T", "column": "金額", "criteria": ">100"},
         ["table", "filter", "T", "金額", ">100"]),
        ({"op": "table", "action": "filter_clear", "name": "T"}, ["table", "filter-clear", "T"]),
        ({"op": "table", "action": "create", "range": "A1:C5", "no_headers": True},
         ["table", "create", "'名簿'!A1:C5", "--no-headers"]),
        ({"op": "pivot_field", "action": "add_row", **P, "field": "地域"}, ["pivot-field", "add-row", "P", "地域"]),
        ({"op": "pivot_field", "action": "add_col", **P, "field": "担当"}, ["pivot-field", "add-col", "P", "担当"]),
        ({"op": "pivot_field", "action": "add_filter", **P, "field": "年"}, ["pivot-field", "add-filter", "P", "年"]),
        ({"op": "pivot_field", "action": "remove", **P, "field": "地域"}, ["pivot-field", "remove", "P", "地域"]),
        ({"op": "pivot_field", "action": "set_func", **P, "field": "sum/売上", "func": "AVERAGE"},
         ["pivot-field", "set-func", "P", "sum/売上", "average"]),
        ({"op": "pivot_field", "action": "set_name", **P, "field": "sum/売上", "name": "売上計"},
         ["pivot-field", "set-name", "P", "sum/売上", "売上計"]),
        ({"op": "pivot_field", "action": "set_format", **P, "field": "sum/売上", "format": "#,##0"},
         ["pivot-field", "set-format", "P", "sum/売上", "#,##0"]),
        ({"op": "pivot_field", "action": "group_date", **P, "by": "Months"},
         ["pivot-field", "group-date", "P", "*", "months"]),               # field を省けば道具が当てる
        ({"op": "pivot_calc", "action": "layout", **P, "layout": "tabular"}, ["pivot-calc", "layout", "P", "tabular"]),
        ({"op": "pivot_calc", "action": "layout", **P}, ["pivot-calc", "layout", "P", "compact"]),
        ({"op": "pivot_calc", "action": "subtotals", **P, "field": "地域", "on": False},
         ["pivot-calc", "subtotals", "P", "地域", "off"]),
        ({"op": "pivot_calc", "action": "calc_field_list", **P}, ["pivot-calc", "calc-field", "list", "P"]),
        ({"op": "slicer", "action": "list"}, ["slicer", "list"]),
        ({"op": "slicer", "action": "add", "source": "P", "field": "地域", "name": "S地域"},
         ["slicer", "add", "P", "地域", "--name", "S地域"]),
        ({"op": "powerquery", "action": "list"}, ["powerquery", "list"]),
        ({"op": "powerquery", "action": "refresh", "name": "Q"}, ["powerquery", "refresh", "Q"]),
        ({"op": "powerquery", "action": "edit", "name": "Q", "m": "let a = 1 in a", "desc": "説明"},
         ["powerquery", "edit", "Q", "--m", "let a = 1 in a", "--desc", "説明"]),
        ({"op": "powerquery", "action": "load", "name": "Q", "to": "model"}, ["powerquery", "load", "Q", "--to", "model"]),
        ({"op": "datamodel", "action": "list"}, ["datamodel", "list"]),
        ({"op": "datamodel", "action": "measure_add", "table": "T", "name": "m", "dax": "SUM(T[金額])",
          "format": "currency", "decimals": 0, "symbol": "¥"},
         ["datamodel", "measure", "add", "T", "m", "--dax", "SUM(T[金額])", "--format", "currency",
          "--decimals", "0", "--symbol", "¥"]),                          # 0 を落とさない
        # グラフ（11 本）
        ({"op": "chart_config", "action": "legend", "chart": "g", "args": ["bottom"]}, ["chart-config", "legend", "g", "bottom"]),
        ({"op": "chart_config", "action": "set-title", "chart": "g", "args": "売上"}, ["chart-config", "set-title", "g", "売上"]),
        ({"op": "chart_config", "action": "set-axis-title", "chart": "g", "args": ["value", "金額"]},
         ["chart-config", "set-axis-title", "g", "value", "金額"]),
        ({"op": "chart_config", "action": "axis-scale", "chart": "g", "min": 0, "max": 100, "major": 20, "value": True},
         ["chart-config", "axis-scale", "g", "--min", "0", "--max", "100", "--major", "20", "--value"]),
        ({"op": "chart_config", "action": "axis-format", "chart": "g", "args": ["#,##0"], "value": True},
         ["chart-config", "axis-format", "g", "#,##0", "--value"]),
        ({"op": "chart_config", "action": "gridlines", "chart": "g", "args": ["on"]}, ["chart-config", "gridlines", "g", "on"]),
        ({"op": "chart_config", "action": "style", "chart": "g", "args": [201]}, ["chart-config", "style", "g", "201"]),
        ({"op": "chart_config", "action": "set-type", "chart": "g", "args": ["line"]}, ["chart-config", "set-type", "g", "line"]),
        ({"op": "chart_config", "action": "set-source", "chart": "g", "args": ["A1:B5"]},
         ["chart-config", "set-source", "g", "A1:B5"]),
        ({"op": "chart_config", "action": "placement", "chart": "g", "args": ["free"]}, ["chart-config", "placement", "g", "free"]),
        ({"op": "chart_config", "action": "data-labels", "chart": "g", "percent": True, "category": True},
         ["chart-config", "data-labels", "g", "--percent", "--category"]),
        # シート
        ({"op": "sheet_op", "action": "activate", "name": "集計"}, ["sheet", "activate", "集計"]),
        ({"op": "sheet_op", "action": "add", "to": "新", "after": "名簿"}, ["sheet", "add", "新", "--after", "名簿"]),
        ({"op": "sheet_op", "action": "copy", "name": "名簿", "to": "名簿の写し", "before": "集計"},
         ["sheet", "copy", "名簿", "名簿の写し", "--before", "集計"]),
        ({"op": "sheet_op", "action": "tab_color", "name": "集計", "to": "#FF0000"},
         ["sheet", "tab-color", "集計", "#FF0000"]),
        # 書式（全項目）
        ({"op": "format", "range": "A1:C1", "bold": True, "italic": True, "color": "#FF0000", "bg": "-1",
          "font": "Meiryo UI", "size": 11, "number_format": "#,##0", "align": "center", "valign": "top", "wrap": True,
          "border": "thin", "col_width": 12, "row_height": 18, "merge": True},
         ["format-range", "--sheet", "名簿", "--show", "A1:C1", "--bold", "--italic", "--color=#FF0000", "--bg=-1",
          "--font=Meiryo UI", "--size", "11", "--number-format=#,##0", "--align", "center", "--valign", "top", "--wrap",
          "--border", "thin", "--col-width", "12", "--row-height", "18", "--merge"]),
        ({"op": "format", "range": "A2:C9", "unbold": True, "unmerge": True},
         ["format-range", "--sheet", "名簿", "--show", "A2:C9", "--unbold", "--unmerge"]),
    ]
    parser = vm.build_parser()
    for act, want in samples:
        _label, toks = va._action_to_tokens(act, "名簿")
        assert toks == want, (act, toks)
        ns, unknown = parser.parse_known_args(toks)
        assert not unknown and ns.command == toks[0], (toks, unknown)
    bads = [
        {"op": "chart_config", "action": "delete", "chart": "g"},
        {"op": "chart_config", "action": "legend", "args": ["bottom"]},                      # グラフ名が無い
        {"op": "sheet_op", "action": "delete", "name": "集計"},
        {"op": "sheet_op", "action": "hide", "name": "集計"},
        {"op": "sheet_op", "action": "rename", "name": "名簿", "to": "新"},                  # 対象シートの改名
        {"op": "sheet_op", "action": "copy", "name": "集計", "to": "名簿"},                  # 対象と同じ名前へ複写
        {"op": "sheet_op", "action": "tab_color", "name": "集計"},                           # 色が無い
        {"op": "sheet_op", "action": "activate"},                                            # 名前が無い
        {"op": "format", "range": "A1", "valign": "middle"},
        {"op": "format", "range": "A1", "border": "double"},
        {"op": "format", "range": "A1", "bold": False},                                      # 当てる項目が無い
        {"op": "pivot_field", "action": "set_func", **P, "field": "x", "func": "median"},
        {"op": "pivot_field", "action": "set_name", **P, "field": "x"},
        {"op": "pivot_field", "action": "remove", **P},                                      # field が無い
        {"op": "pivot_calc", "action": "subtotals", **P},
        {"op": "table", "action": "filter", "name": "T", "column": "金額"},                   # 条件が無い
    ]
    for bad in bads:
        with pytest.raises(ValueError):
            va._action_to_tokens(bad, "名簿")
    # 渡す action の名前には、CLI 側に受け口が本当にある（名前がずれると道具が「不明な action」で落ちる）
    here = os.path.dirname(va.__file__)
    heavy = open(os.path.join(here, 'vbam_heavy.py'), encoding='utf-8').read()
    edit = open(os.path.join(here, 'vbam_edit.py'), encoding='utf-8').read()
    for a in va._CHART_CFG_ACTIONS:
        assert f"'{a}'" in heavy, a
    for op in ('pivot_field', 'pivot_calc'):
        for a in va._HEAVY_ACTIONS[op]:
            name = 'calc-field' if a.startswith('calc_field_') else a.replace('_', '-')
            assert f"'{name}'" in heavy, (op, a)
    for a in va._SHEET_OP_ACTIONS:
        assert f"'{a}'" in edit, a


def test_plan_discovered_items_report_silence_and_ledger_metrics():
    import vbam_agent as va
    old = [{'item': 'A', 'state': '済'}]
    new = va._parse_plan([{'item': 'A', 'state': '済'}, {'item': 'B', 'state': '未', 'why': '文字の数値が混ざっていた'}])
    assert [p['item'] for p in va._plan_added(old, new)] == ['B'] and new[1]['why'] == '文字の数値が混ざっていた'
    assert va._merge_plan(old, new)[1].get('why') == '文字の数値が混ざっていた'
    assert vf._report_is_silent({'report': '【やったこと】A1 に書いた\n【できなかったこと】なし'})
    assert vf._report_is_silent({'report': '【やったこと】A1 に書いた'})
    assert not vf._report_is_silent({'report': '【やったこと】…\n【できなかったこと】写真は手が無い'})
    assert not vf._report_is_silent({'report': '', 'asks': ['どちらの列が正か']})
    rows = [{'time': '2026-09-06 20:00:00', 'book': 'a.xlsx', 'sheet': 'S', 'done': True, 'asks': 1},
            {'time': '2026-09-06 20:10:00', 'book': 'a.xlsx', 'sheet': 'S', 'done': True},
            {'time': '2026-09-06 22:00:00', 'book': 'a.xlsx', 'sheet': 'S', 'done': True}]
    assert va._rework_count(rows) == 1
    s = va._runs_summary(rows)
    assert '聞き返し 1/3' in s and '30 分以内の手直し 1/3' in s
    line = vf._first_shot_line([('x', True, '', {'gates': {'inv': 0}, 'report': 'ok'}),
                                ('y', False, '', {'gates': {'inv': 1}, 'report': '【できなかったこと】なし'})])
    assert '範囲外変化ゼロ 1/2' in line and '黙殺（不合格なのに報告に無い）1/1' in line
    assert 'approval' in va._GATE_LABELS


def test_plan_items_match_when_the_ai_rephrases_them():
    """2026-09-06 夕（シュウさんの初実射）: 言い換えで同じ項目が二重に出て、一覧 6 件・中身 3 件・往復 5 回。"""
    import vbam_agent as va
    g1 = 'A17セルの会員番号が全角数字「１０１２」のままなので半角「1012」に修正してください'
    a1 = 'A17の会員番号「１０１２」を半角「1012」に修正'
    g2 = '16行目のデータ（会員番号1008 中村大輔）が13行目と重複しているため削除してください'
    a2 = '重複している16行目のデータを削除'
    t1 = '表全体（A5:E20）の体裁を整える'
    t2 = '表全体の体裁を整える'
    assert va._plan_same(g1, a1) and va._plan_same(g2, a2) and va._plan_same(t1, t2)
    # 別の項目は別のまま（番地が食い違う・短い別の語）
    assert not va._plan_same('A 列を太字にする', 'A 列を左寄せにする')
    assert not va._plan_same('A17 を半角に直す', 'A18 を半角に直す')
    assert not va._plan_same('B 列の全角を半角に直す', 'C 列の全角を半角に直す')
    assert not va._plan_same('色分け', '並べ替え')
    old = [{'item': t1, 'state': '済'}, {'item': g1, 'state': '未'}, {'item': g2, 'state': '未'}]
    new = [{'item': a1, 'state': '要判断', 'ask': '上書きしてよいですか'},
           {'item': a2, 'state': '要判断', 'ask': '削除してよいですか'}, {'item': t2, 'state': '済'}]
    merged = va._merge_plan(old, new)
    assert len(merged) == 3 and [p['state'] for p in merged] == ['済', '要判断', '要判断']
    assert merged[1]['item'] == g1 and merged[1]['ask'] == '上書きしてよいですか'   # 前の文を残し state と ask を取る
    assert va._plan_pending(merged) == [] and va._plan_added(old, new) == []
    # 完全一致を先に組む＝列の名前だけ違う 2 項目が、順番が入れ替わっても取り違えない
    s1, s2 = '氏名の空白を全角1つに統一', 'フリガナの空白を全角1つに統一'
    m2 = va._merge_plan([{'item': s1, 'state': '未'}, {'item': s2, 'state': '未'}],
                        [{'item': s2, 'state': '未'}, {'item': s1, 'state': '済'}])
    assert [(p['item'], p['state']) for p in m2] == [(s1, '済'), (s2, '未')]


def test_grader_noticed_items_go_to_the_report_not_to_unmet():
    """採点係の「依頼には無いが気づいたこと」は往復を使わず報告に写す（2026-09-06 夕）。"""
    import vbam_agent as va
    t = '{"unmet": [], "noticed": ["C10 が空欄", "A17 が全角"], "note": "ok"}'
    assert va._parse_unmet(t) == [] and va._parse_noticed(t) == ['C10 が空欄', 'A17 が全角']
    assert va._parse_noticed('xx') == [] and va._parse_noticed('{"unmet": ["a"]}') == []
    assert '"noticed"' in va._GRADE_PROMPT and 'noticed に書く' in va._GRADE_PROMPT
    src = inspect.getsource(va.run_agent)
    assert '【気づいたこと】' in src and '_parse_noticed(gtext)' in src
    assert '同じ手をもう一度出さないでください' in src        # 承認の門の後に、同じ手を試させない
    assert 'actions を出さずに' in src                         # 採点の直しでも、承認が無ければ試させない
    assert '往復 {turn}/{limit}' in src                        # 「5/4」を出さない


def test_table_blanks_are_counted_in_materials():
    """テスト用2 の C10（フリガナ）が空欄なのに誰も言わなかった（2026-09-06 夕）。道具が数えて材料に出す。"""
    import vbam_agent as va
    grid = [['タイトル', None, None, None],
            [None, None, None, None],
            ['会員番号', '氏名', 'フリガナ', '備考'],
            ['1001', '佐藤', 'サトウ', None],
            ['1002', '鈴木', None, 'x'],
            ['1003', '田中', 'タナカ', None],
            [None, None, None, None],            # 区切りの空行は飛ばす
            ['1004', '高橋', 'タカハシ', None]]
    assert va._table_blanks(grid, 1, 1, 2) == ['C5']        # 見出し 3 行目・備考のようなまばらな列は数えない
    assert va._table_blanks(grid, 1, 1, None) == []          # 見出しが分からなければ 1 行目＝タイトルの下に列名が無い
    assert va._table_blanks([], 1, 1, 0) == []
    assert '表の中の空欄' in inspect.getsource(va._extra_materials)


def test_the_shape_name_from_the_materials_line_still_finds_the_shape():
    """材料の表示（名前「文字」）をそのまま渡されても図形に当たる（2026-09-06 の実射）。

    テスト用5 を 2 回撃って 2 回とも同じ所で転んだ: 材料が `ロゴ枠「総務課」` と出し、AI が
    その 1 行を name に写し、shape が「図形が無い」で全部の手を差し戻した（1 往復・約 8 秒の損）。
    """
    import vbam_agent as va

    class _Chars:
        def __init__(self, t):
            self.Text = t

    class _Frame:
        def __init__(self, t):
            self._t = t

        def Characters(self):
            return _Chars(self._t)

    class _Shape:
        def __init__(self, name, text=''):
            self.Name = name
            self.TextFrame = _Frame(text)

    class _Shapes:
        def __init__(self, items):
            self._items = items
            self.Count = len(items)

        def __call__(self, key):
            if isinstance(key, int):
                return self._items[key - 1]
            for s in self._items:
                if s.Name == key:
                    return s
            raise Exception('no such shape')

    class _WS:
        def __init__(self, items):
            self.Shapes = _Shapes(items)

    ws = _WS([_Shape('ロゴ枠', '総務課'), _Shape('はみ出した矢印'), _Shape('印刷ボタンもどき', '印刷')])
    assert va._pick_shape(ws, 'ロゴ枠').Name == 'ロゴ枠'                    # 名前そのもの
    assert va._pick_shape(ws, 'ロゴ枠「総務課」').Name == 'ロゴ枠'          # 材料の 1 行をそのまま
    assert va._pick_shape(ws, '印刷ボタンもどき 「印刷」').Name == '印刷ボタンもどき'   # 空白入り
    assert va._pick_shape(ws, '総務課').Name == 'ロゴ枠'                    # 表示テキストだけ
    assert va._pick_shape(ws, '無い図') is None                            # 無いものは無いと言う
    assert va._pick_shape(ws, '') is None


# ----------------------------------------------------------------
# 2026-09-06 夜: 18 項目の回収（手と done の同時受け・道具側の内訳・規則文の畳み込み・採点の門・気づきの畳み・
# 番号列の警告・鍵の伏せ・進み具合の見出し飛ばし・--repeat・控えの覚書の残骸・終わりの検査の使い回し・図形の全部消し）
# ----------------------------------------------------------------

def _loop_fakes(monkeypatch, tmp_path, replies, execute=None):
    """run_agent を API 無しで回す下地（既存のテストと同じ型）。sent＝AI に送った文の一覧。"""
    import vbam_agent as va
    sent = []

    def fake_ask(ai, model, key, history):
        sent.append(history[-1][1])
        return replies[min(len(sent) - 1, len(replies) - 1)], {'in': 1, 'out': 1}, 0.0

    monkeypatch.setattr(va, "_ask", fake_ask)
    monkeypatch.setattr(va, "_ai_setup", lambda ai, model: ('gemini', 'm', 'k'))
    monkeypatch.setattr(va, "_execute", execute or (lambda a, s, wb=None: [("write_cells 1 セル", True, "ok")]))
    monkeypatch.setattr(va, "_audit_now", lambda wb, sheet, targets: [])
    monkeypatch.setattr(va, "_secret_values", lambda wb: set())
    monkeypatch.setattr(va, "_LAST_AGENT_LOG_FILE", str(tmp_path / "log.jsonl"))
    monkeypatch.setattr(va, "_LAST_AGENT_ASK_FILE", str(tmp_path / "ask.txt"))
    wb = type("WB", (), {"Name": "b.xlsx", "Path": "", "Sheets": staticmethod(lambda name: object())})()
    return va, wb, sent


def test_merged_done_skips_the_report_only_turn(tmp_path, monkeypatch):
    """1: 手と done を同じ返事で返し、全部通れば 2 往復目（actions 空の done）を待たない。"""
    replies = ['{"say":"書きます","plan":[{"item":"A1 に書く","state":"済"}],'
               '"actions":[{"op":"write_cells","cells":{"A1":"x"}}],"done":true,"report":"【やったこと】A1 に x"}']
    va, wb, sent = _loop_fakes(monkeypatch, tmp_path, replies)
    r = va.run_agent("A1 に x と書いて", "名簿", wb, max_turns=4, materials="使用範囲: A1",
                     verify=False, backup=False, grade=False, show_image=False)
    assert r['done'] is True and r['turns'] == 1 and r['merged'] == 1 and len(sent) == 1
    assert "A1 に x" in r['report']


def test_clean_first_turn_skips_grading_and_tool_writes_the_did_section(tmp_path, monkeypatch):
    """一発で通った回は採点を省く（--grade で付け直せる）。【やったこと】は道具が通った手から作る（2026-09-11 午後）。"""
    replies = ['{"say":"書きます","plan":[{"item":"A1 に書く","state":"済"}],'
               '"actions":[{"op":"write_cells","cells":{"A1":"x"}}],"done":true,'
               '"report":"【できなかったこと】なし\\n【決めたこと】なし\\n【確認していないこと】なし"}',
               '満たしていない点はありません']
    va, wb, sent = _loop_fakes(monkeypatch, tmp_path, replies)
    r = va.run_agent("A1 に x と書いて", "名簿", wb, max_turns=4, materials="使用範囲: A1",
                     verify=False, backup=False, grade=True, show_image=False)
    assert r['done'] is True and r['turns'] == 1 and len(sent) == 1          # 採点役は呼ばれていない
    assert r['graded'] is False and r['grade_skipped'] is True
    assert r['report'].startswith("【やったこと】（道具が実行した手から）\n・write_cells 1 セル")
    assert "【決めたこと】なし" in r['report']
    # --grade なら今までどおり採点する（2 通目＝採点役）
    va, wb, sent = _loop_fakes(monkeypatch, tmp_path, replies)
    r = va.run_agent("A1 に x と書いて", "名簿", wb, max_turns=4, materials="使用範囲: A1",
                     verify=False, backup=False, grade=True, show_image=False, grade_always=True)
    assert r['graded'] is True and r['grade_skipped'] is False and len(sent) == 2
    # 2 往復かかった回（手の失敗）は採点を省かない
    replies2 = ['{"say":"書きます","plan":[{"item":"A1 に書く","state":"済"}],'
                '"actions":[{"op":"write_cells","cells":{"A1":"x"}}],"done":true,"report":"早い"}',
                '{"plan":[{"item":"A1 に書く","state":"済"}],"actions":[],"done":true,"report":"【やったこと】直した"}',
                '満たしていない点はありません']
    calls = []

    def execute(a, s, wb=None):
        calls.append(a)
        return [("write_cells", len(calls) > 1, "ok" if len(calls) > 1 else "書き先に値がある")]
    va, wb, sent = _loop_fakes(monkeypatch, tmp_path, replies2, execute)
    r = va.run_agent("A1 に x と書いて", "名簿", wb, max_turns=4, materials="使用範囲: A1",
                     verify=False, backup=False, grade=True, show_image=False)
    assert r['graded'] is True and r['grade_skipped'] is False and len(sent) == 3
    assert r['report'].startswith("【やったこと】直した")        # AI が書いたなら道具は足さない
    ns = vm.build_parser().parse_args(["agent", "x", "--grade"])
    assert ns.grade is True
    assert "3 段で書く" in va.RULES and "【やったこと】" in va.RULES


def test_merged_done_is_cancelled_when_a_hand_fails(tmp_path, monkeypatch):
    """1: 手が 1 つでも失敗したら done は取り消し、今までどおり結果を返して続く。"""
    replies = ['{"say":"書きます","plan":[{"item":"A1 に書く","state":"済"}],'
               '"actions":[{"op":"write_cells","cells":{"A1":"x"}}],"done":true,"report":"早すぎる報告"}',
               '{"plan":[{"item":"A1 に書く","state":"済"}],"actions":[],"done":true,"report":"【やったこと】直した"}']
    calls = []

    def execute(a, s, wb=None):
        calls.append(a)
        return [("write_cells", len(calls) > 1, "ok" if len(calls) > 1 else "書き先に値がある")]

    va, wb, sent = _loop_fakes(monkeypatch, tmp_path, replies, execute)
    r = va.run_agent("A1 に x と書いて", "名簿", wb, max_turns=4, materials="使用範囲: A1",
                     verify=False, backup=False, grade=False, show_image=False)
    assert r['merged'] == 0 and r['turns'] == 2 and r['done'] is True
    assert "（失敗）" in sent[1] and "書き先に値がある" in sent[1]
    assert r['report'].startswith("【やったこと】直した")


def test_merged_done_is_not_bounced_by_partially_done_plan(tmp_path, monkeypatch):
    """1: 手と done を同時に受けて手が全部通った回は、一覧に「未」が混じっていても突き返さない。

    2026-09-08: 免除が「一覧が全部未」の回に限られていたので、一部を「済」と書き分けた回＝
    正直に書いた回だけが突き返され、状態を書き換えるだけの往復を 1 回失っていた（本番で 12.7 秒）。
    """
    replies = ['{"say":"書きます","plan":[{"item":"A1 に書く","state":"済"},{"item":"色分け","state":"未"}],'
               '"actions":[{"op":"write_cells","cells":{"A1":"x"}}],"done":true,"report":"やった"}']
    va, wb, sent = _loop_fakes(monkeypatch, tmp_path, replies)
    r = va.run_agent("A1 に x と書いて色分けして", "名簿", wb, max_turns=4, materials="使用範囲: A1",
                     verify=False, backup=False, grade=False, show_image=False)
    assert r['turns'] == 1 and r['done'] is True and len(sent) == 1
    assert va._plan_pending(r['plan']) == []            # 道具が「済」に直している
    assert [p['state'] for p in r['plan']] == ['済', '済']


def test_merged_done_gate_message_carries_the_unseen_results(tmp_path, monkeypatch):
    """1: 同時受けの回を関所が差し戻すとき、AI はまだ結果を見ていないので結果文を先頭に付ける。"""
    replies = ['{"say":"書きます","plan":[{"item":"A1 に書く","state":"済"}],'
               '"actions":[{"op":"write_cells","cells":{"A1":"x"}}],"done":true,"report":"半分"}',
               '{"plan":[{"item":"A1 に書く","state":"済"}],'
               '"actions":[],"done":true,"report":"直しました"}']
    va, wb, sent = _loop_fakes(monkeypatch, tmp_path, replies)
    notes = [["A1:B2 に罫線がありません → tidy を当てる"]]      # 1 回目だけ指摘を返す仕上げ検査
    monkeypatch.setattr(va, '_audit_now', lambda *a, **k: notes.pop(0) if notes else [])
    monkeypatch.setattr(va, '_content_now', lambda *a, **k: ([], []))
    r = va.run_agent("A1 に x と書いて整えて", "名簿", wb, max_turns=4, materials="使用範囲: A1",
                     verify=False, backup=False, grade=False, show_image=False)
    assert r['turns'] == 2 and r['done'] is True
    assert "【前回の結果】" in sent[1] and "write_cells 1 セル" in sent[1] and "罫線がありません" in sent[1]


def test_rules_for_folds_the_heavy_block_unless_needed():
    """3: 重い道具の段は、依頼にも材料にもその語が無ければ 1 行に畳む。"""
    import vbam_agent as va
    plain = va._rules_for("表を整えて", "テーブル（ブック全体）: 0 本\nピボット（ブック全体）: 0 本")
    assert va._HEAVY_STUB in plain and "percent_parent_row" not in plain and '"op":"tidy"' in plain
    assert len(plain) < len(va.RULES) - 2000
    # 2026-09-07 未明: 畳む段が 3 つになった（重い道具・行が多い表の手・表の足回り）。
    # 重い道具の語がある依頼では、重い道具の段は畳まない（他の段は材料しだいで畳む）
    heavy = va._rules_for("ピボットで担当ごとに集計して", "")
    assert va._HEAVY_STUB not in heavy and "percent_parent_row" in heavy
    # 材料にピボットがある回も、重い道具の段は畳まない（他の段は材料しだいで畳む）
    inbook = va._rules_for("整えて", "ピボット（ブック全体）: 2 本" + chr(10) + "  [名簿] P集計 出力=A3:B7")
    assert va._HEAVY_STUB not in inbook and "percent_parent_row" in inbook
    dm = va._rules_for("整えて", "データモデル: テーブル 1 本   メジャー 0 本")
    assert va._HEAVY_STUB not in dm and "percent_parent_row" in dm
    assert va._rules_for("整えて", "", rules="独自の規則") == "独自の規則"     # 目印が無い規則文はそのまま


def test_split_unmet_for_approval_moves_value_changes_to_noticed():
    """6: 承認の言葉が無い走行では、人の値を変える指摘を plan に混ぜず気づきへ回す。"""
    import vbam_agent as va
    unmet = ["A17 の全角「１０１２」を半角に修正してください", "16 行目の重複を削除してください",
             "金額列 G に桁区切りがありません → tidy を当てる"]
    keep, moved = va._split_unmet_for_approval(unmet, approved=False)
    assert keep == [unmet[2]] and moved == unmet[:2]
    keep, moved = va._split_unmet_for_approval(unmet, approved=True)
    assert keep == unmet and moved == []


def test_mask_secrets_hides_key_values_in_prompt_text():
    """12: 鍵の値は AI に送る文から伏せる（長い順に置換・8 字未満は集めない側で弾く）。"""
    import vbam_agent as va
    secrets = {"AIzaSy-abcdefghijklmnop", "abcdefghijk"}
    out = va._mask_secrets("A13: AIzaSy-abcdefghijklmnop / B1: abcdefghijk / C1: ok", secrets)
    assert "AIzaSy" not in out and "abcdefghijk" not in out and out.count(va._SECRET_MASK) == 2 and "C1: ok" in out
    assert va._mask_secrets("そのまま", set()) == "そのまま"


def test_report_head_skips_section_headers():
    """11: 進み具合の note に「【やったこと】」の見出しだけが出ていた。最初の中身の行を取る。"""
    import vbam_agent as va
    assert va._report_head("【やったこと】\nA5:E20 に tidy\n【できなかったこと】") == "A5:E20 に tidy"
    assert va._report_head("【やったこと】") == "【やったこと】"
    assert va._report_head("") == ""


def test_fire_rates_counts_pass_and_first_shot_per_case():
    """10: --repeat N の弾ごとの (合格, 撃った数, 一発合格)。"""
    import vbam_agent as va
    rows = [("a", True, "", {"gates": {}}), ("a", False, "", {"gates": {"hand": 1}}), ("a", True, "", {"gates": {"grade": 1}}),
            ("b", True, "", {"gates": {}})]
    assert vf._fire_rates(rows) == {"a": (2, 3, 1), "b": (1, 1, 1)}


def test_id_columns_and_zenkaku_digits():
    """8: 番号列の見出しは tidy と同じ語で見分ける。全角の数字は番号列でも乱れ。"""
    import vbam_agent as va
    assert va._id_columns(["会員番号", "氏名", "コード", "No.", "金額", None]) == {0, 2, 3}
    assert va._zenkaku_digits("１０１２") and not va._zenkaku_digits("1012")


def test_save_undo_meta_writes_neighbor_only_for_a_real_backup(tmp_path, monkeypatch):
    """15: 控えのファイルが無い覚書（テストの 'p'・'x'）は隣に .json を残さない＝プロジェクト直下の残骸を作らない。"""
    import json
    import os
    import vbam_agent as va
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(vu, "_LAST_AGENT_UNDO_FILE", str(tmp_path / "undo.json"))
    assert va._save_undo_meta({"book": "b.xlsx", "sheet": "名簿", "path": "p"}) is True
    assert not os.path.exists(tmp_path / "p.json")
    real = tmp_path / "b_agent_before_x.xlsx"
    real.write_bytes(b"x")
    va._save_undo_meta({"book": "b.xlsx", "sheet": "名簿", "path": str(real)})
    assert json.loads((tmp_path / "b_agent_before_x.xlsx.json").read_text(encoding="utf-8"))["sheet"] == "名簿"


def test_verify_sheet_reuses_the_grade_materials(monkeypatch, capsys):
    """2: 採点が通った直後の終わりの検査は、材料を読み直さない（materials を渡せば _run_cmd を呼ばない）。"""
    import vbam_agent as va
    calls = []
    monkeypatch.setattr(va, "_run_cmd", lambda tokens, wb=None: (calls.append(tokens) or (True, "エラーセル: 0")))
    assert va._verify_sheet("名簿", None, (0, 0), materials="使用範囲: A1:B3\nエラーセル: 0") is True
    assert calls == []
    va._verify_sheet("名簿", None, (0, 0))
    assert calls and calls[0][0] == "materials"


def test_unbound_shape_names_and_delete_all_without_macro():
    """5: マクロ無しの図形を名前を写さずに全部消す（写し間違いで 1 往復目が止まらない）。"""
    import vbam_agent as va

    class Shp:
        def __init__(self, name, oa):
            self.Name, self.OnAction, self.deleted = name, oa, False
            self.TextFrame = type("TF", (), {"Characters": staticmethod(lambda: type("C", (), {"Text": ""})())})()

        def Delete(self):
            self.deleted = True

    items = [Shp("ロゴ枠", ""), Shp("btn1", "Macro1"), Shp("矢印", "")]

    class Shapes:
        Count = len(items)

        def __call__(self, i):
            if isinstance(i, int):
                return items[i - 1]
            for s in items:
                if s.Name == i:
                    return s
            raise Exception("no shape")

    ws = type("WS", (), {"Shapes": Shapes(), "Name": "名簿"})()
    assert va._unbound_shape_names(ws) == ["ロゴ枠", "矢印"]
    wb = type("WB", (), {"Sheets": staticmethod(lambda name: ws)})()
    label, ok, out = va._direct_action({"op": "shape", "delete": True, "names": "all_without_macro"}, "名簿", wb)
    assert ok and label == "shape delete 2" and [s.deleted for s in items] == [True, False, True]


def test_content_now_feeds_the_audit_gate_and_the_report(monkeypatch):
    """_content_now: 止める指摘は仕上げ検査の関所に、気づきは報告の【気づいたこと】に乗る。"""
    import vbam_agent as va
    import vbam_edit

    class _R:
        Address = "$A$1:$D$6"
        Row, Column = 1, 1
        Rows = type("R", (), {"Count": 6})()
        Columns = type("C", (), {"Count": 4})()
        Value = (("品名", "単価", "数量", "金額"), ("a", 1, 1, 1), ("b", 2, 2, 4), ("c", 3, 3, 9),
                 ("d", None, 4, 16), ("合計", 6, 10, 99))
        Formula = Value
        CurrentRegion = None

    _R.CurrentRegion = _R()
    ws = type("WS", (), {"Range": staticmethod(lambda a: _R()), "Name": "x"})()
    wb = type("WB", (), {"Sheets": staticmethod(lambda name: ws)})()
    block, seen = va._content_now(wb, "x", ["A1"])
    assert any("「合計」行の D6" in b for b in block)
    assert any("表の中の空欄が 1 セル（B5）" in s for s in seen)


def test_grader_sees_the_sheets_the_run_created():
    """採点係に「このループで作った別シート」の現物も渡す（2026-09-06 深夜の実射）。

    「別シートに月別集計を作って」の弾で、AI は正しく別シートに作ったのに、採点係には
    元のシートの材料しか渡していなかったため「別シートが作成されていません」と差し戻していた
    （月別・上位・推移・単価の 4 弾で再現。自己採点の差し戻し 7 件中 6 件がこれ）。
    """
    import vbam_agent as va

    class _WS:
        def __init__(self, name):
            self.Name = name

    class _WB:
        Sheets = [_WS("実射10"), _WS("推移")]

    assert va._other_sheets_now(_WB(), "実射10", ["実射10"]) == ["推移"]
    assert va._other_sheets_now(_WB(), "実射10", ["実射10", "推移"]) == []

    calls = []

    def fake_run(argv, wb):
        calls.append(list(argv))
        return True, f"使用範囲: A1:C3（{argv[1]}）"

    old = va._run_cmd
    va._run_cmd = fake_run
    try:
        mat = va._grade_materials("実射10", object(), ["推移"])
    finally:
        va._run_cmd = old
    assert "【このループで作った別のシート「推移」の現物】" in mat
    assert ["materials", "推移"] in calls
    assert "【このループで作った別のシート】の欄も現物です" in va._GRADE_PROMPT


def test_pivot_field_may_omit_the_field_name():
    """field を書き落とした group_date / top_n で 1 往復まるごと捨てない（2026-09-06 深夜）。

    手の側は '*' を渡し、道具の側が行フィールドから当てる（当てられないときは当てない）。
    """
    import vbam_agent as va
    from vbam_heavy import _guess_row_field
    label, toks = va._heavy_action_to_tokens({'op': 'pivot_field', 'action': 'group_date',
                                    'pivot': 'P推移', 'by': 'months'}, 'S')
    assert toks[:4] == ['pivot-field', 'group-date', 'P推移', '*'] and toks[-1] == 'months'
    label2, toks2 = va._heavy_action_to_tokens({'op': 'pivot_field', 'action': 'add_row',
                                      'pivot': 'P推移', 'field': '地域'}, 'S')
    assert toks2[3] == '地域'

    import datetime

    class _F:
        def __init__(self, name, orient, first=None):
            self.Name = name
            self.Orientation = orient
            self._first = first

        @property
        def DataRange(self):
            first = self._first

            class _C:
                Value = first

            class _R:
                def Cells(self, a, b):
                    return _C()
            return _R()

    class _PT:
        def __init__(self, fields):
            self._f = fields

        def PivotFields(self):
            pt = self

            class _Col:
                Count = len(pt._f)

                def Item(self, i):
                    return pt._f[i - 1]
            return _Col()

    dated = _PT([_F("日付", 1, datetime.datetime(2026, 4, 1)), _F("担当", 2, "佐藤")])
    assert _guess_row_field(dated, 'group-date') == "日付"
    assert _guess_row_field(dated, 'top-n') == "日付"
    two = _PT([_F("日付", 1, datetime.datetime(2026, 4, 1)), _F("地域", 1, "架空県北")])
    assert _guess_row_field(two, 'top-n') is None       # 2 本あるときは当てない


def test_sheet_add_is_idempotent_and_sheet_op_reads_name_aliases():
    """同じ名前の sheet add は失敗にしない／sheet_op の名前は name 以外の置き場からも読む。

    2026-09-06 深夜の実射: 「月別」の弾で sheet add が「既に存在します」で失敗し、同じ返事に並んだ
    ピボットの手が全部取り消されて往復を 1 回まるごと失った。sheet_op に name を書かずに来る往復も 2 件。
    """
    import inspect as _i
    import vbam_edit as ve
    import vbam_agent as va
    src = _i.getsource(ve.cmd_sheet)
    assert "既にあります（作らずにそのまま使います）" in src
    label, toks = va._table_action_to_tokens({'op': 'sheet_op', 'action': 'add', 'sheet': '月別'}, 'S')
    assert toks == ['sheet', 'add', '月別']
    label2, toks2 = va._table_action_to_tokens({'op': 'sheet_op', 'action': 'add', 'to': '推移'}, 'S')
    assert toks2 == ['sheet', 'add', '推移']


def test_key_type_mismatch_is_told_before_the_lookup_not_after():
    """突き合わせのキーの型違いを、書いた後の検査でなく材料の段階で言う（2026-09-06 深夜のループ）。

    実射「会員番号で突き合わせ」は 15 回中 3 回しか一発で通らなかった。練習台は左の表のキーが文字・
    右が数値で、VLOOKUP が全行「申込なし」になる。捕まえるのが書いた後だったので、毎回 1 往復よけいに要った。
    """
    import inspect as _i
    import vbam_agent as va
    from vbam_edit import key_type_mismatch
    grid = [["会員名簿", None, None, None],
            [None, None, None, None],
            ["会員番号", "氏名", None, None],
            ["2001", "青木 誠", "会員番号", "申込コース"],
            ["2002", "石川 恵", 2006, "年間"],
            ["2003", "上田 学", 2012, "半年"]]
    notes = key_type_mismatch(grid)
    assert len(notes) == 1
    assert "「会員番号」の列で型が違います" in notes[0] and "A4＝文字" in notes[0] and "C5＝数値" in notes[0]

    same = [r[:] for r in grid]
    for i in (4, 5):
        same[i][2] = str(same[i][2])
    assert key_type_mismatch(same) == []

    plain = [["支店", "売上"], ["青葉", 10], ["若葉", 20]]      # キーらしい見出しでなければ黙る
    assert key_type_mismatch(plain) == []

    src = _i.getsource(va._extra_materials) if hasattr(va, '_extra_materials') else ''
    assert "キーの型" in (src or io.open(va.__file__, encoding='utf-8').read())


def test_the_total_row_belongs_to_the_table_for_tidy():
    """規則と検査の食い違いを直した（2026-09-06 深夜のループ）。

    規則は「表全体に下の合計行は含めない」と書き、仕上げ検査は「合計行に罫線が無い」と差し戻していた。
    実射「金額の合計」が毎回そこで 1 往復を使っていた。
    """
    import vbam_agent as va
    rules = va.RULES if hasattr(va, 'RULES') else io.open(va.__file__, encoding='utf-8').read()
    assert "自分が足した合計行・合計列は含める" in rules
    assert "上のタイトル行・下の合計行・" not in rules


def test_plan_gate_skips_only_when_the_whole_list_was_written_before_the_hands(tmp_path, monkeypatch):
    """手と done を同時に受けた回で、一覧が全部「未」なら突き返さない（2026-09-07 未明のループ）。

    実射「パワークエリで読み込み」は、1 往復目に手 3 本と done を返して 3 本とも通ったのに、
    同じ返事の plan（実行前に書いたもの）が全部「未」だったため関所が done を突き返し、
    AI は状態を「済」に書き換えるだけの往復を毎回 1 回使っていた。
    ただし「済」と「未」を書き分けている回は AI 自身の申告＝落とし物なので、今までどおり突き返す
    （test_merged_done_gate_message_carries_the_unseen_results がその側を見張っている）。
    """
    replies = ['{"say":"やります","plan":[{"item":"テーブル化","state":"未"},{"item":"クエリ作成","state":"未"}],'
               '"actions":[{"op":"write_cells","cells":{"A1":"x"}}],"done":true,"report":"やった"}',
               '{"plan":[{"item":"テーブル化","state":"済"},{"item":"クエリ作成","state":"済"}],'
               '"actions":[],"done":true,"report":"やった"}']
    va, wb, sent = _loop_fakes(monkeypatch, tmp_path, replies)
    r = va.run_agent("テーブル化してクエリを作って", "名簿", wb, max_turns=4, materials="使用範囲: A1",
                     verify=False, backup=False, grade=False, show_image=False)
    assert r['turns'] == 1 and r['done'] is True      # 状態を書き換えるだけの 2 往復目を使っていない
    assert r['gates'].get('plan', 0) == 0


def test_rules_shrink_for_a_small_plain_sheet():
    """送りを削る（2026-09-07 未明）。費用の 9 割は入力で、その大半が規則文（16,333 字）。

    小さい素の表に「整えて」だけの依頼なら、重い道具・行が多い表の手・表の足回りの 3 段は要らない。
    畳んでも手の名前は見出しに残す（AI に「その手は無い」と誤答させないため）。
    """
    import vbam_agent as va
    mat = ("使用範囲: A1:E9  9行 x 5列\n条件付き書式: 0 本  入力規則: 0 セル\n"
           "テーブル（ブック全体）: 1 本\n  [月別] T名簿 範囲=A3:E9\n"
           "ピボット（ブック全体）: 1 本\n  [月別] P月別 出力=A3:B7\n"
           "スライサー: 0 個\nパワークエリ: 0 本\nデータモデル: テーブル 0 本   メジャー 0 本\n")
    thin = va._rules_for("この表を整えて", mat, sheet="実射1")
    assert len(thin) < len(va.RULES) - 5000
    for stub in (va._HEAVY_STUB, va._ROWS_STUB, va._LEGS_STUB):
        assert stub in thin
    assert '"op":"tidy"' in thin and '"op":"write_grid"' in thin     # 使う手は残っている

    # 別シートのピボットでは重い道具を送らない（実射は 15 弾を 1 冊で回すので、前は毎回送っていた）
    assert va._heavy_needed("この表を整えて", mat, sheet="実射1") is False
    assert va._heavy_needed("この表を整えて", mat, sheet="月別") is True
    assert va._heavy_needed("ピボットにして", mat, sheet="実射1") is True

    # 300 行の表では「行が多い表の手」を送る
    big = mat + "--- 列プロファイル（全 300 行・5 列） ---\n  A 会員番号: 文字 300\n"
    assert va._ROWS_STUB not in va._rules_for("この表を整えて", big, sheet="実射1")

    # 並べ替えを頼まれたら足回りを送る
    assert va._LEGS_STUB not in va._rules_for("売上の多い順に並べ替えて", mat, sheet="実射1")


def test_creating_a_macro_is_not_an_unrequested_change():
    """「作って」の回に、マクロが増えたことを咎めない（2026-09-07 未明の実射）。

    実射「新しいマクロを作る」で、頼まれたとおり Sub を 1 本足したのに
    「頼んでいないマクロが増えた」で done を突き返し、往復を 2 回よけいに使っていた。
    消えた・変わったは、作れと言われた回でも咎める（他のマクロを壊してよい理由にはならない）。
    """
    import vbam_agent as va
    before = {("M1", "既存"): "Sub 既存()\nEnd Sub"}
    added = dict(before)
    added[("M1", "新しい")] = "Sub 新しい()\nEnd Sub"
    assert vmac._code_violations(before, added) == ["頼んでいないマクロが増えた: M1.新しい"]
    assert vmac._code_violations(before, added, allow_new=True) == []

    broke = {("M1", "既存"): "Sub 既存()\n  ' 書き換えた\nEnd Sub",
             ("M1", "新しい"): "Sub 新しい()\nEnd Sub"}
    assert vmac._code_violations(before, broke, allow_new=True) == ["頼んでいないマクロが変わった: M1.既存"]
    assert vmac._code_violations(before, {("M1", "新しい"): "Sub 新しい()\nEnd Sub"},
                               allow_new=True) == ["頼んでいないマクロが消えた: M1.既存"]


def test_agent_claude_code_head(tmp_path, monkeypatch):
    # 3 本目の頭（2026-09-07）: claude -p を stream-json で温める。鍵なし・1 会話 1 プロセス・柵を剥ぐ・画像も送れる
    import base64
    import io
    import json
    import queue
    import pytest
    import vbam_agent as va
    made = []
    killed = []

    class FakeProc:
        """claude -p の偽物: stdin に来た user 行ごとに result 行を返す。stdin を閉じたら EOF"""
        def __init__(self, cmd, **kw):
            self.cmd, self.kw, self.pid = cmd, kw, 4242
            self.q = queue.Queue()
            self.sent = []
            self.closed = False
            proc = self

            class In:
                def write(_s, b):
                    proc.sent.append(b)
                    n = len(proc.sent)
                    proc.q.put(json.dumps({"type": "system", "subtype": "init"}).encode())
                    proc.q.put(json.dumps({"type": "result", "subtype": "success", "is_error": False,
                                           "result": "```json\n{\"actions\":[],\"n\":%d}\n```" % n,
                                           "usage": {"input_tokens": 100, "cache_read_input_tokens": 20,
                                                     "cache_creation_input_tokens": 5, "output_tokens": 7},
                                           "total_cost_usd": 0.001}).encode())

                def flush(_s):
                    pass

                def close(_s):
                    proc.closed = True
                    proc.q.put(b'')

            class Out:
                def readline(_s):
                    return proc.q.get()

            self.stdin, self.stdout, self.stderr = In(), Out(), io.BytesIO(b'')
            made.append(self)

        def poll(self):
            return 0 if self.closed else None

        def wait(self, timeout=None):
            if not self.closed:
                raise RuntimeError("not closed")
            return 0

    monkeypatch.setattr(va.subprocess, "Popen", FakeProc)
    monkeypatch.setattr(va.subprocess, "run", lambda *a, **k: killed.append(a))
    monkeypatch.setattr(va.shutil, "which", lambda name: r"C:\fake\claude.CMD")
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    monkeypatch.setattr(va, "_key_load", lambda ai: None)
    va._cc_close_all()
    # 既定は claude-code（鍵なし・sonnet）。--model が gemini で始まれば gemini に倒す（鍵が無ければ止まる＝それが正しい）
    assert va._ai_setup(None, None) == ('claude-code', 'sonnet', '')
    assert va._ai_setup('claude-code', 'haiku') == ('claude-code', 'haiku', '')
    with pytest.raises(RuntimeError):
        va._ai_setup(None, 'gemini-flash-lite-latest')
    monkeypatch.setenv("GEMINI_API_KEY", "k")
    assert va._ai_setup(None, 'gemini-flash-lite-latest')[0] == 'gemini'
    # 1 通目: 起動して送る。行は ASCII だけ・道具なし・柵は剥がれる・トークンは入力＋キャッシュを畳む
    hist = [('user', "1 通目 日本語", None)]
    text, usage = va._chat_claude_code(hist, 'sonnet', '')
    assert text == '{"actions":[],"n":1}' and usage == {'in': 125, 'out': 7, 'think': 0, 'total': 132}
    p = made[0]
    assert p.cmd[0].endswith("claude.CMD") and p.cmd[1] == '-p' and p.cmd[p.cmd.index('--tools') + 1] == ''
    assert p.cmd[p.cmd.index('--model') + 1] == 'sonnet' and '--strict-mcp-config' in p.cmd
    assert p.sent[0].endswith(b"\n") and all(b < 128 for b in p.sent[0]) and b"\\u65e5" in p.sent[0]
    assert json.loads(p.sent[0]) == {"type": "user", "message": {"role": "user", "content": "1 通目 日本語"}}
    # 2 通目: 同じ履歴の続きなら同じプロセス。画像は content の 2 つ目に base64 で付く
    png = tmp_path / "view.png"
    png.write_bytes(b"\x89PNG\r\n\x1a\n" + b"x" * 20)
    hist += [('model', text), ('user', "2 通目", str(png))]
    text2, _u = va._chat_claude_code(hist, 'sonnet', '')
    assert text2 == '{"actions":[],"n":2}' and len(made) == 1
    c = json.loads(p.sent[1])["message"]["content"]
    assert c[0] == {"type": "text", "text": "2 通目"}
    assert c[1]["source"] == {"type": "base64", "media_type": "image/png",
                              "data": base64.b64encode(png.read_bytes()).decode('ascii')}
    # 別の会話（採点）は新しいプロセス。前のは閉じる（居残らせない）。閉じるのは stdin を閉じるだけ＝taskkill は出ない
    text3, _u = va._chat_claude_code([('user', "採点して")], 'sonnet', '')
    assert len(made) == 2 and p.closed and text3 == '{"actions":[],"n":1}' and killed == []
    va._cc_close_all()
    assert made[1].closed and va._cc_convs == []
    # 予備を先に起こす（2026-09-11 午後）: 本体と採点係の 2 本。1 通目は予備を使う＝新しい Popen は出ない。
    # 続きは長く続いている会話へ（空の予備に流れない）。採点はもう 1 本の予備へ。終わりに使った会話だけ閉じる
    assert va.cc_prewarm(['sonnet', 'sonnet']) == 2 and len(made) == 4
    assert va.cc_prewarm(['sonnet']) == 0                     # 空いていれば起こさない
    hist = [('user', "本体 1 通目", None)]
    t1, _u = va._chat_claude_code(hist, 'sonnet', '')
    hist += [('model', t1), ('user', "本体 2 通目", None)]
    t2, _u = va._chat_claude_code(hist, 'sonnet', '')
    assert len(made) == 4 and json.loads(t2)["n"] == 2 and len(made[2].sent) == 2 and made[3].sent == []
    tg, _u = va._chat_claude_code([('user', "採点して")], 'sonnet', '')
    assert len(made) == 4 and len(made[3].sent) == 1 and not made[2].closed      # 本体は生きたまま
    va.cc_close_used()
    assert made[2].closed and made[3].closed and va._cc_convs == []
    va.cc_prewarm(['sonnet'])
    va.cc_close_used()                                        # 空の予備は残す
    assert len(va._cc_convs) == 1 and not made[4].closed
    va._cc_close_all()
    # _ask の振り分け・--ai と --grade-ai の選択肢・柵剥ぎ
    assert va._CHAT_FN['claude-code'] is va._chat_claude_code
    ns, unknown = vm.build_parser().parse_known_args(["agent", "x", "--ai", "claude-code", "--grade-ai", "claude-code"])
    assert ns.ai == "claude-code" and ns.grade_ai == "claude-code" and not unknown
    assert va._cc_strip_fence("```json\n{\"a\":1}\n```") == '{"a":1}' and va._cc_strip_fence(' {"a":1} ') == '{"a":1}'


def test_inner_cmd_calls_are_pinned():
    """cmd_* を内側から呼ぶ 3 か所が pinned_workbook の中にあること（回帰）。"""
    import vbam_build as vb
    import vbam_agent as va

    src_apply = inspect.getsource(vb._apply)
    assert "with pinned_workbook(" in src_apply and "cmd_tidy(ns)" in src_apply
    assert src_apply.index("with pinned_workbook(") < src_apply.index("cmd_tidy(ns)")

    src_verify = inspect.getsource(vb._verify)
    assert "with pinned_workbook(" in src_verify
    assert src_verify.index("with pinned_workbook(") < src_verify.index("cmd_screenshot(ns)")

    src_clean = inspect.getsource(vc.cmd_clean_table)
    assert "with pinned_workbook(" in src_clean
    assert src_clean.index("with pinned_workbook(") < src_clean.index("cmd_tidy(ns)")


# ================================================================
# plan の関所（2026-09-08）
# 手と done を同じ返事で受け、手が全部通った回は、一覧の「未」で突き返さない。
# 「全部未」の回だけ見逃していたので、一部を「済」と書き分けた回＝正直に書いた回だけが
# 突き返され、状態を書き換えるだけの往復を 1 回失っていた（本番で 12.7 秒）。
# ================================================================

def test_plan_mark_done_only_touches_pending():
    import vbam_agent as va

    plan = [{'item': 'A', 'state': '済'}, {'item': 'B', 'state': '未'},
            {'item': 'C', 'state': '未'}, {'item': 'D', 'state': '要判断'},
            {'item': 'E', 'state': '不可'}]
    assert va._plan_pending(plan) == ['B', 'C']
    va._plan_mark_done(plan, ['B'])                    # 名指しした分だけ
    assert [p['state'] for p in plan] == ['済', '済', '未', '要判断', '不可']
    va._plan_mark_done(plan, va._plan_pending(plan))
    assert va._plan_pending(plan) == []
    assert [p['state'] for p in plan] == ['済', '済', '済', '要判断', '不可']   # 要判断・不可は動かさない


def test_plan_gate_exemption_is_not_all_or_nothing():
    """関所の免除が「一覧が全部未のとき」に限られていないこと（回帰）。"""
    src = inspect.getsource(vbam_agent_run_agent_source())
    i = src.index("pending = _plan_pending(plan)")
    head = src[i:i + 400]
    assert "if pending and merged_now:" in head, head[:200]
    assert "len(pending) == len(plan" not in head, "「全部未」限定の免除が残っている"


def vbam_agent_run_agent_source():
    import vbam_agent as va
    return va.run_agent


def test_merged_flag_does_not_leak_into_the_next_turn(tmp_path, monkeypatch):
    """1: 手と done を同時に受けた次の往復で「手ゼロの done」が来ても、免除を持ち越さない。

    2026-09-08: merged_now を往復ごとに戻していなかったので、1 度立つと以後ずっと立ちっぱなしになり、
    手を 1 本も出していない done が「手は全部通った回」として一覧の「未」を消していた。
    """
    replies = ['{"say":"書きます","plan":[{"item":"A1 に書く","state":"済"}],'
               '"actions":[{"op":"write_cells","cells":{"A1":"x"}}],"done":true,"report":"やった"}',
               '{"say":"消せません","plan":[{"item":"A1 に書く","state":"済"},'
               '{"item":"ボタンを消す","state":"未"}],"actions":[],"done":true,"report":"無理"}',
               '{"plan":[{"item":"A1 に書く","state":"済"},{"item":"ボタンを消す","state":"不可"}],'
               '"actions":[],"done":true,"report":"消せないので不可"}']
    va, wb, sent = _loop_fakes(monkeypatch, tmp_path, replies)
    notes = [["A1:B2 に罫線がありません → tidy を当てる"]]      # 1 回目だけ差し戻して 2 往復目を作る
    monkeypatch.setattr(va, '_audit_now', lambda *a, **k: notes.pop(0) if notes else [])
    monkeypatch.setattr(va, '_content_now', lambda *a, **k: ([], []))
    r = va.run_agent("A1 に x と書いてボタンを消して", "名簿", wb, max_turns=4, materials="使用範囲: A1",
                     verify=False, backup=False, grade=False, show_image=False)
    # 2 往復目の「未」は突き返され（手ゼロ）、3 往復目で「不可」になって初めて done
    assert r['turns'] == 3 and r['done'] is True
    assert [p['state'] for p in r['plan']] == ['済', '不可']


# ================================================================
# 報告の番地照合（2026-09-08）
# 「報告に出た番地のうち、いま空のもの」が拾いすぎていた。
#   ・壊れた式の原因として挙げた E100（=G27/E100 の参照先）を毎回「空です」と言う
#   ・【確認していないこと】の「使用範囲が H31 まで膨張」の H31 を「空です」と言う
# ================================================================

def test_report_addrs_ignores_formula_operands_and_non_did_sections():
    import vbam_agent as va

    report = (
        "【やったこと】\n"
        "・D31: 元の式 =G27/E100 は E100 が空欄のため #DIV/0!。=G27/SUM(E6:E25) に上書き修正。\n"
        "・F5 に「申込コース」と書いた。\n"
        "【できなかったこと】\nなし\n"
        "【確認していないこと】\n"
        "・H列（使用範囲が H31 まで膨張しているが、依頼範囲外のため触らず）\n")
    got = va._report_addrs(report)
    assert "D31" in got and "F5" in got            # 書いたと言っている番地は拾う
    assert "E100" not in got                       # 数式の中の参照先は拾わない
    assert "G27" not in got and "E6" not in got
    assert "H31" not in got                        # 【確認していないこと】は見ない


def test_report_addrs_falls_back_to_whole_text_without_sections():
    import vbam_agent as va

    assert va._report_addrs("A1 と B2 に書きました") == ["A1", "B2"]


def test_mark_undo_stale_only_marks_the_same_book(tmp_path, monkeypatch):
    """控えの「⚠ この後に build が走った」を、別のブックの仕事で付けない（2026-09-08）。"""
    import vbam_agent as va

    meta_file = tmp_path / "undo.json"
    monkeypatch.setattr(vu, '_LAST_AGENT_UNDO_FILE', str(meta_file))
    saved = {}
    monkeypatch.setattr(vu, '_save_undo_meta', lambda m: saved.update(m))
    meta_file.write_text(json.dumps({'book': '売上.xlsx', 'path': 'x'}), encoding='utf-8')

    va._mark_undo_stale('build（別のブック）', '名簿.xlsx')      # 別のブック → 付けない
    assert saved == {}
    va._mark_undo_stale('build（同じブック）', '売上.xlsx')      # 同じブック → 付ける
    assert saved.get('stale') == 'build（同じブック）'
    saved.clear()
    va._mark_undo_stale('book を渡さない古い呼び方')             # 渡さなければ今までどおり
    assert saved.get('stale') == 'book を渡さない古い呼び方'


# ----------------------------------------------------------------
# normalize の手（_do_normalize）を偽のシートで（2026-09-10 夜）。
# 規則そのもの（_normalize_value／_normalize_column）にはテストがあったが、COM に書く側は 1% しか通っていなかった。
# 道具が全行を書き換える手なので、書き先・数式・先頭ゼロ・文字列書式の扱いを間違えると表を丸ごと壊す。
# ----------------------------------------------------------------
_NZ_GENERAL = 'G/標準'


def _nz_col(n):
    s = ''
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _nz_rc(a):
    m = _re.match(r'^\$?([A-Z]+)\$?(\d+)$', a)
    col = 0
    for ch in m.group(1):
        col = col * 26 + ord(ch) - 64
    return int(m.group(2)), col


class _NzAxis:
    """範囲の Rows／Columns: .Count と、(i) で i 行目／i 列目の範囲（1 起点）。"""

    def __init__(self, box, rows):
        self.box, self.rows = box, rows
        self.Count = (box.r2 - box.r1 + 1) if rows else (box.c2 - box.c1 + 1)

    def __call__(self, i):
        b = self.box
        if self.rows:
            return _NzBox(b.ws, b.r1 + i - 1, b.c1, b.r1 + i - 1, b.c2)
        return _NzBox(b.ws, b.r1, b.c1 + i - 1, b.r2, b.c1 + i - 1)


class _NzBox:
    """偽のセル範囲。値・数式・書式は持ち主のシート（_NzSheet）の辞書に置く。"""

    def __init__(self, ws, r1, c1, r2, c2):
        self.ws, self.r1, self.c1, self.r2, self.c2 = ws, r1, c1, r2, c2
        self.Row, self.Column, self.Worksheet = r1, c1, ws
        self.Rows, self.Columns = _NzAxis(self, True), _NzAxis(self, False)

    def _keys(self):
        return [[(r, c) for c in range(self.c1, self.c2 + 1)] for r in range(self.r1, self.r2 + 1)]

    def _grid(self, pick):
        g = tuple(tuple(pick(k) for k in row) for row in self._keys())
        return g[0][0] if len(g) == 1 and len(g[0]) == 1 else g      # COM と同じく 1 セルは素の値

    @property
    def Address(self):
        a = f"${_nz_col(self.c1)}${self.r1}"
        return a if (self.r1, self.c1) == (self.r2, self.c2) else a + f":${_nz_col(self.c2)}${self.r2}"

    @property
    def Value(self):
        return self._grid(lambda k: self.ws.v.get(k))

    @Value.setter
    def Value(self, new):
        keys = self._keys()
        if not isinstance(new, (list, tuple)):
            new = [[new] * len(keys[0]) for _ in keys]
        for row, vals in zip(keys, new):
            for k, x in zip(row, vals):
                self.ws.put(k, x)

    @property
    def Formula(self):
        return self._grid(lambda k: self.ws.f.get(k, '' if self.ws.v.get(k) is None else str(self.ws.v.get(k))))

    @property
    def HasFormula(self):
        n = sum(1 for row in self._keys() for k in row if k in self.ws.f)
        return True if n == self.Rows.Count * self.Columns.Count else (False if n == 0 else None)

    @property
    def NumberFormat(self):
        fmts = {self.ws.fmt.get(k, _NZ_GENERAL) for row in self._keys() for k in row}
        return fmts.pop() if len(fmts) == 1 else None

    @NumberFormat.setter
    def NumberFormat(self, fmt):
        for row in self._keys():
            for k in row:
                self.ws.fmt[k] = fmt

    def ClearContents(self):
        for row in self._keys():
            for k in row:
                self.ws.v.pop(k, None)
                self.ws.f.pop(k, None)
                self.ws.writes.append(k)


class _NzSheet:
    """Excel の読み替えを真似る偽のシート: 標準の書式に数字の文字を書くと数値、yyyy/m/d の文字は日付になる。
    文字列書式（@）のセルには何を書いても文字で入る（数値・日付も文字になる＝CSV 取り込みの列の罠）。"""

    def __init__(self, cells=None, formulas=None, fmts=None):
        self.v, self.f, self.writes = {}, {}, []
        for a, x in (cells or {}).items():
            self.v[_nz_rc(a)] = x
        for a, (form, shown) in (formulas or {}).items():
            self.f[_nz_rc(a)], self.v[_nz_rc(a)] = form, shown
        self.fmt = {_nz_rc(a): x for a, x in (fmts or {}).items()}
        self.Rows, self.Columns = _NzLines(self, 1048576), _NzLines(self, 16384)
        # 見た目・フィルの試験に使うつまみ（既定は「何もしていないシート」）
        self.error_cells, self.f2, self.groups = set(), set(), set()
        self.col_widths, self.row_height, self.borders = {}, 15.0, {}
        self.bold_rows, self.filled_rows, self.selected = set(), set(), None
        outer = self
        count_a = lambda rng: sum(1 for row in rng._keys() for k in row if outer.v.get(k) not in (None, ''))  # noqa: E731
        self.Parent = _InvObj(Application=_InvObj(WorksheetFunction=_InvObj(CountA=count_a)))

    def put(self, k, x):
        self.writes.append(k)
        self.f.pop(k, None)
        if self.fmt.get(k, _NZ_GENERAL) == '@':
            x = None if x in (None, '') else (x if isinstance(x, str) else str(x))
        elif isinstance(x, str):
            m = _re.match(r'^(\d{4})/(\d{1,2})/(\d{1,2})$', x)
            if x == '':
                x = None
            elif x.startswith('='):                   # 数式として入る（値は計算しない＝字面のまま持つ）
                self.f[k] = x
                if k in self.error_cells:
                    x = -2146826281                   # #DIV/0!
            elif _re.match(r'^[-+]?\d+(\.\d+)?$', x.replace(',', '')):
                x = float(x.replace(',', ''))
            elif m:
                x = _dt.datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        self.v[k] = x

    def Range(self, a, b=None):
        if b is not None:
            return _NzBox(self, a.r1, a.c1, b.r2, b.c2)
        parts = str(a).split(':')
        r1, c1 = _nz_rc(parts[0])
        r2, c2 = _nz_rc(parts[-1])
        return _NzBox(self, r1, c1, r2, c2)

    def Cells(self, r, c):
        return _NzBox(self, r, c, r, c)

    def val(self, a):
        return self.v.get(_nz_rc(a))

    def fmt_of(self, a):
        return self.fmt.get(_nz_rc(a), _NZ_GENERAL)


import datetime as _dt                                              # noqa: E402


def test_normalize_to_writes_a_copy_and_keeps_ids_and_formulas_safe():
    """to（右の列へ写す）: 元の列は触らない・数式は写さない・Excel が数値や日付に化かした文字は文字に戻す。"""
    import vbam_agent as va
    ws = _NzSheet({"C1": "電話", "C2": " 0001 ", "C3": "１２,０００", "C6": "abc ", "C7": "2026/01/05"},
                  formulas={"C4": ("=C3*2", 24000)})
    before = dict(ws.v), dict(ws.f)
    act = {"op": "normalize", "range": "C2:C7", "rules": ["trim", "hankaku", "number"],
           "to": "H2", "header": "電話（整形）"}
    label, ok, out = va._do_normalize(ws, act)
    assert ok is True and label == "normalize C2:C7 → H2:H7"
    # 先頭ゼロの番号と日付に見える文字は、Excel が 1・日付に読み替えても文字で入る（@ にして書き直す）
    assert ws.val("H2") == "0001" and ws.fmt_of("H2") == "@"
    assert ws.val("H7") == "2026/01/05" and ws.fmt_of("H7") == "@"
    assert ws.val("H3") == 12000 and ws.fmt_of("H3") == _NZ_GENERAL          # 全角の金額は数値になる
    assert ws.val("H4") is None and (4, 8) not in ws.f                      # 数式は書き先へ動かさない
    assert ws.val("H5") is None and ws.val("H6") == "abc"
    assert ws.val("H1") == "電話（整形）"
    assert (dict(ws.v), dict(ws.f)) != before and all(ws.v[k] == before[0][k] for k in before[0])
    assert ws.f == before[1]                                                 # 元の列（数式も）はそのまま
    assert "書き先: H2:H7（6 行 × 1 列を写した。変えた値 3）・文字のまま守ったセル 2" in out
    assert "前後の空白 2" in out and "全角→半角 1" in out and "数値化 1" in out
    assert "見出し: H1「電話（整形）」" in out and "結果の列の様子:" in out
    assert "例: C2 ' 0001 '→'0001'" in out

    # 見出しのセルが埋まっていれば書かない（人の見出しを潰さない）
    ws2 = _NzSheet({"C2": "a ", "H1": "人の見出し"})
    _, _, out2 = va._do_normalize(ws2, {"op": "normalize", "range": "C2:C2", "rules": ["trim"],
                                        "to": "H2", "header": "新しい見出し"})
    assert ws2.val("H1") == "人の見出し" and "既に「人の見出し」があるので書かなかった" in out2


def test_normalize_to_refuses_a_filled_destination_without_approval():
    """書き先に値があれば、承認（overwrite）が無い限り 1 セルも書かずに断る。承認があれば上書きする。"""
    import vbam_agent as va
    import pytest
    ws = _NzSheet({"C2": " a", "C3": " b", "H3": "人の値"})
    act = {"op": "normalize", "range": "C2:C3", "rules": ["trim"], "to": "H2"}
    with pytest.raises(ValueError, match=r'書き先 H2:H3 に値が 1 セルある.*"overwrite": true'):
        va._do_normalize(ws, act)
    assert ws.writes == [] and ws.val("H3") == "人の値"
    va._do_normalize(ws, dict(act, overwrite=True))
    assert (ws.val("H2"), ws.val("H3")) == ("a", "b")


def test_normalize_to_a_text_formatted_column_writes_real_dates():
    """書き先が文字列書式（@）でも、日付・数値は本物の日付・数値で入る（2026-09-04 の _untext_grid）。"""
    import vbam_agent as va
    ws = _NzSheet({"E2": "2026/01/05", "E3": "2026年2月10日", "E4": "１２,０００"},
                  fmts={"G2": "@", "G3": "@", "G4": "@"})
    va._do_normalize(ws, {"op": "normalize", "range": "E2:E4", "rules": ["date", "number"], "to": "G2"})
    g2, g3 = ws.val("G2"), ws.val("G3")
    assert not isinstance(g2, str) and (g2.year, g2.month, g2.day) == (2026, 1, 5)
    assert not isinstance(g3, str) and (g3.month, g3.day) == (2, 10)
    assert ws.fmt_of("G2") == "yyyy/m/d" and ws.fmt_of("G3") == "yyyy/m/d"
    assert ws.val("G4") == 12000 and ws.fmt_of("G4") == _NZ_GENERAL


def test_normalize_overwrite_touches_only_the_changed_cells():
    """元の列に上書き: 変わったセルだけ書く・数式のセルは書かない・先頭ゼロは文字で残す・空になったら消す。"""
    import vbam_agent as va
    ws = _NzSheet({"B2": "総務", "B4": "　", "B6": " 営業 "}, formulas={"B5": ("=B2", "総務")})
    label, ok, out = va._do_normalize(ws, {"op": "normalize", "range": "B2:B6",
                                           "rules": ["trim", "fill_down"], "overwrite": True})
    assert ok is True and label == "normalize B2:B6（元の列に上書き）"
    assert [ws.val(f"B{r}") for r in range(2, 7)] == ["総務", "総務", "総務", "総務", "営業"]
    assert ws.f == {(5, 2): "=B2"}                                         # 数式は残る
    assert set(ws.writes) == {(3, 2), (4, 2), (6, 2)}                       # 変わった 3 セルだけ
    assert "元の列に書いた: 3 セルを変えた（5 行 × 1 列。数式のセルは触っていない）" in out
    assert "前後の空白 2" in out and "上の値で埋める 2" in out
    # 書き戻し（_cell_write_back）: 空になったら消す・先頭ゼロは @ で守る・@ の列の数値は数値に戻す
    ws2 = _NzSheet({"D2": "　", "D3": "0001 ", "D4": "12,000"}, fmts={"D4": "@"})
    va._do_normalize(ws2, {"op": "normalize", "range": "D2:D4", "rules": ["trim", "number"], "overwrite": True})
    assert ws2.val("D2") is None
    assert ws2.val("D3") == "0001" and ws2.fmt_of("D3") == "@"
    assert ws2.val("D4") == 12000 and ws2.fmt_of("D4") == _NZ_GENERAL


def test_normalize_is_capped_and_routed_through_direct_action():
    """20 万セルを超える範囲は読む前に断る。_direct_action は "sheet" を書けばそのシートの normalize に回す。"""
    import vbam_agent as va
    import pytest
    ws = _NzSheet()
    with pytest.raises(ValueError, match="20 万セルまで"):
        va._do_normalize(ws, {"op": "normalize", "range": "A1:B100001", "rules": ["trim"], "overwrite": True})
    assert ws.writes == []
    other = _NzSheet({"A2": " x"})
    asked = []
    wb = _InvObj(Sheets=lambda name: asked.append(name) or other)
    label, ok, _ = va._direct_action({"op": "normalize", "range": "A2:A2", "rules": ["trim"],
                                      "overwrite": True, "sheet": "集計"}, "名簿", wb)
    assert asked == ["集計"] and ok and other.val("A2") == "x"
    with pytest.raises(ValueError, match="ブックが無い"):
        va._direct_action({"op": "normalize"}, "名簿", None)


# ----------------------------------------------------------------
# undo_agent（戻す手）を偽の Excel で（2026-09-10 夜）。
# 安全網そのもの。これまでは「控えが古い」「番号が無い」で止まる入口と、ソースの文字列検査しか無かった。
# 偽の Excel は、ブックをまたぐ貼り付けで別シート参照が控えへの外部リンクに化ける実機の挙動も真似る。
# ----------------------------------------------------------------
def _nz_formula_set(self, new):
    keys = self._keys()
    if not isinstance(new, (list, tuple)):
        new = [[new]]
    for row, vals in zip(keys, new):
        for k, x in zip(row, vals):
            self.ws.f.pop(k, None)
            if isinstance(x, str) and x.startswith('='):
                self.ws.f[k] = x
            self.ws.v[k] = None if x in (None, '') else x


def _nz_paste_special(self, kind):
    xl = self.ws.xl
    if kind == 8:                                          # xlPasteColumnWidths
        xl.widths_pasted.append(self.Address)
        return
    src, book = xl.clip, xl.clip.ws.book
    for i, row in enumerate(src._keys()):
        for j, k in enumerate(row):
            dst = (self.r1 + i, self.c1 + j)
            self.ws.f.pop(dst, None)
            self.ws.v[dst] = src.ws.v.get(k)
            f = src.ws.f.get(k)
            if f is not None:                              # 別シート参照 → 控えのブックへの外部リンク
                self.ws.f[dst] = _re.sub(r"(?<![\]'\w])([^\W\d]\w*)!",
                                         lambda m: f"'{book.Path}\\[{book.Name}]{m.group(1)}'!", f)


def _nz_copy(self):
    if self.ws.xl.fail_copy:
        raise RuntimeError("コピーできません（偽の失敗）")
    self.ws.xl.clip = self


_NzBox.Formula = _NzBox.Formula.setter(_nz_formula_set)
_NzBox.Cells = property(lambda self: _InvObj(CountLarge=self.Rows.Count * self.Columns.Count))
_NzBox.Clear = _NzBox.ClearContents
_NzBox.Copy = _nz_copy
_NzBox.PasteSpecial = _nz_paste_special
_NzBox.Select = lambda self: None


class _UdShape:
    def __init__(self, sheet, name, left, top, width, height):
        self.sheet, self.Name, self.Left, self.Top, self.Width, self.Height = sheet, name, left, top, width, height

    def Copy(self):
        self.sheet.xl.clip_shape = self


class _UdShapes(list):
    @property
    def Count(self):
        return len(self)

    def __call__(self, key):
        if isinstance(key, int):
            return self[key - 1]
        for s in self:
            if s.Name == key:
                return s
        raise RuntimeError(f"図形 {key} がありません")


class _UdSheet(_NzSheet):
    def __init__(self, xl, name, cells=None, formulas=None, shapes=()):
        super().__init__(cells, formulas)
        self.xl, self.Name, self.book, self.activated = xl, name, None, 0
        self.Shapes = _UdShapes(_UdShape(self, *s) for s in shapes)

    @property
    def UsedRange(self):
        keys = [k for k in set(self.v) | set(self.f) if self.v.get(k) is not None or k in self.f]
        if not keys:
            return _NzBox(self, 1, 1, 1, 1)
        rs, cs = [k[0] for k in keys], [k[1] for k in keys]
        return _NzBox(self, min(rs), min(cs), max(rs), max(cs))

    def Paste(self):                                       # 図形は最前面（最後）に来る
        s = self.xl.clip_shape
        self.Shapes.append(_UdShape(self, "貼った図", 0, 0, s.Width, s.Height))

    def Activate(self):
        self.activated += 1

    def Delete(self):
        self.book._sheets.remove(self)


class _UdSheets:
    def __init__(self, book):
        self._book = book

    def __call__(self, name):
        for s in self._book._sheets:
            if s.Name == name:
                return s
        raise RuntimeError(f"シート {name} がありません")

    def __iter__(self):
        return iter(list(self._book._sheets))

    @property
    def Count(self):
        return len(self._book._sheets)


class _UdBook:
    def __init__(self, name, path, sheets):
        self.Name, self.Path, self.FullName = name, path, os.path.join(path, name)
        self._sheets = list(sheets)
        for s in self._sheets:
            s.book = self
        self.Sheets = self.Worksheets = _UdSheets(self)
        self.closed, self.names_deleted, self.links_changed = [], [], []

    def Names(self, name):
        return _InvObj(Delete=lambda: self.names_deleted.append(name))

    def LinkSources(self, kind):
        found = set()
        for s in self._sheets:
            for f in s.f.values():
                for d, n in _re.findall(r"'([^'\[]*)\\\[([^\]]+)\]", f):
                    found.add(os.path.join(d, n))
        return tuple(sorted(found)) or None

    def ChangeLink(self, old, new, kind):                  # 自分自身へ差し替えると Excel は内部参照に畳む
        self.links_changed.append((old, new, kind))
        d, n = os.path.split(old)
        for s in self._sheets:
            for k, f in list(s.f.items()):
                s.f[k] = _re.sub("'" + _re.escape(f"{d}\\[{n}]") + r"([^']+)'!", r"\1!", f)

    def Close(self, SaveChanges=None):
        self.closed.append(SaveChanges)


class _UdXl:
    def __init__(self):
        self.DisplayAlerts, self.EnableEvents, self.CutCopyMode = True, True, False
        self.clip = self.clip_shape = None
        self.fail_copy = False
        self.widths_pasted, self.opened, self.books = [], [], {}
        outer = self

        class _Workbooks:
            def Open(self, path, UpdateLinks=None, ReadOnly=None):
                outer.opened.append((path, UpdateLinks, ReadOnly, outer.EnableEvents))
                return outer.books[os.path.normcase(path)]
        self.Workbooks = _Workbooks()


def _ud_world(tmp_path, monkeypatch, meta_over=None, src_sheets=('名簿', '集計', '明細')):
    """agent が書き換えた後の「台帳.xlsx」と、書き換える前の控え（別ファイル）と、その覚書を 1 組作る。"""
    import vbam_agent as va
    monkeypatch.setattr(vu, "_LAST_AGENT_UNDO_FILE", str(tmp_path / "_last_agent_undo.json"))
    xl = _UdXl()
    bk = tmp_path / f"台帳{va._AGENT_BACKUP_MARK}20260910_120000.xlsx"
    bk.write_text("x", encoding="utf-8")
    now = [_UdSheet(xl, "名簿", {"A1": "名前", "B1": "点", "A2": "佐藤（書き換え）", "B2": 999, "D9": "AI が足した"}),
           _UdSheet(xl, "集計", {"A1": "合計", "A2": 12345}),
           _UdSheet(xl, "新シート", {"A1": "AI が作った"}),
           _UdSheet(xl, "明細", {"B2": 100})]
    wb = _UdBook("台帳.xlsx", "C:\\work", now)
    before = {"名簿": _UdSheet(xl, "名簿", {"A1": "名前", "B1": "点", "A2": "佐藤"}, {"B2": ("=明細!B2", 100)},
                               shapes=[("ロゴ", 10.0, 20.0, 30.0, 40.0)]),
              "集計": _UdSheet(xl, "集計", {"A1": "合計"}, {"A2": ("=SUM(名簿!B2:B3)", 100)}),
              "明細": _UdSheet(xl, "明細", {"B2": 100})}
    src = _UdBook(bk.name, str(tmp_path), [before[n] for n in src_sheets])
    xl.books[os.path.normcase(str(bk))] = src
    meta = {'book': "台帳.xlsx", 'sheet': "名簿", 'sheets': ["名簿", "集計"], 'path': str(bk),
            'request': "点を直して", 'time': "2026-09-10 12:00:00", 'run_id': "r1",
            'shapes_by_sheet': {"名簿": [{'name': "ロゴ", 'left': 10.0, 'top': 20.0, 'width': 30.0, 'height': 40.0}]},
            'created': [{'kind': 'sheet', 'name': "新シート"}, {'kind': 'name', 'name': "単価"},
                        {'kind': 'sheet', 'name': "集計"}]}
    meta.update(meta_over or {})
    va._save_undo_meta(meta)
    monkeypatch.setattr(vu, "get_workbook", lambda *a, **k: (xl, wb))
    return va, xl, wb, src, str(bk)


def test_undo_puts_every_written_sheet_back_and_cleans_up(tmp_path, monkeypatch, capsys):
    """戻す: 書いたシートを全部控えの姿に・作った物は消す（戻す先のシートは消さない）・消えた図形は貼り戻す。"""
    va, xl, wb, src, bk = _ud_world(tmp_path, monkeypatch)
    va.runs_append({'time': "2026-09-10 12:00:00", 'run_id': "r1", 'mode': 'sheet', 'book': "台帳.xlsx",
                    'request': "点を直して"})
    assert va.undo_agent() is True
    out = capsys.readouterr().out
    meibo, shukei = wb.Sheets("名簿"), wb.Sheets("集計")
    assert [meibo.val(a) for a in ("A1", "B1", "A2")] == ["名前", "点", "佐藤"]
    assert meibo.val("D9") is None                                           # agent が足したセルも消える
    # 貼り付けで ='C:\…\[控え]明細'!B2 に化けた式を、控えの字面で入れ直している（外部リンクを残さない）
    assert meibo.f == {(2, 2): "=明細!B2"} and shukei.f == {(2, 1): "=SUM(名簿!B2:B3)"}
    assert wb.LinkSources(1) is None and wb.links_changed == []
    assert [s.Name for s in wb.Sheets] == ["名簿", "集計", "明細"]              # 作ったシートだけ消えた
    assert wb.names_deleted == ["単価"]
    logo = meibo.Shapes("ロゴ")
    assert (logo.Left, logo.Top, logo.Width, logo.Height) == (10.0, 20.0, 30.0, 40.0)
    # 控えは読み取り専用・リンク更新なし・イベントを止めて開き、保存せずに閉じる。止めたものは元に戻す
    assert xl.opened == [(bk, 0, True, False)] and src.closed == [False]
    assert xl.EnableEvents is True and xl.DisplayAlerts is True and xl.CutCopyMode is False
    assert xl.widths_pasted == ["$A$1:$B$2", "$A$1:$A$2"]
    assert "戻しました: 名簿!A1:B2・集計!A1:A2 を控えの状態に置き換え（書き換わっていたセル 4 個）" in out
    assert "作った物を消しました: 2 個（name 単価 / sheet 新シート）" in out and "消えていた図形を控えから戻しました: 1 個" in out
    # 走行台帳のその走行に「人が戻した」の印が付く
    assert va._runs_load()[0].get('undone')


def test_undo_of_a_big_sheet_still_restores_and_drops_the_backup_link(tmp_path, monkeypatch, capsys):
    """控えが大きすぎて写せないシート（_SNAPSHOT_MAX_CELLS 超）でも戻す。式の外部リンクは最後の受け皿が外す。"""
    va, xl, wb, src, bk = _ud_world(tmp_path, monkeypatch)
    monkeypatch.setattr(vu, "_SNAPSHOT_MAX_CELLS", 1)
    assert va.undo_agent() is True
    out = capsys.readouterr().out
    assert "控えが大きすぎて「名簿」の数式を写せませんでした" in out
    assert wb.links_changed == [(bk, wb.FullName, 1)] and "このブック自身に差し替えました" in out
    assert wb.Sheets("名簿").f == {(2, 2): "=明細!B2"} and wb.Sheets("名簿").val("A2") == "佐藤"
    assert src.closed == [False]


def test_undo_that_fails_midway_still_closes_the_backup_and_restores_events(tmp_path, monkeypatch, capsys):
    va, xl, wb, src, bk = _ud_world(tmp_path, monkeypatch)
    xl.fail_copy = True
    assert va.undo_agent() is False
    assert "戻せませんでした" in capsys.readouterr().out
    assert src.closed == [False] and xl.EnableEvents is True and xl.DisplayAlerts is True


def test_undo_refuses_before_opening_the_backup(tmp_path, monkeypatch, capsys):
    """止まる入口: どれも控えを開く前（Excel の中身に触る前）に返る。dry-run は何も戻さない。"""
    va, xl, wb, src, bk = _ud_world(tmp_path, monkeypatch)
    assert va.undo_agent(dry_run=True) is True
    out = capsys.readouterr().out
    assert "--dry-run: 何も戻していません" in out and "sheet 新シート" in out and xl.opened == []
    wb.Name = "別.xlsx"                                                      # 別のブックが前にいる
    assert va.undo_agent() is False and "控えは「台帳.xlsx」のものです" in capsys.readouterr().out
    wb.Name = "台帳.xlsx"
    wb._sheets = [s for s in wb._sheets if s.Name != "集計"]                  # 戻す先のシートが無い
    assert va.undo_agent() is False and "シート 集計 がこのブックにありません" in capsys.readouterr().out
    assert xl.opened == [] and src.closed == []
    monkeypatch.setattr(vu, "get_workbook",
                        lambda *a, **k: (_ for _ in ()).throw(AssertionError("Excel に触った")))
    os.remove(bk)                                                            # 控えのファイルが消えた
    assert va.undo_agent() is False and "控えのファイルがありません" in capsys.readouterr().out
    va._save_undo_meta({'book': "台帳.xlsx", 'path': bk})                      # 戻す先が書いていない覚書
    assert va.undo_agent() is False and "戻す先のシートがありません" in capsys.readouterr().out
    os.remove(va._LAST_AGENT_UNDO_FILE)
    assert va.undo_agent() is False and "戻せる控えがありません" in capsys.readouterr().out


def test_undo_stops_when_the_backup_lacks_a_sheet(tmp_path, monkeypatch, capsys):
    va, xl, wb, src, bk = _ud_world(tmp_path, monkeypatch, src_sheets=('名簿', '明細'))
    assert va.undo_agent() is False
    assert "控えにシート 集計 がありません" in capsys.readouterr().out
    assert src.closed == [False] and xl.EnableEvents is True
    assert [s.Name for s in wb.Sheets] == ["名簿", "集計", "新シート", "明細"]   # 何も消していない


# ----------------------------------------------------------------
# run_macro_agent（macro モードのループ）を偽の AI・偽のコードで（2026-09-10 夜）。
# これまではソースの文字列検査だけで、関所が本当に done を止めるか・終わりの判定がどうなるかは誰も回していなかった。
# ----------------------------------------------------------------
_MR_CODE = {("Module1", "(宣言部)"): "",
            ("Module1", "集計"): "Sub 集計()\n  Dim n As String\nEnd Sub",
            ("Module1", "印字"): "Sub 印字()\n  On Error Resume Next\nEnd Sub"}
_MR_FIX = ('{"plan":[{"item":"型を直す","state":"未"}],"actions":[{"op":"code_replace","module":"Module1",'
           '"search":"String","replace":"Long"}],"done":false}')
_MR_DONE_PENDING = '{"plan":[{"item":"型を直す","state":"未"}],"actions":[],"done":true,"report":"直した"}'
_MR_DONE = '{"plan":[{"item":"型を直す","state":"済"}],"actions":[],"done":true,"report":"直した"}'
_MR_REVERT = ('{"plan":[{"item":"型を直す","state":"済"}],"actions":[{"op":"replace","name":"印字","code":"x"}],'
              '"done":false}')


def _mr_collateral(code, actions):
    """code_replace が隣のマクロの行にも当たった（search が一意でない）／replace で隣を戻す。"""
    op = actions[0]['op']
    if op == 'code_replace':
        code[("Module1", "集計")] = code[("Module1", "集計")].replace("String", "Long")
        code[("Module1", "印字")] = "Sub 印字()\n  On Error GoTo 0\nEnd Sub"
        return [("code_replace Module1", True, "2 行を書き換えました"), ("compile（書き換え後・道具が自動）", True, "通過")]
    if op == 'replace':
        code[("Module1", "印字")] = _MR_CODE[("Module1", "印字")]
        return [("replace 印字", True, "置き換えました"), ("compile（書き換え後・道具が自動）", True, "通過")]
    if op == 'add':
        code[("Module1", actions[0]['name'])] = actions[0]['code']
        return [(f"add {actions[0]['name']}", True, "足しました"), ("compile（書き換え後・道具が自動）", True, "通過")]
    raise AssertionError(op)


def _mr_clean(code, actions):
    """直すと言った 1 本だけを直す。"""
    code[("Module1", "集計")] = code[("Module1", "集計")].replace("String", "Long")
    return [("code_replace Module1", True, "1 行を書き換えました"), ("compile（書き換え後・道具が自動）", True, "通過")]


def _macro_loop(monkeypatch, tmp_path, replies, on_exec=_mr_clean, names=("集計", "印字"), compile_ok=True):
    import vbam_agent as va
    code = dict(_MR_CODE)
    seen = {'sent': [], 'exec': [], 'stale': [], 'compile': 0}

    def fake_ask(ai, model, key, history):
        seen['sent'].append(history[-1][1])
        return replies[min(len(seen['sent']) - 1, len(replies) - 1)], {'in': 1, 'out': 1}, 0.0

    def fake_exec(actions, rehearse, wb, existing=None, default_module=None, undo=None):   # undo＝直し 1 つ分の控え（2026-09-17 夜）
        seen['exec'].append((actions, list(existing or []), default_module))
        return on_exec(code, actions)

    def fake_run_cmd(toks, wb=None):
        assert toks == ['compile'], toks                    # 終わりの全体コンパイルだけ（他は偽の実行系が持つ）
        seen['compile'] += 1
        return compile_ok, "コンパイル: " + ("通過" if compile_ok else "エラー（型が一致しません）")

    monkeypatch.setattr(va, "_ai_setup", lambda ai, model: ('claude-code', 'sonnet', None))
    monkeypatch.setattr(va, "_ask", fake_ask)
    monkeypatch.setattr(va, "_all_procedure_names", lambda wb: list(names))
    monkeypatch.setattr(vmac, "_macro_materials",
                        lambda macro, wb=None: (f"--- get ---\n{code[('Module1', macro)]}", True, "Module1"))
    monkeypatch.setattr(vmac, "_macro_materials_create", lambda wb=None: ("--- モジュール ---\nModule1", True))
    monkeypatch.setattr(vmac, "_execute_macro", fake_exec)
    monkeypatch.setattr(vmac, "_code_snapshot", lambda wb: dict(code))
    monkeypatch.setattr(va, "_run_cmd", fake_run_cmd)
    monkeypatch.setattr(va, "_mark_undo_stale", lambda what, book=None: seen['stale'].append((what, book)))
    monkeypatch.setattr(va, "_LAST_AGENT_LOG_FILE", str(tmp_path / "log.jsonl"))
    monkeypatch.setattr(va, "_LAST_AGENT_ASK_FILE", str(tmp_path / "ask.txt"))
    return va, code, seen


def test_macro_loop_gates_hold_done_until_the_collateral_rewrite_is_undone(tmp_path, monkeypatch):
    """返事の形 → 手 → 未のまま done（止める）→ 隣のマクロまで変わったまま done（止める）→ 戻す → done で合格。"""
    replies = ['これは JSON ではない', _MR_FIX, _MR_DONE_PENDING, _MR_DONE, _MR_REVERT,
               '{"plan":[{"item":"型を直す","state":"済"}],"actions":[],"done":true,"report":"直して、隣は戻した"}']
    va, code, seen = _macro_loop(monkeypatch, tmp_path, replies, on_exec=_mr_collateral)
    wb = _InvObj(Name="台.xlsm")
    r = vmac.run_macro_agent("集計マクロの型を直して", None, wb, max_turns=6)
    assert r['done'] is True and r['ok'] is True and r['turns'] == 6 and r['wrote'] is True
    assert {k: v for k, v in r['gates'].items() if v} == {'format': 1, 'plan': 1, 'inv': 1}
    assert r['inv'] == [] and r['report'] == "直して、隣は戻した"
    sent = seen['sent']
    assert "返事が規約に合っていません" in sent[1]
    assert "未の項目が残っています: 型を直す" in sent[3]
    assert "頼んでいないマクロが変わった: Module1.印字" in sent[4] and "search は前後を足して一意に" in sent[4]
    # 実行系には「既にあるマクロ」と修理対象のモジュールを渡す。控えの印は同じブックに付ける
    assert len(seen['exec']) == 2 and seen['exec'][0][1:] == (["集計", "印字"], "Module1")
    assert seen['stale'] and seen['stale'][0][0].startswith("macro（マクロ「集計」") and seen['stale'][0][1] == "台.xlsm"
    assert seen['compile'] == 1 and code[("Module1", "印字")] == _MR_CODE[("Module1", "印字")]
    run = va._runs_load()[-1]
    assert run['mode'] == 'macro' and run['ok'] is True and run['path'] is None and run['book'] == "台.xlsm"
    assert run['gates'] == {'format': 1, 'plan': 1, 'inv': 1}


def test_macro_loop_fails_when_the_collateral_rewrite_or_compile_remains(tmp_path, monkeypatch, capsys):
    """関所は 1 回だけ。直さずに done を繰り返せば終われるが、終わりの照合で ok にしない。コンパイル不通過も同じ。"""
    va, code, seen = _macro_loop(monkeypatch, tmp_path, [_MR_FIX, _MR_DONE, _MR_DONE], on_exec=_mr_collateral)
    r = vmac.run_macro_agent("集計マクロの型を直して", None, _InvObj(Name="台.xlsm"), max_turns=3)
    assert r['done'] is True and r['ok'] is False
    assert r['inv'] == ["頼んでいないマクロが変わった: Module1.印字"] and r['gates']['inv'] == 1
    out = capsys.readouterr().out
    assert "頼んでいないコードの変化: 1 件" in out and "＝ 不合格" in out
    va, code, seen = _macro_loop(monkeypatch, tmp_path, [_MR_FIX, _MR_DONE], compile_ok=False)
    r = vmac.run_macro_agent("集計マクロの型を直して", None, _InvObj(Name="台.xlsm"), max_turns=3)
    assert r['done'] is True and r['inv'] == [] and r['ok'] is False
    assert "全体コンパイル 不通過" in capsys.readouterr().out


def test_macro_loop_last_turn_done_with_pending_items_is_dropped(tmp_path, monkeypatch, capsys):
    """最後の往復で「未」のまま done＝続きを頼めない。落とし物として不合格にする。"""
    va, code, seen = _macro_loop(monkeypatch, tmp_path, [_MR_FIX, _MR_DONE_PENDING])
    r = vmac.run_macro_agent("集計マクロの型を直して", None, _InvObj(Name="台.xlsm"), max_turns=2)
    assert r['done'] is False and r['ok'] is False and r['dropped'] == ["型を直す"]
    assert "落とし物: 未の項目が残ったまま往復が尽きました（1 件）" in capsys.readouterr().out


def test_macro_loop_dry_run_and_unknown_macro_touch_nothing(tmp_path, monkeypatch, capsys):
    """--dry-run は手を見せるだけ（実行・控えの印・コンパイル・台帳なし）。直すマクロが分からなければ AI に聞かない。"""
    va, code, seen = _macro_loop(monkeypatch, tmp_path, [_MR_FIX])
    r = vmac.run_macro_agent("集計マクロの型を直して", None, _InvObj(Name="台.xlsm"), dry_run=True)
    assert r['ok'] is True and r['done'] is False and r['turns'] == 1
    assert seen['exec'] == [] and seen['stale'] == [] and seen['compile'] == 0 and va._runs_load() == []
    assert '"op": "code_replace"' in capsys.readouterr().out
    va, code, seen = _macro_loop(monkeypatch, tmp_path, [_MR_FIX])
    r = vmac.run_macro_agent("表の数字がおかしい", None, _InvObj(Name="台.xlsm"))
    assert r == {'done': False, 'ok': False, 'turns': 0, 'report': '', 'usage': {}} and seen['sent'] == []
    assert "直すマクロが分かりません" in capsys.readouterr().out
    r = vmac.run_macro_agent("直して", None, _InvObj(Name="台.xlsm"), macro="無い名前")
    assert r['ok'] is False and seen['sent'] == [] and "マクロ '無い名前' が見つかりません" in capsys.readouterr().out


def test_macro_loop_creating_a_macro_is_not_flagged_as_an_unrequested_change(tmp_path, monkeypatch, capsys):
    """「作って」の回: 材料は作る用・マクロが増えても頼んでいない変化に数えない・add は書き換えに数える。"""
    add = ('{"plan":[{"item":"印刷のマクロを足す","state":"未"}],"actions":[{"op":"add","name":"印刷",'
           '"module":"Module1","code":"Sub 印刷()\\nEnd Sub"}],"done":false}')
    done = '{"plan":[{"item":"印刷のマクロを足す","state":"済"}],"actions":[],"done":true,"report":"足した"}'
    va, code, seen = _macro_loop(monkeypatch, tmp_path, [add, done], on_exec=_mr_collateral, names=("集計",))
    r = vmac.run_macro_agent("印刷用のマクロを新しく作って", None, _InvObj(Name="台.xlsm"))
    assert r['done'] is True and r['ok'] is True and r['wrote'] is True and r['inv'] == []
    assert r['gates']['inv'] == 0 and ("Module1", "印刷") in code
    assert seen['exec'][0][2] is None                                          # 作る回は入れ先を決め打ちしない
    out = capsys.readouterr().out
    assert "新しいマクロを作ります" in out and "頼んでいないコードの変化: なし" in out
    assert seen['stale'][0][0].startswith("macro（マクロ「新規」")


# ----------------------------------------------------------------
# fill の手（_do_fill）と、偽のシートの足回り（2026-09-10 夜）。
# fill は先頭セルに書いて末尾まで伸ばす＝書く手。人の値の上書き・文字で入る式・列全体の指定を道具が止める。
# ----------------------------------------------------------------
class _NzLines:
    """シートの Rows／Columns: .Count と、("4:9")／("B:D") でグループ化の相手を返す。"""

    def __init__(self, ws, count):
        self.ws, self.Count = ws, count

    def __call__(self, spec):
        ws = self.ws

        def ungroup():
            if spec not in ws.groups:
                raise RuntimeError("グループ化されていません")
            ws.groups.discard(spec)
        return _InvObj(Group=lambda: ws.groups.add(spec), Ungroup=ungroup)


def _nz_is_err(v):
    return isinstance(v, int) and not isinstance(v, bool) and v < -2146820000


def _nz_text(self):
    v = self.ws.v.get((self.r1, self.c1))
    fmt = self.ws.fmt.get((self.r1, self.c1), _NZ_GENERAL)
    if v is None:
        return ''
    if _nz_is_err(v):
        return va._XL_ERROR_CODES.get(v, '#ERR')
    if isinstance(v, (int, float)) and not isinstance(v, bool) and '#,##0' in fmt:
        return f"{v:,.0f}"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if hasattr(v, 'year'):
        return f"{v.year}/{v.month}/{v.day}"
    return str(v)


def _nz_shift(f, dr, dc):
    """式を dr 行・dc 列ずらす（$ の付いた側は動かさない）＝Excel のフィルと同じ相対参照。"""
    def rep(m):
        cabs, col, rabs, row = m.groups()
        if not cabs and dc:
            col = _nz_col(_nz_rc(col + '1')[1] + dc)
        if not rabs and dr:
            row = str(int(row) + dr)
        return f"{cabs}{col}{rabs}{row}"
    return _re.sub(r'(\$?)([A-Z]{1,3})(\$?)(\d+)', rep, f)


def _nz_fill(self):
    src = (self.r1, self.c1)
    for row in self._keys():
        for k in row:
            if k == src:
                continue
            self.ws.writes.append(k)
            self.ws.fmt[k] = self.ws.fmt.get(src, _NZ_GENERAL)          # 書式も先頭セルから写る
            f = self.ws.f.get(src)
            if f is not None:
                self.ws.f[k] = self.ws.v[k] = _nz_shift(f, k[0] - src[0], k[1] - src[1])
                if k in self.ws.error_cells:
                    self.ws.v[k] = -2146826281
            else:
                self.ws.f.pop(k, None)
                self.ws.v[k] = self.ws.v.get(src)


def _nz_special(self, kind, sub=None):
    if kind != -4123:
        raise RuntimeError("偽のシートが数えるのは数式のセルだけ")
    hits = _InvList(_NzBox(self.ws, r, c, r, c) for row in self._keys() for (r, c) in row
                    if (r, c) in self.ws.f and (sub != 16 or _nz_is_err(self.ws.v.get((r, c)))))
    if not hits:
        raise RuntimeError("該当するセルが見つかりません")         # COM と同じく「無い」は例外
    return hits


def _nz_formula2_set(self, v):
    self.ws.f2.add((self.r1, self.c1))
    _nz_formula_set(self, v)


_NzBox.Text = property(_nz_text)
_NzBox.FillDown = _nz_fill
_NzBox.FillRight = _nz_fill
_NzBox.SpecialCells = _nz_special
_NzBox.Formula2 = property(lambda self: self.Formula, _nz_formula2_set)
_NzBox.Select = lambda self: setattr(self.ws, 'selected', (self.r1, self.c1))


def test_fill_down_and_right_extend_the_formula_with_relative_references():
    import vbam_agent as va
    ws = _NzSheet({"B2": 100, "C2": 3, "B3": 200, "C3": 0, "B4": 300, "C4": 5})
    ws.error_cells = {(3, 4)}                                             # D3 = B3/C3 は #DIV/0!
    label, ok, out = va._do_fill(ws, {"op": "fill", "range": "D2:D4", "formula": "=B2/C2"})
    assert ok and label == "fill D2:D4"
    assert ws.f == {(2, 4): "=B2/C2", (3, 4): "=B3/C3", (4, 4): "=B4/C4"}
    assert "fill（下）: D2:D4 に '=B2/C2' を伸ばした（3 セル）  エラー 1" in out and "D3='#DIV/0!'" in out
    # 右へ（合計行を横に伸ばす）。$ の付いた側は動かない
    ws2 = _NzSheet({"B2": 1, "C2": 2, "D2": 3})
    _, _, out2 = va._do_fill(ws2, {"op": "fill", "range": "B5:D5", "formula": "=SUM(B$2:B4)*$A$1", "right": True})
    assert ws2.f == {(5, 2): "=SUM(B$2:B4)*$A$1", (5, 3): "=SUM(C$2:C4)*$A$1", (5, 4): "=SUM(D$2:D4)*$A$1"}
    assert "fill（右）: B5:D5" in out2 and "エラー 0" in out2
    # 値でも伸ばせる
    ws3 = _NzSheet()
    va._do_fill(ws3, {"op": "fill", "range": "A2:A4", "value": "未"})
    assert [ws3.val(f"A{r}") for r in (2, 3, 4)] == ["未"] * 3 and ws3.f == {}


def test_fill_refuses_to_overwrite_and_checks_the_shape_of_the_hand(monkeypatch):
    """人の値（先頭セルも）があれば承認なしでは 1 セルも書かない。形の違う手は実行前に断る。"""
    import vbam_agent as va
    import pytest
    ws = _NzSheet({"D2": "人の値", "D4": 7})
    with pytest.raises(ValueError, match=r'fill 先に値がある（先頭 D2 / D3:D5 に 1 セル）.*"overwrite": true'):
        va._do_fill(ws, {"op": "fill", "range": "D2:D5", "formula": "=B2*2"})
    assert ws.writes == []
    # 先頭に同じ数式がある＝人の数式を末尾まで伸ばす頼み方は上書きではない
    ws2 = _NzSheet(formulas={"D2": ("=B2*2", 0)})
    va._do_fill(ws2, {"op": "fill", "range": "D2:D4", "formula": "=B2*2"})
    assert ws2.f[(4, 4)] == "=B4*2"
    # 承認があれば上書きする
    va._do_fill(ws, {"op": "fill", "range": "D2:D5", "formula": "=B2*2", "overwrite": True})
    assert ws.f[(2, 4)] == "=B2*2" and ws.f[(5, 4)] == "=B5*2"
    monkeypatch.setattr(vh, "_FILL_MAX_CELLS", 3)
    for act, msg in (({"range": "D2:D5"}, "formula"),
                     ({"range": "B2:C5", "formula": "=1"}, "1 列（下へ伸ばす）か 1 行"),
                     ({"range": "D2", "formula": "=1"}, "2 セル以上"),
                     ({"range": "D2:D5", "formula": "=1"}, "1 回 3 セルまで"),
                     ({"range": "D2:D3", "formula": "B2*2"}, "'=' で始める")):
        fresh = _NzSheet()
        with pytest.raises(ValueError, match=msg):
            va._do_fill(fresh, dict(act, op="fill"))
        assert fresh.writes == []


def test_fill_into_a_text_formatted_column_writes_real_formulas():
    """書き先が文字列書式（@）でも式は式で入る（2026-09-04 の実機）。動的配列の式は Formula2 で書く。"""
    import vbam_agent as va
    ws = _NzSheet({"B2": 1, "B3": 2}, fmts={"C2": "@", "C3": "@"})
    va._do_fill(ws, {"op": "fill", "range": "C2:C3", "formula": "=B2*10"})
    assert ws.f == {(2, 3): "=B2*10", (3, 3): "=B3*10"}
    assert ws.fmt_of("C2") == _NZ_GENERAL and ws.fmt_of("C3") == _NZ_GENERAL
    ws2 = _NzSheet()
    va._do_fill(ws2, {"op": "fill", "range": "E2:E3", "formula": "=FILTER(A2:A9,B2:B9>0)"})
    assert (2, 5) in ws2.f2 and ws2.f[(3, 5)] == "=FILTER(A3:A10,B3:B10>0)"


# ----------------------------------------------------------------
# 道具が COM で直接やる手（_direct_action）の残りを偽のブックで（2026-09-10 夜）。
# 計算だけの手・画像・グループ化・枠固定・印刷設定・グラフ・CSV 書き出し。これまで通っていたのは 3 割だった。
# ----------------------------------------------------------------
_NzBox.Left = property(lambda self: (self.c1 - 1) * 48.0)          # 列幅 48pt・行高 15pt の方眼
_NzBox.Top = property(lambda self: (self.r1 - 1) * 15.0)
_NzBox.Width = property(lambda self: (self.c2 - self.c1 + 1) * 48.0)
_NzBox.Height = property(lambda self: (self.r2 - self.r1 + 1) * 15.0)


class _XmColl(list):
    """COM のコレクション: .Count・.Item(i)（1 起点）・(番号か名前)。例外を入れておくとその番目で落ちる。"""

    @property
    def Count(self):
        return len(self)

    def Item(self, i):
        x = self[i - 1]
        if isinstance(x, Exception):
            raise x
        return x

    def __call__(self, key):
        if isinstance(key, int):
            return self.Item(key)
        for x in self:
            if getattr(x, 'Name', None) == key:
                return x
        raise RuntimeError(f"{key} がありません")


class _XmPic:
    """図形・画像。縦横比を固定（LockAspectRatio=-1）していれば、幅を変えると高さも追う。"""

    def __init__(self, name, left=0.0, top=0.0, width=200.0, height=100.0, on_action=''):
        self.Name, self.Left, self.Top, self.OnAction = name, left, top, on_action
        self._w, self._h, self.LockAspectRatio, self.coll = width, height, 0, None

    @property
    def Width(self):
        return self._w

    @Width.setter
    def Width(self, w):
        if self.LockAspectRatio == -1:
            self._h = self._h * w / self._w
        self._w = w

    @property
    def Height(self):
        return self._h

    @Height.setter
    def Height(self, h):
        if self.LockAspectRatio == -1:
            self._w = self._w * h / self._h
        self._h = h

    def Delete(self):
        self.coll.remove(self)


class _XmShapes(_UdShapes):
    def __init__(self, items=()):
        super().__init__(items)
        self.pictures = []
        for s in self:
            s.coll = self

    def AddPicture(self, path, link, save, left, top, w, h):
        self.pictures.append((path, link, save, left, top, w, h))
        s = _XmPic(f"図 {len(self) + 1}", left, top)
        s.coll = self
        self.append(s)
        return s


class _XmCells:
    """シートの Cells: (r, c) で 1 セル。条件付き書式・入力規則の数・最後のセル探し（Find）も持つ。"""

    def __init__(self, ws):
        self.ws, self.FormatConditions = ws, _XmColl(ws.cf)

    def __call__(self, r, c):
        return _NzBox(self.ws, r, c, r, c)

    def SpecialCells(self, kind):
        if kind == -4174 and self.ws.n_dv:
            return _InvObj(Count=self.ws.n_dv)
        raise RuntimeError("該当するセルが見つかりません")

    def Find(self, what, after=None, look_in=None, look_at=None, order=1, direction=2):
        keys = [k for k, v in self.ws.v.items() if v not in (None, '')]
        if not keys:
            return None
        r, c = max(keys) if order == 1 else max(keys, key=lambda k: (k[1], k[0]))
        return _NzBox(self.ws, r, c, r, c)


class _XmCharts:
    def __init__(self, ws):
        self.ws = ws

    @property
    def Count(self):
        return len(self.ws.charts)

    def Add(self, left, top, w, h):
        chart = _InvObj(ChartType=None, HasTitle=False, ChartTitle=_InvObj(Text=''), source=None)
        chart.SetSourceData = lambda rng: setattr(chart, 'source', rng)
        co = _InvObj(Name=f"グラフ {len(self.ws.charts) + 1}", Chart=chart, left=left, top=top, width=w, height=h)
        co.Delete = lambda: self.ws.charts.remove(co)
        self.ws.charts.append(co)
        return co


class _XmWin:
    """ウィンドウ: 枠固定は「そのとき選んでいるセルの左上」で決まる（A1 を選んで固定しても固定にならない）。"""

    def __init__(self, book):
        self.book, self.SplitRow, self.SplitColumn, self.Zoom = book, 0, 0, 100

    @property
    def FreezePanes(self):
        return bool(self.SplitRow or self.SplitColumn)

    @FreezePanes.setter
    def FreezePanes(self, on):
        sel = getattr(self.book.ActiveSheet, 'selected', None) or (1, 1)
        self.SplitRow, self.SplitColumn = (sel[0] - 1, sel[1] - 1) if on else (0, 0)


class _XmSheet(_UdSheet):
    def __init__(self, xl, name, cells=None, formulas=None, shapes=()):
        super().__init__(xl, name, cells, formulas)
        self.Shapes = _XmShapes(shapes)
        self.Visible, self.ProtectContents, self.ProtectDrawingObjects = -1, False, False
        self.PageSetup = _InvObj(PrintArea='', Orientation=1, PrintTitleRows='', Zoom=100, FitToPagesWide=1,
                                 FitToPagesTall=1, LeftFooter='', CenterFooter='', RightFooter='',
                                 CenterHorizontally=False)
        self.HPageBreaks, self.VPageBreaks = _InvObj(Count=0), _InvObj(Count=0)
        self.charts, self.cf, self.n_dv, self.evals, self.pivots = [], [], 0, {}, []
        self.ListObjects, self.extra_used, self.AutoFilterMode, self.AutoFilter = _XmColl(), None, False, None

    @property
    def Cells(self):
        return _XmCells(self)

    @property
    def UsedRange(self):
        box = _UdSheet.UsedRange.fget(self)
        if not self.extra_used:
            return box
        r1, c1, r2, c2 = self.extra_used                     # 書式だけ残った行・列（使用範囲の膨張）
        return _NzBox(self, min(box.r1, r1), min(box.c1, c1), max(box.r2, r2), max(box.c2, c2))

    def Activate(self):
        self.activated += 1
        self.book.ActiveSheet = self

    def ChartObjects(self, name=None):
        if name is None:
            return _XmCharts(self)
        for co in self.charts:
            if co.Name == name:
                return co
        raise RuntimeError(f"グラフ {name} がありません")

    def PivotTables(self):
        return list(self.pivots)

    def Evaluate(self, f):
        if f in self.evals:
            return self.evals[f]
        raise RuntimeError("式が正しくありません")


class _XmBook(_UdBook):
    def __init__(self, name, path, sheets):
        super().__init__(name, path, sheets)
        for s in self._sheets:
            s.Parent = self
        self.ActiveSheet, self.pages = self._sheets[0], None
        self.Names, self.SlicerCaches, self.Queries = _XmColl(), _XmColl(), _XmColl()
        self.Model = _InvObj(ModelTables=_XmColl(), ModelMeasures=_XmColl())
        book = self
        count_a = lambda rng: sum(1 for row in rng._keys() for k in row if rng.ws.v.get(k) not in (None, ''))  # noqa: E731
        self.Application = _InvObj(ActiveWindow=_XmWin(self), Calculation=-4105,
                                   ExecuteExcel4Macro=lambda s: book.pages,
                                   WorksheetFunction=_InvObj(CountA=count_a))


def _xm_world():
    """会員名簿（1 行目タイトル・3 行目見出し）と集計と隠しシートのブック。マクロ付きのボタンが 1 つ。"""
    xl = _UdXl()
    meibo = _XmSheet(xl, "名簿", {"A1": "会員名簿", "A3": "会員番号", "B3": "氏名", "C3": "金額", "D3": "入会日",
                                  "A4": "0001", "B4": "佐藤", "C4": 1200.0, "D4": _dt.datetime(2026, 4, 1),
                                  "A5": "0002", "B5": "鈴木", "C5": -2146826246},
                     shapes=[_XmPic("btn印刷", on_action="印刷する")])
    hidden = _XmSheet(xl, "隠し", {"A1": "x"})
    hidden.Visible = 0
    wb = _XmBook("台帳.xlsx", "C:\\work", [meibo, _XmSheet(xl, "集計", {"A1": "合計", "B1": 100}), hidden])
    return meibo, wb


def test_direct_eval_computes_without_writing():
    """eval は式を計算して見るだけ（セルに書かない）。エラー値は名前・配列は先頭・計算できない式は理由つき。"""
    import vbam_agent as va
    import pytest
    ws, wb = _xm_world()
    ws.evals = {"=SUM(C4:C5)": 1500.0, "=1/0": -2146826281, "=C4:C5": ((1.0,), (2.0,))}
    label, ok, out = va._direct_action({"op": "eval", "formulas": ["=SUM(C4:C5)", "1/0", "=C4:C5", "=壊れ("]},
                                       "名簿", wb)
    assert (label, ok) == ("eval 4", True) and ws.writes == []
    assert "=SUM(C4:C5) → 1500" in out and "1/0 → #DIV/0!" in out and "=C4:C5 → 配列 2 個: 1 2" in out
    assert "=壊れ( → 計算できません" in out and "セルには書いていません" in out
    assert va._direct_action({"op": "eval", "formula": "=SUM(C4:C5)"}, "名簿", wb)[0] == "eval 1"
    for act, msg in (({"op": "eval"}, "formula"), ({"op": "eval", "formulas": ["=1", " "]}, "formula"),
                     ({"op": "eval", "formulas": ["=1"] * 21}, "20 本まで")):
        with pytest.raises(ValueError, match=msg):
            va._direct_action(act, "名簿", wb)


def test_direct_image_places_a_picture_with_a_backslash_path(tmp_path):
    """画像はセルの左上に置く。Excel は「/」区切りのパスを見つけられない＝円記号の絶対パスで渡す（9/8 の実射）。"""
    import vbam_agent as va
    import pytest
    ws, wb = _xm_world()
    pic = tmp_path / "写真.png"
    pic.write_bytes(b"\x89PNG")
    label, ok, out = va._direct_action({"op": "image", "path": str(pic).replace("\\", "/"), "at": "C5",
                                        "width": 120, "name": "写真1"}, "名簿", wb)
    assert ok and label == "image 写真.png"
    assert ws.Shapes.pictures[0][0] == os.path.abspath(str(pic)) and "/" not in ws.Shapes.pictures[0][0]
    shp = ws.Shapes("写真1")
    assert (shp.Left, shp.Top, shp.Width, shp.Height) == (96.0, 60.0, 120.0, 60.0)     # 縦横比は保つ
    assert "位置 C5 の左上" in out and "w=120 h=60" in out
    for act, msg in (({"op": "image"}, "path"), ({"op": "image", "path": str(tmp_path / "a.txt")}, "画像ファイルです"),
                     ({"op": "image", "path": str(tmp_path / "無い.png")}, "見つかりません")):
        with pytest.raises(ValueError, match=msg):
            va._direct_action(act, "名簿", wb)


def test_direct_group_folds_rows_and_columns_without_hiding():
    import vbam_agent as va
    import pytest
    ws, wb = _xm_world()
    label, ok, out = va._direct_action({"op": "row_group", "from": 4, "to": 9}, "名簿", wb)
    assert (label, ok) == ("row_group 4:9", True) and ws.groups == {"4:9"} and "隠してはいません" in out
    assert va._direct_action({"op": "col_group", "from": "b", "to": "d"}, "名簿", wb)[0] == "col_group B:D"
    _, _, out = va._direct_action({"op": "row_group", "from": 4, "to": 9, "off": True}, "名簿", wb)
    assert "グループ化を解除: 4:9" in out and ws.groups == {"B:D"}
    with pytest.raises(ValueError, match="解除する所がグループ化されていない"):
        va._direct_action({"op": "row_group", "from": 4, "to": 9, "off": True}, "名簿", wb)
    for act, msg in (({"op": "row_group", "from": 4}, "from と to"), ({"op": "row_group", "from": 9, "to": 4}, "from ≦ to"),
                     ({"op": "col_group", "from": "1", "to": "C"}, "列文字")):
        with pytest.raises(ValueError, match=msg):
            va._direct_action(act, "名簿", wb)


def test_direct_view_freezes_panes_and_reads_them_back():
    """枠固定は言いっぱなしにせず読み戻す。A1 で固定しても固定にならない＝そう言う。"""
    import vbam_agent as va
    import pytest
    ws, wb = _xm_world()
    wb.ActiveSheet = wb.Sheets("集計")
    label, ok, out = va._direct_action({"op": "view", "freeze": "B4", "zoom": 85}, "名簿", wb)
    win = wb.Application.ActiveWindow
    assert (label, ok) == ("view", True) and (win.SplitRow, win.SplitColumn, win.Zoom) == (3, 1, 85)
    assert "枠固定: B4 の左上で固定（上 3 行・左 1 列を固定）" in out and "表示倍率: 85%" in out
    assert wb.ActiveSheet is ws and ws.selected == (1, 1)                  # そのシートで固定し、選択は A1 に戻す
    _, _, out = va._direct_action({"op": "view", "freeze": "A1"}, "名簿", wb)
    assert "⚠ 枠固定: A1 を指定しましたが固定されていません" in out
    _, _, out = va._direct_action({"op": "view", "freeze": "off"}, "名簿", wb)
    assert out == "枠固定: 解除" and win.FreezePanes is False
    with pytest.raises(ValueError, match="freeze か zoom"):
        va._direct_action({"op": "view"}, "名簿", wb)


def test_direct_page_setup_sets_and_reports_the_print_settings():
    """印刷設定は当てた後の現物を 1 行で返す（FitToPage の名前つき＝採点係が読み違えない）。"""
    import vbam_agent as va
    import pytest
    ws, wb = _xm_world()
    wb.pages = 2
    label, ok, out = va._direct_action({"op": "page_setup", "area": "A1:F20", "title_rows": "3", "landscape": True,
                                        "fit_wide": 1, "fit_tall": 0, "center_h": True, "footer_page": True},
                                       "名簿", wb)
    ps = ws.PageSetup
    assert (label, ok) == ("page_setup", True)
    assert (ps.PrintArea, ps.PrintTitleRows, ps.Orientation) == ("$A$1:$F$20", "$3:$3", 2)
    assert ps.Zoom is False and ps.FitToPagesWide == 1 and ps.FitToPagesTall is False and ps.CenterHorizontally
    assert out == ("印刷: 範囲 A1:F20 / 向き 横 / タイトル行 3:3 / 横 1 ページ×縦 自動 ページに収める"
                   "（FitToPage=ON・FitToPagesWide=1・FitToPagesTall=0） / フッター &P / &N / ページ数 2")
    _, _, out = va._direct_action({"op": "page_setup", "zoom": 80, "landscape": False}, "名簿", wb)
    assert "倍率 80%（FitToPage=OFF" in out and "向き 縦" in out
    # 前に出ていないシートのページ数は改ページの数から数える
    wb.ActiveSheet, ws.HPageBreaks = wb.Sheets("集計"), _InvObj(Count=2)
    assert va._direct_action({"op": "page_setup", "center_h": True}, "名簿", wb)[2].endswith("ページ数 3")
    with pytest.raises(ValueError, match="当てる項目がありません"):
        va._direct_action({"op": "page_setup"}, "名簿", wb)


def test_direct_chart_is_recorded_so_undo_can_delete_it(monkeypatch):
    """グラフの結果の文は「作った物」として拾える形（--undo が消す）。ピボットを元にすればピボットグラフ。"""
    import vbam_agent as va
    import vbam_heavy
    import pytest
    ws, wb = _xm_world()
    label, ok, out = va._direct_action({"op": "chart", "range": "A3:C9", "type": "line", "title": "月別"}, "名簿", wb)
    co = ws.charts[0]
    assert (label, ok) == ("chart line", True) and co.Chart.ChartType == 4 and co.Chart.source.Address == "$A$3:$C$9"
    assert co.Chart.HasTitle is True and co.Chart.ChartTitle.Text == "月別"
    assert (co.left, co.top, co.width, co.height) == (154.0, 30.0, 360.0, 216.0)    # 既定はデータの右隣
    assert out == "グラフ作成: [名簿] グラフ 1  種別 line  データ A3:C9  表題 月別  位置 l=154 t=30 w=360 h=216"
    assert va._created_from_results([(label, ok, out)]) == [{'kind': 'chart', 'sheet': '名簿', 'name': 'グラフ 1'}]
    psh = wb.Sheets("集計")
    pt = _InvObj(Name="P集計", TableRange1=psh.Range("A3:C8"))
    monkeypatch.setattr(vbam_heavy, "_find_pivot", lambda wb_, name: (psh, pt) if name == "P集計" else (None, None))
    _, ok, out = va._direct_action({"op": "chart", "pivot": "P集計", "at": "E3"}, "名簿", wb)
    assert "[集計] グラフ 1" in out and "（ピボットグラフ）" in out and "l=192 t=30" in out
    assert psh.charts[0].Chart.ChartType == 51 and len(ws.charts) == 1
    for act, msg in (({"op": "chart", "range": "A3:C9", "type": "bubble"}, "type は"),
                     ({"op": "chart", "pivot": "無い"}, "ピボット '無い' が見つかりません")):
        with pytest.raises(ValueError, match=msg):
            va._direct_action(act, "名簿", wb)
    assert len(ws.charts) == 1 and len(psh.charts) == 1                       # 断った手は何も作らない


def test_direct_export_csv_keeps_ids_dates_and_error_text(tmp_path):
    """CSV: 先頭ゼロは文字のまま・整数は桁を落とさない・日付は指定の形・エラーは #N/A の文字。上書きはしない。"""
    import vbam_agent as va
    import pytest
    ws, wb = _xm_world()
    path = tmp_path / "会員.csv"
    label, ok, out = va._direct_action({"op": "export_csv", "range": "A3:D5", "path": str(path),
                                        "date_format": "yyyy/mm/dd"}, "名簿", wb)
    assert (label, ok) == ("export_csv", True)
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")                                    # 既定は BOM つき UTF-8（Excel で化けない）
    assert raw.decode("utf-8-sig").splitlines() == ["会員番号,氏名,金額,入会日", "0001,佐藤,1200,2026/04/01",
                                                    "0002,鈴木,#N/A,"]
    assert f"CSV 書き出し: {path}（3 行 x 4 列・utf-8-sig・日付 %Y/%m/%d）" in out and "末尾: 0002,鈴木,#N/A," in out
    for act, msg in (({"path": str(path)}, "既にあります"), ({"path": str(tmp_path / "a.txt")}, r"\.csv で指定"),
                     ({"path": str(tmp_path / "別.csv"), "encoding": "latin-1"}, "encoding は")):
        with pytest.raises(ValueError, match=msg):
            va._direct_action(dict(act, op="export_csv", range="A3:D5"), "名簿", wb)
    assert raw == path.read_bytes() and not (tmp_path / "別.csv").exists()


def test_direct_read_sheet_keeps_the_job_clock_and_refuses_unknown_hands(monkeypatch):
    """read_sheet は材料＋追加の材料を返し、仕事の時計を押し直さない。無いシート・知らない手は断る。"""
    import vbam_agent as va
    import pytest
    ws, wb = _xm_world()
    calls, clock = [], []
    monkeypatch.setattr(va, "_run_cmd", lambda toks, wb=None: (calls.append(toks), (True, "使用範囲: A1:D5\n"))[1])
    monkeypatch.setattr(va, "job_clock_get", lambda: "T0")
    monkeypatch.setattr(va, "job_clock_set", lambda c: clock.append(c))
    monkeypatch.setattr(va, "_extra_materials", lambda wb_, ws_: f"追加（{ws_.Name}）\n")
    label, ok, out = va._direct_action({"op": "read_sheet", "sheet": "集計"}, "名簿", wb)
    assert (label, ok, out) == ("read_sheet 集計", True, "使用範囲: A1:D5\n追加（集計）\n")
    assert calls == [["materials", "集計"]] and clock == ["T0"]
    monkeypatch.setattr(va, "_extra_materials", lambda wb_, ws_: 1 / 0)
    label, _, out = va._direct_action({"op": "inspect"}, "名簿", wb)
    assert label == "read_sheet 名簿" and "（追加の材料を読めませんでした: division by zero）" in out
    # 無いシートは実行前の検査が止める（1 手も実行しない）
    res = va._execute([{"op": "read_sheet", "sheet": "無い"}], "名簿", wb)
    assert res[0][1] is False and "シート '無い' がありません（ある: 名簿, 集計, 隠し）" in res[0][2]
    with pytest.raises(ValueError, match="許していない手です: 'format_disk'"):
        va._direct_action({"op": "format_disk"}, "名簿", wb)
    # マクロの付いていない図形が 1 つも無ければ、全部消す手は何もしない（ボタンは残る）
    assert va._direct_action({"op": "shape", "delete": True, "all": True}, "名簿", wb) == \
        ("shape delete 0", True, "マクロの付いていない図形はありません（消すものなし）")
    assert [s.Name for s in ws.Shapes] == ["btn印刷"]


def test_direct_read_file_prefers_the_book_that_is_already_open(tmp_path):
    """読むファイルが同じ Excel で開いたままなら、保存前の直しも見えるようにブックから読む。"""
    import vbam_agent as va
    src = tmp_path / "旧台帳.xlsx"
    src.write_bytes(b"x")                                                     # 中身は読まない（開いているブックから読む）
    rows = (("番号", "氏名"), ("0001", "佐藤"), ("0002", "鈴木"))
    opened = _InvObj(Name="旧台帳.xlsx", FullName=str(src),
                     Sheets=lambda key: _InvObj(Name="一覧", UsedRange=_InvObj(Value=rows)))
    wb = _rf_book("名簿")
    wb.Application = _InvObj(Workbooks=_XmColl([opened]))
    name, ok, out = va._direct_action({"op": "read_file", "path": str(src)}, "名簿", wb)
    assert (name, ok) == ("read_file", True)
    assert "読み込み: 開いているブック 旧台帳.xlsx!一覧 → 取込_旧台帳!A1:B3（3 行 x 2 列）" in out
    assert "先頭ゼロを守った列: A" in out and wb.Sheets("取込_旧台帳").last.Value[1] == ("0001", "佐藤")


# ----------------------------------------------------------------
# 追加の材料（_extra_materials）を偽のブックで（2026-09-10 夜）。
# AI が表を読み違えないための事実（見出しの位置・番号列・膨張・外部リンク・名前・空欄…）を道具が数えて渡す。
# 0% だった＝どれか 1 つが黙って読めなくなっても誰も気づかない。
# ----------------------------------------------------------------
_NzBox.EntireRow = property(lambda self: _InvObj(RowHeight=self.ws.row_height))


def _xm_rich_world():
    """1 行目タイトル・3 行目見出しの会員名簿。右に突き合わせ用の「会員番号」列（数値）と、乱れを一通り。"""
    xl = _UdXl()
    meibo = _XmSheet(xl, "名簿", {
        "A1": "会員名簿",
        "A3": "会員番号", "B3": "氏名", "C3": "金額", "D3": "入会日", "E3": "備考", "F3": "計", "H3": "会員番号",
        "A4": "0001", "B4": "佐藤　一郎", "C4": 1200.0, "D4": _dt.datetime(2026, 4, 1), "E4": "　", "H4": 1.0,
        "A5": "１００２", "B5": " 鈴木 花子", "C5": "12,000", "D5": "2026/5/12", "H5": 2.0,
        "A6": "0003", "B6": "田中 次郎", "C6": 800.0, "D6": _dt.datetime(2026, 6, 3),
        "A7": "0004", "C7": 500.0, "D7": _dt.datetime(2026, 7, 20)},
        formulas={"F4": ("=C4*単価", 0.0), "F5": ("='C:\\受領\\[外.xlsx]Sheet1'!A1", 0.0)})
    meibo.extra_used = (1, 1, 20, 10)                      # 書式だけ残った J20 まで使用範囲が膨らんでいる
    meibo.ProtectContents, meibo.row_height, meibo.n_dv = True, None, 5
    meibo.cf = [_InvObj(), _InvObj()]
    meibo.charts = [_InvObj(Name="グラフ 1")]
    meibo.ListObjects = _XmColl([_InvObj(Name="T会員", Range=_InvObj(Address="$A$3:$F$7"),
                                         ListColumns=[_InvObj(Name="会員番号"), _InvObj(Name="氏名")])])
    shukei = _XmSheet(xl, "集計", {"A1": "合計", "B1": 100})
    shukei.pivots = [_InvObj(Name="P集計", SourceData="名簿!R3C1:R7C6", TableRange2=_InvObj(Address="$A$3:$D$9"),
                             PivotFields=lambda: _XmColl([_InvObj(Name="地域", Orientation=1),
                                                          _InvObj(Name="商品", Orientation=2),
                                                          _InvObj(Name="売上", Orientation=0)]),
                             DataFields=_XmColl([_InvObj(Name="合計 / 売上")]))]
    hidden = _XmSheet(xl, "隠し", {"A1": "x"})
    hidden.Visible = 0
    wb = _XmBook("台帳.xlsx", "C:\\work", [meibo, shukei, hidden])
    wb.ActiveSheet = shukei                                 # 前に出ているのは別のシート
    wb.Application.Calculation = -4135                      # 手動計算
    win = wb.Application.ActiveWindow
    win.SplitRow, win.SplitColumn, win.Zoom = 3, 1, 90
    wb.Names = _XmColl([_InvObj(Name="単価", RefersTo="=集計!$B$1", Visible=True),
                        _InvObj(Name="古い範囲", RefersTo="=#REF!$A$1", Visible=True),
                        _InvObj(Name="未使用名", RefersTo="=名簿!$C$3", Visible=False),
                        _InvObj(Name="名簿!Print_Area", RefersTo="=名簿!$A$1:$F$7", Visible=True),
                        _InvObj(Name="_xlfn.XLOOKUP", RefersTo="=#NAME?", Visible=False)])
    wb.SlicerCaches = _XmColl([_InvObj(SourceName="P集計", Slicers=[_InvObj(Name="地域")])])
    wb.Queries = _XmColl([_InvObj(Name="売上取込")])
    wb.Model = _InvObj(ModelTables=_XmColl([_InvObj(Name="売上")]), ModelMeasures=_XmColl([_InvObj(Name="合計")]))
    return meibo, wb


def test_extra_materials_counts_what_the_ai_would_misread(monkeypatch):
    import vbam_agent as va
    monkeypatch.setattr(vh, "_PROFILE_MIN_ROWS", 5)
    ws, wb = _xm_rich_world()
    lines = va._extra_materials(wb, ws).splitlines()
    text = "\n".join(lines)

    def line(head):
        hit = [s for s in lines if s.startswith(head)]
        assert hit, f"{head!r} の行が無い:\n{text}"
        return hit[0]
    assert lines[0] == "--- 追加の材料（道具が読んだ現物） ---"
    assert line("他のシート: ").startswith("他のシート: 集計(A1:B1), 隠し(A1・非表示)（read_sheet か")
    assert line("見出し行の推定: ").startswith("見出し行の推定: 3 行目（会員番号 氏名 金額 入会日 備考 計 会員番号）　← 1 行目ではありません")
    assert "「会員番号」の列で型が違います（A4＝文字 / H4＝数値）" in line("キーの型: ")
    assert line("計算モード: ").startswith("計算モード: 手動  ← 数式を書いても値が更新されません")
    assert line("保護: ").startswith("保護: このシートは保護されています（セル）")
    assert line("枠固定: ") == "枠固定: B4 の左上で固定  表示倍率: 90%"
    assert wb.ActiveSheet.Name == "集計"                                   # 枠固定を読むために前に出したら戻す
    assert line("行高: ") == "行高: 不揃い"
    assert line("印刷: ") == ("印刷: 範囲 （未設定＝使用範囲） / 向き 縦 / タイトル行 なし / "
                            "倍率 100%（FitToPage=OFF＝ページに収める設定ではない） / フッター なし / ページ数 1")
    assert line("条件付き書式: ") == "条件付き書式: 2 本  入力規則: 5 セル  グラフ: 1 個"
    assert "  [名簿] T会員 範囲=A3:F7 列=会員番号,氏名" in lines
    assert "  [集計] P集計 出力=A3:D9 元=名簿!R3C1:R7C6 行=地域 列=商品 フィルタ=- 値=合計 / 売上" in lines
    assert line("スライサー: ") == "スライサー: 1 個 地域（P集計）"
    assert line("パワークエリ: ").startswith("パワークエリ: 1 本 売上取込")
    assert line("データモデル: ") == "データモデル: テーブル 1 本 売上  メジャー 1 本"
    assert line("外部リンク: ") == "外部リンク: 1 本  外.xlsx  このシートの参照セル: F5"
    assert line("名前定義（ブック全体）: ").startswith("名前定義（ブック全体）: 4 本  壊れ 1  未使用 1")
    assert "  単価 = =集計!$B$1  使用中" in lines and "  古い範囲 = =#REF!$A$1  壊れている(#REF!)" in lines
    assert "  未使用名 = =名簿!$C$3  未使用・非表示" in lines and "XLOOKUP" not in text     # 関数の隠し名は人の名前ではない
    assert "  名簿!Print_Area = =名簿!$A$1:$F$7  Excel が管理（印刷範囲・タイトル行・フィルタ）＝消さない" in lines
    assert line("データの末尾: ").startswith("データの末尾: H7（使用範囲の端 J20 のほうが大きい＝膨張。")
    assert line("空白に見えるゴミ") == "空白に見えるゴミ（全角空白・空白だけ・改行だけ）: 1 セル E4"
    assert line("前後に空白のある値: ") == "前後に空白のある値: 1 セル B5"
    assert line("列の中の文字種").endswith(": A列 全角英数 1／B列 文字の間の空白=全角 1・半角 2")
    # 番号列の文字の数字は正常（直さない側に分ける）。全角の番号は乱れのほうに残す
    assert line("数値に見える文字") == "数値に見える文字（左寄せの数字・全角数字。番号列は除く）: 2 セル A5 C5"
    assert line("番号列の文字の数字").endswith(": 3 セル A4 A6 A7")
    assert line("日付に見える文字: ") == "日付に見える文字: 1 セル D5"
    assert line("表の中の空欄").startswith("表の中の空欄（見出しの下で、同じ列の他の行は埋まっている）: 1 セル B7（")
    assert line("--- 列プロファイル").startswith("--- 列プロファイル（全 20 行・10 列")
    assert lines[-1] == "ファイル: 未保存（大きさは保存後に測れる）"


def test_extra_materials_plain_book_and_one_that_reads_nothing(tmp_path):
    """素直なブックでは「なし」「一致」を言う。何も読めないブックでも落ちずに、読めなかったと言う。"""
    import vbam_agent as va
    ws, wb = _xm_world()
    saved = tmp_path / "台帳.xlsx"
    saved.write_bytes(b"x" * 1234)
    wb.FullName = str(saved)
    lines = va._extra_materials(wb, ws).splitlines()
    for want in ("計算モード: 自動", "枠固定: なし  表示倍率: 100%", "行高: 均一 15", "外部リンク: なし",
                 "名前定義（ブック全体）: なし", "データの末尾: D5（使用範囲と一致）", f"ファイル: 1,234 バイト  {saved}",
                 "テーブル（ブック全体）: 0 本（table create で作れる）"):
        assert want in lines, (want, lines)
    assert any(s.startswith("見出し行の推定: 3 行目（会員番号 氏名 金額 入会日）") for s in lines)
    assert not any(s.startswith(("保護: ", "キーの型: ", "--- 列プロファイル")) for s in lines)
    broken = va._extra_materials(_InvObj(), _InvObj(Name="壊"))
    assert broken.startswith("--- 追加の材料（道具が読んだ現物） ---\n") and broken.count("\n") == 3
    assert "テーブル・ピボットの材料を読めませんでした" in broken and "列プロファイルを読めませんでした" in broken


# ----------------------------------------------------------------
# 採点係に渡す現物（_grade_structure・_grade_looks・_grade_outside_numbers・_looks_of_range）を偽のシートで（2026-09-10 夜）。
# セルに写らない構造と見た目。ここが読めないと、採点係は「確認できない」を unmet に挙げて往復を捨てる（9/8・9/9 の実射）。
# ----------------------------------------------------------------
def _nz_rows_flag(self, rows_set, yes, no):
    got = {r in rows_set for r in range(self.r1, self.r2 + 1)}
    return yes if got == {True} else (no if got == {False} else None)       # COM は「まちまち」を None で返す


def _nz_region(self):
    """CurrentRegion: 空の行と空の列で囲まれるまで広げる（斜めの隣も含む＝Excel と同じ）。"""
    ws = self.ws

    def filled(r, c):
        return ws.v.get((r, c)) not in (None, '') or (r, c) in ws.f
    r1, c1, r2, c2 = self.r1, self.c1, self.r2, self.c2
    while True:
        grown = False
        if r1 > 1 and any(filled(r1 - 1, c) for c in range(max(1, c1 - 1), c2 + 2)):
            r1, grown = r1 - 1, True
        if any(filled(r2 + 1, c) for c in range(max(1, c1 - 1), c2 + 2)):
            r2, grown = r2 + 1, True
        if c1 > 1 and any(filled(r, c1 - 1) for r in range(max(1, r1 - 1), r2 + 2)):
            c1, grown = c1 - 1, True
        if any(filled(r, c2 + 1) for r in range(max(1, r1 - 1), r2 + 2)):
            c2, grown = c2 + 1, True
        if not grown:
            return _NzBox(ws, r1, c1, r2, c2)


_NzBox.Font = property(lambda self: _InvObj(Bold=_nz_rows_flag(self, self.ws.bold_rows, True, False)))
_NzBox.Interior = property(lambda self: _InvObj(ColorIndex=_nz_rows_flag(self, self.ws.filled_rows, 6, -4142)))
_NzBox.Borders = lambda self, kind: _InvObj(LineStyle=self.ws.borders.get(kind, -4142))
_NzBox.ColumnWidth = property(lambda self: self.ws.col_widths.get(self.c1, 8.43))
_NzBox.CurrentRegion = property(_nz_region)


def _looks_world():
    """見出しは太字・塗りあり・格子あり。金額は桁区切り、入会日は書式がまちまち。表の外（F9）に合計。"""
    ws, wb = _xm_world()
    ws.bold_rows, ws.filled_rows = {3}, {3}
    ws.borders = {k: 1 for k in (7, 8, 9, 10, 11, 12)}
    ws.col_widths = {1: 10.0, 2: 12.5, 3: 9.0, 4: 11.0}
    ws.fmt.update({(4, 3): "#,##0", (5, 3): "#,##0", (4, 4): "yyyy/m/d", (9, 6): "#,##0"})
    ws.v[(9, 6)] = 2000.0
    return ws, wb


def test_grade_structure_reads_what_cells_do_not_show():
    import vbam_agent as va
    ws, wb = _xm_world()
    cf = [_InvObj(AppliesTo=_InvObj(Address="$A$4:$A$7"), Type=8, DupeUnique=1),
          _InvObj(AppliesTo=_InvObj(Address="$C$4:$C$7"), Type=2, Formula1="=$C4>1000"),
          _InvObj(AppliesTo=_InvObj(Address="$C$4:$C$7"), Type=3),
          RuntimeError("壊れた条件"),
          _InvObj(AppliesTo=_InvObj(Address="$D$4:$D$7"), Type=99)]
    ws.cf = cf + [_InvObj(AppliesTo=_InvObj(Address=f"$E${r}"), Type=9) for r in (4, 5, 6, 7)]
    ws.n_dv, ws.charts = 3, [_InvObj(), _InvObj()]
    ws.AutoFilterMode = True
    ws.AutoFilter = _InvObj(Range=_InvObj(Address="$A$3:$F$7"),
                            Filters=_XmColl([_InvObj(On=False), _InvObj(On=True), _InvObj(On=True)]))
    ws.Shapes = _XmShapes([_XmPic("btn印刷", on_action="印刷する"), _XmPic("ロゴ"), _XmPic("矢印")])
    win = wb.Application.ActiveWindow
    win.SplitRow, win.SplitColumn = 3, 1
    assert va._grade_structure(ws) == [
        "条件付き書式: 9 本　A4:A7（重複/一意＝重複に色） / C4:C7（数式 =$C4>1000） / C4:C7（カラースケール） / "
        "D4:D7（種類99） / E4（空白） / E5（空白） / E6（空白）　…",          # 8 本まで・読めない 1 本は飛ばす
        "入力規則: 3 セル",
        "枠固定: 上 3 行・左 1 列",
        "オートフィルタ: A3:F7（絞り込み中 2 列）",
        "グラフ: 2 個",
        "図形・ボタン: 3 個　マクロ付き（道具は消さない＝残っていて正しい）: btn印刷　マクロなし: ロゴ・矢印"]
    win.SplitRow = win.SplitColumn = 0
    assert va._grade_structure(wb.Sheets("集計")) == ["条件付き書式: 0 本", "入力規則: 0 セル"]
    assert va._grade_structure(_InvObj()) == ["入力規則: 0 セル"]                 # 何も読めなくても落ちない


def test_grade_looks_hands_borders_header_formats_and_outside_numbers():
    import vbam_agent as va
    ws, wb = _looks_world()
    assert va._grade_looks(ws) == [
        "罫線: A3:D5 に格子あり（外枠と内側の縦横すべて）",
        "見出し行 A3:D3: 太字・塗りあり",
        "列幅: 9.0〜12.5（### は材料に「'###' で読めないセル」の行が無ければ 0）",
        "表示形式（見出しの下・列ごと。カンマの有無はここが現物）: 会員番号=G/標準（例 0001）／氏名=G/標準（例 佐藤）"
        "／金額=#,##0（例 1,200）／入会日=まちまち（例 2026/4/1）",
        "表の外にある数値（離れた集計ブロック等。画面の文字と表示形式）: F9=2,000（#,##0）"]
    # 内側の線が無い・まちまち（None）は格子と言わない
    ws.borders.update({11: -4142, 12: None})
    assert va._grade_looks(ws)[0] == "罫線: A3:D5 は格子になっていない（無い・まちまち: 縦の内側・横の内側）"
    # ピボットの上なら列ごとに読まない（ラベル行と本文で書式が違うのは正常）＝本文の現物を渡す
    ws.pivots = [_InvObj(Name="P集計", TableRange2=ws.Range("A3:D5"),
                         DataBodyRange=_InvObj(NumberFormat="#,##0", Cells=lambda r, c: _InvObj(Text="1,200")))]
    got = va._grade_looks(ws)
    assert "ここはピボット「P集計」です（見出し行・ラベル行と本文で表示形式が違うのは正常＝まちまちではありません）。" \
           "値の表示形式（本文の現物）: #,##0（例 1,200）" in got
    assert not any(s.startswith("表示形式（見出しの下") for s in got) and got[-1].startswith("表の外にある数値")
    # 見出しが日付の表（2026年1月…）は、見出しの表示形式も現物で渡す（9 状態中 7 状態で差し戻していた）
    xl = _UdXl()
    monthly = _XmSheet(xl, "月別", {"A1": "担当", "B1": _dt.datetime(2026, 1, 1), "C1": _dt.datetime(2026, 2, 1),
                                    "A2": "佐藤", "B2": 100.0, "C2": 200.0, "A3": "鈴木", "B3": 150.0, "C3": 250.0})
    monthly.fmt.update({(1, 2): "yyyy年m月", (1, 3): "yyyy年m月"})
    got = va._grade_looks(monthly)
    assert got[0] == "罫線: A1:C3 は格子になっていない（無い・まちまち: 左・上・下・右・縦の内側・横の内側）"
    assert got[1] == "見出し行 A1:C1: 太字でない・塗りなし"
    assert "先頭行に入っている数値・日付（見出しが日付の表・合計ブロック等。表示形式はここが現物）: " \
           "B1=2026/1/1（yyyy年m月）／C1=2026/2/1（yyyy年m月）" in got
    assert va._grade_looks(_InvObj()) == []


def test_looks_of_range_and_grade_materials_sections(monkeypatch):
    """tidy を当てた範囲の見た目 1 行。採点の材料には構造・見た目・仕上げた範囲・作った別シートが足される。"""
    import vbam_agent as va
    ws, wb = _looks_world()
    assert va._looks_of_range(ws, "A3:D5") == "A3:D5: 格子あり（外枠と内側の縦横すべて）／見出し行 太字・塗りあり"
    assert va._looks_of_range(_InvObj(), "A1") == ""
    monkeypatch.setattr(va, "job_clock_get", lambda: None)
    monkeypatch.setattr(va, "job_clock_set", lambda c: None)

    def fake_run(argv, wb_=None):
        if argv[1] == "壊れた":
            raise RuntimeError("読めない")
        return True, f"使用範囲（{argv[1]}）" + ("x" * 5000 if argv[1] == "新シート" else "")
    monkeypatch.setattr(va, "_run_cmd", fake_run)
    monkeypatch.setattr(vg, "_heavy_materials", lambda wb_: ["テーブル（ブック全体）: 0 本"])
    monkeypatch.setattr(va, "_TIDY_RANGES", {"名簿": ["A3:D5"]})
    mat = va._grade_materials("名簿", wb, extra=["新シート", "壊れた"])
    assert mat.startswith("使用範囲（名簿）\n【このシートの構造（セルには写らない現物）】\n条件付き書式: 0 本")
    assert "図形・ボタン: 1 個　マクロ付き（道具は消さない＝残っていて正しい）: btn印刷" in mat
    assert "【見た目（道具が書式を読んだ現物）】\n罫線: A3:D5 に格子あり" in mat
    assert "【この走行で仕上げた範囲（道具が当てた tidy の現物）】\nA3:D5: 格子あり" in mat
    assert "【このループで作った別のシート「新シート」の現物】" in mat and mat.endswith("\n…（以下略）")
    assert "壊れた" not in mat and "【ブックにある重い物" not in mat          # 0 本だけの一覧は足さない
    assert va._grade_materials("名簿", None) == "（材料を読めませんでした）"
    monkeypatch.setattr(va, "_run_cmd", lambda argv, wb_=None: 1 / 0)
    assert va._grade_materials("名簿", wb) == "（材料を読めませんでした: division by zero）"


def test_delete_created_removes_each_kind_in_reverse_and_reports_what_it_could_not(monkeypatch):
    """undo は作った物を逆順で消す（2026-09-10 夜）。見つからない物は黙って飛ばし、消せない物は理由つきで返す。

    それまで通っていたのはシートと名前の 2 種類だけ。表・ピボット・スライサー・グラフ・クエリ・メジャー・
    計算フィールドは、消し方を間違えても --undo が「戻しました」と言うだけで誰も気づかない。
    """
    import vbam_agent as va
    import vbam_heavy
    ws, wb = _xm_world()
    xl = _UdXl()
    log = []

    def gone(tag):
        return lambda: log.append(tag)
    ws.ListObjects = _XmColl([_InvObj(Name="T会員", Delete=gone("table"))])
    pt = _InvObj(Name="P集計", TableRange2=_InvObj(Clear=gone("pivot")),
                 CalculatedFields=lambda: _XmColl([_InvObj(Name="単価", Delete=gone("calc_field"))]))
    monkeypatch.setattr(vbam_heavy, "_find_pivot", lambda wb_, name: (ws, pt) if name == "P集計" else (None, None))
    wb.SlicerCaches = _XmColl([_InvObj(Slicers=[_InvObj(Name="地域", Delete=gone("slicer"))])])
    _XmCharts(ws).Add(0, 0, 10, 10)                                           # グラフ 1
    wb.Queries = _XmColl([_InvObj(Name="売上取込", Delete=gone("query"))])
    wb.Names = _XmColl([_InvObj(Name="単価", Delete=gone("name"))])
    wb.Model = _InvObj(ModelMeasures=[_InvObj(Name="合計", Delete=gone("measure"))])
    items = [{'kind': 'table', 'sheet': '名簿', 'name': 'T会員'},
             {'kind': 'pivot', 'sheet': '集計', 'name': 'P集計'},
             {'kind': 'calc_field', 'pivot': 'P集計', 'name': '単価'},
             {'kind': 'slicer', 'name': '地域'},
             {'kind': 'chart', 'sheet': '名簿', 'name': 'グラフ 1'},
             {'kind': 'query', 'name': '売上取込'},
             {'kind': 'sheet', 'name': '隠し'},
             {'kind': 'name', 'name': '単価'},
             {'kind': 'measure', 'table': '売上', 'name': '合計'},
             {'kind': 'sheet', 'name': '無いシート'},
             {'kind': 'table', 'name': '無いテーブル'},
             {'kind': 'pivot', 'name': '無いピボット'},
             {'kind': 'slicer', 'name': '無いスライサー'},
             {'kind': '知らない種類', 'name': '謎'}]
    done, failed = va._delete_created(xl, wb, items)
    assert done == ["measure 合計", "name 単価", "sheet 隠し", "query 売上取込", "chart グラフ 1", "slicer 地域",
                    "calc_field 単価", "pivot P集計", "table T会員"]                   # 作った順の逆
    assert log == ["measure", "name", "query", "slicer", "calc_field", "pivot", "table"]
    assert failed == ["sheet 無いシート: シート 無いシート がありません"]
    assert [s.Name for s in wb.Sheets] == ["名簿", "集計"] and ws.charts == [] and xl.DisplayAlerts is True
    # 最後の 1 枚のシートは消さない（Excel はシート 0 枚のブックを作れない）
    one = _XmBook("一枚.xlsx", "C:\\work", [_XmSheet(xl, "唯一", {"A1": 1})])
    assert va._delete_created(xl, one, [{'kind': 'sheet', 'name': '唯一'}]) == ([], ["sheet 唯一（最後の 1 枚は消せない）"])
    assert [s.Name for s in one.Sheets] == ["唯一"] and va._delete_created(xl, wb, []) == ([], [])


class _Boom:
    """持っている属性は返し、持っていない属性を読むと落ちる（COM の「読めない」を真似る）。"""

    def __init__(self, **kw):
        self.__dict__.update(kw)

    def __getattr__(self, name):
        raise RuntimeError(f"{name} を読めません")


def test_heavy_materials_lists_every_kind_and_survives_unreadable_parts():
    """重い物の一覧（2026-09-10 夜）。読めない所は飛ばすか「読めない」と書き、一覧は 20 本で切る。

    0 本だけの一覧は採点係に渡さない（_HEAVY_NONZERO_RE）＝「0 本」の書き方が変わると、正しく作ったピボットを
    採点係が「値の貼り付けでは」と疑う日に戻る（2026-09-09）。書き方ごとここで押さえる。
    """
    import vbam_agent as va
    xl = _UdXl()
    odd = _XmSheet(xl, "壊れ", {"A1": 1})
    odd.ListObjects = _XmColl([_InvObj(Name="T列不明", Range=_InvObj(Address="$C$1:$D$9"))])   # 列名が読めない
    many = _XmSheet(xl, "多い", {"A1": 1})
    many.ListObjects = _XmColl([_InvObj(Name=f"T{i:02d}", Range=_InvObj(Address="$A$1:$B$2"),
                                        ListColumns=[_InvObj(Name="列")]) for i in range(25)])
    shukei = _XmSheet(xl, "集計", {"A1": 1})
    shukei.pivots = [_Boom(Name="P部分", TableRange2=_InvObj(Address="$A$1:$C$5"),
                           PivotFields=lambda: _XmColl([_InvObj(Name="地域", Orientation=1), _Boom(Name="壊れた項目"),
                                                        _InvObj(Name="売上", Orientation=4)])),
                     _Boom(Name="P壊れ")]
    wb = _XmBook("台帳.xlsx", "C:\\work", [odd, many, shukei, _InvObj(Name="読めない")])
    wb.SlicerCaches, wb.Queries = _Boom(), _Boom()
    wb.Model = _Boom(ModelTables=_XmColl([_InvObj(Name="売上")]))
    got = va._heavy_materials(wb)
    assert got[0] == "テーブル（ブック全体）: 26 本（table の手で読む・並べ替え・絞り込み・列追加）"
    assert got[1] == "  [壊れ] T列不明 範囲=C1:D9 列=" and got[20] == "  [多い] T18 範囲=A1:B2 列=列"
    assert got[21:] == ["ピボット（ブック全体）: 2 本（pivot_field / pivot_calc の手で。中身は pivot_calc get_data）",
                        "  [集計] P部分 出力=A1:C5 元=? 行=地域 列=- フィルタ=- 値=売上",
                        "  [集計] P壊れ（詳細を読めない: PivotFields を読めません）",
                        "データモデル: テーブル 1 本 売上  メジャー 0 本"]              # スライサーとクエリは読めずに飛ばす
    assert va._HEAVY_NONZERO_RE.search("\n".join(got))
    empty = va._heavy_materials(_XmBook("空.xlsx", "C:\\work", [_XmSheet(xl, "S", {"A1": 1})]))
    assert empty == ["テーブル（ブック全体）: 0 本（table create で作れる）", "ピボット（ブック全体）: 0 本（pivot create で作れる）",
                     "スライサー: 0 個 ", "パワークエリ: 0 本 ", "データモデル: テーブル 0 本   メジャー 0 本"]
    assert va._HEAVY_NONZERO_RE.search("\n".join(empty)) is None                  # 0 本だけ＝採点係に足さない


def test_find_replace_no_longer_refuses_a_replacement_that_contains_the_search_20260911():
    """断って 1 往復を捨てさせる代わりに、CLI 側が既に置き換え後のセルを飛ばす（二重当ては起きない）。"""
    import vbam_agent as va
    _label, toks = va._action_to_tokens({"op": "find_replace", "range": "H6:H30", "search": "受注済",
                                         "replace": "受注済み", "overwrite": True}, "顧客")
    assert "--whole" not in toks and toks[-3:] == ["受注済", "受注済み", "H6:H30"]
    # validation の候補は source でも受ける（AI は 5 回とも source と書いた）
    _label, toks = va._action_to_tokens({"op": "validation", "range": "H6:H31", "type": "list",
                                         "source": "相談中,受注済み,失注"}, "顧客")
    assert toks[-2:] == ["--list", "相談中,受注済み,失注"]
    # 材料の --full（エージェント・採点係は 200 行まで全体を見せる）
    import vba_manager as vm
    ns = vm.build_parser().parse_args(["materials", "--full"])
    assert ns.full is True
    assert "keys は省くのが既定" in va.RULES


def test_delete_hands_are_moved_last_and_rows_go_bottom_first_20260911():
    import vbam_agent as va
    old = dict(va._APPROVAL_CTX)
    va._APPROVAL_CTX.update(text="重複行は削除してよい。値は上書きしてよい", answer=False)
    try:
        acts = [{"op": "row_delete", "at": 17, "count": 2, "overwrite": True},
                {"op": "row_delete", "at": 20, "count": 2, "overwrite": True},
                {"op": "write_cells", "cells": {"D6": "03-1111-2222"}, "overwrite": True},
                {"op": "tidy", "ranges": ["A5:I30"]}]
        p = va._prevalidate(acts, "顧客", None)
        assert p == {}
        assert [a["op"] for a in acts] == ["write_cells", "row_delete", "row_delete", "tidy"]
        assert [a["at"] for a in acts[1:3]] == [20, 17]      # 下の行から
        assert acts[3]["ranges"] == ["A5"]                  # 消した後の tidy は表全体（末尾は道具が決める）
        acts = [{"op": "dedupe", "range": "A5:I56", "keys": ["会社名"], "overwrite": True},
                {"op": "normalize", "range": "E6:E56", "rules": ["lower"], "overwrite": True},
                {"op": "tidy", "ranges": ["A5:I31"]}]
        assert va._prevalidate(acts, "顧客", None) == {}
        assert [a["op"] for a in acts] == ["normalize", "dedupe", "tidy"]
        assert "dedupe" in va._ALLOWED_OPS and "dedupe" in va._DIRECT_OPS and va._needs_approval({"op": "dedupe"})
        assert "dedupe" in va.RULES and "lower" in va.RULES and '"plain": true' in va.RULES
        _label, toks = va._action_to_tokens({"op": "format", "range": "A6:I31", "plain": True, "bg": "none"}, "顧客")
        assert "--plain" in toks and "--bg=none" in toks
        # 飛び飛びの行番号は道具が連番のかたまりに分けて、下から消す（はねて 2 往復捨てていた・2026-09-11 午後）
        acts = [{"op": "row_delete", "rows": [17, 18, 20, 26, 28, 29], "overwrite": True},
                {"op": "tidy", "ranges": ["A5:I30"]}]
        assert va._prevalidate(acts, "顧客", None) == {}
        assert [(a["at"], a["count"]) for a in acts if a["op"] == "row_delete"] == [(28, 2), (26, 1), (20, 1), (17, 2)]
        assert acts[-1]["op"] == "tidy"
    finally:
        va._APPROVAL_CTX.clear()
        va._APPROVAL_CTX.update(old)


def test_ranges_are_stretched_to_data_end_when_the_same_reply_deletes_rows_20260911():
    """消す手は道具が最後に回すのに、AI は消した後の番地（A6:A35）で normalize を書いた＝36〜56 行が直らず不合格。"""
    import types
    import vbam_agent as va
    vals = [["会社名", "金額"]] + [[f"社{i}", i] for i in range(51)]          # 5 行目見出し・6〜56 行に値
    ur = types.SimpleNamespace(Row=5, Value=tuple(tuple(r) for r in vals),      # COM はタプルのタプル
                               Rows=types.SimpleNamespace(Count=52), Columns=types.SimpleNamespace(Count=2))
    wb = types.SimpleNamespace(Sheets=lambda name: types.SimpleNamespace(UsedRange=ur))
    acts = [{"op": "normalize", "range": "A6:A35", "rules": ["trim"], "overwrite": True},
            {"op": "format", "range": "A6:B300", "bold": True},
            {"op": "row_delete", "at": 40, "count": 3, "overwrite": True}]
    va._clip_to_data(acts, "顧客", wb)
    assert acts[0]["range"] == "A6:A56"          # 伸ばした（消す前の番地で全行に当てる）
    assert acts[1]["range"] == "A6:B56"          # はみ出しは今までどおり切る
    # 消す手が無い返事では、手前で終わる範囲はそのまま（AI が範囲を絞った意図を潰さない）
    acts = [{"op": "normalize", "range": "A6:A35", "rules": ["trim"], "overwrite": True}]
    va._clip_to_data(acts, "顧客", wb)
    assert acts[0]["range"] == "A6:A35"


def test_regex_dollar_backrefs_and_validation_aliases_20260911():
    """$1 の書き方で "$1-$2-$3" の文字が 13 セルに入った・validation の候補を values で書かれて 1 往復を捨てた。"""
    import vbam_hands as vh
    import vbam_agent as va
    _rng, rules, *_ = vh._normalize_spec({"op": "normalize", "range": "D6:D56", "overwrite": True,
                                          "rules": [{"regex": {"search": r"^(\d{2})(\d{4})(\d{4})$",
                                                               "replace": "$1-$2-$3"}}]})
    assert vh._normalize_value("0322221111", rules)[0] == "03-2222-1111"
    import pytest
    with pytest.raises(ValueError, match="replace"):
        vh._normalize_spec({"op": "normalize", "range": "D6:D9", "overwrite": True,
                            "rules": [{"regex": {"search": r"^(\d+)$", "replace": "$2"}}]})
    for key in ("list", "values", "items", "options", "choices"):
        _label, toks = va._action_to_tokens({"op": "validation", "range": "H6:H31", key: ["相談中", "受注済み", "失注"]}, "顧客")
        assert toks[-2:] == ["--list", "相談中,受注済み,失注"]


def test_validation_candidates_are_picked_up_whatever_the_name_20260911():
    """AI は候補を list・source・value と往復ごとに違う名前で書いた（7 回で 3 通り）。名前を問わず拾う。"""
    import vbam_agent as va
    for act in ({"value": "相談中,受注済み,失注"}, {"candidates": ["相談中", "受注済み", "失注"]},
                {"allowed": "相談中、受注済み、失注"}):
        _label, toks = va._action_to_tokens(dict({"op": "validation", "range": "H6:H31", "type": "list"}, **act), "顧客")
        assert toks[-2:] == ["--list", "相談中,受注済み,失注"], act
    import pytest
    with pytest.raises(ValueError, match="list"):
        va._action_to_tokens({"op": "validation", "range": "H6:H31", "type": "list"}, "顧客")


def test_row_delete_accepts_a_row_range_20260911():
    """Haiku は row_delete を "range": "18:18" と書き、「row の at は行番号（数）です」ではねていた（揺らしの直す候補 1）。
    行の範囲の書き方は at と count に畳む。消す順（下から）を決める _row_span も同じ読み方をする。"""
    import vbam_agent as va
    for act, want in (({"range": "18:18"}, ["18", "1"]), ({"range": "18:20"}, ["18", "3"]),
                      ({"range": "A18:I20"}, ["18", "3"]), ({"at": "$18:$19"}, ["18", "2"]),
                      ({"at": "18", "count": 2}, ["18", "2"]), ({"rows": [18, 19]}, ["18", "2"])):
        _label, toks = va._action_to_tokens(dict({"op": "row_delete", "overwrite": True}, **act), "顧客")
        assert toks[:4] == ["row", "delete"] + want, act
        assert va._row_span(act) == (int(want[0]), int(want[1])), act
    import pytest
    with pytest.raises(ValueError, match="行番号"):
        va._action_to_tokens({"op": "row_delete", "range": "A:A", "overwrite": True}, "顧客")
