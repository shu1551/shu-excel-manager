# -*- coding: utf-8 -*-
"""vbam_structure.py — ブックの「構造だけ」を値を出さずに 1 枚にまとめる（2026-10-01・structure）。

オフィス田中のワークシート診断ツール（http://officetanaka.net/library/）に学んだ。田中さんの主張は
「AI に要るのはブックの値ではなく構造」。値を含まない診断結果なら、機密を含む実務のブックでも AI に出せる。

  structure [excel_file] [--out 出力.md] [--sheet 名]

出すもの（田中さんの診断ツールの 17 シートの項目を覆う）:
  概要・シート・表示形式（ユーザー定義）・スタイル・名前・数式（同じ形は R1C1 で 1 行に）・関数・参照（他シート・他ブック）・
  条件付き書式・入力規則・テーブル・ピボット・オブジェクト・クエリ（M 言語）・データ接続・ハイパーリンク・VBA（点検つき）
出さないもの: **セルの値**（式の結果も。型とエラーの種類だけ）・コメントと図形の文字（字数だけ）・ブックのプロパティの中身
  （作成者・会社・タイトルは「あり」だけ）・接続文字列のパスワード。
出すと決めたもの（2026-10-01 shu「伏せなくていいでしょ」）: 式の中の文字列リテラル・M 言語・入力規則のリストと
  メッセージ・テーブルの列名とピボットのフィールド名（式の構造化参照に出る＝表の形そのもの）・リンク先。

純 Python の部品（COM なし・pytest で見る）: fml_functions・fml_sheet_refs・vba_checklist・numfmts_from_styles_xml・
mask_connection・value_kind。
"""
import os
import re
import html
import zipfile

# ----------------------------------------------------------------
# 1. 純 Python の部品
# ----------------------------------------------------------------

_STR_LIT_RE = re.compile(r'"(?:[^"]|"")*"')
_FUNC_RE = re.compile(r'(?<![A-Za-z0-9_.\]])((?:_xlfn\.|_xlws\.|_xludf\.)*[A-Za-z][A-Za-z0-9_.]*)\(')
# 'シート名'!  または  シート名!  （[ブック]シート! も含む）
_SHEET_REF_RE = re.compile(r"(?:'((?:[^']|'')+)'|((?:\[[^\]]+\])?[^\s'!=(),:+\-*/&^<>;{}\"\[\]]+))!")
_VOLATILE = {'NOW', 'TODAY', 'RAND', 'RANDBETWEEN', 'OFFSET', 'INDIRECT', 'CELL', 'INFO', 'RANDARRAY'}

# COM の Value がエラーを返すときの int（0x800A0000 + Excel のエラー番号）
XL_ERRORS = {-2146826288: '#NULL!', -2146826281: '#DIV/0!', -2146826273: '#VALUE!', -2146826265: '#REF!',
             -2146826259: '#NAME?', -2146826252: '#NUM!', -2146826246: '#N/A', -2146826245: '#GETTING_DATA',
             -2146826243: '#SPILL!', -2146826242: '#CONNECT!', -2146826241: '#BLOCKED!', -2146826240: '#UNKNOWN!',
             -2146826239: '#FIELD!', -2146826238: '#CALC!'}


def _strip_strings(formula):
    """式の中の文字列リテラルを "" に潰す（関数名・参照を拾うときに中の語を拾わない）。"""
    return _STR_LIT_RE.sub('""', formula or '')


def fml_functions(formula):
    """式で使っている関数名の集合（大文字・_xlfn. などの接頭辞を外す）。"""
    out = set()
    for m in _FUNC_RE.finditer(_strip_strings(formula)):
        name = re.sub(r'^(?:_xlfn\.|_xlws\.|_xludf\.)+', '', m.group(1)).upper()
        if name and not name.startswith('_XL'):
            out.add(name)
    return out


def fml_sheet_refs(formula):
    """式が参照している他シート・他ブック → {'sheets': {シート名}, 'books': {'[ブック]シート' か 'パス[ブック]シート'}}。"""
    sheets, books = set(), set()
    for m in _SHEET_REF_RE.finditer(_strip_strings(formula)):
        name = (m.group(1) or m.group(2) or '').replace("''", "'")
        if not name:
            continue
        if '[' in name and ']' in name:
            books.add(name)
        else:
            sheets.add(name)
    return {'sheets': sheets, 'books': books}


def value_kind(v):
    """式の結果の「型」だけ（値は出さない）。数・日付・文字・論理・空・エラー名。"""
    if v is None:
        return '空'
    if isinstance(v, bool):
        return '論理'
    if isinstance(v, int) and v in XL_ERRORS:
        return XL_ERRORS[v]
    if isinstance(v, (int, float)):
        return '数'
    if isinstance(v, str):
        return '空' if v == '' else '文字'
    if hasattr(v, 'year'):
        return '日付'
    return '他'


def numfmts_from_styles_xml(xml):
    """xlsx の styles.xml の <numFmt> → [(番号, 書式)]。164 以降＝ユーザー定義（VBA からは一覧で取れない）。"""
    out = set()
    for tag in re.findall(r'<(?:\w+:)?numFmt\b[^>]*>', xml or ''):
        i = re.search(r'\bnumFmtId="(\d+)"', tag)
        f = re.search(r'\bformatCode="([^"]*)"', tag)
        if i and f:
            out.add((int(i.group(1)), html.unescape(f.group(1))))
    return sorted(out)


def mask_connection(s):
    """接続文字列のパスワードを伏せる（Password=・Pwd=）。"""
    return re.sub(r'(?i)\b(password|pwd)\s*=\s*("[^"]*"|[^;]*)', r'\1=***', s or '')


# VBA の点検（田中さんの診断ツール [VBA CheckList] の 14 規則。レベル 0=留意・1=好ましくない・2=互換性）
_VBA_RULES = [
    (0, 'CreateObject', re.compile(r'\bCreateObject\s*\(', re.I)),
    (1, 'Selection', re.compile(r'(?<![.\w])Selection\b', re.I)),
    (1, 'Range の中で文字列結合', re.compile(r'\bRange\s*\([^)]*&', re.I)),
    (1, 'GoTo（On Error を除く）', re.compile(r'^(?!.*\bOn\s+Error\b).*\bGoTo\b', re.I)),
    (2, 'FileSearch（2007 で廃止）', re.compile(r'\bFileSearch\b', re.I)),
    (2, 'Cells.Count（2007 以降は CountLarge）', re.compile(r'\bCells\.Count\b(?!Large)', re.I)),
    (2, '古いブック形式で保存（Excel 95/97 形式）', re.compile(r'\bSaveAs\b.*(?:\bxlExcel9795\b|FileFormat\s*:=\s*43\b)', re.I)),
    (2, 'ChartObjects', re.compile(r'\bChartObjects\b', re.I)),
    (2, 'Shapes', re.compile(r'\bShapes\b', re.I)),
    (2, 'CommandBars（2007 以降はリボン）', re.compile(r'\bCommandBars\b', re.I)),
    (2, 'PivotCaches.Add（2007 以降は Create）', re.compile(r'\bPivotCaches\s*(?:\(\s*\))?\.Add\b', re.I)),
]
_AUTO_RE = re.compile(r'^\s*(?:Public\s+|Private\s+)?Sub\s+(Auto_Open|Auto_Close)\b', re.I)
_PROC_RE = re.compile(r'^\s*(?:(Public|Private|Friend)\s+)?(?:Static\s+)?'
                      r'(Sub|Function|Property\s+(?:Get|Let|Set))\s+([^\W\d]\w*)', re.I)
_END_RE = re.compile(r'^\s*End\s+(Sub|Function|Property)\b', re.I)
_DECLARE_RE = re.compile(r'^\s*(?:(?:Public|Private)\s+)?Declare\s+(?:PtrSafe\s+)?(?:Sub|Function)\s+(\w+)', re.I)


def _code_only(line):
    """コメントを外し、文字列リテラルを潰した 1 行（点検の誤検知を防ぐ）。"""
    inq, cut = False, len(line)
    for i, ch in enumerate(line):
        if ch == '"':
            inq = not inq
        elif ch == "'" and not inq:
            cut = i
            break
    code = line[:cut]
    if re.match(r'^\s*Rem\b', code, re.I):
        return ''
    return _STR_LIT_RE.sub('""', code)


def vba_procs(code):
    """モジュールのコード → (宣言部の情報, [手続き])。手続き = {name, kind, scope, lines, hits: [(レベル, 規則, 行)]}。"""
    lines = (code or '').splitlines()
    procs, cur = [], None
    head = {'option_explicit': False, 'declares': []}
    for i, raw in enumerate(lines, start=1):
        line = _code_only(raw)
        if cur is None:
            m = _PROC_RE.match(line)
            if m:
                kind = re.sub(r'\s+', ' ', m.group(2)).title()
                cur = {'name': m.group(3), 'kind': kind, 'scope': (m.group(1) or 'Public').title(),
                       'start': i, 'hits': []}
                if _AUTO_RE.match(line):
                    cur['hits'].append((0, '古い自動実行（Auto_Open/Auto_Close）', i))
                continue
            if re.match(r'^\s*Option\s+Explicit\b', line, re.I):
                head['option_explicit'] = True
            d = _DECLARE_RE.match(line)
            if d:
                head['declares'].append(d.group(1))
            continue
        if _END_RE.match(line):
            cur['lines'] = i - cur['start'] + 1
            procs.append(cur)
            cur = None
            continue
        for lvl, label, rx in _VBA_RULES:
            if rx.search(line):
                cur['hits'].append((lvl, label, i))
    if cur is not None:
        cur['lines'] = len(lines) - cur['start'] + 1
        procs.append(cur)
    return head, procs


# ---- 値なしの切り替え（materials・seiri・shelf-run の返事から値を外し、値を返す手を止める）----
#   中身を外に出せない職場で、AI に渡るのを構造だけにする。棚のマクロ（Excel の中で動く）はそのまま値を見て直す。
#   置き場は鍵の金庫と同じ %LOCALAPPDATA%\vba-manager。環境変数 EXCEL_MANAGER_NO_VALUES=1/0 が設定より先。
_NV_FILE = os.path.join(os.environ.get('LOCALAPPDATA') or os.path.expanduser('~'), 'vba-manager', 'no_values.flag')
# 値（セルの中身・画面の文字）を AI に返す手。値なしの間は止める
VALUE_COMMANDS = {'read-range', 'read-selection', 'snapshot', 'snapshot-diff', 'trace', 'screenshot', 'diagnose',
                  'audit', 'find', 'style-map', 'build-sheet'}


def no_values_mode():
    """値なしの切り替えが入っているか。"""
    env = str(os.environ.get('EXCEL_MANAGER_NO_VALUES', '')).strip().lower()
    if env in ('1', 'on', 'true', 'yes'):
        return True
    if env in ('0', 'off', 'false', 'no'):
        return False
    return os.path.isfile(_NV_FILE)


def no_values_refusal(name):
    """値なしの間に値を返す手を撃たれたときの返事（None＝撃ってよい）。"""
    if name in VALUE_COMMANDS and no_values_mode():
        return (f"値なしの切り替え中のため、値を返す手（{name}）は止めています。"
                "表の形は structure（ブック全体）か materials（1 シート）で見られます。"
                "表を直すのは seiri（棚のマクロが Excel の中で値を見て直す）。切り替えを外すのは no-values off")
    return None


def cmd_no_values(args):
    """値なしの切り替え: no-values [on|off]（引数なしは今の状態）。"""
    word = (args.posargs[0].lower() if getattr(args, 'posargs', None) else '')
    if word in ('on', 'オン', '入れる'):
        os.makedirs(os.path.dirname(_NV_FILE), exist_ok=True)
        with open(_NV_FILE, 'w', encoding='utf-8') as f:
            f.write('on\n')
    elif word in ('off', 'オフ', '外す'):
        try:
            os.remove(_NV_FILE)
        except FileNotFoundError:
            pass
    elif word:
        print("使い方: no-values [on|off]")
        return False
    on = no_values_mode()
    env = os.environ.get('EXCEL_MANAGER_NO_VALUES')
    print(f"値なし: {'入っています' if on else '外れています'}"
          + (f"（環境変数 EXCEL_MANAGER_NO_VALUES={env} が設定より先）" if env else ''))
    if on:
        print("  materials・seiri・shelf-run は値を出さず（型・番地・式だけ）、値を返す手（"
              + '・'.join(sorted(VALUE_COMMANDS)) + "）は止めます。ブック全体の形は structure")
    return True


def type_grid(rng, max_rows=200, max_cols=30):
    """範囲を「型」の格子にする（値の代わり）。数・文・日・論・誤（#名）・式→型。空は空けておく。"""
    try:
        fa, fv = _grid(rng.Formula), _grid(rng.Value)
    except Exception:
        try:
            fa = None
            fv = _grid(rng.Value)
        except Exception:
            return "(読めません)"
    r0, c0 = int(rng.Row), int(rng.Column)
    short = {'数': '数', '文字': '文', '日付': '日', '論理': '論', '空': '', '他': '?'}
    rows = []
    for i, row in enumerate(fv[:max_rows]):
        cells = []
        for j, v in enumerate(row[:max_cols]):
            k = value_kind(v)
            k = short.get(k, k)
            f = fa[i][j] if fa and i < len(fa) and j < len(fa[i]) else None
            if isinstance(f, str) and f.startswith('='):
                k = '式→' + (k or '空')
            cells.append(k)
        rows.append(cells)
    if not rows:
        return "(空の範囲です)"
    ncols = max(len(r) for r in rows)
    w = [max([len(_col(c0 + j))] + [len(r[j]) * 2 if j < len(r) else 0 for r in rows]) for j in range(ncols)]
    rw = len(str(r0 + len(rows) - 1))

    def line(cells, label):
        out = [label.rjust(rw)]
        for j in range(ncols):
            s = cells[j] if j < len(cells) else ''
            out.append(s + ' ' * max(0, w[j] - len(s) * (2 if not s.isascii() else 1)))
        return ' | '.join(out)
    out = [line([_col(c0 + j) for j in range(ncols)], ''), '-' * (rw + sum(w) + 3 * ncols)]
    out += [line(r, str(r0 + i)) for i, r in enumerate(rows)]
    return '\n'.join(out)


# ----------------------------------------------------------------
# 2. COM で集める
# ----------------------------------------------------------------

_VIS = {-1: '表示', 0: '非表示', 2: '完全に非表示（VBA からしか戻せない）'}
_CF_TYPE = {1: 'セルの値', 2: '数式', 3: 'カラースケール', 4: 'データバー', 5: '上位/下位', 6: 'アイコン',
            8: '一意/重複', 9: '文字列', 10: '空白', 11: '日付', 12: '平均より上/下', 13: '空白以外',
            16: 'エラー', 17: 'エラー以外'}
_CF_OP = {1: '間', 2: '間以外', 3: '等しい', 4: '等しくない', 5: 'より大きい', 6: 'より小さい',
          7: '以上', 8: '以下'}
_DV_TYPE = {0: 'すべての値', 1: '整数', 2: '小数', 3: 'リスト', 4: '日付', 5: '時刻', 6: '文字列（長さ）',
            7: 'ユーザー設定'}
_DV_ALERT = {1: '停止', 2: '注意', 3: '情報'}
_DV_IME = {0: 'コントロールなし', 1: 'オン', 2: 'オフ', 3: '無効', 4: 'ひらがな', 5: '全角カタカナ',
           6: '半角カタカナ', 7: '全角英数', 8: '半角英数'}
_SHAPE_TYPE = {1: 'オートシェイプ', 2: '吹き出し', 3: 'グラフ', 4: 'コメント', 5: 'フリーフォーム', 6: 'グループ',
               7: '埋め込み OLE', 8: 'フォームコントロール', 9: '線', 10: 'リンク OLE', 11: 'リンク図',
               12: 'ActiveX コントロール', 13: '図', 14: 'プレースホルダー', 15: 'ワードアート', 16: 'メディア',
               17: 'テキストボックス', 24: 'グラフィック', 28: 'グラフィック'}
_CONN_TYPE = {1: 'OLEDB', 2: 'ODBC', 3: 'XML', 4: 'テキスト', 5: 'Web', 6: 'データフィード', 7: 'データモデル',
              8: 'ワークシート', 9: '接続なし'}
_COMP_TYPE = {1: '標準モジュール', 2: 'クラス', 3: 'フォーム', 100: 'シート/ブック'}
# 新しいブックに最初から付いている参照（OLE Automation・Office）＝「標準でない」に数えない
_STD_REFS = {'{00020430-0000-0000-C000-000000000046}', '{2DF8D04C-5BFA-101B-BDE5-00AA0044DE52}',
             '{0D452EE1-E08F-101A-852E-02608C4D0BB4}'}   # Forms 2.0 はユーザーフォームを足すと Excel が自分で付ける
_FMT_NAME = {51: 'xlsx', 52: 'xlsm', 50: 'xlsb', 56: 'xls', 55: 'xlam'}


def _addr(rng):
    try:
        return str(rng.Address).replace('$', '')
    except Exception:
        return '?'


def _grid(v):
    if v is None or not isinstance(v, tuple):
        return [[v]]
    return [list(r) if isinstance(r, tuple) else [r] for r in v]


def _col(n):
    s = ''
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _count_mixed(rng, attr, cap=200000):
    """rng の各セルの attr（Locked・FormulaHidden）が True の数。まとめて聞いて混在の列だけ数える。"""
    try:
        v = getattr(rng, attr)
    except Exception:
        return None
    n = int(rng.CountLarge)
    if v is True:
        return n
    if v is False:
        return 0
    if n > cap:
        return None
    total = 0
    for j in range(1, int(rng.Columns.Count) + 1):      # 列ごとにまとめて聞き、混在の列だけ 1 セルずつ
        col = rng.Columns(j)
        try:
            cv = getattr(col, attr)
        except Exception:
            cv = None
        if cv is True:
            total += int(col.Cells.CountLarge)
        elif cv is None:
            for c in col.Cells:
                try:
                    if getattr(c, attr):
                        total += 1
                except Exception:
                    pass
    return total


def _formulas_from_file(path, sheet, cells):
    """保存済みの xlsx/xlsm から、指定セルの式を読む → {(行, 列): '=式'}（保護で隠した式用。値は読まない）。"""
    if not path or not cells or not path.lower().endswith(('.xlsx', '.xlsm', '.xltx', '.xltm')):
        return {}
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    except Exception:
        return {}
    out = {}
    try:
        if sheet not in wb.sheetnames:
            return {}
        ws = wb[sheet]
        rows = [r for r, _c in cells]
        cols = [c for _r, c in cells]
        want = set(cells)
        for row in ws.iter_rows(min_row=min(rows), max_row=max(rows), min_col=min(cols), max_col=max(cols)):
            for cell in row:
                key = (getattr(cell, 'row', None), getattr(cell, 'column', None))
                v = getattr(cell, 'value', None)
                if key in want and isinstance(v, str) and v.startswith('='):
                    out[key] = v
    except Exception:
        pass
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return out


def _sheet_part(ws, sheet_info, lines_fml, funcs_by_sheet, refs_rows, totals):
    """1 枚のワークシートの構造を集める。"""
    name = ws.Name
    info = {'name': name, 'type': 'ワークシート', 'vis': _VIS.get(int(ws.Visible), str(ws.Visible)),
            'protect': bool(ws.ProtectContents), 'used': '', 'unlocked': 0, 'hidden_fml': 0,
            'merged': 0, 'cf': 0, 'dv': 0}
    try:
        ur = ws.UsedRange
        nr, nc = int(ur.Rows.Count), int(ur.Columns.Count)
        info['used'] = f"{_addr(ur)}（{nr}行×{nc}列）"
    except Exception:
        ur, nr, nc = None, 0, 0
    # 結合
    try:
        from vbam_view import _merged_areas_in_range
        areas, skipped = _merged_areas_in_range(ur) if ur is not None else ([], None)
        info['merged'] = len(areas) if not skipped else f"未走査（{skipped}セル）"
    except Exception:
        pass
    # 式（一括で A1・R1C1・値の型を読む）
    fcells, locked = [], []
    bulk_ok = False
    if ur is not None and nr * nc <= 1_000_000:
        try:
            fa, fr = _grid(ur.Formula), _grid(ur.FormulaR1C1)
            fv = _grid(ur.Value)
            r0, c0 = int(ur.Row), int(ur.Column)
            for i, row in enumerate(fa):
                for j, f in enumerate(row):
                    if isinstance(f, str) and f.startswith('='):
                        rc = fr[i][j] if i < len(fr) and j < len(fr[i]) else f
                        v = fv[i][j] if i < len(fv) and j < len(fv[i]) else None
                        fcells.append((r0 + i, c0 + j, f, str(rc), v))
            bulk_ok = True
        except Exception:
            pass
    if ur is not None and not bulk_ok and info['protect'] and nr * nc <= 200_000:
        # 保護シートで「数式を表示しない」にした式があると、Excel は範囲の式の読み出しを丸ごと拒む（2026-10-01 試しブック）。
        # 1 セルずつ聞き、読めない式は保存済みのファイルから読む（保護は外さない。田中さんの診断ツールも中身を出していた）
        for c in ur.Cells:
            try:
                if not c.HasFormula:
                    continue
            except Exception:
                continue
            try:
                fcells.append((int(c.Row), int(c.Column), str(c.Formula), str(c.FormulaR1C1), c.Value))
            except Exception:
                locked.append((int(c.Row), int(c.Column), c.Value))
        bulk_ok = True
    if ur is not None and not bulk_ok:
        try:
            for c in ur.SpecialCells(-4123):
                fcells.append((int(c.Row), int(c.Column), str(c.Formula), str(c.FormulaR1C1), c.Value))
                if len(fcells) >= 20000:
                    break
        except Exception:
            pass
    from_file = set()
    if locked:
        got = _formulas_from_file(totals.get('path'), name, [(r, c) for r, c, _v in locked])
        for r, c, v in locked:
            f = got.get((r, c))
            if f:
                fcells.append((r, c, f, f, v))
                from_file.add((r, c))
    locked_fml = len(locked) - len(from_file)
    totals['formulas'] += len(fcells)
    # 配列数式（CSE）とスピル：式のある所だけ聞く（R1C1 の型ごとに代表の 1 セル）
    pats = {}
    for r, c, a1, rc, v in fcells:
        p = pats.setdefault(rc, {'n': 0, 'first': (r, c), 'last': (r, c), 'a1': a1, 'kinds': {}})
        p['n'] += 1
        p['last'] = (r, c)
        k = value_kind(v)
        p['kinds'][k] = p['kinds'].get(k, 0) + 1
        funcs_by_sheet.setdefault(name, set()).update(fml_functions(a1))
        refs = fml_sheet_refs(a1)
        for s in sorted(refs['sheets']):
            if s != name:
                refs_rows.add((name, s, ''))
        for b in sorted(refs['books']):
            refs_rows.add((name, '', b))
            totals['ext_cells'] += 1
    for rc, p in sorted(pats.items(), key=lambda kv: (kv[1]['first'][1], kv[1]['first'][0])):
        r, c = p['first']
        tags = []
        try:
            cell = ws.Cells(r, c)
            if cell.HasArray:
                if bool(getattr(cell, 'HasSpill', False)) is False:
                    tags.append('配列数式（CSE）')
            if getattr(cell, 'HasSpill', False):
                tags.append(f"スピル {_addr(cell.SpillingToRange)}")
        except Exception:
            pass
        if (r, c) in from_file:
            tags.append('保護で隠した式（保存済みのファイルから読んだ）')
        fn = fml_functions(p['a1'])
        vol = sorted(fn & _VOLATILE)
        if vol:
            tags.append('揮発関数 ' + '・'.join(vol))
        where = f"{_col(c)}{r}" if p['n'] == 1 else f"{_col(c)}{r}〜{_col(p['last'][1])}{p['last'][0]}"
        kinds = '・'.join(f"{k}{'' if n == 1 and p['n'] == 1 else ' ' + str(n)}" for k, n in p['kinds'].items())
        lines_fml.append(f"| {name} | {where} | {p['n']} | `{p['a1']}` | {kinds} | {'・'.join(tags)} |")
    # 保護シートのロック解除・非表示の式
    if ur is not None:                             # ロック解除は保護していないシートでも数える（入力欄の用意＝作り手の意図。田中さんも数える）
        n_locked = _count_mixed(ur, 'Locked')
        info['unlocked'] = int(ur.CountLarge) - n_locked if isinstance(n_locked, int) else '?'
    if info['protect'] and ur is not None:
        hid = len(locked)
        totals['formulas'] += locked_fml
        if locked_fml:
            lines_fml.append(f"| {name} | （{locked_fml} セル） | {locked_fml} | （保護で中身を隠した式＝未保存か xls のため読めません。保護は外していません） |  | 非表示の式 |")
        for r, c, *_ in fcells:
            if (r, c) in from_file:                # 上の len(locked) で数え済み
                continue
            try:
                if ws.Cells(r, c).FormulaHidden:
                    hid += 1
            except Exception:
                pass
        info['hidden_fml'] = hid
    sheet_info.append(info)
    return info


def _cf_rows(ws, out):
    n = 0
    try:
        fcs = ws.Cells.FormatConditions
        cnt = int(fcs.Count)
    except Exception:
        return 0
    for i in range(1, cnt + 1):
        try:
            fc = fcs(i)
        except Exception:
            continue
        n += 1
        row = {'sheet': ws.Name, 'addr': '?', 'type': '?', 'op': '', 'f1': '', 'f2': '', 'look': []}
        try:
            row['addr'] = _addr(fc.AppliesTo)
        except Exception:
            pass
        try:
            t = int(fc.Type)
            row['type'] = _CF_TYPE.get(t, str(t))
        except Exception:
            t = None
        try:
            row['op'] = _CF_OP.get(int(fc.Operator), '')
        except Exception:
            pass
        for k in ('Formula1', 'Formula2'):
            try:
                row['f' + k[-1]] = str(getattr(fc, k) or '')
            except Exception:
                pass
        try:
            if fc.NumberFormat:
                row['look'].append('表示形式')
        except Exception:
            pass
        try:
            if fc.Font.Bold is not None or fc.Font.Italic is not None or fc.Font.ColorIndex not in (None, -4142, -4105):
                row['look'].append('フォント')
        except Exception:
            pass
        try:
            if any(int(fc.Borders(b).LineStyle or -4142) not in (-4142, 0) for b in (1, 2, 3, 4)):
                row['look'].append('罫線')
        except Exception:
            pass
        try:
            if fc.Interior.ColorIndex not in (None, -4142, -4105):
                row['look'].append('塗り')
        except Exception:
            pass
        whole = bool(re.search(r'(^|[,:])\$?[A-Z]{1,3}:\$?[A-Z]{1,3}($|,)', row['addr'])) or \
            bool(re.search(r'(^|[,:])\$?\d+:\$?\d+($|,)', row['addr']))
        row['whole'] = whole
        out.append(row)
    return n


def _dv_rows(ws, out):
    try:
        cells = ws.Cells.SpecialCells(-4174)          # xlCellTypeAllValidation
    except Exception:
        return 0
    total = 0
    try:
        ur = ws.UsedRange
    except Exception:
        ur = None
    for area in cells.Areas:
        n = int(area.CountLarge)
        used = n
        if ur is not None:                       # 件数は使用範囲の中で数える（シートまるごとの規則で 171 億と出た・2026-10-01 仕事のブック）
            try:
                isect = ws.Application.Intersect(area, ur)
                used = int(isect.CountLarge) if isect is not None else 0
            except Exception:
                pass
        total += used
        row = {'sheet': ws.Name, 'addr': _addr(area), 'cells': n, 'used': used}
        try:
            v = area.Cells(1, 1).Validation
            row['type'] = _DV_TYPE.get(int(v.Type), str(v.Type))
            f1 = str(v.Formula1 or '')
            try:
                f2 = str(v.Formula2 or '')
            except Exception:
                f2 = ''
            row['formula'] = f1 + (f' 〜 {f2}' if f2 else '')
            for k in ('InputTitle', 'InputMessage', 'ErrorTitle', 'ErrorMessage'):
                try:
                    row[k] = str(getattr(v, k) or '')
                except Exception:
                    row[k] = ''
            try:
                row['style'] = _DV_ALERT.get(int(v.AlertStyle), '')
            except Exception:
                row['style'] = ''
            try:
                row['ime'] = _DV_IME.get(int(v.IMEMode), '')
            except Exception:
                row['ime'] = ''
        except Exception:
            row['type'] = '?'
        row['whole'] = n >= 1048576 or bool(re.search(r'^\$?[A-Z]{1,3}:\$?[A-Z]{1,3}$', row['addr']))
        # 種類「すべての値」でメッセージも日本語入力の指定も無い＝何もしない規則（コピーで残った残骸）
        row['empty'] = (row.get('type') == 'すべての値' and not any(row.get(k) for k in
                        ('InputTitle', 'InputMessage', 'ErrorTitle', 'ErrorMessage')) and row.get('ime') in ('', 'コントロールなし'))
        out.append(row)
    return total


def _shape_rows(ws, out, counts):
    def walk(shapes, parent):
        for shp in shapes:
            try:
                t = int(shp.Type)
            except Exception:
                t = -1
            row = {'sheet': ws.Name, 'parent': parent or '', 'name': '', 'type': _SHAPE_TYPE.get(t, str(t)),
                   'at': '', 'h': '', 'w': '', 'macro': '', 'link': '', 'chars': 0, 'series': [], 'hidden': False}
            try:
                row['name'] = shp.Name
            except Exception:
                pass
            try:
                row['at'] = _addr(shp.TopLeftCell)
                row['h'], row['w'] = round(float(shp.Height)), round(float(shp.Width))
            except Exception:
                pass
            try:
                row['hidden'] = not bool(shp.Visible)
            except Exception:
                pass
            try:
                oa = shp.OnAction
                if oa:
                    row['macro'] = str(oa)
            except Exception:
                pass
            try:
                row['link'] = str(shp.Hyperlink.Address or shp.Hyperlink.SubAddress or '')
            except Exception:
                pass
            try:
                from vbam_view import _shape_text
                row['chars'] = len(_shape_text(shp) or '')
            except Exception:
                pass
            if t == 3:
                try:
                    for s in shp.Chart.SeriesCollection():
                        row['series'].append(str(s.Formula))
                except Exception:
                    pass
            if t == 12:
                try:
                    row['type'] += f"（{shp.OLEFormat.progID}）"
                except Exception:
                    pass
            key = {3: 'chart', 13: 'picture', 17: 'textbox', 4: 'comment', 12: 'activex'}.get(t, 'other')
            counts[key] = counts.get(key, 0) + 1
            out.append(row)
            if t == 6:
                try:
                    walk(shp.GroupItems, row['name'])
                except Exception:
                    pass
    try:
        walk(ws.Shapes, None)
    except Exception:
        pass


def collect(wb, only_sheet=None):
    """ブックの構造を集めて dict にする（値は入れない）。"""
    d = {'book': wb.Name, 'sheets': [], 'fml': [], 'funcs': {}, 'refs': set(), 'cf': [], 'dv': [], 'tables': [],
         'pivots': [], 'shapes': [], 'shape_counts': {}, 'comments': [], 'links': [], 'names': [], 'styles': [],
         'numfmts': None, 'queries': [], 'conns': [], 'vba': [], 'ext_links': [], 'totals': {'formulas': 0, 'ext_cells': 0}}
    # 概要
    try:
        d['format'] = _FMT_NAME.get(int(wb.FileFormat), str(wb.FileFormat))
    except Exception:
        d['format'] = '?'
    try:
        full = str(wb.FullName)
        d['saved_path'] = full if os.path.isfile(full) else None
        d['size'] = os.path.getsize(full) if d['saved_path'] else None
    except Exception:
        d['saved_path'], d['size'] = None, None
    d['totals']['path'] = d['saved_path']
    props = []
    for k, label in (('Author', '作成者'), ('Last Author', '最終更新者'), ('Company', '会社'), ('Title', 'タイトル')):
        try:
            if str(wb.BuiltinDocumentProperties(k).Value or '').strip():
                props.append(label)
        except Exception:
            pass
    d['props'] = props
    # シート
    for sh in wb.Sheets:
        if only_sheet and sh.Name != only_sheet:
            continue
        try:
            is_ws = int(sh.Type) == -4167        # xlWorksheet
        except Exception:
            is_ws = False
        if not is_ws:
            d['sheets'].append({'name': sh.Name, 'type': 'グラフシート', 'vis': _VIS.get(int(sh.Visible), ''),
                                'protect': False, 'used': '', 'unlocked': 0, 'hidden_fml': 0, 'merged': 0,
                                'cf': 0, 'dv': 0})
            continue
        info = _sheet_part(sh, d['sheets'], d['fml'], d['funcs'], d['refs'], d['totals'])
        info['cf'] = _cf_rows(sh, d['cf'])
        info['dv'] = _dv_rows(sh, d['dv'])
        # テーブル
        try:
            for lo in sh.ListObjects:
                t = {'sheet': sh.Name, 'name': lo.Name, 'addr': _addr(lo.Range), 'cols': [], 'conn': '',
                     'totals': False}
                try:
                    t['cols'] = [str(c.Name) for c in lo.ListColumns]
                except Exception:
                    pass
                try:
                    t['totals'] = bool(lo.ShowTotals)
                except Exception:
                    pass
                try:
                    t['conn'] = str(lo.QueryTable.WorkbookConnection.Name)
                except Exception:
                    pass
                d['tables'].append(t)
        except Exception:
            pass
        # ピボット
        try:
            for pt in sh.PivotTables():
                p = {'sheet': sh.Name, 'name': pt.Name, 'addr': _addr(pt.TableRange2), 'src': '', 'conn': '',
                     'refresh': '', 'rows': [], 'cols': [], 'data': [], 'page': []}
                try:
                    p['src'] = str(pt.SourceData)
                except Exception:
                    pass
                try:
                    p['conn'] = str(pt.PivotCache().WorkbookConnection.Name)
                except Exception:
                    pass
                try:
                    p['refresh'] = '開くとき更新' if pt.PivotCache().RefreshOnFileOpen else ''
                except Exception:
                    pass
                for key, attr in (('rows', 'RowFields'), ('cols', 'ColumnFields'), ('data', 'DataFields'),
                                  ('page', 'PageFields')):
                    try:
                        p[key] = [str(f.Name) for f in getattr(pt, attr)]
                    except Exception:
                        pass
                d['pivots'].append(p)
        except Exception:
            pass
        # 図形・グラフ・コメント
        _shape_rows(sh, d['shapes'], d['shape_counts'])
        try:
            for cm in sh.Comments:
                d['comments'].append((sh.Name, _addr(cm.Parent), 'メモ'))
        except Exception:
            pass
        try:
            for cm in sh.CommentsThreaded:
                d['comments'].append((sh.Name, _addr(cm.Parent), 'スレッド コメント'))
        except Exception:
            pass
        # ハイパーリンク（表示の文字はセルの値なので出さない）
        try:
            for hl in sh.Hyperlinks:
                try:
                    at = _addr(hl.Range)
                except Exception:
                    at = '(図形)'
                d['links'].append((sh.Name, at, str(hl.Address or ''), str(hl.SubAddress or '')))
        except Exception:
            pass
    # 名前
    try:
        for nm in wb.Names:
            try:
                n = str(nm.Name)
                ref = str(nm.RefersTo)
            except Exception:
                continue
            scope = '(ブック)'
            try:
                par = nm.Parent
                if par.Name != wb.Name:
                    scope = par.Name
            except Exception:
                pass
            d['names'].append({'name': n, 'ref': ref, 'visible': bool(nm.Visible), 'scope': scope,
                               'error': '#REF!' in ref or '#NAME?' in ref,
                               # Excel が自分で作る名前（新しい関数・オートフィルター・クエリの読み込み先）は数えない
                               'system': bool(re.match(r'^(_xl|_FilterDatabase$|ExternalData_\d+$)', n.split('!')[-1]))})
    except Exception:
        pass
    # スタイル（ユーザー定義）
    try:
        for st in wb.Styles:
            try:
                if not st.BuiltIn:
                    d['styles'].append(str(st.NameLocal))
            except Exception:
                pass
    except Exception:
        pass
    # ユーザー定義の表示形式（VBA では一覧で取れない＝保存済みのファイルの styles.xml から）
    if d['saved_path'] and d['saved_path'].lower().endswith(('.xlsx', '.xlsm', '.xltx', '.xltm', '.xlam')):
        try:
            with zipfile.ZipFile(d['saved_path']) as z:
                xml = z.read('xl/styles.xml').decode('utf-8', 'replace')
            d['numfmts'] = [f for i, f in numfmts_from_styles_xml(xml) if i >= 164]
        except Exception:
            d['numfmts'] = None
    # クエリ・接続
    try:
        for q in wb.Queries:
            d['queries'].append({'name': str(q.Name), 'm': str(q.Formula)})
    except Exception:
        pass
    try:
        for c in wb.Connections:
            row = {'name': str(c.Name), 'type': '', 'conn': '', 'cmd': '', 'refresh': ''}
            try:
                row['type'] = _CONN_TYPE.get(int(c.Type), str(c.Type))
            except Exception:
                pass
            for sub in ('OLEDBConnection', 'ODBCConnection'):
                try:
                    o = getattr(c, sub)
                    row['conn'] = mask_connection(str(o.Connection))
                    try:
                        ct = o.CommandText
                        row['cmd'] = ' '.join(ct) if isinstance(ct, (tuple, list)) else str(ct or '')
                    except Exception:
                        pass
                    try:
                        row['refresh'] = '開くとき更新' if o.RefreshOnFileOpen else ''
                    except Exception:
                        pass
                    break
                except Exception:
                    continue
            d['conns'].append(row)
    except Exception:
        pass
    # 他ブックへのリンク
    try:
        ls = wb.LinkSources(1)                  # xlExcelLinks
        if ls:
            d['ext_links'] = [str(x) for x in ls]
    except Exception:
        pass
    # VBA
    try:
        for comp in wb.VBProject.VBComponents:
            cm = comp.CodeModule
            n = int(cm.CountOfLines)
            code = cm.Lines(1, n) if n else ''
            head, procs = vba_procs(code)
            if not n and not procs:
                continue
            d['vba'].append({'module': comp.Name, 'type': _COMP_TYPE.get(int(comp.Type), str(comp.Type)),
                             'lines': n, 'head': head, 'procs': procs})
        refs = []
        for r in wb.VBProject.References:
            try:
                if not r.BuiltIn and str(r.Guid).upper() not in _STD_REFS:
                    refs.append(str(r.Description or r.Name) + (' ★参照不可' if r.IsBroken else ''))
            except Exception:
                pass
        d['vba_refs'] = refs
    except Exception as e:
        d['vba_error'] = str(e)
    return d


# ----------------------------------------------------------------
# 3. 書き出し（Markdown）
# ----------------------------------------------------------------

def _cell(s):
    return str(s).replace('|', '｜').replace('\r', '').replace('\n', '⏎')


def points(d):
    """集めた構造 → 気をつける所 [(見つけたこと, なぜ, 確かめる所)]（純 Python）。"""
    out = []
    sc = d.get('shape_counts') or {}
    sheets = d.get('sheets') or []
    for s in sheets:
        if s['vis'].startswith('完全に'):
            out.append((f"シート「{s['name']}」が完全に非表示",
                        "シートの見出しを右クリックした「再表示」の一覧にも出ない＝あることに誰も気づかない",
                        "開発 → Visual Basic → プロジェクトのシート → プロパティの Visible"))
    hid = [s['name'] for s in sheets if s['vis'] == '非表示']
    if hid:
        out.append((f"非表示のシート {len(hid)} 枚（{'・'.join(hid[:5])}{'…' if len(hid) > 5 else ''}）",
                    "見えない所の式や値が、見えている表の結果を左右していることがある",
                    "シートの見出しを右クリック → 再表示"))
    for s in sheets:
        if s['protect'] and s.get('unlocked') == 0:
            out.append((f"シート「{s['name']}」は保護されていて、ロック解除のセルが 0",
                        "入力欄を空けずに保護している＝どのセルにも入力できない",
                        "ホーム → 書式 → セルのロック／校閲 → シート保護の解除"))
        if s['protect'] and s.get('hidden_fml'):
            out.append((f"シート「{s['name']}」に保護で隠した式 {s['hidden_fml']} セル",
                        "数式バーに式が出ない＝計算の中身を確かめられない（引き継ぎの時に困る）",
                        "セルの書式設定 → 保護 → 表示しない"))
    nm = [n for n in d.get('names') or [] if not n['system']]
    for n in nm:
        if n['error']:
            out.append((f"名前「{n['name']}」の参照先が壊れている（{n['ref']}）",
                        "この名前を使う式・入力規則・印刷範囲が #REF! になる",
                        "数式 → 名前の管理"))
    hidden_names = [n['name'] for n in nm if not n['visible']]
    if hidden_names:
        out.append((f"非表示の名前 {len(hidden_names)} 個（{'・'.join(hidden_names[:5])}{'…' if len(hidden_names) > 5 else ''}）",
                    "名前の管理にも出ない。他のブックから写ってきた古い名前が多い",
                    "VBA（Name.Visible）でしか見えない"))
    for r in d.get('cf') or []:
        if r['whole']:
            out.append((f"条件付き書式が列・行まるごと（{r['sheet']}!{r['addr']}）",
                        "100 万行ぶん判定する＝スクロールや入力が重くなる元",
                        "ホーム → 条件付き書式 → ルールの管理"))
    cf_by = {}
    for r in d.get('cf') or []:
        cf_by[r['sheet']] = cf_by.get(r['sheet'], 0) + 1
    for s, n in cf_by.items():
        if n >= 50:
            out.append((f"シート「{s}」の条件付き書式が {n} ルール",
                        "コピーや行の挿入で分かれて増えたもの＝重さの元",
                        "ホーム → 条件付き書式 → ルールの管理（このワークシート）"))
    empties = [r for r in d.get('dv') or [] if r.get('empty')]
    if empties:
        out.append((f"中身の無い入力規則 {len(empties)} か所（{'・'.join(r['sheet'] + '!' + r['addr'] for r in empties[:4])}"
                    f"{'…' if len(empties) > 4 else ''}）",
                    "何も制限しない規則がコピーや貼り付けで残っている＝ファイルを重くするだけ",
                    "範囲を選んで データ → データの入力規則 → すべてクリア"))
    for r in d.get('dv') or []:
        if r['whole'] and not r.get('empty'):
            out.append((f"入力規則が列まるごと（{r['sheet']}!{r['addr']}・{r['cells']:,} セル）",
                        "使っていない行まで規則が付く＝ファイルが重くなり、どこまでが表か分からなくなる",
                        "データ → データの入力規則"))
    if d.get('ext_links') or (d.get('totals') or {}).get('ext_cells'):
        out.append((f"他のブックを参照している（{len(d.get('ext_links') or [])} ファイル）",
                    "相手のファイルが動く・名前が変わると #REF! になり、開くたびに更新を聞かれる",
                    "データ → リンクの編集"))
    vol = sorted({v.strip() for ln in d.get('fml') or [] for t in ln.split('|') if '揮発関数 ' in t
                  for v in t.split('揮発関数 ')[1].split('・')[:1] + t.split('揮発関数 ')[1].split('・')[1:]})
    if vol:
        out.append((f"揮発関数（{'・'.join(vol)}）",
                    "開くたびに再計算される＝何も変えていなくても閉じるときに保存を聞かれ、大きい表では遅くなる",
                    "数式 → 数式の表示"))
    ghosts = [s for s in d.get('shapes') or [] if s['type'] not in ('コメント',)
              and (s.get('hidden') or (isinstance(s['h'], int) and isinstance(s['w'], int) and (s['h'] <= 1 or s['w'] <= 1)))]
    if ghosts:
        out.append((f"見えないオブジェクト {len(ghosts)} 個（{'・'.join(s['sheet'] + '!' + s['name'] for s in ghosts[:5])}"
                    f"{'…' if len(ghosts) > 5 else ''}）",
                    "非表示か大きさ 0 の図形。行のコピーで増え続け、ファイルを重くする",
                    "ホーム → 検索と選択 → オブジェクトの選択と表示"))
    if sc.get('activex'):
        out.append((f"ActiveX コントロール {sc['activex']} 個",
                    "Office の更新や別の PC で動かなくなることがある",
                    "開発 → デザインモード（フォームコントロールに替える）"))
    for c in d.get('conns') or []:
        if c.get('refresh') and not str(c.get('conn', '')).startswith(('OLEDB;Provider=Microsoft.Mashup', 'WORKSHEET')):
            out.append((f"開くときに更新する接続「{c['name']}」",
                        "接続先に届かない PC では開くたびに待たされ、エラーが出る",
                        "データ → クエリと接続 → プロパティ"))
    refs = d.get('vba_refs') or []
    if refs:
        broken = [r for r in refs if '参照不可' in r]
        out.append((f"標準でない参照設定 {len(refs)} 本" + (f"（うち参照不可 {len(broken)}）" if broken else ''),
                    "その部品が入っていない PC では、関係ないマクロまでコンパイルエラーで止まる",
                    "Visual Basic → ツール → 参照設定"))
    for m in d.get('vba') or []:
        for p in m['procs']:
            for lvl, label, _ln in p['hits']:
                if label.startswith('古い自動実行'):
                    out.append((f"{m['module']}.{p['name']} が古い自動実行",
                                "Auto_Open／Auto_Close は古い形の自動実行＝別のブックから開いたときなどに動かないことがある",
                                "ThisWorkbook の Workbook_Open に移す"))
                elif label.startswith(('FileSearch', 'Cells.Count', '古いブック形式')):     # CommandBars 等は今も動く
                    out.append((f"{m['module']}.{p['name']}: {label}",
                                "今の Excel では動かないか、意図と違う形で動く",
                                f"{m['module']} の {_ln} 行目"))
    # 同じ手続きの同じ規則は行をまとめる（古い形式の SaveAs が 2 行あると 2 本並んでいた）
    merged = {}
    for what, why, where in out:
        m = re.match(r'^(.+) の (\d+) 行目$', where or '')
        key = (what, why, m.group(1) if m else where)
        if key in merged and m:
            merged[key].append(m.group(2))
        else:
            merged.setdefault(key, [m.group(2)] if m else [])
    return [(w, y, (f"{h} の {'・'.join(ln)} 行目" if ln else h)) for (w, y, h), ln in merged.items()]


def render(d):
    L = []
    add = L.append
    T = d['totals']
    sc = d['shape_counts']
    hidden_ws = [s for s in d['sheets'] if s['vis'] != '表示']
    add(f"# ブックの構造: {d['book']}")
    add("（値は出していません。式・名前・書式・規則などの「形」だけ。オフィス田中のワークシート診断ツールの項目に合わせています）")
    add("")
    add("## 概要")
    add(f"- 形式: {d.get('format', '?')}" + (f"・{d['size']:,} バイト" if d.get('size') else '・未保存（保存してから撃つと表示形式も読めます）'))
    if d['props']:
        add(f"- プロパティに書かれているもの: {'・'.join(d['props'])}（中身は出しません）")
    add(f"- シート {len(d['sheets'])} 枚（非表示 {len(hidden_ws)}）・保護 {sum(1 for s in d['sheets'] if s['protect'])} 枚")
    add(f"- 数式 {T['formulas']} セル・他ブックを参照する式 {T['ext_cells']} セル・名前 {len([n for n in d['names'] if not n['system']])}"
        f"・条件付き書式 {len(d['cf'])} ルール・入力規則 {sum(r.get('used', r['cells']) for r in d['dv']):,} セル（使用範囲の中"
        + (f"・範囲の全体では {sum(r['cells'] for r in d['dv']):,}" if sum(r['cells'] for r in d['dv']) != sum(r.get('used', r['cells']) for r in d['dv']) else '') + "）")
    add(f"- テーブル {len(d['tables'])}・ピボット {len(d['pivots'])}・クエリ {len(d['queries'])}・接続 {len(d['conns'])}"
        f"・ハイパーリンク {len(d['links'])}・コメント {len(d['comments'])} セル")
    add(f"- グラフ {sc.get('chart', 0)}・図 {sc.get('picture', 0)}・テキストボックス {sc.get('textbox', 0)}"
        f"・ActiveX {sc.get('activex', 0)}・図形ほか {sc.get('other', 0)}")
    nmod = sum(1 for m in d['vba'] if m['procs'])
    add(f"- マクロ: {'あり（' + str(nmod) + ' モジュール）' if nmod else 'なし'}"
        + (f"・標準でない参照設定 {len(d.get('vba_refs') or [])}" if d.get('vba_refs') else ''))
    # 気をつける所（道具が先に言う）。田中さんの診断ツールの「診断ポイント」に学び、見つけたことに「なぜ」と
    # 「Excel のどこで確かめるか」を添える（2026-10-01 shu「取り入れられるところは取り入れたほうがいい」）
    pts = points(d)
    if pts:
        add("")
        add("### 気をつける所（見つけたこと → なぜ → Excel のどこで確かめるか）")
        for what, why, where in pts:
            add(f"- **{what}** → {why}" + (f"（確かめる所: {where}）" if where else ''))    # シート
    add("")
    add("## シート")
    add("| シート | 種類 | 使用範囲 | 表示 | 保護 | ロック解除 | 非表示の式 | 結合 | 条件付き書式 | 入力規則 |")
    add("|---|---|---|---|---|---|---|---|---|---|")
    for s in d['sheets']:
        add(f"| {_cell(s['name'])} | {s['type']} | {s['used']} | {s['vis']} | {'保護' if s['protect'] else ''} | "
            f"{s['unlocked'] or ''} | {s['hidden_fml'] if s['protect'] else ''} | {s['merged'] or ''} | "
            f"{s['cf'] or ''} | {format(s['dv'], ',') if s['dv'] else ''} |")
    # 表示形式・スタイル
    add("")
    add("## ユーザー定義の表示形式")
    if d['numfmts'] is None:
        add("（保存された xlsx/xlsm から読みます。未保存か xls/xlsb のため読めません）")
    elif d['numfmts']:
        for f in d['numfmts']:
            add(f"- `{f}`")
    else:
        add("なし")
    add("")
    add("## ユーザー定義のスタイル")
    add('・'.join(d['styles']) if d['styles'] else "なし")
    # 名前
    add("")
    add("## 名前")
    names = [n for n in d['names'] if not n['system']]
    sysn = [n for n in d['names'] if n['system']]
    if names:
        add("| 名前 | 参照 | 表示 | 範囲 | 壊れ |")
        add("|---|---|---|---|---|")
        for n in names:
            add(f"| {_cell(n['name'])} | `{_cell(n['ref'])}` | {'' if n['visible'] else '非表示'} | {_cell(n['scope'])} | "
                f"{'★' if n['error'] else ''} |")
    else:
        add("なし")
    if sysn:
        add(f"（Excel が自分で作る隠れた名前 {len(sysn)} 個は省略: {'・'.join(n['name'] for n in sysn[:5])}{'…' if len(sysn) > 5 else ''}）")
    # 数式
    add("")
    add("## 数式（同じ形の式は R1C1 で 1 行にまとめています。値は出さず、結果の型だけ）")
    if d['fml']:
        add("| シート | 番地 | 個数 | 式（代表の A1） | 結果の型 | 印 |")
        add("|---|---|---|---|---|---|")
        L.extend(d['fml'][:400])
        if len(d['fml']) > 400:
            add(f"…ほか {len(d['fml']) - 400} 型")
    else:
        add("なし")
    add("")
    add("## 関数（シートごと）")
    if d['funcs']:
        for s, fs in d['funcs'].items():
            add(f"- {s}: {'・'.join(sorted(fs))}")
    else:
        add("なし")
    add("")
    add("## 参照（他シート・他ブック）")
    if d['refs'] or d['ext_links']:
        for s, to, book in sorted(d['refs']):
            add(f"- {s} → {to or book}{'（他ブック）' if book else ''}")
        for x in d['ext_links']:
            add(f"- リンク元のファイル: {x}")
    else:
        add("なし")
    # 条件付き書式・入力規則
    add("")
    add("## 条件付き書式")
    if d['cf']:
        add("| シート | 範囲 | 種類 | 演算子 | 式1 | 式2 | 書式 |")
        add("|---|---|---|---|---|---|---|")
        for r in d['cf']:
            add(f"| {r['sheet']} | {r['addr']}{' ★まるごと' if r['whole'] else ''} | {r['type']} | {r['op']} | "
                f"`{_cell(r['f1'])}` | {('`' + _cell(r['f2']) + '`') if r['f2'] else ''} | {'・'.join(r['look'])} |")
    else:
        add("なし")
    add("")
    add("## 入力規則")
    empties = [r for r in d['dv'] if r.get('empty')]
    rules = [r for r in d['dv'] if not r.get('empty')]
    if rules:
        add("| シート | 範囲 | セル数 | 種類 | 式・リスト | 入力時タイトル／メッセージ | エラー時（スタイル・タイトル／メッセージ） | 日本語入力 |")
        add("|---|---|---|---|---|---|---|---|")
        for r in rules:
            add(f"| {r['sheet']} | {r['addr']} | {r.get('used', r['cells']):,} | {r.get('type', '')} | `{_cell(r.get('formula', ''))}` | "
                f"{_cell(r.get('InputTitle', ''))}／{_cell(r.get('InputMessage', ''))} | "
                f"{r.get('style', '')}・{_cell(r.get('ErrorTitle', ''))}／{_cell(r.get('ErrorMessage', ''))} | {r.get('ime', '')} |")
    if empties:
        add("中身の無い規則（種類「すべての値」でメッセージも指定も無い＝コピーで残った残骸）: "
            + "・".join(f"{r['sheet']}!{r['addr']}" for r in empties[:20]) + ("…" if len(empties) > 20 else ""))
    if not d['dv']:
        add("なし")
    # テーブル・ピボット
    add("")
    add("## テーブル（列名は表の形として出します）")
    if d['tables']:
        for t in d['tables']:
            add(f"- {t['sheet']}!{t['addr']} **{t['name']}**: 列 {'・'.join(t['cols'])}"
                + ('・集計行あり' if t['totals'] else '') + (f"・接続 {t['conn']}" if t['conn'] else ''))
    else:
        add("なし")
    add("")
    add("## ピボットテーブル")
    if d['pivots']:
        for p in d['pivots']:
            add(f"- {p['sheet']}!{p['addr']} **{p['name']}**: 元 {p['src'] or p['conn']}"
                + (f"・行 {'・'.join(p['rows'])}" if p['rows'] else '') + (f"・列 {'・'.join(p['cols'])}" if p['cols'] else '')
                + (f"・値 {'・'.join(p['data'])}" if p['data'] else '') + (f"・フィルター {'・'.join(p['page'])}" if p['page'] else '')
                + (f"・{p['refresh']}" if p['refresh'] else ''))
    else:
        add("なし")
    # オブジェクト
    add("")
    add("## オブジェクト（図形・グラフ・コメント。中の文字は字数だけ）")
    if d['shapes']:
        add("| シート | 親グループ | 名前 | 種類 | 位置 | 高さ×幅 | マクロ | リンク | 文字 |")
        add("|---|---|---|---|---|---|---|---|---|")
        for s in d['shapes']:
            add(f"| {s['sheet']} | {_cell(s['parent'])} | {_cell(s['name'])} | {s['type']} | {s['at']} | {s['h']}×{s['w']} | "
                f"{_cell(s['macro'])} | {_cell(s['link'])} | {str(s['chars']) + ' 字' if s['chars'] else ''} |")
            for f in s['series']:
                add(f"|  |  |  | 系列 |  |  |  |  | `{_cell(f)}` |")
    else:
        add("なし")
    if d['comments']:
        add("")
        add("コメントのあるセル: " + "・".join(f"{s}!{a}" + ('' if k == 'メモ' else f'（{k}）') for s, a, k in d['comments']))
    # クエリ・接続・リンク
    add("")
    add("## クエリ（Power Query の M 言語）")
    if d['queries']:
        for q in d['queries']:
            where = [f"{t['sheet']}!{t['addr']}" for t in d['tables'] if t['conn'] in (f"クエリ - {q['name']}", f"Query - {q['name']}")]
            add(f"### {q['name']}" + (f"（読み込み先 {'・'.join(where)}）" if where else '（接続のみ・データモデル）'))
            add("```")
            add(q['m'])
            add("```")
    else:
        add("なし")
    add("")
    add("## データ接続")
    if d['conns']:
        for c in d['conns']:
            add(f"- **{c['name']}**（{c['type']}）{('・' + c['refresh']) if c['refresh'] else ''}: `{_cell(c['conn'])}`"
                + (f"・コマンド `{_cell(c['cmd'])}`" if c['cmd'] else ''))
    else:
        add("なし")
    add("")
    add("## ハイパーリンク（表示の文字はセルの値なので出しません）")
    if d['links']:
        for s, a, addr, sub in d['links']:
            add(f"- {s}!{a} → {addr}{('#' + sub) if sub else ''}")
    else:
        add("なし")
    # VBA
    add("")
    add("## VBA")
    if d.get('vba_error'):
        add(f"（読めません: {d['vba_error']}。「VBA プロジェクト オブジェクト モデルへのアクセスを信頼する」が要ります）")
    elif d['vba']:
        add("| モジュール | 種類 | 行数 | Option Explicit | API 宣言 | 手続き（種類・範囲・行数） |")
        add("|---|---|---|---|---|---|")
        for m in d['vba']:
            procs = '・'.join(f"{p['name']}（{p['kind']}・{p['scope']}・{p['lines']}行）" for p in m['procs'])
            add(f"| {m['module']} | {m['type']} | {m['lines']} | {'あり' if m['head']['option_explicit'] else 'なし'} | "
                f"{'・'.join(m['head']['declares'])} | {_cell(procs)} |")
        hits = [(m['module'], p['name'], h) for m in d['vba'] for p in m['procs'] for h in p['hits']]
        if hits:
            add("")
            add("点検（田中さんの VBA CheckList の規則。0=留意・1=好ましくない・2=互換性）:")
            lab = {0: '留意', 1: '好ましくない', 2: '互換性'}
            seen = {}
            for mod, proc, (lvl, label, ln) in hits:
                seen.setdefault((mod, proc, lvl, label), []).append(ln)
            for (mod, proc, lvl, label), lns in seen.items():
                add(f"- {lab[lvl]}: {mod}.{proc} — {label}（{'・'.join(str(x) + ' 行目' for x in lns[:6])}{'…' if len(lns) > 6 else ''}）")
        if d.get('vba_refs'):
            add("")
            add("標準でない参照設定: " + '・'.join(d['vba_refs']))
        add("")
        add("（呼び出し関係は call-graph、手続きの中の流れは flow、詳しい点検は check で見られます）")
    else:
        add("なし")
    return '\n'.join(L) + '\n'


def cmd_structure(args):
    """構造だけの診断: structure [excel_file] [--out 出力.md] [--sheet 名]

    値を出さずに、ブックの形（式・名前・書式・規則・テーブル・ピボット・図形・クエリ・接続・リンク・VBA）を 1 枚にまとめる。
    機密を含むブックでも、これなら AI に渡せる（オフィス田中のワークシート診断ツールに学んだ・2026-10-01）。"""
    import vbam_core as vc
    target_file, rest = vc.parse_target_and_rest(args.posargs)
    xl, wb = vc.get_workbook(target_file, readonly=True)
    only = getattr(args, 'sheet_opt', None) or (rest[0] if rest else None)
    try:
        vc.cmd_progress_note(f"{wb.Name}: 構造を読んでいます")
    except Exception:
        pass
    d = collect(wb, only)
    md = render(d)
    out = getattr(args, 'out', None) or os.path.join(os.path.dirname(os.path.abspath(__file__)), '_last_structure.md')
    out = os.path.abspath(out)
    with open(out, 'w', encoding='utf-8') as f:
        f.write(md)
    print(md)
    print(f"（同じ内容を {out} に保存しました。値は入っていません）")
    return True
