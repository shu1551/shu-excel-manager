# -*- coding: utf-8 -*-
"""試験（vbam_exam・2026-09-11 夜）: 正解の表を持つお題の作り方と、値で突き合わせる採点。

Excel にも AI にも触らない（お題は Python で作り、終わった表は偽のセルで渡す）。本物は agent --exam で撃つ。
"""
import unicodedata

import pytest

import vbam_exam as ve


def _cells_from(exj, rows, title=None):
    """正解の形の行 → probe が返す形のセル（直し終わった理想の表）。"""
    out = []
    if title:
        out += [[{'v': title, 't': title}], [{'v': None, 't': ''}]]
    out.append([{'v': n, 't': n} for n, _t in exj['columns']])
    for r in rows:
        line = []
        for (_n, t), v in zip(exj['columns'], r):
            if isinstance(v, str) and v.startswith('D:'):
                y, m, d = v[2:].split('-')
                line.append({'v': v, 't': f"{int(y)}/{int(m)}/{int(d)}"})
            elif t in ('money', 'int'):
                line.append({'v': float(v), 't': f"{v:,}" if t == 'money' else str(v)})
            else:
                line.append({'v': v, 't': v})
        out.append(line)
    return out


def _perfect(ex):
    exj = ve._truth_json(ex)
    return exj, _cells_from(exj, exj['truth'], ex.get('title'))


def _dirty_cells(ex):
    """書いたままのお題（汚れたまま）→ セル。"""
    exj = ve._truth_json(ex)
    out = []
    if ex.get('title'):
        out += [[{'v': ex['title'], 't': ex['title']}], [{'v': None, 't': ''}]]
    out.append([{'v': n, 't': n} for n, _t in exj['columns']])
    for d in ex['dirty']:
        line = []
        for c in d['cells']:
            v = c['v']
            if hasattr(v, 'isoformat'):
                line.append({'v': f"D:{v.isoformat()}", 't': v.isoformat()})
            elif isinstance(v, int):
                line.append({'v': float(v), 't': str(v)})
            else:
                line.append({'v': v, 't': v})
        out.append(line)
    return exj, out


def test_same_seed_same_exam_and_other_seed_differs():
    a = ve.make_exams(7)
    b = ve.make_exams(7)
    c = ve.make_exams(8)
    assert [x['name'] for x in a] == ['顧客名簿', '売上明細', '会員名簿', '経費精算', '在庫表']
    assert [ve._truth_json(x)['truth'] for x in a] == [ve._truth_json(x)['truth'] for x in b]
    assert [ve._truth_json(x)['truth'] for x in a] != [ve._truth_json(x)['truth'] for x in c]


def test_only_filters_exams():
    assert [x['name'] for x in ve.make_exams(3, '在庫表,会員名簿')] == ['会員名簿', '在庫表']


def test_hankaku_kana_round_trips_by_nfkc():
    assert len(ve._KANA_F) == len(ve._KANA_H)
    for s in ('ダイスケ', 'ワタナベ ヨウコ', 'パーク'):
        h = ve._hankaku_kana(s)
        assert ve._HALF_KANA_RE.search(h)
        assert unicodedata.normalize('NFKC', h) == s


@pytest.mark.parametrize('seed', [1, 7, 42, 2026])
def test_perfect_table_passes(seed):
    for ex in ve.make_exams(seed):
        exj, cells = _perfect(ex)
        g = ve.grade(exj, cells, fonts={n: {'bold': False, 'italic': False, 'size': 11.0, 'color': 0.0, 'name': '游ゴシック'}
                                        for n, _t in exj['columns']})
        assert g['pass'], (ex['name'], g['problems'])
        assert g['rows_ok'] == g['truth_rows'] == len(ex['rows'])


@pytest.mark.parametrize('seed', [1, 7, 42])
def test_untouched_exam_fails_with_duplicates_and_formats(seed):
    for ex in ve.make_exams(seed):
        exj, cells = _dirty_cells(ex)
        g = ve.grade(exj, cells)
        assert not g['pass'], ex['name']
        assert g['format_errors'] > 0 or g['uneven'], ex['name']
        if ex.get('dups'):
            assert g['dup_left'] == exj['dup_rows'], (ex['name'], g['problems'])
        assert g['missing'] == 0, (ex['name'], g['problems'])      # 汚れていても値は全部ある


def test_deleting_a_legit_row_is_missing():
    ex = ve.make_exams(5, '経費精算')[0]
    exj = ve._truth_json(ex)
    rows = list(exj['truth'])
    twin = next(i for i in range(len(rows) - 1) if rows[i] == rows[i + 1])     # 本当に 2 回あった支払い
    del rows[twin]
    g = ve.grade(exj, _cells_from(exj, rows))
    assert not g['pass'] and g['missing'] == 1
    assert '消しすぎ' in g['problems'][0]


def test_same_company_other_contact_is_not_a_duplicate():
    ex = ve.make_exams(11, '顧客名簿')[0]
    exj = ve._truth_json(ex)
    names = [r[0] for r in exj['truth']]
    dup_company = next(n for n in names if names.count(n) > 1)
    rows = [r for i, r in enumerate(exj['truth']) if not (r[0] == dup_company and names.index(dup_company) != i)]
    g = ve.grade(exj, _cells_from(exj, rows))
    assert g['missing'] == 1


def test_wrong_value_is_paired_by_key():
    ex = ve.make_exams(5, '在庫表')[0]
    exj = ve._truth_json(ex)
    rows = [list(r) for r in exj['truth']]
    rows[3][3] = rows[3][3] + 1                 # 数量を 1 つ違える
    g = ve.grade(exj, _cells_from(exj, rows, ex['title']))
    assert not g['pass'] and g['value_errors'] == 1 and g['missing'] == 0 and g['extra'] == 0
    assert '数量' in [p for p in g['problems'] if '値の誤り' in p][0]


def test_truth_phones_do_not_collide_with_longer_area_codes():
    """018-505-… は 0185-05-… と読める＝正解の表のほうが間違う（9/11 夜の初回の試験で踏んだ）。"""
    codes = [a for a, _x, _y in ve._AREAS + ve._MOBILE]
    for seed in range(40):
        for ex in ve.make_exams(seed):
            for j, (_n, t) in enumerate(ex['columns']):
                if t != 'phone':
                    continue
                for r in ex['rows']:
                    own = r[j].split('-')[0]
                    flat = r[j].replace('-', '')
                    assert not any(len(x) > len(own) and flat.startswith(x) for x in codes), r[j]


def test_phone_split_and_number_checks():
    assert ve.strict('phone', {'v': '0185-12-3456'}, '0185-12-3456') is None
    assert ve.strict('phone', {'v': '0185123456'}, '0185-12-3456') is None
    assert '正しくは' in ve.strict('phone', {'v': '018-512-3456'}, '0185-12-3456')
    assert '0' in ve.strict('phone', {'v': 185123456.0}, '0185-12-3456')
    assert ve.strict('money', {'v': '¥12,000'}, 12000) == "数値になっていない（文字のまま）"
    assert ve.strict('date', {'v': '2026/9/1'}, 'D:2026-09-01') == "日付になっていない（文字のまま）"
    assert ve.strict('code', {'v': 'Ｐ－１００２'}, 'P-1002') == "英数字が全角のまま"
    assert ve.strict('kana', {'v': 'ｻﾄｳ ﾏｺﾄ'}, 'サトウ マコト') == "半角カナが残っている"


def test_idents_read_through_formats():
    assert ve.ident('date', {'v': 'R8.9.1'}) == ve.ident('date', {'v': 'D:2026-09-01'})
    assert ve.ident('date', {'v': '令和8年9月1日'}) == ('d', '2026-09-01')
    assert ve.ident('money', {'v': '１２０００'}) == ve.ident('money', {'v': 12000.0})
    assert ve.ident('int', {'v': '12個'}) == ('n', 12.0)
    assert ve.ident('company', {'v': '㈱アクアテック'}) == ve.ident('company', {'v': '株式会社アクアテック'})


def test_uneven_columns():
    assert ve.uneven('company', [{'v': '株式会社A'}, {'v': '(株)B'}])
    assert not ve.uneven('company', [{'v': '(株)A'}, {'v': '(株)B'}])
    assert ve.uneven('phone', [{'v': '03-1234-5678'}, {'v': '0312345678'}])
    assert ve.uneven('money', [{'v': 12000.0, 't': '12,000'}, {'v': 3400.0, 't': '3400'}])
    assert ve.uneven('person', [{'v': '佐藤 誠'}, {'v': '鈴木　恵'}])


def test_mixed_font_is_reported_only_when_asked():
    ex = ve.make_exams(3, '在庫表')[0]
    exj, cells = _perfect(ex)
    fonts = {n: {'bold': False, 'italic': False, 'size': 11.0, 'color': 0.0, 'name': '游ゴシック'} for n, _t in exj['columns']}
    fonts['品名']['bold'] = None
    g = ve.grade(exj, cells, fonts)
    assert not g['pass'] and any('太字' in u for u in g['uneven'])
    exj2 = dict(exj, font_dirt=False)
    assert ve.grade(exj2, cells, fonts)['pass']


def test_header_not_found():
    ex = ve.make_exams(3, '会員名簿')[0]
    exj = ve._truth_json(ex)
    g = ve.grade(exj, [[{'v': 'x'}]])
    assert not g['pass'] and '見出し' in g['problems'][0]


# ----------------------------------------------------------------
# 試験で見つけた道具の穴の直し（9/11 夜・種 7 の初回: 合格 1/5・嘘の合格 3）
# ----------------------------------------------------------------

def test_dup_pairs_do_not_match_other_deals_by_numbers_only():
    """商品・数量・金額が同じでも、日付も担当も違う取引は重複ではない（売上明細で 4 行消しすぎた）。"""
    import datetime as dt
    import vbam_hands as vh
    rows = [["日付", "担当", "商品", "数量", "単価", "金額"]]
    staff = ["木村 翔", "田中 彩", "山田 陽子", "小松 真理"]
    goods = ["ボールペン赤", "B5ノート", "電卓", "付箋", "封筒", "トナー", "電池", "ファイル"]
    for k in range(24):
        rows.append([dt.datetime(2026, 4 + k // 8, 1 + k), staff[k % 4], goods[k % 8], 1 + k, 100 + 10 * k, (1 + k) * (100 + 10 * k)])
    rows.append([dt.datetime(2026, 8, 9), "山田 陽子", rows[5][2], rows[5][3], rows[5][4], rows[5][5]])   # 別の日・別の人
    rows.append(["2026/4/1", " 木村 翔", " " + rows[1][2], str(rows[1][3]), rows[1][4], f"{rows[1][5]:,}"])  # 本当の重複（表記ゆれ）
    pairs = vh._dup_pairs(rows, 0)
    assert [(i, k) for i, k, _e in pairs] == [(26, 1)]


def test_normalize_postal_and_units():
    import vbam_hands as vh
    rules = [('postal', None)]
    for src, want in (('0100951', '010-0951'), ('〒010-0951', '010-0951'), ('０１０－０９５１', '010-0951'),
                      ('010 0951', '010-0951'), ('010-0951', '010-0951'), ('12345', '12345')):
        assert vh._normalize_value(src, rules)[0] == want, src
    assert vh._normalize_value('12個', [('number', None)])[0] == 12
    assert vh._normalize_value('１２０００円', [('hankaku', None), ('number', None)])[0] == 12000


def test_format_request_lets_format_only_normalize_through():
    """「書き方をそろえて」の依頼なら、書き方だけの規則は承認の言葉なしで通す（消す・regex は今までどおり止める）。"""
    import vbam_agent as va
    va._approval_set("経費精算の表の書き方をそろえてください。行は消さないでください。")
    try:
        assert not va._needs_approval({"op": "normalize", "range": "A2:A9", "rules": ["trim", "date"], "overwrite": True})
        assert va._needs_approval({"op": "normalize", "range": "A2:A9", "rules": [{"regex": {"search": "a", "replace": "b"}}],
                                   "overwrite": True})
        assert va._needs_approval({"op": "normalize", "range": "A2:A9", "rules": ["fill_down"], "overwrite": True})
        assert va._needs_approval({"op": "row_delete", "at": 3})
        va._approval_set("この表を見て")
        assert va._needs_approval({"op": "normalize", "range": "A2:A9", "rules": ["trim"], "overwrite": True})
    finally:
        va._approval_set("")


def test_column_misfits_catch_what_the_exam_found():
    import datetime as dt
    import vbam_grade as vg
    rows = [["会社名", "電話", "郵便番号", "品番", "氏名", "登録日", "金額"],
            ["株式会社A", "03-1234-5678", "010-0951", "P-1001", "佐藤 誠", dt.datetime(2026, 1, 1), 1000.0],
            ["(株)B", "090 0568 3136", "0100952", "Ｐ－１００２", "鈴木　恵", "2026/1/2", "１２０００"],
            ["株式会社C", "0312345678", "〒010-0953", "P-1003", "ｻﾄｳ", dt.datetime(2026, 1, 3), 3000.0]]
    req = ("会社名の表記ゆれ（株式会社と（株））・全角半角・前後の空白・電話番号・郵便番号・登録日・金額の書き方をそろえて。"
           "フリガナ（半角カナは全角に）")
    got = vg._column_misfits(rows, 0, range(7), 1, 1, req)
    text = "\n".join(got)
    for want in ("（株）と株式会社", "電話番号の書き方", "B3 '090 0568 3136'", "郵便番号が 000-0000", "C3 '0100952'",
                 "全角の英数字", "D3", "空白に全角と半角", "半角カナ", "文字のままの日付", "F3", "数値になっていない", "G3"):
        assert want in text, (want, text)
    clean = [rows[0], ["株式会社A", "03-1234-5678", "010-0951", "P-1001", "佐藤 誠", dt.datetime(2026, 1, 1), 1000.0],
             ["株式会社B", "090-0568-3136", "010-0952", "P-1002", "鈴木 恵", dt.datetime(2026, 1, 2), 12000.0]]
    assert vg._column_misfits(clean, 0, range(7), 1, 1, req) == []
    assert vg._column_misfits(rows, 0, range(7), 1, 1, "この表を見て") == []     # 依頼に出ない語は見ない


def test_normalize_corp_rule():
    import vbam_hands as vh
    for src, want in (('（株）アクアテック', '株式会社アクアテック'), ('㈱北斗精機', '株式会社北斗精機'),
                      ('株式会社 アクアテック', '株式会社アクアテック'), ('やまびこ観光 株式会社', 'やまびこ観光株式会社'),
                      ('有限会社　竿燈企画', '有限会社竿燈企画'), ('(有)白神ソフト', '有限会社白神ソフト'),
                      ('こまち製菓(株)', 'こまち製菓株式会社'), ('株式会社横手精密', '株式会社横手精密')):
        assert vh._normalize_value(src, [('corp', None)])[0] == want, src
    assert ve.strict('company', {'v': '株式会社 アクアテック'}, '株式会社アクアテック') == "会社の種類の前後に空白が残っている"


def test_grader_duplicate_claim_is_dropped_when_tool_finds_none():
    import vbam_grade as vg

    class _Sheet:
        class UsedRange:
            Value = (("会社名", "担当者", "電話"), ("株式会社A", "佐藤 誠", "03-1111-2222"),
                     ("株式会社A", "鈴木 恵", "03-1333-4444"), ("株式会社B", "高橋 学", "03-1555-6666"),
                     ("株式会社C", "田中 彩", "03-1777-8888"))

    class _Book:
        def Sheets(self, name):
            return _Sheet()
    keep, dropped = vg._unmet_drop_satisfied(["重複行が削除されていない：行3（株式会社A）", "見出しが太字でない"],
                                             _Book(), "s", "重複行は削除してよい")
    assert keep == ["見出しが太字でない"] and len(dropped) == 1 and '重複は 0 組' in dropped[0]


def test_materials_name_the_phone_and_postal_rules():
    import vbam_view as vv
    grid = [["氏名", "電話", "郵便番号"],
            ["佐藤 誠", "03-1234-5678", "010-0951"],
            ["鈴木 恵", "090 0568 3136", "〒010-0952"],
            ["高橋 学", "018(720)5698", "0100953"]]
    text = "\n".join(vv.dirt_notes(grid, 1, 1, 0))
    assert "電話の書き方がそろっていない（B列 2 セル" in text and "normalize の phone" in text
    assert "郵便番号の書き方がそろっていない（C列 2 セル" in text and "normalize の postal" in text


# ----------------------------------------------------------------
# マクロの先撃ち（vbam_prefire・9/11 深夜）: どの依頼で撃つか・重複行を消すも撃つか（純 Python）
# ----------------------------------------------------------------

@pytest.mark.parametrize('seed', [7, 2026])
def test_prefire_plan_for_exam_requests(seed):
    import vbam_prefire as vp
    want = {'顧客名簿': True, '売上明細': True, '会員名簿': True, '経費精算': False, '在庫表': True}
    for ex in ve.make_exams(seed):
        fire, dedupe, other = vp.plan_of(ex['request'])
        assert fire and not other, ex['name']
        assert dedupe == want[ex['name']], ex['name']


def test_prefire_plan_words():
    import vbam_prefire as vp
    assert vp.plan_of("この表を見て") == (False, False, False)                      # 整理の依頼でない＝撃たない
    assert vp.plan_of("重複を消して") == (False, False, False)                      # 整える語が無い＝撃たない（AI の往復）
    assert vp.plan_of("書き方をそろえて。重複行は削除してよい") == (True, True, False)
    assert vp.plan_of("書き方をそろえて。重複はそのまま残して") == (True, False, False)
    assert vp.plan_of("表を整えて、下に合計を出して") == (True, False, True)          # 合計はマクロの外＝撃った後に AI
    assert vp.plan_of("重複もあるけど消さないで、表記をそろえて") == (True, False, False)


def test_harder_dirt_appears_and_is_read_leniently_20260912():
    """9/12 に足した汚れ（改行・ひらがなのフリガナ・長音の記号・数値／8 桁の日付・負の数・率・途中の空行）が
    お題に出て、採点のゆるい読みでは同じ行と分かり、strict は書き方の誤りとして咎める。"""
    kinds, blanks = set(), 0
    for seed in range(12):
        for ex in ve.make_exams(seed):
            blanks += sum(1 for d in ex['dirty'] if d.get('blank'))
            for d in ex['dirty']:
                kinds |= {c.get('kind') for c in d['cells'] if c.get('kind')}
    for k in ('linebreak', 'hiragana', 'choon', 'serial', 'digits8', 'text_pct', 'zen_pct'):
        assert k in kinds, k
    assert kinds & {'tri', 'blk', 'paren', 'minus_text'}
    assert blanks > 0
    assert ve.ident('date', {'v': 46266.0}) == ('d', '2026-09-01')
    assert ve.ident('date', {'v': '20260901'}) == ('d', '2026-09-01')
    assert ve.ident('money', {'v': '△3,000'}) == ve.ident('money', {'v': '(3,000)'}) == ('n', -3000.0)
    assert ve.ident('pct', {'v': '５％'}) == ve.ident('pct', {'v': 0.05}) == ('n', 0.05)
    assert ve.ident('kana', {'v': 'さとう まこと'}) == ve.ident('kana', {'v': 'サトウ マコト'})
    assert ve.ident('text', {'v': 'マスキングテ-プ'}) == ve.ident('text', {'v': 'マスキングテープ'})
    assert ve.ident('text', {'v': 'トナ‐TN-27'}) == ve.ident('text', {'v': 'トナーTN-27'})
    assert ve.strict('person', {'v': '佐藤\n誠'}, '佐藤 誠') == "セルの中の改行が残っている"
    assert ve.strict('kana', {'v': 'さとう'}, 'サトウ') == "ひらがなが残っている"
    assert ve.strict('text', {'v': 'コピ-用紙'}, 'コピー用紙').startswith("長音が記号のまま")
    assert ve.strict('text', {'v': 'ルーム-2'}, 'ルーム-2') is None           # 後ろが数字の記号は区切り
    assert ve.strict('pct', {'v': '5%'}, 0.05) == "数値になっていない（文字のまま）"


def test_body_rows_skip_single_blank_rows_and_stop_at_two():
    exj = {'columns': [['名', 'text'], ['数', 'int']]}
    cells = [[{'v': '名'}, {'v': '数'}], [{'v': 'a'}, {'v': 1.0}], [{'v': None}, {'v': None}],
             [{'v': 'b'}, {'v': 2.0}], [{'v': None}, {'v': None}], [{'v': None}, {'v': None}], [{'v': '注記'}, {'v': None}]]
    hdr = ve.find_header(cells, ['名', '数'])
    assert [r[0]['v'] for r in ve.body_rows(cells, hdr, ['名', '数'])] == ['a', 'b']
    assert ve.body_span(cells, hdr, ['名', '数']) == (1, 3)


def test_write_book_keeps_text_as_text(tmp_path):
    from openpyxl import load_workbook
    ex = ve.make_exams(7, '顧客名簿')[0]
    path = tmp_path / 'c.xlsx'
    ve.write_book(ex, str(path))
    ws = load_workbook(str(path))['顧客名簿']
    assert [ws.cell(1, j).value for j in range(1, 8)] == [n for n, _t in ex['columns']]
    phones = [ws.cell(r, 3).value for r in range(2, 2 + len(ex['dirty']))]
    assert all(isinstance(p, str) for p in phones if p is not None)       # None は途中の空行
    assert sum(1 for d in ex['dirty'] if d.get('blank')) >= 1
    assert any(p.isdigit() for p in phones if p)     # 区切りなしの電話も文字のまま（先頭の 0 が残る）
