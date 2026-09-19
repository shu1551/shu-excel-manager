# -*- coding: utf-8 -*-
"""課別支出額の集合縦棒グラフ（鍛える回路の題材・グラフ 1 本目・2026-09-17 深夜）。

仕事: 「課別支出」シートの表（課名・支出額）から、集合縦棒グラフを作る（デザインの指定つき）。
正解の決まり（依頼文に全部書く）:
  - 系列は 1 本: 名前＝支出額の見出しのセル・値＝支出額・項目＝課名。合計の行・空の行は入れない。
  - タイトル「課別支出額」・凡例なし・棒に値のデータラベル・数値軸の表示形式 #,##0・数値軸の目盛線なし・
    棒の間隔（GapWidth）60・棒の色 #1F4E79・グラフの文字 Meiryo UI。
  - 置き場所: 表の右に 1 列空けた列の、見出しの行の高さ（位置は比べない）。
  - 前に作ったグラフがこのシートに残っていたら消してから作る（2 回撃ってもグラフは 1 つ）。
正解のブックは、直す前のブックを自分の Excel（非表示）で開き、この決まりどおりにグラフを作って保存する（人の Excel に触らない）。
組: 本番 6 課／A 12 課／B 合計行／C 表題・B3 から／D 見出しの改行と単位／E 前のグラフが残っている／F 途中に空行／
  G 金額が数式／H 撃った後にもう一度
  py make_pairs.py            → *_前.xlsx と *_正解.xlsx
  py make_pairs.py --unseen N → 未見_種N
"""
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
KA = ['総務課', '財政課', '企画課', '税務課', '福祉課', '建設課', '農林課', '教育総務課', '観光課', '環境課', '会計課', '議会事務局',
      '商工課', '保健課', '子育て支援課', '市民課', '防災課', '都市計画課', '上下水道課', '文化スポーツ課']
SHEET = '課別支出'
TITLE = '課別支出額'
COLOR = (0x1F, 0x4E, 0x79)
REQUEST = ('課別支出シートの表から、課ごとの支出額の集合縦棒グラフを、表の右に1列空けた列の見出しの行の高さに作ってください。'
           'タイトルは「課別支出額」、凡例なし、棒に値のデータラベル、縦軸の表示形式は #,##0、縦軸の目盛線なし、棒の間隔は 60、'
           '棒の色は #1F4E79、グラフの文字は Meiryo UI。合計の行と空の行はグラフに入れない。'
           '前に作ったグラフが残っていたら消して作り直す')


def build(seed, n):
    rng = random.Random(seed)
    kas = rng.sample(KA, n) if n <= len(KA) else [f"{rng.choice(KA)}{i + 1}" for i in range(n)]
    return [(k, rng.randrange(50, 5000) * 1000) for k in kas]


def write_before(path, rows, title=False, r0=1, c0=1, total=False, head='支出額', blank_at=None, formula=False, old_chart=False):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    h = r0
    if title:
        ws.cell(r0, c0, '令和8年度 課別支出（9月末）').font = Font(bold=True, size=14)
        h = r0 + 2
    ws.cell(h, c0, '課名').font = Font(bold=True)
    ws.cell(h, c0 + 1, head).font = Font(bold=True)
    if '\n' in head:
        ws.cell(h, c0 + 1).alignment = Alignment(wrap_text=True)
    r = h + 1
    data_rows = []
    for i, (k, a) in enumerate(rows):
        if blank_at is not None and i == blank_at:
            r += 1
        ws.cell(r, c0, k)
        ws.cell(r, c0 + 1, f"={a // 2}+{a - a // 2}" if formula else a).number_format = '#,##0'
        data_rows.append(r)
        r += 1
    if total:
        ws.cell(r, c0, '合計').font = Font(bold=True)
        ws.cell(r, c0 + 1, f"=SUM({chr(64 + c0 + 1)}{h + 1}:{chr(64 + c0 + 1)}{r - 1})").number_format = '#,##0'
    wb.save(path)
    return {'h': h, 'c0': c0, 'rows': data_rows, 'old_chart': old_chart}


def build_truth(xl, before, truth, layout):
    """直す前のブックを開き、決まりどおりのグラフを作って truth に保存する（COM）。old_chart なら前のグラフも入れてから消す。"""
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Sheets(SHEET)
        h, c0, rows = layout['h'], layout['c0'], layout['rows']
        for i in range(int(ws.ChartObjects().Count), 0, -1):
            ws.ChartObjects(i).Delete()
        left = ws.Cells(h, c0 + 3).Left
        top = ws.Cells(h, c0 + 3).Top
        co = ws.ChartObjects().Add(left, top, 480, 288)
        co.Name = '表の整理_グラフ'      # マクロは「表の整理_」で始まる名前のグラフ（自分が作った物）だけを消す
        ch = co.Chart
        ch.ChartType = 51
        for i in range(int(ch.SeriesCollection().Count), 0, -1):
            ch.SeriesCollection(i).Delete()
        s = ch.SeriesCollection().NewSeries()
        col = c0 + 1

        def union(c):
            return ws.Range(",".join(ws.Cells(r, c).Address for r in rows))
        s.Name = f"='{SHEET}'!{ws.Cells(h, col).Address}"
        s.Values = union(col)
        s.XValues = union(c0)
        ch.HasTitle = True
        ch.ChartTitle.Text = TITLE
        ch.HasLegend = False
        s.HasDataLabels = True
        ax = ch.Axes(2)
        ax.TickLabels.NumberFormat = '#,##0'
        ax.HasMajorGridlines = False
        ch.ChartGroups(1).GapWidth = 60
        s.Format.Fill.ForeColor.RGB = COLOR[0] + COLOR[1] * 256 + COLOR[2] * 65536
        ch.ChartArea.Format.TextFrame2.TextRange.Font.Name = 'Meiryo UI'
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def add_old_chart(xl, path, layout):
    """直す前のブックに「前に作ったグラフ」（折れ線・題「旧」）を置いて保存する（COM）。"""
    wb = xl.Workbooks.Open(path)
    try:
        ws = wb.Sheets(SHEET)
        co = ws.ChartObjects().Add(ws.Cells(layout['h'], layout['c0'] + 3).Left, ws.Cells(layout['h'], layout['c0'] + 3).Top, 300, 200)
        co.Name = '表の整理_前のグラフ'      # マクロは「表の整理_」で始まる名前のグラフ（自分が作った物）だけを消す
        co.Chart.ChartType = 4
        s = co.Chart.SeriesCollection().NewSeries()
        s.Values = ws.Range(ws.Cells(layout['rows'][0], layout['c0'] + 1), ws.Cells(layout['rows'][-1], layout['c0'] + 1))
        co.Chart.HasTitle = True
        co.Chart.ChartTitle.Text = '旧'
        wb.Save()
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, rows, kw in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            layout = write_before(before, rows, **kw)
            if kw.get('old_chart'):
                add_old_chart(xl, before, layout)
            build_truth(xl, before, truth, layout)
            print('  ', stem)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', build(s(1), rng.randint(3, 15)), {}),
          ('未見2_合計行と表題', build(s(2), 8), {'title': True, 'r0': 2, 'c0': 2, 'total': True}),
          ('未見3_見出しと空行', build(s(3), 9), {'head': '支出額\n（円）', 'blank_at': 4}),
          ('未見4_前のグラフと数式', build(s(4), 7), {'old_chart': True, 'formula': True}),
          ('未見5_大量', build(s(5), 40), {'total': True})], out)
    return out


if __name__ == '__main__':
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        sys.exit(0)
    make([('グラフ_本番', build(3, 6), {}),
          ('グラフ_試験A', build(13, 12), {}),
          ('グラフ_試験B', build(23, 7), {'total': True}),
          ('グラフ_試験C', build(33, 6), {'title': True, 'r0': 3, 'c0': 2}),
          ('グラフ_試験D', build(43, 8), {'head': '支出額\n（円）'}),
          ('グラフ_試験E', build(53, 6), {'old_chart': True}),
          ('グラフ_試験F', build(63, 9), {'blank_at': 5}),
          ('グラフ_試験G', build(73, 7), {'formula': True})], HERE)
    shutil.copyfile(os.path.join(HERE, 'グラフ_本番_正解.xlsx'), os.path.join(HERE, 'グラフ_試験H_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, 'グラフ_本番_正解.xlsx'), os.path.join(HERE, 'グラフ_試験H_正解.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    print('ok')
