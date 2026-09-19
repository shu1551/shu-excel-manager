# -*- coding: utf-8 -*-
"""支出明細のテーブルに、課マスタのテーブルから課名を引く計算列を足す（鍛える回路の題材・テーブル 2 本目・2026-09-18）。

仕事: テーブル「T支出明細」の課コードの列のすぐ右に列「課名」を足し、テーブル「T課マスタ」から課コードで課名を引く計算列にする。
正解の決まり（依頼文に全部書く）:
  - 式は構造化参照（[@課コード]・T課マスタ[課コード]・T課マスタ[課名]）。マスタに無いコードは「未登録」。
  - 既に「課名」の列があれば、その列を計算列に入れ直す（列を増やさない）。
  - ほかの列・テーブルの名前・スタイルは変えない。
組: 本番 30 行／A 3,000 行／B マスタに無いコード／C 課名の列が既にある（手で打った古い値）／D 課コードが右端の並び／
  E 表題の下にテーブル（A3）／F 撃った後にもう一度／G マスタが同じシートの右にある
"""
import datetime
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
KA = [(110, '総務課'), (120, '財政課'), (130, '企画課'), (140, '税務課'), (210, '福祉課'), (310, '建設課'), (320, '農林課'),
      (410, '教育総務課'), (510, '観光課'), (520, '環境課')]
TEKIYO = ['コピー用紙', '郵送料', '清掃業務', '複合機リース', '修繕工事', 'パソコン購入', '出張旅費']
HEAD = ['伝票番号', '支出日', '課コード', '摘要', '支出額']
SWAP = ['伝票番号', '支出日', '摘要', '支出額', '課コード']
SHEET = '支出明細'
MASTER = '課マスタ'
FORMULA = '=IFERROR(INDEX(T課マスタ[課名],MATCH([@課コード],T課マスタ[課コード],0)),"未登録")'
REQUEST = ('テーブル T支出明細 の課コードの列のすぐ右に列「課名」を足し、テーブル T課マスタ から課コードで課名を引く計算列にしてください'
           '（構造化参照で。マスタに無いコードは「未登録」）。既に課名の列があれば、その列を計算列に入れ直す（列を増やさない）')


def build(seed, n, unknown=False):
    rng = random.Random(seed)
    kas = rng.sample(KA, rng.randint(4, len(KA)))
    rows = []
    for i in range(n):
        code, name = rng.choice(kas)
        if unknown and rng.random() < 0.1:
            code, name = rng.choice([990, 999]), '未登録'
        rows.append({'伝票番号': 26000 + i + 1, '支出日': datetime.datetime(2026, rng.randint(4, 9), rng.randint(1, 28)),
                     '課コード': code, '摘要': rng.choice(TEKIYO), '支出額': rng.randrange(1, 900) * 100, '課名': name})
    return rows, sorted(set(kas))


def write_before(path, data, cols=HEAD, title=False, stale_name=False, master_same_sheet=False):
    rows, kas = data
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    h = 1
    if title:
        ws.cell(1, 1, '令和8年度 支出明細').font = Font(bold=True, size=14)
        h = 3
    cols = list(cols)
    if stale_name:
        cols.insert(cols.index('課コード') + 1, '課名')
    for j, c in enumerate(cols, 1):
        ws.cell(h, j, c)
    rng = random.Random(len(rows))
    for i, r in enumerate(rows, 1):
        for j, c in enumerate(cols, 1):
            v = r[c]
            if c == '課名':                                    # 手で打った古い値（ところどころ違う）
                v = v if rng.random() < 0.7 else '（旧）' + v
            cell = ws.cell(h + i, j, v)
            if c == '支出日':
                cell.number_format = 'yyyy/m/d'
            if c == '支出額':
                cell.number_format = '#,##0'
    t = Table(displayName='T支出明細', ref=f"A{h}:{get_column_letter(len(cols))}{h + len(rows)}")
    t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(t)
    if master_same_sheet:
        mws, mc = ws, len(cols) + 3
    else:
        mws, mc = wb.create_sheet(MASTER), 1
    mws.cell(h if master_same_sheet else 1, mc, '課コード')
    mws.cell(h if master_same_sheet else 1, mc + 1, '課名')
    top = h if master_same_sheet else 1
    for i, (code, name) in enumerate(kas, 1):
        mws.cell(top + i, mc, code)
        mws.cell(top + i, mc + 1, name)
    mt = Table(displayName='T課マスタ', ref=f"{get_column_letter(mc)}{top}:{get_column_letter(mc + 1)}{top + len(kas)}")
    mt.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
    mws.add_table(mt)
    wb.active = 0
    wb.save(path)


def build_truth(wb):
    """決まりどおりの計算列（COM）。"""
    ws = wb.Sheets(SHEET)
    lo = ws.ListObjects('T支出明細')
    names = [str(lo.ListColumns(i).Name) for i in range(1, int(lo.ListColumns.Count) + 1)]
    if '課名' in names:
        lc = lo.ListColumns('課名')
    else:
        lc = lo.ListColumns.Add(names.index('課コード') + 2)
        lc.Name = '課名'
    lc.DataBodyRange.Formula = FORMULA


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, data, kw in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            write_before(before, data, **kw)
            wb = xl.Workbooks.Open(before)
            try:
                build_truth(wb)
                wb.Worksheets(SHEET).Activate()
                wb.SaveAs(truth, FileFormat=51)
            finally:
                wb.Close(SaveChanges=False)
            print('  ', stem)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', build(s(1), rng.randint(10, 80)), {}),
          ('未見2_無いコードと古い課名', build(s(2), 40, unknown=True), {'stale_name': True}),
          ('未見3_並び違いと表題', build(s(3), 25), {'cols': SWAP, 'title': True}),
          ('未見4_マスタが同じシート', build(s(4), 30, unknown=True), {'master_same_sheet': True})], out)
    return out


if __name__ == '__main__':
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        sys.exit(0)
    make([('参照_本番', build(3, 30), {}),
          ('参照_試験A', build(13, 3000), {}),
          ('参照_試験B', build(23, 40, unknown=True), {}),
          ('参照_試験C', build(33, 25), {'stale_name': True}),
          ('参照_試験D', build(43, 20), {'cols': SWAP}),
          ('参照_試験E', build(53, 22), {'title': True}),
          ('参照_試験G', build(63, 18), {'master_same_sheet': True})], HERE)
    shutil.copyfile(os.path.join(HERE, '参照_本番_正解.xlsx'), os.path.join(HERE, '参照_試験F_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, '参照_本番_正解.xlsx'), os.path.join(HERE, '参照_試験F_正解.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    lines = ['支出明細に課マスタから課名を引いて', '課コードから課名をマスタ参照で出して', 'T支出明細に課名の列を足してマスタから引く',
             '課マスタを見て課名を入れる計算列を', '明細のテーブルに課名を VLOOKUP みたいに引いて', '課コードの右に課名をマスタから']
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(lines) + '\n')
    print('ok')
