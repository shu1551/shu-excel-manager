# -*- coding: utf-8 -*-
"""表の右端に「推移」の列を足し、行ごとにスパークライン（折れ線）を入れる（E グラフ 45・2026-09-18 第二期）。

列の決まり: 選ぶ列は無い（表全体）。表は推移グラフのお題の部品（項目の列＋月・四半期・年度の数の列＋合計の列と行）。
正解の決まり（依頼文に全部書く）:
  - 見出しの行＝空でないセルが 2 つ以上並ぶ最初の行。項目の列＝左端の文字の列。表の右端＝見出しの行で左端から右へ空のセルに当たる手前。
  - 元の範囲＝その行の、項目の列より右で見出しが「合計」「計」でない数の列が続く所（左端の番号・コードの列、右の備考の列は入れない）。
  - 表の右端の次の列の見出しに「推移」。本文の行ごとに 1 本（合計・計の行と空の行には入れない）。
  - 折れ線・最高点に印・線の色 #1F4E79。既に「推移」の列があれば、そこのスパークラインを消して入れ直す（列は足さない）。
組: 本番 月／A 8 件・合計の行と列／B 表題・下半期／C 合計なし・四半期／D 年度／E 日付の見出し／F 撃った後にもう一度／
  G 表題・コードの列が左に／H 途中に空行・備考の列が右に
"""
import importlib.util
import os
import shutil
import sys

from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
_spec = importlib.util.spec_from_file_location('suii', os.path.join(HERE, '..', '推移グラフ', 'make_pairs.py'))
suii = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(suii)

REQUEST = ('表の右端に「推移」の列を足して、行ごとの推移をスパークライン（折れ線）で入れてください。'
           '見出しの行＝空でないセルが 2 つ以上並ぶ最初の行（文字でも数でも日付でもよい）。項目の列＝左端の文字の列。'
           '表の右端＝見出しの行で、表のいちばん左の列から右へ見て空のセルに当たる手前の列。'
           '元の範囲＝その行の、項目の列より右にあって、見出しが「合計」「計」でない数の列が続く所（項目の列より左の番号・コードの列と、右の文字の列は入れない）。'
           '表の右端の次の列の見出しに「推移」と書き、本文の行ごとに 1 本入れる（合計・計・総計の行と、行がまるごと空の行には入れない）。'
           '種類は折れ線（xlSparkLine）、最高点に印を付け（Points.Highpoint.Visible）、線の色は #1F4E79。'
           '見出しの行に既に「推移」の列があれば、列は足さずに、その列のスパークラインを消して入れ直す。表の値は触らない。')
PHRASES = ['スパークラインの列を足して', '各行にスパークラインを入れて', '推移をスパークラインで見せて',
           'セル内の小さなグラフで推移を', '行ごとの傾向をミニグラフで', '表の右にスパークラインをお願い']
COLOR = 0x794E1F      # #1F4E79（Excel の色の数は B*65536+G*256+R）


def build_truth(xl, before, truth, t, lay):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(t['sheet'])
        h, ic = lay['h'], lay['item_c']
        c = ic + 1
        while ws.Cells(h, c).Value is not None:
            c += 1
        right = c - 1                                             # 表の右端
        first = ic + 1
        last = first + len(t['heads']) - 1                        # 月などの数の列（合計の列の手前まで）
        out = right + 1
        ws.Cells(h, out).Value = '推移'
        for r in lay['rows']:
            src = f"{get_column_letter(first)}{r}:{get_column_letter(last)}{r}"
            g = ws.Cells(r, out).SparklineGroups.Add(1, src)
            g.Points.Highpoint.Visible = True
            g.SeriesColor.Color = COLOR
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t in jobs:
            before, truth = os.path.join(out, f"{stem}_前.xlsx"), os.path.join(out, f"{stem}_正解.xlsx")
            lay = suii.write_before(before, t)
            build_truth(xl, before, truth, t, lay)
            print('  ', stem, t['sheet'], t['ka'], t['period'], len(t['rows']))
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_月', suii.build(s(1), 5)),
          ('未見2_表題と四半期', suii.build(s(2), 6, 'quarter', title=True, totals=False)),
          ('未見3_コードと空行', suii.build(s(3), 7, code_col=True, blank_at=3)),
          ('未見4_年度と備考', suii.build(s(4), 4, 'year', note_col=True))], out)
    return out


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('推移線_本番', suii.build(2601, 4)),
          ('推移線_試験A', suii.build(2602, 8)),
          ('推移線_試験B', suii.build(2603, 5, 'half', title=True)),
          ('推移線_試験C', suii.build(2604, 7, 'quarter', totals=False)),
          ('推移線_試験D', suii.build(2605, 4, 'year')),
          ('推移線_試験E', suii.build(2606, 3, 'date')),
          ('推移線_試験G', suii.build(2607, 12, title=True, code_col=True)),
          ('推移線_試験H', suii.build(2608, 6, blank_at=2, note_col=True))], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '推移線_本番_正解.xlsx'), os.path.join(HERE, f'推移線_試験F_{x}.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
