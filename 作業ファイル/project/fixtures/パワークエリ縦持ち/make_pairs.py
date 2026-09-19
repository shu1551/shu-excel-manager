# -*- coding: utf-8 -*-
"""横持ちの月別支出をパワークエリで縦持ちにする（鍛える回路の題材・パワークエリ 2 本目・2026-09-18）。

仕事: テーブル「T月別支出」（課名と月ごとの列）を、列が 課名・月・支出額 の縦持ちに変換し、新しいシート「PQ縦持ち」の A1 に
      テーブル「Q縦持ち」として読み込む。クエリ名も Q縦持ち。
正解の決まり（依頼文に全部書く）:
  - 課コードの列と合計の列は入れない（無い表もある）。課名が「合計」の行も入れない。
  - 空欄の月は行を作らない（0 は行を作る）。月の列の数は表ごとに違う（4月〜9月・4月〜3月など）。
  - 同じ名前のクエリやシートが既にあれば消して作り直す。
組: 本番 5 課×6 か月／A 10 課×12 か月／B 合計の列あり／C 課コードの列なし／D 空欄の月と 0／E 前のクエリとシート／
  F 撃った後にもう一度／G 課名が合計の行がテーブルの中にある
"""
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
KA = ['総務課', '財政課', '企画課', '税務課', '福祉課', '建設課', '農林課', '教育総務課', '観光課', '環境課', '会計課', '議会事務局']
MONTHS = ['4月', '5月', '6月', '7月', '8月', '9月', '10月', '11月', '12月', '1月', '2月', '3月']
QNAME = 'Q縦持ち'
OUT_SHEET = 'PQ縦持ち'
SHEET = '月別支出'
M = '''let
    Source = Excel.CurrentWorkbook(){[Name="T月別支出"]}[Content],
    Removed = Table.RemoveColumns(Source, {"課コード", "合計"}, MissingField.Ignore),
    NoTotal = Table.SelectRows(Removed, each [課名] <> "合計"),
    Unpivoted = Table.UnpivotOtherColumns(NoTotal, {"課名"}, "月", "支出額"),
    Typed = Table.TransformColumnTypes(Unpivoted, {{"課名", type text}, {"月", type text}, {"支出額", Int64.Type}})
in
    Typed'''
REQUEST = ('パワークエリで、テーブル T月別支出（課名と月ごとの列）を、列が 課名・月・支出額 の縦持ちに変換し、新しいシート「PQ縦持ち」の A1 に'
           'テーブル「Q縦持ち」として読み込んでください。クエリ名も Q縦持ち。課コードの列と合計の列は入れない。課名が合計の行も入れない。'
           '空欄の月は行を作らない。同じ名前のクエリやシートが既にあれば消して作り直す')


def build(seed, n_ka, n_month, blanks=False):
    rng = random.Random(seed)
    kas = rng.sample(KA, n_ka)
    rows = []
    for i, k in enumerate(kas):
        vals = {}
        for m in MONTHS[:n_month]:
            v = rng.randrange(0, 900) * 1000
            if blanks and rng.random() < 0.2:
                v = None if rng.random() < 0.6 else 0
            vals[m] = v
        rows.append({'課コード': f"{110 + i * 10}", '課名': k, **vals})
    return rows, MONTHS[:n_month]


def write_before(path, data, code=True, total_col=False, total_row=False):
    rows, months = data
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    cols = (['課コード'] if code else []) + ['課名'] + months + (['合計'] if total_col else [])
    for j, c in enumerate(cols, 1):
        ws.cell(1, j, c).font = Font(bold=True)
    body = list(rows)
    if total_row:
        body.append({'課コード': None, '課名': '合計', **{m: sum((r[m] or 0) for r in rows) for m in months}})
    for i, r in enumerate(body, 2):
        for j, c in enumerate(cols, 1):
            if c == '合計':
                v = sum((r[m] or 0) for m in months)
            else:
                v = r.get(c)
            if v is not None:
                ws.cell(i, j, v)
    t = Table(displayName='T月別支出', ref=f"A1:{get_column_letter(len(cols))}{len(body) + 1}")
    t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(t)
    wb.save(path)


def load_query(wb, name=QNAME, formula=M, sheet=OUT_SHEET):
    """クエリを作り（同じ名前があれば消す）、シート sheet の A1 に表として読み込む（COM）。"""
    xl = wb.Application
    for i in range(int(wb.Queries.Count), 0, -1):
        if str(wb.Queries.Item(i).Name) == name:
            wb.Queries.Item(i).Delete()
    for sh in list(wb.Worksheets):
        if str(sh.Name) == sheet:
            xl.DisplayAlerts = False
            sh.Delete()
    wb.Queries.Add(name, formula)
    ws = wb.Worksheets.Add(After=wb.Worksheets(wb.Worksheets.Count))
    ws.Name = sheet
    conn = f"OLEDB;Provider=Microsoft.Mashup.OleDb.1;Data Source=$Workbook$;Location={name};Extended Properties=\"\""
    lo = ws.ListObjects.Add(SourceType=0, Source=conn, Destination=ws.Range("$A$1"))
    lo.QueryTable.CommandType = 2
    lo.QueryTable.CommandText = f"SELECT * FROM [{name}]"
    lo.QueryTable.Refresh(False)
    lo.Name = name
    return lo


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, data, kw in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            stale = kw.pop('stale', False)
            write_before(before, data, **kw)
            if stale:                                        # 前のクエリ（列の名前が 金額 の古い形）とシートを残しておく
                wb = xl.Workbooks.Open(before)
                try:
                    load_query(wb, formula=M.replace('"月", "支出額"', '"月", "金額"').replace('{"支出額", Int64.Type}', '{"金額", Int64.Type}'))
                    wb.Worksheets(1).Activate()
                    wb.Save()
                finally:
                    wb.Close(SaveChanges=False)
            wb = xl.Workbooks.Open(before)
            try:
                load_query(wb)
                wb.Worksheets(1).Activate()
                wb.SaveAs(truth, FileFormat=51)
            finally:
                wb.Close(SaveChanges=False)
            print('  ', stem)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', build(s(1), rng.randint(3, 8), rng.randint(3, 12)), {}),
          ('未見2_合計の列と行・空欄', build(s(2), 7, 9, blanks=True), {'total_col': True, 'total_row': True}),
          ('未見3_課コードなしと前のクエリ', build(s(3), 5, 12), {'code': False, 'stale': True})], out)
    return out


if __name__ == '__main__':
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        sys.exit(0)
    make([('縦持ち_本番', build(3, 5, 6), {}),
          ('縦持ち_試験A', build(13, 10, 12), {}),
          ('縦持ち_試験B', build(23, 6, 6), {'total_col': True}),
          ('縦持ち_試験C', build(33, 5, 8), {'code': False}),
          ('縦持ち_試験D', build(43, 6, 6, blanks=True), {}),
          ('縦持ち_試験E', build(53, 4, 6), {'stale': True}),
          ('縦持ち_試験G', build(63, 5, 6), {'total_row': True})], HERE)
    shutil.copyfile(os.path.join(HERE, '縦持ち_本番_正解.xlsx'), os.path.join(HERE, '縦持ち_試験F_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, '縦持ち_本番_正解.xlsx'), os.path.join(HERE, '縦持ち_試験F_正解.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    lines = ['パワークエリで月別支出を縦持ちに', '横持ちの月別の表をパワークエリで縦に並べ替えて', 'T月別支出をピボット解除して縦持ちにして',
             'パワークエリで列のピボット解除をお願い', '月の列を縦持ちに変換するクエリを作って', '月別支出を課名・月・支出額の縦持ちに']
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(lines) + '\n')
    print('ok')
