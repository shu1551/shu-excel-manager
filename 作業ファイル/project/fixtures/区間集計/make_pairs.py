# -*- coding: utf-8 -*-
"""選んでいる数の列を区間（金額帯）に分けて、区間ごとの件数の表を表の右に作る（B 集計 19・2026-09-18）。

列の決まり: 選んでいる 1 列＝数の列。選んでいる列は <名前>_選択.txt に書く。
正解の決まり（依頼文に全部書く）:
  - 幅＝本文の最大値を 10 で割った値以上で、いちばん小さい 1・2・5×10 の累乗（最大 8,731,000 → 1,000,000）。
  - 区間は 0 から幅ずつ、最大値を含む区間まで（件数 0 の区間も出す）。
  - 置き場所＝表の右に 1 列空けた所。見出しの行は選んでいるセルの行。見出しは「以上」「未満」「件数」。
  - 以上・未満は値（#,##0）。件数は =COUNTIFS(本文, ">="&以上のセル, 本文, "<"&未満のセル)（本文は $ 付きの絶対参照）。
  - 本文＝見出しの次の行から最後の本文の行まで。合計（合計・計）の行は入れない。既に同じ表が右にあれば作り直す。
組: 本番／A 40 行／B 表題と単位／C 合計の行／D 途中の空行／E 列の並び違い／F 撃った後にもう一度／G 番号の列と摘要／H 千円単位の小さい数
"""
import os
import random
import shutil
import sys

from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import shukei   # noqa: E402

REQUEST = ('選んでいる数の列を金額の区間に分けて、区間ごとの件数の表を、表の右に 1 列空けて作ってください。'
           '見出しの行は必ず「選んでいるセルの行」＝Selection.Cells(1,1).Row をそのまま使う（自分で見出しの行を探さない）。'
           '表の右端＝見出しの行で、表のいちばん左の列から右へ見て空のセルに当たる手前の列。その 2 つ右の列から 3 列の表を書く。'
           '区間の幅＝本文の最大値を 10 で割った値以上で、いちばん小さい「1・2・5 × 10 の累乗」'
           '（例: 最大 8,731,000 なら 1,000,000／最大 4,200,000 なら 500,000／最大 8,731 なら 1,000）。'
           '区間は 0 から幅ずつ、最大値を含む区間まで（件数が 0 の区間も出す）。'
           '見出しは「以上」「未満」「件数」。以上・未満は値で書き、表示形式は #,##0。'
           '件数は =COUNTIFS(本文, ">="&以上のセル, 本文, "<"&未満のセル) の式（本文の範囲は $ を付けた絶対参照）。'
           '本文は見出しの次の行から最後の本文の行まで。表の下に合計（合計・計）の行があればそれは入れない'
           '（合計の行＝行のどこかのセル、ふつうは左の項目の列に「合計」「計」「総計」とある行。選んだ数の列には合計の数が入っているので、数の列だけ見ても見分けられない）。'
           '既に同じ表が右にあれば作り直す。')
PHRASES = ['金額帯ごとの件数を出して', '金額の区間ごとに何件あるか表にして', '度数分布の表を右に作って',
           '選んだ列を階級に分けて件数を数えて', '金額の分布を表にして', '価格帯別の件数を出して']


def nice(x):
    """x 以上でいちばん小さい 1・2・5×10 の累乗（純 Python）。"""
    p = 1
    while True:
        for m in (1, 2, 5):
            if m * p >= x:
                return m * p
        p *= 10


def small(seed, n, **kw):
    """千円単位の表（数を 1/1000 に）。"""
    t = shukei.design(seed, n, **kw)
    for row in t['rows']:
        for nm, r in t['cols']:
            if r == 'num' and isinstance(row.get(nm), int):
                row[nm] = row[nm] // 1000
    return t


def build_truth(xl, before, truth, t, lay):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(t['sheet'])
        c0, h, last = lay['c0'], lay['h'], lay['last']
        nc = c0 + shukei.col_of(t, 'num') - 1
        nl = get_column_letter(nc)
        nrng = f"${nl}${h + 1}:${nl}${last}"
        vals = [ws.Cells(r, nc).Value for r in range(h + 1, last + 1)]
        mx = max(v for v in vals if isinstance(v, (int, float)))
        w = nice(mx / 10)
        out = c0 + lay['ncols'] + 1
        lo_l, hi_l = get_column_letter(out), get_column_letter(out + 1)
        for j, nm in enumerate(('以上', '未満', '件数')):
            ws.Cells(h, out + j).Value = nm
        k = 0
        while k * w <= mx:
            r = h + 1 + k
            ws.Cells(r, out).Value2 = k * w
            ws.Cells(r, out + 1).Value2 = (k + 1) * w
            ws.Cells(r, out).NumberFormat = '#,##0'
            ws.Cells(r, out + 1).NumberFormat = '#,##0'
            ws.Cells(r, out + 2).Formula = f'=COUNTIFS({nrng},">="&{lo_l}{r},{nrng},"<"&{hi_l}{r})'
            k += 1
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            lay = shukei.write_before(before, t)
            sel = f"{get_column_letter(shukei.col_of(t, 'num'))}{lay['h']}"
            build_truth(xl, before, truth, t, lay)
            open(os.path.join(out, f"{stem}_選択.txt"), 'w', encoding='utf-8').write(sel + '\n')
            print('  ', stem, t['sheet'], [n for n, _r in t['cols']], '選ぶ列', sel)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', shukei.design(s(1), rng.randint(6, 14))),
          ('未見2_表題と合計行', shukei.design(s(2), 9, title=True, unit=True, total=True)),
          ('未見3_並び違いの千円', small(s(3), 8, swap=True)),
          ('未見4_空行と番号列', shukei.design(s(4), 10, blanks=2, no_col=True))], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('区間_本番', shukei.design(2001, 12)),
          ('区間_試験A', shukei.design(2002, 40)),
          ('区間_試験B', shukei.design(2003, 8, title=True, unit=True)),
          ('区間_試験C', shukei.design(2004, 10, total=True)),
          ('区間_試験D', shukei.design(2005, 11, blanks=2)),
          ('区間_試験E', shukei.design(2006, 9, swap=True)),
          ('区間_試験G', shukei.design(2007, 7, no_col=True, text=True)),
          ('区間_試験H', small(2008, 10))], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '区間_本番_正解.xlsx'), os.path.join(HERE, f'区間_試験F_{x}.xlsx'))
    shutil.copyfile(os.path.join(HERE, '区間_本番_選択.txt'), os.path.join(HERE, '区間_試験F_選択.txt'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
