# -*- coding: utf-8 -*-
"""選んでいる列の式を、見た目の値のまま値にする（2026-09-18 第二期）。

列の決まり: 選んでいる列（1〜2 本）。<名前>_選択.txt に書く。表は集計表（shukei）に式の列（差額・率）と合計の行の式を足したもの。
正解の決まり（依頼文に全部書く）:
  - 見出しの行＝選んでいるセルの行。選んだ列の、見出しの次の行から表の最後の行（下の合計の行も入れる）までの式を、その値にする。
  - 値のセル・空のセルはそのまま。ほかの列は変えない（ほかの列の式は式のまま）。
組: 本番 差額の列／A 40 行／B 表題と単位／C 合計の行の式／D 途中の空行／E 2 列（差額と率）／F 撃った後にもう一度／G 式と値が混ざった列／H 式の無い列
"""
import os
import random
import shutil
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import shukei   # noqa: E402

REQUEST = ('選んでいる列の式を、今の値のまま値にしてください（式を消して値だけ残す）。'
           '見出しの行は必ず「選んでいるセルの行」＝Selection.Cells(1,1).Row をそのまま使う（自分で見出しの行を探さない）。選んでいる列は Selection.Areas の各列。'
           '選んだ列の、見出しの次の行から表の最後の行（下の合計の行も入れる）までの式を、その値にする（rng.Value = rng.Value）。'
           '値のセル・空のセルはそのまま。ほかの列は変えない（ほかの列の式は式のまま残す）。')
PHRASES = ['選んだ列の式を値にして', '数式を値に変換して', 'この列を値貼り付けにして',
           '選んでいる列の計算式を消して値だけに', '式を値で固定して', '関数を値にしておいて']


def add_formulas(path, t, lay, mixed=False, total=False):
    """差額（2 つ目の数−1 つ目の数）と率（2 つ目÷1 つ目）の列を右に足し、合計の行にも SUM の式を入れる（openpyxl）。"""
    wb = load_workbook(path)
    ws = wb.worksheets[0]
    h, last, c0 = lay['h'], lay['last'], lay['c0']
    a = get_column_letter(c0 + shukei.col_of(t, 'num', 0) - 1)
    b = get_column_letter(c0 + shukei.col_of(t, 'num', 1) - 1)
    dc, rc = c0 + lay['ncols'], c0 + lay['ncols'] + 1
    ws.cell(h, dc, '差額')
    ws.cell(h, rc, '率')
    rng = random.Random(t['seed'] + 9)
    for r in range(h + 1, last + 1):
        if ws.cell(r, c0 + shukei.col_of(t, 'num', 0) - 1).value is None:
            continue
        if mixed and rng.random() < 0.3:
            ws.cell(r, dc, rng.randrange(1, 900) * 1000).number_format = '#,##0'
        else:
            ws.cell(r, dc, f"={b}{r}-{a}{r}").number_format = '#,##0'
        ws.cell(r, rc, f"=IF({a}{r}=0,0,{b}{r}/{a}{r})").number_format = '0.0%'
    if total and lay.get('total_at'):
        tr = lay['total_at']
        for col in (c0 + shukei.col_of(t, 'num', 0) - 1, c0 + shukei.col_of(t, 'num', 1) - 1, dc):
            L = get_column_letter(col)
            ws.cell(tr, col, f"=SUM({L}{h + 1}:{L}{last})").number_format = '#,##0'
    wb.save(path)
    return dc, rc


def build_truth(xl, before, truth, t, lay, cols, total):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(t['sheet'])
        last = lay['total_at'] if total and lay.get('total_at') else lay['last']
        for c in cols:
            rng = ws.Range(ws.Cells(lay['h'] + 1, c), ws.Cells(last, c))
            rng.Value = rng.Value
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t, which, mixed in jobs:
            before, truth = os.path.join(out, f"{stem}_前.xlsx"), os.path.join(out, f"{stem}_正解.xlsx")
            lay = shukei.write_before(before, t)
            dc, rc = add_formulas(before, t, lay, mixed=mixed, total=t['total'])
            cols = {'diff': [dc], 'both': [dc, rc], 'plain': [lay['c0'] + shukei.col_of(t, 'item') - 1]}[which]
            build_truth(xl, before, truth, t, lay, cols, t['total'])
            sel = ",".join(f"{get_column_letter(c)}{lay['h']}" for c in cols)
            open(os.path.join(out, f"{stem}_選択.txt"), 'w', encoding='utf-8').write(sel + '\n')
            print('  ', stem, t['sheet'], [n for n, _r in t['cols']], '選ぶ列', sel)
    finally:
        xl = None
        vf._quit(pid)


def d(seed, n, **kw):
    kw.setdefault('nums', 2)
    return shukei.design(seed, n, **kw)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', d(s(1), rng.randint(5, 14)), 'diff', False),
          ('未見2_表題と合計', d(s(2), 8, title=True, unit=True, total=True), 'both', False),
          ('未見3_前年度今年度', d(s(3), 7, prev=True), 'both', True),
          ('未見4_空行', d(s(4), 9, blanks=2), 'diff', False)], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('式値_本番', d(4001, 8), 'diff', False),
          ('式値_試験A', d(4002, 40), 'diff', False),
          ('式値_試験B', d(4003, 7, title=True, unit=True), 'diff', False),
          ('式値_試験C', d(4004, 9, total=True), 'diff', False),
          ('式値_試験D', d(4005, 10, blanks=2), 'diff', False),
          ('式値_試験E', d(4006, 8), 'both', False),
          ('式値_試験G', d(4007, 9), 'diff', True),
          ('式値_試験H', d(4008, 7), 'plain', False)], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '式値_本番_正解.xlsx'), os.path.join(HERE, f'式値_試験F_{x}.xlsx'))
    shutil.copyfile(os.path.join(HERE, '式値_本番_選択.txt'), os.path.join(HERE, '式値_試験F_選択.txt'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
