# -*- coding: utf-8 -*-
"""表を「いつもの見やすい推移グラフ」にする（どの表でも動く形・2026-09-18 作り直し）。
推移折れ線（デザインを依頼に書いた版）と見やすい推移グラフ（型を差分から学んだ版）を 1 本にまとめる。

列の決まり: 左端の文字の列が系列の名前、その右の数の列（合計・計の列は除く）が項目（見出しが項目の名前）。
項目は月（4月〜3月・10月〜3月・日付の見出し）・四半期・年度など表ごとに違う。見出しの語・シート名・項目の名前も違う（vocab.py）。
見せ方の型（正解の作り方）:
  - 見出しの行＝値の入ったセルが 2 つ以上ある最初の行。合計・計の行と空の行は入れない。
  - マーカー付き折れ線。系列は行ごと。行の合計の大きい行から 5 つまで（5 以下なら全部）、大きい順に並べる。
  - 線の色は順に #1F4E79・#C00000・#548235・#BF8F00・#7030A0、線の太さ 2.25、マーカーの大きさ 6。
  - タイトル: 表の上に表題があればその文字、無ければ「（左端の列の見出し）ごとの推移」。14pt・太字。
  - 凡例は右・数値軸 #,##0・目盛線あり・グラフの文字 Meiryo UI。前に作ったグラフは消して作り直す（置き場所は比べない）。
組: 本番 4 件・月／A 8 件（上位 5）・合計の行と列／B 表題 5 件・下半期／C 合計なし 7 件・四半期／D 前のグラフ・年度／
  E 3 件・日付の見出し／F 撃った後にもう一度／G 表題 12 件・コードの列が左に／H 途中に空行・備考の列が右に
"""
import datetime
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
from vocab import Vocab   # noqa: E402

REQUEST = ('表を、いつもの見やすい推移グラフにしてください。左端の文字の列を系列、その右の数の列（合計の列は除く）を項目にする。'
           '前に作ったグラフが残っていたら消して作り直す')
PHRASES = ['この表をいつもの推移グラフに', '見やすい推移グラフにして', '推移をいつもの見せ方の折れ線で',
           'いつもの形で推移のグラフを', '見栄えのいい推移グラフにしてください', '推移を見やすいグラフでお願い']
COLORS = [(0x1F, 0x4E, 0x79), (0xC0, 0x00, 0x00), (0x54, 0x82, 0x35), (0xBF, 0x8F, 0x00), (0x70, 0x30, 0xA0)]
TITLES = ['令和8年度 月次集計（一般会計）', '{k}別の実績の推移', '令和8年度 {k}ごとの執行状況']
PERIODS = {
    'month': [f"{m}月" for m in (4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3)],
    'half': [f"{m}月" for m in (10, 11, 12, 1, 2, 3)],
    'quarter': ['第1四半期', '第2四半期', '第3四半期', '第4四半期'],
    'year': ['令和4年度', '令和5年度', '令和6年度', '令和7年度', '令和8年度'],
    'date': 'date',
}


def rgb(c):
    return c[0] + c[1] * 256 + c[2] * 65536


def build(seed, n, period='month', **kw):
    rng = random.Random(seed)
    v = Vocab(seed)
    heads = [datetime.datetime(2026 if m >= 4 else 2027, m, 1) for m in (4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3)] \
        if period == 'date' else PERIODS[period]
    items = v.items(n, rng)
    rows = [(k, [rng.randrange(0, 900) * 1000 for _ in heads]) for k in items]
    t = {'sheet': v.sheet('getsuji'), 'ka': v.word('ka'), 'heads': heads, 'period': period, 'rows': rows,
         'title': None, 'totals': True, 'old_chart': False, 'code_col': False, 'note_col': False, 'blank_at': None,
         'code': v.word('code'), 'note': v.word('tekiyo')}
    if kw.pop('title', False):
        t['title'] = rng.choice(TITLES).format(k=t['ka'])
    t.update(kw)
    return t


def write_before(path, t):
    wb = Workbook()
    ws = wb.active
    ws.title = t['sheet']
    h = 1
    if t['title']:
        ws.cell(1, 1, t['title']).font = Font(bold=True, size=14)
        h = 3
    c = 1
    if t['code_col']:
        ws.cell(h, c, t['code']).font = Font(bold=True)
        c += 1
    item_c = c
    ws.cell(h, item_c, t['ka']).font = Font(bold=True)
    for j, hd in enumerate(t['heads'], 1):
        cell = ws.cell(h, item_c + j, hd)
        cell.font = Font(bold=True)
        if isinstance(hd, datetime.datetime):
            cell.number_format = 'm"月"'
    last_c = item_c + len(t['heads'])
    if t['totals']:
        last_c += 1
        ws.cell(h, last_c, '合計').font = Font(bold=True)
    if t['note_col']:
        ws.cell(h, last_c + 1, t['note']).font = Font(bold=True)
    r = h + 1
    data_rows = []
    for i, (k, vals) in enumerate(t['rows']):
        if t['blank_at'] is not None and i == t['blank_at']:
            r += 1
        if t['code_col']:
            ws.cell(r, 1, 100 + i * 10)
        ws.cell(r, item_c, k)
        for j, val in enumerate(vals, 1):
            ws.cell(r, item_c + j, val).number_format = '#,##0'
        if t['totals']:
            ws.cell(r, item_c + len(vals) + 1, sum(vals)).number_format = '#,##0'
        if t['note_col']:
            ws.cell(r, last_c + 1, random.Random(i).choice(['', '見込み', '確定']))
        data_rows.append(r)
        r += 1
    if t['totals']:
        ws.cell(r, item_c, '合計').font = Font(bold=True)
        for j in range(len(t['heads']) + 1):
            ws.cell(r, item_c + 1 + j, sum((vals + [sum(vals)])[j] for _k, vals in t['rows'])).number_format = '#,##0'
    wb.save(path)
    return {'h': h, 'item_c': item_c, 'rows': data_rows, 'last': r}


def add_old_chart(xl, path, t, lay):
    wb = xl.Workbooks.Open(path)
    try:
        ws = wb.Worksheets(t['sheet'])
        co = ws.ChartObjects().Add(ws.Cells(lay['last'] + 2, 1).Left, ws.Cells(lay['last'] + 2, 1).Top, 300, 200)
        co.Name = '表の整理_前のグラフ'      # マクロは「表の整理_」で始まる名前のグラフ（自分が作った物）だけを消す
        co.Chart.ChartType = 51
        s = co.Chart.SeriesCollection().NewSeries()
        s.Values = ws.Range(ws.Cells(lay['rows'][0], lay['item_c'] + 1), ws.Cells(lay['rows'][0], lay['item_c'] + 3))
        wb.Save()
    finally:
        wb.Close(SaveChanges=False)


def build_truth(xl, before, truth, t, lay):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(t['sheet'])
        h, ic = lay['h'], lay['item_c']
        nh = len(t['heads'])
        for i in range(int(ws.ChartObjects().Count), 0, -1):
            ws.ChartObjects(i).Delete()
        order = sorted(range(len(t['rows'])), key=lambda i: -sum(t['rows'][i][1]))[:5]
        co = ws.ChartObjects().Add(ws.Cells(lay['last'] + 2, 1).Left, ws.Cells(lay['last'] + 2, 1).Top, 640, 320)
        co.Name = '表の整理_グラフ'      # マクロは「表の整理_」で始まる名前のグラフ（自分が作った物）だけを消す
        ch = co.Chart
        ch.ChartType = 65
        for i in range(int(ch.SeriesCollection().Count), 0, -1):
            ch.SeriesCollection(i).Delete()
        cats = ws.Range(ws.Cells(h, ic + 1), ws.Cells(h, ic + nh))
        for k, i in enumerate(order):
            r = lay['rows'][i]
            s = ch.SeriesCollection().NewSeries()
            s.Name = f"='{t['sheet']}'!{ws.Cells(r, ic).GetAddress(True, True, 1)}"
            s.Values = ws.Range(ws.Cells(r, ic + 1), ws.Cells(r, ic + nh))
            s.XValues = cats
            s.Format.Line.ForeColor.RGB = rgb(COLORS[k])
            s.Format.Line.Weight = 2.25
            s.MarkerSize = 6
        ch.HasTitle = True
        ch.ChartTitle.Text = t['title'] or f"{t['ka']}ごとの推移"
        ch.HasLegend = True
        ch.Legend.Position = -4152
        ax = ch.Axes(2)
        ax.TickLabels.NumberFormat = '#,##0'
        ax.HasMajorGridlines = True
        ch.ChartArea.Format.TextFrame2.TextRange.Font.Name = 'Meiryo UI'
        ch.ChartTitle.Format.TextFrame2.TextRange.Font.Size = 14
        ch.ChartTitle.Format.TextFrame2.TextRange.Font.Bold = True
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t in jobs:
            before, truth = os.path.join(out, f"{stem}_前.xlsx"), os.path.join(out, f"{stem}_正解.xlsx")
            lay = write_before(before, t)
            if t['old_chart']:
                add_old_chart(xl, before, t, lay)
            build_truth(xl, before, truth, t, lay)
            print('  ', stem, t['sheet'], t['ka'], t['period'], len(t['rows']))
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', build(s(1), rng.randint(2, 5), rng.choice(['month', 'quarter']))),
          ('未見2_表題と多い件数', build(s(2), rng.randint(6, 12), rng.choice(['half', 'year']), title=True, code_col=True)),
          ('未見3_合計なしと前のグラフ', build(s(3), 6, 'date', totals=False, old_chart=True)),
          ('未見4_ちょうど5件と空行', build(s(4), 5, 'month', blank_at=2, note_col=True))], out)
    return out


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('推移_本番', build(3, 4, 'month')),
          ('推移_試験A', build(13, 8, 'month')),
          ('推移_試験B', build(23, 5, 'half', title=True)),
          ('推移_試験C', build(33, 7, 'quarter', totals=False)),
          ('推移_試験D', build(43, 4, 'year', old_chart=True)),
          ('推移_試験E', build(53, 3, 'date')),
          ('推移_試験G', build(63, 12, 'month', title=True, code_col=True)),
          ('推移_試験H', build(73, 6, 'month', blank_at=3, note_col=True))], HERE)
    shutil.copyfile(os.path.join(HERE, '推移_本番_正解.xlsx'), os.path.join(HERE, '推移_試験F_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, '推移_本番_正解.xlsx'), os.path.join(HERE, '推移_試験F_正解.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
