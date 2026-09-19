# -*- coding: utf-8 -*-
"""表をテーブルにして集計行を出す（どの表でも動く形・2026-09-18 作り直し。前の版は職員名簿テーブル）。

列の決まり: 数の列（見出しに 番号・No・コード・年度 を含む列と率（%）の列は除く）だけ集計行で合計。ほかの列は集計なし。
テーブル名はシート名から作る（T＋シート名）。見出しの語・シート名・列の並びは表ごとに違う（vocab.py）。
正解の決まり:
  - 見出しの行＝値の入ったセルが 2 つ以上ある最初の行。範囲は見出しから表の最後の行まで（表の下の※注記と空行は入れない）。
  - スタイル TableStyleMedium2・行の縞模様あり・フィルタのボタンあり・集計行あり。
  - 既にテーブルがあればそれを直す（2 つ作らない・名前とスタイルを決まりどおりに）。計算列は足さない。
組: 本番 名簿 12 行／A 明細 300 行・率の列／B 表題つき／C 表の下に※注記／D 既にテーブル（別の名前・別のスタイル）／
  E 列の並びが違う・番号の列／F 撃った後にもう一度／G 途中に空行・日付の列
"""
import datetime
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
from vocab import Vocab   # noqa: E402

SEI = ['佐藤', '鈴木', '高橋', '田中', '伊藤', '渡辺', '山本', '中村', '小林', '加藤', '吉田', '山田']
MEI = ['一郎', '花子', '誠', '恵', '学', '悟', '由美', '健太', '真理', '修', '陽子', '大輔']
REQUEST = ('アクティブなシートの表をテーブルにしてください。テーブル名は T とシート名をつないだ名前、スタイルは TableStyleMedium2、'
           '行の縞模様あり、フィルタのボタンあり。範囲は見出しから表の最後の行まで（表の下の※注記と空行は入れない）。'
           '集計行を出し、数の列は合計（見出しに 番号・No・コード・年度 を含む列と率（%）の列、文字・日付の列は集計なし）。'
           '計算列は足さない。既にテーブルがあればそれを直す（2 つ作らない・名前とスタイルを決まりどおりに）')
PHRASES = ['この表をテーブルにして集計行も', '表をテーブル化して合計の集計行を出して', 'テーブルにしてください（集計行つき）',
           'テーブル化と集計行をお願い', 'いつものテーブルの形にして', '表をテーブルにして下に集計行']
TITLES = ['令和8年4月1日現在 {k}一覧', '令和8年度 {k}の明細', '{k}別の一覧表']


def build(seed, n, kind='meibo', **kw):
    rng = random.Random(seed)
    v = Vocab(seed)
    if kind == 'meibo':
        cols = [(v.word('ka'), 'text', 'items'), (v.word('name'), 'text', 'person'), ('採用日', 'date', ''),
                (v.word('kihon'), 'num', ''), (v.word('teate'), 'num', '')]
        sheet = v.sheet('meibo')
    else:
        cols = [(v.word('no'), 'no', ''), (v.word('date'), 'date', ''), (v.word('ka'), 'text', 'items'),
                (v.word('kamoku'), 'text', 'kamoku'), (v.word('amount'), 'num', '')]
        sheet = v.sheet('meisai')
    if kw.pop('rate_col', False):
        cols.append((v.word('rate'), 'rate', ''))
    if kw.pop('swap', False):
        cols = [cols[1], cols[0]] + cols[2:][::-1]
    rows = []
    items = v.items(max(3, min(n, 8)), rng)
    kamokus = v.kamokus(4, rng)
    for i in range(n):
        row = {}
        for name, role, pool in cols:
            if role == 'no':
                row[name] = 100 + i
            elif role == 'date':
                row[name] = datetime.datetime(rng.randint(2020, 2026), rng.randint(1, 12), rng.randint(1, 28))
            elif role == 'num':
                row[name] = rng.randrange(10, 900) * 1000
            elif role == 'rate':
                row[name] = round(rng.uniform(0.1, 1.0), 3)
            elif pool == 'person':
                row[name] = rng.choice(SEI) + ' ' + rng.choice(MEI)
            else:
                row[name] = rng.choice(kamokus if pool == 'kamoku' else items)
        rows.append(row)
    cols = [(name, role) for name, role, _p in cols]
    t = {'sheet': sheet, 'cols': cols, 'rows': rows, 'title': None, 'note': False, 'existing': False,
         'blank_at': None, 'kind': kind, 'ka': v.word('ka')}
    if kw.pop('title', False):
        t['title'] = rng.choice(TITLES).format(k=t['ka'])
    t.update(kw)
    return t


def write_before(path, t):
    wb = Workbook()
    ws = wb.active
    ws.title = t['sheet']
    h = 1
    if t['title']:
        ws.cell(1, 1, t['title']).font = Font(bold=True, size=14)
        h = 3
    for j, (name, _role) in enumerate(t['cols'], 1):
        ws.cell(h, j, name).font = Font(bold=True)
    r = h
    for i, row in enumerate(t['rows']):
        if t['blank_at'] is not None and i == t['blank_at']:
            r += 1
        r += 1
        for j, (name, role) in enumerate(t['cols'], 1):
            cell = ws.cell(r, j, row[name])
            if role == 'date':
                cell.number_format = 'yyyy/m/d'
            elif role == 'num':
                cell.number_format = '#,##0'
            elif role == 'rate':
                cell.number_format = '0.0%'
    last = r
    if t['note']:
        ws.cell(last + 2, 1, '※ 手当は扶養手当と住居手当の合計')
    wb.save(path)
    return {'h': h, 'last': last, 'ncols': len(t['cols'])}


def build_table(ws, t, lay, name=None, style='TableStyleMedium2', totals=True):
    rng = ws.Range(ws.Cells(lay['h'], 1), ws.Cells(lay['last'], lay['ncols']))
    lo = ws.ListObjects(1) if int(ws.ListObjects.Count) else ws.ListObjects.Add(1, rng, None, 1)
    lo.Name = name or ('T' + str(ws.Name))
    lo.TableStyle = style
    lo.ShowTableStyleRowStripes = True
    lo.ShowAutoFilter = True
    if not totals:
        return lo
    lo.ShowTotals = True
    sum_cols = [nm for nm, role in t['cols'] if role == 'num']
    for i in range(1, int(lo.ListColumns.Count) + 1):
        c = lo.ListColumns(i)
        c.TotalsCalculation = 1 if str(c.Name) in sum_cols else 0
    return lo


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            lay = write_before(before, t)
            if t['existing']:
                wb = xl.Workbooks.Open(before)
                try:
                    build_table(wb.Worksheets(t['sheet']), t, lay, name='テーブル1', style='TableStyleLight9', totals=False)
                    wb.Save()
                finally:
                    wb.Close(SaveChanges=False)
            wb = xl.Workbooks.Open(before)
            try:
                build_table(wb.Worksheets(t['sheet']), t, lay)
                wb.SaveAs(truth, FileFormat=51)
            finally:
                wb.Close(SaveChanges=False)
            print('  ', stem, t['sheet'], [c for c, _r in t['cols']])
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_名簿', build(s(1), rng.randint(5, 20))),
          ('未見2_明細と率', build(s(2), 30, 'meisai', rate_col=True, title=True)),
          ('未見3_注記と既存テーブル', build(s(3), 10, note=True, existing=True)),
          ('未見4_並び違いと空行', build(s(4), 12, 'meisai', swap=True, blank_at=4))], out)
    return out


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('テーブル_本番', build(3, 12)),
          ('テーブル_試験A', build(13, 300, 'meisai', rate_col=True)),
          ('テーブル_試験B', build(23, 10, title=True)),
          ('テーブル_試験C', build(33, 9, note=True)),
          ('テーブル_試験D', build(43, 8, existing=True)),
          ('テーブル_試験E', build(53, 11, 'meisai', swap=True)),
          ('テーブル_試験G', build(63, 10, blank_at=4))], HERE)
    shutil.copyfile(os.path.join(HERE, 'テーブル_本番_正解.xlsx'), os.path.join(HERE, 'テーブル_試験F_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, 'テーブル_本番_正解.xlsx'), os.path.join(HERE, 'テーブル_試験F_正解.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
