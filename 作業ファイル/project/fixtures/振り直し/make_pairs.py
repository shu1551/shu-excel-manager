# -*- coding: utf-8 -*-
"""決算統計の振り直し（鍛える回路の題材・2026-09-17）: 直す前と正解の 2 冊ずつを作る。

仕事: 「支出明細」の予算科目コードを「対応表」で決算統計区分に振り直し、区分ごとの支出額を集計する。
正解の決まり（人が手でやるときの形）:
  - 明細の右の空いた列に「決算統計区分」を足す（摘要など人の列は上書きしない）。対応表に無いコードは「対応なし」。
    既に「決算統計区分」の列があればそこを使う（2 回撃っても同じ結果）。
  - コードの前後の空白は無視して引く（明細のコードのセルは書き換えない）。
  - コードが空の行（途中の空行・最後の合計行）には区分を書かず、集計にも入れない。
  - 足した列の右に 1 列空けて集計表。見出しは明細の見出しと同じ行に「決算統計区分」「支出額」。
  - 区分は対応表に出てくる順（重複は 1 回）。明細に 1 件も無い区分も 0 で並べる。
  - その下に「対応なし」（無くても 0 で出す）、最後に「合計」（明細の行の合計。合計行は入れない）。
  - 対応表は見出しの語で列を探す（列の並び・ほかの列があっても同じ）。
組（依頼文は同じ）:
  本番 30 行／A 45 行・区分の順が違う・対応なし 3／B 表題つき・対応なし 0・0 円の区分／C 支出額の右に摘要／
  D B3 から始まる・途中に空行 2／E 最後に合計行／F 対応表が「区分・備考・コード」の並び／
  G コードの前後に空白／H 撃った後の表にもう一度（前＝本番の正解）
  py make_pairs.py  → このフォルダに *_前.xlsx と *_正解.xlsx
"""
import os
import random
import shutil

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))

# 節・細節 → (科目名, 区分)。同じ需用費でも修繕料は維持補修費＝節の名前では振れない（振り直しの本題）
SETSU = [
    ('01-00', '報酬', '人件費'), ('02-00', '給料', '人件費'), ('03-00', '職員手当等', '人件費'), ('04-00', '共済費', '人件費'),
    ('07-00', '報償費', '物件費'), ('08-00', '旅費', '物件費'), ('10-01', '需用費（消耗品費）', '物件費'),
    ('10-02', '需用費（燃料費）', '物件費'), ('10-05', '需用費（修繕料）', '維持補修費'), ('11-01', '役務費（通信運搬費）', '物件費'),
    ('12-00', '委託料', '物件費'), ('13-00', '使用料及び賃借料', '物件費'), ('14-00', '工事請負費', '普通建設事業費'),
    ('17-00', '備品購入費', '物件費'), ('18-01', '負担金', '補助費等'), ('18-02', '補助金', '補助費等'), ('19-00', '扶助費', '扶助費'),
]
MOKU = ['02-01-01', '02-01-02', '03-02-01', '07-01-03', '08-02-01']
ORDER1 = ['人件費', '物件費', '維持補修費', '扶助費', '補助費等', '普通建設事業費']
ORDER2 = ['物件費', '補助費等', '人件費', '普通建設事業費', '維持補修費', '扶助費']


def build(seed, rows, kubun_order, missing, zero_kubun=None):
    rng = random.Random(seed)
    pairs = [(f"{m}-{s}", name, k) for m in rng.sample(MOKU, 3) for s, name, k in SETSU]
    order = {k: i for i, k in enumerate(kubun_order)}
    pairs = [p for p in pairs if p[2] in order]
    pairs.sort(key=lambda p: (order[p[2]], p[0]))
    table = [(c, k) for c, _n, k in pairs]
    usable = [p for p in pairs if p[2] != zero_kubun]
    body = []
    for _ in range(rows - missing):
        c, name, k = rng.choice(usable)
        body.append([c, name, rng.randrange(3, 900) * 1000])
    for _ in range(missing):                       # 対応表に無いコード（新設の細節など）
        body.append([f"{rng.choice(MOKU)}-10-09", '需用費（その他）', rng.randrange(3, 90) * 1000])
    rng.shuffle(body)
    return table, body


def write_pair(stem, table, body, title=None, memo=False, r0=1, c0=1, blanks=(), total=False,
               map_layout='std', spaces=False, out=HERE, sheet='支出明細', map_sheet='対応表'):
    kmap = dict(table)
    kubun = []
    for _c, k in table:
        if k not in kubun:
            kubun.append(k)
    rng = random.Random(stem)
    shown = []                                     # 明細に書くコード（spaces なら前後に空白）
    for c, _n, _a in body:
        shown.append((' ' + c if rng.random() < 0.5 else c + '  ') if spaces and rng.random() < 0.6 else c)
    for truth in (False, True):
        wb = Workbook()
        ws = wb.active
        ws.title = sheet
        h = r0
        if title:
            ws.cell(r0, c0, title).font = Font(bold=True, size=14)
            h = r0 + 2
        base = ['予算科目コード', '科目名', '支出額'] + (['摘要'] if memo else [])
        kcol = c0 + len(base)                      # 足す列＝表の右の空いた列
        heads = base + (['決算統計区分'] if truth else [])
        for j, v in enumerate(heads):
            ws.cell(h, c0 + j, v).font = Font(bold=True)
        r = h + 1
        for i, ((c, name, amt), code) in enumerate(zip(body, shown)):
            if i in blanks:
                r += 1                             # 途中の空行
            ws.cell(r, c0, code).number_format = '@'
            ws.cell(r, c0 + 1, name)
            ws.cell(r, c0 + 2, amt).number_format = '#,##0'
            if memo:
                ws.cell(r, c0 + 3, f"{name[:2]}の支払 {i + 1}")
            if truth:
                ws.cell(r, kcol, kmap.get(c, '対応なし'))
            r += 1
        if total:
            ws.cell(r, c0 + 1, '合計').font = Font(bold=True)
            ws.cell(r, c0 + 2, sum(a for _c, _n, a in body)).number_format = '#,##0'
        if truth:
            sums = {k: 0 for k in kubun}
            none = 0
            for c, _n, amt in body:
                if c in kmap:
                    sums[kmap[c]] += amt
                else:
                    none += amt
            sc = kcol + 2                          # 集計表＝足した列の右に 1 列空けて
            ws.cell(h, sc, '決算統計区分').font = Font(bold=True)
            ws.cell(h, sc + 1, '支出額').font = Font(bold=True)
            rr = h + 1
            for k in kubun + ['対応なし']:
                ws.cell(rr, sc, k)
                ws.cell(rr, sc + 1, sums.get(k, none if k == '対応なし' else 0)).number_format = '#,##0'
                rr += 1
            ws.cell(rr, sc, '合計')
            ws.cell(rr, sc + 1, sum(a for _c, _n, a in body)).number_format = '#,##0'
        mt = wb.create_sheet(map_sheet)
        if map_layout == 'swapped':
            cols = ['決算統計区分', '備考', '予算科目コード']
        else:
            cols = ['予算科目コード', '決算統計区分']
        for j, v in enumerate(cols, 1):
            mt.cell(1, j, v).font = Font(bold=True)
        for i, (c, k) in enumerate(table, 2):
            row = {'予算科目コード': c, '決算統計区分': k, '備考': '' if i % 3 else '令和7年度から'}
            for j, v in enumerate(cols, 1):
                cell = mt.cell(i, j, row[v] or None)
                if v == '予算科目コード':
                    cell.number_format = '@'
        wb.save(os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx"))


def unseen(seed):
    """AI に見せない表（種ごと）。同じ癖を種を変えて作り、癖を全部重ねた表とシート名の違う表を足す。
    置き場は 未見_種N（鍛えるときの --test に入れない＝出来たマクロの本当の力を測る）。"""
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    o = [ORDER1, ORDER2]
    s = lambda k: seed * 10 + k      # noqa: E731
    write_pair('未見1_標準', *build(s(1), rng.randint(15, 60), o[rng.randrange(2)], rng.randint(0, 4)), out=out)
    write_pair('未見2_表題', *build(s(2), rng.randint(15, 40), o[rng.randrange(2)], rng.randint(0, 3), zero_kubun='扶助費'),
               title='令和8年度 支出明細（建設課）', out=out)
    write_pair('未見3_摘要', *build(s(3), rng.randint(15, 40), o[rng.randrange(2)], rng.randint(0, 3)), memo=True, out=out)
    write_pair('未見4_位置と空行', *build(s(4), 30, o[rng.randrange(2)], 2), r0=rng.randint(2, 5), c0=rng.randint(2, 4),
               blanks=tuple(sorted(rng.sample(range(2, 28), 3))), out=out)
    write_pair('未見5_合計行', *build(s(5), rng.randint(15, 40), o[rng.randrange(2)], 1), total=True, out=out)
    write_pair('未見6_対応表の並び', *build(s(6), 25, o[rng.randrange(2)], 2), map_layout='swapped', out=out)
    write_pair('未見7_空白', *build(s(7), 25, o[rng.randrange(2)], 2), spaces=True, out=out)
    write_pair('未見8_全部入り', *build(s(8), 35, o[rng.randrange(2)], 3), title='令和8年度 支出明細（福祉課）', memo=True,
               r0=2, c0=3, blanks=(4, 19), total=True, map_layout='swapped', spaces=True, out=out)
    write_pair('未見9_シート名', *build(s(9), 20, o[rng.randrange(2)], 1), sheet='R8支出', map_sheet='対応表（R8）', out=out)
    write_pair('未見10_大量', *build(s(10), 3000, o[rng.randrange(2)], 25), out=out)
    return out


if __name__ == '__main__':
    import sys
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        sys.exit(0)
    main = build(17, 30, ORDER1, 2)
    write_pair('振り直し_本番', *main)
    write_pair('振り直し_試験A', *build(29, 45, ORDER2, 3))
    write_pair('振り直し_試験B', *build(41, 18, ORDER1, 0, zero_kubun='扶助費'), title='令和7年度 支出明細（総務課）')
    write_pair('振り直し_試験C', *build(53, 24, ORDER2, 1), memo=True)
    write_pair('振り直し_試験D', *build(67, 26, ORDER1, 2), r0=3, c0=2, blanks=(7, 15))
    write_pair('振り直し_試験E', *build(79, 22, ORDER2, 1), total=True)
    write_pair('振り直し_試験F', *build(83, 28, ORDER1, 2), map_layout='swapped')
    write_pair('振り直し_試験G', *build(97, 25, ORDER2, 2), spaces=True)
    # H: 撃った後の表にもう一度撃つ（前＝本番の正解・正解＝同じ）
    shutil.copyfile(os.path.join(HERE, '振り直し_本番_正解.xlsx'), os.path.join(HERE, '振り直し_試験H_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, '振り直し_本番_正解.xlsx'), os.path.join(HERE, '振り直し_試験H_正解.xlsx'))
    print('ok')
