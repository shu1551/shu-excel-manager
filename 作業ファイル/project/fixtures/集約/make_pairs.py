# -*- coding: utf-8 -*-
"""各課のシートを「集約」シートにまとめる（鍛える回路の題材・2026-09-17。照会の回答を課ごとのシートで集める仕事）。

仕事: ブックの各課のシート（事業名・予算額・執行額の表）を、「集約」シートに 1 行ずつまとめる。
正解の決まり:
  - 「集約」シートの見出し（課名・事業名・予算額・執行額・執行率）の下に、シートのタブの順・各シートの行の順で並べる。
    課名＝シート名。執行率＝執行額÷予算額（丸めない。予算額が 0 なら 0）。
  - 課のシート＝見出しの行に「事業名」「予算額」「執行額」がそろうシート。「記入例」「説明」のシートは除く。
    列の並び・表題の行はシートごとに違ってよい（見出しの語で探す）。
  - 事業名が空の行・事業名が「合計」「計」の行は入れない。金額はカンマ・全角・前後の空白があっても数で読む。
  - 最後に「合計」の行（課名の列に「合計」・予算額と執行額の合計・その執行率）。
  - 集約シートに前の結果が残っていたら消してから書く（2 回撃っても同じ）。見出しの右や下の注記には触らない…は求めない。
組（依頼文は同じ）: 本番 4 課／A 6 課・行数ばらばら／B 列の並びが違う課／C 課のシートに表題の行／D 記入例と説明のシート／
  E 集約に前の結果（行が多い）が残っている／F 金額が文字（カンマ・全角）／G 行の無い課と合計行のある課／
  H 集約の見出しが 3 行目（表題つき）／I 撃った後にもう一度
  py make_pairs.py            → *_前.xlsx と *_正解.xlsx
  py make_pairs.py --unseen N → 未見_種N
"""
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
KA = ['総務課', '財政課', '企画課', '税務課', '福祉課', '建設課', '農林課', '教育総務課', '観光課', '環境課']
JIGYO = ['庁舎管理', '広報発行', '防災訓練', '道路補修', '除雪', '健康診断', '子育て支援', '観光案内所', '農道整備',
         '図書購入', '学校給食', '水道管更新', '移住相談', '文化祭', 'ごみ収集', '公園整備', '空き家対策', '高齢者見守り']
_ZEN = str.maketrans('0123456789,', '０１２３４５６７８９，')
HEAD = ['課名', '事業名', '予算額', '執行額', '執行率']


def build(seed, n_ka, rows=(2, 8)):
    rng = random.Random(seed)
    data = []
    for ka in rng.sample(KA, n_ka):
        items = []
        k = rng.randint(*rows)
        names = rng.sample(JIGYO, k) if k <= len(JIGYO) else [f"{rng.choice(JIGYO)}{i + 1}" for i in range(k)]
        for j in names:
            b = rng.randrange(1, 500) * 10000
            items.append((j, b, min(b, rng.randrange(0, 520) * 10000) if rng.random() > 0.1 else 0))
        data.append((ka, items))
    return data


def write_pair(stem, data, out=HERE, swapped=(), titled=(), extras=False, stale=0, text_amount=False,
               empty_ka=False, total_rows=(), sum_title=False, rng_seed=None):
    rng = random.Random(rng_seed or stem)
    if empty_ka and data:
        data = data[:1] + [(data[1][0], [])] + data[2:]
    for truth in (False, True):
        wb = Workbook()
        sm = wb.active
        sm.title = '集約'
        h = 3 if sum_title else 1
        if sum_title:
            sm.cell(1, 1, '令和8年度 事業執行状況（全課）').font = Font(bold=True, size=14)
        for j, v in enumerate(HEAD, 1):
            sm.cell(h, j, v).font = Font(bold=True)
        r = h + 1
        if truth:
            tb = te = 0
            for ka, items in data:
                for jg, b, e in items:
                    sm.cell(r, 1, ka)
                    sm.cell(r, 2, jg)
                    sm.cell(r, 3, b)
                    sm.cell(r, 4, e)
                    sm.cell(r, 5, (e / b) if b else 0)
                    tb += b
                    te += e
                    r += 1
            sm.cell(r, 1, '合計')
            sm.cell(r, 3, tb)
            sm.cell(r, 4, te)
            sm.cell(r, 5, (te / tb) if tb else 0)
        elif stale:
            for i in range(stale):                   # 前の結果（今回より多い行）
                sm.cell(r + i, 1, KA[i % len(KA)])
                sm.cell(r + i, 2, f'旧事業{i + 1}')
                sm.cell(r + i, 3, 1000 * (i + 1))
                sm.cell(r + i, 4, 500 * (i + 1))
                sm.cell(r + i, 5, 0.5)
            sm.cell(r + stale, 1, '合計')
        if extras:
            ex = wb.create_sheet('説明')
            ex.cell(1, 1, '各課は自分のシートに事業ごとの予算額と執行額を記入してください。')
            ex2 = wb.create_sheet('記入例')
            for j, v in enumerate(['事業名', '予算額', '執行額'], 1):
                ex2.cell(1, j, v)
            ex2.cell(2, 1, '（例）庁舎清掃')
            ex2.cell(2, 2, 1200000)
            ex2.cell(2, 3, 800000)
        for k, (ka, items) in enumerate(data):
            ws = wb.create_sheet(ka)
            hh = 1
            if k in titled:
                ws.cell(1, 1, f'{ka} 事業一覧（令和8年度）').font = Font(bold=True)
                ws.cell(2, 1, '単位：円')
                hh = 4
            cols = ['執行額', '事業名', '予算額'] if k in swapped else ['事業名', '予算額', '執行額']
            for j, v in enumerate(cols, 1):
                ws.cell(hh, j, v).font = Font(bold=True)
            rr = hh + 1
            for jg, b, e in items:
                vals = {'事業名': jg, '予算額': b, '執行額': e}
                if text_amount:
                    for key in ('予算額', '執行額'):
                        vals[key] = f"{vals[key]:,}".translate(_ZEN) if rng.random() < 0.5 else f" {vals[key]:,} "
                for j, v in enumerate(cols, 1):
                    ws.cell(rr, j, vals[v])
                rr += 1
            if k in total_rows:
                ws.cell(rr, cols.index('事業名') + 1, '合計')
                ws.cell(rr, cols.index('予算額') + 1, sum(b for _j, b, _e in items))
                ws.cell(rr, cols.index('執行額') + 1, sum(e for _j, _b, e in items))
        wb.save(os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx"))


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    write_pair('未見1_標準', build(s(1), rng.randint(2, 7)), out=out)
    write_pair('未見2_並びと表題', build(s(2), 5), swapped=(1, 3), titled=(0, 3), out=out)
    write_pair('未見3_説明と記入例', build(s(3), 4), extras=True, out=out)
    write_pair('未見4_前の結果', build(s(4), 3), stale=rng.randint(20, 40), out=out)
    write_pair('未見5_文字の金額', build(s(5), 4), text_amount=True, out=out)
    write_pair('未見6_空と合計行', build(s(6), 5), empty_ka=True, total_rows=(0, 3), out=out)
    write_pair('未見7_全部入り', build(s(7), 6), swapped=(2,), titled=(1, 4), extras=True, stale=30, text_amount=True,
               empty_ka=True, total_rows=(0, 5), sum_title=True, out=out)
    write_pair('未見8_大量', build(s(8), 10, rows=(150, 300)), out=out)
    return out


if __name__ == '__main__':
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        sys.exit(0)
    write_pair('集約_本番', build(5, 4))
    write_pair('集約_試験A', build(15, 6, rows=(1, 12)))
    write_pair('集約_試験B', build(25, 4), swapped=(1, 2))
    write_pair('集約_試験C', build(35, 3), titled=(0, 2))
    write_pair('集約_試験D', build(45, 4), extras=True)
    write_pair('集約_試験E', build(55, 3), stale=25)
    write_pair('集約_試験F', build(65, 4), text_amount=True)
    write_pair('集約_試験G', build(75, 5), empty_ka=True, total_rows=(2,))
    write_pair('集約_試験H', build(85, 3), sum_title=True)
    shutil.copyfile(os.path.join(HERE, '集約_本番_正解.xlsx'), os.path.join(HERE, '集約_試験I_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, '集約_本番_正解.xlsx'), os.path.join(HERE, '集約_試験I_正解.xlsx'))
    print('ok')
