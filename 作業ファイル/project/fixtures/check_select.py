# -*- coding: utf-8 -*-
"""入口の「列を選ぶ」が当たるかを確かめる（純 Python・Excel なし・2026-09-18）。

お題の <名前>_選択.txt（人が選んでいる列の番地）と、そのお題の表の見出しから作った人の言い方の依頼文で、
vbam_prefire.pick_columns が同じ列を選ぶかを見る。
  py check_select.py <フォルダ> …          （省略時は 選択.txt を持つフォルダを全部）
"""
import glob
import os
import re
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'python_scripts')))
import vbam_core   # noqa: E402
vbam_core.setup_encoding()
import vbam_prefire as vp   # noqa: E402
from openpyxl import load_workbook   # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))


def header_cells(path):
    """前の表の見出しの行 → [(列番号, 見出し)]（値が 2 つ以上並ぶ最初の行）。"""
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    for row in ws.iter_rows(min_row=1, max_row=12):
        cells = [(c.column, str(c.value).strip()) for c in row if c.value not in (None, '')]
        if len(cells) >= 2:
            return cells
    return []


def main():
    folders = sys.argv[1:] or sorted({os.path.dirname(p) for p in glob.glob(os.path.join(HERE, '*', '*_選択.txt'))})
    ok = bad = 0
    for folder in folders:
        for sel_path in sorted(glob.glob(os.path.join(folder, '*_選択.txt'))):
            stem = sel_path[:-len('_選択.txt')]
            before = stem + '_前.xlsx'
            if not os.path.isfile(before):
                continue
            want = [vp._col_of_letter(a) if hasattr(vp, '_col_of_letter') else
                    sum((ord(ch) - 64) * 26 ** i for i, ch in enumerate(reversed(re.match(r'[A-Z]+', a).group())))
                    for a in [x.strip() for x in open(sel_path, encoding='utf-8').read().split(',') if x.strip()]]
            cells = header_cells(before)
            names = {c: n for c, n in cells}
            # 人の言い方（選んだ列の見出しを、選んだ順に並べた依頼文）
            req = 'と'.join(names.get(c, '?') for c in want) + 'で作って'
            got = vp.pick_columns(req, cells, (len(want), len(want)))
            mark = 'OK ' if got == want else 'NG '
            if got == want:
                ok += 1
            else:
                bad += 1
            print(f"  {mark}{os.path.basename(stem)}: 依頼「{req}」→ {got}（期待 {want}）")
    print(f"入口の列選び: {ok + bad} 件　当たり {ok}・外れ {bad}")
    return 1 if bad else 0


if __name__ == '__main__':
    sys.exit(main())
