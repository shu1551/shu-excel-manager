# -*- coding: utf-8 -*-
"""実物らしい表・第 2 弾（2026-09-17 夜）: 第 1 弾（make_real.py）で鍛え直したマクロに、まだ当てていない実物の癖を当てる。
第 1 弾の癖に合わせて直しただけか、実物の表に通用する直し方になったかを見る。

  振り直し: 金額に「円」の文字・空行が 2 行続く・コードの前後に NBSP（Web の画面から貼った空白）・対応表の区分が結合セル・
            見出しが 2 行目で 1 行目に単位・全部入り
  突合: 新の金額に「円」・番号に NBSP・旧の番号が数で 000000 表示・旧に空行 2 行・新の見出しが 2 行目・全部入り
  集約: 課のシートがテーブル・途中に空行・金額に「円」・全部入り
  帳票: 受付番号などが結合なし・電話のある件とない件・備考の列・申請額に「円」の文字（帳票のまま写す）・全部入り
  月次: ISO の日付文字（2026-04-15）・空行が 2 行続く・下に「集計」の行（SUBTOTAL）・金額に「円」・全部入り

  py make_real2.py [種]  → fixtures\\<型>\\未見_実物2（種 1）／未見_実物2種N
"""
import datetime
import importlib.util
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
_spec = importlib.util.spec_from_file_location('make_real', os.path.join(HERE, 'make_real.py'))
MR = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(MR)
FR, TG, SY, CH, GJ = MR.FR, MR.TG, MR.SY, MR.CH, MR.GJ
post, data_rows, L = MR.post, MR.data_rows, MR.L
NBSP = ' '
SEED = 1


def _out(kind):
    p = os.path.join(HERE, kind, '未見_実物2' if SEED == 1 else f'未見_実物2種{SEED}')
    os.makedirs(p, exist_ok=True)
    for f in os.listdir(p):
        if f.endswith('.xlsx'):
            os.remove(os.path.join(p, f))
    return p


def yen(ws, rows, col, key, p=0.5):
    r2 = random.Random(key)
    for r in rows:
        v = ws.cell(r, col).value
        if isinstance(v, (int, float)) and r2.random() < p:
            ws.cell(r, col).value = f"{int(v):,}円"


def nbsp(ws, rows, col, key, p=0.4):
    r2 = random.Random(key)
    for r in rows:
        v = ws.cell(r, col).value
        if v is not None and r2.random() < p:
            ws.cell(r, col).value = (NBSP + str(v)) if r2.random() < 0.3 else (str(v) + NBSP)


def double_blank(ws, row):
    """row（空行）の下にもう 1 行の空行を入れる（その行より下が 1 行ずつ下がる）。"""
    ws.insert_rows(row + 1)


def merge_kubun(mt, first, col=2):
    """対応表の区分の列で、同じ区分が続く所を縦に結合する（値は先頭だけ）。"""
    r = first
    last = mt.max_row
    while r <= last:
        v = mt.cell(r, col).value
        e = r
        while e + 1 <= last and mt.cell(e + 1, col).value == v and v is not None:
            e += 1
        if e > r:
            mt.merge_cells(start_row=r, start_column=col, end_row=e, end_column=col)
        r = e + 1


# ============================================================ 振り直し
def furi(seed):
    out = _out('振り直し')
    rng = random.Random(seed)
    s = lambda k: seed * 1000 + k      # noqa: E731

    def B(k, n=None, miss=None):
        order = (FR.ORDER1, FR.ORDER2)[rng.randrange(2)]
        return FR.build(s(k), n or rng.randint(25, 50), order, rng.randint(1, 3) if miss is None else miss)

    stem = '実物2_01_金額に円'
    t, b = B(1)
    FR.write_pair(stem, t, b, out=out)
    rows1, _ = data_rows(1, len(b), ())
    post(out, stem, lambda wb, truth: yen(wb['支出明細'], rows1, 3, stem))

    stem = '実物2_02_空行が2行続く'
    t, b = B(2, n=rng.randint(30, 45))
    bl = (14,)
    FR.write_pair(stem, t, b, blanks=bl, out=out)
    rows2, _ = data_rows(1, len(b), bl)
    post(out, stem, lambda wb, truth: double_blank(wb['支出明細'], rows2[14] - 1))

    stem = '実物2_03_コードにNBSP'
    t, b = B(3)
    FR.write_pair(stem, t, b, out=out)
    rows3, _ = data_rows(1, len(b), ())
    post(out, stem, lambda wb, truth: nbsp(wb['支出明細'], rows3, 1, stem))

    stem = '実物2_04_対応表の区分が結合'
    t, b = B(4)
    FR.write_pair(stem, t, b, out=out)
    post(out, stem, lambda wb, truth: merge_kubun(wb['対応表'], 2))

    stem = '実物2_05_見出しが2行目'
    t, b = B(5)
    FR.write_pair(stem, t, b, out=out)

    def f05(wb, truth):
        ws = wb['支出明細']
        ws.insert_rows(1)
        ws.cell(1, 3, '（単位：円）')
    post(out, stem, f05)

    stem = '実物2_06_全部入り'
    t, b = B(6, n=40, miss=3)
    bl6 = (18,)
    FR.write_pair(stem, t, b, blanks=bl6, total=True, out=out)
    rows6, _ = data_rows(1, len(b), bl6)

    def f06(wb, truth):
        ws = wb['支出明細']
        yen(ws, rows6, 3, stem)
        nbsp(ws, rows6, 1, stem)
        MR.zen(ws, rows6, stem + 'z', 1, p=0.2)
        merge_kubun(wb['対応表'], 2)
        double_blank(ws, rows6[18] - 1)                  # ここから下が 1 行下がる
        ws.insert_rows(1)
        ws.cell(1, 3, '（単位：円）')
    post(out, stem, f06)
    return out


# ============================================================ 突合
def totsu(seed):
    out = _out('突合')
    rng = random.Random(seed)
    s = lambda k: seed * 1000 + k      # noqa: E731

    def B(k, n=None):
        return TG.build(s(k), n or rng.randint(22, 40), rng.randint(1, 4), rng.randint(1, 3), rng.randint(1, 3))

    stem = '実物2_01_新の金額に円'
    old, new = B(1)
    TG.write_pair(stem, old, new, out=out)
    post(out, stem, lambda wb, truth: yen(wb['新システム'], range(2, 2 + len(new)), 2, stem))

    stem = '実物2_02_番号にNBSP'
    old, new = B(2)
    TG.write_pair(stem, old, new, out=out)
    rows2, _ = data_rows(1, len(old), ())
    post(out, stem, lambda wb, truth: nbsp(wb['旧システム'], rows2, 1, stem))

    stem = '実物2_03_旧の番号が数で0詰め表示'
    old, new = B(3)
    TG.write_pair(stem, old, new, out=out)
    rows3, _ = data_rows(1, len(old), ())

    def f03(wb, truth):
        ws = wb['旧システム']
        for r in rows3:
            c = ws.cell(r, 1)
            c.value = int(c.value)
            c.number_format = '000000'
    post(out, stem, f03)

    stem = '実物2_04_旧に空行2行'
    old, new = B(4, n=rng.randint(26, 40))
    bl = (12,)
    TG.write_pair(stem, old, new, blanks=bl, out=out)
    rows4, _ = data_rows(1, len(old), bl)
    post(out, stem, lambda wb, truth: double_blank(wb['旧システム'], rows4[12] - 1))

    stem = '実物2_05_新の見出しが2行目'
    old, new = B(5)
    TG.write_pair(stem, old, new, out=out)

    def f05(wb, truth):
        ns = wb['新システム']
        ns.insert_rows(1)
        ns.cell(1, 2, '（単位：円）')
    post(out, stem, f05)

    stem = '実物2_06_全部入り'
    old, new = B(6, n=36)
    bl6 = (15,)
    TG.write_pair(stem, old, new, blanks=bl6, total=True, memo=True, out=out)
    rows6, _ = data_rows(1, len(old), bl6)

    def f06(wb, truth):
        ws = wb['旧システム']
        nbsp(ws, rows6, 1, stem)
        r2 = random.Random(stem + 'n')
        for r in rows6:
            if r2.random() < 0.3:
                c = ws.cell(r, 1)
                if str(c.value).isdigit():
                    c.value = int(c.value)
                    c.number_format = '000000'
        double_blank(ws, rows6[15] - 1)
        ns = wb['新システム']
        yen(ns, range(2, 2 + len(new)), 2, stem)
        ns.insert_rows(1)
        ns.cell(1, 2, '（単位：円）')
    post(out, stem, f06)
    return out


# ============================================================ 集約
def ka_sheet2(ws, ka, items, key, table=False, gap=False, yen_p=0.0, name=None):
    rng = random.Random(key)
    cols = ['事業名', '予算額', '執行額']
    for j, c in enumerate(cols, 1):
        ws.cell(1, j, c).font = Font(bold=True)
    r = 2
    for i, (jg, b, e) in enumerate(items):
        if gap and i == len(items) // 2 and len(items) >= 2:
            r += 1
        for j, v in enumerate((jg, b, e), 1):
            if j > 1 and rng.random() < yen_p:
                v = f"{v:,}円"
            ws.cell(r, j, v)
        r += 1
    if table and items:
        MR.add_table(ws, f"A1:C{r - 1}", name or f"T{abs(hash(key)) % 100000}")


def shuyaku(seed):
    out = _out('集約')
    rng = random.Random(seed)
    s = lambda k: seed * 1000 + k      # noqa: E731
    D = lambda k, n=None: SY.build(s(k), n or rng.randint(3, 6), rows=(3, 9))  # noqa: E731

    def write(stem, data, opts, extras=False):
        for truth in (False, True):
            wb = Workbook()
            sm = wb.active
            sm.title = '集約'
            for j, v in enumerate(SY.HEAD, 1):
                sm.cell(1, j, v).font = Font(bold=True)
            if truth:
                r, tb, te = 2, 0, 0
                for ka, items in data:
                    for jg, b, e in items:
                        for j, v in enumerate((ka, jg, b, e, (e / b) if b else 0), 1):
                            sm.cell(r, j, v)
                        tb += b
                        te += e
                        r += 1
                for j, v in ((1, '合計'), (3, tb), (4, te), (5, (te / tb) if tb else 0)):
                    sm.cell(r, j, v)
            if extras:
                wb.create_sheet('説明').cell(1, 1, '各課は自分のシートに記入してください。')
            for k, (ka, items) in enumerate(data):
                ka_sheet2(wb.create_sheet(ka), ka, items, f"{stem}/{k}", name=f"T課{k + 1}", **opts(k))
            wb.save(os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx"))

    write('実物2_01_課のシートがテーブル', D(1), lambda k: {'table': True})
    write('実物2_02_途中に空行', D(2), lambda k: {'gap': True})
    write('実物2_03_金額に円', D(3), lambda k: {'yen_p': 0.5})
    write('実物2_04_全部入り', D(4, n=6), lambda k: {'table': k % 2 == 0, 'gap': k % 3 != 1, 'yen_p': 0.3}, extras=True)
    return out


# ============================================================ 帳票
def ch_write2(stem, recs, out, merged=True, note_col=False, phone_some=False, yen_amount=False):
    """電話のある件とない件（件ごとに行数が違う）・備考の列・結合なし・申請額に「円」の文字。"""
    layout = ('受付番号', '申請者', '申請額', '受付日') + (('備考',) if note_col else ())
    sub = ('団体名', '代表者')
    pos, last_col = MR.ch_cols(layout, 1, sub)
    lab_c, val_c = pos['団体名'], pos['代表者']
    top = sorted(pos, key=lambda x: pos[x])
    labels = ['所在地'] + (['電話'] if phone_some else [])
    fields = top + labels
    for truth in (False, True):
        wb = Workbook()
        ws = wb.active
        ws.title = CH.SHEET
        ws.cell(1, 1, '令和8年度 地域づくり活動補助金 申請受付簿').font = Font(bold=True, size=14)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
        h = 2
        MR.ch_head(ws, h, layout, 1, sub)
        r = h + 2
        for rec in recs:
            labs = ['所在地'] + (['電話'] if phone_some and rec.get('電話') else [])
            k = 1 + len(labs)
            for name, c in pos.items():
                v = rec.get(name, '')
                if v != '':
                    if name == '申請額' and yen_amount:
                        v = f"{v:,}円"
                    cell = ws.cell(r, c, v)
                    if name == '受付日':
                        cell.number_format = 'yyyy/m/d'
                if name in sub or not merged:
                    continue
                if k > 1:
                    ws.merge_cells(start_row=r, start_column=c, end_row=r + k - 1, end_column=c)
            for j, lab in enumerate(labs, 1):
                ws.cell(r + j, lab_c, lab)
                ws.cell(r + j, val_c, rec.get(lab, ''))
            for rr in range(r, r + k):
                for cc in range(1, last_col + 1):
                    ws.cell(rr, cc).border = MR.BOX
            r += k
        lc = last_col + 2
        if truth:
            for j, f in enumerate(fields):
                ws.cell(h, lc + j, f)
            for i, rec in enumerate(recs, 1):
                for j, f in enumerate(fields):
                    v = rec.get(f, '')
                    if f == '申請額' and yen_amount and v != '':
                        v = f"{v:,}円"
                    if v != '':
                        ws.cell(h + i, lc + j, v)
        for c in range(1, last_col + 1):
            ws.column_dimensions[L(c)].width = 16
        wb.save(os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx"))


def chohyo(seed):
    out = _out('帳票')
    rng = random.Random(seed)
    s = lambda k: seed * 1000 + k      # noqa: E731
    R = lambda k, n=None, **kw: CH.build(s(k), n or rng.randint(5, 14), **kw)  # noqa: E731

    ch_write2('実物2_01_結合なし', R(1), out, merged=False)
    recs = R(2, phone=True)
    for x in recs:
        if random.Random(x['受付番号']).random() < 0.45:
            x['電話'] = ''
    ch_write2('実物2_02_電話のある件とない件', recs, out, phone_some=True)
    recs = R(3)
    for x in recs:
        v = random.Random(x['受付番号'] + 'b').random()
        x['備考'] = '書類不備（再提出）' if v < 0.2 else ('電話で確認済み' if v < 0.4 else '')
    ch_write2('実物2_03_備考の列', recs, out, note_col=True)
    ch_write2('実物2_04_申請額に円の文字', R(4), out, yen_amount=True)
    recs = R(5, 16, phone=True)
    for x in recs:
        r2 = random.Random(x['受付番号'] + 'all')
        if r2.random() < 0.4:
            x['電話'] = ''
        x['備考'] = '取下げの相談あり' if r2.random() < 0.2 else ''
    ch_write2('実物2_05_全部入り', recs, out, merged=False, note_col=True, phone_some=True, yen_amount=True)
    return out


# ============================================================ 月次
def getsuji(seed):
    out = _out('月次')
    rng = random.Random(seed)
    s = lambda k: seed * 1000 + k      # noqa: E731
    Bd = lambda k, n=None, **kw: GJ.build(s(k), n or rng.randint(40, 90), **kw)  # noqa: E731

    lines = MR.plain(Bd(1))
    r2 = random.Random('iso')
    for ln in lines:
        if r2.random() < 0.6:
            d = ln[2][0]
            ln[1]['支出日'] = f"{d:%Y-%m-%d}"
    MR.gj_write('実物2_01_ISOの日付', lines, out)

    lines = MR.plain(Bd(2))
    k = len(lines) // 2
    lines = lines[:k] + [('b',), ('b',)] + lines[k:]
    MR.gj_write('実物2_02_空行が2行続く', lines, out)

    rows = Bd(3)
    lines = MR.plain(rows)
    n = len(lines)
    lines.append(('x', {'摘要': '集計', '支出額': f"=SUBTOTAL(9,E2:E{n + 1})"}))
    MR.gj_write('実物2_03_下に集計の行', lines, out)

    lines = MR.plain(Bd(4))
    r2 = random.Random('yen')
    for ln in lines:
        if r2.random() < 0.5:
            ln[1]['支出額'] = f"{ln[2][2]:,}円"
    MR.gj_write('実物2_04_金額に円', lines, out)

    rows = Bd(5, 90, n_ka=6)
    lines = MR.plain(rows)
    r2 = random.Random('all2')
    for ln in lines:
        d, _k, a = ln[2]
        x = r2.random()
        if x < 0.3:
            ln[1]['支出日'] = f"{d:%Y-%m-%d}"
        if r2.random() < 0.3:
            ln[1]['支出額'] = f"{a:,}円"
    k = len(lines) // 3
    lines = lines[:k] + [('b',), ('b',)] + lines[k:]
    lines.append(('x', {'摘要': '集計', '支出額': f"=SUBTOTAL(9,E2:E{len(lines) + 1})"}))
    MR.gj_write('実物2_05_全部入り', lines, out)
    return out


if __name__ == '__main__':
    seed = SEED = int(sys.argv[1]) if len(sys.argv) > 1 and sys.argv[1].isdigit() else 1
    only = [a for a in sys.argv[1:] if not a.isdigit()]
    for name, fn in (('振り直し', furi), ('突合', totsu), ('集約', shuyaku), ('帳票', chohyo), ('月次', getsuji)):
        if only and name not in only:
            continue
        p = fn(seed)
        print(f"{name}: {len([f for f in os.listdir(p) if f.endswith('_前.xlsx')])} 組 → {p}")
