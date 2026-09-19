# -*- coding: utf-8 -*-
"""パワークエリで、テーブルの文字の列の前後の空白・全角の空白・改行を掃除した表を作る（G パワークエリ 61・2026-09-18 第二期）。

列の決まり: 選ぶ列は無い。元＝アクティブなシートのテーブル（最初のもの）。
正解の決まり（依頼文に全部書く）:
  - 文字の列＝本文の値が全部文字の列（空は数えない）。その列だけ、全角の空白を半角の空白に替え（Text.Replace）、改行などの制御文字を取り（Text.Clean）、
    前後の空白を取る（Text.Trim）。数の列はそのまま。
  - クエリ名・シート名・テーブル名＝元のテーブル名＋「_掃除」。新しいシートの A1 に読み込む。同じ名前のクエリ・シートがあれば消して作り直す。
組: 本番／A 100 行／B 改行の入った値／C 全角の空白だけ／D 先頭 0 の文字の番号／E 列の並び違い／F 撃った後にもう一度／G 空欄／H 前のクエリとシート
"""
import importlib.util
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
_spec = importlib.util.spec_from_file_location('tate', os.path.join(HERE, '..', '縦持ち', 'make_pairs.py'))
tate = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(tate)
from vocab import Vocab   # noqa: E402

REQUEST = ('パワークエリで、アクティブなシートのテーブル（最初のもの）の文字を掃除した表を作ってください。'
           '文字の列＝本文の値が全部文字の列（空のセルは数えない）。その列だけ、全角の空白を半角の空白に替え（Text.Replace）、'
           '改行などの制御文字を取り（Text.Clean）、前後の空白を取る（Text.Trim）。数の列はそのまま。'
           'クエリの名前・新しいシートの名前・読み込むテーブルの名前は、元のテーブル名の後ろに「_掃除」。新しいシートの A1 に読み込む。'
           '同じ名前のクエリやシートが既にあれば消して作り直す。M の式は Table.TransformColumns で書く。')
PHRASES = ['パワークエリで文字の掃除をして', 'クエリで前後の空白と改行を取って', 'パワークエリでトリムとクリーンをかけて',
           '文字のゴミをクエリで取り除いて', 'パワークエリで余分な空白を消した表を', 'クエリで文字列をきれいにして']
SEI = ['佐藤', '鈴木', '高橋', '田中', '伊藤', '小松', '工藤', '三浦']
MEI = ['一郎', '花子', '健', '美咲', '大輔', '由美']


def dirty(rng, x, kind):
    r = rng.random()
    if kind == 'nl' and r < 0.4:
        return x + '\n'
    if kind == 'zen' and r < 0.5:
        return '　' + x + '　'
    if r < 0.25:
        return ' ' + x
    if r < 0.5:
        return x + '  '
    if r < 0.6:
        return x + '　'
    return x


def build(seed, n, kind='mix', **kw):
    rng = random.Random(seed)
    v = Vocab(seed)
    kas = v.items(5, rng)
    rows = []
    for i in range(n):
        name = f"{rng.choice(SEI)}　{rng.choice(MEI)}" if kind == 'zen' else f"{rng.choice(SEI)} {rng.choice(MEI)}"
        rows.append({'no': f"{i + 1:04d}" if kw.get('zero_no') else 1000 + i, 'name': dirty(rng, name, kind),
                     'ka': dirty(rng, rng.choice(kas), kind), 'amount': rng.randrange(10, 900) * 100,
                     'note': dirty(rng, rng.choice(['継続', '新規', '見込み', '確定']), kind)})
    t = {'sheet': v.sheet('meisai'), 'table': rng.choice(['T一覧', 'T名簿', 'T明細', 'Tデータ']) + str(seed % 7),
         'heads': [v.word('no'), v.word('name'), v.word('ka'), v.word('amount').replace('（円）', ''), v.word('tekiyo')],
         'rows': rows, 'swap': False, 'stale': False, 'blanks': False}
    t.update({k: val for k, val in kw.items() if k != 'zero_no'})
    if t['blanks']:
        rows[1]['note'] = None
        rows[3]['ka'] = None
    return t


def write_before(path, t):
    wb = Workbook()
    ws = wb.active
    ws.title = t['sheet']
    keys = ['no', 'name', 'ka', 'amount', 'note']
    order = [0, 3, 1, 2, 4] if t['swap'] else [0, 1, 2, 3, 4]
    for j, k in enumerate(order, 1):
        ws.cell(1, j, t['heads'][k]).font = Font(bold=True)
    for i, row in enumerate(t['rows'], 2):
        for j, k in enumerate(order, 1):
            if row[keys[k]] is not None:
                ws.cell(i, j, row[keys[k]])
    tb = Table(displayName=t['table'], ref=f"A1:{get_column_letter(5)}{len(t['rows']) + 1}")
    tb.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(tb)
    wb.save(path)
    return [t['heads'][k] for k in order], [keys[k] for k in order]


def m_formula(t, heads, keys):
    text_cols = [h for h, k in zip(heads, keys) if k in ('name', 'ka', 'note') or (k == 'no' and isinstance(t['rows'][0]['no'], str))]
    ops = ", ".join(f'{{"{h}", each if _ = null then null else Text.Trim(Text.Clean(Text.Replace(_, "　", " "))), type text}}'
                    for h in text_cols)
    return f'''let
    Source = Excel.CurrentWorkbook(){{[Name="{t['table']}"]}}[Content],
    Cleaned = Table.TransformColumns(Source, {{{ops}}})
in
    Cleaned'''


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t in jobs:
            before, truth = os.path.join(out, f"{stem}_前.xlsx"), os.path.join(out, f"{stem}_正解.xlsx")
            heads, keys = write_before(before, t)
            name = t['table'] + '_掃除'
            if t['stale']:
                wb = xl.Workbooks.Open(before)
                try:
                    tate.load_query(wb, name, f'let Source = Excel.CurrentWorkbook(){{[Name="{t["table"]}"]}}[Content] in Source', name)
                    wb.Worksheets(1).Activate()
                    wb.Save()
                finally:
                    wb.Close(SaveChanges=False)
            wb = xl.Workbooks.Open(before)
            try:
                tate.load_query(wb, name, m_formula(t, heads, keys), name)
                wb.Worksheets(1).Activate()
                wb.SaveAs(truth, FileFormat=51)
            finally:
                wb.Close(SaveChanges=False)
            print('  ', stem, t['sheet'], t['table'], heads)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', build(s(1), rng.randint(5, 20))),
          ('未見2_改行と空欄', build(s(2), 8, 'nl', blanks=True)),
          ('未見3_全角と並び違い', build(s(3), 7, 'zen', swap=True)),
          ('未見4_先頭0と前のクエリ', build(s(4), 6, zero_no=True, stale=True))], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('掃除_本番', build(3801, 8)),
          ('掃除_試験A', build(3802, 100)),
          ('掃除_試験B', build(3803, 8, 'nl')),
          ('掃除_試験C', build(3804, 8, 'zen')),
          ('掃除_試験D', build(3805, 8, zero_no=True)),
          ('掃除_試験E', build(3806, 7, swap=True)),
          ('掃除_試験G', build(3807, 7, blanks=True)),
          ('掃除_試験H', build(3808, 6, stale=True))], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '掃除_本番_正解.xlsx'), os.path.join(HERE, f'掃除_試験F_{x}.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
