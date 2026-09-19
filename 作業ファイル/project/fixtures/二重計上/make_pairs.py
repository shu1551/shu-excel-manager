# -*- coding: utf-8 -*-
"""同じ日・同じ金額の行（二重計上の疑い）を表の右に一覧にする（I 検査 76・2026-09-18）。

列の決まり: 選ぶ列は無い（表全体）。日付の列・金額の列は表の形で決める。
正解の決まり（依頼文に全部書く）:
  - 見出しの行＝空でないセルが 2 つ以上並ぶ最初の行。表の右端＝見出しの行で左端から右へ空のセルに当たる手前。
  - 日付の列＝本文の 8 割以上が日付の値の列（いちばん左）。金額の列＝本文の 8 割以上が数の列のうち、いちばん右の列。
  - 日付と金額の両方が同じ行が 2 行以上ある組を、表の右に 1 列空けて値で書く。見出しは「元の行」＋表の見出しをそのまま。
    組は最初に出た順、組の中は上から順。「元の行」はシートの行番号。日付は yyyy/m/d・金額は #,##0。
  - 空行・日付か金額が空の行（合計の行など）は見ない。組が無ければ見出しの行だけ書く。既に一覧が右にあれば作り直す。
組: 本番 2 組／A 800 行 6 組／B 表題と単位／C 合計の行／D 途中の空行と 3 行の組／E 列の並び違い／F 撃った後にもう一度／
  G 番号の列（数の列が 2 本＝右が金額）／H 組なし
"""
import datetime
import os
import random
import shutil
import sys

from openpyxl import load_workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import detail   # noqa: E402

REQUEST = ('同じ日付・同じ金額の行が 2 行以上ある組（二重計上の疑い）を、表の右に 1 列空けて一覧にしてください。'
           '見出しの行＝空でないセルが 2 つ以上並ぶ最初の行（文字でも数でもよい）。表の右端＝見出しの行で、表のいちばん左の列から右へ見て空のセルに当たる手前の列。'
           '日付の列＝本文の 8 割以上が日付の値の列（いちばん左のもの）。金額の列＝本文の 8 割以上が数の列のうち、いちばん右の列（見出しの語では決めない）。'
           '日付と金額の両方が同じ行を組にし、2 行以上ある組だけを、表の右端の 2 つ右の列から値で書く。'
           '見出しは「元の行」と表の見出しをそのまま並べる。組は表の中で最初に出た順、組の中は上から順。「元の行」はその行のシートの行番号。'
           '日付の表示形式は yyyy/m/d、金額は #,##0。空行と、日付か金額が空の行（下の合計の行など）は見ない。'
           '組が 1 つも無ければ見出しの行だけ書く。表の値は触らない。既に一覧が右にあれば消して作り直す。')
PHRASES = ['二重計上の疑いがある行を出して', '同じ日に同じ金額の支出がないか調べて', '二重払いのチェックをして',
           '同日同額の伝票を洗い出して', '二重に支払っていないか確認して', '日付と金額が同じ組を一覧に']


def dup(seed, n, groups=2, triple=False, **kw):
    t = detail.design(seed, n, **kw)
    rng = random.Random(seed + 81)
    dn, nn = detail.name_of(t, 'date'), detail.name_of(t, 'num')
    others = [nm for nm, r in t['cols'] if r == 'text']
    # 元から同じ日・同じ額の行があれば額をずらして消す（組の数を種で決めるため）
    seen = set()
    for row in t['rows']:
        while (row[dn], row[nn]) in seen:
            row[nn] += 100
        seen.add((row[dn], row[nn]))
    for g in range(groups):
        i = rng.randrange(len(t['rows']))
        src = t['rows'][i]
        for _k in range(2 if (triple and g == 0) else 1):
            cp = dict(src)
            for nm in others:
                if rng.random() < 0.5:
                    cp[nm] = rng.choice(['消耗品購入', '会議費用', '委託料支払', '修繕工事', '備品購入', '印刷代'])
            t['rows'].insert(min(len(t['rows']), i + 1 + rng.randrange(0, 4)), cp)
    t['n'] = len(t['rows'])
    if t['blanks']:
        rb = random.Random(seed + 13)
        t['blank_rows'] = sorted(rb.sample(range(1, t['n'] - 1), t['blanks']))
    return t


def truth_of(before, truth, t, lay):
    wb = load_workbook(before)
    ws = wb.worksheets[0]
    h, last, nc = lay['h'], lay['last'], lay['ncols']
    dc, mc = detail.col_of(t, 'date'), detail.col_of(t, 'num')
    groups, order = {}, []
    for r in range(h + 1, last + 1):
        d, m = ws.cell(r, dc).value, ws.cell(r, mc).value
        if not isinstance(d, datetime.datetime) or not isinstance(m, (int, float)):
            continue
        k = (d, m)
        if k not in groups:
            groups[k] = []
            order.append(k)
        groups[k].append(r)
    out = nc + 2
    ws.cell(h, out, '元の行')
    for j in range(1, nc + 1):
        ws.cell(h, out + j, ws.cell(h, j).value)
    w = h
    for k in order:
        if len(groups[k]) < 2:
            continue
        for r in groups[k]:
            w += 1
            ws.cell(w, out, r)
            for j in range(1, nc + 1):
                c = ws.cell(w, out + j, ws.cell(r, j).value)
                if j == dc:
                    c.number_format = 'yyyy/m/d'
                elif j == mc:
                    c.number_format = '#,##0'
    wb.save(truth)
    return sum(1 for k in order if len(groups[k]) >= 2)


def make(jobs, out):
    for stem, t in jobs:
        before = os.path.join(out, f"{stem}_前.xlsx")
        truth = os.path.join(out, f"{stem}_正解.xlsx")
        lay = detail.write_before(before, t)
        g = truth_of(before, truth, t, lay)
        print('  ', stem, t['sheet'], [n for n, _r in t['cols']], '組', g)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', dup(s(1), rng.randint(20, 60), groups=rng.randint(1, 4))),
          ('未見2_表題と合計行', dup(s(2), 30, title=True, unit=True, total=True)),
          ('未見3_並び違い', dup(s(3), 25, groups=3, triple=True, swap=True)),
          ('未見4_空行と番号列', dup(s(4), 28, blanks=2, no_col=True))], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('二重_本番', dup(2301, 30)),
          ('二重_試験A', dup(2302, 800, groups=6, n_items=10)),
          ('二重_試験B', dup(2303, 25, title=True, unit=True)),
          ('二重_試験C', dup(2304, 24, total=True)),
          ('二重_試験D', dup(2305, 26, groups=2, triple=True, blanks=2)),
          ('二重_試験E', dup(2306, 28, swap=True)),
          ('二重_試験G', dup(2307, 22, no_col=True)),
          ('二重_試験H', dup(2308, 20, groups=0))], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '二重_本番_正解.xlsx'), os.path.join(HERE, f'二重_試験F_{x}.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
