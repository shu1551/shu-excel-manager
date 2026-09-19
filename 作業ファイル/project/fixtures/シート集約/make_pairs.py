# -*- coding: utf-8 -*-
"""ほかのシートの表を 1 枚に集める（どの表でも動く形・2026-09-18 作り直し。前の版は集約）。

形で決める: アクティブなシート＝集める先（見出しの行だけがある）。集めるシート＝その見出し（1 列目とその率の列を除く）が
そろうシート。1 列目にはシート名を入れる。見出しの語・シート名・項目の名前・列の並びは表ごとに違う（vocab.py）。
正解の決まり:
  - 集める先の見出しのとおりに、ほかのシートの表を 1 行ずつ。1 列目はシート名、2 列目からは各シートの同じ見出しの列の値。
  - 率（表示形式が %）の列は、その左の 2 つの数の列の 右÷左（左が 0 なら 0）。
  - シートのタブの順・各シートの行の順。1 列目（項目）が空の行・合計や計の行は入れない。
  - 「記入例」「例」「説明」「様式」を名前に含むシート、見出しがそろわないシートは集めない。
  - 金額はカンマ・全角・空白つきの文字でも数で読む。最後に「合計」の行（1 列目に「合計」・数の列の合計・率はその合計から）。
  - 集める先に前の結果が残っていたら消してから書く（2 回撃っても同じ）。
組: 本番 4 枚／A 6 枚・行数ばらばら／B 列の並びが違うシート／C 表題の行があるシート／D 記入例と説明のシート／
  E 前の結果が残っている／F 金額が文字／G 行の無いシートと合計行のあるシート／H 集める先の見出しが 3 行目／I 撃った後にもう一度
"""
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
from vocab import Vocab, ITEMS   # noqa: E402

_ZEN = str.maketrans('0123456789,', '０１２３４５６７８９，')
ROW_HEADS = ['事業名', '細事業名', '項目', '内容', '事業']
SUM_SHEETS = ['集約', 'まとめ', '全体', '一覧', '集計']
JIGYO = ITEMS['事業']
REQUEST = ('アクティブなシート（集める先）の見出しの行のとおりに、ほかのシートの表を 1 行ずつ集めてください。'
           '1 列目にはシート名、2 列目からは各シートの表の同じ見出しの列の値を入れる。'
           '率（表示形式が %）の列は、その左の 2 つの数の列の 右÷左（左が 0 なら 0）。'
           '集めるシートは、集める先の 2 列目以降の見出し（率の列を除く）がそろうシート。'
           '「記入例」「例」「説明」「様式」を名前に含むシートと、見出しがそろわないシートは集めない。'
           'シートのタブの順・各シートの行の順に並べ、項目の列が空の行と 合計・計 の行は入れない。'
           '金額はカンマ・全角・空白つきの文字でも数で読む。'
           '最後に「合計」の行（1 列目に「合計」・数の列の合計・率はその合計から計算）。'
           '集める先に前の結果が残っていたら消してから書く')
PHRASES = ['各シートの表を1枚にまとめて', 'ほかのシートの回答を集めて', 'シートごとの表を集約して',
           '全シートを1つの表にまとめて', '課ごとのシートを集めて合計も', '各シートから集めて一覧に']


def build(seed, n_sheets, rows=(2, 8)):
    rng = random.Random(seed)
    v = Vocab(seed, item_kind='課')
    data = []
    for ka in v.items(n_sheets, rng):
        items = []
        k = rng.randint(*rows)
        names = rng.sample(JIGYO, k) if k <= len(JIGYO) else [f"{rng.choice(JIGYO)}{i + 1}" for i in range(k)]
        for j in names:
            b = rng.randrange(1, 500) * 10000
            items.append((j, b, min(b, rng.randrange(0, 520) * 10000) if rng.random() > 0.1 else 0))
        data.append((ka, items))
    t = {'data': data, 'sum_sheet': rng.choice(SUM_SHEETS), 'ka': v.word('ka'), 'row_head': rng.choice(ROW_HEADS),
         'yosan': v.word('yosan'), 'shikko': v.word('shikko'), 'rate': v.word('rate'),
         'swapped': (), 'titled': (), 'extras': False, 'stale': 0, 'text_amount': False, 'empty_ka': False,
         'total_rows': (), 'sum_title': False}
    return t


def write_pair(stem, t, out=HERE):
    rng = random.Random(stem)
    data = t['data']
    if t['empty_ka'] and len(data) > 1:
        data = data[:1] + [(data[1][0], [])] + data[2:]
    head = [t['ka'], t['row_head'], t['yosan'], t['shikko'], t['rate']]
    for truth in (False, True):
        wb = Workbook()
        sm = wb.active
        sm.title = t['sum_sheet']
        h = 3 if t['sum_title'] else 1
        if t['sum_title']:
            sm.cell(1, 1, '令和8年度 事業執行状況（全体）').font = Font(bold=True, size=14)
        for j, v in enumerate(head, 1):
            sm.cell(h, j, v).font = Font(bold=True)
        sm.cell(h, 5).number_format = '0.0%'
        r = h + 1
        if truth:
            tb = te = 0
            for ka, items in data:
                for jg, b, e in items:
                    sm.cell(r, 1, ka)
                    sm.cell(r, 2, jg)
                    sm.cell(r, 3, b).number_format = '#,##0'
                    sm.cell(r, 4, e).number_format = '#,##0'
                    sm.cell(r, 5, (e / b) if b else 0).number_format = '0.0%'
                    tb += b
                    te += e
                    r += 1
            sm.cell(r, 1, '合計')
            sm.cell(r, 3, tb).number_format = '#,##0'
            sm.cell(r, 4, te).number_format = '#,##0'
            sm.cell(r, 5, (te / tb) if tb else 0).number_format = '0.0%'
        elif t['stale']:
            for i in range(t['stale']):
                sm.cell(r + i, 1, f"前の課{i + 1}")
                sm.cell(r + i, 2, f'旧事業{i + 1}')
                sm.cell(r + i, 3, 1000 * (i + 1))
                sm.cell(r + i, 4, 500 * (i + 1))
                sm.cell(r + i, 5, 0.5).number_format = '0.0%'
            sm.cell(r + t['stale'], 1, '合計')
        if t['extras']:
            ex = wb.create_sheet('説明')
            ex.cell(1, 1, '各課は自分のシートに事業ごとの予算額と執行額を記入してください。')
            ex2 = wb.create_sheet('記入例')
            for j, v in enumerate([t['row_head'], t['yosan'], t['shikko']], 1):
                ex2.cell(1, j, v)
            ex2.cell(2, 1, '（例）庁舎清掃')
            ex2.cell(2, 2, 1200000)
            ex2.cell(2, 3, 800000)
        for k, (ka, items) in enumerate(data):
            ws = wb.create_sheet(ka)
            hh = 1
            if k in t['titled']:
                ws.cell(1, 1, f'{ka} 事業一覧（令和8年度）').font = Font(bold=True)
                ws.cell(2, 1, '単位：円')
                hh = 4
            cols = ([t['shikko'], t['row_head'], t['yosan']] if k in t['swapped']
                    else [t['row_head'], t['yosan'], t['shikko']])
            for j, v in enumerate(cols, 1):
                ws.cell(hh, j, v).font = Font(bold=True)
            rr = hh + 1
            for jg, b, e in items:
                vals = {t['row_head']: jg, t['yosan']: b, t['shikko']: e}
                if t['text_amount']:
                    for key in (t['yosan'], t['shikko']):
                        vals[key] = f"{vals[key]:,}".translate(_ZEN) if rng.random() < 0.5 else f" {vals[key]:,} "
                for j, v in enumerate(cols, 1):
                    ws.cell(rr, j, vals[v])
                rr += 1
            if k in t['total_rows'] and items:
                ws.cell(rr, cols.index(t['row_head']) + 1, '合計')
                ws.cell(rr, cols.index(t['yosan']) + 1, sum(b for _j, b, _e in items))
                ws.cell(rr, cols.index(t['shikko']) + 1, sum(e for _j, _b, e in items))
        wb.save(os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx"))
    print('  ', stem, t['sum_sheet'], head, len(data))


def kw(t, **over):
    t.update(over)
    return t


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    write_pair('未見1_標準', build(s(1), rng.randint(2, 6)), out)
    write_pair('未見2_並びと表題', kw(build(s(2), 4), swapped=(1,), titled=(0, 2)), out)
    write_pair('未見3_記入例と前の結果', kw(build(s(3), 3), extras=True, stale=rng.randint(4, 9)), out)
    write_pair('未見4_文字の金額と合計行', kw(build(s(4), 4), text_amount=True, total_rows=(0, 2)), out)
    write_pair('未見5_全部入り', kw(build(s(5), 5, rows=(1, 6)), swapped=(2,), titled=(1,), extras=True,
                                text_amount=True, total_rows=(3,), empty_ka=True, sum_title=True), out)
    return out


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    write_pair('集約_本番', build(7, 4))
    write_pair('集約_試験A', build(17, 6, rows=(1, 10)))
    write_pair('集約_試験B', kw(build(27, 4), swapped=(1, 3)))
    write_pair('集約_試験C', kw(build(37, 4), titled=(0, 2)))
    write_pair('集約_試験D', kw(build(47, 3), extras=True))
    write_pair('集約_試験E', kw(build(57, 3), stale=7))
    write_pair('集約_試験F', kw(build(67, 4), text_amount=True))
    write_pair('集約_試験G', kw(build(77, 4), empty_ka=True, total_rows=(0, 2)))
    write_pair('集約_試験H', kw(build(87, 3), sum_title=True))
    shutil.copyfile(os.path.join(HERE, '集約_本番_正解.xlsx'), os.path.join(HERE, '集約_試験I_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, '集約_本番_正解.xlsx'), os.path.join(HERE, '集約_試験I_正解.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
