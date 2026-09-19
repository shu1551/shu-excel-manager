# -*- coding: utf-8 -*-
"""テーブルの範囲を、すぐ下に足した行とすぐ右に足した列まで広げる（H テーブル 66・2026-09-18 第二期）。

列の決まり: 選ぶ列は無い。アクティブなシートのテーブル全部。
正解の決まり（依頼文に全部書く）:
  - 行: テーブルのすぐ下の行から、テーブルの列（広げた後の列）のどれかに値がある行が続くところまで（行がまるごと空の行で止まる）。
  - 列: テーブルの見出しの行のすぐ右のセルから、値のある見出しが続くところまで（空のセルで止まる）。
  - ListObject.Resize で広げる。値・名前・スタイルは変えない。足す物が無ければ何もしない。
組: 本番 3 行足した／A 800 行＋20 行／B 表題の下のテーブル／C 列だけ足した／D 行と列／E 空行の後ろの行は入れない／F 撃った後にもう一度／
  G テーブルが 2 つ（両方広げる）／H 足した物なし
"""
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import detail   # noqa: E402

REQUEST = ('アクティブなシートのテーブルを全部、すぐ下に足した行と、すぐ右に足した列まで広げてください。'
           '行: テーブルのすぐ下の行から、テーブルの列（右に広げた後の列）のどれかに値がある行が続くところまで入れる（行がまるごと空の行で止まる）。'
           '列: テーブルの見出しの行で、テーブルのすぐ右のセルから、値のある見出しが続くところまで入れる（空のセルで止まる）。'
           'ListObject.Resize で広げる（先に列を決め、その列で下の行を数える）。値・テーブルの名前・スタイルは変えない。足す物が無ければ何もしない。')
PHRASES = ['テーブルの範囲を広げて', '足した行もテーブルに入れて', 'テーブルを新しい行まで拡張して',
           'テーブルの範囲が足りないので直して', '追加した列もテーブルに含めて', 'テーブルに下の行を取り込んで']


def write_before(path, t, add_rows, add_col, gap, second, title):
    wb = Workbook()
    ws = wb.active
    ws.title = t['sheet']
    if title:
        ws['A1'] = '令和8年度 支出明細'
        ws['A1'].font = Font(bold=True, size=14)
    h0 = 3 if title else 1
    nc = len(t['cols'])
    for j, (nm, _r) in enumerate(t['cols'], 1):
        ws.cell(h0, j, nm).font = Font(bold=True)
    rows = t['rows']
    n_in = len(rows) - add_rows
    r = h0
    for i, row in enumerate(rows):
        r += 1
        if gap and i == n_in + 1:
            r += 1                                            # 足した行の途中に空行（その後ろは入れない）
        for j, (nm, role) in enumerate(t['cols'], 1):
            c = ws.cell(r, j, row[nm])
            if role == 'date':
                c.number_format = 'yyyy/m/d'
    if add_col:
        ws.cell(h0, nc + 1, '確認').font = Font(bold=True)
        for rr in range(h0 + 1, h0 + 1 + n_in):
            ws.cell(rr, nc + 1, random.Random(rr).choice(['済', '未', '']) or None)
    tb = Table(displayName=t.get('tname', 'T明細'), ref=f"A{h0}:{get_column_letter(nc)}{h0 + n_in}")
    tb.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(tb)
    if second:
        c2 = nc + 4
        ws.cell(h0, c2, '区分').font = Font(bold=True)
        ws.cell(h0, c2 + 1, '件数').font = Font(bold=True)
        for i in range(4):
            ws.cell(h0 + 1 + i, c2, f"区分{i + 1}")
            ws.cell(h0 + 1 + i, c2 + 1, i * 3 + 1)
        tb2 = Table(displayName='T区分', ref=f"{get_column_letter(c2)}{h0}:{get_column_letter(c2 + 1)}{h0 + 2}")
        tb2.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
        ws.add_table(tb2)
    wb.save(path)


def build_truth(xl, before, truth, t):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(t['sheet'])
        for i in range(1, int(ws.ListObjects.Count) + 1):
            lo = ws.ListObjects(i)
            r0, c0 = int(lo.Range.Row), int(lo.Range.Column)
            r1, c1 = r0 + int(lo.Range.Rows.Count) - 1, c0 + int(lo.Range.Columns.Count) - 1
            while ws.Cells(r0, c1 + 1).Value is not None:
                c1 += 1
            while any(ws.Cells(r1 + 1, c).Value is not None for c in range(c0, c1 + 1)):
                r1 += 1
            lo.Resize(ws.Range(ws.Cells(r0, c0), ws.Cells(r1, c1)))
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t, add_rows, add_col, gap, second, title in jobs:
            before, truth = os.path.join(out, f"{stem}_前.xlsx"), os.path.join(out, f"{stem}_正解.xlsx")
            write_before(before, t, add_rows, add_col, gap, second, title)
            build_truth(xl, before, truth, t)
            print('  ', stem, t['sheet'], '行', add_rows, '列' if add_col else '', '空行' if gap else '', '2 つ' if second else '')
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', detail.design(s(1), rng.randint(20, 50)), rng.randint(1, 5), False, False, False, False),
          ('未見2_表題と列', detail.design(s(2), 30), 2, True, False, False, True),
          ('未見3_空行の後ろ', detail.design(s(3), 25), 4, False, True, False, False),
          ('未見4_2つ', detail.design(s(4), 28), 3, True, False, True, False)], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('拡張_本番', detail.design(4501, 30), 3, False, False, False, False),
          ('拡張_試験A', detail.design(4502, 820), 20, False, False, False, False),
          ('拡張_試験B', detail.design(4503, 25), 2, False, False, False, True),
          ('拡張_試験C', detail.design(4504, 24), 0, True, False, False, False),
          ('拡張_試験D', detail.design(4505, 26), 3, True, False, False, False),
          ('拡張_試験E', detail.design(4506, 28), 5, False, True, False, False),
          ('拡張_試験G', detail.design(4507, 22), 2, False, False, True, False),
          ('拡張_試験H', detail.design(4508, 20), 0, False, False, False, False),
          ('拡張_試験I', detail.design(4509, 24, no_date=True), 3, False, False, False, False)], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '拡張_本番_正解.xlsx'), os.path.join(HERE, f'拡張_試験F_{x}.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
