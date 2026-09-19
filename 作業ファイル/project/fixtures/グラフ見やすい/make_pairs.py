# -*- coding: utf-8 -*-
"""課別支出を「いつもの見やすい棒グラフ」にする（鍛える回路の題材・グラフのデザイン・2026-09-18 未明）。

依頼は短い（見せ方の決まりは書かない）。正解のグラフが、課で決めている見せ方の型を持つ＝AI は正解との違いから型を学ぶ。
見せ方の型（正解の作り方）:
  - 支出額の大きい順に並べる（表の並びは変えない・グラフの中だけ）。合計の行は入れない。
  - 課が 8 以下なら集合縦棒、9 以上なら集合横棒（ラベルが読める）で、横棒は大きい順を上から（項目軸を反転）。
  - タイトル: 表の上に表題があればその文字、無ければ「課別支出額」。14pt・太字。
  - 凡例なし・棒に値のデータラベル・数値軸 #,##0・目盛線なし・棒の間隔 50・グラフの文字 Meiryo UI。
  - 棒の色 #1F4E79、いちばん大きい棒だけ #C00000。
  - 前に作ったグラフは消して作り直す。置き場所は表の右（位置は比べない）。
組: 本番 5 課／A 7 課／B 12 課（横棒）／C 表題つき 6 課／D 表題つき 10 課（横棒）／E 合計行つき 9 課（横棒）／F 前のグラフ／G 撃った後にもう一度／
  H 合計行つき 8 課（縦棒＝合計の行は課に数えない）
"""
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
KA = ['総務課', '財政課', '企画課', '税務課', '福祉課', '建設課', '農林課', '教育総務課', '観光課', '環境課', '会計課', '議会事務局',
      '商工課', '保健課', '子育て支援課', '市民課']
SHEET = '課別支出'
REQUEST = '課別支出の表を、いつもの見やすい棒グラフにしてください。前に作ったグラフが残っていたら消して作り直す'
BLUE, RED = (0x1F, 0x4E, 0x79), (0xC0, 0x00, 0x00)
TITLES = ['令和8年度 課別支出（9月末）', '令和8年度上半期 課別の支出', '課別支出の状況']


def rgb(c):
    return c[0] + c[1] * 256 + c[2] * 65536


def build(seed, n):
    rng = random.Random(seed)
    kas = rng.sample(KA, n)
    amounts = rng.sample(range(50, 5000), n)                  # 同じ額を作らない（大きい順が一つに決まる）
    return [(k, a * 1000) for k, a in zip(kas, amounts)]


def write_before(path, rows, title=None, total=False):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    h = 1
    if title:
        ws.cell(1, 1, title).font = Font(bold=True, size=14)
        h = 3
    ws.cell(h, 1, '課名').font = Font(bold=True)
    ws.cell(h, 2, '支出額').font = Font(bold=True)
    for i, (k, a) in enumerate(rows, 1):
        ws.cell(h + i, 1, k)
        ws.cell(h + i, 2, a).number_format = '#,##0'
    if total:
        ws.cell(h + len(rows) + 1, 1, '合計').font = Font(bold=True)
        ws.cell(h + len(rows) + 1, 2, sum(a for _k, a in rows)).number_format = '#,##0'
    wb.save(path)
    return {'h': h, 'rows': rows, 'title': title}


def build_truth(xl, before, truth, lay):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Sheets(SHEET)
        for i in range(int(ws.ChartObjects().Count), 0, -1):
            ws.ChartObjects(i).Delete()
        data = sorted(lay['rows'], key=lambda x: -x[1])
        horizontal = len(data) >= 9
        co = ws.ChartObjects().Add(ws.Cells(lay['h'], 4).Left, ws.Cells(lay['h'], 4).Top, 520, 300)
        co.Name = '表の整理_グラフ'      # マクロは「表の整理_」で始まる名前のグラフ（自分が作った物）だけを消す
        ch = co.Chart
        ch.ChartType = 57 if horizontal else 51
        for i in range(int(ch.SeriesCollection().Count), 0, -1):
            ch.SeriesCollection(i).Delete()
        s = ch.SeriesCollection().NewSeries()
        s.Name = f"='{SHEET}'!{ws.Cells(lay['h'], 2).Address}"
        s.Values = [a for _k, a in data]
        s.XValues = [k for k, _a in data]
        ch.HasTitle = True
        ch.ChartTitle.Text = lay['title'] or '課別支出額'
        ch.ChartTitle.Format.TextFrame2.TextRange.Font.Size = 14
        ch.ChartTitle.Format.TextFrame2.TextRange.Font.Bold = True
        ch.HasLegend = False
        s.HasDataLabels = True
        ax = ch.Axes(2)
        ax.TickLabels.NumberFormat = '#,##0'
        ax.HasMajorGridlines = False
        if horizontal:
            ch.Axes(1).ReversePlotOrder = True
        ch.ChartGroups(1).GapWidth = 50
        s.Format.Fill.ForeColor.RGB = rgb(BLUE)
        s.Points(1).Format.Fill.ForeColor.RGB = rgb(RED)
        ch.ChartArea.Format.TextFrame2.TextRange.Font.Name = 'Meiryo UI'
        ch.ChartTitle.Format.TextFrame2.TextRange.Font.Size = 14      # 文字の書式を全体に当てた後でもう一度
        ch.ChartTitle.Format.TextFrame2.TextRange.Font.Bold = True
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def add_old_chart(xl, path, lay):
    wb = xl.Workbooks.Open(path)
    try:
        ws = wb.Sheets(SHEET)
        co = ws.ChartObjects().Add(ws.Cells(lay['h'], 4).Left, ws.Cells(lay['h'], 4).Top, 300, 200)
        co.Name = '表の整理_前のグラフ'      # マクロは「表の整理_」で始まる名前のグラフ（自分が作った物）だけを消す
        co.Chart.ChartType = 4
        s = co.Chart.SeriesCollection().NewSeries()
        s.Values = ws.Range(ws.Cells(lay['h'] + 1, 2), ws.Cells(lay['h'] + len(lay['rows']), 2))
        wb.Save()
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, rows, kw in jobs:
            before, truth = os.path.join(out, f"{stem}_前.xlsx"), os.path.join(out, f"{stem}_正解.xlsx")
            old = kw.pop('old_chart', False)
            lay = write_before(before, rows, **kw)
            if old:
                add_old_chart(xl, before, lay)
            build_truth(xl, before, truth, lay)
            print('  ', stem)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_縦棒', build(s(1), rng.randint(3, 8)), {}),
          ('未見2_横棒と表題', build(s(2), rng.randint(9, 16)), {'title': rng.choice(TITLES)}),
          ('未見3_合計行と前のグラフ', build(s(3), 8), {'total': True, 'old_chart': True}),
          ('未見4_ちょうど9課', build(s(4), 9), {})], out)
    return out


if __name__ == '__main__':
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        sys.exit(0)
    make([('見やすい_本番', build(3, 5), {}),
          ('見やすい_試験A', build(13, 7), {}),
          ('見やすい_試験B', build(23, 12), {}),
          ('見やすい_試験C', build(33, 6), {'title': TITLES[0]}),
          ('見やすい_試験D', build(43, 10), {'title': TITLES[1]}),
          ('見やすい_試験E', build(53, 9), {'total': True}),
          ('見やすい_試験F', build(63, 6), {'old_chart': True}),
          ('見やすい_試験H', build(73, 8), {'total': True})], HERE)
    shutil.copyfile(os.path.join(HERE, '見やすい_本番_正解.xlsx'), os.path.join(HERE, '見やすい_試験G_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, '見やすい_本番_正解.xlsx'), os.path.join(HERE, '見やすい_試験G_正解.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    lines = ['課別支出を見やすいグラフにして', 'いつもの型で課別支出のグラフを', '課別支出を見栄えのいい棒グラフに',
             '課別の支出を見やすく棒グラフで', '課別支出のグラフをいつもの見せ方で', '見やすい課別支出グラフをお願い']
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(lines) + '\n')
    print('ok')
