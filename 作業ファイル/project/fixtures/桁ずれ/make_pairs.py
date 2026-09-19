# -*- coding: utf-8 -*-
"""選んでいる数の列で、桁のずれた値（千円と円の混在など）を表の右に一覧にする（I 検査 75・2026-09-18 第二期）。

列の決まり: 選んでいる 1 列＝数の列。<名前>_選択.txt に書く。表は明細（detail）。
正解の決まり（依頼文に全部書く）:
  - 見出しの行＝選んでいるセルの行。本文＝見出しの次の行から最後の行まで（空の行と「合計」「計」「総計」の行は見ない）。
  - 本文の数（0 と空は除く）の中央値を求め、中央値の 100 倍以上なら「大きすぎ」、100 分の 1 以下なら「小さすぎ」。
  - 表の右に 1 列空けて「元の行」「値」「判定」の一覧（上から順・元の行はシートの行番号・値は #,##0）。無ければ見出しの行だけ。
  - 表の値は触らない。既に一覧が右にあれば消して作り直す。
組: 本番 千円で入れた値 2 つ（小さすぎ）／A 800 行／B 表題と単位／C 合計の行（大きいが見ない）／D 大きすぎ（円を千倍）／E 列の並び違い／
  F 撃った後にもう一度／G 番号の列と 0／H ずれなし
"""
import os
import random
import shutil
import statistics
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import detail   # noqa: E402

REQUEST = ('選んでいる数の列で、桁のずれた値（千円と円が混ざったなど）を表の右に 1 列空けて一覧にしてください。'
           '見出しの行は必ず「選んでいるセルの行」＝Selection.Cells(1,1).Row をそのまま使う（自分で見出しの行を探さない）。'
           '表の列＝見出しの行で、表のいちばん左の列から右へ見て空のセルに当たる手前の列まで。本文は見出しの次の行から最後の行まで。'
           '行がまるごと空の行と、行のどこかに「合計」「計」「総計」とある行は見ない。'
           '本文の数（0 と空は除く）の中央値（WorksheetFunction.Median）を求め、中央値の 100 倍以上の値は「大きすぎ」、100 分の 1 以下の値は「小さすぎ」とする。'
           '表の右端の 2 つ右の列から、見出しの行に「元の行」「値」「判定」と書き、その下に上から順に 1 行ずつ（元の行＝シートの行番号・値の表示形式 #,##0）。'
           '無ければ見出しの行だけ書く。表の値は触らない。既に一覧が右にあれば消して作り直す。')
PHRASES = ['選んだ列の桁ずれを探して', '千円と円が混ざっていないか調べて', '桁が違う金額を洗い出して',
           '単位の間違いがないかチェックして', '金額の桁ずれを一覧に', '選んでいる列で桁が合わない値を出して']


def spoil(path, t, lay, small=2, big=0, zeros=0):
    wb = load_workbook(path)
    ws = wb.worksheets[0]
    nc = detail.col_of(t, 'num')
    rng = random.Random(t['seed'] + 41)
    body = [r for r in range(lay['h'] + 1, lay['last'] + 1) if isinstance(ws.cell(r, nc).value, (int, float))]
    picks = rng.sample(body, min(len(body), small + big + zeros))
    for k, r in enumerate(picks):
        v = ws.cell(r, nc).value
        ws.cell(r, nc).value = (max(1, v // 1000) if k < small else v * 1000 if k < small + big else 0)
    wb.save(path)


def truth_of(before, truth, t, lay):
    wb = load_workbook(before)
    ws = wb.worksheets[0]
    h, nc, ncols = lay['h'], detail.col_of(t, 'num'), lay['ncols']
    rows = []
    for r in range(h + 1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ncols + 1)]
        if all(v is None for v in vals) or any(str(v).strip() in ('合計', '計', '総計') for v in vals if v is not None):
            continue
        v = ws.cell(r, nc).value
        if isinstance(v, (int, float)) and v != 0:
            rows.append((r, v))
    med = statistics.median([v for _r, v in rows])
    out = ncols + 2
    for j, nm in enumerate(('元の行', '値', '判定')):
        ws.cell(h, out + j, nm)
    w = h
    for r, v in rows:
        why = '大きすぎ' if v >= med * 100 else '小さすぎ' if v <= med / 100 else None
        if why:
            w += 1
            ws.cell(w, out, r)
            ws.cell(w, out + 1, v).number_format = '#,##0'
            ws.cell(w, out + 2, why)
    wb.save(truth)
    return w - h


def make(jobs, out):
    for stem, t, small, big, zeros in jobs:
        before, truth = os.path.join(out, f"{stem}_前.xlsx"), os.path.join(out, f"{stem}_正解.xlsx")
        lay = detail.write_before(before, t)
        spoil(before, t, lay, small, big, zeros)
        k = truth_of(before, truth, t, lay)
        sel = f"{get_column_letter(detail.col_of(t, 'num'))}{lay['h']}"
        open(os.path.join(out, f"{stem}_選択.txt"), 'w', encoding='utf-8').write(sel + '\n')
        print('  ', stem, t['sheet'], [n for n, _r in t['cols']], '選ぶ列', sel, 'ずれ', k)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', detail.design(s(1), rng.randint(20, 60)), 2, 0, 0),
          ('未見2_表題と合計行', detail.design(s(2), 30, title=True, unit=True, total=True), 1, 1, 0),
          ('未見3_並び違い', detail.design(s(3), 25, swap=True), 0, 2, 0),
          ('未見4_空行と番号列', detail.design(s(4), 28, blanks=2, no_col=True), 2, 0, 2)], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('桁_本番', detail.design(4301, 30), 2, 0, 0),
          ('桁_試験A', detail.design(4302, 800, n_items=10), 6, 3, 0),
          ('桁_試験B', detail.design(4303, 25, title=True, unit=True), 2, 0, 0),
          ('桁_試験C', detail.design(4304, 24, total=True), 1, 0, 0),
          ('桁_試験D', detail.design(4305, 26), 0, 2, 0),
          ('桁_試験E', detail.design(4306, 28, swap=True), 2, 1, 0),
          ('桁_試験G', detail.design(4307, 22, no_col=True), 1, 0, 3),
          ('桁_試験H', detail.design(4308, 20), 0, 0, 0)], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '桁_本番_正解.xlsx'), os.path.join(HERE, f'桁_試験F_{x}.xlsx'))
    shutil.copyfile(os.path.join(HERE, '桁_本番_選択.txt'), os.path.join(HERE, '桁_試験F_選択.txt'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
