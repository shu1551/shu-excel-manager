# -*- coding: utf-8 -*-
"""選んでいる日付の列で、おかしい日付（空・日付でない・未来・年度外）の行を表の右に一覧にする（I 検査 72・2026-09-18）。

列の決まり: 選んでいる 1 列＝日付の列。選んでいる列は <名前>_選択.txt に書く。
正解の決まり（依頼文に全部書く）:
  - 理由は順に 1 つだけ: 空／日付でない（日付の値でない＝文字など）／未来の日付（今日より後）／年度外（表の年度の 4/1〜翌 3/31 の外）。
  - 表の年度＝本文の日付の値でいちばん多い年度（1〜3 月は前の年）。
  - 表の右に 1 列空けて「元の行」「（日付の列の見出し）」「理由」。上から順。無ければ見出しの行だけ。
  - 空行・合計の行は見ない。既に一覧が右にあれば作り直す。
明細は 2024 年度（未来の日付は 2031 年＝今日の日付で答えが揺れない）。
組: 本番／A 800 行／B 表題と単位／C 合計の行／D 途中の空行／E 列の並び違いで 2025 年度／F 撃った後にもう一度／G 番号の列／H おかしい日付なし
"""
import datetime
import os
import random
import shutil
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import detail   # noqa: E402

REQUEST = ('選んでいる日付の列で、おかしい日付の行を、表の右に 1 列空けて一覧にしてください。'
           '見出しの行は必ず「選んでいるセルの行」＝Selection.Cells(1,1).Row をそのまま使う（自分で見出しの行を探さない）。'
           '表の右端＝見出しの行で、表のいちばん左の列から右へ見て空のセルに当たる手前の列。その 2 つ右の列から書く。'
           '本文は見出しの次の行から最後の本文の行まで。行がまるごと空の行と、下の合計の行（行のどこかのセルに「合計」「計」「総計」とある行）は見ない。'
           '理由は次の順に 1 つだけ: 日付のセルが空＝「空」／日付の値でない（文字など）＝「日付でない」／今日（Date）より後＝「未来の日付」／'
           '表の年度（4 月 1 日〜翌年 3 月 31 日）の外＝「年度外」。表の年度＝本文の日付の値でいちばん多い年度（4〜12 月はその年、1〜3 月は前の年）。'
           '一覧の見出しは「元の行」・日付の列の見出し・「理由」。元の行はシートの行番号、日付の列の値はそのまま写す（日付は yyyy/m/d）。上から順。'
           'おかしい行が無ければ見出しの行だけ書く。表の値は触らない。既に一覧が右にあれば消して作り直す。')
PHRASES = ['日付がおかしい行を洗い出して', '年度外の日付がないか確認して', '選んだ列の日付をチェックして',
           '日付の入力ミスを探して', '未来の日付や空の日付を一覧に', '日付の妥当性を検査して']

TODAY = datetime.datetime(2026, 9, 18)


def fy(d):
    return d.year if d.month >= 4 else d.year - 1


def bad(seed, n, k=1, year=2024, **kw):
    t = detail.design(seed, n, year=year, **kw)
    if not k:
        return t
    dn = detail.name_of(t, 'date')
    rng = random.Random(seed + 101)
    kinds = [datetime.datetime(year, 3, rng.randint(1, 28)), datetime.datetime(year + 1, 4, rng.randint(1, 28)),
             datetime.datetime(2031, 5, rng.randint(1, 28)), None, rng.choice(['未定', f'{year}/13/40', '不明'])]
    idx = rng.sample(range(len(t['rows'])), len(kinds) * k)
    for j, i in enumerate(idx):
        t['rows'][i][dn] = kinds[j % len(kinds)]
    return t


def reason(v, year):
    if v is None or (isinstance(v, str) and not v.strip()):
        return '空'
    if not isinstance(v, datetime.datetime):
        return '日付でない'
    if v > TODAY:
        return '未来の日付'
    if fy(v) != year:
        return '年度外'
    return None


def truth_of(before, truth, t, lay):
    wb = load_workbook(before)
    ws = wb.worksheets[0]
    h, last, nc = lay['h'], lay['last'], lay['ncols']
    dc = detail.col_of(t, 'date')
    body = [r for r in range(h + 1, last + 1) if any(ws.cell(r, c).value is not None for c in range(1, nc + 1))]
    years = {}
    for r in body:
        v = ws.cell(r, dc).value
        if isinstance(v, datetime.datetime):
            years[fy(v)] = years.get(fy(v), 0) + 1
    year = max(years, key=years.get)
    out = nc + 2
    ws.cell(h, out, '元の行')
    ws.cell(h, out + 1, ws.cell(h, dc).value)
    ws.cell(h, out + 2, '理由')
    w = h
    for r in body:
        v = ws.cell(r, dc).value
        why = reason(v, year)
        if why:
            w += 1
            ws.cell(w, out, r)
            c = ws.cell(w, out + 1, v)
            if isinstance(v, datetime.datetime):
                c.number_format = 'yyyy/m/d'
            ws.cell(w, out + 2, why)
    wb.save(truth)
    return w - h


def make(jobs, out):
    for stem, t in jobs:
        before = os.path.join(out, f"{stem}_前.xlsx")
        truth = os.path.join(out, f"{stem}_正解.xlsx")
        lay = detail.write_before(before, t)
        k = truth_of(before, truth, t, lay)
        sel = f"{get_column_letter(detail.col_of(t, 'date'))}{lay['h']}"
        open(os.path.join(out, f"{stem}_選択.txt"), 'w', encoding='utf-8').write(sel + '\n')
        print('  ', stem, t['sheet'], [n for n, _r in t['cols']], '選ぶ列', sel, 'おかしい行', k)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', bad(s(1), rng.randint(20, 60))),
          ('未見2_表題と合計行', bad(s(2), 30, title=True, unit=True, total=True)),
          ('未見3_並び違い2025', bad(s(3), 25, year=2025, swap=True)),
          ('未見4_空行と番号列', bad(s(4), 28, blanks=2, no_col=True))], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('日付検査_本番', bad(2401, 30)),
          ('日付検査_試験A', bad(2402, 800, k=3, n_items=10)),
          ('日付検査_試験B', bad(2403, 25, title=True, unit=True)),
          ('日付検査_試験C', bad(2404, 24, total=True)),
          ('日付検査_試験D', bad(2405, 26, blanks=2)),
          ('日付検査_試験E', bad(2406, 28, year=2025, swap=True)),
          ('日付検査_試験G', bad(2407, 22, no_col=True)),
          ('日付検査_試験H', bad(2408, 20, k=0))], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '日付検査_本番_正解.xlsx'), os.path.join(HERE, f'日付検査_試験F_{x}.xlsx'))
    shutil.copyfile(os.path.join(HERE, '日付検査_本番_選択.txt'), os.path.join(HERE, '日付検査_試験F_選択.txt'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
