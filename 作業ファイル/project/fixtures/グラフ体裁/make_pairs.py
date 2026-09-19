# -*- coding: utf-8 -*-
"""シートにある既存のグラフを、いつもの見せ方にそろえる（E グラフ 46・2026-09-18 第二期）。

列の決まり: 選ぶ列は無い。アクティブなシートのグラフ全部が対象。データ・種類・系列の色・置き場所は変えない。
正解の決まり（依頼文に全部書く）:
  - グラフ全体の文字 Meiryo UI。タイトルがあれば 14pt・太字（無いグラフにタイトルは足さない）。
  - 凡例: 円・ドーナツは右／系列が 1 本なら凡例なし／2 本以上なら下。
  - 円・ドーナツ以外: 数値軸の表示形式は #,##0（今の表示形式に % があればそのまま）・数値軸の目盛線なし。
組: 本番 棒 1 本／A 折れ線 2 系列・タイトルなし／B 円と棒／C 棒・折れ線・円の 3 つ／D 率の棒（% の軸）／E 積み上げ 2 系列／
  F 撃った後にもう一度／G 横棒とマーカー付き折れ線／H グラフなし
"""
import os
import random
import shutil
import sys

from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import shukei   # noqa: E402

REQUEST = ('このシートにある既存のグラフを全部、いつもの見せ方にそろえてください。データ・グラフの種類・系列の色・置き場所・大きさは変えない。'
           'グラフ全体の文字は Meiryo UI（ChartArea に当てる）。タイトルがあるグラフは、全体の文字を当てた後でタイトルを 14pt・太字にする（タイトルの無いグラフに足さない）。'
           '凡例: 円・ドーナツのグラフは右／系列が 1 本のグラフは凡例なし／系列が 2 本以上のグラフは下。'
           '円・ドーナツ以外のグラフ: 数値軸（Axes(xlValue)）の表示形式は #,##0（今の表示形式に % が入っていればそのまま）、数値軸の目盛線なし。'
           'グラフが 1 つも無ければ何もしない。')
PHRASES = ['グラフの見せ方をいつもの形にそろえて', 'このシートのグラフを全部いつもの見た目に', '既存のグラフの体裁を整えて',
           'グラフのフォントと凡例をそろえて', '作ってあるグラフの見た目を直して', 'グラフを全部同じ体裁に']
KINDS = {'bar1': (51, 1), 'line2': (4, 2), 'pie': (5, 1), 'rate': (51, 'rate'), 'stack2': (52, 2), 'hbar1': (57, 1), 'mline2': (65, 2)}


def add_charts(xl, path, t, lay, kinds, titled):
    wb = xl.Workbooks.Open(path)
    try:
        ws = wb.Worksheets(t['sheet'])
        h, last, c0 = lay['h'], lay['last'], lay['c0']
        ic = c0 + shukei.col_of(t, 'item') - 1
        il = get_column_letter(ic)
        for k, kind in enumerate(kinds):
            ctype, ns = KINDS[kind]
            if ns == 'rate':
                cols = [c0 + shukei.col_of(t, 'rate') - 1]
            else:
                cols = [c0 + shukei.col_of(t, 'num', j) - 1 for j in range(ns)]
            addr = f"{il}{h}:{il}{last}," + ",".join(f"{get_column_letter(c)}{h}:{get_column_letter(c)}{last}" for c in cols)
            co = ws.ChartObjects().Add(ws.Cells(h, c0 + lay['ncols'] + 1).Left + (k % 2) * 330,
                                       ws.Cells(h, 1).Top + (k // 2) * 230, 320, 220)
            ch = co.Chart
            ch.ChartType = ctype
            ch.SetSourceData(ws.Range(addr), 2)
            if titled[k]:
                ch.HasTitle = True
                ch.ChartTitle.Text = f"{t['cols'][ic - c0][0]}別の{ws.Cells(h, cols[0]).Value}"
            else:
                ch.HasTitle = False
        wb.Save()
    finally:
        wb.Close(SaveChanges=False)


def build_truth(xl, before, truth):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(1)
        for i in range(1, int(ws.ChartObjects().Count) + 1):
            ch = ws.ChartObjects(i).Chart
            ch.ChartArea.Format.TextFrame2.TextRange.Font.Name = 'Meiryo UI'
            if ch.HasTitle:
                ch.ChartTitle.Format.TextFrame2.TextRange.Font.Size = 14
                ch.ChartTitle.Format.TextFrame2.TextRange.Font.Bold = True
            pie = int(ch.ChartType) in (5, -4120, 69, 70)
            if pie:
                ch.HasLegend = True
                ch.Legend.Position = -4152
            elif int(ch.SeriesCollection().Count) == 1:
                ch.HasLegend = False
            else:
                ch.HasLegend = True
                ch.Legend.Position = -4107
            if not pie:
                ax = ch.Axes(2)
                if '%' not in str(ax.TickLabels.NumberFormat):
                    ax.TickLabels.NumberFormat = '#,##0'
                ax.HasMajorGridlines = False
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t, kinds, titled in jobs:
            before, truth = os.path.join(out, f"{stem}_前.xlsx"), os.path.join(out, f"{stem}_正解.xlsx")
            lay = shukei.write_before(before, t)
            if kinds:
                add_charts(xl, before, t, lay, kinds, titled)
            build_truth(xl, before, truth)
            print('  ', stem, t['sheet'], [n for n, _r in t['cols']], kinds)
    finally:
        xl = None
        vf._quit(pid)


def d(seed, n, **kw):
    kw.setdefault('nums', 2)
    return shukei.design(seed, n, **kw)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_棒と折れ線', d(s(1), rng.randint(4, 9)), ['bar1', 'line2'], [True, False]),
          ('未見2_円と率', d(s(2), 6, rate=True, title=True), ['pie', 'rate'], [True, True]),
          ('未見3_積み上げと横棒', d(s(3), 7, total=True), ['stack2', 'hbar1'], [False, True]),
          ('未見4_マーカー折れ線', d(s(4), 5, no_col=True), ['mline2'], [True])], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('体裁_本番', d(3101, 6), ['bar1'], [True]),
          ('体裁_試験A', d(3102, 8), ['line2'], [False]),
          ('体裁_試験B', d(3103, 5), ['pie', 'bar1'], [True, True]),
          ('体裁_試験C', d(3104, 6, title=True, unit=True), ['bar1', 'line2', 'pie'], [True, True, False]),
          ('体裁_試験D', d(3105, 7, rate=True), ['rate'], [True]),
          ('体裁_試験E', d(3106, 6, total=True), ['stack2'], [False]),
          ('体裁_試験G', d(3107, 9, no_col=True), ['hbar1', 'mline2'], [True, True]),
          ('体裁_試験H', d(3108, 5), [], [])], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '体裁_本番_正解.xlsx'), os.path.join(HERE, f'体裁_試験F_{x}.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
