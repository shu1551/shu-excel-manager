# -*- coding: utf-8 -*-
"""支出明細をパワークエリで課×月のクロス集計にする（鍛える回路の題材・パワークエリ 4 本目・グループ化と列のピボット・2026-09-18）。

仕事: テーブル T支出明細 から、行＝課名（昇順）・列＝月（年度の順 4月→3月・明細にある月だけ）・値＝支出額の合計 の表を、
      新しいシート「PQ課別月別」の A1 にテーブル「Q課別月別」として読み込む。クエリ名も Q課別月別。
正解の決まり（依頼文に全部書く）:
  - 月の見出しは「4月」の形。明細の無い課×月は空欄。同じ名前のクエリやシートが既にあれば消して作り直す。
組: 本番 60 行（4〜9 月）／A 3,000 行（4〜3 月＝年度をまたぐ）／B 4〜6 月だけ／C 列が多い・並びが違う／D 前のクエリとシート／
  E 撃った後にもう一度／G 1〜3 月と 4 月（並びが 4,1,2,3）
"""
import datetime
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
KA = ['総務課', '財政課', '企画課', '税務課', '福祉課', '建設課', '農林課', '教育総務課', '観光課', '環境課']
KAMOKU = ['需用費', '役務費', '委託料', '旅費', '備品購入費']
QNAME = 'Q課別月別'
OUT_SHEET = 'PQ課別月別'
SHEET = '支出明細'
HEAD = ['支出日', '課名', '支出額']
WIDE = ['伝票番号', '課名', '科目', '支出額', '支出日', '摘要']
M = '''let
    Source = Excel.CurrentWorkbook(){[Name="T支出明細"]}[Content],
    Typed = Table.TransformColumnTypes(Source, {{"支出日", type date}, {"課名", type text}, {"支出額", Int64.Type}}),
    AddM = Table.AddColumn(Typed, "月", each Date.Month([支出日]), Int64.Type),
    Grouped = Table.Group(AddM, {"課名", "月"}, {{"支出額", each List.Sum([支出額]), Int64.Type}}),
    Order = List.Select({4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3}, each List.Contains(Grouped[月], _)),
    Named = Table.TransformColumns(Grouped, {{"月", each Text.From(_) & "月", type text}}),
    Pivoted = Table.Pivot(Named, List.Transform(Order, each Text.From(_) & "月"), "月", "支出額", List.Sum),
    Sorted = Table.Sort(Pivoted, {{"課名", Order.Ascending}})
in
    Sorted'''
REQUEST = ('パワークエリで、テーブル T支出明細 から、行は課名（昇順）、列は月（年度の順に 4月から3月まで・明細にある月だけ・見出しは「4月」の形）、'
           '値は支出額の合計のクロス集計表を作り、新しいシート「PQ課別月別」の A1 にテーブル「Q課別月別」として読み込んでください。'
           'クエリ名も Q課別月別。明細の無い課と月は空欄。同じ名前のクエリやシートが既にあれば消して作り直す')


def build(seed, n, months=(4, 9)):
    rng = random.Random(seed)
    kas = rng.sample(KA, rng.randint(3, 7))
    lo, hi = months
    rows = []
    for i in range(n):
        m = rng.randint(lo, hi)
        mm = (m - 1) % 12 + 1
        d = datetime.datetime(2026 + (m > 12), mm, rng.randint(1, 28))
        rows.append({'伝票番号': 26000 + i, '支出日': d, '課名': rng.choice(kas), '科目': rng.choice(KAMOKU),
                     '支出額': rng.randrange(1, 900) * 100, '摘要': '―'})
    return rows


def write_before(path, rows, cols=HEAD):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    for j, c in enumerate(cols, 1):
        ws.cell(1, j, c).font = Font(bold=True)
    for i, r in enumerate(rows, 2):
        for j, c in enumerate(cols, 1):
            cell = ws.cell(i, j, r[c])
            if c == '支出日':
                cell.number_format = 'yyyy/m/d'
    t = Table(displayName='T支出明細', ref=f"A1:{get_column_letter(len(cols))}{len(rows) + 1}")
    t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(t)
    wb.save(path)


def load_query(wb, formula=M):
    xl = wb.Application
    for i in range(int(wb.Queries.Count), 0, -1):
        if str(wb.Queries.Item(i).Name) == QNAME:
            wb.Queries.Item(i).Delete()
    for sh in list(wb.Worksheets):
        if str(sh.Name) == OUT_SHEET:
            xl.DisplayAlerts = False
            sh.Delete()
    wb.Queries.Add(QNAME, formula)
    ws = wb.Worksheets.Add(After=wb.Worksheets(wb.Worksheets.Count))
    ws.Name = OUT_SHEET
    conn = f"OLEDB;Provider=Microsoft.Mashup.OleDb.1;Data Source=$Workbook$;Location={QNAME};Extended Properties=\"\""
    lo = ws.ListObjects.Add(SourceType=0, Source=conn, Destination=ws.Range("$A$1"))
    lo.QueryTable.CommandType = 2
    lo.QueryTable.CommandText = f"SELECT * FROM [{QNAME}]"
    lo.QueryTable.Refresh(False)
    lo.Name = QNAME
    return lo


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, rows, kw in jobs:
            before, truth = os.path.join(out, f"{stem}_前.xlsx"), os.path.join(out, f"{stem}_正解.xlsx")
            stale = kw.pop('stale', False)
            write_before(before, rows, **kw)
            if stale:
                wb = xl.Workbooks.Open(before)
                try:
                    load_query(wb, M.replace('Order.Ascending', 'Order.Descending'))
                    wb.Worksheets(1).Activate()
                    wb.Save()
                finally:
                    wb.Close(SaveChanges=False)
            wb = xl.Workbooks.Open(before)
            try:
                load_query(wb)
                wb.Worksheets(1).Activate()
                wb.SaveAs(truth, FileFormat=51)
            finally:
                wb.Close(SaveChanges=False)
            print('  ', stem)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', build(s(1), rng.randint(20, 120), months=(4, rng.randint(5, 12))), {}),
          ('未見2_年度またぎと列が多い', build(s(2), 200, months=(7, 15)), {'cols': WIDE}),
          ('未見3_前のクエリ', build(s(3), 40, months=(9, 13)), {'stale': True})], out)
    return out


if __name__ == '__main__':
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        sys.exit(0)
    make([('クロス_本番', build(3, 60), {}),
          ('クロス_試験A', build(13, 3000, months=(4, 15)), {}),
          ('クロス_試験B', build(23, 50, months=(4, 6)), {}),
          ('クロス_試験C', build(33, 70), {'cols': WIDE}),
          ('クロス_試験D', build(43, 45), {'stale': True}),
          ('クロス_試験G', build(63, 40, months=(13, 16)), {})], HERE)
    shutil.copyfile(os.path.join(HERE, 'クロス_本番_正解.xlsx'), os.path.join(HERE, 'クロス_試験E_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, 'クロス_本番_正解.xlsx'), os.path.join(HERE, 'クロス_試験E_正解.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    lines = ['パワークエリで課別月別のクロス集計を', '支出明細をクエリで課と月のクロス表に', 'T支出明細から月を列にした課別の表をパワークエリで',
             'パワークエリの列のピボットで課別月別の集計', '課×月の支出額をクエリで集計して横に月を並べて', 'PQ課別月別のシートにクロス集計を読み込んで']
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(lines) + '\n')
    print('ok')
