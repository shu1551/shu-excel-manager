# -*- coding: utf-8 -*-
"""旧システムと新システムの出力の突合（鍛える回路の題材・2026-09-17。財務会計の移行で必ず出る仕事）。

仕事: 「旧システム」の伝票ごとに「新システム」を引き、新の支出額と突合の結果を右に足し、結果ごとの件数を出す。
正解の決まり:
  - 旧の表の右の空いた列に「新の支出額」「突合」を足す（既にその見出しがあればそこを使う＝2 回撃っても同じ）。
    突合＝ 一致／金額違い／新に無し。新に無しの行の「新の支出額」は空。
  - 伝票番号は先頭の 0・数値と文字の違い・前後の空白を無視して比べる。金額はカンマ・全角・前後の空白を無視して数で比べる。
  - 伝票番号が空の行（途中の空行・最後の合計行）には書かず、数えない。
  - 「突合」の列の右に 1 列空けて件数の表。見出しは旧の表の見出しと同じ行に「突合」「件数」。
    行は 一致・金額違い・新に無し・旧に無し の順で 0 件も出す（旧に無し＝新にだけある伝票の数）。
  - 新システムのシートは見出しの語で列を探す（表題の行・列の並び・ほかの列があっても同じ）。
組（依頼文は同じ）:
  本番／A 新の伝票番号が数値（旧は 0 詰めの文字）／B 新に表題の行と列の並び違い・起票日の列／C 旧に摘要の列／
  D 旧が B3 から・途中に空行／E 旧の最後に合計行／F 新の金額が文字（カンマ・全角）／G シート名が「旧システム出力」「新システム出力」／
  H 撃った後の表にもう一度
  py make_pairs.py                → このフォルダに *_前.xlsx と *_正解.xlsx
  py make_pairs.py --unseen N     → 未見_種N（鍛えるときに AI に見せない表）
"""
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
KAMOKU = ['旅費', '需用費', '役務費', '委託料', '使用料及び賃借料', '工事請負費', '備品購入費', '負担金', '補助金', '報償費']
_ZEN = str.maketrans('0123456789,', '０１２３４５６７８９，')


def build(seed, n, diff, miss, extra):
    rng = random.Random(seed)
    start = rng.randint(100, 5000)
    keys = sorted(rng.sample(range(start, start + n * 3), n + extra))
    old = [(k, rng.choice(KAMOKU), rng.randrange(1, 900) * 1000) for k in keys[:n]]
    new = {k: a for k, _m, a in old}
    for k, _m, a in rng.sample(old, miss):
        del new[k]
    for k in rng.sample([k for k in new], diff):
        new[k] = new[k] + rng.choice([-1, 1]) * rng.randrange(1, 50) * 100
    for k in keys[n:]:
        new[k] = rng.randrange(1, 900) * 1000
    return old, new


def write_pair(stem, old, new, out=HERE, key_num=False, new_title=False, new_swapped=False, memo=False,
               r0=1, c0=1, blanks=(), total=False, amount_text=False, names=('旧システム', '新システム')):
    rng = random.Random(stem)
    kubun = []
    for truth in (False, True):
        wb = Workbook()
        ws = wb.active
        ws.title = names[0]
        h = r0
        base = ['伝票番号', '科目名', '支出額'] + (['摘要'] if memo else [])
        add = c0 + len(base)
        heads = base + (['新の支出額', '突合'] if truth else [])
        for j, v in enumerate(heads):
            ws.cell(h, c0 + j, v).font = Font(bold=True)
        counts = {'一致': 0, '金額違い': 0, '新に無し': 0, '旧に無し': len([k for k in new if k not in {o[0] for o in old}])}
        r = h + 1
        for i, (k, m, a) in enumerate(old):
            if i in blanks:
                r += 1
            ws.cell(r, c0, f"{k:06d}").number_format = '@'
            ws.cell(r, c0 + 1, m)
            ws.cell(r, c0 + 2, a).number_format = '#,##0'
            if memo:
                ws.cell(r, c0 + 3, f"{m}の支払")
            if truth:
                if k in new:
                    ws.cell(r, add, new[k]).number_format = '#,##0'
                    res = '一致' if new[k] == a else '金額違い'
                else:
                    res = '新に無し'
                ws.cell(r, add + 1, res)
                counts[res] += 1
            r += 1
        if total:
            ws.cell(r, c0 + 1, '合計').font = Font(bold=True)
            ws.cell(r, c0 + 2, sum(a for _k, _m, a in old)).number_format = '#,##0'
        if truth:
            sc = add + 3
            ws.cell(h, sc, '突合').font = Font(bold=True)
            ws.cell(h, sc + 1, '件数').font = Font(bold=True)
            for i, lab in enumerate(['一致', '金額違い', '新に無し', '旧に無し'], h + 1):
                ws.cell(i, sc, lab)
                ws.cell(i, sc + 1, counts[lab])
        ns = wb.create_sheet(names[1])
        nh = 1
        if new_title:
            ns.cell(1, 1, '新財務会計システム 支出負担行為一覧（出力）').font = Font(bold=True)
            ns.cell(2, 1, '出力日 2027/01/15')
            nh = 4
        cols = ['支出額', '起票日', '伝票番号'] if new_swapped else ['伝票番号', '支出額']
        for j, v in enumerate(cols, 1):
            ns.cell(nh, j, v).font = Font(bold=True)
        items = list(new.items())
        rng.shuffle(items)
        for i, (k, a) in enumerate(items, nh + 1):
            vals = {'伝票番号': k if key_num else f"{k:06d}", '支出額': a, '起票日': f"2027/1/{(k % 28) + 1}"}
            if amount_text:
                vals['支出額'] = f"{a:,}".translate(_ZEN) if rng.random() < 0.5 else f" {a:,} "
            for j, v in enumerate(cols, 1):
                c = ns.cell(i, j, vals[v])
                if v == '伝票番号' and not key_num:
                    c.number_format = '@'
        wb.save(os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx"))


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    b = lambda k: build(s(k), rng.randint(15, 50), rng.randint(0, 5), rng.randint(0, 4), rng.randint(0, 4))  # noqa: E731
    write_pair('未見1_標準', *b(1), out=out)
    write_pair('未見2_番号が数値', *b(2), key_num=True, out=out)
    write_pair('未見3_新の表題と並び', *b(3), new_title=True, new_swapped=True, out=out)
    write_pair('未見4_摘要', *b(4), memo=True, out=out)
    write_pair('未見5_位置と空行', *b(5), r0=rng.randint(2, 5), c0=rng.randint(2, 4), blanks=(3, 9), out=out)
    write_pair('未見6_合計行', *b(6), total=True, out=out)
    write_pair('未見7_金額が文字', *b(7), amount_text=True, out=out)
    write_pair('未見8_全部入り', *build(s(8), 40, 4, 3, 3), key_num=True, new_title=True, new_swapped=True, memo=True,
               r0=2, c0=3, blanks=(5, 20), total=True, amount_text=True, names=('旧システム出力', '新システム出力'), out=out)
    write_pair('未見9_大量', *build(s(9), 3000, 60, 40, 30), key_num=True, out=out)
    return out


if __name__ == '__main__':
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        sys.exit(0)
    write_pair('突合_本番', *build(11, 30, 3, 2, 2))
    write_pair('突合_試験A', *build(23, 40, 4, 3, 1), key_num=True)
    write_pair('突合_試験B', *build(37, 25, 2, 1, 3), new_title=True, new_swapped=True)
    write_pair('突合_試験C', *build(43, 22, 2, 2, 0), memo=True)
    write_pair('突合_試験D', *build(59, 28, 3, 1, 2), r0=3, c0=2, blanks=(6, 17))
    write_pair('突合_試験E', *build(61, 20, 1, 0, 1), total=True)
    write_pair('突合_試験F', *build(71, 26, 3, 2, 2), amount_text=True)
    write_pair('突合_試験G', *build(89, 18, 0, 2, 2), names=('旧システム出力', '新システム出力'))
    shutil.copyfile(os.path.join(HERE, '突合_本番_正解.xlsx'), os.path.join(HERE, '突合_試験H_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, '突合_本番_正解.xlsx'), os.path.join(HERE, '突合_試験H_正解.xlsx'))
    print('ok')
