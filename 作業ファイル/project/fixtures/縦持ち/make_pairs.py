# -*- coding: utf-8 -*-
"""横持ちの表をパワークエリで縦持ちにする（どの表でも動く形・2026-09-18 作り直し。前の版は PQ縦持ち）。

列の決まり: キーの列＝テーブルの左端の文字の列（見出しはそのまま）、値の列＝その右の数の列。
作る物の名前は元のテーブル名から（「（元のテーブル名）_縦持ち」）。見出しの語・シート名・テーブル名・項目の名前は表ごとに違う。
正解の決まり:
  - 見出しに 番号・No・コード を含む列と、合計・計・総計の列は入れない。キーが合計・計・総計の行も入れない。
  - 空欄の組は行を作らない（0 は行を作る）。
  - 縦持ちの列は キーの列の見出し・「項目」・「値」。新しいシートの A1 に同じ名前のテーブルとして読み込む。
  - 同じ名前のクエリ・シートが既にあれば消して作り直す。
組: 本番 5 件×6 か月／A 10 件×12 か月／B 合計の列／C コードの列なし／D 空欄と 0／E 前のクエリとシート／
  F 撃った後にもう一度／G 合計の行／H 四半期の列・備考の列
"""
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
from vocab import Vocab   # noqa: E402

MONTHS = ['4月', '5月', '6月', '7月', '8月', '9月', '10月', '11月', '12月', '1月', '2月', '3月']
QUARTERS = ['第1四半期', '第2四半期', '第3四半期', '第4四半期']
REQUEST = ('パワークエリで、アクティブなシートのテーブルを縦持ちに変換してください。キーの列は左端の文字の列（見出しはそのまま）、'
           '値の列はその右の数の列。縦持ちの列は キーの列の見出し・「項目」・「値」の 3 列。'
           '見出しに 番号・No・コード を含む列と、合計・計・総計の列は入れない。キーが合計・計・総計の行も入れない。'
           '空欄は行を作らない。新しいシート（名前は 元のテーブル名 と _縦持ち をつないだ名前）の A1 に同じ名前のテーブルとして'
           '読み込む。クエリ名も同じ。同じ名前のクエリやシートが既にあれば消して作り直す')
PHRASES = ['パワークエリでこの表を縦持ちに', '横持ちの表をパワークエリで縦に並べ替えて', 'テーブルをピボット解除して縦持ちにして',
           'パワークエリで列のピボット解除をお願い', '横の列を縦持ちに変換するクエリを作って', 'クエリで縦持ちの表にして']


def build(seed, n, m, period='month', blanks=False, **kw):
    rng = random.Random(seed)
    v = Vocab(seed)
    heads = (MONTHS if period == 'month' else QUARTERS)[:m]
    items = v.items(n, rng)
    rows = []
    for i, k in enumerate(items):
        vals = {}
        for hd in heads:
            val = rng.randrange(0, 900) * 1000
            if blanks and rng.random() < 0.2:
                val = None if rng.random() < 0.6 else 0
            vals[hd] = val
        rows.append({'code': f"{110 + i * 10}", 'key': k, **vals})
    t = {'sheet': v.sheet('getsuji'), 'table': v.table('getsuji'), 'ka': v.word('ka'), 'code': v.word('code'),
         'biko': v.word('tekiyo'), 'heads': heads, 'rows': rows, 'code_col': True, 'total_col': False,
         'total_row': False, 'note_col': False, 'stale': False, 'period': period}
    t.update(kw)
    return t


def write_before(path, t):
    wb = Workbook()
    ws = wb.active
    ws.title = t['sheet']
    cols = ([t['code']] if t['code_col'] else []) + [t['ka']] + t['heads'] + (['合計'] if t['total_col'] else []) \
        + ([t['biko']] if t['note_col'] else [])
    for j, c in enumerate(cols, 1):
        ws.cell(1, j, c).font = Font(bold=True)
    body = list(t['rows'])
    if t['total_row']:
        body.append({'code': None, 'key': '合計', **{m: sum((r[m] or 0) for r in t['rows']) for m in t['heads']}})
    for i, r in enumerate(body, 2):
        for j, c in enumerate(cols, 1):
            if c == '合計':
                v = sum((r[m] or 0) for m in t['heads'])
            elif c == t['code']:
                v = r.get('code')
            elif c == t['ka']:
                v = r.get('key')
            elif c == t['biko']:
                v = random.Random(i).choice(['', '確定', '見込'])
            else:
                v = r.get(c)
            if v not in (None, ''):
                ws.cell(i, j, v)
    tb = Table(displayName=t['table'], ref=f"A1:{get_column_letter(len(cols))}{len(body) + 1}")
    tb.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(tb)
    wb.save(path)


def m_formula(t, value_name='値'):
    drop = [c for c in ([t['code']] if t['code_col'] else []) + (['合計'] if t['total_col'] else [])
            + ([t['biko']] if t['note_col'] else [])]
    drop_txt = ", ".join(f'"{c}"' for c in drop)
    return f'''let
    Source = Excel.CurrentWorkbook(){{[Name="{t['table']}"]}}[Content],
    Removed = Table.RemoveColumns(Source, {{{drop_txt}}}, MissingField.Ignore),
    NoTotal = Table.SelectRows(Removed, each [{t['ka']}] <> "合計" and [{t['ka']}] <> "計" and [{t['ka']}] <> "総計"),
    Unpivoted = Table.UnpivotOtherColumns(NoTotal, {{"{t['ka']}"}}, "項目", "{value_name}"),
    Typed = Table.TransformColumnTypes(Unpivoted, {{{{"{t['ka']}", type text}}, {{"項目", type text}}, {{"{value_name}", Int64.Type}}}})
in
    Typed'''


def load_query(wb, name, formula, sheet):
    xl = wb.Application
    for i in range(int(wb.Queries.Count), 0, -1):
        if str(wb.Queries.Item(i).Name) == name:
            wb.Queries.Item(i).Delete()
    for sh in list(wb.Worksheets):
        if str(sh.Name) == sheet:
            xl.DisplayAlerts = False
            sh.Delete()
    wb.Queries.Add(name, formula)
    ws = wb.Worksheets.Add(After=wb.Worksheets(wb.Worksheets.Count))
    ws.Name = sheet
    conn = f"OLEDB;Provider=Microsoft.Mashup.OleDb.1;Data Source=$Workbook$;Location={name};Extended Properties=\"\""
    lo = ws.ListObjects.Add(SourceType=0, Source=conn, Destination=ws.Range("$A$1"))
    lo.QueryTable.CommandType = 2
    lo.QueryTable.CommandText = f"SELECT * FROM [{name}]"
    lo.QueryTable.Refresh(False)
    lo.Name = name
    return lo


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            name = t['table'] + '_縦持ち'
            write_before(before, t)
            if t['stale']:                                   # 前のクエリ（値の列の名前が古い形）とシートを残す
                wb = xl.Workbooks.Open(before)
                try:
                    load_query(wb, name, m_formula(t, value_name='金額'), name)
                    wb.Worksheets(1).Activate()
                    wb.Save()
                finally:
                    wb.Close(SaveChanges=False)
            wb = xl.Workbooks.Open(before)
            try:
                load_query(wb, name, m_formula(t), name)
                wb.Worksheets(1).Activate()
                wb.SaveAs(truth, FileFormat=51)
            finally:
                wb.Close(SaveChanges=False)
            print('  ', stem, t['sheet'], t['table'], t['ka'], len(t['heads']))
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', build(s(1), rng.randint(3, 8), rng.randint(3, 12))),
          ('未見2_合計の列と行・空欄', build(s(2), 7, 9, blanks=True, total_col=True, total_row=True)),
          ('未見3_コードなしと前のクエリ', build(s(3), 5, 12, code_col=False, stale=True)),
          ('未見4_四半期と備考', build(s(4), 6, 4, 'quarter', note_col=True))], out)
    return out


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('縦持ち_本番', build(3, 5, 6)),
          ('縦持ち_試験A', build(13, 10, 12)),
          ('縦持ち_試験B', build(23, 6, 6, total_col=True)),
          ('縦持ち_試験C', build(33, 5, 8, code_col=False)),
          ('縦持ち_試験D', build(43, 6, 6, blanks=True)),
          ('縦持ち_試験E', build(53, 4, 6, stale=True)),
          ('縦持ち_試験G', build(63, 5, 6, total_row=True)),
          ('縦持ち_試験H', build(73, 6, 4, 'quarter', note_col=True))], HERE)
    shutil.copyfile(os.path.join(HERE, '縦持ち_本番_正解.xlsx'), os.path.join(HERE, '縦持ち_試験F_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, '縦持ち_本番_正解.xlsx'), os.path.join(HERE, '縦持ち_試験F_正解.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
