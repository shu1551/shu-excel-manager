# -*- coding: utf-8 -*-
"""職員名簿をテーブルにする（鍛える回路の題材・テーブル 1 本目・2026-09-17 深夜）。

仕事: 「職員名簿」シートの表（所属・氏名・採用日・基本給・手当）をテーブルにし、計算列と集計行を付ける。
正解の決まり（依頼文に全部書く）:
  - テーブル名 T職員名簿・スタイル TableStyleMedium2・行の縞模様あり・フィルタのボタンあり。
  - 範囲は見出しから最後の職員の行まで（表の下の※注記と空行は入れない）。
  - 右端に計算列「支給額」＝ =[@基本給]+[@手当]。
  - 集計行を出し、基本給・手当・支給額は合計（ほかの列は集計なし）。
  - 既にテーブルがあれば、それを直す（同じ表に 2 つ作らない・名前とスタイルを決まりどおりに）。
正解のブックは、直す前のブックを自分の Excel（非表示）で開き、この決まりどおりにテーブルを作って保存する。
組: 本番 12 人／A 300 人／B 表題・見出しが 3 行目／C 表の下に※注記／D 既にテーブル（テーブル1・別のスタイル）／
  E 列の並びが違う（手当が先）／F 撃った後にもう一度
"""
import datetime
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
SHOZOKU = ['総務課', '財政課', '企画課', '税務課', '福祉課', '建設課', '農林課']
SEI = ['佐藤', '鈴木', '高橋', '田中', '伊藤', '渡辺', '山本', '中村', '小林', '加藤', '吉田', '山田']
MEI = ['一郎', '花子', '誠', '恵', '学', '悟', '由美', '健太', '真理', '修', '陽子', '大輔']
HEAD = ['所属', '氏名', '採用日', '基本給', '手当']
SWAP = ['氏名', '所属', '手当', '採用日', '基本給']
SHEET = '職員名簿'
REQUEST = ('職員名簿の表をテーブルにしてください。テーブル名は T職員名簿、スタイルは TableStyleMedium2、行の縞模様あり、フィルタのボタンあり。'
           '範囲は見出しから最後の職員の行まで（表の下の※注記と空行は入れない）。右端に計算列「支給額」（=[@基本給]+[@手当]）を足し、'
           '集計行を出して基本給・手当・支給額は合計にする（ほかの列は集計なし）。既にテーブルがあればそれを直す（2 つ作らない）')


def build(seed, n):
    rng = random.Random(seed)
    return [{'所属': rng.choice(SHOZOKU), '氏名': rng.choice(SEI) + ' ' + rng.choice(MEI),
             '採用日': datetime.datetime(rng.randint(1990, 2025), rng.randint(1, 12), 1),
             '基本給': rng.randrange(180, 450) * 1000, '手当': rng.randrange(0, 80) * 1000} for _ in range(n)]


def write_before(path, rows, cols=HEAD, title=False, note=False):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    h = 1
    if title:
        ws.cell(1, 1, '令和8年4月1日現在 職員名簿').font = Font(bold=True, size=14)
        h = 3
    for j, c in enumerate(cols, 1):
        ws.cell(h, j, c).font = Font(bold=True)
    for i, r in enumerate(rows, 1):
        for j, c in enumerate(cols, 1):
            cell = ws.cell(h + i, j, r[c])
            if c == '採用日':
                cell.number_format = 'yyyy/m/d'
            if c in ('基本給', '手当'):
                cell.number_format = '#,##0'
    last = h + len(rows)
    if note:
        ws.cell(last + 2, 1, '※ 手当は扶養手当と住居手当の合計')
    wb.save(path)
    return {'h': h, 'last': last, 'ncols': len(cols)}


def build_table(ws, lay, name='T職員名簿', style='TableStyleMedium2', calc=True):
    """決まりどおりのテーブル（COM）。既存のテーブルがあればそれを直す。"""
    rng = ws.Range(ws.Cells(lay['h'], 1), ws.Cells(lay['last'], lay['ncols']))
    lo = None
    if int(ws.ListObjects.Count):
        lo = ws.ListObjects(1)
    else:
        lo = ws.ListObjects.Add(1, rng, None, 1)
    lo.Name = name
    lo.TableStyle = style
    lo.ShowTableStyleRowStripes = True
    lo.ShowAutoFilter = True
    if not calc:
        return lo
    names = [str(lo.ListColumns(i).Name) for i in range(1, int(lo.ListColumns.Count) + 1)]
    if '支給額' not in names:
        lc = lo.ListColumns.Add()
        lc.Name = '支給額'
    lc = lo.ListColumns('支給額')
    lc.DataBodyRange.Formula = '=[@基本給]+[@手当]'
    lc.DataBodyRange.NumberFormat = '#,##0'
    lo.ShowTotals = True
    for i in range(1, int(lo.ListColumns.Count) + 1):
        c = lo.ListColumns(i)
        c.TotalsCalculation = 1 if str(c.Name) in ('基本給', '手当', '支給額') else 0
    return lo


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, rows, kw in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            existing = kw.pop('existing', False)
            lay = write_before(before, rows, **kw)
            if existing:                                     # 既にテーブル（名前テーブル1・Light9・計算列も集計行も無し）
                wb = xl.Workbooks.Open(before)
                try:
                    build_table(wb.Sheets(SHEET), lay, name='テーブル1', style='TableStyleLight9', calc=False)
                    wb.Save()
                finally:
                    wb.Close(SaveChanges=False)
            wb = xl.Workbooks.Open(before)
            try:
                build_table(wb.Sheets(SHEET), lay)
                wb.SaveAs(truth, FileFormat=51)
            finally:
                wb.Close(SaveChanges=False)
            print('  ', stem)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', build(s(1), rng.randint(5, 40)), {}),
          ('未見2_表題と注記', build(s(2), 20), {'title': True, 'note': True}),
          ('未見3_並び違いと既存', build(s(3), 15), {'cols': SWAP, 'existing': True}),
          ('未見4_大量', build(s(4), 2000), {})], out)
    return out


if __name__ == '__main__':
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        sys.exit(0)
    make([('テーブル_本番', build(3, 12), {}),
          ('テーブル_試験A', build(13, 300), {}),
          ('テーブル_試験B', build(23, 10), {'title': True}),
          ('テーブル_試験C', build(33, 14), {'note': True}),
          ('テーブル_試験D', build(43, 11), {'existing': True}),
          ('テーブル_試験E', build(53, 9), {'cols': SWAP})], HERE)
    shutil.copyfile(os.path.join(HERE, 'テーブル_本番_正解.xlsx'), os.path.join(HERE, 'テーブル_試験F_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, 'テーブル_本番_正解.xlsx'), os.path.join(HERE, 'テーブル_試験F_正解.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    lines = ['職員名簿をテーブル化して', '名簿をテーブルにして支給額の列と集計行も', '職員名簿を T職員名簿 のテーブルに',
             'テーブルにして計算列と集計行を付けて', '職員名簿のテーブルをお願いします', '名簿の範囲をテーブルに変換して集計行を表示']
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(lines) + '\n')
    print('ok')
