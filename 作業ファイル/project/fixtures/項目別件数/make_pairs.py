# -*- coding: utf-8 -*-
"""選んでいる 1 列で「項目ごとの件数の表」を明細の右に作る（B 集計 13・2026-09-18）。

列の決まり: 選んでいる 1 列＝数える項目の列。選んでいる列は <名前>_選択.txt に書く。
正解の決まり（依頼文に全部書く）:
  - 置き場所＝明細の右に 1 列空けた所。見出しの行は明細の見出しの行と同じ行。
  - 見出しは［元の列の見出し］と「件数」。
  - 項目は明細に出てくる順（初めて出た順）。値は =COUNTIF(項目の列の本文, 項目のセル)。
  - いちばん下に「合計」の行＝右に =SUM（作った件数の範囲）。
  - 同じ表が既に右にあれば作り直す。
組: 本番 60 行／A 800 行・項目 10／B 表題と単位の行／C 明細の下に合計の行／D 途中に空行／E 列の並びが違う／
  F 撃った後にもう一度／G 番号の列／H 項目が 2 つだけ
"""
import os
import random
import shutil
import sys

from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import detail   # noqa: E402

REQUEST = ('選んでいる 1 列（数える項目の列）で、項目ごとの件数の表を明細の右に 1 列空けて作ってください。'
           '明細の右端＝見出しの行で、明細のいちばん左の列から右へ見て空のセルに当たる手前の列（UsedRange の右端ではない。前に作った件数の表が右にあると UsedRange はそこまで広がる）。'
           '見出しの行は必ず「選んでいるセルの行」＝Selection.Cells(1,1).Row をそのまま使う（自分で見出しの行を探さない。数の列が 1 本だけの表もあるので、数の数では見分けられない）。'
           '件数の表の見出しの行も同じ行で、左の見出しは元の列の見出しをそのまま、右の見出しは「件数」。'
           '項目は明細に出てくる順（初めて出た順）に 1 行ずつ並べ、値は =COUNTIF(項目の列の本文, 項目のセル) の式'
           '（本文の範囲は $ を付けた絶対参照）。いちばん下に合計の行を足し、左に「合計」・右に =SUM で件数の合計を入れる。'
           '明細の本文は見出しの次の行から最後の明細の行まで（明細の下に合計の行があればそれは入れない）。'
           '同じ表が既に右にあれば作り直す。')
PHRASES = ['選んだ列で項目ごとの件数表を右に作って', '項目別件数の表を明細の横に付けて', 'COUNTIFで項目ごとの件数を数えた表がほしい',
           '選んでいる列で件数の表を作成', '項目ごとに何件あるかの表を明細の右に', '項目別の件数表をCOUNTIFで']


def build_truth(xl, before, truth, t, lay):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(t['sheet'])
        c0, h, last = lay['c0'], lay['h'], lay['last']
        ic = c0 + detail.col_of(t, 'item') - 1
        iname = detail.name_of(t, 'item')
        il = get_column_letter(ic)
        irng = f"${il}${h + 1}:${il}${last}"
        out = lay['ncols'] + c0 + 1
        ws.Cells(h, out).Value = iname
        ws.Cells(h, out + 1).Value = '件数'
        seen = []
        for row in t['rows']:
            if row[iname] not in seen:
                seen.append(row[iname])
        for k, v in enumerate(seen, 1):
            ws.Cells(h + k, out).Value = v
            key = ws.Cells(h + k, out).GetAddress(False, False)
            ws.Cells(h + k, out + 1).Formula = f"=COUNTIF({irng},{key})"
        at = h + len(seen) + 1
        ws.Cells(at, out).Value = '合計'
        a = ws.Cells(h + 1, out + 1).GetAddress(False, False)
        b = ws.Cells(h + len(seen), out + 1).GetAddress(False, False)
        ws.Cells(at, out + 1).Formula = f"=SUM({a}:{b})"
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            lay = detail.write_before(before, t)
            sel = f"{get_column_letter(detail.col_of(t, 'item'))}{lay['h']}"
            build_truth(xl, before, truth, t, lay)
            open(os.path.join(out, f"{stem}_選択.txt"), 'w', encoding='utf-8').write(sel + '\n')
            print('  ', stem, t['sheet'], [n for n, _r in t['cols']], '選ぶ列', sel)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', detail.design(s(1), rng.randint(30, 200))),
          ('未見2_表題と合計行', detail.design(s(2), 60, title=True, unit=True, total=True)),
          ('未見3_並び違い', detail.design(s(3), 50, swap=True)),
          ('未見4_空行と番号列', detail.design(s(4), 40, blanks=2, no_col=True))], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('項目別件数_本番', detail.design(201, 60)),
          ('項目別件数_試験A', detail.design(202, 800, n_items=10)),
          ('項目別件数_試験B', detail.design(203, 50, title=True, unit=True)),
          ('項目別件数_試験C', detail.design(204, 40, total=True)),
          ('項目別件数_試験D', detail.design(205, 45, blanks=2)),
          ('項目別件数_試験E', detail.design(206, 50, swap=True)),
          ('項目別件数_試験G', detail.design(207, 35, no_col=True)),
          ('項目別件数_試験H', detail.design(208, 30, n_items=2)),
          ('項目別件数_試験I', detail.design(209, 26, no_date=True))], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '項目別件数_本番_正解.xlsx'), os.path.join(HERE, f'項目別件数_試験F_{x}.xlsx'))
    shutil.copyfile(os.path.join(HERE, '項目別件数_本番_選択.txt'), os.path.join(HERE, '項目別件数_試験F_選択.txt'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
