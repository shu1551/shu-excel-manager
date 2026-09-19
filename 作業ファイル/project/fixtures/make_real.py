# -*- coding: utf-8 -*-
"""実物らしい表（2026-09-17 夜）: 鍛えた 5 本（振り直し・突合・集約・帳票一覧・月次集計）に、職場の表で本当に出る癖を当てる。

make_pairs.py の組（表題・並び違い・空行・合計行・文字の金額…）は、マクロを鍛えたときと同じ頭で作った癖。
ここでは別の頭で、実物の表に出る癖を型ごとに 12〜16 種ずつ作る。正解の決まりは各 make_pairs.py の決まりのまま
（人が手でやったらこうなる、が一つに決まる癖だけを入れる）。

  共通: オートフィルタと枠固定・テーブル（ListObject）・見出しの改行と全角の空白と単位・表の下の※注記・小計の行・
        戻入（△の負の数・△の文字）・金額が数式・非表示の列・全角のコード・システム出力の広い表・大量
  型ごと: 款の見出し行・対応表の表題/注記/重複/全角/区分名の空白（振り直し）／新の合計行・新が B 列から・番号が数（突合）／
        記入例の行・見出し 2 段・記入者の行・シート名に番号・例だけの課（集約）／所在地の改行・和暦の表示・余白の列・
        途中の空の枠・合計行・表題ごと改ページ・申請者 3 列（帳票）／8 桁の数の日付・和暦の文字・日時・課ごとの小計（月次）

  py make_real.py [種]  → fixtures\\<型>\\未見_実物\\*_前.xlsx と *_正解.xlsx
"""
import datetime
import importlib.util
import os
import random
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
ZEN = str.maketrans('0123456789-,', '０１２３４５６７８９－，')
MINUS = '#,##0;[Red]△#,##0'
WAREKI = '[$-ja-JP-x-gannen]ggge"年"m"月"d"日"'
WRAP = Alignment(wrap_text=True, vertical='center')
THIN = Side(style='thin')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _load(kind):
    spec = importlib.util.spec_from_file_location(f"mp_{kind}", os.path.join(HERE, kind, 'make_pairs.py'))
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


FR, TG, SY, CH, GJ = (_load(k) for k in ('振り直し', '突合', '集約', '帳票', '月次'))


SEED = 1


def _out(kind):
    p = os.path.join(HERE, kind, '未見_実物' if SEED == 1 else f'未見_実物種{SEED}')
    os.makedirs(p, exist_ok=True)
    for f in os.listdir(p):
        if f.endswith('.xlsx'):
            os.remove(os.path.join(p, f))
    return p


def post(out, stem, fn):
    """write_pair が作った 前・正解 の 2 冊に同じ手を当てる（fn(wb, truth)）。"""
    for truth in (False, True):
        p = os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx")
        wb = load_workbook(p)
        fn(wb, truth)
        wb.save(p)


def data_rows(h, n, blanks):
    """write_pair と同じ数え方で、明細の行番号と、最後の明細の次の行を返す。"""
    rows, r = [], h + 1
    for i in range(n):
        if i in blanks:
            r += 1
        rows.append(r)
        r += 1
    return rows, r


def split(a, unit=1000):
    """金額を 2 つの数の足し算の式に（=6000+6000）。"""
    if a <= 0:
        return f"={a}"
    a1 = (a // 2) // unit * unit
    return f"={a1}+{a - a1}"


def zen(ws, rows, key, col, p=0.4):
    r2 = random.Random(key)
    for r in rows:
        v = ws.cell(r, col).value
        if v is not None and r2.random() < p:
            ws.cell(r, col).value = str(v).translate(ZEN)


def subtotals(ws, rows, blanks, name_col, amt_col, label='小計'):
    """空行（blanks の位置）に小計の行を入れる（金額は直前の小計からの SUM）。"""
    start = rows[0]
    for i in sorted(blanks):
        br = rows[i] - 1
        ws.cell(br, name_col, label).font = Font(bold=True)
        ws.cell(br, amt_col, f"=SUM({L(amt_col)}{start}:{L(amt_col)}{br - 1})").number_format = '#,##0'
        start = br + 1


def add_table(ws, ref, name):
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(t)


# ============================================================ 振り直し
def furi(seed):
    out = _out('振り直し')
    rng = random.Random(seed)
    s = lambda k: seed * 100 + k      # noqa: E731

    def B(k, n=None, miss=None):
        order = (FR.ORDER1, FR.ORDER2)[rng.randrange(2)]
        return FR.build(s(k), n or rng.randint(20, 50), order, rng.randint(1, 3) if miss is None else miss)

    # 01 オートフィルタと枠固定
    stem = '実物01_フィルタと枠固定'
    t, b = B(1)
    FR.write_pair(stem, t, b, out=out)
    rows, end = data_rows(1, len(b), ())

    def f01(wb, truth):
        ws = wb['支出明細']
        ws.auto_filter.ref = f"A1:C{rows[-1]}"
        ws.freeze_panes = 'A2'
        mt = wb['対応表']
        mt.auto_filter.ref = f"A1:B{mt.max_row}"
        mt.freeze_panes = 'A2'
    post(out, stem, f01)

    # 02 明細のコードが全角（手入力）
    stem = '実物02_全角のコード'
    t, b = B(2)
    FR.write_pair(stem, t, b, out=out)
    rows2, _ = data_rows(1, len(b), ())
    post(out, stem, lambda wb, truth: zen(wb['支出明細'], rows2, stem, 1))

    # 03 表の下に※注記（コードの列に書いてある）
    stem = '実物03_表の下の注記'
    t, b = B(3)
    FR.write_pair(stem, t, b, out=out)
    _r, end3 = data_rows(1, len(b), ())

    def f03(wb, truth):
        ws = wb['支出明細']
        ws.cell(end3 + 1, 1, '※ 支出額は税込みです。')
        ws.cell(end3 + 2, 1, '※ 令和8年9月30日現在')
    post(out, stem, f03)

    # 04 戻入（負の数・△の表示）
    stem = '実物04_戻入の負の数'
    t, b = B(4, n=rng.randint(25, 45))
    for i in random.Random(stem).sample(range(len(b)), 4):
        b[i][2] = -b[i][2]
    FR.write_pair(stem, t, b, out=out)
    rows4, _ = data_rows(1, len(b), ())

    def f04(wb, truth):
        for r in rows4:
            wb['支出明細'].cell(r, 3).number_format = MINUS
    post(out, stem, f04)

    # 05 金額が数式・合計行が SUM
    stem = '実物05_金額が数式'
    t, b = B(5)
    FR.write_pair(stem, t, b, total=True, out=out)
    rows5, end5 = data_rows(1, len(b), ())

    def f05(wb, truth):
        ws = wb['支出明細']
        for r in rows5:
            ws.cell(r, 3).value = split(ws.cell(r, 3).value)
        ws.cell(end5, 3).value = f"=SUM(C{rows5[0]}:C{rows5[-1]})"
    post(out, stem, f05)

    # 06 小計の行（科目名の列に「小計」・コードは空）
    stem = '実物06_小計の行'
    t, b = B(6, n=rng.randint(30, 45))
    bl = (10, 20)
    FR.write_pair(stem, t, b, blanks=bl, total=True, out=out)
    rows6, _ = data_rows(1, len(b), bl)
    post(out, stem, lambda wb, truth: subtotals(wb['支出明細'], rows6, bl, 2, 3))

    # 07 テーブル（ListObject）
    stem = '実物07_テーブル'
    t, b = B(7)
    FR.write_pair(stem, t, b, out=out)
    rows7, _ = data_rows(1, len(b), ())
    post(out, stem, lambda wb, truth: add_table(wb['支出明細'], f"A1:C{rows7[-1]}", 'T支出明細'))

    # 08 見出しの改行・全角の空白・単位
    stem = '実物08_見出しの改行と空白'
    t, b = B(8)
    FR.write_pair(stem, t, b, out=out)

    def f08(wb, truth):
        ws = wb['支出明細']
        for c, v in ((1, '予算科目\nコード'), (2, '科　目　名'), (3, '支出額\n（円）')):
            ws.cell(1, c).value = v
            ws.cell(1, c).alignment = WRAP
        ws.row_dimensions[1].height = 32
    post(out, stem, f08)

    # 09 非表示の列（科目名と支出額の間に所属）
    stem = '実物09_非表示の列'
    t, b = B(9)
    FR.write_pair(stem, t, b, out=out)
    rows9, _ = data_rows(1, len(b), ())

    def f09(wb, truth):
        ws = wb['支出明細']
        ws.insert_cols(3)
        ws.cell(1, 3, '所属').font = Font(bold=True)
        for r in rows9:
            ws.cell(r, 3, '総務課')
        ws.column_dimensions['C'].hidden = True
    post(out, stem, f09)

    # 10 システム出力の広い表（12 列・支出額の半分が文字）
    stem = '実物10_システム出力'
    t, b = B(10, n=rng.randint(30, 60))
    fr_system(stem, t, b, out)

    # 11 款の見出し行（コードの列に「02 総務費」・金額なし）
    stem = '実物11_款の見出し行'
    t, b = B(11, n=rng.randint(30, 45))
    bl11 = (0, 11, 22)
    FR.write_pair(stem, t, b, blanks=bl11, out=out)
    rows11, _ = data_rows(1, len(b), bl11)

    def f11(wb, truth):
        ws = wb['支出明細']
        for i, lab in zip(bl11, ('02 総務費', '03 民生費', '08 土木費')):
            ws.cell(rows11[i] - 1, 1, lab).font = Font(bold=True)
    post(out, stem, f11)

    # 12 対応表に表題・重複・※注記
    stem = '実物12_対応表の表題と注記'
    t, b = B(12)
    FR.write_pair(stem, t, b, out=out)

    def f12(wb, truth):
        mt = wb['対応表']
        n = mt.max_row
        mt.insert_rows(1, 2)
        mt.cell(1, 1, '決算統計区分 対応表（令和8年度）').font = Font(bold=True, size=13)
        dup = 3 + max(1, (n - 1) // 2)                          # 真ん中あたりの行をもう一度（同じ対応）
        last = mt.max_row
        mt.cell(last + 1, 1, mt.cell(dup, 1).value).number_format = '@'
        mt.cell(last + 1, 2, mt.cell(dup, 2).value)
        mt.cell(last + 3, 1, '※ 新設の細節は財政課に確認すること')
    post(out, stem, f12)

    # 13 対応表のコードが全角・後ろに空白
    stem = '実物13_対応表の全角'
    t, b = B(13)
    FR.write_pair(stem, t, b, out=out)

    def f13(wb, truth):
        mt = wb['対応表']
        r2 = random.Random(stem)
        for r in range(2, mt.max_row + 1):
            v = mt.cell(r, 1).value
            x = r2.random()
            if v and x < 0.35:
                mt.cell(r, 1).value = str(v).translate(ZEN)
            elif v and x < 0.55:
                mt.cell(r, 1).value = str(v) + ' '
    post(out, stem, f13)

    # 14 対応表の区分名の後ろに空白（半角・全角）
    stem = '実物14_区分名の空白'
    t, b = B(14)
    FR.write_pair(stem, t, b, out=out)

    def f14(wb, truth):
        mt = wb['対応表']
        r2 = random.Random(stem)
        for r in range(3, mt.max_row + 1):
            v = mt.cell(r, 2).value
            if v and r2.random() < 0.3:
                mt.cell(r, 2).value = v + r2.choice((' ', '　'))
    post(out, stem, f14)

    # 15 全部入り（表題・摘要・非表示の列・小計・合計・戻入・全角・注記・対応表の表題と注記・フィルタ）
    stem = '実物15_全部入り'
    t, b = B(15, n=40, miss=3)
    for i in random.Random(stem).sample(range(len(b)), 3):
        b[i][2] = -b[i][2]
    bl15 = (13, 27)
    FR.write_pair(stem, t, b, title='令和8年度 支出明細（観光課）', memo=True, blanks=bl15, total=True, out=out)
    rows15, end15 = data_rows(3, len(b), bl15)

    def f15(wb, truth):
        ws = wb['支出明細']
        ws.insert_cols(3)                                     # 先に列を入れる（式の番地がずれないように）
        ws.cell(3, 3, '所属').font = Font(bold=True)
        for r in rows15:
            ws.cell(r, 3, '観光課')
            ws.cell(r, 4).number_format = MINUS
        ws.column_dimensions['C'].hidden = True
        subtotals(ws, rows15, bl15, 2, 4)
        zen(ws, rows15, stem, 1)
        ws.cell(end15 + 2, 1, '※ 支出額は税込みです。')
        ws.cell(end15 + 3, 1, '※ △は戻入')
        ws.auto_filter.ref = f"A3:E{end15}"
        ws.freeze_panes = 'A4'
        f12(wb, truth)
    post(out, stem, f15)

    # 16 大量（5,000 行・全角・戻入・注記）
    stem = '実物16_大量'
    t, b = B(16, n=5000, miss=30)
    for i in random.Random(stem).sample(range(len(b)), 40):
        b[i][2] = -b[i][2]
    FR.write_pair(stem, t, b, out=out)
    rows16, end16 = data_rows(1, len(b), ())

    def f16(wb, truth):
        ws = wb['支出明細']
        zen(ws, rows16, stem, 1, p=0.1)
        ws.cell(end16 + 1, 1, '※ 支出額は税込みです。')
        ws.freeze_panes = 'A2'
    post(out, stem, f16)
    return out


def fr_system(stem, table, body, out):
    kmap = dict(table)
    kubun = list(dict.fromkeys(k for _c, k in table))
    heads = ['年度', '会計', '所属コード', '所属名', '支出命令番号', '支払日', '予算科目コード', '科目名', '債権者名', '支出額', '摘要', '処理状況']
    rng = random.Random(stem)
    lines = []
    for i, (c, name, amt) in enumerate(body):
        d = datetime.datetime(2026, rng.randint(4, 9), rng.randint(1, 28))
        lines.append(['2026', '一般会計', '0101', '総務課', f"S{i + 1:06d}", d, c, name,
                      rng.choice(['株式会社秋田商事', '有限会社北都', '個人（非公表）', '架空市役所']),
                      str(amt) if rng.random() < 0.5 else amt, f"{name}の支払", '支払済'])
    for truth in (False, True):
        wb = Workbook()
        ws = wb.active
        ws.title = '支出明細'
        for j, v in enumerate(heads + (['決算統計区分'] if truth else []), 1):
            ws.cell(1, j, v).font = Font(bold=True)
        for i, line in enumerate(lines, 2):
            for j, v in enumerate(line, 1):
                cell = ws.cell(i, j, v)
                if j in (1, 3, 5, 7) or isinstance(v, str) and j == 10:
                    cell.number_format = '@'
                if j == 6:
                    cell.number_format = 'yyyy/m/d'
            if truth:
                ws.cell(i, 13, kmap.get(body[i - 2][0], '対応なし'))
        if truth:
            sums = {k: 0 for k in kubun}
            none = 0
            for c, _n, a in body:
                if c in kmap:
                    sums[kmap[c]] += a
                else:
                    none += a
            ws.cell(1, 15, '決算統計区分')
            ws.cell(1, 16, '支出額')
            rr = 2
            for k in kubun + ['対応なし']:
                ws.cell(rr, 15, k)
                ws.cell(rr, 16, sums.get(k, none if k == '対応なし' else 0))
                rr += 1
            ws.cell(rr, 15, '合計')
            ws.cell(rr, 16, sum(a for _c, _n, a in body))
        mt = wb.create_sheet('対応表')
        mt.cell(1, 1, '予算科目コード')
        mt.cell(1, 2, '決算統計区分')
        for i, (c, k) in enumerate(table, 2):
            mt.cell(i, 1, c).number_format = '@'
            mt.cell(i, 2, k)
        wb.save(os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx"))


# ============================================================ 突合
def totsu(seed):
    out = _out('突合')
    rng = random.Random(seed)
    s = lambda k: seed * 100 + k      # noqa: E731

    def B(k, n=None, diff=None, miss=None, extra=None):
        return TG.build(s(k), n or rng.randint(20, 45), rng.randint(1, 4) if diff is None else diff,
                        rng.randint(1, 3) if miss is None else miss, rng.randint(1, 3) if extra is None else extra)

    def key_rows(ns, first=2):
        out_ = {}
        for r in range(first, ns.max_row + 1):
            v = ns.cell(r, 1).value
            if v is not None:
                out_[str(v)] = r
        return out_

    # 01 フィルタと枠固定（旧・新）
    stem = '実物01_フィルタと枠固定'
    old, new = B(1)
    TG.write_pair(stem, old, new, out=out)
    rows, _ = data_rows(1, len(old), ())

    def f01(wb, truth):
        ws = wb['旧システム']
        ws.auto_filter.ref = f"A1:C{rows[-1]}"
        ws.freeze_panes = 'A2'
        ns = wb['新システム']
        ns.auto_filter.ref = f"A1:B{ns.max_row}"
        ns.freeze_panes = 'A2'
    post(out, stem, f01)

    # 02 旧の伝票番号が全角
    stem = '実物02_旧の番号が全角'
    old, new = B(2)
    TG.write_pair(stem, old, new, out=out)
    rows2, _ = data_rows(1, len(old), ())
    post(out, stem, lambda wb, truth: zen(wb['旧システム'], rows2, stem, 1))

    # 03 旧の下に※注記（伝票番号の列）
    stem = '実物03_旧の下の注記'
    old, new = B(3)
    TG.write_pair(stem, old, new, out=out)
    _r, end3 = data_rows(1, len(old), ())

    def f03(wb, truth):
        ws = wb['旧システム']
        ws.cell(end3 + 1, 1, '※ 令和9年1月15日の出力です。')
        ws.cell(end3 + 2, 1, '※ 取消伝票は含みません。')
    post(out, stem, f03)

    # 04 新の下に合計行と注記
    stem = '実物04_新の合計行と注記'
    old, new = B(4)
    TG.write_pair(stem, old, new, out=out)

    def f04(wb, truth):
        ns = wb['新システム']
        r = 1 + len(new) + 1
        ns.cell(r, 1, '合計').font = Font(bold=True)
        ns.cell(r, 2, sum(new.values())).number_format = '#,##0'
        ns.cell(r + 2, 1, f'※ 出力件数 {len(new)} 件')
    post(out, stem, f04)

    # 05 戻入（旧は負の数に△の表示・新は「△12,000」の文字）
    stem = '実物05_戻入'
    old, new = B(5, n=30)
    both = [i for i, (k, _m, _a) in enumerate(old) if k in new]
    negs = set()
    for i in random.Random(stem).sample(both, 4):
        k, m, a = old[i]
        old[i] = (k, m, -a)
        new[k] = -new[k]
        negs.add(k)
    TG.write_pair(stem, old, new, out=out)
    rows5, _ = data_rows(1, len(old), ())

    def f05(wb, truth):
        ws = wb['旧システム']
        for r in rows5:
            ws.cell(r, 3).number_format = MINUS
        ns = wb['新システム']
        for key, r in key_rows(ns).items():
            if key.isdigit() and int(key) in negs:
                ns.cell(r, 2).value = f"△{-new[int(key)]:,}"
    post(out, stem, f05)

    # 06 見出しの改行・全角の空白・単位
    stem = '実物06_見出しの改行と空白'
    old, new = B(6)
    TG.write_pair(stem, old, new, out=out)

    def f06(wb, truth):
        ws = wb['旧システム']
        for c, v in ((1, '伝票\n番号'), (2, '科　目　名'), (3, '支出額\n（円）')):
            ws.cell(1, c).value = v
            ws.cell(1, c).alignment = WRAP
        ns = wb['新システム']
        ns.cell(1, 1).value = '伝票\n番号'
        ns.cell(1, 2).value = '支出額（円）'
    post(out, stem, f06)

    # 07 新がシステム出力の広い表（10 列・伝票番号が 3 列目）
    stem = '実物07_新が広い表'
    old, new = B(7)
    TG.write_pair(stem, old, new, out=out)
    post(out, stem, lambda wb, truth: tg_new_wide(wb, new, stem))

    # 08 旧に非表示の列
    stem = '実物08_旧に非表示の列'
    old, new = B(8)
    TG.write_pair(stem, old, new, out=out)
    rows8, _ = data_rows(1, len(old), ())

    def f08(wb, truth):
        ws = wb['旧システム']
        ws.insert_cols(3)
        ws.cell(1, 3, '所属').font = Font(bold=True)
        for r in rows8:
            ws.cell(r, 3, '建設課')
        ws.column_dimensions['C'].hidden = True
    post(out, stem, f08)

    # 09 旧の金額が数式
    stem = '実物09_金額が数式'
    old, new = B(9)
    TG.write_pair(stem, old, new, out=out)
    rows9, _ = data_rows(1, len(old), ())

    def f09(wb, truth):
        ws = wb['旧システム']
        for r in rows9:
            ws.cell(r, 3).value = split(ws.cell(r, 3).value)
    post(out, stem, f09)

    # 10 新が 2 行目・B 列から（A 列が空・1 行目に表題）
    stem = '実物10_新がB列から'
    old, new = B(10)
    TG.write_pair(stem, old, new, out=out)

    def f10(wb, truth):
        ns = wb['新システム']
        ns.insert_rows(1)
        ns.insert_cols(1)
        ns.cell(1, 2, '新財務会計システム 支出一覧').font = Font(bold=True)
        ns.column_dimensions['A'].width = 2
    post(out, stem, f10)

    # 11 旧に小計の行
    stem = '実物11_旧に小計'
    old, new = B(11, n=rng.randint(26, 40))
    bl = (8, 17)
    TG.write_pair(stem, old, new, blanks=bl, total=True, out=out)
    rows11, _ = data_rows(1, len(old), bl)
    post(out, stem, lambda wb, truth: subtotals(wb['旧システム'], rows11, bl, 2, 3))

    # 12 旧の伝票番号が数（先頭の 0 なし）と文字の混在
    stem = '実物12_旧の番号が数'
    old, new = B(12)
    TG.write_pair(stem, old, new, out=out)
    rows12, _ = data_rows(1, len(old), ())

    def f12(wb, truth):
        ws = wb['旧システム']
        r2 = random.Random(stem)
        for r in rows12:
            if r2.random() < 0.5:
                c = ws.cell(r, 1)
                c.value = int(c.value)
                c.number_format = 'General'
    post(out, stem, f12)

    # 13 新に同じ見出しの別の表（件数の表が上にある・伝票の表は 6 行目から）
    stem = '実物13_新の上に件数の表'
    old, new = B(13)
    TG.write_pair(stem, old, new, out=out)

    def f13(wb, truth):
        ns = wb['新システム']
        ns.insert_rows(1, 5)
        ns.cell(1, 1, '出力条件').font = Font(bold=True)
        ns.cell(2, 1, '会計年度')
        ns.cell(2, 2, '令和8年度')
        ns.cell(3, 1, '出力件数')
        ns.cell(3, 2, len(new))
    post(out, stem, f13)

    # 14 大量（5,000 件・全角・注記）
    stem = '実物14_大量'
    old, new = TG.build(s(14), 5000, 80, 50, 40)
    TG.write_pair(stem, old, new, key_num=True, out=out)
    rows14, end14 = data_rows(1, len(old), ())

    def f14(wb, truth):
        ws = wb['旧システム']
        zen(ws, rows14, stem, 1, p=0.1)
        ws.cell(end14 + 1, 1, '※ 令和9年1月15日の出力です。')
    post(out, stem, f14)

    # 15 全部入り（摘要・非表示の列・小計・合計・戻入・全角・注記・新は広い表に合計行と注記と△の文字・フィルタ）
    stem = '実物15_全部入り'
    old, new = B(15, n=36, diff=4, miss=3, extra=3)
    both = [i for i, (k, _m, _a) in enumerate(old) if k in new]
    negs15 = set()
    for i in random.Random(stem).sample(both, 3):
        k, m, a = old[i]
        old[i] = (k, m, -a)
        new[k] = -new[k]
        negs15.add(k)
    bl15 = (9, 21)
    TG.write_pair(stem, old, new, memo=True, blanks=bl15, total=True, out=out)
    rows15, end15 = data_rows(1, len(old), bl15)

    def f15(wb, truth):
        ws = wb['旧システム']
        ws.insert_cols(3)
        ws.cell(1, 3, '所属').font = Font(bold=True)
        for r in rows15:
            ws.cell(r, 3, '福祉課')
            ws.cell(r, 4).number_format = MINUS
        ws.column_dimensions['C'].hidden = True
        subtotals(ws, rows15, bl15, 2, 4)
        zen(ws, rows15, stem, 1)
        ws.cell(end15 + 2, 1, '※ 令和9年1月15日の出力です。')
        ws.auto_filter.ref = f"A1:E{end15}"
        ws.freeze_panes = 'A2'
        tg_new_wide(wb, new, stem, negs=negs15, total=True)
    post(out, stem, f15)
    return out


def tg_new_wide(wb, new, stem, negs=(), total=False):
    del wb['新システム']
    ns = wb.create_sheet('新システム')
    heads = ['会計年度', '伝票区分', '伝票番号', '起票日', '所属', '科目', '債権者', '支出額', '摘要', '状態']
    for j, v in enumerate(heads, 1):
        ns.cell(1, j, v).font = Font(bold=True)
    items = sorted(new.items())
    rng = random.Random(stem + '新')
    rng.shuffle(items)
    r = 2
    for k, a in items:
        vals = ['2026', '支出負担行為', f"{k:06d}", datetime.datetime(2027, 1, (k % 28) + 1), '総務課', rng.choice(TG.KAMOKU),
                rng.choice(['株式会社秋田商事', '有限会社北都', '架空市役所']), a, '', '確定']
        if k in negs:
            vals[7] = f"△{-a:,}"
        for j, v in enumerate(vals, 1):
            c = ns.cell(r, j, v)
            if j in (1, 3):
                c.number_format = '@'
            if j == 4:
                c.number_format = 'yyyy/m/d'
        r += 1
    if total:
        ns.cell(r, 3, '合計').font = Font(bold=True)
        ns.cell(r, 8, sum(new.values()))
        ns.cell(r + 2, 1, f'※ 出力件数 {len(new)} 件')


# ============================================================ 集約
def ka_sheet(ws, ka, items, key, order=('事業名', '予算額', '執行額'), title=False, heads=None, filt=False, example=False,
             subtotal=False, note=False, rate=False, memo=False, formula=False, col0=1, writer=False, two_tier=False,
             text_amount=False):
    rng = random.Random(key)
    r = 1
    if writer:
        ws.cell(1, col0, '記入者：')
        ws.cell(1, col0 + 1, rng.choice(CH.SEI) + ' ' + rng.choice(CH.MEI))
        r = 3
    if title:
        ws.cell(r, col0, f'{ka} 事業一覧（令和8年度）').font = Font(bold=True)
        ws.cell(r + 1, col0, '単位：円')
        r += 3
    cols = list(order) + (['執行率'] if rate else []) + (['備考'] if memo else [])
    C = {name: col0 + j for j, name in enumerate(cols)}
    hh = r
    if two_tier:
        ws.cell(hh, C['事業名'], '事業名')
        ws.merge_cells(start_row=hh, start_column=C['事業名'], end_row=hh + 1, end_column=C['事業名'])
        ws.cell(hh, C['予算額'], '金額（円）')
        ws.merge_cells(start_row=hh, start_column=C['予算額'], end_row=hh, end_column=C['執行額'])
        ws.cell(hh + 1, C['予算額'], '予算額')
        ws.cell(hh + 1, C['執行額'], '執行額')
        for other in cols[3:]:
            ws.cell(hh, C[other], other)
            ws.merge_cells(start_row=hh, start_column=C[other], end_row=hh + 1, end_column=C[other])
        r = hh + 2
    else:
        for name in cols:
            ws.cell(hh, C[name], (heads or {}).get(name, name)).font = Font(bold=True)
            if heads:
                ws.cell(hh, C[name]).alignment = WRAP
        r = hh + 1
    first = r

    def amount(v, rr):
        if formula:
            return split(v, 10000)
        if text_amount and rng.random() < 0.5:
            return f"{v:,}".translate(SY._ZEN) if rng.random() < 0.5 else f" {v:,} "
        return v

    def put(rr, jg, b, e, raw=False):
        ws.cell(rr, C['事業名'], jg)
        ws.cell(rr, C['予算額'], b if raw else amount(b, rr)).number_format = '#,##0'
        ws.cell(rr, C['執行額'], e if raw else amount(e, rr)).number_format = '#,##0'
        if rate:
            bl, el = L(C['予算額']), L(C['執行額'])
            ws.cell(rr, C['執行率'], f'=IFERROR(IF({bl}{rr}=0,0,{el}{rr}/{bl}{rr}),"")').number_format = '0.0%'
        if memo and rng.random() < 0.3:
            ws.cell(rr, C['備考'], rng.choice(['年度内完了予定', '繰越あり', '入札不調', '']))

    if example:
        put(r, '（例）庁舎清掃', 1200000, 800000, raw=True)
        r += 1
    mid = len(items) // 2 if subtotal and len(items) >= 4 else None
    seg = []
    for i, (jg, b, e) in enumerate(items):
        if mid is not None and i == mid:
            put(r, '小計', sum(x[1] for x in seg), sum(x[2] for x in seg), raw=True)
            r += 1
            seg = []
        put(r, jg, b, e)
        seg.append((jg, b, e))
        r += 1
    if subtotal and items:
        put(r, '小計', sum(x[1] for x in seg), sum(x[2] for x in seg), raw=True)
        r += 1
    last = r - 1
    if note:
        ws.cell(r + 1, C['事業名'], '※ 金額は税込みで記入してください。')
        ws.cell(r + 2, C['事業名'], '※ 9月30日までに財政課へ提出')
    if filt:
        ws.auto_filter.ref = f"{L(col0)}{hh + (1 if two_tier else 0)}:{L(col0 + len(cols) - 1)}{max(last, first)}"
        ws.freeze_panes = f"{L(col0)}{first}"
    return first


def sy_write(stem, data, out, opts=lambda k, ka: {}, extras=False, sum_title=False):
    for truth in (False, True):
        wb = Workbook()
        sm = wb.active
        sm.title = '集約'
        h = 3 if sum_title else 1
        if sum_title:
            sm.cell(1, 1, '令和8年度 事業執行状況（全課）').font = Font(bold=True, size=14)
        for j, v in enumerate(SY.HEAD, 1):
            sm.cell(h, j, v).font = Font(bold=True)
        if truth:
            r = h + 1
            tb = te = 0
            for ka, items in data:
                for jg, b, e in items:
                    for j, v in enumerate((ka, jg, b, e, (e / b) if b else 0), 1):
                        sm.cell(r, j, v)
                    tb += b
                    te += e
                    r += 1
            for j, v in ((1, '合計'), (3, tb), (4, te), (5, (te / tb) if tb else 0)):
                sm.cell(r, j, v)
        if extras:
            ex = wb.create_sheet('説明')
            ex.cell(1, 1, '各課は自分のシートに事業ごとの予算額と執行額を記入してください。')
            ex2 = wb.create_sheet('記入例')
            for j, v in enumerate(['事業名', '予算額', '執行額'], 1):
                ex2.cell(1, j, v)
            ex2.cell(2, 1, '（例）庁舎清掃')
            ex2.cell(2, 2, 1200000)
            ex2.cell(2, 3, 800000)
        for k, (ka, items) in enumerate(data):
            ka_sheet(wb.create_sheet(ka), ka, items, key=f"{stem}/{k}", **opts(k, ka))
        wb.save(os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx"))


def shuyaku(seed):
    out = _out('集約')
    rng = random.Random(seed)
    s = lambda k: seed * 100 + k      # noqa: E731
    D = lambda k, n=None, rows=(3, 9): SY.build(s(k), n or rng.randint(3, 6), rows=rows)  # noqa: E731
    sy_write('実物01_フィルタと枠固定', D(1), out, lambda k, ka: {'filt': True})
    sy_write('実物02_注記と小計', D(2), out, lambda k, ka: {'note': True, 'subtotal': True})
    sy_write('実物03_見出しの空白と改行', D(3), out,
             lambda k, ka: {'heads': {'事業名': '事　業　名', '予算額': '予算額\n（円）', '執行額': '執行額（円）'}})
    sy_write('実物04_執行率と備考の列', D(4), out, lambda k, ka: {'rate': True, 'memo': True})
    sy_write('実物05_金額が数式', D(5), out, lambda k, ka: {'formula': True})
    d6 = [(f"{i + 1:02d}_{ka}", items) for i, (ka, items) in enumerate(D(6))]
    sy_write('実物06_シート名に番号', d6, out)
    sy_write('実物07_記入例の行', D(7), out, lambda k, ka: {'example': True})
    sy_write('実物08_見出しが2段', D(8), out, lambda k, ka: {'two_tier': True, 'rate': k % 2 == 1})
    sy_write('実物09_記入者の行とB列から', D(9), out, lambda k, ka: {'writer': True, 'col0': 2})
    d10 = D(10, n=4)
    d10 = d10[:2] + [(d10[2][0], [])] + d10[3:]
    sy_write('実物10_例だけの課', d10, out, lambda k, ka: {'example': True, 'note': True})

    def o11(k, ka):
        r2 = random.Random(f"11/{seed}/{k}")
        o = {'example': r2.random() < 0.5, 'subtotal': r2.random() < 0.5, 'note': r2.random() < 0.5,
             'filt': r2.random() < 0.5, 'title': r2.random() < 0.4, 'memo': r2.random() < 0.4,
             'rate': r2.random() < 0.4, 'col0': r2.choice((1, 1, 2))}
        if r2.random() < 0.3:
            o['two_tier'] = True
        elif r2.random() < 0.5:
            o['order'] = ('執行額', '事業名', '予算額')
            o['heads'] = {'事業名': '事　業　名', '予算額': '予算額\n（円）'}
        else:
            o['text_amount'] = True
        return o
    d11 = [(f"{i + 1:02d}_{ka}", items) for i, (ka, items) in enumerate(D(11, n=7))]
    sy_write('実物11_全部入り', d11, out, o11, extras=True, sum_title=True)
    d12 = [(f"{i + 1:02d}_{ka}", items) for i, (ka, items) in enumerate(SY.build(s(12), 10, rows=(60, 120)))]
    sy_write('実物12_大量', d12, out, lambda k, ka: {'note': True, 'subtotal': True, 'filt': True})
    return out


# ============================================================ 帳票
def ch_cols(layout, c0, sub):
    pos, c = {}, c0
    for name in layout:
        if name == '申請者':
            for j, sname in enumerate(sub):
                pos[sname] = c + j
            c += len(sub)
        else:
            pos[name] = c
            c += 1
    return pos, c - 1


def ch_head(ws, r, layout, c0, sub):
    c = c0
    for name in layout:
        if name == '申請者':
            ws.cell(r, c, '申請者')
            ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + len(sub) - 1)
            for j, sname in enumerate(sub):
                ws.cell(r + 1, c + j, sname)
            for cc in range(c, c + len(sub)):
                for rr in (r, r + 1):
                    ws.cell(rr, cc).border = BOX
                    ws.cell(rr, cc).font = Font(bold=True)
                    ws.cell(rr, cc).alignment = Alignment(horizontal='center')
            c += len(sub)
        else:
            ws.cell(r, c, name)
            ws.merge_cells(start_row=r, start_column=c, end_row=r + 1, end_column=c)
            for rr in (r, r + 1):
                ws.cell(rr, c).border = BOX
                ws.cell(rr, c).font = Font(bold=True)
            ws.cell(r, c).alignment = Alignment(horizontal='center', vertical='center')
            c += 1


def ch_write(stem, recs, out, layout=('受付番号', '申請者', '申請額', '受付日'), c0=1, row0=1, title_rows=1, gaps=(),
             note=False, total=False, page=0, repeat_title=False, wareki=False, formula=False, sub=('団体名', '代表者')):
    labels = ['所在地'] + (['電話'] if recs and '電話' in recs[0] else [])
    k = 1 + len(labels)
    pos, last_col = ch_cols(layout, c0, sub)
    lab_c, val_c = pos[sub[0]], pos[sub[1]]
    fields = sorted(pos, key=lambda x: pos[x]) + labels
    frames, ri = [], 0
    for i in range(len(recs) + len(gaps)):
        if i in gaps:
            frames.append(None)
        else:
            frames.append(recs[ri])
            ri += 1
    TITLE = '令和8年度 地域づくり活動補助金 申請受付簿'
    for truth in (False, True):
        wb = Workbook()
        ws = wb.active
        ws.title = CH.SHEET
        h = row0
        if title_rows:
            ws.cell(row0, c0, TITLE).font = Font(bold=True, size=14)
            ws.merge_cells(start_row=row0, start_column=c0, end_row=row0, end_column=last_col)
            if title_rows == 2:
                ws.cell(row0 + 1, c0, '受付期間：令和8年4月1日～令和8年5月31日')
            h = row0 + title_rows
        ch_head(ws, h, layout, c0, sub)
        r = h + 2
        for i, rec in enumerate(frames):
            if page and i and i % page == 0:
                if repeat_title:
                    ws.cell(r, c0, TITLE).font = Font(bold=True, size=14)
                    ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=last_col)
                    r += 1
                ch_head(ws, r, layout, c0, sub)
                r += 2
            for name, c in pos.items():
                v = rec.get(name, '') if rec else ''
                if v != '':
                    cell = ws.cell(r, c, split(v) if (formula and name == '申請額') else v)
                    if name == '受付日':
                        cell.number_format = WAREKI if wareki else 'yyyy/m/d'
                    if name == '申請額':
                        cell.number_format = '#,##0'
                    if name == '受付番号':
                        cell.number_format = '@'
                    if name == '所在地':
                        cell.alignment = WRAP
                if name in sub:
                    continue
                ws.merge_cells(start_row=r, start_column=c, end_row=r + k - 1, end_column=c)
            for j, lab in enumerate(labels, 1):
                ws.cell(r + j, lab_c, lab)
                v = rec.get(lab, '') if rec else ''
                if v != '':
                    ws.cell(r + j, val_c, v).alignment = WRAP
            for rr in range(r, r + k):
                for cc in range(c0, last_col + 1):
                    ws.cell(rr, cc).border = BOX
            r += k
        if total:
            ws.cell(r, pos['受付番号'], '合計').font = Font(bold=True)
            ws.cell(r, lab_c, f"{len(recs)}件")
            ws.cell(r, pos['申請額'], sum(x['申請額'] for x in recs)).number_format = '#,##0'
            for cc in range(c0, last_col + 1):
                ws.cell(r, cc).border = BOX
            r += 1
        if note:
            ws.cell(r, c0, '※ 受付番号は受付順に付けています。申請額は千円未満を切り捨てています。')
            ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=last_col)
        lc = last_col + 2
        if truth:
            for j, f in enumerate(fields):
                ws.cell(h, lc + j, f)
            for i, rec in enumerate(recs, 1):
                for j, f in enumerate(fields):
                    v = rec.get(f, '')
                    if v != '':
                        ws.cell(h + i, lc + j, v)
        if c0 > 1:
            ws.column_dimensions['A'].width = 2
        for c in range(c0, last_col + 1):
            ws.column_dimensions[L(c)].width = 16
        wb.save(os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx"))


def chohyo(seed):
    out = _out('帳票')
    rng = random.Random(seed)
    s = lambda k: seed * 100 + k      # noqa: E731
    R = lambda k, n=None, **kw: CH.build(s(k), n or rng.randint(5, 14), **kw)  # noqa: E731

    recs = R(1)
    for x in recs:
        if random.Random(x['受付番号']).random() < 0.5:
            x['所在地'] += '\n（' + random.choice(['公民館内', '会館2階', '集会所', '事務局']) + '）'
    ch_write('実物01_所在地の改行', recs, out)
    ch_write('実物02_受付日の和暦表示', R(2), out, wareki=True)
    ch_write('実物03_申請額が数式', R(3), out, formula=True)
    ch_write('実物04_余白の列と行', R(4), out, c0=2, row0=2)
    n5 = rng.randint(6, 12)
    ch_write('実物05_途中の空の枠', R(5, n5), out, gaps=tuple(sorted(rng.sample(range(1, n5), 3))))
    ch_write('実物06_合計行と注記', R(6), out, total=True, note=True)
    ch_write('実物07_表題が2行', R(7, phone=True), out, title_rows=2)
    recs8 = R(8, phone=True)
    for x in recs8:
        if random.Random(x['受付番号'] + 'p').random() < 0.35:
            x['電話'] = random.Random(x['受付番号']).choice(['なし', '－', '（携帯のみ）'])
    ch_write('実物08_電話がなし', recs8, out)
    ch_write('実物09_表題ごと改ページ', R(9, rng.randint(16, 30)), out, page=6, repeat_title=True)
    ch_write('実物10_申請者が3列', [dict(x, 担当者=random.Random(x['受付番号']).choice(CH.SEI)) for x in R(10)], out,
             sub=('団体名', '代表者', '担当者'))
    recs11 = R(11, 28, phone=True, blank_phone=True, zero_id=True)
    for x in recs11:
        if random.Random(x['受付番号']).random() < 0.3:
            x['所在地'] += '\n（公民館内）'
    ch_write('実物11_全部入り', recs11, out, layout=CH.SWAP, c0=2, row0=2, title_rows=2, gaps=(4, 13, 22), total=True,
             note=True, page=8, repeat_title=True, wareki=True, formula=True)
    ch_write('実物12_大量', R(12, 500, zero_id=True), out, page=20, gaps=(50, 51, 300), total=True, note=True)
    return out


# ============================================================ 月次
def gj_write(stem, lines, out, cols=tuple(GJ.HEAD), title=False, heads=None, filt=False, table=False, hidden=(), fmt=None):
    """lines: ('d', 見せる値 {列: 値}, (日付, 課名, 金額)) ／ ('x', {列: 値})（小計・注記など集計に入れない行）／ ('b',)"""
    order, tot = [], {}
    for ln in lines:
        if ln[0] != 'd':
            continue
        d, ka, amt = ln[2]
        key = ka.strip().strip('　').strip()
        if key not in tot:
            order.append(key)
            tot[key] = {m: 0 for m in GJ.MONTHS}
        tot[key][d.month] += amt
    cols = list(cols)
    for truth in (False, True):
        wb = Workbook()
        ws = wb.active
        ws.title = GJ.SHEET
        h = 1
        if title:
            ws.cell(1, 1, '令和8年度 支出明細（一般会計）').font = Font(bold=True, size=14)
            ws.cell(2, 1, '単位：円')
            h = 4
        for j, c in enumerate(cols, 1):
            ws.cell(h, j, (heads or {}).get(c, c)).font = Font(bold=True)
            if heads:
                ws.cell(h, j).alignment = WRAP
        r = h + 1
        last = h
        for ln in lines:
            if ln[0] == 'b':
                r += 1
                continue
            for j, c in enumerate(cols, 1):
                v = ln[1].get(c)
                if v is None or v == '':
                    continue
                cell = ws.cell(r, j, v)
                if c == '支出日' and isinstance(v, datetime.datetime):
                    cell.number_format = 'yyyy/m/d h:mm' if (v.hour or v.minute) else 'yyyy/m/d'
                if c == '支出額' and not isinstance(v, str) or (c == '支出額' and str(v).startswith('=')):
                    cell.number_format = (fmt or {}).get('支出額', '#,##0')
                if c in ('所属コード', '伝票番号', '会計年度'):
                    cell.number_format = '@'
            last = r
            r += 1
        for c in hidden:
            ws.column_dimensions[L(cols.index(c) + 1)].hidden = True
        if filt:
            ws.auto_filter.ref = f"A{h}:{L(len(cols))}{last}"
            ws.freeze_panes = f"A{h + 1}"
        if table:
            add_table(ws, f"A{h}:{L(len(cols))}{last}", 'T支出明細')
            ws.freeze_panes = f"A{h + 1}"
        lc = len(cols) + 2
        if truth:
            ws.cell(h, lc, '課名')
            for j, m in enumerate(GJ.MONTHS, 1):
                ws.cell(h, lc + j, f"{m}月")
            ws.cell(h, lc + 13, '合計')
            for i, ka in enumerate(order, 1):
                ws.cell(h + i, lc, ka)
                for j, m in enumerate(GJ.MONTHS, 1):
                    ws.cell(h + i, lc + j, tot[ka][m])
                ws.cell(h + i, lc + 13, sum(tot[ka].values()))
            rr = h + len(order) + 1
            ws.cell(rr, lc, '合計')
            for j, m in enumerate(GJ.MONTHS, 1):
                ws.cell(rr, lc + j, sum(tot[ka][m] for ka in order))
            ws.cell(rr, lc + 13, sum(sum(tot[ka].values()) for ka in order))
        wb.save(os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx"))


def _dt(d):
    return datetime.datetime(d.year, d.month, d.day)


def plain(rows):
    return [('d', {'支出日': _dt(d), '課名': ka, '科目': km, '摘要': tk, '支出額': a}, (d, ka, a)) for d, ka, km, tk, a in rows]


def getsuji(seed):
    out = _out('月次')
    rng = random.Random(seed)
    s = lambda k: seed * 100 + k      # noqa: E731
    Bd = lambda k, n=None, **kw: GJ.build(s(k), n or rng.randint(40, 90), **kw)  # noqa: E731

    # 01 支出日が 8 桁の数（システムの CSV）
    lines = plain(Bd(1))
    r2 = random.Random('01')
    for ln in lines:
        if r2.random() < 0.5:
            d = ln[2][0]
            ln[1]['支出日'] = int(f"{d:%Y%m%d}")
    gj_write('実物01_8桁の数の日付', lines, out)

    # 02 支出日が和暦の文字
    lines = plain(Bd(2))
    r2 = random.Random('02')
    for ln in lines:
        d = ln[2][0]
        x = r2.random()
        if x < 0.25:
            ln[1]['支出日'] = f"R{d.year - 2018}.{d.month}.{d.day}"
        elif x < 0.5:
            ln[1]['支出日'] = f"令和{d.year - 2018}年{d.month}月{d.day}日"
        elif x < 0.6:
            ln[1]['支出日'] = f"R{d.year - 2018:02d}.{d.month:02d}.{d.day:02d}"
    gj_write('実物02_和暦の文字の日付', lines, out)

    # 03 表の下に※注記（支出日の列）
    lines = plain(Bd(3)) + [('b',), ('x', {'支出日': '※ 令和8年9月30日現在'}), ('x', {'支出日': '※ 戻入は含みません。'})]
    gj_write('実物03_表の下の注記', lines, out)

    # 04 課ごとに並べて小計（課名＋摘要「小計」）・最後に合計
    rows = sorted(Bd(4), key=lambda x: (x[1], x[0]))
    lines = []
    for ka in list(dict.fromkeys(x[1] for x in rows)):
        grp = [x for x in rows if x[1] == ka]
        lines += plain(grp)
        lines.append(('x', {'課名': ka, '摘要': '小計', '支出額': sum(x[4] for x in grp)}))
    lines.append(('x', {'摘要': '合計', '支出額': sum(x[4] for x in rows)}))
    gj_write('実物04_課ごとの小計', lines, out)

    # 05 戻入（負の数に△の表示・「△1,200」の文字）
    rows = Bd(5)
    r2 = random.Random('05')
    for x in r2.sample(rows, 6):
        x[4] = -x[4]
    lines = plain(rows)
    for ln in lines:
        if ln[2][2] < 0 and r2.random() < 0.5:
            ln[1]['支出額'] = f"△{-ln[2][2]:,}"
    gj_write('実物05_戻入', lines, out, fmt={'支出額': MINUS})

    # 06 見出しの全角の空白・改行・単位
    gj_write('実物06_見出しの空白と改行', plain(Bd(6)), out,
             heads={'課名': '課　名', '摘要': '摘　要', '支出額': '支出額\n（円）', '支出日': '支出\n年月日'})

    # 07 テーブルと枠固定
    gj_write('実物07_テーブル', plain(Bd(7)), out, table=True)

    # 08 金額が数式
    lines = plain(Bd(8))
    for ln in lines:
        ln[1]['支出額'] = split(ln[2][2], 100)
    gj_write('実物08_金額が数式', lines, out)

    # 09 非表示の列（課名と科目の間に所属コード）
    lines = plain(Bd(9))
    for ln in lines:
        ln[1]['所属コード'] = f"{(sum(map(ord, ln[2][1])) % 90) + 10:02d}01"
    gj_write('実物09_非表示の列', lines, out, cols=('支出日', '課名', '所属コード', '科目', '摘要', '支出額'), hidden=('所属コード',))

    # 10 支出日が日時（時刻つき）
    lines = plain(Bd(10))
    r2 = random.Random('10')
    for ln in lines:
        d = ln[2][0]
        ln[1]['支出日'] = datetime.datetime(d.year, d.month, d.day, r2.randint(8, 17), r2.randint(0, 59))
    gj_write('実物10_日時', lines, out)

    # 11 システム出力の広い表（13 列）
    lines = plain(Bd(11))
    r2 = random.Random('11')
    for i, ln in enumerate(lines):
        ln[1].update({'会計年度': '2026', '伝票番号': f"{i + 1:06d}", '所属コード': '0101', '款': '総務費', '項': '総務管理費',
                      '目': '一般管理費', '債権者': r2.choice(['株式会社秋田商事', '有限会社北都']), '支払方法': '口座振替'})
    gj_write('実物11_システム出力', lines, out,
             cols=('会計年度', '伝票番号', '支出日', '所属コード', '課名', '款', '項', '目', '科目', '摘要', '債権者', '支出額', '支払方法'))

    # 12 大量（10,000 行）
    gj_write('実物12_大量', plain(Bd(12, 10000, n_ka=12)), out, filt=True)

    # 13 全部入り
    rows = sorted(Bd(13, 120, n_ka=6, gap_months=(8,)), key=lambda x: (x[1], x[0]))
    r2 = random.Random('13')
    for x in r2.sample(rows, 5):
        x[4] = -x[4]
    lines = []
    for ka in list(dict.fromkeys(x[1] for x in rows)):
        grp = [x for x in rows if x[1] == ka]
        for ln in plain(grp):
            d, _k, a = ln[2]
            y = r2.random()
            if y < 0.2:
                ln[1]['支出日'] = int(f"{d:%Y%m%d}")
            elif y < 0.35:
                ln[1]['支出日'] = f"R{d.year - 2018}.{d.month}.{d.day}"
            if a < 0 and r2.random() < 0.5:
                ln[1]['支出額'] = f"△{-a:,}"
            elif a > 0 and r2.random() < 0.2:
                ln[1]['支出額'] = split(a, 100)
            if r2.random() < 0.2:
                ln[1]['課名'] = ka + r2.choice((' ', '　'))
            ln[1]['所属コード'] = '0101'
            lines.append(ln)
        lines.append(('x', {'課名': ka, '摘要': '小計', '支出額': sum(x[4] for x in grp)}))
    lines += [('x', {'摘要': '合計', '支出額': sum(x[4] for x in rows)}), ('b',), ('x', {'支出日': '※ 令和9年3月31日現在'})]
    gj_write('実物13_全部入り', lines, out, cols=('支出日', '課名', '所属コード', '科目', '摘要', '支出額'), title=True,
             heads={'課名': '課　名', '支出額': '支出額\n（円）'}, hidden=('所属コード',), filt=True, fmt={'支出額': MINUS})
    return out


if __name__ == '__main__':
    seed = SEED = int(sys.argv[1]) if len(sys.argv) > 1 and sys.argv[1].isdigit() else 1
    only = [a for a in sys.argv[1:] if not a.isdigit()]
    for name, fn in (('振り直し', furi), ('突合', totsu), ('集約', shuyaku), ('帳票', chohyo), ('月次', getsuji)):
        if only and name not in only:
            continue
        p = fn(seed)
        n = len([f for f in os.listdir(p) if f.endswith('_前.xlsx')])
        print(f"{name}: {n} 組 → {p}")
