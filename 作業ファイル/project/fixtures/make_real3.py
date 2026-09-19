# -*- coding: utf-8 -*-
"""実物らしい表・第 3 弾（2026-09-17 夜）: 似た名前の列の取り違え（見出しを「含む」で探すマクロの罠）と、按分の実物の癖。

  振り直し: 旧予算科目コードの列が左に並ぶ／対応表に旧決算統計区分の列
  突合: 新に元伝票番号の列が左に並ぶ（別の伝票の番号）／新に支出額（税抜）の列が左に並ぶ
  集約: 課のシートに前年度予算額の列が左に並ぶ
  月次: 支出日に曜日（2026/04/15(水)）／起票日の列が左に並ぶ／支出額（税抜）の列が左に並ぶ
  帳票: 連絡先が 2 段（電話・メール）
  按分: 人数が「12人」の文字／人数の内訳（正職員・会計年度任用職員・人数（計））／金額（税抜）の列が左に並ぶ／
        注記と小計の行／見出しの改行と全角の空白

  py make_real3.py [種]  → fixtures\\<型>\\未見_実物3（種 1）／未見_実物3種N
"""
import datetime
import importlib.util
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

HERE = os.path.dirname(os.path.abspath(__file__))


def _mod(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


MR = _mod('make_real', os.path.join(HERE, 'make_real.py'))
AB = _mod('mp_按分', os.path.join(HERE, '按分', 'make_pairs.py'))
FR, TG, SY, CH, GJ = MR.FR, MR.TG, MR.SY, MR.CH, MR.GJ
post, data_rows, L = MR.post, MR.data_rows, MR.L
SEED = 1
WEEK = '月火水木金土日'


def _out(kind):
    p = os.path.join(HERE, kind, '未見_実物3' if SEED == 1 else f'未見_実物3種{SEED}')
    os.makedirs(p, exist_ok=True)
    for f in os.listdir(p):
        if f.endswith('.xlsx'):
            os.remove(os.path.join(p, f))
    return p


def furi(seed):
    out = _out('振り直し')
    rng = random.Random(seed)
    s = lambda k: seed * 5000 + k      # noqa: E731

    def B(k):
        return FR.build(s(k), rng.randint(25, 45), (FR.ORDER1, FR.ORDER2)[rng.randrange(2)], rng.randint(1, 3))

    stem = '実物3_01_旧コードの列'
    t, b = B(1)
    FR.write_pair(stem, t, b, out=out)
    rows, _ = data_rows(1, len(b), ())

    def f01(wb, truth):
        ws = wb['支出明細']
        ws.insert_cols(1)
        ws.cell(1, 1, '旧予算科目コード').font = Font(bold=True)
        r2 = random.Random(stem)
        other = [c for c, _k in t]
        for r in rows:
            ws.cell(r, 1, r2.choice(other)).number_format = '@'     # 別の科目の古いコード（対応表に当たってしまう）
    post(out, stem, f01)

    stem = '実物3_02_対応表に旧区分の列'
    t, b = B(2)
    FR.write_pair(stem, t, b, out=out)

    def f02(wb, truth):
        mt = wb['対応表']
        mt.insert_cols(2)
        mt.cell(1, 2, '旧決算統計区分').font = Font(bold=True)
        r2 = random.Random(stem)
        kinds = list(dict.fromkeys(k for _c, k in t))
        for r in range(2, mt.max_row + 1):
            if mt.cell(r, 1).value:
                mt.cell(r, 2, r2.choice(kinds + ['その他']))
    post(out, stem, f02)
    return out


def totsu(seed):
    out = _out('突合')
    rng = random.Random(seed)
    s = lambda k: seed * 5000 + k      # noqa: E731

    def B(k):
        return TG.build(s(k), rng.randint(22, 40), rng.randint(1, 4), rng.randint(1, 3), rng.randint(1, 3))

    stem = '実物3_01_新に元伝票番号の列'
    old, new = B(1)
    TG.write_pair(stem, old, new, out=out)

    def f01(wb, truth):
        ns = wb['新システム']
        ns.insert_cols(1)
        ns.cell(1, 1, '元伝票番号').font = Font(bold=True)
        r2 = random.Random(stem)
        keys = [k for k, _m, _a in old]
        for r in range(2, 2 + len(new)):
            if r2.random() < 0.6:
                ns.cell(r, 1, f"{r2.choice(keys):06d}").number_format = '@'
    post(out, stem, f01)

    stem = '実物3_02_新に税抜の列'
    old, new = B(2)
    TG.write_pair(stem, old, new, out=out)

    def f02(wb, truth):
        ns = wb['新システム']
        ns.insert_cols(2)
        ns.cell(1, 2, '支出額（税抜）').font = Font(bold=True)
        for r in range(2, 2 + len(new)):
            a = ns.cell(r, 3).value
            if isinstance(a, (int, float)):
                ns.cell(r, 2, int(a * 10 / 11)).number_format = '#,##0'
    post(out, stem, f02)
    return out


def shuyaku(seed):
    out = _out('集約')
    rng = random.Random(seed)
    data = SY.build(seed * 5000 + 1, rng.randint(3, 6), rows=(3, 9))
    stem = '実物3_01_前年度予算額の列'
    for truth in (False, True):
        wb = Workbook()
        sm = wb.active
        sm.title = '集約'
        for j, v in enumerate(SY.HEAD, 1):
            sm.cell(1, j, v).font = Font(bold=True)
        if truth:
            r, tb, te = 2, 0, 0
            for ka, items in data:
                for jg, b, e in items:
                    for j, v in enumerate((ka, jg, b, e, (e / b) if b else 0), 1):
                        sm.cell(r, j, v)
                    tb, te, r = tb + b, te + e, r + 1
            for j, v in ((1, '合計'), (3, tb), (4, te), (5, (te / tb) if tb else 0)):
                sm.cell(r, j, v)
        for k, (ka, items) in enumerate(data):
            ws = wb.create_sheet(ka)
            r2 = random.Random(f"{stem}/{k}")
            for j, v in enumerate(['事業名', '前年度予算額', '予算額', '執行額'], 1):
                ws.cell(1, j, v).font = Font(bold=True)
            for i, (jg, b, e) in enumerate(items, 2):
                for j, v in enumerate((jg, int(b * r2.uniform(0.7, 1.3)) // 10000 * 10000, b, e), 1):
                    ws.cell(i, j, v)
        wb.save(os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx"))
    return out


def getsuji(seed):
    out = _out('月次')
    rng = random.Random(seed)
    Bd = lambda k, **kw: GJ.build(seed * 5000 + k, rng.randint(40, 90), **kw)  # noqa: E731

    lines = MR.plain(Bd(1))
    r2 = random.Random('week')
    for ln in lines:
        d = ln[2][0]
        if r2.random() < 0.5:
            ln[1]['支出日'] = f"{d:%Y/%m/%d}({WEEK[d.weekday()]})"
    MR.gj_write('実物3_01_曜日つきの日付', lines, out)

    lines = MR.plain(Bd(2))
    r2 = random.Random('kihyo')
    for ln in lines:
        d = ln[2][0]
        k = d - datetime.timedelta(days=r2.randint(3, 25))           # 起票は前の月になることがある
        ln[1]['起票日'] = datetime.datetime(k.year, k.month, k.day)
    MR.gj_write('実物3_02_起票日の列', lines, out, cols=('起票日', '支出日', '課名', '科目', '摘要', '支出額'))

    lines = MR.plain(Bd(3))
    for ln in lines:
        ln[1]['支出額（税抜）'] = int(ln[2][2] * 10 / 11)
    MR.gj_write('実物3_03_税抜の列', lines, out, cols=('支出日', '課名', '科目', '摘要', '支出額（税抜）', '支出額'))
    return out


def chohyo(seed):
    out = _out('帳票')
    rng = random.Random(seed)
    recs = CH.build(seed * 5000 + 1, rng.randint(5, 12), phone=True)
    for x in recs:
        r2 = random.Random(x['受付番号'])
        x['メール'] = f"info{r2.randrange(100, 999)}@example.jp" if r2.random() < 0.6 else ''
    subs = {'申請者': ('団体名', '代表者'), '連絡先': ('電話', 'メール')}
    layout = ('受付番号', '申請者', '連絡先', '申請額', '受付日')
    pos, c = {}, 1
    for name in layout:
        for j, sn in enumerate(subs.get(name, (name,))):
            pos[sn] = c + j
        c += len(subs.get(name, (name,)))
    last_col = c - 1
    fields = sorted(pos, key=lambda x: pos[x]) + ['所在地']
    stem = '実物3_01_連絡先が2段'
    for truth in (False, True):
        wb = Workbook()
        ws = wb.active
        ws.title = CH.SHEET
        ws.cell(1, 1, '令和8年度 地域づくり活動補助金 申請受付簿').font = Font(bold=True, size=14)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
        h, cc = 2, 1
        for name in layout:
            sub = subs.get(name)
            if sub:
                ws.cell(h, cc, name)
                ws.merge_cells(start_row=h, start_column=cc, end_row=h, end_column=cc + len(sub) - 1)
                for j, sn in enumerate(sub):
                    ws.cell(h + 1, cc + j, sn)
                cc += len(sub)
            else:
                ws.cell(h, cc, name)
                ws.merge_cells(start_row=h, start_column=cc, end_row=h + 1, end_column=cc)
                cc += 1
        for rr in (h, h + 1):
            for c2 in range(1, last_col + 1):
                ws.cell(rr, c2).border = MR.BOX
                ws.cell(rr, c2).font = Font(bold=True)
                ws.cell(rr, c2).alignment = Alignment(horizontal='center', vertical='center')
        r = h + 2
        for rec in recs:
            for name, col in pos.items():
                v = rec.get(name, '')
                if v != '':
                    cell = ws.cell(r, col, v)
                    if name == '受付日':
                        cell.number_format = 'yyyy/m/d'
                if name not in ('団体名', '代表者'):
                    ws.merge_cells(start_row=r, start_column=col, end_row=r + 1, end_column=col)
            ws.cell(r + 1, pos['団体名'], '所在地')
            ws.cell(r + 1, pos['代表者'], rec['所在地'])
            for rr in (r, r + 1):
                for c2 in range(1, last_col + 1):
                    ws.cell(rr, c2).border = MR.BOX
            r += 2
        lc = last_col + 2
        if truth:
            for j, f in enumerate(fields):
                ws.cell(h, lc + j, f)
            for i, rec in enumerate(recs, 1):
                for j, f in enumerate(fields):
                    if rec.get(f, '') != '':
                        ws.cell(h + i, lc + j, rec[f])
        wb.save(os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx"))
    return out


def anbun(seed):
    out = _out('按分')
    rng = random.Random(seed)
    s = lambda k: seed * 5000 + k      # noqa: E731

    stem = '実物3_01_人数が文字'
    items, staff = AB.build(s(1), rng.randint(4, 9), rng.randint(4, 8))
    AB.write_pair(stem, items, staff, out=out)

    def f01(wb, truth):
        st = wb['課別人数']
        for r in range(2, st.max_row + 1):
            v = st.cell(r, 2).value
            if isinstance(v, int) and (r % 2 == 0):
                st.cell(r, 2).value = f"{v}人"
    post(out, stem, f01)

    stem = '実物3_02_人数の内訳の列'
    items, staff = AB.build(s(2), rng.randint(4, 9), rng.randint(4, 8))
    AB.write_pair(stem, items, staff, out=out)

    def f02(wb, truth):
        st = wb['課別人数']
        st.insert_cols(2, 2)
        st.cell(1, 2, '正職員').font = Font(bold=True)
        st.cell(1, 3, '会計年度任用職員').font = Font(bold=True)
        st.cell(1, 4, '人数（計）').font = Font(bold=True)
        r2 = random.Random(stem)
        for r in range(2, st.max_row + 1):
            p = st.cell(r, 4).value
            if isinstance(p, int):
                a = r2.randint(0, p)
                st.cell(r, 2, a)
                st.cell(r, 3, p - a)
    post(out, stem, f02)

    stem = '実物3_03_税抜の列'
    items, staff = AB.build(s(3), rng.randint(4, 9), rng.randint(4, 8))
    AB.write_pair(stem, items, staff, out=out)

    def f03(wb, truth):
        ws = wb['共通経費']
        ws.insert_cols(2)
        ws.cell(1, 2, '金額（税抜）').font = Font(bold=True)
        for r in range(2, 2 + len(items)):
            a = ws.cell(r, 3).value
            if isinstance(a, (int, float)):
                ws.cell(r, 2, int(a * 10 / 11))
    post(out, stem, f03)

    # 共通経費の 4 件目の下に小計の行・表の下に※注記。按分表は見出しの行から詰めて書く決まり＝左の表だけ行を入れて書き直す
    # （insert_rows だと右の按分表の正解まで 1 行下がる）
    stem = '実物3_04_注記と小計'
    items, staff = AB.build(s(4), 8, rng.randint(4, 8))
    AB.write_pair(stem, items, staff, out=out)
    for truth in (False, True):
        p = os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx")
        from openpyxl import load_workbook
        wb = load_workbook(p)
        ws = wb['共通経費']
        vals = [[ws.cell(r, c).value for c in range(1, 3)] for r in range(1, 2 + len(items))]
        right = [[ws.cell(r, c).value for c in range(4, ws.max_column + 1)] for r in range(1, ws.max_row + 1)]
        for r in range(1, ws.max_row + 3):
            for c in range(1, ws.max_column + 1):
                ws.cell(r, c).value = None
        rows = vals[:5] + [['小計', sum(a for _h, a in items[:4])]] + vals[5:]
        for i, row in enumerate(rows, 1):
            for j, v in enumerate(row, 1):
                ws.cell(i, j, v)
        ws.cell(len(rows) + 2, 1, '※ 金額は税込み・年間見込み')
        for i, row in enumerate(right, 1):
            for j, v in enumerate(row, 4):
                if v is not None:
                    ws.cell(i, j, v)
        st = wb['課別人数']
        st.cell(st.max_row + 2, 1, '※ 4月1日現在の人数')
        wb.save(p)

    stem = '実物3_05_見出しの改行と空白'
    items, staff = AB.build(s(5), rng.randint(4, 9), rng.randint(4, 8))
    AB.write_pair(stem, items, staff, out=out)

    def f05(wb, truth):
        ws = wb['共通経費']
        ws.cell(1, 1).value = '費　目'
        ws.cell(1, 2).value = '金額\n（円）'
        st = wb['課別人数']
        st.cell(1, 1).value = '課　名'
        st.cell(1, 2).value = '人数\n（人）'
    post(out, stem, f05)
    return out


if __name__ == '__main__':
    seed = SEED = int(sys.argv[1]) if len(sys.argv) > 1 and sys.argv[1].isdigit() else 1
    only = [a for a in sys.argv[1:] if not a.isdigit()]
    for name, fn in (('振り直し', furi), ('突合', totsu), ('集約', shuyaku), ('月次', getsuji), ('帳票', chohyo), ('按分', anbun)):
        if only and name not in only:
            continue
        p = fn(seed)
        print(f"{name}: {len([f for f in os.listdir(p) if f.endswith('_前.xlsx')])} 組 → {p}")
