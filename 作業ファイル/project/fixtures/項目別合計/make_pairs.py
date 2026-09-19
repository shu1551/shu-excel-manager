# -*- coding: utf-8 -*-
"""選んでいる 2 列で「項目ごとの合計の表」を明細の右に作る（B 集計 11・2026-09-18）。

列の決まり: 選んでいる 2 列（1 つ目＝項目の列、2 つ目＝数の列）。選んでいる列は <名前>_選択.txt に書く。
正解の決まり（依頼文に全部書く）:
  - 置き場所＝明細の右に 1 列空けた所。見出しの行は明細の見出しの行と同じ行。
  - 見出しは元の 2 列の見出しをそのまま（語を作らない）。
  - 項目は明細に出てくる順（初めて出た順）。重複はまとめる。
  - 値は =SUMIF(項目の列の範囲, 項目のセル, 数の列の範囲)（本文の範囲は絶対参照）。
  - いちばん下に「合計」の行＝左に「合計」・右に =SUM（作った表の値の範囲）。
  - 同じ表が既に右にあれば作り直す（2 回撃っても結果が変わらない）。
組: 本番 60 行／A 800 行・項目 10／B 表題と単位の行／C 明細の下に合計の行／D 途中に空行／E 列の並びが違う／
  F 撃った後にもう一度／G 番号の列／H 項目が 2 つだけ
  py make_pairs.py            → *_前.xlsx と *_正解.xlsx
  py make_pairs.py --unseen N → 未見_種N
"""
import os
import random
import shutil
import sys

from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import detail   # noqa: E402

REQUEST = ('選んでいる 2 列（1 つ目＝項目の列、2 つ目＝数の列）で、項目ごとの合計の表を明細の右に 1 列空けて作ってください。'
           '見出しの行は明細の見出しの行と同じ行で、見出しは元の 2 列の見出しをそのまま使う。'
           '項目は明細に出てくる順（初めて出た順）に 1 行ずつ並べ、値は =SUMIF(項目の列の本文, 項目のセル, 数の列の本文) の式'
           '（本文の範囲は $ を付けた絶対参照）。いちばん下に合計の行を足し、左に「合計」・右に =SUM で作った値の合計を入れる。'
           '明細の本文は見出しの次の行から最後の明細の行まで（明細の下に合計の行があればそれは入れない）。'
           '同じ表が既に右にあれば作り直す。')
PHRASES = ['選んだ2列で項目ごとの合計表を右に作って', '項目別合計の表を明細の横に付けて', 'SUMIFで項目ごとに合計した表がほしい',
           '選んでいる列で項目ごとの合計表を作成', '項目ごとに合計した表を明細の右に', '項目別の合計表をSUMIFで']


def build_truth(xl, before, truth, t, lay):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(t['sheet'])
        c0, h, last = lay['c0'], lay['h'], lay['last']
        ic = c0 + detail.col_of(t, 'item') - 1
        nc = c0 + detail.col_of(t, 'num') - 1
        iname, nname = detail.name_of(t, 'item'), detail.name_of(t, 'num')
        il, nl = get_column_letter(ic), get_column_letter(nc)
        irng = f"${il}${h + 1}:${il}${last}"
        nrng = f"${nl}${h + 1}:${nl}${last}"
        out = lay['ncols'] + c0 + 1                       # 明細の右に 1 列空ける
        ws.Cells(h, out).Value = iname
        ws.Cells(h, out + 1).Value = nname
        seen = []
        for row in t['rows']:
            v = row[iname]
            if v not in seen:
                seen.append(v)
        for k, v in enumerate(seen, 1):
            ws.Cells(h + k, out).Value = v
            key = ws.Cells(h + k, out).GetAddress(False, False)
            ws.Cells(h + k, out + 1).Formula = f"=SUMIF({irng},{key},{nrng})"
            ws.Cells(h + k, out + 1).NumberFormat = '#,##0'
        at = h + len(seen) + 1
        ws.Cells(at, out).Value = '合計'
        a = ws.Cells(h + 1, out + 1).GetAddress(False, False)
        b = ws.Cells(h + len(seen), out + 1).GetAddress(False, False)
        ws.Cells(at, out + 1).Formula = f"=SUM({a}:{b})"
        ws.Cells(at, out + 1).NumberFormat = '#,##0'
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            lay = detail.write_before(before, t)
            sel = ",".join(f"{get_column_letter(detail.col_of(t, role))}{lay['h']}" for role in ('item', 'num'))
            build_truth(xl, before, truth, t, lay)
            open(os.path.join(out, f"{stem}_選択.txt"), 'w', encoding='utf-8').write(sel + '\n')
            print('  ', stem, t['sheet'], [n for n, _r in t['cols']], '選ぶ列', sel)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', detail.design(s(1), rng.randint(30, 200))),
          ('未見2_表題と合計行', detail.design(s(2), 60, title=True, unit=True, total=True)),
          ('未見3_並び違い', detail.design(s(3), 50, swap=True)),
          ('未見4_空行と番号列', detail.design(s(4), 40, blanks=2, no_col=True))], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('項目別合計_本番', detail.design(101, 60)),
          ('項目別合計_試験A', detail.design(102, 800, n_items=10)),
          ('項目別合計_試験B', detail.design(103, 50, title=True, unit=True)),
          ('項目別合計_試験C', detail.design(104, 40, total=True)),
          ('項目別合計_試験D', detail.design(105, 45, blanks=2)),
          ('項目別合計_試験E', detail.design(106, 50, swap=True)),
          ('項目別合計_試験G', detail.design(107, 35, no_col=True)),
          ('項目別合計_試験H', detail.design(108, 30, n_items=2)),
          ('項目別合計_試験I', detail.design(109, 26, no_date=True))], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '項目別合計_本番_正解.xlsx'), os.path.join(HERE, f'項目別合計_試験F_{x}.xlsx'))
    shutil.copyfile(os.path.join(HERE, '項目別合計_本番_選択.txt'), os.path.join(HERE, '項目別合計_試験F_選択.txt'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
