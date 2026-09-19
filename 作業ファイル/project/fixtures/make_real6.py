# -*- coding: utf-8 -*-
"""実物らしい表・第 6 弾（2026-09-17 深夜）: シートの並び・グラフシート・和暦の文字の受付日・システム出力の大量行・シートの保護。

  集約: 集約シートがいちばん後ろ／課のシートの間にグラフシート
  帳票: 受付日が和暦の文字（「令和8年4月3日」＝帳票のまま写す＝文字）
  大量: 突合 旧 7,000 行・新 約 20,000 行・月次 明細 30,000 行・振り直し 明細 20,000 行（1 セルずつ読むマクロの 60 秒打ち切り）
  壊れ（撃ってはいけない）: 対象のシートに保護がかかっている（振り直し・突合・集約・月次・帳票・按分・比較）

  py make_real6.py [種]  → fixtures\\<型>\\未見_実物6（種 1）／未見_実物6種N、壊れ\\*_保護
"""
import importlib.util
import os
import random
import shutil
import sys

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

HERE = os.path.dirname(os.path.abspath(__file__))


def _mod(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


MR = _mod('make_real', os.path.join(HERE, 'make_real.py'))
AB = _mod('mp_按分', os.path.join(HERE, '按分', 'make_pairs.py'))
HK = _mod('mp_比較', os.path.join(HERE, '比較', 'make_pairs.py'))
FR, TG, SY, CH, GJ = MR.FR, MR.TG, MR.SY, MR.CH, MR.GJ
post = MR.post
SEED = 1


def _out(kind):
    p = os.path.join(HERE, kind, '未見_実物6' if SEED == 1 else f'未見_実物6種{SEED}')
    os.makedirs(p, exist_ok=True)
    for f in os.listdir(p):
        if f.endswith('.xlsx'):
            os.remove(os.path.join(p, f))
    return p


def shuyaku(seed):
    out = _out('集約')
    rng = random.Random(seed)
    data = SY.build(seed * 11000 + 1, rng.randint(3, 6), rows=(3, 8))
    stem = '実物6_01_集約シートが後ろ'
    SY.write_pair(stem, data, out=out)
    post(out, stem, lambda wb, truth: wb.move_sheet('集約', offset=len(wb.sheetnames) - 1))

    data = SY.build(seed * 11000 + 2, rng.randint(3, 6), rows=(3, 8))
    stem = '実物6_02_グラフシートが混ざる'
    SY.write_pair(stem, data, out=out)

    def f(wb, truth):
        ka = data[0][0]
        ws = wb[ka]
        ch = BarChart()
        ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=1 + len(data[0][1])), titles_from_data=True)
        cs = wb.create_chartsheet('グラフ', index=2)
        cs.add_chart(ch)
    post(out, stem, f)
    return out


def chohyo(seed):
    out = _out('帳票')
    rng = random.Random(seed)
    recs = CH.build(seed * 11000 + 1, rng.randint(5, 12))
    for x in recs:
        d = x['受付日']
        x['受付日'] = f"令和{d.year - 2018}年{d.month}月{d.day}日"
    CH.write_pair('実物6_01_受付日が和暦の文字', recs, out=out)
    return out


def big(seed):
    rng = random.Random(seed)
    out = _out('突合')
    old, new = TG.build(seed * 11000 + 1, 7000, 60, 40, 13000)
    TG.write_pair('実物6_01_新が2万行', old, new, key_num=True, out=out)
    out = _out('月次')
    GJ.write_pair('実物6_01_明細3万行', GJ.build(seed * 11000 + 2, 30000, n_ka=12), out=out)
    out = _out('振り直し')
    FR.write_pair('実物6_01_明細2万行', *FR.build(seed * 11000 + 3, 20000, (FR.ORDER1, FR.ORDER2)[rng.randrange(2)], 60), out=out)


def protected():
    """壊れ\\<仕事>_保護: 各型の本番の表の対象シートに保護（正解＝直す前と同じ）。"""
    dst = os.path.join(HERE, '壊れ')
    for task, rel, sheet in (('振り直し', '振り直し/振り直し_本番_前.xlsx', '支出明細'), ('突合', '突合/突合_本番_前.xlsx', '旧システム'),
                             ('集約', '集約/集約_本番_前.xlsx', '集約'), ('月次集計', '月次/月次_本番_前.xlsx', '支出明細'),
                             ('帳票一覧', '帳票/帳票_本番_前.xlsx', '受付簿'), ('按分', '按分/按分_本番_前.xlsx', '共通経費'),
                             ('比較', '比較/比較_本番_前.xlsx', '今年度')):
        wb = load_workbook(os.path.join(HERE, rel))
        wb[sheet].protection.sheet = True
        b = os.path.join(dst, f"{task}_保護_前.xlsx")
        wb.save(b)
        shutil.copyfile(b, os.path.join(dst, f"{task}_保護_正解.xlsx"))


if __name__ == '__main__':
    seed = SEED = int(sys.argv[1]) if len(sys.argv) > 1 and sys.argv[1].isdigit() else 1
    for name, fn in (('集約', shuyaku), ('帳票', chohyo), ('大量', big)):
        fn(seed)
        print(f"{name}: 作りました")
    if seed == 1:
        protected()
        print("壊れ: シートの保護 7 枚")
