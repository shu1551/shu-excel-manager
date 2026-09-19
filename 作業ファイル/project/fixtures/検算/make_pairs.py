# -*- coding: utf-8 -*-
"""表の下の合計の行が明細の和と合っているかを、表の右に検算の表で出す（I 検査 69・2026-09-18）。

列の決まり: 選ぶ列は無い（表全体）。
正解の決まり（依頼文に全部書く）:
  - 見出しの行＝空でないセルが 2 つ以上並ぶ最初の行。合計の行＝見出しより下で「合計」「計」「総計」と書いたセルがある最後の行。
  - 検算する列＝合計の行に数が入っている列（左から順）。明細＝見出しの次の行から合計の行の 1 つ上まで。
  - 置き場所＝表の右に 1 列空けた所・見出しの行から。見出しは「検算の列」「合計の行」「明細の和」「差」「判定」。
  - 1 列につき 1 行: 列の見出し（値）／=合計の行のそのセル／=SUM(明細の範囲・$ 付き)／=合計の行−明細の和／=IF(差=0,"一致","不一致")。
    数の表示形式は #,##0。合計の行が無ければ何もしない。既に同じ表が右にあれば作り直す。
組: 本番 2 列の 1 つが違う／A 40 行で全部合う／B 表題と単位／C 前年度・今年度が両方違う／D 途中の空行／E 列の並び違い／
  F 撃った後にもう一度／G 番号の列と率の列（率は合計の行が空＝見ない）／H 合計の行が SUM の式
"""
import os
import random
import shutil
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import shukei   # noqa: E402

REQUEST = ('表の下の合計の行が、明細の和と合っているかを検算する表を、表の右に 1 列空けて作ってください。'
           '見出しの行＝空でないセルが 2 つ以上並ぶ最初の行（文字でも数でもよい）。表の右端＝見出しの行で、表のいちばん左の列から右へ見て空のセルに当たる手前の列。'
           '合計の行＝見出しの行より下で、「合計」「計」「総計」と書いたセルがある最後の行。合計の行が無ければ何もしない。'
           '検算する列＝合計の行に数が入っている列（左から順）。明細＝見出しの次の行から合計の行の 1 つ上の行まで。'
           '表の右端の 2 つ右の列から、見出しの行に「検算の列」「合計の行」「明細の和」「差」「判定」の 5 つの見出しを書き、'
           'その下に検算する列 1 本につき 1 行: 列の見出し（値）／=合計の行のそのセル（式で参照）／=SUM(明細の範囲)（$ を付けた絶対参照）／'
           '=合計の行−明細の和 の式／=IF(差=0,"一致","不一致") の式。数の表示形式は #,##0。'
           '表の値は触らない。既に同じ表が右にあれば作り直す。')
PHRASES = ['合計の行を検算して', '合計が明細の足し算と合っているか確かめて', '縦計の検算をお願い',
           '合計欄が正しいかチェックして', '合計と内訳の和が一致するか見て', '集計表の合計を検算してほしい']


def build(seed, n, wrong=(0,), formula=False, **kw):
    kw.setdefault('total', True)
    t = shukei.design(seed, n, **kw)
    t['wrong'] = wrong
    t['formula'] = formula
    return t


def spoil(before, t, lay):
    """合計の行の数を一部わざと違えておく（wrong＝何本目の数の列か）／formula なら SUM の式にする。"""
    wb = load_workbook(before)
    ws = wb.worksheets[0]
    r = lay['total_at']
    nums = [j for j, (_n, role) in enumerate(t['cols'], lay['c0']) if role == 'num']
    rng = random.Random(t['seed'] + 91)
    for k, c in enumerate(nums):
        if t['formula']:
            L = get_column_letter(c)
            ws.cell(r, c).value = f"=SUM({L}{lay['h'] + 1}:{L}{lay['last']})"
        elif k in t['wrong']:
            ws.cell(r, c).value = ws.cell(r, c).value + rng.choice((1000, -1000, 90000, 100))
    wb.save(before)
    return nums


def build_truth(xl, before, truth, t, lay, nums):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(t['sheet'])
        c0, h, last, tot = lay['c0'], lay['h'], lay['last'], lay['total_at']
        out = c0 + lay['ncols'] + 1
        for j, nm in enumerate(('検算の列', '合計の行', '明細の和', '差', '判定')):
            ws.Cells(h, out + j).Value = nm
        for k, c in enumerate(nums, 1):
            r = h + k
            L = get_column_letter(c)
            a, b, d = (get_column_letter(out + 1), get_column_letter(out + 2), get_column_letter(out + 3))
            ws.Cells(r, out).Value = ws.Cells(h, c).Value
            ws.Cells(r, out + 1).Formula = f"={L}{tot}"
            ws.Cells(r, out + 2).Formula = f"=SUM(${L}${h + 1}:${L}${tot - 1})"
            ws.Cells(r, out + 3).Formula = f"={a}{r}-{b}{r}"
            ws.Cells(r, out + 4).Formula = f'=IF({d}{r}=0,"一致","不一致")'
            for j in (1, 2, 3):
                ws.Cells(r, out + j).NumberFormat = '#,##0'
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            lay = shukei.write_before(before, t)
            nums = spoil(before, t, lay)
            build_truth(xl, before, truth, t, lay, nums)
            print('  ', stem, t['sheet'], [n for n, _r in t['cols']], '違えた', t['wrong'], '式' if t['formula'] else '')
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', build(s(1), rng.randint(6, 14), nums=2, wrong=(1,))),
          ('未見2_表題と単位', build(s(2), 9, title=True, unit=True, wrong=())),
          ('未見3_並び違い', build(s(3), 8, swap=True, prev=True, wrong=(0, 1))),
          ('未見4_空行と番号列', build(s(4), 10, blanks=2, no_col=True, wrong=(0,)))], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('検算_本番', build(2201, 9, nums=2, wrong=(1,))),
          ('検算_試験A', build(2202, 40, nums=2, wrong=())),
          ('検算_試験B', build(2203, 8, title=True, unit=True, wrong=(0,))),
          ('検算_試験C', build(2204, 10, prev=True, wrong=(0, 1))),
          ('検算_試験D', build(2205, 11, blanks=2, nums=2, wrong=(0,))),
          ('検算_試験E', build(2206, 9, swap=True, wrong=(0,))),
          ('検算_試験G', build(2207, 7, no_col=True, rate=True, nums=2, wrong=(1,))),
          ('検算_試験H', build(2208, 8, formula=True, wrong=()))], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '検算_本番_正解.xlsx'), os.path.join(HERE, f'検算_試験F_{x}.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
