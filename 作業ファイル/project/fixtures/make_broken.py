# -*- coding: utf-8 -*-
"""見出しは合っているのに、マクロの前提が欠けた表（入口の試し撃ちの確かめ用・2026-09-17 夕）。
正解＝直す前と同じ（撃ってはいけない。撃つなら AI に回す）。
  py make_broken.py → 壊れ フォルダに *_前.xlsx と *_正解.xlsx"""
import glob
import os
import shutil

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, '壊れ')
os.makedirs(OUT, exist_ok=True)


def one(src, stem, drop_sheets):
    wb = load_workbook(src)
    for name in drop_sheets:
        if name in wb.sheetnames:
            del wb[name]
    before = os.path.join(OUT, f"{stem}_前.xlsx")
    wb.save(before)
    shutil.copyfile(before, os.path.join(OUT, f"{stem}_正解.xlsx"))


def pick(folder):
    return sorted(glob.glob(os.path.join(HERE, folder, '*_前.xlsx')))[0]


one(pick('振り直し'), '振り直し_対応表なし', ['対応表'])
one(pick('突合'), '突合_新システムなし', ['新システム'])
wb = load_workbook(pick('集約'))
ka = [n for n in wb.sheetnames if n != '集約']
one(pick('集約'), '集約_課のシートなし', ka)
one(pick('按分'), '按分_課別人数なし', ['課別人数'])                  # 2026-09-17 夜


def made(folder, stem, **kw):
    """make_pairs の write_pair で直す前を作り、正解＝直す前にする（件の無い表＝何も書かない）。"""
    import importlib.util
    spec = importlib.util.spec_from_file_location(f"mp_{folder}", os.path.join(HERE, folder, 'make_pairs.py'))
    mp = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mp)
    mp.write_pair(stem, [], out=OUT, **kw)
    shutil.copyfile(os.path.join(OUT, f"{stem}_前.xlsx"), os.path.join(OUT, f"{stem}_正解.xlsx"))


made('帳票', '帳票一覧_件なし', empty_frames=6, note=True)      # 印刷用の空の枠だけ
made('月次', '月次集計_明細なし')                                 # 見出しだけの明細
print(OUT)
