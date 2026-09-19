# -*- coding: utf-8 -*-
"""集計表の値を、項目の名前と列の見出しで突き合わせて、別のシートの様式に転記する（D 帳票 36/37・2026-09-18 第二期）。

列の決まり: 選ぶ列は無い。アクティブなシート＝集計表（項目の列＋数の列）。様式＝ほかのシートで、項目の名前が 2 つ以上並ぶシート。
正解の決まり（依頼文に全部書く）:
  - 集計表の見出しの行＝空でないセルが 2 つ以上並ぶ最初の行。項目の列＝左端の文字の列。数の列＝見出しがあり本文が数の列。
  - 集計表の本文の行（空の行と「合計」「計」「総計」の行は除く）ごとに、様式で同じ項目の名前の行（前後の空白・全角空白は無視）と、
    同じ見出しの列が交わるセルに、その値を書く。式の入ったセルには書かない。様式に無い項目・見出しは飛ばす。
  - 集計表と、様式のほかのセルは変えない。
組: 本番／A 30 項目／B 表題と単位の集計表／C 集計表の合計の行・様式の合計は式／D 様式の項目に全角空白の字下げ／E 前年度・今年度の並び違い／
  F 撃った後にもう一度／G 様式に列が 1 本だけ／H 様式に無い項目と様式だけの項目
"""
import os
import random
import shutil
import sys

from openpyxl import load_workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import shukei   # noqa: E402

REQUEST = ('アクティブなシートの集計表の値を、別のシートの様式に転記してください。'
           '様式＝アクティブなシート以外で、集計表の項目の名前が 2 つ以上並んでいるシート（最初のもの）。'
           '集計表の見出しの行＝空でないセルが 2 つ以上並ぶ最初の行（文字でも数でもよい）。項目の列＝左端の文字の列。'
           '数の列＝見出しがあって本文が数の列。本文は見出しの次の行から最後の行まで（行がまるごと空の行と、行のどこかに「合計」「計」「総計」とある行は転記しない）。'
           '本文の行ごとに、様式の中で項目の名前が同じ行（前後の空白・全角の空白は無視して比べる）と、見出しが同じ列が交わるセルに、その数を書く。'
           '様式の見出しの行＝集計表の数の列の見出しが入っている行。式が入っているセル（合計の式など）には書かない。'
           '様式に無い項目・見出しは飛ばす。集計表と、様式のほかのセルは変えない。')
PHRASES = ['様式に転記して', '集計表の数字を様式に写して', '決算統計の様式に転記', '項目名で突き合わせて様式に書き込んで',
           '報告様式に数字を入れて', 'この表の値を様式シートへ転記']
FORMS = ['様式', '報告様式', '様式第1号', '決算統計様式', '提出用']


def spec(seed, n, **kw):
    kw.setdefault('nums', 2)
    form = {k: kw.pop(k) for k in ('indent', 'one_col', 'missing', 'extra', 'form_total') if k in kw}
    t = shukei.design(seed, n, **kw)
    t['form'] = form
    t['fsheet'] = random.Random(seed + 5).choice(FORMS)
    return t


def add_form(path, t, lay):
    rng = random.Random(t['seed'] + 17)
    wb = load_workbook(path)
    ws2 = wb.create_sheet(t['fsheet'])
    f = t['form']
    iname = shukei.name_of(t, 'item')
    items = [r[iname] for r in t['rows']]
    if f.get('missing'):
        items = items[1:]                                   # 集計表の先頭の項目は様式に無い
    items = items[:]
    rng.shuffle(items)
    if f.get('extra'):
        items += ['その他', '予備費']
    nums = [nm for nm, r in t['cols'] if r == 'num']
    heads = nums[::-1] if not f.get('one_col') else nums[-1:]
    ws2['A1'] = f"{t['fsheet']}　{iname}別の状況"
    ws2['A1'].font = Font(bold=True, size=14)
    ws2.cell(4, 1, 'No')
    ws2.cell(4, 2, '区分')
    for j, hd in enumerate(heads, 3):
        ws2.cell(4, j, hd)
    ws2.cell(4, 3 + len(heads), '備考')
    for i, it in enumerate(items):
        ws2.cell(5 + i, 1, i + 1)
        ws2.cell(5 + i, 2, ('　' + it + ' ') if f.get('indent') and i % 2 == 0 else it)
    end = 4 + len(items)
    if f.get('form_total', True):
        ws2.cell(end + 1, 2, '合計')
        for j in range(3, 3 + len(heads)):
            L = 'ABCDEFGH'[j - 1]
            ws2.cell(end + 1, j, f"=SUM({L}5:{L}{end})")
    wb.save(path)


def build_truth(xl, before, truth, t, lay):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(t['sheet'])
        f2 = wb.Worksheets(t['fsheet'])
        h, c0 = lay['h'], lay['c0']
        ic = c0 + shukei.col_of(t, 'item') - 1
        ncols = [(c0 + j - 1, nm) for j, (nm, r) in enumerate(t['cols'], 1) if r == 'num']
        ur = f2.UsedRange
        nr, nc = int(ur.Row) + int(ur.Rows.Count), int(ur.Column) + int(ur.Columns.Count)
        norm = lambda v: str(v or '').replace('　', ' ').strip()      # noqa: E731
        head_col = {norm(f2.Cells(4, c).Value): c for c in range(1, nc)}
        item_row = {norm(f2.Cells(r, 2).Value): r for r in range(5, nr)}
        for r in range(h + 1, lay['last'] + 1):
            it = ws.Cells(r, ic).Value
            if it is None or norm(it) in ('合計', '計', '総計'):
                continue
            rr = item_row.get(norm(it))
            if not rr:
                continue
            for c, nm in ncols:
                cc = head_col.get(nm)
                v = ws.Cells(r, c).Value
                if cc and isinstance(v, (int, float)) and not f2.Cells(rr, cc).HasFormula:
                    f2.Cells(rr, cc).Value2 = v
        ws.Activate()
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t in jobs:
            before, truth = os.path.join(out, f"{stem}_前.xlsx"), os.path.join(out, f"{stem}_正解.xlsx")
            lay = shukei.write_before(before, t)
            add_form(before, t, lay)
            build_truth(xl, before, truth, t, lay)
            print('  ', stem, t['sheet'], t['fsheet'], [n for n, _r in t['cols']], t['form'])
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', spec(s(1), rng.randint(5, 12))),
          ('未見2_表題と合計', spec(s(2), 7, title=True, unit=True, total=True, indent=True)),
          ('未見3_前年度今年度', spec(s(3), 6, prev=True, swap=True, extra=True)),
          ('未見4_空行と無い項目', spec(s(4), 8, blanks=1, missing=True, one_col=True))], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('様式_本番', spec(3401, 6, extra=True)),
          ('様式_試験A', spec(3402, 30)),
          ('様式_試験B', spec(3403, 6, title=True, unit=True)),
          ('様式_試験C', spec(3404, 7, total=True)),
          ('様式_試験D', spec(3405, 6, indent=True)),
          ('様式_試験E', spec(3406, 6, prev=True, swap=True)),
          ('様式_試験G', spec(3407, 5, one_col=True, form_total=False)),
          ('様式_試験H', spec(3408, 6, missing=True, extra=True))], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '様式_本番_正解.xlsx'), os.path.join(HERE, f'様式_試験F_{x}.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
