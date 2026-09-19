# -*- coding: utf-8 -*-
"""選んでいる数の列で「上位 5 件」を表の右に抜き出す（B 集計 21・2026-09-18）。

列の決まり: 選んでいる 1 列＝順位を付ける数の列。項目の列は表のいちばん左の文字の列。
正解の決まり（依頼文に全部書く）:
  - 置き場所＝表の右に 1 列空けた所。見出しの行は選んでいるセルの行。
  - 見出しは［項目の列の見出し］と［数の列の見出し］をそのまま。
  - 大きい順に 5 行（本文が 5 行より少なければその数）。
  - 数は =LARGE(数の列の本文, k)、項目は =INDEX(項目の列の本文, MATCH(LARGE(数の列の本文, k), 数の列の本文, 0)) の式
    （どちらも本文の範囲は $ を付けた絶対参照）。数の表示形式は元の数の列と同じ #,##0。
  - 既に同じ表が右にあれば作り直す。
組: 本番／A 40 行／B 表題と単位／C 合計の行／D 途中の空行／E 列の並び違い／F 撃った後にもう一度／G 番号の列と摘要／H 本文が 3 行だけ
"""
import os
import random
import shutil
import sys

from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import shukei   # noqa: E402

N = 5
REQUEST = ('選んでいる数の列で上位 5 件を抜き出した表を、表の右に 1 列空けて作ってください。'
           '項目の列は表のいちばん左の文字の列（数でも日付でもない列）。'
           '見出しの行は必ず「選んでいるセルの行」＝Selection.Cells(1,1).Row をそのまま使う（自分で見出しの行を探さない）。'
           '見出しは項目の列の見出しと数の列の見出しをそのまま使う。大きい順に 5 行（本文が 5 行より少なければその数だけ）。'
           '数は =LARGE(数の列の本文, k)、項目は =INDEX(項目の列の本文, MATCH(LARGE(数の列の本文, k), 数の列の本文, 0)) の式で'
           '（本文の範囲はどちらも $ を付けた絶対参照。k は 1 から順）。数の表示形式は元の数の列と同じ #,##0。'
           '本文は見出しの次の行から最後の本文の行まで。表の下に合計（合計・計）の行があればそれは入れない。'
           '既に同じ表が右にあれば作り直す。')
PHRASES = ['選んだ列で上位5件を右に抜き出して', '金額の大きい順にベスト5の表を作って', '上位5件の一覧を表の右に',
           '選んでいる列で上位5件を出して', 'トップ5を抜き出した表がほしい', '大きい順に5件だけ右に出して']


def build_truth(xl, before, truth, t, lay):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(t['sheet'])
        c0, h, last = lay['c0'], lay['h'], lay['last']
        nc = c0 + shukei.col_of(t, 'num') - 1
        ic = c0 + shukei.col_of(t, 'item') - 1
        nl, il = get_column_letter(nc), get_column_letter(ic)
        nrng = f"${nl}${h + 1}:${nl}${last}"
        irng = f"${il}${h + 1}:${il}${last}"
        out = c0 + lay['ncols'] + 1
        ws.Cells(h, out).Value = shukei.name_of(t, 'item')
        ws.Cells(h, out + 1).Value = shukei.name_of(t, 'num')
        n = min(N, len(t['rows']))
        for k in range(1, n + 1):
            ws.Cells(h + k, out).Formula = f"=INDEX({irng},MATCH(LARGE({nrng},{k}),{nrng},0))"
            ws.Cells(h + k, out + 1).Formula = f"=LARGE({nrng},{k})"
            ws.Cells(h + k, out + 1).NumberFormat = '#,##0'
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            lay = shukei.write_before(before, t)
            sel = f"{get_column_letter(shukei.col_of(t, 'num'))}{lay['h']}"
            build_truth(xl, before, truth, t, lay)
            open(os.path.join(out, f"{stem}_選択.txt"), 'w', encoding='utf-8').write(sel + '\n')
            print('  ', stem, t['sheet'], [n for n, _r in t['cols']], '選ぶ列', sel)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', shukei.design(s(1), rng.randint(6, 14))),
          ('未見2_表題と合計行', shukei.design(s(2), 9, title=True, unit=True, total=True)),
          ('未見3_並び違い', shukei.design(s(3), 8, swap=True)),
          ('未見4_空行と番号列', shukei.design(s(4), 10, blanks=2, no_col=True))], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('上位_本番', shukei.design(701, 9)),
          ('上位_試験A', shukei.design(702, 40)),
          ('上位_試験B', shukei.design(703, 8, title=True, unit=True)),
          ('上位_試験C', shukei.design(704, 10, total=True)),
          ('上位_試験D', shukei.design(705, 11, blanks=2)),
          ('上位_試験E', shukei.design(706, 9, swap=True)),
          ('上位_試験G', shukei.design(707, 7, no_col=True, text=True)),
          ('上位_試験H', shukei.design(708, 3))], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '上位_本番_正解.xlsx'), os.path.join(HERE, f'上位_試験F_{x}.xlsx'))
    shutil.copyfile(os.path.join(HERE, '上位_本番_選択.txt'), os.path.join(HERE, '上位_試験F_選択.txt'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
