# -*- coding: utf-8 -*-
"""課別・科目別の集計表を積み上げ縦棒にする（鍛える回路の題材・グラフ 6 本目・2026-09-18）。

仕事: 「課別科目別」シートの集計表（行＝課名・列＝科目・右端に合計の列・下に合計の行）から、表の右に積み上げ縦棒グラフを作る。
正解の決まり（依頼文に全部書く）:
  - 項目＝課名（合計の行は入れない）、系列＝科目（合計の列は入れない）。系列の名前は科目の見出し。
  - タイトル「課別・科目別支出」、凡例は右、数値軸 #,##0。前に作ったグラフは消して作り直す。
組: 本番 5 課×4 科目／A 12 課×7 科目／B 表題・見出しが 3 行目／C 合計の行と列が無い／D 前のグラフ／E 科目が 1 つ／F 撃った後にもう一度
"""
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
KA = ['総務課', '財政課', '企画課', '税務課', '福祉課', '建設課', '農林課', '教育総務課', '観光課', '環境課', '会計課', '議会事務局']
KAMOKU = ['需用費', '役務費', '委託料', '使用料及び賃借料', '工事請負費', '備品購入費', '旅費']
SHEET = '課別科目別'
TITLE = '課別・科目別支出'
REQUEST = ('課別科目別の集計表から、表の右に積み上げ縦棒グラフを作ってください。項目は課名（合計の行は入れない）、系列は科目（合計の列は入れない）。'
           'タイトルは「課別・科目別支出」、凡例は右、数値軸は #,##0。前に作ったグラフは消して作り直す')


def build(seed, n_ka, n_km):
    rng = random.Random(seed)
    kas, kms = rng.sample(KA, n_ka), rng.sample(KAMOKU, n_km)
    return kas, kms, [[rng.randrange(0, 900) * 1000 for _ in kms] for _ in kas]


def write_before(path, data, title=False, totals=True):
    kas, kms, vals = data
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    h = 1
    if title:
        ws.cell(1, 1, '令和8年度 課別・科目別支出（9月末）').font = Font(bold=True, size=14)
        h = 3
    ws.cell(h, 1, '課名').font = Font(bold=True)
    for j, k in enumerate(kms, 2):
        ws.cell(h, j, k).font = Font(bold=True)
    if totals:
        ws.cell(h, len(kms) + 2, '合計').font = Font(bold=True)
    for i, (ka, row) in enumerate(zip(kas, vals), 1):
        ws.cell(h + i, 1, ka)
        for j, v in enumerate(row, 2):
            ws.cell(h + i, j, v).number_format = '#,##0'
        if totals:
            c = len(kms) + 2
            ws.cell(h + i, c, f"=SUM(B{h + i}:{chr(64 + c - 1)}{h + i})").number_format = '#,##0'
    if totals:
        r = h + len(kas) + 1
        ws.cell(r, 1, '合計').font = Font(bold=True)
        for j in range(2, len(kms) + 3):
            col = chr(64 + j)
            ws.cell(r, j, f"=SUM({col}{h + 1}:{col}{h + len(kas)})").number_format = '#,##0'
    wb.save(path)
    return {'h': h, 'n_ka': len(kas), 'n_km': len(kms)}


def build_truth(xl, before, truth, lay):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Sheets(SHEET)
        for i in range(int(ws.ChartObjects().Count), 0, -1):
            ws.ChartObjects(i).Delete()
        h, n_ka, n_km = lay['h'], lay['n_ka'], lay['n_km']
        co = ws.ChartObjects().Add(ws.Cells(h, n_km + 4).Left, ws.Cells(h, 1).Top, 520, 300)
        co.Name = '表の整理_グラフ'      # マクロは「表の整理_」で始まる名前のグラフ（自分が作った物）だけを消す
        ch = co.Chart
        ch.ChartType = 52
        for i in range(int(ch.SeriesCollection().Count), 0, -1):
            ch.SeriesCollection(i).Delete()
        cats = ws.Range(ws.Cells(h + 1, 1), ws.Cells(h + n_ka, 1))
        for j in range(2, n_km + 2):
            s = ch.SeriesCollection().NewSeries()
            s.Name = f"='{SHEET}'!{ws.Cells(h, j).GetAddress(True, True, 1)}"
            s.Values = ws.Range(ws.Cells(h + 1, j), ws.Cells(h + n_ka, j))
            s.XValues = cats
        ch.HasTitle = True
        ch.ChartTitle.Text = TITLE
        ch.HasLegend = True
        ch.Legend.Position = -4152
        ch.Axes(2).TickLabels.NumberFormat = '#,##0'
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def add_old_chart(xl, path, lay):
    wb = xl.Workbooks.Open(path)
    try:
        ws = wb.Sheets(SHEET)
        co = ws.ChartObjects().Add(400, 10, 300, 200)
        co.Name = '表の整理_前のグラフ'      # マクロは「表の整理_」で始まる名前のグラフ（自分が作った物）だけを消す
        co.Chart.ChartType = 51
        s = co.Chart.SeriesCollection().NewSeries()
        s.Values = ws.Range(ws.Cells(lay['h'] + 1, 2), ws.Cells(lay['h'] + lay['n_ka'], 2))
        wb.Save()
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, data, kw in jobs:
            before, truth = os.path.join(out, f"{stem}_前.xlsx"), os.path.join(out, f"{stem}_正解.xlsx")
            old = kw.pop('old_chart', False)
            lay = write_before(before, data, **kw)
            if old:
                add_old_chart(xl, before, lay)
            build_truth(xl, before, truth, lay)
            print('  ', stem)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', build(s(1), rng.randint(3, 10), rng.randint(2, 7)), {}),
          ('未見2_表題と合計なし', build(s(2), 6, 5), {'title': True, 'totals': False}),
          ('未見3_前のグラフ', build(s(3), 8, 3), {'old_chart': True}),
          ('未見4_12課7科目', build(s(4), 12, 7), {'title': True})], out)
    return out


if __name__ == '__main__':
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        sys.exit(0)
    make([('積み上げ_本番', build(3, 5, 4), {}),
          ('積み上げ_試験A', build(13, 12, 7), {}),
          ('積み上げ_試験B', build(23, 6, 5), {'title': True}),
          ('積み上げ_試験C', build(33, 7, 4), {'totals': False}),
          ('積み上げ_試験D', build(43, 5, 3), {'old_chart': True}),
          ('積み上げ_試験E', build(53, 6, 1), {})], HERE)
    shutil.copyfile(os.path.join(HERE, '積み上げ_本番_正解.xlsx'), os.path.join(HERE, '積み上げ_試験F_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, '積み上げ_本番_正解.xlsx'), os.path.join(HERE, '積み上げ_試験F_正解.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    lines = ['課別科目別の表を積み上げ棒グラフに', '科目ごとに積み上げた課別のグラフを', '課別・科目別支出の積み上げ縦棒をお願い',
             '積み上げ縦棒で課ごとの科目の内訳を', '課別の支出を科目で積み上げて見せて', '集計表から積み上げグラフを作って']
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(lines) + '\n')
    print('ok')
