# -*- coding: utf-8 -*-
"""課別の予算執行状況を、棒と折れ線の複合グラフにする（鍛える回路の題材・グラフ 5 本目・2026-09-18）。

仕事: 「予算執行」シートの表（課名・予算額・執行額・執行率）から、表の右に複合グラフを作る。
正解の決まり（依頼文に全部書く）:
  - 予算額と執行額は集合縦棒（主軸・#,##0）、執行率はマーカー付き折れ線で第 2 軸（0%・最大 100%）。
  - タイトル「課別予算執行状況」・凡例は下。合計の行は入れない。前に作ったグラフは消して作り直す。
  - 執行率の列は式（=執行額/予算額）のことも値のこともある。
組: 本番 6 課／A 15 課／B 表題・見出しが 3 行目／C 合計の行／D 前のグラフ／E 列の並びが違う（執行率が 2 列目）／F 撃った後にもう一度
"""
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
KA = ['総務課', '財政課', '企画課', '税務課', '福祉課', '建設課', '農林課', '教育総務課', '観光課', '環境課', '会計課', '議会事務局',
      '商工課', '保健課', '子育て支援課', '市民課']
HEAD = ['課名', '予算額', '執行額', '執行率']
SWAP = ['課名', '執行率', '予算額', '執行額']
SHEET = '予算執行'
TITLE = '課別予算執行状況'
REQUEST = ('予算執行の表から、表の右に複合グラフを作ってください。予算額と執行額は集合縦棒（数値軸 #,##0）、執行率はマーカー付き折れ線で'
           '第2軸（表示形式 0%・最大 100%）。タイトルは「課別予算執行状況」、凡例は下。合計の行は入れない。前に作ったグラフは消して作り直す')


def build(seed, n):
    rng = random.Random(seed)
    out = []
    for k in rng.sample(KA, n):
        yosan = rng.randrange(100, 5000) * 1000
        out.append((k, yosan, int(yosan * rng.uniform(0.3, 1.0) / 1000) * 1000))
    return out


def write_before(path, rows, cols=HEAD, title=False, total=False, formula=True):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    h = 1
    if title:
        ws.cell(1, 1, '令和8年度 予算執行状況（9月末）').font = Font(bold=True, size=14)
        h = 3
    for j, c in enumerate(cols, 1):
        ws.cell(h, j, c).font = Font(bold=True)
    ci = {c: j for j, c in enumerate(cols, 1)}
    L = lambda c: chr(64 + ci[c])      # noqa: E731
    body = list(rows) + ([('合計', sum(r[1] for r in rows), sum(r[2] for r in rows))] if total else [])
    for i, (k, y, s) in enumerate(body, 1):
        r = h + i
        ws.cell(r, ci['課名'], k)
        ws.cell(r, ci['予算額'], y).number_format = '#,##0'
        ws.cell(r, ci['執行額'], s).number_format = '#,##0'
        ws.cell(r, ci['執行率'], f"={L('執行額')}{r}/{L('予算額')}{r}" if formula else round(s / y, 4)).number_format = '0.0%'
    wb.save(path)
    return {'h': h, 'n': len(rows), 'cols': list(cols)}


def build_truth(xl, before, truth, lay):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Sheets(SHEET)
        for i in range(int(ws.ChartObjects().Count), 0, -1):
            ws.ChartObjects(i).Delete()
        h, n, cols = lay['h'], lay['n'], lay['cols']
        col = lambda c: cols.index(c) + 1      # noqa: E731
        rng = lambda c: ws.Range(ws.Cells(h + 1, col(c)), ws.Cells(h + n, col(c)))      # noqa: E731
        co = ws.ChartObjects().Add(ws.Cells(h, len(cols) + 2).Left, ws.Cells(h, 1).Top, 520, 300)
        co.Name = '表の整理_グラフ'      # マクロは「表の整理_」で始まる名前のグラフ（自分が作った物）だけを消す
        ch = co.Chart
        ch.ChartType = 51
        for i in range(int(ch.SeriesCollection().Count), 0, -1):
            ch.SeriesCollection(i).Delete()
        for name in ('予算額', '執行額', '執行率'):
            s = ch.SeriesCollection().NewSeries()
            s.Name = f"='{SHEET}'!{ws.Cells(h, col(name)).GetAddress(True, True, 1)}"
            s.Values = rng(name)
            s.XValues = rng('課名')
        s3 = ch.SeriesCollection(3)
        s3.ChartType = 65
        s3.AxisGroup = 2
        ch.Axes(2, 1).TickLabels.NumberFormat = '#,##0'
        ax2 = ch.Axes(2, 2)
        ax2.TickLabels.NumberFormat = '0%'
        ax2.MinimumScale = 0
        ax2.MaximumScale = 1
        ch.HasTitle = True
        ch.ChartTitle.Text = TITLE
        ch.HasLegend = True
        ch.Legend.Position = -4107
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def add_old_chart(xl, path, lay):
    wb = xl.Workbooks.Open(path)
    try:
        ws = wb.Sheets(SHEET)
        co = ws.ChartObjects().Add(300, 10, 300, 200)
        co.Name = '表の整理_前のグラフ'      # マクロは「表の整理_」で始まる名前のグラフ（自分が作った物）だけを消す
        co.Chart.ChartType = 4
        s = co.Chart.SeriesCollection().NewSeries()
        s.Values = ws.Range(ws.Cells(lay['h'] + 1, 2), ws.Cells(lay['h'] + lay['n'], 2))
        wb.Save()
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, rows, kw in jobs:
            before, truth = os.path.join(out, f"{stem}_前.xlsx"), os.path.join(out, f"{stem}_正解.xlsx")
            old = kw.pop('old_chart', False)
            lay = write_before(before, rows, **kw)
            if old:
                add_old_chart(xl, before, lay)
            build_truth(xl, before, truth, lay)
            print('  ', stem)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', build(s(1), rng.randint(3, 12)), {}),
          ('未見2_表題と合計と値の執行率', build(s(2), 8), {'title': True, 'total': True, 'formula': False}),
          ('未見3_並び違いと前のグラフ', build(s(3), 7), {'cols': SWAP, 'old_chart': True}),
          ('未見4_16課', build(s(4), 16), {})], out)
    return out


if __name__ == '__main__':
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        sys.exit(0)
    make([('複合_本番', build(3, 6), {}),
          ('複合_試験A', build(13, 15), {}),
          ('複合_試験B', build(23, 7), {'title': True}),
          ('複合_試験C', build(33, 9), {'total': True}),
          ('複合_試験D', build(43, 5), {'old_chart': True}),
          ('複合_試験E', build(53, 8), {'cols': SWAP})], HERE)
    shutil.copyfile(os.path.join(HERE, '複合_本番_正解.xlsx'), os.path.join(HERE, '複合_試験F_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, '複合_本番_正解.xlsx'), os.path.join(HERE, '複合_試験F_正解.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    lines = ['予算と執行額の棒に執行率の折れ線を重ねて', '予算執行状況の複合グラフを作って', '執行率を第2軸にした課別のグラフ',
             '棒と折れ線の組み合わせで予算執行を', '予算額・執行額・執行率を 1 つのグラフに', '課別予算執行状況のグラフをお願い']
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(lines) + '\n')
    print('ok')
