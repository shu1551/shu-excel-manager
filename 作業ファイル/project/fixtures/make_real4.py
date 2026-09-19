# -*- coding: utf-8 -*-
"""実物らしい表・第 4 弾（2026-09-17 夜）: 合計の行の書き方の揺れ（「合　計」「総　計」「小　計」「計」を表の先頭の列に書く）と、
全角数字の和暦。合計の行を「合計」との完全一致で除いているマクロの取りこぼしを見る。

  振り直し: 表の下にコードの列へ「合　計」／突合: 旧の伝票番号の列に「合　計」・新の番号の列に「総　計」／
  集約: 課のシートの事業名に「小　計」「合　計」／月次: 課名の列に「合　計」・支出日が全角数字の和暦／
  按分: 費目に「合　計」・課名に「計」／帳票: 受付番号の列に「合　計」

  py make_real4.py [種]  → fixtures\\<型>\\未見_実物4（種 1）／未見_実物4種N
"""
import importlib.util
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))


def _mod(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


MR = _mod('make_real', os.path.join(HERE, 'make_real.py'))
AB = _mod('mp_按分', os.path.join(HERE, '按分', 'make_pairs.py'))
FR, TG, SY, CH, GJ = MR.FR, MR.TG, MR.SY, MR.CH, MR.GJ
post, data_rows = MR.post, MR.data_rows
SEED = 1
ZEN_D = str.maketrans('0123456789', '０１２３４５６７８９')


def _out(kind):
    p = os.path.join(HERE, kind, '未見_実物4' if SEED == 1 else f'未見_実物4種{SEED}')
    os.makedirs(p, exist_ok=True)
    for f in os.listdir(p):
        if f.endswith('.xlsx'):
            os.remove(os.path.join(p, f))
    return p


def furi(seed):
    out = _out('振り直し')
    rng = random.Random(seed)
    t, b = FR.build(seed * 7000 + 1, rng.randint(25, 45), FR.ORDER1, 2)
    stem = '実物4_01_コードの列に合計'
    FR.write_pair(stem, t, b, out=out)
    rows, end = data_rows(1, len(b), ())

    def f(wb, truth):
        ws = wb['支出明細']
        ws.cell(end, 1, '合　計').font = Font(bold=True)
        ws.cell(end, 3, f"=SUM(C2:C{end - 1})").number_format = '#,##0'
    post(out, stem, f)
    return out


def totsu(seed):
    out = _out('突合')
    rng = random.Random(seed)
    old, new = TG.build(seed * 7000 + 1, rng.randint(22, 40), 2, 2, 2)
    stem = '実物4_01_番号の列に合計'
    TG.write_pair(stem, old, new, out=out)
    rows, end = data_rows(1, len(old), ())

    def f(wb, truth):
        ws = wb['旧システム']
        ws.cell(end, 1, '合　計').font = Font(bold=True)
        ws.cell(end, 3, sum(a for _k, _m, a in old)).number_format = '#,##0'
        ns = wb['新システム']
        r = 2 + len(new)
        ns.cell(r, 1, '総　計')
        ns.cell(r, 2, sum(new.values()))
    post(out, stem, f)
    return out


def shuyaku(seed):
    out = _out('集約')
    rng = random.Random(seed)
    data = SY.build(seed * 7000 + 1, rng.randint(3, 6), rows=(4, 9))
    stem = '実物4_01_小計と合計の空白'
    for truth in (False, True):
        wb = Workbook()
        sm = wb.active
        sm.title = '集約'
        for j, v in enumerate(SY.HEAD, 1):
            sm.cell(1, j, v).font = Font(bold=True)
        if truth:
            r, tb, te = 2, 0, 0
            for ka, items in data:
                for jg, b, e in items:
                    for j, v in enumerate((ka, jg, b, e, (e / b) if b else 0), 1):
                        sm.cell(r, j, v)
                    tb, te, r = tb + b, te + e, r + 1
            for j, v in ((1, '合計'), (3, tb), (4, te), (5, (te / tb) if tb else 0)):
                sm.cell(r, j, v)
        for k, (ka, items) in enumerate(data):
            ws = wb.create_sheet(ka)
            for j, v in enumerate(['事業名', '予算額', '執行額'], 1):
                ws.cell(1, j, v).font = Font(bold=True)
            r = 2
            half = len(items) // 2
            for i, (jg, b, e) in enumerate(items):
                if i == half:
                    ws.cell(r, 1, '小　計')
                    ws.cell(r, 2, sum(x[1] for x in items[:half]))
                    ws.cell(r, 3, sum(x[2] for x in items[:half]))
                    r += 1
                for j, v in enumerate((jg, b, e), 1):
                    ws.cell(r, j, v)
                r += 1
            ws.cell(r, 1, ('合　計', '総 計', '計')[k % 3])
            ws.cell(r, 2, sum(x[1] for x in items))
            ws.cell(r, 3, sum(x[2] for x in items))
        wb.save(os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx"))
    return out


def getsuji(seed):
    out = _out('月次')
    rng = random.Random(seed)
    lines = MR.plain(GJ.build(seed * 7000 + 1, rng.randint(40, 80)))
    r2 = random.Random('zen-wareki')
    for ln in lines:
        d = ln[2][0]
        if r2.random() < 0.5:
            ln[1]['支出日'] = f"令和{d.year - 2018}年{d.month}月{d.day}日".translate(ZEN_D)
    lines.append(('x', {'課名': '合　計', '支出額': sum(ln[2][2] for ln in lines if ln[0] == 'd')}))
    MR.gj_write('実物4_01_全角の和暦と課名の列に合計', lines, out)
    return out


def anbun(seed):
    out = _out('按分')
    rng = random.Random(seed)
    items, staff = AB.build(seed * 7000 + 1, rng.randint(4, 8), rng.randint(4, 7))
    stem = '実物4_01_費目と課名に合計'
    AB.write_pair(stem, items, staff, out=out)

    def f(wb, truth):
        ws = wb['共通経費']
        r = 2 + len(items)
        ws.cell(r, 1, '合　計')
        ws.cell(r, 2, sum(a for _h, a in items))
        st = wb['課別人数']
        rr = st.max_row + 1
        st.cell(rr, 1, '計')
        st.cell(rr, 2, sum(p for _k, p in staff))
    post(out, stem, f)
    return out


def chohyo(seed):
    out = _out('帳票')
    rng = random.Random(seed)
    recs = CH.build(seed * 7000 + 1, rng.randint(5, 12))
    stem = '実物4_01_受付番号の列に合計'
    CH.write_pair(stem, recs, out=out)

    def f(wb, truth):
        ws = wb[CH.SHEET]
        r = 4 + 2 * len(recs)                                   # 表題 1・見出し 2 段・1 件 2 行
        ws.cell(r, 1, '合　計')
        ws.cell(r, 4, sum(x['申請額'] for x in recs))
    post(out, stem, f)
    return out


if __name__ == '__main__':
    seed = SEED = int(sys.argv[1]) if len(sys.argv) > 1 and sys.argv[1].isdigit() else 1
    for name, fn in (('振り直し', furi), ('突合', totsu), ('集約', shuyaku), ('月次', getsuji), ('按分', anbun), ('帳票', chohyo)):
        p = fn(seed)
        print(f"{name}: {len([f for f in os.listdir(p) if f.endswith('_前.xlsx')])} 組 → {p}")
