# -*- coding: utf-8 -*-
"""性質別の構成比の円グラフ（鍛える回路の題材・グラフ 3 本目・2026-09-17 深夜）。

仕事: 「性質別」シートの表（決算統計区分・支出額＋合計の行）から、構成比の円グラフを作る。
正解の決まり（依頼文に全部書く）:
  - 系列 1 本: 名前＝支出額の見出しのセル・値＝支出額・項目＝決算統計区分。合計の行と、支出額が 0 の区分は入れない。
  - タイトル「性質別の構成比」・凡例は右・データラベルはパーセント（値は出さない）・グラフの文字 Meiryo UI。
  - 置き場所: 表の右に 1 列空けた列の見出しの行（位置は比べない）。前に作ったグラフが残っていたら消して作り直す。
組: 本番 6 区分／A 9 区分／B 0 の区分が 2 つ／C 表題・B3 から／D 合計の行が無い／E 前のグラフ／F 撃った後にもう一度
"""
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
KUBUN = ['人件費', '物件費', '維持補修費', '扶助費', '補助費等', '普通建設事業費', '災害復旧事業費', '公債費', '積立金', '繰出金']
SHEET = '性質別'
TITLE = '性質別の構成比'
REQUEST = ('性質別シートの表から、決算統計区分ごとの支出額の構成比を円グラフにして、表の右に1列空けた列の見出しの行の高さに作ってください。'
           'タイトルは「性質別の構成比」、凡例は右、データラベルはパーセント（値は出さない）、グラフの文字は Meiryo UI。'
           '合計の行と支出額が0の区分はグラフに入れない。前に作ったグラフが残っていたら消して作り直す')


def build(seed, n, zeros=0):
    rng = random.Random(seed)
    ks = rng.sample(KUBUN, n)
    rows = [(k, rng.randrange(100, 90000) * 1000) for k in ks]
    for i in rng.sample(range(n), zeros):
        rows[i] = (rows[i][0], 0)
    return rows


def write_before(path, rows, title=False, r0=1, c0=1, total=True):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    h = r0
    if title:
        ws.cell(r0, c0, '令和7年度 決算 性質別歳出').font = Font(bold=True, size=14)
        h = r0 + 2
    ws.cell(h, c0, '決算統計区分').font = Font(bold=True)
    ws.cell(h, c0 + 1, '支出額').font = Font(bold=True)
    data = []
    for i, (k, a) in enumerate(rows, 1):
        ws.cell(h + i, c0, k)
        ws.cell(h + i, c0 + 1, a).number_format = '#,##0'
        data.append((h + i, a))
    last = h + len(rows)
    if total:
        ws.cell(last + 1, c0, '合計').font = Font(bold=True)
        ws.cell(last + 1, c0 + 1, sum(a for _k, a in rows)).number_format = '#,##0'
    wb.save(path)
    return {'h': h, 'c0': c0, 'rows': [r for r, a in data if a != 0]}


def build_truth(xl, before, truth, lay):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Sheets(SHEET)
        h, c0, rows = lay['h'], lay['c0'], lay['rows']
        for i in range(int(ws.ChartObjects().Count), 0, -1):
            ws.ChartObjects(i).Delete()
        co = ws.ChartObjects().Add(ws.Cells(h, c0 + 3).Left, ws.Cells(h, c0 + 3).Top, 420, 300)
        co.Name = '表の整理_グラフ'      # マクロは「表の整理_」で始まる名前のグラフ（自分が作った物）だけを消す
        ch = co.Chart
        ch.ChartType = 5
        for i in range(int(ch.SeriesCollection().Count), 0, -1):
            ch.SeriesCollection(i).Delete()
        s = ch.SeriesCollection().NewSeries()
        s.Name = f"='{SHEET}'!{ws.Cells(h, c0 + 1).Address}"
        s.Values = ws.Range(",".join(ws.Cells(r, c0 + 1).Address for r in rows))
        s.XValues = ws.Range(",".join(ws.Cells(r, c0).Address for r in rows))
        ch.HasTitle = True
        ch.ChartTitle.Text = TITLE
        ch.HasLegend = True
        ch.Legend.Position = -4152
        s.HasDataLabels = True
        s.DataLabels().ShowPercentage = True
        s.DataLabels().ShowValue = False
        ch.ChartArea.Format.TextFrame2.TextRange.Font.Name = 'Meiryo UI'
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def add_old_chart(xl, path, lay):
    wb = xl.Workbooks.Open(path)
    try:
        ws = wb.Sheets(SHEET)
        co = ws.ChartObjects().Add(ws.Cells(lay['h'], lay['c0'] + 3).Left, ws.Cells(lay['h'], lay['c0'] + 3).Top, 300, 200)
        co.Name = '表の整理_前のグラフ'      # マクロは「表の整理_」で始まる名前のグラフ（自分が作った物）だけを消す
        co.Chart.ChartType = 51
        s = co.Chart.SeriesCollection().NewSeries()
        s.Values = ws.Range(ws.Cells(lay['rows'][0], lay['c0'] + 1), ws.Cells(lay['rows'][-1], lay['c0'] + 1))
        wb.Save()
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, rows, kw in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            old = kw.pop('old_chart', False)
            lay = write_before(before, rows, **kw)
            if old:
                add_old_chart(xl, before, lay)
            build_truth(xl, before, truth, lay)
            print('  ', stem)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', build(s(1), rng.randint(3, 10)), {}),
          ('未見2_0と表題', build(s(2), 8, zeros=3), {'title': True, 'r0': 2, 'c0': 2}),
          ('未見3_合計なしと前のグラフ', build(s(3), 7, zeros=1), {'total': False, 'old_chart': True})], out)
    return out


if __name__ == '__main__':
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        sys.exit(0)
    make([('構成比_本番', build(3, 6), {}),
          ('構成比_試験A', build(13, 9), {}),
          ('構成比_試験B', build(23, 8, zeros=2), {}),
          ('構成比_試験C', build(33, 6), {'title': True, 'r0': 3, 'c0': 2}),
          ('構成比_試験D', build(43, 7), {'total': False}),
          ('構成比_試験E', build(53, 6), {'old_chart': True})], HERE)
    shutil.copyfile(os.path.join(HERE, '構成比_本番_正解.xlsx'), os.path.join(HERE, '構成比_試験F_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, '構成比_本番_正解.xlsx'), os.path.join(HERE, '構成比_試験F_正解.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    lines = ['性質別の構成比を円グラフにして', '決算統計区分の割合の円グラフを作って', '性質別歳出の構成比のグラフ',
             '区分ごとの構成比を円グラフで表示', '性質別の円グラフをお願いします', '支出額の構成比を円グラフに']
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(lines) + '\n')
    print('ok')
