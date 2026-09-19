# -*- coding: utf-8 -*-
"""表の中のエラー値（#DIV/0!・#N/A など）のセルを、表の右に一覧にする（I 検査 71・2026-09-18 第二期）。

列の決まり: 選ぶ列は無い（表全体）。表は集計表（shukei）に式の列（率・単価）を足し、0 で割る行・見つからない行を混ぜたもの。
正解の決まり（依頼文に全部書く）:
  - 見出しの行＝空でないセルが 2 つ以上並ぶ最初の行。表の列＝見出しの行で左端から右へ空に当たる手前まで。表の最後の行＝表の列のどれかに値がある最後の行。
  - 見出しの次の行から表の最後の行まで、上の行から・左の列から順に、エラー値のセルを探す。
  - 表の右に 1 列空けて「番地」「見出し」「項目」の一覧（番地＝A1 形式の $ なし・見出し＝その列の見出し・項目＝その行の左端の文字の列の値）。
  - エラーが無ければ見出しの行だけ書く。既に一覧が右にあれば消して作り直す。
組: 本番 率の #DIV/0! 2 つ／A 40 行／B 表題と単位／C 合計の行／D #N/A（VLOOKUP）も／E 列の並び違い／F 撃った後にもう一度／G エラー 1 つ／H エラーなし
"""
import os
import random
import shutil
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import shukei   # noqa: E402

REQUEST = ('表の中のエラー値（#DIV/0!・#N/A・#VALUE! など）のセルを、表の右に 1 列空けて一覧にしてください。'
           '見出しの行＝空でないセルが 2 つ以上並ぶ最初の行（文字でも数でもよい）。表の列＝見出しの行で、表のいちばん左の列から右へ見て空のセルに当たる手前の列まで。'
           '表の最後の行＝表の列のどれかに値がある最後の行（下の合計の行も入れる）。'
           '見出しの次の行から表の最後の行まで、上の行から順に、同じ行は左の列から順に、エラー値のセル（IsError）を探す。'
           '表の右端の 2 つ右の列から、見出しの行に「番地」「見出し」「項目」と書き、その下に 1 セル 1 行で: 番地＝A1 形式で $ なし（B5 など）・'
           '見出し＝そのセルの列の見出し・項目＝その行の、表の左端の文字の列の値。エラーが 1 つも無ければ見出しの行だけ書く。'
           '表の値は触らない。既に一覧が右にあれば消して作り直す。')
PHRASES = ['エラーのセルを一覧にして', '表の中のエラー値を探して', '#DIV/0!や#N/Aになっている所を洗い出して',
           'エラーになっている式の場所を教えて', '計算エラーのセルを右に一覧で', 'エラー値のある所を一覧にしてほしい']


def add_error_cols(path, t, lay, n_div=2, lookup=False):
    wb = load_workbook(path)
    ws = wb.worksheets[0]
    h, last, c0 = lay['h'], lay['last'], lay['c0']
    a = c0 + shukei.col_of(t, 'num', 0) - 1
    rng = random.Random(t['seed'] + 31)
    body = [r for r in range(h + 1, last + 1) if ws.cell(r, a).value is not None]
    zero = set(rng.sample(body, min(n_div, len(body))))
    kc = c0 + lay['ncols']                                  # 件数の列（0 の行で率が #DIV/0!）
    rc = kc + 1
    ws.cell(h, kc, '件数')
    ws.cell(h, rc, '1件あたり')
    A, K = get_column_letter(a), get_column_letter(kc)
    for r in body:
        ws.cell(r, kc, 0 if r in zero else rng.randint(1, 30))
        ws.cell(r, rc, f"={A}{r}/{K}{r}").number_format = '#,##0'
    if lookup:
        lc = rc + 1
        ws.cell(h, lc, '区分')
        I = get_column_letter(c0 + shukei.col_of(t, 'item') - 1)
        ws.cell(h, lc + 9, '対応')                          # 一覧（表の右 2 列目から 3 列）と重ならない遠くに小さな対応表（先頭 2 項目だけ・見出しの行から＝表題の行を見出しに見せない）
        items = [ws.cell(r, c0 + shukei.col_of(t, 'item') - 1).value for r in body[:2]]
        for i, it in enumerate(items, h + 1):
            ws.cell(i, lc + 9, it)
            ws.cell(i, lc + 10, '重点')
        L1, L2 = get_column_letter(lc + 9), get_column_letter(lc + 10)
        for r in body:
            ws.cell(r, lc, f"=VLOOKUP({I}{r},${L1}${h + 1}:${L2}${h + 2},2,FALSE)")
    wb.save(path)


def build_truth(xl, before, truth, t):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(t['sheet'])
        ur = ws.UsedRange
        h = next(r for r in range(1, 60) if sum(1 for c in range(1, 30) if ws.Cells(r, c).Value is not None) >= 2)
        left = next(c for c in range(1, 30) if ws.Cells(h, c).Value is not None)
        right = left
        while ws.Cells(h, right + 1).Value is not None:
            right += 1
        last = max(int(ws.Cells(ws.Rows.Count, c).End(-4162).Row) for c in range(left, right + 1))
        item_c = next(c for c in range(left, right + 1)
                      if isinstance(ws.Cells(h + 1, c).Value, str) or isinstance(ws.Cells(h + 2, c).Value, str))
        out = right + 2
        for j, nm in enumerate(('番地', '見出し', '項目')):
            ws.Cells(h, out + j).Value = nm
        w = h
        for r in range(h + 1, last + 1):
            for c in range(left, right + 1):
                cell = ws.Cells(r, c)
                if isinstance(cell.Value, int) and cell.Value < -2146820000 and cell.Text.startswith('#'):
                    w += 1
                    ws.Cells(w, out).Value = cell.Address.replace('$', '')
                    ws.Cells(w, out + 1).Value = ws.Cells(h, c).Value
                    ws.Cells(w, out + 2).Value = ws.Cells(r, item_c).Value
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t, n_div, lookup in jobs:
            before, truth = os.path.join(out, f"{stem}_前.xlsx"), os.path.join(out, f"{stem}_正解.xlsx")
            lay = shukei.write_before(before, t)
            add_error_cols(before, t, lay, n_div, lookup)
            build_truth(xl, before, truth, t)
            print('  ', stem, t['sheet'], [n for n, _r in t['cols']], '0 割り', n_div, 'VLOOKUP' if lookup else '')
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', shukei.design(s(1), rng.randint(6, 14)), 2, False),
          ('未見2_表題と合計', shukei.design(s(2), 9, title=True, unit=True, total=True), 1, True),
          ('未見3_並び違い', shukei.design(s(3), 8, swap=True), 3, False),
          ('未見4_エラーなし', shukei.design(s(4), 7, no_col=True), 0, False)], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('エラー_本番', shukei.design(4201, 9), 2, False),
          ('エラー_試験A', shukei.design(4202, 40), 4, False),
          ('エラー_試験B', shukei.design(4203, 8, title=True, unit=True), 2, False),
          ('エラー_試験C', shukei.design(4204, 9, total=True), 2, False),
          ('エラー_試験D', shukei.design(4205, 8), 1, True),
          ('エラー_試験E', shukei.design(4206, 8, swap=True), 2, False),
          ('エラー_試験G', shukei.design(4207, 7, no_col=True), 1, False),
          ('エラー_試験H', shukei.design(4208, 7), 0, False)], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, 'エラー_本番_正解.xlsx'), os.path.join(HERE, f'エラー_試験F_{x}.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
