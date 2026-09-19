# -*- coding: utf-8 -*-
"""共通経費を課ごとの人数で按分する（鍛える回路の題材・2026-09-17 夜。庁舎管理費・電気料・複合機代を課に割り振る仕事）。

仕事: 「共通経費」の費目ごとの金額を、「課別人数」の人数で各課に按分した表を、共通経費の表の右に 1 列空けて作る。
正解の決まり:
  - 按分表の見出しは 1 行: 費目・各課名（課別人数の表の順）・合計。左上＝共通経費の見出しの行・表の右端の 2 つ右。
  - 各課の按分額＝金額×その課の人数÷総人数 の 1 円未満切り捨て。端数（金額−按分額の合計）は人数のいちばん多い課
    （同じ人数なら表の上の課）に足す。合計の列＝その費目の金額（按分額の合計と同じ）。
  - 最後の行に「合計」（各課の按分額の合計と総合計）。
  - 共通経費: 費目が空の行・費目が「合計」「計」の行は入れない。金額はカンマ・全角・前後の空白・「円」があっても数で読む。
    見出しの語（費目・金額）で列を探す（表題の行・列の並び・ほかの列があっても同じ）。
  - 課別人数: 課名が空の行・「合計」「計」の行は入れない。人数 0 の課も列に出す（按分額は 0）。見出しの語（課名・人数）で探す。
  - 按分表の場所に前の按分表が残っていたら消してから書く（2 回撃っても同じ）。
組（依頼文は同じ）: 本番 費目 6・課 5／A 課 10・費目 12／B 人数表に合計行と空行／C 人数 0 の課／D 金額が文字（カンマ・全角・円）／
  E 共通経費に表題・B3 から／F 共通経費の下に合計行・備考の列／G 最多の人数が同じ課が 2 つ／H 前の按分表（列が多い）が残っている／
  I 撃った後にもう一度／J 人数のシート名が「人数一覧」
  py make_pairs.py            → *_前.xlsx と *_正解.xlsx
  py make_pairs.py --unseen N → 未見_種N
"""
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
HIMOKU = ['庁舎清掃委託料', '電気料', '上下水道料', 'ガス料', '複合機賃借料', 'コピー用紙代', '警備委託料', '空調保守点検',
          'エレベーター保守', '植栽管理', '廃棄物処理', '郵便料', '電話料', '新聞購読料', '消防設備点検', '駐車場賃借料',
          '庁内LAN保守', 'トイレットペーパー', '消耗品（共通）', '修繕料（共用部）']
KA = ['総務課', '財政課', '企画課', '税務課', '福祉課', '建設課', '農林課', '教育総務課', '観光課', '環境課', '会計課', '議会事務局',
      '商工課', '保健課', '子育て支援課']
_ZEN = str.maketrans('0123456789,', '０１２３４５６７８９，')


def build(seed, n_himoku, n_ka, zero_ka=0, tie=False):
    rng = random.Random(seed)
    items = [(h, rng.randrange(10, 3000) * 1000 + rng.choice((0, 0, rng.randrange(1, 999)))) for h in
             (rng.sample(HIMOKU, n_himoku) if n_himoku <= len(HIMOKU) else [f"{rng.choice(HIMOKU)}{i + 1}" for i in range(n_himoku)])]
    kas = rng.sample(KA, n_ka) if n_ka <= len(KA) else [f"{rng.choice(KA)}{i + 1}" for i in range(n_ka)]
    people = [rng.randint(3, 40) for _ in kas]
    for i in rng.sample(range(n_ka), zero_ka):
        people[i] = 0
    if tie:
        top = max(people)
        others = [i for i in range(n_ka) if people[i] != top]
        if others:
            people[others[-1]] = top                 # 下のほうに同じ最多の課
    return items, list(zip(kas, people))


def allocate(amount, staff):
    total = sum(p for _k, p in staff)
    if not total:
        return [0 for _ in staff]
    shares = [amount * p // total for _k, p in staff]
    rest = amount - sum(shares)
    top = max(p for _k, p in staff)
    i = next(i for i, (_k, p) in enumerate(staff) if p == top)
    shares[i] += rest
    return shares


def write_pair(stem, items, staff, out=HERE, title=False, r0=1, c0=1, text_amount=False, total=False, memo=False,
               staff_total=False, staff_blank=False, stale=0, cost_sheet='共通経費', staff_sheet='課別人数'):
    rng = random.Random(stem)
    for truth in (False, True):
        wb = Workbook()
        ws = wb.active
        ws.title = cost_sheet
        h = r0
        if title:
            ws.cell(r0, c0, '令和8年度 共通経費（本庁舎）').font = Font(bold=True, size=14)
            h = r0 + 2
        heads = ['費目', '金額'] + (['備考'] if memo else [])
        for j, v in enumerate(heads):
            ws.cell(h, c0 + j, v).font = Font(bold=True)
        rng2 = random.Random(stem + 'v')
        for i, (hm, amt) in enumerate(items, 1):
            ws.cell(h + i, c0, hm)
            v = amt
            if text_amount and rng2.random() < 0.6:
                x = rng2.random()
                v = f"{amt:,}".translate(_ZEN) if x < 0.3 else (f" {amt:,} " if x < 0.6 else f"{amt:,}円")
            ws.cell(h + i, c0 + 1, v).number_format = '#,##0'
            if memo and rng2.random() < 0.3:
                ws.cell(h + i, c0 + 2, rng2.choice(['年間契約', '前年度比増', '単価改定']))
        last = h + len(items)
        if total:
            ws.cell(last + 1, c0, '合計').font = Font(bold=True)
            ws.cell(last + 1, c0 + 1, sum(a for _h, a in items)).number_format = '#,##0'
        lc = c0 + len(heads) + 1                                # 按分表＝表の右端の 2 つ右
        names = [k for k, _p in staff]
        if truth:
            for j, v in enumerate(['費目'] + names + ['合計']):
                ws.cell(h, lc + j, v)
            sums = [0] * len(staff)
            for i, (hm, amt) in enumerate(items, 1):
                shares = allocate(amt, staff)
                ws.cell(h + i, lc, hm)
                for j, s in enumerate(shares, 1):
                    ws.cell(h + i, lc + j, s)
                    sums[j - 1] += s
                ws.cell(h + i, lc + len(staff) + 1, amt)
            rr = h + len(items) + 1
            ws.cell(rr, lc, '合計')
            for j, s in enumerate(sums, 1):
                ws.cell(rr, lc + j, s)
            ws.cell(rr, lc + len(staff) + 1, sum(a for _h, a in items))
        elif stale:
            for j, v in enumerate(['費目'] + [f'旧課{k + 1}' for k in range(stale)] + ['合計']):
                ws.cell(h, lc + j, v)
            for i in range(1, len(items) + 4):
                for j in range(stale + 2):
                    ws.cell(h + i, lc + j, f'旧{i}' if j == 0 else 1000 * i)
        st = wb.create_sheet(staff_sheet)
        st.cell(1, 1, '課名').font = Font(bold=True)
        st.cell(1, 2, '人数').font = Font(bold=True)
        r = 2
        for i, (k, p) in enumerate(staff):
            if staff_blank and i == len(staff) // 2:
                r += 1
            st.cell(r, 1, k)
            st.cell(r, 2, p)
            r += 1
        if staff_total:
            st.cell(r, 1, '合計').font = Font(bold=True)
            st.cell(r, 2, sum(p for _k, p in staff))
        wb.save(os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx"))


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    write_pair('未見1_標準', *build(s(1), rng.randint(3, 12), rng.randint(3, 9)), out=out)
    write_pair('未見2_人数表の合計と空行', *build(s(2), 6, 7), staff_total=True, staff_blank=True, out=out)
    write_pair('未見3_人数0と同数', *build(s(3), 8, 6, zero_ka=2, tie=True), out=out)
    write_pair('未見4_文字の金額', *build(s(4), 9, 5), text_amount=True, out=out)
    write_pair('未見5_表題と位置と合計行', *build(s(5), 7, 6), title=True, r0=2, c0=2, total=True, memo=True, out=out)
    write_pair('未見6_前の按分表', *build(s(6), 5, 4), stale=9, out=out)
    write_pair('未見7_全部入り', *build(s(7), 10, 8, zero_ka=1, tie=True), title=True, r0=3, c0=2, text_amount=True, total=True,
               memo=True, staff_total=True, staff_blank=True, stale=11, out=out)
    write_pair('未見8_シート名', *build(s(8), 6, 5), cost_sheet='共通経費（R8）', staff_sheet='人数表', out=out)
    write_pair('未見9_大量', *build(s(9), 200, 30, zero_ka=3, tie=True), out=out)
    return out


if __name__ == '__main__':
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        sys.exit(0)
    write_pair('按分_本番', *build(3, 6, 5))
    write_pair('按分_試験A', *build(13, 12, 10))
    write_pair('按分_試験B', *build(23, 7, 6), staff_total=True, staff_blank=True)
    write_pair('按分_試験C', *build(33, 6, 6, zero_ka=2))
    write_pair('按分_試験D', *build(43, 8, 5), text_amount=True)
    write_pair('按分_試験E', *build(53, 5, 4), title=True, r0=3, c0=2)
    write_pair('按分_試験F', *build(63, 7, 5), total=True, memo=True)
    write_pair('按分_試験G', *build(73, 6, 6, tie=True))
    write_pair('按分_試験H', *build(83, 5, 4), stale=8)
    write_pair('按分_試験J', *build(93, 6, 5), staff_sheet='人数一覧')
    shutil.copyfile(os.path.join(HERE, '按分_本番_正解.xlsx'), os.path.join(HERE, '按分_試験I_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, '按分_本番_正解.xlsx'), os.path.join(HERE, '按分_試験I_正解.xlsx'))
    print('ok')
