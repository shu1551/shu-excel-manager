# -*- coding: utf-8 -*-
"""選んでいる日付の列から「年度」「月」「四半期」の 3 列を表の右端に足す（B 集計 20・2026-09-18）。

列の決まり: 選んでいる 1 列＝日付の列。選んでいる列は <名前>_選択.txt に書く。
正解の決まり（依頼文に全部書く）:
  - 表のいちばん右の列のすぐ右に 3 列。見出しは「年度」「月」「四半期」。
  - 年度＝=YEAR(日付)-IF(MONTH(日付)<4,1,0)（4 月始まり）。月＝=MONTH(日付)。四半期＝=INT(MOD(MONTH(日付)-4,12)/3)+1。
  - 見出しの行は選んでいるセルの行。本文＝その次の行から最後の本文の行まで。合計の行と空行には書かない。
  - 既に「年度」「月」「四半期」の列があれば足さずに書き直す。
組: 本番 60 行／A 800 行／B 表題と単位／C 明細の下に合計の行／D 空行／E 列の並び違い／F 撃った後にもう一度／G 番号の列／H 一部の月だけ
"""
import os
import random
import shutil
import sys

from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import detail   # noqa: E402

REQUEST = ('選んでいる日付の列から、年度・月・四半期の 3 列を表のいちばん右の列のすぐ右に足してください。'
           '見出しは「年度」「月」「四半期」。年度は =YEAR(日付のセル)-IF(MONTH(日付のセル)<4,1,0) の式（4 月始まりの年度）、'
           '月は =MONTH(日付のセル)、四半期は =INT(MOD(MONTH(日付のセル)-4,12)/3)+1 の式（4〜6 月が 1）。'
           '見出しの行は必ず「選んでいるセルの行」＝Selection.Cells(1,1).Row をそのまま使う（自分で見出しの行を探さない）。'
           '本文は見出しの次の行から最後の本文の行まで。表の下に合計（合計・計）の行があればそこには書かない。'
           '途中の空行にも書かない。3 列とも表示形式は標準（G/標準）。'
           '既に「年度」「月」「四半期」の列があれば、足さずにそこを書き直す。')
PHRASES = ['選んだ日付の列から年度と月と四半期の列を足して', '日付から年度・月・四半期の列を作って', '年度と四半期の列を右に付けて',
           '選んでいる日付で年度月四半期の列を', '日付を年度と月と四半期に分けた列がほしい', '四半期の列を日付から足して']


def build_truth(xl, before, truth, t, lay):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(t['sheet'])
        c0, h, last = lay['c0'], lay['h'], lay['last']
        dc = c0 + detail.col_of(t, 'date') - 1
        dl = get_column_letter(dc)
        out = c0 + lay['ncols']
        for k, name in enumerate(('年度', '月', '四半期')):
            ws.Cells(h, out + k).Value = name
        for r in range(h + 1, last + 1):
            if ws.Cells(r, dc).Value is None:
                continue
            ws.Cells(r, out).Formula = f"=YEAR({dl}{r})-IF(MONTH({dl}{r})<4,1,0)"
            ws.Cells(r, out + 1).Formula = f"=MONTH({dl}{r})"
            ws.Cells(r, out + 2).Formula = f"=INT(MOD(MONTH({dl}{r})-4,12)/3)+1"
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            lay = detail.write_before(before, t)
            sel = f"{get_column_letter(detail.col_of(t, 'date'))}{lay['h']}"
            build_truth(xl, before, truth, t, lay)
            open(os.path.join(out, f"{stem}_選択.txt"), 'w', encoding='utf-8').write(sel + '\n')
            print('  ', stem, t['sheet'], [n for n, _r in t['cols']], '選ぶ列', sel)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', detail.design(s(1), rng.randint(30, 200))),
          ('未見2_表題と合計行', detail.design(s(2), 60, title=True, unit=True, total=True)),
          ('未見3_並び違い', detail.design(s(3), 50, swap=True)),
          ('未見4_空行と番号列', detail.design(s(4), 40, blanks=2, no_col=True))], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('年度月_本番', detail.design(601, 60)),
          ('年度月_試験A', detail.design(602, 800, n_items=10)),
          ('年度月_試験B', detail.design(603, 50, title=True, unit=True)),
          ('年度月_試験C', detail.design(604, 40, total=True)),
          ('年度月_試験D', detail.design(605, 45, blanks=2)),
          ('年度月_試験E', detail.design(606, 50, swap=True)),
          ('年度月_試験G', detail.design(607, 35, no_col=True)),
          ('年度月_試験H', detail.design(608, 30, gap_months=(5, 6, 7, 8, 9, 10, 11, 1, 2, 3)))], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '年度月_本番_正解.xlsx'), os.path.join(HERE, f'年度月_試験F_{x}.xlsx'))
    shutil.copyfile(os.path.join(HERE, '年度月_本番_選択.txt'), os.path.join(HERE, '年度月_試験F_選択.txt'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
