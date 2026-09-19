# -*- coding: utf-8 -*-
"""依頼の言い方のお題（2026-09-17 夜）: 鍛えるときに見せていない言い方で、入口がどの登録マクロを撃つかを確かめる（AI も Excel も使わない）。
配った .xlam の登録簿（表の整理の頭の行）を読み、vbam_prefire.plan_full に依頼文と、各型の本番の表の見出しの語を渡す（入口と同じ絞り込み）。

  py check_phrases.py [--xlam パス]

- 当てたい: その仕事の言い方 → そのマクロ「だけ」を撃つ
- 当ててはいけない: ほかの仕事・似た語の依頼 → 鍛えたマクロを撃たない（表を整える・重複・合計行は道具の既定なので見ない）
"""
import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'python_scripts')))
import vbam_core
vbam_core.setup_encoding()
import vbam_lineage as vl
import vbam_prefire as vp

# 仕事 → その Sub の名前は台帳から引く（2026-09-18: 作り直しで Sub の名前が変わるので、手で書かない）
import vbam_forge as _vf   # noqa: E402
FORGED = {k: ((v.get('sub') or None) if v.get('passed') else None) for k, v in _vf._forge_load().items() if not v.get('retired')}

HIT1 = {   # 2026-09-17 夜に各型の phrases.txt へ足した＝鍛えた言い方（答え合わせには HIT2 を使う）
    '振り直し': ['決算統計の区分ごとに支出額を出して', 'この明細を決算統計用に分類して集計', '予算科目を決算統計区分に置き換えて区分別の表を',
               '決算統計の性質別に振り分けてほしい', '対応表を見て区分を付けて、区分ごとの合計も'],
    '突合': ['旧システムと新システムの伝票を照らし合わせて', '移行した伝票が合っているか確認して', '新旧システムの支出額を比べて一致しているか見て',
           '伝票番号で新システムと突き合わせ、金額違いを出して', 'データ移行の検証をしてください'],
    '集約': ['各課から来た回答シートを1枚にまとめて', '課ごとのシートを集計シートに集めて', '照会の回答を取りまとめてください',
           '全課の事業を集約シートに一覧で', '各課シートの事業を集めて合計も出して'],
    '帳票一覧': ['この受付簿をデータの形にして', '結合セルの帳票を1行1件のリストに', '申請受付簿を集計できる一覧表に変換',
             '様式の受付簿から一覧を作って', '帳票形式を表形式に直して'],
    '月次集計': ['支出明細を月別・課別にまとめて', '課ごとの月ごとの支出額の表を作って', '4月から3月までの月別の支出を課ごとに',
             '月次の執行状況の表を課別で', '明細から課別月別の集計表'],
    '按分': ['共通経費を各課に人数按分して', '庁舎の光熱水費を職員数で割り振って', '共通経費の課別負担額を人数比で出して',
           '人数に応じて共通経費を課に配分', '按分計算をお願いします。端数は人数の多い課に'],
}
HIT2 = {
    '振り直し': ['支出を決算統計の区分で集計しなおして', '決算統計の区分別の支出額の表がほしい', '予算科目コードから決算統計区分を引いて集計',
               '決算統計向けに科目を区分へ振り替えて'],
    '突合': ['新システムの出力と旧システムの伝票を突合して', '旧と新で伝票の金額が一致するかチェックして', '新システムに無い伝票を洗い出して',
           '移行後のデータと突き合わせ'],
    '集約': ['各課シートの内容を集約シートに転記して', '課別の回答を集約して合計行も', '全課分の事業を一本化して集約',
           '照会回答を集約シートへ'],
    '帳票一覧': ['受付簿の帳票を横一列の一覧に', '2段見出しの受付簿を一覧形式に直して', '帳票になっている申請を一覧化',
             '受付簿から1件1行のリストを作成'],
    '月次集計': ['課別月別の支出集計表をお願い', '支出明細を月次集計して', '月別集計表を課ごとに作成して',
             '年度の月別集計を課単位で'],
    '按分': ['共通経費を課別人数で按分', '共通経費の按分表を作成してください', '人数割りで共通経費を各課へ', '共通経費を人数比で案分して'],
    '比較': ['前年度との比較表を作成', '今年度と前年度の増減を科目別に出して', '予算の前年度比較をお願い', '対前年度の増減率も出して'],
    '見やすい棒グラフ': ['課別の支出を棒グラフで見せて', '課ごとの支出額を縦棒グラフに', '課別支出のグラフを作成'],
    '見やすい推移グラフ': ['月別支出の推移を課ごとに折れ線で', '課別の月次推移グラフ', '月ごとの支出の推移を折れ線グラフに'],
    '構成比円': ['性質別の割合を円グラフに', '決算統計区分の構成比のグラフ', '歳出の性質別構成比を円で'],
    '課別科目別ピボット': ['課別科目別にピボットで集計', '行と列でピボットの集計表を作って', '課と科目のクロス集計をピボットで'],
    # 「支出明細からピボットテーブルを作成して」はピボットの仕事が 3 本ある今は AI へが正しい（9/18 夜・期待を今の姿に）
    '職員名簿テーブル': ['職員名簿をテーブルに変換', '名簿をテーブルにして集計行を出して', '職員名簿のテーブル化をお願い'],
    'PQ課別集計': ['パワークエリで月の明細をまとめて課別集計', 'T支出のテーブルをパワークエリで追加して集計', 'パワークエリで課別の合計を読み込んで'],
}
HIT = HIT2
MISS = ['この表を整えて', '重複している行を消して', '伝票番号で並べ替えて', '課ごとの人数を数えて', '支出額の合計を下に足して',
        '予算額と執行額のグラフを作って', '月の列を足して',    # 「受付日の書き方をそろえて」は日付直しそのもの（9/18 外した）
        '新しいシートを作って', '課名の空白を消して', 'この表をまとめて見やすく', '前年度の資料を開いて',
        '共通の書式にそろえて',   # 「人数の列を数値にして」は数値直しそのもの（9/18 外した）
        '回答が届いた課に印を付けて', '月ごとの件数を数えて', '伝票の一覧を作って', '課の一覧を別のシートに', 'この明細を課ごとに色分け',
        '受付日の順に並べて']   # 「決算書の様式に貼り付けて」は様式転記そのもの（9/18 第二期で外した）


# 1 つの依頼で 2 つの仕事: 鍛えたマクロを撃ち、残りの仕事（other）は AI に回す＝(仕事, 依頼, AI に回す仕事が残るか)
COMBO = [
    ('振り直し+構成比円', '決算統計区分に振り直して、区分ごとの円グラフも作って', False),   # 円グラフも登録済み＝2 本とも撃つ（9/18 夜）
    ('振り直し', '決算統計の区分で集計しなおして、支出額の大きい順に並べ替えて', True),
    ('突合', '新システムと突合して、金額違いの行に色を付けて', True),
    ('突合', '新システムと突合したら、新に無しの伝票だけ抽出して', True),
    ('集約', '各課シートを集約シートにまとめて、執行率のグラフも', True),
    ('帳票一覧', '受付簿を一覧に直して、受付日の順に並べ替えて', True),
    ('月次集計+見やすい推移グラフ', '月次集計表を作って、課ごとの折れ線グラフも', False),   # 2026-09-18: 月次集計の後に推移の折れ線も撃つ（鍛えた 2 本で片づく）
    ('按分+印刷設定', '共通経費を課別人数で按分して、印刷の設定もして', False),   # 9/18 第二期: 印刷設定ができた＝2 本とも撃つ
    ('比較', '前年度比較表を作って、増減率の大きい順に並べ替え', True),
    ('振り直し', '決算統計区分に振り直して集計してください', False),
    ('課別科目別ピボット', '行と列でピボットを作って、ピボットグラフも', True),   # 「課別科目別」は表の語＝依頼の語にしない（9/18 夜）
    ('課名を引く', '課マスタから課名を引いて、支出額の大きい順に並べ替えて', True),
    ('PQ縦持ち', '月別支出を縦持ちにして、課名で色分けして', True),
    ('見やすい棒グラフ+印刷設定', '課別支出を棒グラフにして、印刷の設定もして', False),
    ('月次集計', '月次集計表を作ってください', False),
]


# 表は台帳の弾（_agent_forge の「撃つ前」の写し）から引く（2026-09-18: 作り直しでお題のフォルダが変わるので手で書かない）
BOOKS = {k: (v.get('before') or '') for k, v in _vf._forge_load().items() if v.get('before') and not v.get('retired')}
SHEETS = {k: v.get('sheet') for k, v in _vf._forge_load().items() if v.get('before') and not v.get('retired')}


def sheet_words_of(path):
    from openpyxl import load_workbook
    full = path if os.path.isabs(path) else os.path.join(os.path.dirname(os.path.abspath(__file__)), path)
    ws = load_workbook(full, read_only=True).worksheets[0]
    return vp.sheet_words([list(r) for r in ws.iter_rows(min_row=1, max_row=15, values_only=True)])


def main():
    args = sys.argv[1:]
    xlam = os.path.join(os.environ.get('APPDATA', ''), 'Microsoft', 'AddIns', '秀コンボ.xlam')
    if '--xlam' in args:
        xlam = args[args.index('--xlam') + 1]
    text = vl._read_closed_book(xlam).get('表の整理') or ''
    reg = vp.registry_from_text(text, owner=os.path.basename(xlam))
    import vbam_forge as vf
    ledger = vf._forge_load()
    for task in list(FORGED):                              # Sub の名前は鍛え直しで変わる＝台帳から読む
        FORGED[task] = (ledger.get(task) or {}).get('sub') or FORGED[task] or f'（{task} は未登録）'
    forged = set(FORGED.values())
    words = {task: sheet_words_of(path) for task, path in BOOKS.items()}
    # 表の形（' 形:）でも絞る＝入口と同じ（2026-09-18）
    shapes = {task: vf.shape_of_file(path, SHEETS.get(task)) for task, path in BOOKS.items()}
    bad = 0
    for task, reqs in [(t, r) for t, r in HIT.items() if t in words and FORGED.get(t)]:
        for req in reqs:
            got = [e['name'] for e in vp.plan_full(req, reg, words[task], shapes[task])['extras'] if e['name'] in forged]
            ok = got == [FORGED[task]]
            bad += not ok
            if not ok:
                print(f"  当たらない: {task}「{req}」→ {got or '撃たない（AI へ）'}")
    for req in MISS:
        for task, w in words.items():                  # どの仕事の表で頼まれても、鍛えたマクロは撃たない
            got = [e['name'] for e in vp.plan_full(req, reg, w, shapes[task])['extras'] if e['name'] in forged]
            if got:
                bad += 1
                print(f"  誤爆: {task}の表で「{req}」→ {got}")
    for task, req, want_other in [c for c in COMBO if all(x in words and FORGED.get(x) for x in c[0].split('+'))]:
        tasks = task.split('+')                            # 「月次集計+推移折れ線」＝表は先頭の仕事の表・撃つのは並べた仕事ぜんぶ
        p = vp.plan_full(req, reg, words[tasks[0]], shapes[tasks[0]])
        got = [e['name'] for e in p['extras'] if e['name'] in forged]
        want = [FORGED[x] for x in tasks]
        if sorted(got) != sorted(want) or bool(p['other']) != want_other:
            bad += 1
            print(f"  2 つの仕事: {task}「{req}」→ 撃つ {got or 'なし'}・AI に回す仕事 {p['left'] or 'なし'}"
                  f"（期待: {'・'.join(want)} を撃ち、AI に回す仕事 {'あり' if want_other else 'なし'}）")
    n = sum(len(v) for v in HIT.values()) + len(MISS) * len(words) + len(COMBO)
    print(f"言い方のお題 {n} 本: 合格 {n - bad}・外れ {bad}")


if __name__ == '__main__':
    main()
