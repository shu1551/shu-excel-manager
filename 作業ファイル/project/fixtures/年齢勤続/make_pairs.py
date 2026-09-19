# -*- coding: utf-8 -*-
"""名簿の生年月日・採用日から、基準日の年齢・勤続年数の列を表の右端に足す（J 名簿 77・2026-09-18）。

列の決まり: 選んでいる 2 列（1 つ目＝生年月日の列、2 つ目＝採用日の列）。<名前>_選択.txt に書く。
正解の決まり（依頼文に全部書く）:
  - 表の右端の次に「年齢」、その次に「勤続年数」。既に見出しの行にあれば足さずにそこを書き直す。
  - 基準日＝見出しより上で「基準日」で始まる文字のセルのすぐ右の日付（$ 付きで指す）。無ければ TODAY()。
  - =DATEDIF(生年月日, 基準日, "Y")／=DATEDIF(採用日, 基準日, "Y")。元が空なら空。行がまるごと空の行には書かない。
組: 本番／A 400 行／B 表題の下に基準日／C 基準日が 3 月 31 日／D 途中の空行と空の生年月日・採用日／E 採用日が左（並び違い）／
  F 撃った後にもう一度／G 基準日が表題の右／H 「基準日：」と書いたラベル
"""
import datetime
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
from vocab import Vocab   # noqa: E402

REQUEST = ('選んでいる 2 列（1 つ目＝生年月日の列、2 つ目＝採用日の列）から、年齢と勤続年数の列を表の右端に足してください。'
           '見出しの行は必ず「選んでいるセルの行」＝Selection.Cells(1,1).Row をそのまま使う（自分で見出しの行を探さない）。'
           '表の右端＝見出しの行で、表のいちばん左の列から右へ見て空のセルに当たる手前の列。その次の列の見出しに「年齢」、その次に「勤続年数」'
           '（見出しの行に「年齢」「勤続年数」の列が既にあれば、足さずにそこを書き直す）。'
           '基準日＝見出しの行より上で、「基準日」で始まる文字のセルのすぐ右のセルの日付（式では $ を付けた絶対参照で指す）。'
           '基準日のセルが無ければ TODAY()。値は =DATEDIF(生年月日のセル, 基準日, "Y") と =DATEDIF(採用日のセル, 基準日, "Y") の式。'
           '生年月日・採用日が空の行はその列を空のままにする。行がまるごと空の行には書かない。本文は見出しの次の行から最後の本文の行まで。'
           'ほかの列は触らない。')
PHRASES = ['年齢と勤続年数の列を足して', '生年月日から年齢を出して勤続年数も', '基準日時点の年齢と勤続を計算して',
           '職員の年齢・勤続年数を入れて', '名簿に年齢と勤続の列を', '採用日から勤続年数を出して年齢も']

SEI = ['佐藤', '鈴木', '高橋', '田中', '伊藤', '渡辺', '小松', '工藤', '佐々木', '斎藤', '三浦', '菅原', '加藤', '藤田']
MEI = ['一郎', '花子', '健', '美咲', '大輔', '由美', '翔', '恵', '誠', '陽子', '亮', '真理', '拓也', '直美']
BIRTH = ['生年月日', '誕生日']
HIRE = ['採用日', '採用年月日', '入庁日']
NO = ['職員番号', '番号', '職員No']


def build(seed, n, base=(2026, 4, 1), where='top', blanks=0, holes=False, swap=False):
    rng = random.Random(seed)
    v = Vocab(seed)
    base_d = datetime.datetime(*base)
    cols = [(rng.choice(NO), 'no'), (v.word('name'), 'name'), (v.word('ka'), 'ka'),
            (rng.choice(BIRTH), 'birth'), (rng.choice(HIRE), 'hire'), (v.word('kihon'), 'kihon')]
    if swap:
        cols = [cols[0], cols[1], cols[4], cols[2], cols[3], cols[5]]
    kas = v.items(5, rng)
    rows = []
    for i in range(n):
        b = datetime.datetime(rng.randint(1964, 2003), rng.randint(1, 12), rng.randint(1, 28))
        hy = min(base_d.year - 1, b.year + rng.randint(22, 35))
        h = datetime.datetime(hy, 4, 1)
        rows.append({'no': 100 + i, 'name': rng.choice(SEI) + ' ' + rng.choice(MEI), 'ka': rng.choice(kas),
                     'birth': b, 'hire': h, 'kihon': rng.randrange(180, 450) * 1000})
    if holes:
        rows[rng.randrange(n)]['birth'] = None
        rows[rng.randrange(n)]['hire'] = None
    blank_rows = sorted(random.Random(seed + 13).sample(range(1, n - 1), blanks)) if blanks else []
    return {'sheet': v.sheet('meibo'), 'cols': cols, 'rows': rows, 'base': base_d, 'where': where,
            'blank_rows': blank_rows, 'seed': seed}


def write_before(path, t):
    wb = Workbook()
    ws = wb.active
    ws.title = t['sheet']
    w = t['where']
    if w == 'top':
        ws.cell(1, 1, '基準日'); ws.cell(1, 2, t['base']).number_format = 'yyyy/m/d'; h = 3; bref = '$B$1'
    elif w == 'colon':
        ws.cell(1, 1, '基準日：'); ws.cell(1, 2, t['base']).number_format = 'yyyy/m/d'; h = 3; bref = '$B$1'
    elif w == 'under_title':
        ws.cell(1, 1, '職員名簿（令和8年度）').font = Font(bold=True, size=14)
        ws.cell(2, 1, '基準日'); ws.cell(2, 2, t['base']).number_format = 'yyyy/m/d'; h = 4; bref = '$B$2'
    else:  # right_of_title
        ws.cell(1, 1, '職員一覧').font = Font(bold=True, size=14)
        ws.cell(1, 4, '基準日'); ws.cell(1, 5, t['base']).number_format = 'yyyy/m/d'; h = 3; bref = '$E$1'
    for j, (nm, _r) in enumerate(t['cols'], 1):
        ws.cell(h, j, nm).font = Font(bold=True)
    r = h
    for i, row in enumerate(t['rows']):
        if i in t['blank_rows']:
            r += 1
        r += 1
        for j, (_nm, role) in enumerate(t['cols'], 1):
            c = ws.cell(r, j, row[role])
            if role in ('birth', 'hire'):
                c.number_format = 'yyyy/m/d'
            elif role == 'kihon':
                c.number_format = '#,##0'
    wb.save(path)
    return {'h': h, 'last': r, 'ncols': len(t['cols']), 'bref': bref}


def build_truth(xl, before, truth, t, lay):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(t['sheet'])
        h, last, nc = lay['h'], lay['last'], lay['ncols']
        bc = next(j for j, (_n, r) in enumerate(t['cols'], 1) if r == 'birth')
        hc = next(j for j, (_n, r) in enumerate(t['cols'], 1) if r == 'hire')
        ws.Cells(h, nc + 1).Value = '年齢'
        ws.Cells(h, nc + 2).Value = '勤続年数'
        for r in range(h + 1, last + 1):
            if all(ws.Cells(r, c).Value is None for c in range(1, nc + 1)):
                continue
            for k, src in enumerate((bc, hc), 1):
                if ws.Cells(r, src).Value is not None:
                    ws.Cells(r, nc + k).Formula = f'=DATEDIF({get_column_letter(src)}{r},{lay["bref"]},"Y")'
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            lay = write_before(before, t)
            build_truth(xl, before, truth, t, lay)
            bc = next(j for j, (_n, r) in enumerate(t['cols'], 1) if r == 'birth')
            hc = next(j for j, (_n, r) in enumerate(t['cols'], 1) if r == 'hire')
            sel = f"{get_column_letter(bc)}{lay['h']},{get_column_letter(hc)}{lay['h']}"
            open(os.path.join(out, f"{stem}_選択.txt"), 'w', encoding='utf-8').write(sel + '\n')
            print('  ', stem, t['sheet'], [n for n, _r in t['cols']], '選ぶ列', sel, '基準日', lay['bref'])
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', build(s(1), rng.randint(20, 60))),
          ('未見2_表題の下', build(s(2), 30, base=(2027, 3, 31), where='under_title')),
          ('未見3_並び違い', build(s(3), 25, swap=True, where='right_of_title')),
          ('未見4_空行と穴', build(s(4), 28, blanks=2, holes=True, where='colon'))], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('年齢_本番', build(2501, 30)),
          ('年齢_試験A', build(2502, 400)),
          ('年齢_試験B', build(2503, 25, where='under_title')),
          ('年齢_試験C', build(2504, 24, base=(2026, 3, 31))),
          ('年齢_試験D', build(2505, 26, blanks=2, holes=True)),
          ('年齢_試験E', build(2506, 28, swap=True)),
          ('年齢_試験G', build(2507, 22, where='right_of_title', base=(2025, 10, 1))),
          ('年齢_試験H', build(2508, 20, where='colon'))], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '年齢_本番_正解.xlsx'), os.path.join(HERE, f'年齢_試験F_{x}.xlsx'))
    shutil.copyfile(os.path.join(HERE, '年齢_本番_選択.txt'), os.path.join(HERE, '年齢_試験F_選択.txt'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
