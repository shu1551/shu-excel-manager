# -*- coding: utf-8 -*-
"""選んでいる数の列の右に「順位」の列を足す（B 集計 18・2026-09-18）。

列の決まり: 選んでいる 1 列＝順位を付ける数の列。選んでいる列は <名前>_選択.txt に書く。
正解の決まり（依頼文に全部書く）:
  - 選んだ数の列のすぐ右に 1 列挿入する（右の列を押し出す）。見出しは「順位」。
  - 値は =RANK.EQ(その行のセル, 本文の範囲, 0)＝大きい順・範囲は $ を付けた絶対参照。
  - 本文＝見出しの次の行から最後の本文の行まで。合計（合計・計）の行には書かない。途中の空行にも書かない。
  - すぐ右が既に「順位」の列なら、挿入せずにそこを書き直す。
組: 本番／A 40 行／B 表題と単位／C 合計の行／D 途中の空行／E 列の並び違い／F 撃った後にもう一度／G 番号の列と摘要／H 数の列が 2 本（右の列を選ぶ）
"""
import os
import random
import shutil
import sys

from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import shukei   # noqa: E402

REQUEST = ('選んでいる数の列のすぐ右に 1 列挿入して、順位の列にしてください（右の列は押し出す）。'
           '見出しは「順位」。値は =RANK.EQ(その行のセル, 本文の範囲, 0) の式（大きい順）で、範囲は $ を付けた絶対参照。'
           '見出しの行は必ず「選んでいるセルの行」＝Selection.Cells(1,1).Row をそのまま使う（自分で見出しの行を探さない。表題や「単位：円」の行が上にあり、列が 2 つだけの表もあるので、値の数では見分けられない）。本文は見出しの次の行から最後の本文の行まで。'
           '表の下に合計（合計・計）の行があればそこには書かない。途中の空行にも書かない。'
           'すぐ右が既に「順位」の列なら、挿入せずにそこを書き直す。')
PHRASES = ['選んだ列の順位の列を足して', '金額の大きい順に順位を付けて', '順位の列を右に入れて', '選んでいる列でランキングを出して',
           '何位かを列で出して', '額の多い順に番付の列を']


def build_truth(xl, before, truth, t, lay, nth):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(t['sheet'])
        c0, h, last = lay['c0'], lay['h'], lay['last']
        col = c0 + shukei.col_of(t, 'num', nth) - 1
        ws.Columns(col + 1).Insert()
        ws.Cells(h, col + 1).Value = '順位'
        src = get_column_letter(col)
        rng = f"${src}${h + 1}:${src}${last}"
        for r in range(h + 1, last + 1):
            if ws.Cells(r, col).Value is None:
                continue
            ws.Cells(r, col + 1).Formula = f"=RANK.EQ({src}{r},{rng},0)"
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t, nth in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            lay = shukei.write_before(before, t)
            sel = f"{get_column_letter(shukei.col_of(t, 'num', nth))}{lay['h']}"
            build_truth(xl, before, truth, t, lay, nth)
            open(os.path.join(out, f"{stem}_選択.txt"), 'w', encoding='utf-8').write(sel + '\n')
            print('  ', stem, t['sheet'], [n for n, _r in t['cols']], '選ぶ列', sel)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', shukei.design(s(1), rng.randint(5, 12)), 0),
          ('未見2_表題と合計行', shukei.design(s(2), 8, title=True, unit=True, total=True), 0),
          ('未見3_並び違い', shukei.design(s(3), 7, swap=True, nums=2), 1),
          ('未見4_空行と番号列', shukei.design(s(4), 9, blanks=2, no_col=True), 0)], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('順位_本番', shukei.design(1801, 8), 0),
          ('順位_試験A', shukei.design(1802, 40), 0),
          ('順位_試験B', shukei.design(1803, 7, title=True, unit=True), 0),
          ('順位_試験C', shukei.design(1804, 9, total=True), 0),
          ('順位_試験D', shukei.design(1805, 10, blanks=2), 0),
          ('順位_試験E', shukei.design(1806, 8, swap=True), 0),
          ('順位_試験G', shukei.design(1807, 6, no_col=True, text=True), 0),
          ('順位_試験H', shukei.design(1808, 7, nums=2), 1)], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '順位_本番_正解.xlsx'), os.path.join(HERE, f'順位_試験F_{x}.xlsx'))
    shutil.copyfile(os.path.join(HERE, '順位_本番_選択.txt'), os.path.join(HERE, '順位_試験F_選択.txt'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
