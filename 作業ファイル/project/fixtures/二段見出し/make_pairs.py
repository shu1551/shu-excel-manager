# -*- coding: utf-8 -*-
"""2 段の見出しを 1 段に畳む（A 掃除 5・2026-09-18）。

列の決まり: 選ぶ列は無い（表全体）。
正解の決まり（依頼文に全部書く）:
  - 見出しの行＝空でないセルが 2 つ以上並ぶ最初の行。その行（上段）に空のセルか結合があり、次の行（下段）に数も日付も無いとき、
    その 2 行を 2 段の見出しとみなす。そうでなければ何もしない。
  - 列の名前＝上段の語＋下段の語（間に何も入れない）。上段が空の列は左の列の上段の語を引き継ぐ（結合していた範囲と同じ）。
    下段が空の列は上段の語だけ。上段と下段が同じ語なら 1 つだけ。
  - 見出しの 2 行の結合を解いて、下段の行に列の名前を書き、上段の行を行ごと削除する。表題（見出しより上）と本文は触らない。
作り方: 正解は 1 段の見出しの表を openpyxl でそのまま書く（結合なし）。
組: 本番／A 800 行／B 表題の結合つき／C 結合なし（上段は組の先頭の列だけ）／D 3 組と番号の列／E 合計の行／
  F 撃った後にもう一度／G 左の列が 3 本／H 2 段の見出しが 1 組だけ
"""
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
from vocab import Vocab   # noqa: E402

REQUEST = ('アクティブなシートの表の 2 段になっている見出しを、1 段の見出しに畳んでください。'
           '見出しの行＝空でないセルが 2 つ以上並ぶ最初の行（上段）。上段に空のセルか結合セルがあり、次の行（下段）に数も日付も無いとき、'
           'その 2 行を 2 段の見出しとみなす。そうでなければ何もしない（2 回撃っても結果が変わらない）。'
           '表の列＝上段か下段のどちらかに語がある列（左端〜右端）。'
           '列の名前＝上段の語＋下段の語（間に何も入れない）。上段が空の列は左の列の上段の語を引き継ぐ（結合していた範囲と同じ）。'
           '下段が空の列は上段の語だけ。上段と下段が同じ語なら 1 つだけ。'
           '見出しの 2 行の結合を解き、下段の行に列の名前を書いてから、上段の行を行ごと削除する。'
           '表題（見出しより上の行）と本文の値は触らない。')
PHRASES = ['2段の見出しを1段にまとめて', '二段になっている項目名を一行にして', '見出しの上下2行をつなげて1行に',
           '上段と下段の見出しを結合して1行の見出しに', '階層になった見出しを平らにして', '見出しを1行に畳んで']

YEARS = [('前年度', '今年度'), ('令和7年度', '令和8年度'), ('R7', 'R8'), ('当初', '補正後'), ('上半期', '下半期')]
LOWS = [('予算額', '執行額'), ('件数', '金額'), ('予算', '決算'), ('計画', '実績')]


def build(seed, n, kind='std'):
    rng = random.Random(seed)
    v = Vocab(seed)
    left = [v.word('ka'), v.word('kamoku')]
    if kind == 'left3':
        left.append(v.word('tekiyo'))
    if kind == 'no':
        left = [v.word('no')] + left
    ups = list(rng.choice(YEARS))
    if kind == 'no':
        ups.append('増減')
    if kind == 'one':
        ups = ups[:1]
    lows = list(rng.choice(LOWS))
    t = {'sheet': v.sheet('yosan'), 'seed': seed, 'kind': kind, 'left': left, 'ups': ups, 'lows': lows,
         'title': f"令和8年度 {left[0]}別の状況" if kind == 'title' else None,
         'total': kind == 'total', 'merged': kind != 'nomerge'}
    items = v.items(max(2, n // 3 + 1), rng)
    kams = v.kamokus(5, rng)
    t['rows'] = []
    for i in range(n):
        row = []
        for j, _nm in enumerate(left):
            if kind == 'no' and j == 0:
                row.append(1000 + i)
            elif j == (1 if kind == 'no' else 0):
                row.append(items[i % len(items)])
            elif j == (2 if kind == 'no' else 1):
                row.append(rng.choice(kams))
            else:
                row.append(rng.choice(['消耗品', '委託', '修繕', '備品']))
        row += [rng.randrange(1, 900) * 1000 for _ in range(len(ups) * len(lows))]
        t['rows'].append(row)
    return t


def names(t):
    return list(t['left']) + [u + lo for u in t['ups'] for lo in t['lows']]


def write(path, t, flat):
    wb = Workbook()
    ws = wb.active
    ws.title = t['sheet']
    r = 1
    width = len(t['left']) + len(t['ups']) * len(t['lows'])
    if t['title']:
        ws.cell(1, 1, t['title']).font = Font(bold=True, size=14)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
        r = 3
    if flat:
        for j, nm in enumerate(names(t), 1):
            ws.cell(r, j, nm).font = Font(bold=True)
        r += 1
    else:
        for j, nm in enumerate(t['left'], 1):
            ws.cell(r, j, nm).font = Font(bold=True)
            if t['merged']:
                ws.merge_cells(start_row=r, start_column=j, end_row=r + 1, end_column=j)
        c = len(t['left']) + 1
        for u in t['ups']:
            ws.cell(r, c, u).font = Font(bold=True)
            if t['merged']:
                ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + len(t['lows']) - 1)
            for k, lo in enumerate(t['lows']):
                ws.cell(r + 1, c + k, lo).font = Font(bold=True)
            c += len(t['lows'])
        r += 2
    first = r
    for row in t['rows']:
        for j, x in enumerate(row, 1):
            cell = ws.cell(r, j, x)
            if isinstance(x, int) and x >= 1000 and j > len(t['left']):
                cell.number_format = '#,##0'
        r += 1
    if t['total']:
        ws.cell(r, 1, '合計').font = Font(bold=True)
        for j in range(len(t['left']) + 1, width + 1):
            ws.cell(r, j, sum(row[j - 1] for row in t['rows'])).number_format = '#,##0'
    wb.save(path)
    return first


def make(jobs, out):
    for stem, t in jobs:
        write(os.path.join(out, f"{stem}_前.xlsx"), t, flat=False)
        write(os.path.join(out, f"{stem}_正解.xlsx"), t, flat=True)
        print('  ', stem, t['sheet'], t['kind'], names(t))


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', build(s(1), rng.randint(12, 40))),
          ('未見2_表題', build(s(2), 20, 'title')),
          ('未見3_結合なし', build(s(3), 16, 'nomerge')),
          ('未見4_番号と3組', build(s(4), 18, 'no'))], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('二段_本番', build(1701, 20)),
          ('二段_試験A', build(1702, 800)),
          ('二段_試験B', build(1703, 18, 'title')),
          ('二段_試験C', build(1704, 20, 'nomerge')),
          ('二段_試験D', build(1705, 16, 'no')),
          ('二段_試験E', build(1706, 18, 'total')),
          ('二段_試験G', build(1707, 18, 'left3')),
          ('二段_試験H', build(1708, 15, 'one'))], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '二段_本番_正解.xlsx'), os.path.join(HERE, f'二段_試験F_{x}.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
