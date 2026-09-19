# -*- coding: utf-8 -*-
"""選んでいる 2 列から「増減額」「増減率」の列を表の右端に足す（B 集計 17・2026-09-18）。

列の決まり: 選んでいる 2 列（1 つ目＝前の数の列、2 つ目＝後の数の列）。選んでいる列は <名前>_選択.txt に書く。
正解の決まり（依頼文に全部書く）:
  - 表のいちばん右の列のすぐ右に 2 列足す。見出しは「増減額」「増減率」。
  - 増減額＝(2 つ目のセル)−(1 つ目のセル) の式・表示形式 #,##0。増減率＝増減額の式÷(1 つ目のセル)・表示形式 0.0%。
  - 本文＝見出しの次の行から最後の本文の行まで。合計（合計・計）の行と途中の空行には書かない。
  - 既に「増減額」「増減率」の列があれば、足さずにそこを書き直す。
組: 本番／A 40 行／B 表題と単位／C 合計の行／D 途中の空行／E 列の並び違い／F 撃った後にもう一度／G 番号の列と摘要／H 率の列が右端
"""
import os
import random
import shutil
import sys

from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import shukei   # noqa: E402

REQUEST = ('選んでいる 2 列（1 つ目＝前の数の列、2 つ目＝後の数の列）から、増減額の列と増減率の列を表のいちばん右の列の'
           'すぐ右に 2 列足してください。見出しは「増減額」「増減率」。'
           '増減額＝(2 つ目のセル)-(1 つ目のセル) の式で表示形式 #,##0、増減率＝((2 つ目のセル)-(1 つ目のセル))/(1 つ目のセル) の式で'
           '表示形式 0.0%。見出しの行は必ず「選んでいるセルの行」＝Selection.Cells(1,1).Row をそのまま使う（自分で見出しの行を探さない。表題や「単位：円」の行が上にあり、列が 2 つだけの表もあるので、値の数では見分けられない）。本文は見出しの次の行から最後の本文の行まで。'
           '表の下に合計（合計・計）の行があればそこには書かない。途中の空行にも書かない。'
           '既に「増減額」「増減率」の列があれば、足さずにそこを書き直す。')
PHRASES = ['選んだ2列で増減の列を足して', '増減額の列と増減率の列を右に足して', '2つの列の差と伸び率の列を付けて',
           '選んでいる2列から増減の列を作って', '前年と今年の増減の列を足して', '差額と伸び率の列を右に']


def build_truth(xl, before, truth, t, lay):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(t['sheet'])
        c0, h, last = lay['c0'], lay['h'], lay['last']
        a = c0 + shukei.col_of(t, 'num', 0) - 1
        b = c0 + shukei.col_of(t, 'num', 1) - 1
        out = c0 + lay['ncols']
        ws.Cells(h, out).Value = '増減額'
        ws.Cells(h, out + 1).Value = '増減率'
        al, bl = get_column_letter(a), get_column_letter(b)
        for r in range(h + 1, last + 1):
            if ws.Cells(r, a).Value is None:
                continue
            ws.Cells(r, out).Formula = f"={bl}{r}-{al}{r}"
            ws.Cells(r, out).NumberFormat = '#,##0'
            ws.Cells(r, out + 1).Formula = f"=({bl}{r}-{al}{r})/{al}{r}"
            ws.Cells(r, out + 1).NumberFormat = '0.0%'
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            lay = shukei.write_before(before, t)
            sel = ",".join(f"{get_column_letter(shukei.col_of(t, 'num', k))}{lay['h']}" for k in (0, 1))
            build_truth(xl, before, truth, t, lay)
            open(os.path.join(out, f"{stem}_選択.txt"), 'w', encoding='utf-8').write(sel + '\n')
            print('  ', stem, t['sheet'], [n for n, _r in t['cols']], '選ぶ列', sel)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', shukei.design(s(1), rng.randint(5, 12), prev=True)),
          ('未見2_表題と合計行', shukei.design(s(2), 8, prev=True, title=True, unit=True, total=True)),
          ('未見3_並び違い', shukei.design(s(3), 7, prev=True, swap=True)),
          ('未見4_空行と番号列', shukei.design(s(4), 9, prev=True, blanks=2, no_col=True))], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('増減_本番', shukei.design(501, 8, prev=True)),
          ('増減_試験A', shukei.design(502, 40, prev=True)),
          ('増減_試験B', shukei.design(503, 7, prev=True, title=True, unit=True)),
          ('増減_試験C', shukei.design(504, 9, prev=True, total=True)),
          ('増減_試験D', shukei.design(505, 10, prev=True, blanks=2)),
          ('増減_試験E', shukei.design(506, 8, prev=True, swap=True)),
          ('増減_試験G', shukei.design(507, 6, prev=True, no_col=True, text=True)),
          ('増減_試験H', shukei.design(508, 7, prev=True, rate=True))], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '増減_本番_正解.xlsx'), os.path.join(HERE, f'増減_試験F_{x}.xlsx'))
    shutil.copyfile(os.path.join(HERE, '増減_本番_選択.txt'), os.path.join(HERE, '増減_試験F_選択.txt'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
