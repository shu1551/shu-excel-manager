# -*- coding: utf-8 -*-
"""選んでいる 2 列（項目・数）で、平均・最大・最小の表を表の右に作る（B 集計 14・2026-09-18）。

列の決まり: 選んでいる 2 列（1 つ目＝項目の列、2 つ目＝数の列）。選んでいる列は <名前>_選択.txt に書く。
正解の決まり（依頼文に全部書く）:
  - 置き場所＝表の右に 1 列空けた所（表の右端＝見出しの行で、いちばん左の列から右へ空のセルに当たる手前まで）。見出しの行は選んでいるセルの行。
  - 3 列: 見出しは「指標」［数の列の見出し］［項目の列の見出し］。下に 平均・最大・最小 の 3 行。
  - 数は =AVERAGE／=MAX／=MIN(数の列の本文)。項目は最大・最小の行だけ =INDEX(項目の列の本文, MATCH(MAX(…), 数の列の本文, 0))（平均の行は空）。
    本文の範囲は $ を付けた絶対参照。数の表示形式は #,##0。
  - 本文＝見出しの次の行から最後の本文の行まで。合計（合計・計）の行は入れない。既に同じ表が右にあれば作り直す。
組: 本番／A 40 行／B 表題と単位／C 合計の行／D 途中の空行／E 列の並び違い／F 撃った後にもう一度／G 番号の列と摘要／H 数の列が 2 本（右の列）
"""
import os
import random
import shutil
import sys

from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import shukei   # noqa: E402

REQUEST = ('選んでいる 2 列（1 つ目＝項目の列、2 つ目＝数の列）で、平均・最大・最小の表を、表の右に 1 列空けて作ってください。'
           '見出しの行は必ず「選んでいるセルの行」＝Selection.Cells(1,1).Row をそのまま使う（自分で見出しの行を探さない）。'
           '表の右端＝見出しの行で、表のいちばん左の列から右へ見て空のセルに当たる手前の列。その 2 つ右の列から 3 列の表を書く。'
           '見出しは「指標」・数の列の見出し・項目の列の見出しをそのまま。その下に「平均」「最大」「最小」の 3 行。'
           '数は =AVERAGE(数の列の本文)・=MAX(数の列の本文)・=MIN(数の列の本文) の式。'
           '項目は最大と最小の行だけ =INDEX(項目の列の本文, MATCH(MAX(数の列の本文), 数の列の本文, 0))（最小は MIN）の式で、平均の行の項目は空。'
           '本文の範囲は $ を付けた絶対参照。数の表示形式は #,##0。'
           '本文は見出しの次の行から最後の本文の行まで。表の下に合計（合計・計）の行があればそれは入れない。'
           '既に同じ表が右にあれば作り直す。')
PHRASES = ['選んだ列の平均と最大と最小を出して', '平均・最大・最小の表を右に作って', '最大値と最小値と平均値を出して',
           'いちばん多い所と少ない所と平均を表にして', '選んでいる列の最大最小平均を出して', '金額の平均と最高・最低を出して']


def build_truth(xl, before, truth, t, lay, nth):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(t['sheet'])
        c0, h, last = lay['c0'], lay['h'], lay['last']
        nc = c0 + shukei.col_of(t, 'num', nth) - 1
        ic = c0 + shukei.col_of(t, 'item') - 1
        nl, il = get_column_letter(nc), get_column_letter(ic)
        nrng = f"${nl}${h + 1}:${nl}${last}"
        irng = f"${il}${h + 1}:${il}${last}"
        out = c0 + lay['ncols'] + 1
        ws.Cells(h, out).Value = '指標'
        ws.Cells(h, out + 1).Value = shukei.name_of(t, 'num', nth)
        ws.Cells(h, out + 2).Value = shukei.name_of(t, 'item')
        for k, (lab, fn) in enumerate((('平均', 'AVERAGE'), ('最大', 'MAX'), ('最小', 'MIN')), 1):
            ws.Cells(h + k, out).Value = lab
            ws.Cells(h + k, out + 1).Formula = f"={fn}({nrng})"
            ws.Cells(h + k, out + 1).NumberFormat = '#,##0'
            if fn != 'AVERAGE':
                ws.Cells(h + k, out + 2).Formula = f"=INDEX({irng},MATCH({fn}({nrng}),{nrng},0))"
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t, nth in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            lay = shukei.write_before(before, t)
            sel = (f"{get_column_letter(shukei.col_of(t, 'item'))}{lay['h']},"
                   f"{get_column_letter(shukei.col_of(t, 'num', nth))}{lay['h']}")
            build_truth(xl, before, truth, t, lay, nth)
            open(os.path.join(out, f"{stem}_選択.txt"), 'w', encoding='utf-8').write(sel + '\n')
            print('  ', stem, t['sheet'], [n for n, _r in t['cols']], '選ぶ列', sel)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', shukei.design(s(1), rng.randint(5, 12)), 0),
          ('未見2_表題と合計行', shukei.design(s(2), 8, title=True, unit=True, total=True), 0),
          ('未見3_並び違い', shukei.design(s(3), 7, swap=True, nums=2), 1),
          ('未見4_空行と番号列', shukei.design(s(4), 9, blanks=2, no_col=True), 0)], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('平均_本番', shukei.design(1901, 8), 0),
          ('平均_試験A', shukei.design(1902, 40), 0),
          ('平均_試験B', shukei.design(1903, 7, title=True, unit=True), 0),
          ('平均_試験C', shukei.design(1904, 9, total=True), 0),
          ('平均_試験D', shukei.design(1905, 10, blanks=2), 0),
          ('平均_試験E', shukei.design(1906, 8, swap=True), 0),
          ('平均_試験G', shukei.design(1907, 6, no_col=True, text=True), 0),
          ('平均_試験H', shukei.design(1908, 7, nums=2), 1)], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '平均_本番_正解.xlsx'), os.path.join(HERE, f'平均_試験F_{x}.xlsx'))
    shutil.copyfile(os.path.join(HERE, '平均_本番_選択.txt'), os.path.join(HERE, '平均_試験F_選択.txt'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
