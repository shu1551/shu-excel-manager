# -*- coding: utf-8 -*-
"""金額を別のシートの基準（人数など）で按分する（どの表でも動く形・2026-09-18 作り直し。前の版は按分）。

形で決める: アクティブなシート＝金額の表（項目の列と数の列）。基準の表＝別のシートの表（項目の列と数の列）。
見出しの語・シート名・項目の名前・列の並びは表ごとに違う（vocab.py）。
正解の決まり:
  - 按分表の見出しは 1 行: 金額の表の項目の列の見出し・基準の表の項目（その表の順）・合計。
    左上＝金額の表の見出しの行・表の右端の 2 つ右（1 列空ける）。
  - 各項目の按分額＝金額×その項目の基準÷基準の合計 の 1 円未満切り捨て。端数（金額−按分額の合計）は
    基準がいちばん大きい項目（同じなら表の上）に足す。合計の列＝その行の金額。
  - 最後の行に「合計」（各列の按分額の合計と総合計）。
  - 両方の表で、項目が空の行と 合計・計 の行は入れない。基準が 0 の項目も列に出す（按分額は 0）。
  - 金額はカンマ・全角・前後の空白・「円」つきの文字でも数で読む。前の按分表が残っていたら消してから書く。
組: 本番 項目 6・基準 5／A 基準 10・項目 12／B 基準の表に合計行と空行／C 基準 0 の項目／D 金額が文字／
  E 表題と位置／F 合計行と備考の列／G 基準の最大が 2 つ／H 前の按分表（列が多い）が残る／I 撃った後にもう一度
"""
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
from vocab import Vocab   # noqa: E402

_ZEN = str.maketrans('0123456789,', '０１２３４５６７８９，')
HIMOKU = ['庁舎清掃委託料', '電気料', '上下水道料', 'ガス料', '複合機賃借料', 'コピー用紙代', '警備委託料', '空調保守点検',
          'エレベーター保守', '植栽管理', '廃棄物処理', '郵便料', '電話料', '新聞購読料', '消防設備点検', '駐車場賃借料']
COST_SHEETS = ['共通経費', '共通費', '庁舎経費', '共用経費', '経費']
BASE_SHEETS = ['課別人数', '人数一覧', '職員数', '配分基準', '基準']
COST_HEADS = ['費目', '経費項目', '項目', '内容']
REQUEST = ('アクティブなシートの表の金額を、別のシートの表（項目の列と数の列＝按分の基準）の数で按分した表を、'
           'アクティブなシートの表の右に 1 列空けて作ってください。'
           '見出しは 1 行で アクティブなシートの項目の列の見出し・基準の表の項目（その表の順）・合計。'
           '各項目の按分額＝切り捨て（その行の金額 × その項目の基準 ÷ 基準の合計）。'
           '基準の合計は「基準の表の数の列の合計」（金額の表の数は使わない・行ごとに変わらない 1 つの数）。'
           '端数（金額−按分額の合計）は基準がいちばん大きい項目（同じなら表の上）に足す。合計の列はその行の金額。'
           '最後の行に「合計」（各列の按分額の合計と総合計）。'
           '両方の表で、項目が空の行と 合計・計 の行は入れない。基準が 0 の項目も列に出す（按分額は 0）。'
           '金額はカンマ・全角・前後の空白・「円」つきの文字でも数で読む。前の按分表が残っていたら消してから書く')
PHRASES = ['共通経費を人数で按分して', '別のシートの人数で割り振って', '基準の数で按分表を作って',
           '人数比で配分した表を', '按分して課別の表にして', '人数割りで按分額を出して']


def build(seed, n_himoku, n_base, zero=0, tie=False):
    rng = random.Random(seed)
    v = Vocab(seed, item_kind='課')
    names = (rng.sample(HIMOKU, n_himoku) if n_himoku <= len(HIMOKU)
             else [f"{rng.choice(HIMOKU)}{i + 1}" for i in range(n_himoku)])
    items = [(h, rng.randrange(10, 3000) * 1000 + rng.choice((0, 0, rng.randrange(1, 999)))) for h in names]
    kas = v.items(n_base, rng)
    people = [rng.randint(3, 40) for _ in kas]
    for i in rng.sample(range(n_base), zero):
        people[i] = 0
    if tie:
        top = max(people)
        others = [i for i in range(n_base) if people[i] != top]
        if others:
            people[others[-1]] = top
    t = {'items': items, 'base': list(zip(kas, people)), 'cost_sheet': rng.choice(COST_SHEETS),
         'base_sheet': rng.choice(BASE_SHEETS), 'cost_head': rng.choice(COST_HEADS), 'amount': v.word('amount'),
         'base_item': v.word('ka'), 'base_num': v.word('ninzu'), 'memo_head': v.word('tekiyo'),
         'title': False, 'r0': 1, 'c0': 1, 'text_amount': False, 'total': False, 'memo': False,
         'base_total': False, 'base_blank': False, 'stale': 0}
    return t


def allocate(amount, base):
    total = sum(p for _k, p in base)
    if not total:
        return [0 for _ in base]
    shares = [amount * p // total for _k, p in base]
    rest = amount - sum(shares)
    top = max(p for _k, p in base)
    i = next(i for i, (_k, p) in enumerate(base) if p == top)
    shares[i] += rest
    return shares


def write_pair(stem, t, out=HERE):
    items, base = t['items'], t['base']
    for truth in (False, True):
        wb = Workbook()
        ws = wb.active
        ws.title = t['cost_sheet']
        h, c0 = t['r0'], t['c0']
        if t['title']:
            ws.cell(t['r0'], c0, '令和8年度 共通経費（本庁舎）').font = Font(bold=True, size=14)
            h = t['r0'] + 2
        heads = [t['cost_head'], t['amount']] + ([t['memo_head']] if t['memo'] else [])
        for j, v in enumerate(heads):
            ws.cell(h, c0 + j, v).font = Font(bold=True)
        rng2 = random.Random(stem + 'v')
        for i, (hm, amt) in enumerate(items, 1):
            ws.cell(h + i, c0, hm)
            v = amt
            if t['text_amount'] and rng2.random() < 0.6:
                x = rng2.random()
                v = f"{amt:,}".translate(_ZEN) if x < 0.3 else (f" {amt:,} " if x < 0.6 else f"{amt:,}円")
            ws.cell(h + i, c0 + 1, v).number_format = '#,##0'
            if t['memo'] and rng2.random() < 0.3:
                ws.cell(h + i, c0 + 2, rng2.choice(['年間契約', '前年度比増', '単価改定']))
        last = h + len(items)
        if t['total']:
            ws.cell(last + 1, c0, '合計').font = Font(bold=True)
            ws.cell(last + 1, c0 + 1, sum(a for _h, a in items)).number_format = '#,##0'
        lc = c0 + len(heads) + 1
        names = [k for k, _p in base]
        if truth:
            for j, v in enumerate([t['cost_head']] + names + ['合計']):
                ws.cell(h, lc + j, v).font = Font(bold=True)
            sums = [0] * len(base)
            for i, (hm, amt) in enumerate(items, 1):
                shares = allocate(amt, base)
                ws.cell(h + i, lc, hm)
                for j, sv in enumerate(shares, 1):
                    ws.cell(h + i, lc + j, sv).number_format = '#,##0'
                    sums[j - 1] += sv
                ws.cell(h + i, lc + len(base) + 1, amt).number_format = '#,##0'
            rr = h + len(items) + 1
            ws.cell(rr, lc, '合計')
            for j, sv in enumerate(sums, 1):
                ws.cell(rr, lc + j, sv).number_format = '#,##0'
            ws.cell(rr, lc + len(base) + 1, sum(a for _h, a in items)).number_format = '#,##0'
        elif t['stale']:
            for j, v in enumerate([t['cost_head']] + [f'旧{k + 1}' for k in range(t['stale'])] + ['合計']):
                ws.cell(h, lc + j, v)
            for i in range(1, len(items) + 4):
                for j in range(t['stale'] + 2):
                    ws.cell(h + i, lc + j, f'旧{i}' if j == 0 else 1000 * i)
        st = wb.create_sheet(t['base_sheet'])
        st.cell(1, 1, t['base_item']).font = Font(bold=True)
        st.cell(1, 2, t['base_num']).font = Font(bold=True)
        r = 2
        for i, (k, p) in enumerate(base):
            if t['base_blank'] and i == len(base) // 2:
                r += 1
            st.cell(r, 1, k)
            st.cell(r, 2, p)
            r += 1
        if t['base_total']:
            st.cell(r, 1, '合計').font = Font(bold=True)
            st.cell(r, 2, sum(p for _k, p in base))
        wb.save(os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx"))
    print('  ', stem, t['cost_sheet'], t['base_sheet'], t['cost_head'], t['base_num'], len(items), len(base))


def kw(t, **over):
    t.update(over)
    return t


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    write_pair('未見1_標準', build(s(1), rng.randint(3, 10), rng.randint(3, 8)), out)
    write_pair('未見2_基準の合計行と空行', kw(build(s(2), 6, 6), base_total=True, base_blank=True), out)
    write_pair('未見3_0と文字の金額', kw(build(s(3), 7, 5, zero=2), text_amount=True), out)
    write_pair('未見4_表題と合計行', kw(build(s(4), 5, 4), title=True, r0=2, c0=2, total=True, memo=True), out)
    write_pair('未見5_最大が2つと前の表', kw(build(s(5), 6, 6, tie=True), stale=8), out)
    write_pair('未見6_大きい', build(s(6), 15, 12), out)
    return out


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    write_pair('按分_本番', build(7, 6, 5))
    write_pair('按分_試験A', build(17, 12, 10))
    write_pair('按分_試験B', kw(build(27, 6, 6), base_total=True, base_blank=True))
    write_pair('按分_試験C', build(37, 5, 6, zero=2))
    write_pair('按分_試験D', kw(build(47, 6, 5), text_amount=True))
    write_pair('按分_試験E', kw(build(57, 5, 4), title=True, r0=3, c0=2))
    write_pair('按分_試験F', kw(build(67, 6, 5), total=True, memo=True))
    write_pair('按分_試験G', build(77, 5, 6, tie=True))
    write_pair('按分_試験H', kw(build(87, 4, 4), stale=7))
    shutil.copyfile(os.path.join(HERE, '按分_本番_正解.xlsx'), os.path.join(HERE, '按分_試験I_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, '按分_本番_正解.xlsx'), os.path.join(HERE, '按分_試験I_正解.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
