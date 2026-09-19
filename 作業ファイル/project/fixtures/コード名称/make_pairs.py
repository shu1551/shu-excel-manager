# -*- coding: utf-8 -*-
"""選んでいるコードの列を、別のシートの対応表で名称に置き換える（C 突合 28・2026-09-18）。

列の決まり: 選んでいる 1 列＝コードの列。対応表は別のシート（コードの列と名称の列の 2 列）。
正解の決まり（依頼文に全部書く）:
  - 選んだ列の値を、対応表の名称に置き換える（元の値は置き換えてよい＝上書きしてよい）。見出しは変えない。
  - 対応表に無いコードはそのまま残す。空のセルは触らない。
  - コードは前後の空白・先頭の 0・数値と文字の違いを無視して突き合わせる。
  - 2 回撃っても結果が変わらない（もう名称になっている値はそのまま）。
組: 本番／A 800 行／B 表題と単位／C 合計の行／D 空行／E 列の並び違い／F 撃った後にもう一度／G 対応表が数値のコード／
  H 対応表に無いコードが混ざる
"""
import os
import random
import shutil
import sys

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import detail   # noqa: E402
from vocab import SHEETS   # noqa: E402

REQUEST = ('選んでいる列（コードの列）の値を、別のシートの対応表（コードの列と名称の列が並ぶ 2 列の表）を見て名称に'
           '置き換えてください（元の値は上書きしてよい）。見出しは変えない。'
           '見出しの行は必ず「選んでいるセルの行」＝Selection.Cells(1,1).Row をそのまま使う（自分で見出しの行を探さない）。'
           '対応表のシートは、選んだ列の値と同じ値が並ぶ列を持つシートを形で探す（シート名で探さない）。'
           '対応表は 2 列で、左がコード・右が名称（コードを探すのは左の列だけ）。'
           '対応表に無いコードはそのまま残し、空のセルは触らない。'
           'コードは前後の空白・先頭の 0・数値と文字の違いを無視して突き合わせる。'
           'もう名称になっている値はそのまま（2 回撃っても結果が変わらない）。')
PHRASES = ['選んだ列のコードを名称に置き換えて', '対応表を見てコードを名前に直して', 'コード列を名称に一括変換して',
           '選んでいる列のコードを対応表の名前にして', 'コードを読み替えて名称にして', 'コードのままだと分からないので名称に']


def build(seed, n, unknown=0, code_num=False, **kw):
    t = detail.design(seed, n, **kw)
    rng = random.Random(seed + 77)
    v = t['vocab']
    iname = detail.name_of(t, 'item')
    names = sorted({row[iname] for row in t['rows']})
    codes = {nm: (k + 1) * 10 + 1 for k, nm in enumerate(names)}
    t['map'] = {(c if code_num else f"{c:04d}"): nm for nm, c in codes.items()}
    t['code_head'] = v.word('code')
    t['name_head'] = iname
    t['master_sheet'] = rng.choice([s for s in SHEETS['master']])
    t['code_num'] = code_num
    unknown_codes = [(9000 + i) if code_num else f"{9000 + i:04d}" for i in range(unknown)]
    for i, row in enumerate(t['rows']):
        c = codes[row[iname]]
        row[iname] = c if code_num else f"{c:04d}"
        if unknown_codes and i % max(1, n // max(1, unknown)) == 0 and unknown_codes:
            row[iname] = unknown_codes.pop()
    t['cols'] = [(t['code_head'] if nm == iname else nm, r) for nm, r in t['cols']]
    for row in t['rows']:                                   # 列の名前を変えたので値のキーも付け替える
        row[t['code_head']] = row.pop(iname)
    return t


def write_before(path, t):
    lay = detail.write_before(path, t)
    wb = load_workbook(path)
    ms = wb.create_sheet(t['master_sheet'])
    ms.cell(1, 1, t['code_head']).font = Font(bold=True)
    ms.cell(1, 2, t['name_head']).font = Font(bold=True)
    for i, (c, nm) in enumerate(sorted(t['map'].items(), key=lambda kv: str(kv[0])), 2):
        ms.cell(i, 1, c)
        ms.cell(i, 2, nm)
    wb.save(path)
    return lay


def build_truth(xl, before, truth, t, lay):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(t['sheet'])
        c0, h, last = lay['c0'], lay['h'], lay['last']
        cc = c0 + detail.col_of(t, 'item') - 1
        table = {str(k).strip().lstrip('0'): v for k, v in t['map'].items()}
        for r in range(h + 1, last + 1):
            v = ws.Cells(r, cc).Value
            if v is None or str(v).strip() == '':
                continue
            key = str(v).strip()
            if isinstance(v, float) and v.is_integer():
                key = str(int(v))
            nm = table.get(key.lstrip('0'))
            if nm:
                ws.Cells(r, cc).Value = nm
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            lay = write_before(before, t)
            sel = f"{get_column_letter(detail.col_of(t, 'item'))}{lay['h']}"
            build_truth(xl, before, truth, t, lay)
            open(os.path.join(out, f"{stem}_選択.txt"), 'w', encoding='utf-8').write(sel + '\n')
            print('  ', stem, t['sheet'], '/', t['master_sheet'], [n for n, _r in t['cols']], '選ぶ列', sel)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', build(s(1), rng.randint(20, 60))),
          ('未見2_表題と合計行', build(s(2), 30, title=True, unit=True, total=True)),
          ('未見3_並び違いと数値コード', build(s(3), 25, code_num=True, swap=True)),
          ('未見4_空行と未登録', build(s(4), 28, unknown=3, blanks=2))], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('コード_本番', build(1001, 30)),
          ('コード_試験A', build(1002, 800, n_items=10)),
          ('コード_試験B', build(1003, 25, title=True, unit=True)),
          ('コード_試験C', build(1004, 24, total=True)),
          ('コード_試験D', build(1005, 26, blanks=2)),
          ('コード_試験E', build(1006, 28, swap=True)),
          ('コード_試験G', build(1007, 22, code_num=True)),
          ('コード_試験H', build(1008, 20, unknown=3)),
          ('コード_試験I', build(1009, 24, no_date=True))], HERE)   # 日付の無い表（2026-09-18 夜）
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, 'コード_本番_正解.xlsx'), os.path.join(HERE, f'コード_試験F_{x}.xlsx'))
    shutil.copyfile(os.path.join(HERE, 'コード_本番_選択.txt'), os.path.join(HERE, 'コード_試験F_選択.txt'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
