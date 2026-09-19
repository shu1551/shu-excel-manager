# -*- coding: utf-8 -*-
"""選んだ列を、選んだ順に表の左から並べ直す（A 掃除 8・2026-09-18）。

列の決まり: 選んでいる列（見出しのセル・何本でも）。選んだ順＝Selection.Areas の順。<名前>_選択.txt に書く。
正解の決まり（依頼文に全部書く）:
  - 選んだ列を選んだ順に表の左端から並べ、選ばなかった列はその右に元の順のまま続ける。
  - 動かすのは見出しの行から表の最後の行（下の合計の行を含む）まで。表題・単位の行は動かさない。
  - 値・表示形式・太字は列と一緒に動かす。既に選んだ順に並んでいれば何もしない。
作り方: 正解は「前」を openpyxl で読み、表の塊（見出しの行〜最後の行）のセルを並べ替えて書き直す。
組: 本番 3 列／A 800 行 2 列／B 表題と単位／C 合計の行／D 途中の空行／E 全部の列を逆順／F 撃った後にもう一度／G 番号の列／H 1 列だけ
"""
import copy
import os
import random
import shutil
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import detail   # noqa: E402

REQUEST = ('選んでいる列（見出しのセル）を、選んだ順に表の左端から並べ直してください。選んだ順＝Selection.Areas の順（Areas(1) を左端に）。'
           '選ばなかった列は、その右に元の順のまま続ける。'
           '見出しの行は必ず「選んでいるセルの行」＝Selection.Cells(1,1).Row をそのまま使う（自分で見出しの行を探さない）。'
           '表の列＝見出しの行で、表のいちばん左の列から右へ見て空のセルに当たる手前の列まで。'
           '動かすのは見出しの行から表の最後の行（下の合計の行を含む）までで、表題・単位の行（見出しより上）は動かさない。'
           '値・表示形式・太字は列と一緒に動かす（列ごと切り取って挿入しない＝表題が動く）。既に選んだ順に並んでいれば何もしない。')
PHRASES = ['選んだ順に列を並べ替えて', '列の順番を選んだ順にして', '選んだ列を左に寄せて並べて', 'この順番に列を入れ替えて',
           '列の並びを選んだ順に変えて', '選んだ列を先頭に持ってきて']


def reorder_file(before, truth, h, order):
    """before の表の塊（見出しの行〜最後の行）を order（元の列番号の並び・1 から）に並べ替えて truth に保存。"""
    wb = load_workbook(before)
    ws = wb.worksheets[0]
    last = ws.max_row
    n = len(order)
    for r in range(h, last + 1):
        src = [(ws.cell(r, c).value, ws.cell(r, c).number_format, copy.copy(ws.cell(r, c).font)) for c in range(1, n + 1)]
        for j, c in enumerate(order, 1):
            v, nf, ft = src[c - 1]
            cell = ws.cell(r, j)
            cell.value = v
            cell.number_format = nf
            cell.font = ft
    wb.save(truth)


def make(jobs, out):
    for stem, t, roles in jobs:
        before = os.path.join(out, f"{stem}_前.xlsx")
        truth = os.path.join(out, f"{stem}_正解.xlsx")
        lay = detail.write_before(before, t)
        picked = [detail.col_of(t, r) for r in roles]
        order = picked + [c for c in range(1, lay['ncols'] + 1) if c not in picked]
        reorder_file(before, truth, lay['h'], order)
        sel = ",".join(f"{get_column_letter(c)}{lay['h']}" for c in picked)
        open(os.path.join(out, f"{stem}_選択.txt"), 'w', encoding='utf-8').write(sel + '\n')
        print('  ', stem, t['sheet'], [n for n, _r in t['cols']], '選ぶ列', sel)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', detail.design(s(1), rng.randint(20, 60)), ['num', 'item']),
          ('未見2_表題と合計行', detail.design(s(2), 30, title=True, unit=True, total=True), ['kamoku', 'num', 'date']),
          ('未見3_並び違い', detail.design(s(3), 25, swap=True), ['date', 'kamoku']),
          ('未見4_空行と番号列', detail.design(s(4), 28, blanks=2, no_col=True), ['item', 'no'])], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('並替_本番', detail.design(2101, 30), ['num', 'date', 'item']),
          ('並替_試験A', detail.design(2102, 800, n_items=10), ['item', 'num']),
          ('並替_試験B', detail.design(2103, 25, title=True, unit=True), ['kamoku', 'date']),
          ('並替_試験C', detail.design(2104, 24, total=True), ['num', 'text', 'item']),
          ('並替_試験D', detail.design(2105, 26, blanks=2), ['text', 'num']),
          ('並替_試験E', detail.design(2106, 28, swap=True), ['text', 'kamoku', 'date', 'num', 'item']),
          ('並替_試験G', detail.design(2107, 22, no_col=True), ['item', 'no', 'num']),
          ('並替_試験H', detail.design(2108, 20), ['num']),
          ('並替_試験I', detail.design(2109, 24, no_date=True), ['num', 'item'])], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '並替_本番_正解.xlsx'), os.path.join(HERE, f'並替_試験F_{x}.xlsx'))
    # 撃った後にもう一度＝並べ終わった位置（左から 3 列）を同じ順で選ぶ → 何も変わらない
    open(os.path.join(HERE, '並替_試験F_選択.txt'), 'w', encoding='utf-8').write('A1,B1,C1\n')
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
