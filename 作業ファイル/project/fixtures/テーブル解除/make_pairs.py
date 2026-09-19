# -*- coding: utf-8 -*-
"""アクティブなシートのテーブルを、普通のセルの範囲に戻す（H テーブル 68・2026-09-18 第二期）。

列の決まり: 選ぶ列は無い。アクティブなシートのテーブル全部（ほかのシートのテーブルは触らない）。
正解の決まり（依頼文に全部書く）:
  - ListObject.Unlist で範囲に戻す。値・書式（色・罫線）・式はそのまま残す（テーブルの参照の式は普通の番地の式になる）。
  - テーブルが無ければ何もしない。
組: 本番 1 つ／A 800 行／B 2 つ／C 集計行つき／D 計算列（式）つき／E 表題の下／F 撃った後にもう一度／G ほかのシートにもテーブル（それは残す）／H テーブルなし
"""
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import detail   # noqa: E402

REQUEST = ('アクティブなシートのテーブルを全部、普通のセルの範囲に戻してください（テーブルの解除）。'
           'ListObject.Unlist で範囲に戻す。値・色・罫線・式はそのまま残す。ほかのシートのテーブルは触らない。テーブルが無ければ何もしない。')
PHRASES = ['テーブルを普通の範囲に戻して', 'テーブルを解除して', 'テーブル化をやめて普通の表にして',
           'このシートのテーブルを範囲に変換', 'テーブルの設定を外して', 'テーブルをただの表に戻して']


def write_before(path, t, kind, n_tables, other, title):
    wb = Workbook()
    ws = wb.active
    ws.title = t['sheet']
    if title:
        ws['A1'] = '令和8年度 支出明細'
        ws['A1'].font = Font(bold=True, size=14)
    h0 = 3 if title else 1
    nc = len(t['cols'])
    for j, (nm, _r) in enumerate(t['cols'], 1):
        ws.cell(h0, j, nm).font = Font(bold=True)
    for i, row in enumerate(t['rows'], h0 + 1):
        for j, (nm, role) in enumerate(t['cols'], 1):
            c = ws.cell(i, j, row[nm])
            if role == 'date':
                c.number_format = 'yyyy/m/d'
    last = h0 + len(t['rows'])
    if kind == 'calc':
        a = get_column_letter(detail.col_of(t, 'num'))
        ws.cell(h0, nc + 1, '税込').font = Font(bold=True)
        for i in range(h0 + 1, last + 1):
            ws.cell(i, nc + 1, f"=ROUND({a}{i}*1.1,0)")
        nc += 1
    tb = Table(displayName='T明細', ref=f"A{h0}:{get_column_letter(nc)}{last}")
    tb.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(tb)
    if n_tables > 1:
        c2 = nc + 3
        ws.cell(h0, c2, '区分').font = Font(bold=True)
        ws.cell(h0, c2 + 1, '件数').font = Font(bold=True)
        for i in range(3):
            ws.cell(h0 + 1 + i, c2, f"区分{i + 1}")
            ws.cell(h0 + 1 + i, c2 + 1, i + 2)
        tb2 = Table(displayName='T区分', ref=f"{get_column_letter(c2)}{h0}:{get_column_letter(c2 + 1)}{h0 + 3}")
        tb2.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
        ws.add_table(tb2)
    if other:
        ws2 = wb.create_sheet('マスタ')
        ws2['A1'], ws2['B1'] = 'コード', '名称'
        for i in range(3):
            ws2.cell(2 + i, 1, 100 + i)
            ws2.cell(2 + i, 2, f"名称{i + 1}")
        tb3 = Table(displayName='Tマスタ', ref="A1:B4")
        tb3.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
        ws2.add_table(tb3)
    wb.save(path)


def build_truth(xl, before, truth, t, kind):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(t['sheet'])
        if kind == 'totals' and int(ws.ListObjects.Count):
            ws.ListObjects(1).ShowTotals = True             # 集計行つき（前にも付ける）
            wb.Save()
        for i in range(int(ws.ListObjects.Count), 0, -1):
            ws.ListObjects(i).Unlist()
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t, kind, n_tables, other, title in jobs:
            before, truth = os.path.join(out, f"{stem}_前.xlsx"), os.path.join(out, f"{stem}_正解.xlsx")
            if n_tables:
                write_before(before, t, kind, n_tables, other, title)
            else:
                detail.write_before(before, t)
            build_truth(xl, before, truth, t, kind)
            print('  ', stem, t['sheet'], kind, 'テーブル', n_tables, 'ほかのシートにも' if other else '')
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', detail.design(s(1), rng.randint(20, 50)), '', 1, False, False),
          ('未見2_集計行と表題', detail.design(s(2), 30), 'totals', 1, False, True),
          ('未見3_計算列と2つ', detail.design(s(3), 25), 'calc', 2, False, False),
          ('未見4_ほかのシート', detail.design(s(4), 28), '', 1, True, False)], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('解除_本番', detail.design(4601, 30), '', 1, False, False),
          ('解除_試験A', detail.design(4602, 800), '', 1, False, False),
          ('解除_試験B', detail.design(4603, 25), '', 2, False, False),
          ('解除_試験C', detail.design(4604, 24), 'totals', 1, False, False),
          ('解除_試験D', detail.design(4605, 26), 'calc', 1, False, False),
          ('解除_試験E', detail.design(4606, 28), '', 1, False, True),
          ('解除_試験G', detail.design(4607, 22), '', 1, True, False),
          ('解除_試験H', detail.design(4608, 20), '', 0, False, False)], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '解除_本番_正解.xlsx'), os.path.join(HERE, f'解除_試験F_{x}.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
