# -*- coding: utf-8 -*-
"""実物らしい表・第 5 弾（2026-09-17 夜）: 見えていない行と、表の上の集計欄、按分の小数の人数。

  共通（振り直し・突合・集約・月次・帳票）: グループ化で折りたたんだ行（非表示の行）＝人は「全部の行」のつもりで頼む
  振り直し: 明細の上に「支出合計」の小さな集計欄（見出しは 5 行目）
  按分: 人数が小数（会計年度任用職員の 0.5 人換算）・金額 0 の費目

  py make_real5.py [種]  → fixtures\\<型>\\未見_実物5（種 1）／未見_実物5種N
"""
import importlib.util
import os
import random
import sys
from fractions import Fraction

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))


def _mod(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


MR = _mod('make_real', os.path.join(HERE, 'make_real.py'))
AB = _mod('mp_按分', os.path.join(HERE, '按分', 'make_pairs.py'))
FR, TG, SY, CH, GJ = MR.FR, MR.TG, MR.SY, MR.CH, MR.GJ
post, data_rows = MR.post, MR.data_rows
SEED = 1


def _out(kind):
    p = os.path.join(HERE, kind, '未見_実物5' if SEED == 1 else f'未見_実物5種{SEED}')
    os.makedirs(p, exist_ok=True)
    for f in os.listdir(p):
        if f.endswith('.xlsx'):
            os.remove(os.path.join(p, f))
    return p


def hide_rows(ws, rows, key, p=0.3):
    """行をグループ化して折りたたむ（outline_level 1・hidden）。"""
    r2 = random.Random(key)
    for r in rows:
        if r2.random() < p:
            ws.row_dimensions[r].outline_level = 1
            ws.row_dimensions[r].hidden = True
    ws.sheet_properties.outlinePr.summaryBelow = True


def furi(seed):
    out = _out('振り直し')
    rng = random.Random(seed)
    t, b = FR.build(seed * 9000 + 1, rng.randint(25, 45), FR.ORDER2, 2)
    stem = '実物5_01_折りたたんだ行'
    FR.write_pair(stem, t, b, out=out)
    rows, _ = data_rows(1, len(b), ())
    post(out, stem, lambda wb, truth: hide_rows(wb['支出明細'], rows, stem))

    t, b = FR.build(seed * 9000 + 2, rng.randint(25, 45), FR.ORDER1, 2)
    stem = '実物5_02_上に集計欄'
    FR.write_pair(stem, t, b, r0=5, out=out)

    def f(wb, truth):
        ws = wb['支出明細']
        ws.cell(1, 1, '支出合計').font = Font(bold=True)
        ws.cell(1, 2, sum(x[2] for x in b)).number_format = '#,##0'
        ws.cell(2, 1, '件数')
        ws.cell(2, 2, len(b))
        ws.cell(3, 1, '（令和8年9月30日現在）')
    post(out, stem, f)
    return out


def totsu(seed):
    out = _out('突合')
    rng = random.Random(seed)
    old, new = TG.build(seed * 9000 + 1, rng.randint(22, 40), 3, 2, 2)
    stem = '実物5_01_旧に折りたたんだ行'
    TG.write_pair(stem, old, new, out=out)
    rows, _ = data_rows(1, len(old), ())
    post(out, stem, lambda wb, truth: hide_rows(wb['旧システム'], rows, stem))
    return out


def shuyaku(seed):
    out = _out('集約')
    rng = random.Random(seed)
    data = SY.build(seed * 9000 + 1, rng.randint(3, 6), rows=(4, 9))
    stem = '実物5_01_課のシートに折りたたんだ行'
    SY.write_pair(stem, data, out=out)

    def f(wb, truth):
        for k, (ka, items) in enumerate(data):
            hide_rows(wb[ka], range(2, 2 + len(items)), f"{stem}/{k}", p=0.4)
    post(out, stem, f)
    return out


def getsuji(seed):
    out = _out('月次')
    rng = random.Random(seed)
    rows = GJ.build(seed * 9000 + 1, rng.randint(40, 80))
    stem = '実物5_01_折りたたんだ行'
    GJ.write_pair(stem, rows, out=out)
    post(out, stem, lambda wb, truth: hide_rows(wb['支出明細'], range(2, 2 + len(rows)), stem))
    return out


def chohyo(seed):
    out = _out('帳票')
    rng = random.Random(seed)
    recs = CH.build(seed * 9000 + 1, rng.randint(6, 12))
    stem = '実物5_01_折りたたんだ件'
    CH.write_pair(stem, recs, out=out)

    def f(wb, truth):
        ws = wb[CH.SHEET]
        r2 = random.Random(stem)
        for i in range(len(recs)):
            if r2.random() < 0.3:
                for r in (4 + 2 * i, 5 + 2 * i):                   # 表題 1・見出し 2 段・1 件 2 行
                    ws.row_dimensions[r].outline_level = 1
                    ws.row_dimensions[r].hidden = True
    post(out, stem, f)
    return out


def anbun(seed):
    out = _out('按分')
    rng = random.Random(seed)
    items, staff = AB.build(seed * 9000 + 1, rng.randint(4, 8), rng.randint(4, 7))
    r2 = random.Random('half')
    staff = [(k, p + (Fraction(1, 2) if r2.random() < 0.5 else 0)) for k, p in staff]
    items = items[:-1] + [(items[-1][0], 0)]                           # 金額 0 の費目
    stem = '実物5_01_小数の人数と金額0'
    for truth in (False, True):
        wb = Workbook()
        ws = wb.active
        ws.title = '共通経費'
        ws.cell(1, 1, '費目')
        ws.cell(1, 2, '金額')
        for i, (h, a) in enumerate(items, 2):
            ws.cell(i, 1, h)
            ws.cell(i, 2, a)
        if truth:
            names = [k for k, _p in staff]
            for j, v in enumerate(['費目'] + names + ['合計']):
                ws.cell(1, 4 + j, v)
            total = sum(p for _k, p in staff)
            top = max(p for _k, p in staff)
            ti = next(i for i, (_k, p) in enumerate(staff) if p == top)
            sums = [0] * len(staff)
            for i, (h, a) in enumerate(items, 2):
                shares = [int(Fraction(a) * p / total) for _k, p in staff]    # 1 円未満切り捨て（正の数）
                shares[ti] += a - sum(shares)
                ws.cell(i, 4, h)
                for j, sv in enumerate(shares, 1):
                    ws.cell(i, 4 + j, sv)
                    sums[j - 1] += sv
                ws.cell(i, 4 + len(staff) + 1, a)
            rr = 2 + len(items)
            ws.cell(rr, 4, '合計')
            for j, sv in enumerate(sums, 1):
                ws.cell(rr, 4 + j, sv)
            ws.cell(rr, 4 + len(staff) + 1, sum(a for _h, a in items))
        st = wb.create_sheet('課別人数')
        st.cell(1, 1, '課名')
        st.cell(1, 2, '人数')
        for i, (k, p) in enumerate(staff, 2):
            st.cell(i, 1, k)
            st.cell(i, 2, float(p) if p != int(p) else int(p))
        wb.save(os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx"))
    return out


if __name__ == '__main__':
    seed = SEED = int(sys.argv[1]) if len(sys.argv) > 1 and sys.argv[1].isdigit() else 1
    for name, fn in (('振り直し', furi), ('突合', totsu), ('集約', shuyaku), ('月次', getsuji), ('帳票', chohyo), ('按分', anbun)):
        p = fn(seed)
        print(f"{name}: {len([f for f in os.listdir(p) if f.endswith('_前.xlsx')])} 組 → {p}")
