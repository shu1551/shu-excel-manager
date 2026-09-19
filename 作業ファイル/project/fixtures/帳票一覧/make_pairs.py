# -*- coding: utf-8 -*-
"""帳票（結合セルの受付簿）を 1 行 1 件の一覧に直す（どの表でも動く形・2026-09-18 作り直し。前の版は帳票一覧）。

形で決める: 帳票の見出しは 2 段、1 件は 2〜3 行で、縦に結合したセルと、件の中のラベル＋値の行でできている。
項目の名前・シート名・列の並び・件の行数は表ごとに違う（名前の組を種で選ぶ）。
正解の決まり:
  - 一覧の見出しは 1 行。2 段の見出しは下段の名前、下段が無い所は上段の名前、件の中のラベルも見出し。
    並びは帳票に出る順（件の 1 行目の左から右、次の行）。
  - 一覧の左上＝帳票の見出しの上段の行・帳票の右端の列の 2 つ右（1 列空ける）。元の帳票は触らない。
  - 値は帳票のまま（番号の先頭の 0 を落とさない・日付は日付）。値の無い欄は空のまま。
  - 入れない: 値の無い空の枠（印刷用）・※で始まる注記・改ページで繰り返した見出し。
  - 一覧の場所に前の一覧が残っていたら消してから書く（2 回撃っても同じ）。
組: 本番 8 件／A 30 件／B 1 件 3 行（空の欄あり）／C 表題なし／D 空の枠と※注記／E 前の一覧が残る／
  F 列の並びが違う／G 改ページで見出しを繰り返す／H 番号が先頭 0 の文字／I 撃った後にもう一度
"""
import datetime
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
from vocab import Vocab   # noqa: E402

# 項目の名前の組（表ごとに選ぶ）: no=番号・group=上段・name1/name2=下段の 2 つ・amount=金額・date=日付・labels=件の中のラベル
NAME_SETS = [
    {'no': '受付番号', 'group': '申請者', 'name1': '団体名', 'name2': '代表者', 'amount': '申請額', 'date': '受付日',
     'labels': ['所在地', '電話']},
    {'no': '整理番号', 'group': '事業者', 'name1': '名称', 'name2': '代表者名', 'amount': '要望額', 'date': '申請日',
     'labels': ['住所', '連絡先']},
    {'no': '申請番号', 'group': '申込者', 'name1': '団体', 'name2': '担当者', 'amount': '申請金額', 'date': '受付年月日',
     'labels': ['所在', '電話番号']},
    {'no': '管理番号', 'group': '相手方', 'name1': '相手方名', 'name2': '代表', 'amount': '契約金額', 'date': '契約日',
     'labels': ['住所地', 'ＦＡＸ']},
]
SHEETS = ['受付簿', '申請受付', '受付一覧', '様式', '受付台帳']
DANTAI = ['上町自治会', '下町子ども会', '本町商店会', '緑ヶ丘老人クラブ', '川口地区防災会', '中央公園愛護会', '北山婦人会',
          '南台スポーツ少年団', '旭町青年会', '若葉保存会', '東野農業者組合', '桜通り町内会', '西浜漁業者の会',
          'みどり子育てサークル', '青葉文化協会', '栄町交通安全協会', '大沢集落営農組合', '山手ボランティア会']
SEI = ['佐藤', '鈴木', '高橋', '田中', '伊藤', '渡辺', '山本', '中村', '小林', '加藤', '吉田', '山田']
MEI = ['一郎', '花子', '誠', '恵', '学', '悟', '由美', '健太', '真理', '修', '陽子', '大輔']
MACHI = ['本町', '上町', '下町', '旭町', '栄町', '緑町', '青葉台', '桜木', '川原', '新田']
THIN = Side(style='thin')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
REQUEST = ('アクティブなシートの帳票（結合セルで 1 件が複数行）を、帳票の右に 1 列空けて 1 行 1 件の一覧に直してください。'
           '見出しは 1 行（2 段の見出しは下段の名前、下段が無い所は上段の名前、件の中のラベルも見出し）、'
           '並びは帳票に出る順（件の 1 行目の左から右、そのあと件の中のラベルの順）。'
           '一覧の左上は帳票の見出しの上段の行・帳票の右端の列の 2 つ右。'
           '帳票の右端の列は「2 段の見出し（結合セル・枠線）の右端の列」（前に作った一覧の見出しは 1 行で結合も枠線も無い＝帳票ではない）。'
           'その 2 つ右から必ず書き、そこから右に何か残っていれば消してから書く。'
           '値は帳票のまま'
           '（番号の先頭の 0 を落とさない・日付は日付）。値の無い欄は空のまま。'
           '値の無い空の枠・※で始まる注記・改ページで繰り返した見出しは入れない。元の帳票は触らない。'
           '一覧の場所に前の一覧が残っていたら消してから書く')
PHRASES = ['帳票を1行1件の一覧に直して', '結合セルの帳票をデータベース形式に', '受付簿をリストにして',
           '様式の帳票を一覧表に変換して', '1件複数行の帳票を1行にまとめて', '帳票から一覧を作って']


def build(seed, n, phone=False, zero_id=False, blank_phone=False):
    rng = random.Random(seed)
    v = Vocab(seed)
    N = dict(rng.choice(NAME_SETS))
    if not phone:
        N['labels'] = N['labels'][:1]
    recs = []
    day = datetime.date(2026, 4, 1) + datetime.timedelta(days=rng.randrange(0, 20))
    for i in range(n):
        no = f"{i + 1:04d}" if zero_id else f"R8-{i + 1:03d}"
        dantai = rng.choice(DANTAI) if n > len(DANTAI) else DANTAI[(seed + i * 7) % len(DANTAI)]
        rep = rng.choice(SEI) + ' ' + rng.choice(MEI)
        amt = rng.randrange(5, 300) * 1000
        day = day + datetime.timedelta(days=rng.randrange(0, 3))
        addr = f"架空市{rng.choice(MACHI)}{rng.randrange(1, 9)}-{rng.randrange(1, 30)}"
        rec = {N['no']: no, N['name1']: dantai, N['name2']: rep, N['amount']: amt,
               N['date']: datetime.datetime(day.year, day.month, day.day), N['labels'][0]: addr}
        if len(N['labels']) > 1:
            rec[N['labels'][1]] = ('' if (blank_phone and rng.random() < 0.3)
                                   else f"0186-{rng.randrange(10, 99)}-{rng.randrange(1000, 9999)}")
        recs.append(rec)
    return {'recs': recs, 'N': N, 'sheet': v.sheet('meisai') if rng.random() < 0.3 else rng.choice(SHEETS),
            'layout': (N['no'], N['group'], N['amount'], N['date']), 'title': True, 'empty_frames': 0,
            'note': False, 'stale': 0, 'page': 0}


def _head(ws, r, t, first_col=1):
    N = t['N']
    c = first_col
    for name in t['layout']:
        if name == N['group']:
            ws.cell(r, c, N['group'])
            ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 1)
            ws.cell(r + 1, c, N['name1'])
            ws.cell(r + 1, c + 1, N['name2'])
            for cc in (c, c + 1):
                for rr in (r, r + 1):
                    ws.cell(rr, cc).border = BOX
                    ws.cell(rr, cc).font = Font(bold=True)
                    ws.cell(rr, cc).alignment = Alignment(horizontal='center')
            c += 2
        else:
            ws.cell(r, c, name)
            ws.merge_cells(start_row=r, start_column=c, end_row=r + 1, end_column=c)
            for rr in (r, r + 1):
                ws.cell(rr, c).border = BOX
                ws.cell(rr, c).font = Font(bold=True)
                ws.cell(r, c).alignment = Alignment(horizontal='center', vertical='center')
            c += 1
    return c - 1


def _columns(t):
    N = t['N']
    pos, c = {}, 1
    for name in t['layout']:
        if name == N['group']:
            pos[N['name1']], pos[N['name2']] = c, c + 1
            c += 2
        else:
            pos[name] = c
            c += 1
    return pos, pos[N['name1']], pos[N['name2']]


def list_fields(t):
    pos, _l, _v = _columns(t)
    top = sorted(pos, key=lambda k: pos[k])
    return top + list(t['N']['labels'])


def write_pair(stem, t, out=HERE):
    N = t['N']
    labels = N['labels']
    k = 1 + len(labels)
    pos, lab_c, val_c = _columns(t)
    fields = list_fields(t)
    for truth in (False, True):
        wb = Workbook()
        ws = wb.active
        ws.title = t['sheet']
        h = 1
        last_col = len(t['layout']) + 1
        if t['title']:
            ws.cell(1, 1, '令和8年度 地域づくり活動補助金 申請受付簿').font = Font(bold=True, size=14)
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
            h = 2
        _head(ws, h, t)
        r = h + 2
        frames = [dict(x) for x in t['recs']] + [None] * t['empty_frames']
        for i, rec in enumerate(frames):
            if t['page'] and i and i % t['page'] == 0:
                _head(ws, r, t)
                r += 2
            for name, c in pos.items():
                v = rec.get(name, '') if rec else ''
                if v != '':
                    ws.cell(r, c, v)
                if name in (N['name1'], N['name2']):
                    continue
                ws.merge_cells(start_row=r, start_column=c, end_row=r + k - 1, end_column=c)
            for j, lab in enumerate(labels, 1):
                ws.cell(r + j, lab_c, lab)
                v = rec.get(lab, '') if rec else ''
                if v != '':
                    ws.cell(r + j, val_c, v)
            for rr in range(r, r + k):
                for cc in range(1, last_col + 1):
                    ws.cell(rr, cc).border = BOX
            r += k
        if t['note']:
            ws.cell(r, 1, '※ 番号は受付順に付けています。金額は千円未満を切り捨てています。')
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last_col)
        lc = last_col + 2
        if truth:
            for j, f in enumerate(fields):
                ws.cell(h, lc + j, f)
            for i, rec in enumerate(t['recs'], 1):
                for j, f in enumerate(fields):
                    v = rec.get(f, '')
                    if v != '':
                        ws.cell(h + i, lc + j, v)
        elif t['stale']:
            for j, f in enumerate(fields):
                ws.cell(h, lc + j, f)
            for i in range(1, t['stale'] + 1):
                ws.cell(h + i, lc, f"旧-{i:03d}")
                ws.cell(h + i, lc + 1, f"旧団体{i}")
                ws.cell(h + i, lc + 3, 1000 * i)
        for c in range(1, last_col + 1):
            ws.column_dimensions[chr(64 + c)].width = 16
        wb.save(os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx"))
    print('  ', stem, t['sheet'], t['layout'], labels, len(t['recs']))


def kw(t, **over):
    if 'swap' in over:
        over.pop('swap')
        N = t['N']
        t['layout'] = (N['date'], N['no'], N['amount'], N['group'])
    t.update(over)
    return t


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    write_pair('未見1_標準', build(s(1), rng.randint(3, 15)), out)
    write_pair('未見2_3行と表題なし', kw(build(s(2), rng.randint(4, 12), phone=True, blank_phone=True), title=False), out)
    write_pair('未見3_空の枠と注記', kw(build(s(3), rng.randint(2, 9)), empty_frames=rng.randint(2, 6), note=True), out)
    write_pair('未見4_前の一覧', kw(build(s(4), rng.randint(2, 6)), stale=rng.randint(12, 30)), out)
    write_pair('未見5_並び違い', kw(build(s(5), rng.randint(3, 10), phone=True), swap=True), out)
    write_pair('未見6_改ページ', kw(build(s(6), rng.randint(15, 40)), page=rng.choice((5, 8, 10)), note=True), out)
    write_pair('未見7_全部入り', kw(build(s(7), 25, phone=True, zero_id=True, blank_phone=True), swap=True, title=False,
                                empty_frames=3, note=True, stale=40, page=10), out)
    return out


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    write_pair('帳票_本番', build(5, 8))
    write_pair('帳票_試験A', build(15, 30))
    write_pair('帳票_試験B', build(25, 6, phone=True, blank_phone=True))
    write_pair('帳票_試験C', kw(build(35, 5), title=False))
    write_pair('帳票_試験D', kw(build(45, 4), empty_frames=5, note=True))
    write_pair('帳票_試験E', kw(build(55, 3), stale=20))
    write_pair('帳票_試験F', kw(build(65, 7, phone=True), swap=True))
    write_pair('帳票_試験G', kw(build(75, 23), page=10))
    write_pair('帳票_試験H', build(85, 9, zero_id=True))
    shutil.copyfile(os.path.join(HERE, '帳票_本番_正解.xlsx'), os.path.join(HERE, '帳票_試験I_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, '帳票_本番_正解.xlsx'), os.path.join(HERE, '帳票_試験I_正解.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
