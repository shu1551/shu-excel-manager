# -*- coding: utf-8 -*-
"""一覧の 1 行ごとに、ひな形のシートを写して差し込み、1 件 1 枚のシートを作る（D 帳票 31・2026-09-18 第二期）。

列の決まり: 選ぶ列は無い。アクティブなシート＝一覧。ひな形＝「{見出し}」の文字が入っているほかのシート。
正解の決まり（依頼文に全部書く）:
  - 一覧の見出しの行＝空でないセルが 2 つ以上並ぶ最初の行。本文の行ごとに（空の行と「合計」「計」の行は除く）ひな形を最後のシートの後ろへ写す。
  - シートの名前＝その行の左端の列の値。同じ名前のシートがあれば消して作り直す。
  - 写したシートのセルで、文字がちょうど「{見出し}」なら一覧のその列の値そのもの（数・日付のまま）を入れる。
    文の中の「{見出し}」は、一覧のそのセルの見た目の文字（.Text）に置き換える。一覧に無い見出しの「{…}」はそのまま。
  - 一覧とひな形は変えない。
組: 本番 5 件／A 25 件／B 表題つきの一覧／C 合計の行／D 途中の空行／E 一覧とひな形の間に別のシート／F 撃った後にもう一度／G 別のひな形（一覧に無い {…}）／H 1 件
"""
import datetime
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
from vocab import Vocab   # noqa: E402

REQUEST = ('アクティブなシートの一覧の 1 行ごとに、ひな形のシートを写して差し込み、1 件 1 枚のシートを作ってください。'
           'ひな形＝アクティブなシート以外で、「{見出し}」の形の文字（波かっこで囲んだ見出しの語）が入っているシート。'
           '一覧の見出しの行＝空でないセルが 2 つ以上並ぶ最初の行（文字でも数でもよい）。本文は見出しの次の行から最後の行まで。'
           '行がまるごと空の行と、行のどこかのセルに「合計」「計」「総計」とある行は作らない。'
           '本文の行ごとに、ひな形のシートを最後のシートの後ろへ写し（上から順に後ろへ足す）、シートの名前はその行の左端の列の値にする。'
           '同じ名前のシートが既にあれば、消してから作り直す。'
           '写したシートのセルで、文字がちょうど「{見出し}」だけのセルには、一覧のその列の値そのもの（数・日付は数・日付のまま）を入れる。'
           '文の中にある「{見出し}」は、一覧のそのセルの見た目の文字（Range.Text）に置き換える。一覧に無い見出しの「{…}」はそのまま残す。'
           '一覧のシートとひな形のシートは変えない。')
PHRASES = ['一覧から1件1枚の帳票を作って', 'ひな形に差し込んで人ごとのシートを作って', '差し込み印刷のように1行ずつシートにして',
           '一覧の行ごとに通知書を作って', 'ひな形のシートに一覧を差し込んで', '1件ずつ様式に流し込んだシートを']
TEMPLATES = ['ひな形', '様式', '通知書ひな形', 'テンプレート']


def build(seed, n, **kw):
    rng = random.Random(seed)
    v = Vocab(seed)
    t = {'sheet': rng.choice(['一覧', '対象者一覧', '交付一覧', '名簿']), 'tsheet': rng.choice(TEMPLATES),
         'no': rng.choice(['番号', '整理番号', '交付番号']), 'name': v.word('name'), 'ka': v.word('ka'),
         'amount': v.word('amount').replace('（円）', ''), 'date': rng.choice(['交付日', '決定日', '支払日']),
         'title': False, 'total': False, 'blank_at': None, 'tmpl_first': False, 'alt': False, 'no_date': False}
    t.update(kw)
    sei = ['佐藤', '鈴木', '高橋', '田中', '伊藤', '渡辺', '小松', '工藤', '三浦', '菅原']
    mei = ['一郎', '花子', '健', '美咲', '大輔', '由美', '翔', '恵']
    kas = v.items(4, rng)
    t['rows'] = [(1001 + i, rng.choice(sei) + ' ' + rng.choice(mei), rng.choice(kas), rng.randrange(10, 900) * 100,
                  datetime.datetime(2026, rng.randint(4, 12), rng.randint(1, 28))) for i in range(n)]
    return t


def write_before(path, t):
    wb = Workbook()
    ws = wb.active
    ws.title = t['sheet']
    h = 3 if t['title'] else 1
    if t['title']:
        ws.cell(1, 1, '令和8年度 交付対象者一覧').font = Font(bold=True, size=14)
    heads = [t['no'], t['name'], t['ka'], t['amount']] + ([] if t['no_date'] else [t['date']])   # 日付の無い一覧も（2026-09-18 夜）
    for j, nm in enumerate(heads, 1):
        ws.cell(h, j, nm).font = Font(bold=True)
        ws.column_dimensions['ABCDE'[j - 1]].width = 16
    r = h
    for i, row in enumerate(t['rows']):
        if t['blank_at'] is not None and i == t['blank_at']:
            r += 1
        r += 1
        for j, x in enumerate(row[:len(heads)], 1):
            c = ws.cell(r, j, x)
            if j == 4:
                c.number_format = '#,##0'
            elif j == 5:
                c.number_format = 'yyyy/m/d'
    if t['total']:
        ws.cell(r + 1, 2, '合計').font = Font(bold=True)
        ws.cell(r + 1, 4, sum(x[3] for x in t['rows'])).number_format = '#,##0'
    if t['tmpl_first']:
        memo = wb.create_sheet('メモ')                 # 一覧とひな形の間に別のシート（ひな形は位置でなく {…} で探す）
        memo['A1'], memo['A2'] = '作業メモ', '9月分は差し替え予定'
    tp = wb.create_sheet(t['tsheet'])
    if not t['alt']:
        tp['A1'] = '交付決定通知書'
        tp['A1'].font = Font(bold=True, size=16)
        tp['A3'] = '{%s} 様' % t['name']
        tp['A5'], tp['B5'] = t['ka'], '{%s}' % t['ka']
        tp['A6'], tp['B6'] = '交付額', '{%s}' % t['amount']
        tp['B6'].number_format = '#,##0"円"'
        if not t['no_date']:
            tp['A7'], tp['B7'] = t['date'], '{%s}' % t['date']
            tp['B7'].number_format = 'yyyy"年"m"月"d"日"'
        tp['A9'] = '上記のとおり {%s} 円の交付を決定しました（番号 {%s}）。' % (t['amount'], t['no'])
    else:
        tp['B2'] = '支払通知'
        tp['B4'] = '{%s}' % t['no']
        tp['C4'] = '{%s} 殿' % t['name']
        tp['B6'] = '{%s} に {%s} を支払います。' % (t['date'], t['amount'])
        tp['B8'] = '備考: {備考}'
    tp.column_dimensions['A'].width = 14
    tp.column_dimensions['B'].width = 18
    wb.save(path)
    return h


def build_truth(xl, before, truth, t, h):
    import re
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(t['sheet'])
        tp = wb.Worksheets(t['tsheet'])
        heads = [str(ws.Cells(h, j).Value) for j in range(1, 6)]
        r = h + 1
        last = ws.Cells(ws.Rows.Count, 1).End(-4162).Row
        last = max(last, ws.Cells(ws.Rows.Count, 2).End(-4162).Row)
        for r in range(h + 1, last + 1):
            vals = [ws.Cells(r, j).Value for j in range(1, 6)]
            if all(v is None for v in vals) or any(str(v) in ('合計', '計', '総計') for v in vals if v is not None):
                continue
            key = str(int(vals[0])) if isinstance(vals[0], float) and vals[0] == int(vals[0]) else str(vals[0])
            tp.Copy(None, wb.Worksheets(wb.Worksheets.Count))
            ns = wb.Worksheets(wb.Worksheets.Count)
            ns.Name = key
            ur = ns.UsedRange
            for rr in range(1, int(ur.Row) + int(ur.Rows.Count)):
                for cc in range(1, int(ur.Column) + int(ur.Columns.Count)):
                    cell = ns.Cells(rr, cc)
                    v = cell.Value
                    if not isinstance(v, str) or '{' not in v:
                        continue
                    m = re.fullmatch(r'\{([^{}]+)\}', v)
                    if m and m.group(1) in heads:
                        x = ws.Cells(r, heads.index(m.group(1)) + 1)
                        cell.Value2 = x.Value2
                        continue
                    cell.Value = re.sub(r'\{([^{}]+)\}', lambda mm: (str(ws.Cells(r, heads.index(mm.group(1)) + 1).Text)
                                                                    if mm.group(1) in heads else mm.group(0)), v)
        ws.Activate()
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t in jobs:
            before, truth = os.path.join(out, f"{stem}_前.xlsx"), os.path.join(out, f"{stem}_正解.xlsx")
            h = write_before(before, t)
            build_truth(xl, before, truth, t, h)
            print('  ', stem, t['sheet'], t['tsheet'], len(t['rows']))
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', build(s(1), rng.randint(3, 8))),
          ('未見2_表題と合計', build(s(2), 5, title=True, total=True)),
          ('未見3_間に別のシート', build(s(3), 4, tmpl_first=True, alt=True)),
          ('未見4_空行', build(s(4), 6, blank_at=2))], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('差込_本番', build(3201, 5)),
          ('差込_試験A', build(3202, 25)),
          ('差込_試験B', build(3203, 4, title=True)),
          ('差込_試験C', build(3204, 5, total=True)),
          ('差込_試験D', build(3205, 6, blank_at=3)),
          ('差込_試験E', build(3206, 4, tmpl_first=True)),
          ('差込_試験G', build(3207, 5, alt=True)),
          ('差込_試験H', build(3208, 1)),
          ('差込_試験I', build(3209, 4, no_date=True))], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '差込_本番_正解.xlsx'), os.path.join(HERE, f'差込_試験F_{x}.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
