# -*- coding: utf-8 -*-
"""同じ見出しのテーブルを全部追加して集計する（どの表でも動く形・2026-09-18 作り直し。前の版は PQ課別集計）。

形で決める: アクティブなシートのテーブルと **見出しの集まりが同じ**テーブルを、このブックから全部足す
（見出しが違うテーブルは入れない）。集計する 2 列は選んでいる列（1 つ目＝文字の列＝行、2 つ目＝数の列＝値）。
見出しの語・シート名・テーブル名・項目の名前は表ごとに違う（vocab.py）。作る物の名前は元のテーブル名から。
正解の決まり:
  - 新しいシート（名前は 元のテーブル名 と _集計 をつないだ名前）の A1 に同じ名前のテーブルとして読み込む。クエリ名も同じ。
  - 列は 1 つ目の列・2 つ目の列（合計）、1 つ目の列の昇順。
  - 同じ名前のクエリやシートが既にあれば消して作り直す（2 回撃っても 1 つ）。
組: 本番 3 表／A 6 表・大量／B 1 表だけ／C 列の並びが違う表がある／D 見出しが違う表がある（入れない）／
  E 前のクエリとシートが残る（中身が古い）／F 撃った後にもう一度／G 空行と余計な列の表／H 項目が 1 つの表
"""
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import importlib.util   # noqa: E402

from vocab import Vocab   # noqa: E402

_spec = importlib.util.spec_from_file_location('mp_tate', os.path.join(HERE, '..', '縦持ち', 'make_pairs.py'))
T = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(T)

MONTHS = ['4月', '5月', '6月', '7月', '8月', '9月', '10月', '11月', '12月', '1月', '2月', '3月']
REQUEST = ('パワークエリで、アクティブなシートのテーブルと見出しの集まりが同じテーブルを、このブックから全部追加して 1 つにし、'
           '選んでいる 2 列（1 つ目＝文字の列を行、2 つ目＝数の列を値）で集計してください。'
           '見出しが違うテーブル（見出しの数や語が違うもの）は入れない。列の並びが違うだけなら見出しで合わせて足す。'
           '列は 1 つ目の列・2 つ目の列の合計の 2 列、1 つ目の列の昇順。'
           '新しいシート（名前は 元のテーブル名 と _集計 をつないだ名前）の A1 に同じ名前のテーブルとして読み込む。'
           'クエリ名も同じ。同じ名前のクエリやシートが既にあれば消して作り直す')
PHRASES = ['同じ見出しのテーブルを全部足して集計して', 'パワークエリで月ごとの表をまとめて集計',
           '月別のテーブルを追加して合計を出して', 'クエリで複数の表を1つにして集計',
           '同じ形の表を全部つないで集計表に', 'パワークエリで追加して集計']


def design(seed, n_tables, rows, swap=0, other=False, diff_heads=False, blanks=0, n_items=6):
    rng = random.Random(seed)
    v = Vocab(seed)
    item = v.word('ka')
    num = v.word('amount').split('（')[0]        # M の列参照 [名前] は全角の括弧で壊れる
    kamoku = v.word('kamoku')
    items = v.items(min(n_items, 20), rng)
    kamokus = v.kamokus(5, rng)
    base = v.table('meisai')
    tables = []
    first_cols = [item, kamoku, num]
    rng.shuffle(first_cols)                      # 選ぶ列は表ごとに違う番地になる（A1,C1 に固定しない）
    for k in range(n_tables):
        m = MONTHS[k]
        cols = list(first_cols)
        if k == swap and k > 0:
            cols = [num, item, kamoku]
        body = [{item: rng.choice(items), kamoku: rng.choice(kamokus), num: rng.randrange(1, 900) * 100}
                for _ in range(rows)]
        tables.append({'sheet': f"{m}" if k else v.sheet('meisai'), 'name': base if k == 0 else f"{base}{m}",
                       'cols': cols, 'body': body, 'blanks': blanks if k == 0 else 0})
    others = []
    if other:
        yosan = v.word('yosan')
        others.append({'sheet': '予算', 'name': base + '予算',
                       'cols': [item, yosan], 'body': [{item: k, yosan: 999900} for k in items[:4]]})
    if diff_heads:
        others.append({'sheet': 'ほかの表', 'name': base + '別',
                       'cols': [item, kamoku, num, v.word('no')],
                       'body': [{item: rng.choice(items), kamoku: rng.choice(kamokus),
                                 num: rng.randrange(1, 900) * 100, v.word('no'): f"A{i:03d}"} for i in range(5)]})
    return {'item': item, 'num': num, 'kamoku': kamoku, 'tables': tables, 'others': others,
            'base': base, 'stale': False}


def _put(ws, t, tb):
    for j, c in enumerate(tb['cols'], 1):
        ws.cell(1, j, c).font = Font(bold=True)
    r = 2
    for i, row in enumerate(tb['body']):
        if tb.get('blanks') and i and i % max(1, len(tb['body']) // (tb['blanks'] + 1)) == 0 and tb['blanks']:
            r += 1
        for j, c in enumerate(tb['cols'], 1):
            ws.cell(r, j, row[c])
        r += 1
    last = r - 1
    tab = Table(displayName=tb['name'], ref=f"A1:{get_column_letter(len(tb['cols']))}{last}")
    tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(tab)


def write_before(path, t):
    wb = Workbook()
    wb.remove(wb.active)
    for tb in t['tables'] + t['others']:
        ws = wb.create_sheet(tb['sheet'])
        _put(ws, t, tb)
    wb.save(path)


def m_formula(t, order='Order.Ascending'):
    names = ", ".join(f'"{tb["name"]}"' for tb in t['tables'])
    return f'''let
    Source = Excel.CurrentWorkbook(),
    Rows = Table.SelectRows(Source, each List.Contains({{{names}}}, [Name])),
    Combined = Table.Combine(Rows[Content]),
    Typed = Table.TransformColumnTypes(Combined, {{{{"{t['item']}", type text}}, {{"{t['num']}", type number}}}}),
    Grouped = Table.Group(Typed, {{"{t['item']}"}}, {{{{"{t['num']}", each List.Sum([{t['num']}]), type nullable number}}}}),
    Sorted = Table.Sort(Grouped, {{{{"{t['item']}", {order}}}}})
in
    Sorted'''


def make(jobs, out=HERE):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            name = t['base'] + '_集計'
            write_before(before, t)
            first = t['tables'][0]
            sel = ",".join(f"{get_column_letter(first['cols'].index(c) + 1)}1" for c in (t['item'], t['num']))
            if t['stale']:                                   # 前のクエリ（降順・古い）とシートを残す
                wb = xl.Workbooks.Open(before)
                try:
                    T.load_query(wb, name, m_formula(t, 'Order.Descending'), name)
                    wb.Worksheets(first['sheet']).Activate()
                    wb.Save()
                finally:
                    wb.Close(SaveChanges=False)
            wb = xl.Workbooks.Open(before)
            try:
                T.load_query(wb, name, m_formula(t), name)
                wb.Worksheets(first['sheet']).Activate()
                wb.SaveAs(truth, FileFormat=51)
            finally:
                wb.Close(SaveChanges=False)
            open(os.path.join(out, f"{stem}_選択.txt"), 'w', encoding='utf-8').write(sel + '\n')
            print('  ', stem, first['sheet'], name, len(t['tables']), '表　選ぶ列', sel)
    finally:
        xl = None
        vf._quit(pid)


def kw(t, **over):
    t.update(over)
    return t


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', design(s(1), rng.randint(2, 5), rng.randint(10, 60))),
          ('未見2_並び違いとほかの表', design(s(2), 4, 30, swap=2, other=True)),
          ('未見3_見出し違いと前のクエリ', kw(design(s(3), 3, 25, diff_heads=True), stale=True)),
          ('未見4_1表と空行', design(s(4), 1, 20, blanks=2)),
          ('未見5_大量', design(s(5), 6, 400, n_items=12))], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('追加_本番', design(3, 3, 20)),
          ('追加_試験A', design(13, 6, 500, n_items=12)),
          ('追加_試験B', design(23, 1, 25)),
          ('追加_試験C', design(33, 4, 30, swap=2)),
          ('追加_試験D', design(43, 3, 20, other=True, diff_heads=True)),
          ('追加_試験E', kw(design(53, 3, 18), stale=True)),
          ('追加_試験G', design(63, 3, 24, blanks=2, other=True)),
          ('追加_試験H', design(73, 2, 15, n_items=1))])
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '追加_本番_正解.xlsx'), os.path.join(HERE, f'追加_試験F_{x}.xlsx'))
    shutil.copyfile(os.path.join(HERE, '追加_本番_選択.txt'), os.path.join(HERE, '追加_試験F_選択.txt'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
