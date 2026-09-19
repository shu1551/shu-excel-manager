# -*- coding: utf-8 -*-
"""明細から月次集計表（項目×月）を作る（どの表でも動く形・2026-09-18 作り直し。前の版は月次集計）。

列の決まり: 選んでいる 3 列（1 つ目＝項目の列・2 つ目＝日付の列・3 つ目＝金額の列）。
見出しの語・シート名・項目の名前・列の並びは表ごとに違う（vocab.py・detail.py）。選んでいる列は <名前>_選択.txt に書く。
正解の決まり:
  - 集計表の左上＝明細の見出しの行・明細の右端の列の 2 つ右（1 列空ける）。明細は触らない。
  - 見出しは 1 行: 項目の列の見出し・4月・5月・…・3月（年度の順に 12 か月）・合計。支出の無い月も列を出して 0。
  - 項目は明細に出てくる順（重複なし・前後の空白は落として同じ項目）。各月＝その項目・その月の金額の合計（無ければ 0）。
    右端に行の合計、最後の行に「合計」（月ごとの合計と総合計）。
  - 日付が文字（2026/04/15・2026年4月15日）でも日付として読む。金額がカンマ・全角・空白つきの文字でも数で読む。
  - 入れない: 明細の空行・明細の下の合計の行。前の集計が残っていたら消してから書く（2 回撃っても同じ）。
組: 本番 60 行／A 2,000 行／B 文字の日付と金額／C 表題と単位の行／D 列の並び違い／E 前の集計が残る／
  F 空行と合計行／G 月の抜けと項目の空白／H 日付順でない／I 撃った後にもう一度
"""
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import detail   # noqa: E402

MONTHS = [4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3]
REQUEST = ('選んでいる 3 列（1 つ目が項目の列、2 つ目が日付の列、3 つ目が金額の列）で、明細の右に 1 列空けて月次集計表を'
           '作ってください。見出しは 1 行で 項目の列の見出し・4月から3月（年度の順に 12 か月）・合計。'
           '項目は明細に出てくる順（前後の空白は落として同じ項目）、各月はその項目・その月の金額の合計（無ければ 0）、'
           '右端に行の合計、最後に「合計」の行。明細の空行と明細の下の合計の行は入れない。'
           '日付が文字（2026/04/15・2026年4月15日）でも日付として読み、金額がカンマ・全角・空白つきの文字でも数で読む。'
           '集計表の場所に前の集計が残っていたら消してから書く。明細は触らない')
PHRASES = ['月ごとの集計表を明細の右に作って', '項目ごと月ごとに集計した表を作って', '月次集計表を作ってください',
           '明細から月別の集計表を', '月別の集計を右に出して', '毎月の執行状況の表を作って']


def summary(t):
    """明細 → (項目の順, {項目: {月: 合計}})（純 Python）。"""
    item = detail.name_of(t, 'item')
    date = detail.name_of(t, 'date')
    num = detail.name_of(t, 'num')
    order, tot = [], {}
    for row in t['rows']:
        key = str(row[item]).strip().strip('　').strip()
        d = detail.date_of(row[date])
        if key not in tot:
            order.append(key)
            tot[key] = {m: 0 for m in MONTHS}
        tot[key][d.month] += detail.num_of(row[num])
    return order, tot


def write_pair(stem, t, out=HERE, stale=0):
    order, tot = summary(t)
    sel = None
    for truth in (False, True):
        wb = Workbook()
        ws = wb.active
        ws.title = t['sheet']
        lay = detail.write_into(ws, t)
        h, lc = lay['h'], lay['ncols'] + 2
        if sel is None:
            sel = ",".join(f"{get_column_letter(detail.col_of(t, role))}{h}" for role in ('item', 'date', 'num'))
        if truth:
            ws.cell(h, lc, detail.name_of(t, 'item')).font = Font(bold=True)
            for j, m in enumerate(MONTHS, 1):
                ws.cell(h, lc + j, f"{m}月").font = Font(bold=True)
            ws.cell(h, lc + 13, '合計').font = Font(bold=True)
            for i, k in enumerate(order, 1):
                ws.cell(h + i, lc, k)
                for j, m in enumerate(MONTHS, 1):
                    ws.cell(h + i, lc + j, tot[k][m])
                ws.cell(h + i, lc + 13, sum(tot[k].values()))
            rr = h + len(order) + 1
            ws.cell(rr, lc, '合計').font = Font(bold=True)
            for j, m in enumerate(MONTHS, 1):
                ws.cell(rr, lc + j, sum(tot[k][m] for k in order))
            ws.cell(rr, lc + 13, sum(sum(tot[k].values()) for k in order))
        elif stale:
            ws.cell(h, lc, detail.name_of(t, 'item')).font = Font(bold=True)
            for j, m in enumerate(MONTHS, 1):
                ws.cell(h, lc + j, f"{m}月").font = Font(bold=True)
            ws.cell(h, lc + 13, '合計').font = Font(bold=True)
            for i in range(1, stale + 1):
                ws.cell(h + i, lc, f"前の{i}")
                for j in range(1, 14):
                    ws.cell(h + i, lc + j, 1000 * i)
            ws.cell(h + stale + 1, lc, '合計')
        wb.save(os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx"))
    open(os.path.join(out, f"{stem}_選択.txt"), 'w', encoding='utf-8').write(sel + '\n')
    print('  ', stem, t['sheet'], [c for c, _r in t['cols']], '選ぶ列', sel)


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    write_pair('未見1_標準', detail.design(s(1), rng.randint(20, 120)), out)
    write_pair('未見2_文字の日付と金額', detail.design(s(2), 80, text_date=True, text_amount=True), out)
    write_pair('未見3_表題と並び違い', detail.design(s(3), 50, swap=True, title=True, unit=True), out)
    write_pair('未見4_前の集計', detail.design(s(4), 40), out, stale=rng.randint(6, 11))
    write_pair('未見5_空行と合計行', detail.design(s(5), 60, blanks=3, total=True), out)
    write_pair('未見6_月の抜けと空白', detail.design(s(6), 30, gap_months=(8, 1, 2), item_space=True), out)
    write_pair('未見7_全部入り', detail.design(s(7), 150, swap=True, title=True, text_date=True, text_amount=True,
                                            blanks=2, total=True, item_space=True, sort=False), out, stale=10)
    return out


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    write_pair('月次_本番', detail.design(5, 60))
    write_pair('月次_試験A', detail.design(15, 2000, n_items=10))
    write_pair('月次_試験B', detail.design(25, 70, text_date=True, text_amount=True))
    write_pair('月次_試験C', detail.design(35, 40, title=True, unit=True))
    write_pair('月次_試験D', detail.design(45, 50, swap=True))
    write_pair('月次_試験E', detail.design(55, 30), stale=9)
    write_pair('月次_試験F', detail.design(65, 45, blanks=2, total=True))
    write_pair('月次_試験G', detail.design(75, 25, gap_months=(6, 11, 3), item_space=True))
    write_pair('月次_試験H', detail.design(85, 60, sort=False, no_col=True))
    shutil.copyfile(os.path.join(HERE, '月次_本番_正解.xlsx'), os.path.join(HERE, '月次_試験I_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, '月次_本番_正解.xlsx'), os.path.join(HERE, '月次_試験I_正解.xlsx'))
    shutil.copyfile(os.path.join(HERE, '月次_本番_選択.txt'), os.path.join(HERE, '月次_試験I_選択.txt'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
