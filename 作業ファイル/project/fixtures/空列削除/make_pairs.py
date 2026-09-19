# -*- coding: utf-8 -*-
"""表の中の、見出しも本文も全部空の列を削除して詰める（2026-09-18 第二期）。

列の決まり: 選ぶ列は無い（表全体）。
正解の決まり（依頼文に全部書く）:
  - 見出しの行＝空でないセルが 2 つ以上並ぶ最初の行。表の列＝見出しの行で最初に値のある列から、最後に値のある列まで。
  - 表の最後の行＝表の列のどれかに値がある最後の行。見出しの行から表の最後の行まで全部空の列を、列ごと削除して左へ詰める（右の列から消す）。
  - 見出しだけある列・表の外の列は消さない。空の列が無ければ何もしない。
組: 本番 1 列／A 800 行・2 列／B 表題と単位／C 合計の行／D 途中の空行と 2 列続き／E 番号の列の右／F 撃った後にもう一度／G 見出しだけの列（残す）／H 空の列なし
"""
import os
import random
import shutil
import sys

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import detail   # noqa: E402

REQUEST = ('アクティブなシートの表の中にある、見出しも本文も全部空の列を、列ごと削除して左へ詰めてください。'
           '見出しの行＝空でないセルが 2 つ以上並ぶ最初の行（文字でも数でもよい）。表の列＝見出しの行で、最初に値のある列から最後に値のある列まで。'
           '表の最後の行＝表の列のどれかに値がある最後の行。見出しの行から表の最後の行まで全部空の列を、右の列から順に EntireColumn.Delete で消す。'
           '見出しだけある列と、表の外の列は消さない。空の列が無ければ何もしない。値は変えない。')
PHRASES = ['表の中の空の列を消して', '空白の列を削除して詰めて', '何も入っていない列を消して',
           '間の空いた列を詰めて', '空列を取り除いて', '使っていない空の列を削除']


def insert_empty(path, t, cols, head_only=None):
    """before に空の列を差し込む（openpyxl・右から差し込めば位置はずれない）。正解は空の列の無い元の表。"""
    wb = load_workbook(path)
    ws = wb.worksheets[0]
    for c in sorted(cols, reverse=True):
        ws.insert_cols(c)
    if head_only:
        h = 3 if t.get('title') else 1
        if t.get('unit'):
            h = 4
        ws.insert_cols(head_only)
        ws.cell(h, head_only, '備考2')
    wb.save(path)


def make(jobs, out):
    for stem, t, cols, head_only in jobs:
        before, truth = os.path.join(out, f"{stem}_前.xlsx"), os.path.join(out, f"{stem}_正解.xlsx")
        detail.write_before(truth, t)
        shutil.copyfile(truth, before)
        insert_empty(before, t, cols)
        if head_only:                                        # 見出しだけの列は前にも正解にもある（消さない）
            insert_empty(before, t, [], head_only)
            insert_empty(truth, t, [], head_only - len([c for c in cols if c < head_only]))
        print('  ', stem, t['sheet'], [n for n, _r in t['cols']], '空の列', cols, '見出しだけ' if head_only else '')


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', detail.design(s(1), rng.randint(20, 50)), [rng.randint(2, 5)], None),
          ('未見2_表題と合計行', detail.design(s(2), 30, title=True, unit=True, total=True), [3], None),
          ('未見3_2列続きと空行', detail.design(s(3), 25, blanks=2), [4, 4], None),
          ('未見4_番号と見出しだけ', detail.design(s(4), 28, no_col=True), [2], 5)], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('空列_本番', detail.design(4701, 30), [3], None),
          ('空列_試験A', detail.design(4702, 800), [2, 5], None),
          ('空列_試験B', detail.design(4703, 25, title=True, unit=True), [4], None),
          ('空列_試験C', detail.design(4704, 24, total=True), [2], None),
          ('空列_試験D', detail.design(4705, 26, blanks=2), [3, 3], None),
          ('空列_試験E', detail.design(4706, 28, no_col=True), [2], None),
          ('空列_試験G', detail.design(4707, 22), [2], 5),
          ('空列_試験H', detail.design(4708, 20), [], None),
          ('空列_試験I', detail.design(4709, 24, no_date=True), [3], None)], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '空列_本番_正解.xlsx'), os.path.join(HERE, f'空列_試験F_{x}.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
