# -*- coding: utf-8 -*-
"""課別・科目別の支出額のピボットテーブル（鍛える回路の題材・ピボット 1 本目・2026-09-17 深夜）。

仕事: 「支出明細」シートの明細（支出日・課名・科目・摘要・支出額）から、新しいシート「課別科目別」の A3 にピボットテーブルを作る。
正解の決まり（依頼文に全部書く）:
  - 行＝課名・列＝科目・値＝支出額の合計（表示形式 #,##0）・行と列の総計あり。
  - 元の範囲は明細の見出しから最後の明細の行まで（明細の下の合計の行は入れない）。見出しの語で列を探す（列の並びは表ごとに違う）。
  - 「課別科目別」シートが既にあれば消して作り直す（2 回撃ってもピボットは 1 つ）。
正解のブックは、直す前のブックを自分の Excel（非表示）で開き、この決まりどおりにピボットを作って保存する。
組: 本番 60 行／A 1,500 行／B 表題・見出しが 4 行目／C 前の結果のシートが残っている／D 明細の下に合計の行／
  E 列の並びが違う／F 撃った後にもう一度
"""
import datetime
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
KA = ['総務課', '財政課', '企画課', '税務課', '福祉課', '建設課', '農林課', '教育総務課', '観光課', '環境課']
KAMOKU = ['需用費', '役務費', '委託料', '使用料及び賃借料', '工事請負費', '備品購入費', '旅費']
TEKIYO = ['コピー用紙', '郵送料', '清掃業務', '複合機リース', '修繕工事', 'パソコン購入', '出張旅費']
HEAD = ['支出日', '課名', '科目', '摘要', '支出額']
SWAP = ['課名', '支出額', '科目', '支出日', '摘要']
SHEET = '支出明細'
OUT_SHEET = '課別科目別'
REQUEST = ('支出明細から、新しいシート「課別科目別」の A3 にピボットテーブルを作ってください。行は課名、列は科目、値は支出額の合計'
           '（表示形式 #,##0）、行と列の総計あり。元の範囲は明細の見出しから最後の明細の行まで（明細の下の合計の行は入れない）。'
           '「課別科目別」シートが既にあれば消して作り直す')


def build(seed, n, n_ka=5):
    rng = random.Random(seed)
    kas = rng.sample(KA, n_ka)
    rows = []
    for _ in range(n):
        d = datetime.datetime(2026, rng.randint(4, 12), rng.randint(1, 28))
        rows.append({'支出日': d, '課名': rng.choice(kas), '科目': rng.choice(KAMOKU), '摘要': rng.choice(TEKIYO),
                     '支出額': rng.randrange(1, 800) * 100})
    return rows


def write_before(path, rows, cols=HEAD, title=False, total=False):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    h = 1
    if title:
        ws.cell(1, 1, '令和8年度 支出明細（一般会計）').font = Font(bold=True, size=14)
        ws.cell(2, 1, '単位：円')
        h = 4
    for j, c in enumerate(cols, 1):
        ws.cell(h, j, c).font = Font(bold=True)
    for i, r in enumerate(rows, 1):
        for j, c in enumerate(cols, 1):
            cell = ws.cell(h + i, j, r[c])
            if c == '支出日':
                cell.number_format = 'yyyy/m/d'
            if c == '支出額':
                cell.number_format = '#,##0'
    last = h + len(rows)
    if total:
        ws.cell(last + 1, cols.index('摘要') + 1, '合計').font = Font(bold=True)
        ws.cell(last + 1, cols.index('支出額') + 1, sum(r['支出額'] for r in rows))
    wb.save(path)
    return {'h': h, 'last': last, 'ncols': len(cols)}


def build_pivot(wb, lay, sheet_name=OUT_SHEET, rows_field='課名', cols_field='科目'):
    """決まりどおりのピボットを作る（COM）。前の結果シートがあれば消す。"""
    xl = wb.Application
    for sh in list(wb.Worksheets):
        if str(sh.Name) == sheet_name:
            xl.DisplayAlerts = False
            sh.Delete()
    src = wb.Sheets(SHEET)
    rng = src.Range(src.Cells(lay['h'], 1), src.Cells(lay['last'], lay['ncols']))
    out = wb.Worksheets.Add(After=wb.Worksheets(wb.Worksheets.Count))
    out.Name = sheet_name
    pc = wb.PivotCaches().Create(SourceType=1, SourceData=f"'{SHEET}'!{rng.GetAddress(True, True, -4150)}")
    pt = pc.CreatePivotTable(TableDestination=f"'{sheet_name}'!R3C1", TableName='P' + sheet_name)
    pt.PivotFields(rows_field).Orientation = 1
    pt.PivotFields(cols_field).Orientation = 2
    df = pt.AddDataField(pt.PivotFields('支出額'), '合計 / 支出額', -4157)
    df.NumberFormat = '#,##0'
    pt.RowGrand = True
    pt.ColumnGrand = True
    return pt


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, rows, kw in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            stale = kw.pop('stale', False)
            lay = write_before(before, rows, **kw)
            if stale:                                         # 前の結果（古い並びのピボット＝行が科目）を置いて保存
                wb = xl.Workbooks.Open(before)
                try:
                    build_pivot(wb, lay, rows_field='科目', cols_field='課名')
                    wb.Worksheets(SHEET).Activate()
                    wb.Save()
                finally:
                    wb.Close(SaveChanges=False)
            wb = xl.Workbooks.Open(before)
            try:
                build_pivot(wb, lay)
                wb.Worksheets(SHEET).Activate()
                wb.SaveAs(truth, FileFormat=51)
            finally:
                wb.Close(SaveChanges=False)
            print('  ', stem)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', build(s(1), rng.randint(30, 120), n_ka=rng.randint(3, 8)), {}),
          ('未見2_表題と合計行', build(s(2), 80), {'title': True, 'total': True}),
          ('未見3_並び違いと前の結果', build(s(3), 50), {'cols': SWAP, 'stale': True}),
          ('未見4_大量', build(s(4), 3000, n_ka=10), {})], out)
    return out


if __name__ == '__main__':
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        sys.exit(0)
    make([('ピボット_本番', build(3, 60), {}),
          ('ピボット_試験A', build(13, 1500, n_ka=10), {}),
          ('ピボット_試験B', build(23, 50), {'title': True}),
          ('ピボット_試験C', build(33, 40), {'stale': True}),
          ('ピボット_試験D', build(43, 45), {'total': True}),
          ('ピボット_試験E', build(53, 55), {'cols': SWAP})], HERE)
    shutil.copyfile(os.path.join(HERE, 'ピボット_本番_正解.xlsx'), os.path.join(HERE, 'ピボット_試験F_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, 'ピボット_本番_正解.xlsx'), os.path.join(HERE, 'ピボット_試験F_正解.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    lines = ['課別・科目別のピボットテーブルを作って', '支出明細を課名と科目でピボット集計して', '課ごと科目ごとの支出額をピボットで',
             'ピボットで課別科目別の表をお願いします', '支出明細のピボットテーブルを別シートに', '課名を行・科目を列にしたピボット']
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(lines) + '\n')
    print('ok')
