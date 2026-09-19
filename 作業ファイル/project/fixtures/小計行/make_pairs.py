# -*- coding: utf-8 -*-
"""選んでいる列の値が変わる所に「小計」の行を入れる（B 集計 22・2026-09-18）。

列の決まり: 選んでいる 1 列＝区分の列（その列で並んでいる表）。選んでいる列は <名前>_選択.txt に書く。
正解の決まり（依頼文に全部書く）:
  - 区分の値が変わる直前の行の下に 1 行挿入し、区分の列に「（その区分の値） 小計」と書く。
  - 数の列（見出しが 番号・No・コード・年度 を含む列と率の列は除く）に =SUM(そのかたまりの範囲) を入れる。
  - 最後のかたまりの下にも小計の行を入れる。表の下に合計（合計・計）の行があれば、その上に入れる。
  - 既に「小計」の行があれば作り直す（2 回撃っても増えない）。
組: 本番 3 区分／A 4 区分・多い行／B 表題と単位／C 下に合計の行／D 数の列が 2 本／E 列の並び違い／F 撃った後にもう一度／
  G 番号の列と摘要／H 区分が 2 つだけ
"""
import os
import random
import shutil
import sys

from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import shukei   # noqa: E402

REQUEST = ('選んでいる列（区分の列）の値が変わる所に小計の行を入れてください。'
           '区分の値が変わる直前の行の下に 1 行挿入し、区分の列に「（その区分の値） 小計」（値・半角空白・小計）と書く。'
           '数の列には =SUM(そのかたまりの範囲) の式を入れる（見出しに 番号・No・コード・年度 を含む列と、表示形式が % の率の列には入れない）。'
           '最後のかたまりの下にも小計の行を入れる。表の下に合計（合計・計）の行があれば、小計はその上に入れる。'
           '見出しの行は必ず「選んでいるセルの行」＝Selection.Cells(1,1).Row をそのまま使う（自分で見出しの行を探さない）。'
           '本文は見出しの次の行から最後の本文の行まで。既に「小計」の行があれば作り直す（行を増やさない）。')
PHRASES = ['選んだ列の区分ごとに小計行を入れて', '区分が変わる所に小計を挿入して', '小計の行を区切りごとに足して',
           '選んでいる列で小計行を作って', 'グループごとの小計行を入れてほしい', '区分の切れ目に小計を入れて']


def blocks(t):
    """[(かたまりの最初の添字, 最後の添字, 区分の値)]（純 Python・行は区分の列で並んでいる）。"""
    gname = shukei.name_of(t, 'group')
    out = []
    for i, row in enumerate(t['rows']):
        if out and out[-1][2] == row[gname]:
            out[-1][1] = i
        else:
            out.append([i, i, row[gname]])
    return [tuple(b) for b in out]


def build_truth(xl, before, truth, t, lay):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(t['sheet'])
        c0, h = lay['c0'], lay['h']
        gc = c0 + shukei.col_of(t, 'group') - 1
        ncols = [c0 + j - 1 for j, (_n, r) in enumerate(t['cols'], 1) if r == 'num']
        for first, lastk, val in reversed(blocks(t)):
            r1, r2 = h + 1 + first, h + 1 + lastk
            at = r2 + 1
            ws.Rows(at).Insert()
            ws.Cells(at, gc).Value = f"{val} 小計"
            for c in ncols:
                cl = get_column_letter(c)
                ws.Cells(at, c).Formula = f"=SUM({cl}{r1}:{cl}{r2})"
                ws.Cells(at, c).NumberFormat = '#,##0'
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            lay = shukei.write_before(before, t)
            sel = f"{get_column_letter(shukei.col_of(t, 'group'))}{lay['h']}"
            build_truth(xl, before, truth, t, lay)
            open(os.path.join(out, f"{stem}_選択.txt"), 'w', encoding='utf-8').write(sel + '\n')
            print('  ', stem, t['sheet'], [n for n, _r in t['cols']], '選ぶ列', sel)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', shukei.design(s(1), rng.randint(7, 14), group=3)),
          ('未見2_表題と合計行', shukei.design(s(2), 9, group=3, title=True, unit=True, total=True)),
          ('未見3_並び違い', shukei.design(s(3), 8, group=2, swap=True, nums=2)),
          ('未見4_番号列と摘要', shukei.design(s(4), 10, group=4, no_col=True, text=True))], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('小計_本番', shukei.design(801, 9, group=3)),
          ('小計_試験A', shukei.design(802, 16, group=4)),
          ('小計_試験B', shukei.design(803, 9, group=3, title=True, unit=True)),
          ('小計_試験C', shukei.design(804, 9, group=3, total=True)),
          ('小計_試験D', shukei.design(805, 9, group=3, nums=2)),
          ('小計_試験E', shukei.design(806, 8, group=2, swap=True)),
          ('小計_試験G', shukei.design(807, 9, group=3, no_col=True, text=True)),
          ('小計_試験H', shukei.design(808, 6, group=2))], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '小計_本番_正解.xlsx'), os.path.join(HERE, f'小計_試験F_{x}.xlsx'))
    shutil.copyfile(os.path.join(HERE, '小計_本番_選択.txt'), os.path.join(HERE, '小計_試験F_選択.txt'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
