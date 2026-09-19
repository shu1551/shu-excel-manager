# -*- coding: utf-8 -*-
"""一覧表を複数行の帳票ブロックに展開するお題ペア生成スクリプト（マクロの鍛え方 7番対応・第2版）。

穴の解消:
1. 帳票が並ぶ列の「真下」にある人のメモが消えないことを検査（bottom_memo）。
2. 数も日付も無い名簿（氏名・フリガナ・所属・役職・住所・電話・メール）でも動く（no_date_no_num）。
3. vocab.py を使って組ごとに見出しの語とシート名を変化させる。
"""
import datetime
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
from vocab import Vocab

THIN = Side(style='thin')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

REQUEST = ('アクティブなシートにある1行1件の一覧表を、表の右に1列空けて1件複数行のカード型帳票ブロックに展開してください。'
           '一覧の見出しの行＝空でないセルが2つ以上並ぶ最初の行（文字だけでもよい。数や日付が無くてもよい）。'
           '本文は見出しの次の行から最後の行まで（空行と合計行は除く）。'
           '帳票ブロックは、見出しの左端列から右端列に向かって、1行あたり「ラベル・値」を3組（計6列）ずつ並べます。'
           '1件あたりの行数は「見出しの列数 ÷ 3 の切り上げ」行（例えば4〜6列なら2行、7〜9列なら3行）。'
           '各件のブロック全体（切り上げ行数 × 6列）を細線罫線で囲み、件と件の間は1行空けてください。'
           '見出しラベルは太字、値は元のセルと同じ値と表示形式（先頭ゼロ・日付・カンマ等）を維持してください。'
           '展開先の右側や、展開した帳票の「下」にある人のメモや別の表は絶対に消さないでください。'
           '消してよいのは「前に作った帳票ブロック」の範囲だけです。'
           '元の一覧表は変更しないでください。')

PHRASES = ['一覧を複数行の帳票ブロックに展開して', '一覧表のデータを3列組のカード形式に並べ替えて',
           '1行1件のリストを複数行の伝票型に展開して', '一覧から3組並びの帳票ブロックを作って',
           '明細表を複数行のカード型帳票に直して', '表の右に1件複数行の帳票を展開して']


def build(seed, n, col_count=7, zero_id=False, start_pos=(2, 1), title=False,
          bottom_memo=False, no_date_no_num=False, total=False, blank_row=False):
    rng = random.Random(seed)
    v = Vocab(seed)
    
    sei = ['佐藤', '鈴木', '高橋', '田中', '伊藤', '渡辺', '山本', '中村']
    mei = ['一郎', '花子', '誠', '恵', '学', '悟', '由美', '健太']
    machi = ['本町', '上町', '下町', '旭町', '栄町', '緑町', '青葉台']
    
    if no_date_no_num:
        # 数も日付も一切無い名簿（テキストのみ）
        all_fields = [
            ('氏名', 'name'),
            ('フリガナ', 'kana'),
            (v.word('ka'), 'ka'),
            ('役職', 'pos'),
            ('住所地', 'addr'),
            ('連絡先', 'tel'),
            ('メール', 'email')
        ]
        fields = all_fields[:col_count]
    else:
        # 通常の一覧（vocab で語を変える）
        all_fields = [
            (v.word('code'), 'code'),
            (v.word('name'), 'name'),
            (v.word('ka'), 'ka'),
            (v.word('date'), 'date'),
            (v.word('amount').replace('（円）', ''), 'amount'),
            ('連絡先', 'tel'),
            ('所在地', 'addr'),
            (v.word('tekiyo'), 'tekiyo'),
            ('区分', 'kubun')
        ]
        fields = all_fields[:col_count]

    recs = []
    base_date = datetime.date(2026, 4, 1) + datetime.timedelta(days=rng.randrange(0, 15))
    for i in range(n):
        rec = {}
        day = base_date + datetime.timedelta(days=i * 2 + rng.randrange(0, 2))
        s_sei = rng.choice(sei)
        s_mei = rng.choice(mei)
        for name, kind in fields:
            if kind == 'code':
                rec[name] = f"{i + 1:04d}" if zero_id else f"R8-{i + 1:03d}"
            elif kind == 'name':
                rec[name] = s_sei + ' ' + s_mei
            elif kind == 'kana':
                rec[name] = 'サトウ イチロウ'
            elif kind == 'ka':
                rec[name] = rng.choice(v.items(4, rng))
            elif kind == 'pos':
                rec[name] = rng.choice(['主査', '課長補佐', '主任', '係長', '担当員'])
            elif kind == 'date':
                rec[name] = datetime.datetime(day.year, day.month, day.day)
            elif kind == 'amount':
                rec[name] = rng.randrange(10, 500) * 1000
            elif kind == 'tel':
                rec[name] = f"0186-{rng.randrange(10, 99)}-{rng.randrange(1000, 9999)}"
            elif kind == 'addr':
                rec[name] = f"架空市{rng.choice(machi)}{rng.randrange(1, 5)}-{rng.randrange(1, 20)}"
            elif kind == 'email':
                rec[name] = f"user{i+1}@example.lg.jp"
            elif kind == 'tekiyo':
                rec[name] = rng.choice(['通常申請分', '年度末繰越分', '至急要請案件', '定期更新手続き'])
            elif kind == 'kubun':
                rec[name] = rng.choice(['新規', '継続', '変更', '更新'])
        recs.append(rec)
        
    sheet_name = '名簿' if no_date_no_num else v.sheet('meisai')
    return {
        'fields': fields, 'recs': recs, 'start_pos': start_pos,
        'title': title, 'bottom_memo': bottom_memo, 'total': total,
        'blank_row': blank_row, 'sheet': sheet_name,
        'no_date_no_num': no_date_no_num
    }


def write_pair(stem, t, out=HERE):
    fields = t['fields']
    recs = t['recs']
    start_r, start_c = t['start_pos']
    col_count = len(fields)
    block_rows = (col_count + 2) // 3
    
    for truth in (False, True):
        wb = Workbook()
        ws = wb.active
        ws.title = t['sheet']
        
        # 表題（title=True の場合）
        head_r = start_r
        if t['title']:
            ws.cell(1, start_c, '令和8年度 業務受付明細一覧表').font = Font(bold=True, size=14)
            head_r = max(start_r, 3)
            
        # 一覧の見出し
        for j, (fname, _) in enumerate(fields):
            c = ws.cell(head_r, start_c + j, fname)
            c.font = Font(bold=True)
            c.border = BOX
            
        # 一覧の本文データ
        curr_r = head_r
        for i, rec in enumerate(recs):
            if t['blank_row'] and i == len(recs) // 2:
                curr_r += 1  # 途中の空行
            curr_r += 1
            for j, (fname, kind) in enumerate(fields):
                val = rec[fname]
                c = ws.cell(curr_r, start_c + j, val)
                c.border = BOX
                if kind == 'amount':
                    c.number_format = '#,##0'
                elif kind == 'date':
                    c.number_format = 'yyyy/m/d'
                elif kind == 'code' and str(val).startswith('0'):
                    c.number_format = '@'
                    
        # 合計行（total=True の場合）
        if t['total']:
            curr_r += 1
            ws.cell(curr_r, start_c, '合計').font = Font(bold=True)
            for j, (fname, kind) in enumerate(fields):
                cell = ws.cell(curr_r, start_c + j)
                cell.border = BOX
                if kind == 'amount':
                    tot = sum(r[fname] for r in recs)
                    cell.value = tot
                    cell.number_format = '#,##0'
                    
        # 展開先の列
        target_c = start_c + col_count + 1  # 1列空ける
        
        # 帳票が並ぶ列の「真下」にある人のメモ（bottom_memo=True）
        # 帳票は target_r から (block_rows + 1) * len(recs) 行まで並ぶ。
        # その下（30行目など）に置いたメモは、マクロ実行後も正解でも絶対に消えてはならない！
        if t['bottom_memo']:
            memo_r = head_r + (block_rows + 1) * len(recs) + 8  # 帳票の数行下
            ws.cell(memo_r, target_c, '【決裁・審査メモ】').font = Font(bold=True)
            ws.cell(memo_r + 1, target_c, '※未決裁につき押印確認を要する')
            ws.cell(memo_r + 1, target_c + 1, '担当: 総務課長')
            
        # 正解の展開（truth=True の場合）
        if truth:
            target_r = head_r
            for rec in recs:
                for idx, (fname, kind) in enumerate(fields):
                    r_off = idx // 3
                    c_off = (idx % 3) * 2
                    
                    # ラベル
                    lbl_cell = ws.cell(target_r + r_off, target_c + c_off, fname)
                    lbl_cell.font = Font(bold=True)
                    
                    # 値
                    val_cell = ws.cell(target_r + r_off, target_c + c_off + 1, rec[fname])
                    if kind == 'amount':
                        val_cell.number_format = '#,##0'
                    elif kind == 'date':
                        val_cell.number_format = 'yyyy/m/d'
                    elif kind == 'code' and str(rec[fname]).startswith('0'):
                        val_cell.number_format = '@'
                        
                # 外枠・細線罫線
                for rr in range(target_r, target_r + block_rows):
                    for cc in range(target_c, target_c + 6):
                        ws.cell(rr, cc).border = BOX
                        
                # 次のブロックへ（1行空ける）
                target_r += block_rows + 1

        # 列幅設定
        for col in ws.columns:
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = 14
            
        file_path = os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx")
        wb.save(file_path)
    print(f"Generated: {stem} ({len(recs)}件, {col_count}列, sheet={t['sheet']})")


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k
    
    write_pair('未見1_標準7列', build(s(1), rng.randint(4, 12), col_count=7), out)
    write_pair('未見2_数日付なし名簿', build(s(2), rng.randint(3, 8), col_count=6, no_date_no_num=True), out)
    write_pair('未見3_下メモあり', build(s(3), 4, col_count=5, bottom_memo=True), out)
    write_pair('未見4_9列先頭ゼロ', build(s(4), rng.randint(4, 10), col_count=9, zero_id=True), out)
    write_pair('未見5_表題つき合計あり', build(s(5), rng.randint(5, 12), col_count=6, title=True, total=True), out)
    write_pair('未見6_5列開始B3', build(s(6), rng.randint(3, 7), col_count=5, start_pos=(3, 2)), out)
    return out


def main():
    os.makedirs(HERE, exist_ok=True)
    
    # 手順書7番に対応したお題の生成（本番 ＋ 試験A〜H）
    write_pair('複数行帳票_本番', build(101, 6, col_count=7))
    write_pair('複数行帳票_試験A', build(102, 15, col_count=7))                             # 多件数
    write_pair('複数行帳票_試験B', build(103, 5, col_count=7, no_date_no_num=True))         # 【穴2解消】数・日付なし名簿
    write_pair('複数行帳票_試験C', build(104, 8, col_count=9, zero_id=True))                 # 9列・先頭ゼロ
    write_pair('複数行帳票_試験D', build(105, 5, col_count=6, title=True))                   # 表題つき
    write_pair('複数行帳票_試験E', build(106, 4, col_count=5, bottom_memo=True))             # 【穴1解消】帳票列の下にメモあり
    write_pair('複数行帳票_試験F', build(107, 7, col_count=8, blank_row=True))               # 途中に空行
    write_pair('複数行帳票_試験G', build(108, 6, col_count=7, total=True))                   # 合計行あり
    write_pair('複数行帳票_試験H', build(109, 1, col_count=6))                              # 1件のみ
    
    # 試験Iは冪等性テスト（本番の正解を前として渡す）
    shutil.copyfile(os.path.join(HERE, '複数行帳票_本番_正解.xlsx'), os.path.join(HERE, '複数行帳票_試験I_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, '複数行帳票_本番_正解.xlsx'), os.path.join(HERE, '複数行帳票_試験I_正解.xlsx'))
    
    with open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8') as f:
        f.write(REQUEST + '\n')
    with open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8') as f:
        f.write('\n'.join(PHRASES) + '\n')
    print('make_pairs.py (v2): All pairs generated successfully.')


if __name__ == '__main__':
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        unseen(int(sys.argv[2]))
    else:
        main()
