# -*- coding: utf-8 -*-
"""課別・月別の支出の推移の折れ線グラフ（鍛える回路の題材・グラフ 2 本目・2026-09-17 深夜）。

仕事: 「月次集計」シートの表（課名・4月〜3月・合計＋合計の行）から、課ごとの推移のマーカー付き折れ線グラフを作る。
正解の決まり（依頼文に全部書く）:
  - 系列＝課ごと（行ごと）: 名前＝課名のセル・値＝その課の 4月〜3月・項目＝見出しの 4月〜3月。合計の行と合計の列は入れない。
  - タイトル「課別月別支出の推移」・凡例は下・数値軸のタイトル「支出額（円）」・数値軸の表示形式 #,##0・目盛線あり・
    グラフの文字 Meiryo UI。
  - 置き場所: 表の下に 1 行空けた行の、表の左端（位置は比べない）。前に作ったグラフが残っていたら消して作り直す。
組: 本番 4 課／A 8 課／B 表題・B3 から／C 合計の行と列が無い表／D 月の見出しが「4月」でなく日付（2026/4/1 を m月 表示）／
  E 前のグラフ／F 空の課（全部 0 の課も系列に入れる）／G 撃った後にもう一度
"""
import datetime
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
KA = ['総務課', '財政課', '企画課', '税務課', '福祉課', '建設課', '農林課', '教育総務課', '観光課', '環境課', '会計課', '議会事務局']
MONTHS = [4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3]
SHEET = '月次集計'
TITLE = '課別月別支出の推移'
REQUEST = ('月次集計シートの表から、課ごとの月別支出の推移をマーカー付き折れ線グラフにして、表の下に1行空けた行の表の左端に作ってください。'
           '系列は課ごと（行ごと）、項目は4月から3月。合計の行と合計の列は入れない。タイトルは「課別月別支出の推移」、凡例は下、'
           '縦軸のタイトルは「支出額（円）」、縦軸の表示形式は #,##0、目盛線あり、グラフの文字は Meiryo UI。'
           '前に作ったグラフが残っていたら消して作り直す')


def build(seed, n, zero=False):
    rng = random.Random(seed)
    kas = rng.sample(KA, n)
    rows = [(k, [rng.randrange(0, 900) * 1000 for _ in MONTHS]) for k in kas]
    if zero:
        rows[len(rows) // 2] = (rows[len(rows) // 2][0], [0] * 12)
    return rows


def write_before(path, rows, title=False, r0=1, c0=1, totals=True, date_heads=False):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    h = r0
    if title:
        ws.cell(r0, c0, '令和8年度 月次集計（一般会計）').font = Font(bold=True, size=14)
        h = r0 + 2
    ws.cell(h, c0, '課名').font = Font(bold=True)
    for j, m in enumerate(MONTHS, 1):
        if date_heads:
            c = ws.cell(h, c0 + j, datetime.datetime(2026 if m >= 4 else 2027, m, 1))
            c.number_format = 'm"月"'
        else:
            ws.cell(h, c0 + j, f"{m}月")
        ws.cell(h, c0 + j).font = Font(bold=True)
    if totals:
        ws.cell(h, c0 + 13, '合計').font = Font(bold=True)
    for i, (k, vals) in enumerate(rows, 1):
        ws.cell(h + i, c0, k)
        for j, v in enumerate(vals, 1):
            ws.cell(h + i, c0 + j, v).number_format = '#,##0'
        if totals:
            ws.cell(h + i, c0 + 13, sum(vals)).number_format = '#,##0'
    last = h + len(rows)
    if totals:
        ws.cell(last + 1, c0, '合計').font = Font(bold=True)
        for j in range(1, 14 if totals else 13):
            ws.cell(last + 1, c0 + j, sum((r[1][j - 1] if j <= 12 else sum(r[1])) for r in rows)).number_format = '#,##0'
        last += 1
    wb.save(path)
    return {'h': h, 'c0': c0, 'n': len(rows), 'last': last}


def build_truth(xl, before, truth, lay):
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Sheets(SHEET)
        h, c0, n = lay['h'], lay['c0'], lay['n']
        for i in range(int(ws.ChartObjects().Count), 0, -1):
            ws.ChartObjects(i).Delete()
        top = ws.Cells(lay['last'] + 2, c0).Top
        co = ws.ChartObjects().Add(ws.Cells(lay['last'] + 2, c0).Left, top, 640, 320)
        co.Name = '表の整理_グラフ'      # マクロは「表の整理_」で始まる名前のグラフ（自分が作った物）だけを消す
        ch = co.Chart
        ch.ChartType = 65
        for i in range(int(ch.SeriesCollection().Count), 0, -1):
            ch.SeriesCollection(i).Delete()
        cats = ws.Range(ws.Cells(h, c0 + 1), ws.Cells(h, c0 + 12))
        for i in range(1, n + 1):
            s = ch.SeriesCollection().NewSeries()
            s.Name = f"='{SHEET}'!{ws.Cells(h + i, c0).Address}"
            s.Values = ws.Range(ws.Cells(h + i, c0 + 1), ws.Cells(h + i, c0 + 12))
            s.XValues = cats
        ch.HasTitle = True
        ch.ChartTitle.Text = TITLE
        ch.HasLegend = True
        ch.Legend.Position = -4107
        ax = ch.Axes(2)
        ax.HasTitle = True
        ax.AxisTitle.Text = '支出額（円）'
        ax.TickLabels.NumberFormat = '#,##0'
        ax.HasMajorGridlines = True
        ch.ChartArea.Format.TextFrame2.TextRange.Font.Name = 'Meiryo UI'
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def add_old_chart(xl, path, lay):
    wb = xl.Workbooks.Open(path)
    try:
        ws = wb.Sheets(SHEET)
        co = ws.ChartObjects().Add(ws.Cells(lay['last'] + 2, lay['c0']).Left, ws.Cells(lay['last'] + 2, lay['c0']).Top, 300, 200)
        co.Name = '表の整理_前のグラフ'      # マクロは「表の整理_」で始まる名前のグラフ（自分が作った物）だけを消す
        co.Chart.ChartType = 51
        s = co.Chart.SeriesCollection().NewSeries()
        s.Values = ws.Range(ws.Cells(lay['h'] + 1, lay['c0'] + 1), ws.Cells(lay['h'] + lay['n'], lay['c0'] + 1))
        wb.Save()
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, rows, kw in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            old = kw.pop('old_chart', False)
            lay = write_before(before, rows, **kw)
            if old:
                add_old_chart(xl, before, lay)
            build_truth(xl, before, truth, lay)
            print('  ', stem)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', build(s(1), rng.randint(2, 9)), {}),
          ('未見2_表題と日付の見出し', build(s(2), 5), {'title': True, 'r0': 2, 'c0': 2, 'date_heads': True}),
          ('未見3_合計なしと0の課', build(s(3), 6, zero=True), {'totals': False}),
          ('未見4_前のグラフ', build(s(4), 7), {'old_chart': True}),
          ('未見5_大量', build(s(5), 12), {})], out)
    return out


if __name__ == '__main__':
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        sys.exit(0)
    make([('推移_本番', build(3, 4), {}),
          ('推移_試験A', build(13, 8), {}),
          ('推移_試験B', build(23, 5), {'title': True, 'r0': 3, 'c0': 2}),
          ('推移_試験C', build(33, 4), {'totals': False}),
          ('推移_試験D', build(43, 6), {'date_heads': True}),
          ('推移_試験E', build(53, 5), {'old_chart': True}),
          ('推移_試験F', build(63, 6, zero=True), {})], HERE)
    shutil.copyfile(os.path.join(HERE, '推移_本番_正解.xlsx'), os.path.join(HERE, '推移_試験G_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, '推移_本番_正解.xlsx'), os.path.join(HERE, '推移_試験G_正解.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    lines = ['課ごとの月別推移を折れ線グラフにして', '月次集計の推移グラフを作って', '課別の月ごとの支出を折れ線で',
             '月別支出の推移を課ごとにグラフ化', '推移の折れ線グラフをお願いします', '4月から3月までの課別推移のグラフ']
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(lines) + '\n')
    print('ok')
