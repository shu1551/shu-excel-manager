# -*- coding: utf-8 -*-
"""お題の「集計表」を種で作る（B 集計の群で共有・2026-09-18）。

detail.py が「明細（1 行 1 件・日付つき）」を作るのに対して、こちらは**もう集計された小さい表**
（項目の列＋数の列 1〜2 本）を作る。累計・構成比・増減・小計行・上位 N の仕事はこの形の表に当てる。
見出しの語・シート名・項目の名前・列の並び・表題・単位の行・合計の行・空行は種で変わる（vocab.py）。

  t = design(seed, n, **kw)            … 表の設計
  lay = write_before(path, t)          … openpyxl で 1 シートのブックを書く → {'h','last','ncols','c0','total_at','note_at'}
kw: nums=1|2（数の列の数）／prev=True（前年度・今年度の 2 列）／rate=True（率の列）／text=True（摘要の列）／
    no_col・code_col（左端に番号・コードの列）／title・unit（表題・単位の行）／total=True（下に合計の行）／
    note=True（表の下に※の注記）／blanks=N（途中の空行）／swap=True（列の並びを変える）／group=N（区分の列＝N 個の組）
"""
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from vocab import Vocab   # noqa: E402

TITLES = ['令和8年度 {k}別の状況', '令和8年度 {k}別集計表', '{k}別の実績（9月末現在）']
GROUPS = {
    '区分': ['義務的経費', '投資的経費', 'その他経費'],
    '性質': ['人件費', '物件費', '補助費等', '普通建設事業費'],
    '会計': ['一般会計', '特別会計', '企業会計'],
    '分類': ['経常', '臨時', '特定'],
}
GROUP_HEADS = {'区分': ['区分', '区分名'], '性質': ['性質別区分', '性質'], '会計': ['会計区分', '会計'], '分類': ['分類', '分類名']}


def design(seed, n=8, **kw):
    rng = random.Random(f"shukei-{seed}")
    v = Vocab(seed)
    t = {'sheet': v.sheet(kw.pop('sheet_role', 'shukei')), 'seed': seed, 'n': n, 'vocab': v,
         'title': None, 'unit': False, 'total': False, 'note': None, 'blanks': 0, 'swap': False,
         'no_col': False, 'code_col': False, 'rate': False, 'text': False, 'nums': 1, 'prev': False, 'group': 0}
    want_title = kw.pop('title', False)
    t.update(kw)
    cols = []
    if t['no_col']:
        cols.append((v.word('no'), 'no'))
    if t['code_col']:
        cols.append((v.word('code'), 'code'))
    if t['group']:
        gk = rng.choice(sorted(GROUPS))
        t['group_kind'] = gk
        t['group_names'] = GROUPS[gk][:max(2, min(t['group'], len(GROUPS[gk])))]
        cols.append((rng.choice(GROUP_HEADS[gk]), 'group'))
    cols.append((v.word('ka'), 'item'))
    if t['prev']:
        cols.append((v.word('prev'), 'num'))
        cols.append((v.word('cur'), 'num'))
    elif t['nums'] >= 2:
        cols.append((v.word('yosan'), 'num'))
        cols.append((v.word('shikko'), 'num'))
    else:
        cols.append((v.word('amount'), 'num'))
    if t['rate']:
        cols.append((v.word('rate'), 'rate'))
    if t['text']:
        cols.append((v.word('tekiyo'), 'text'))
    if t['swap']:
        # 数の列を左、項目の列を右に（列の並びで決め打ちしているマクロを落とす）
        head = [c for c in cols if c[1] in ('no', 'code', 'group')]
        rest = [c for c in cols if c[1] not in ('no', 'code', 'group')]
        cols = head + rest[1:] + rest[:1]
    t['cols'] = cols
    items = v.items(n, rng)
    rows = []
    used = set()
    for i, it in enumerate(items):
        row = {}
        for name, role in cols:
            if role == 'no':
                row[name] = i + 1
            elif role == 'code':
                row[name] = f"{100 + (i + 1) * 10}"
            elif role == 'group':
                row[name] = t['group_names'][i % len(t['group_names'])]
            elif role == 'item':
                row[name] = it
            elif role == 'text':
                row[name] = rng.choice(['継続', '新規', '一部', ''])
        nums = [nm for nm, r in cols if r == 'num']
        y = _unique(rng, used)
        row[nums[0]] = y
        if len(nums) > 1:
            row[nums[1]] = _unique(rng, used, int(y * rng.uniform(0.3, 1.4) / 1000) * 1000)
        if t['rate']:
            rn = next(nm for nm, r in cols if r == 'rate')
            row[rn] = round(row[nums[-1]] / row[nums[0]], 4) if row[nums[0]] else 0
        rows.append(row)
    if t['group']:
        # 区分の列があれば、その列で並べておく（小計行はかたまりの切れ目に入れる）
        gname = next(nm for nm, r in cols if r == 'group')
        order = {g: k for k, g in enumerate(t['group_names'])}
        rows.sort(key=lambda r0: order[r0[gname]])
    t['rows'] = rows
    if want_title:
        t['title'] = rng.choice(TITLES).format(k=next(nm for nm, r in cols if r == 'item'))
    t['blank_rows'] = sorted(random.Random(seed + 13).sample(range(1, max(2, n - 1)),
                                                             min(t['blanks'], max(0, n - 2)))) if t['blanks'] else []
    return t


def _unique(rng, used, v=None):
    """同じ数を 2 度出さない（上位 N の並びが種で揺れないように）。"""
    for _ in range(400):
        x = v if v is not None else rng.randrange(100, 9000) * 1000
        if x not in used and x > 0:
            used.add(x)
            return x
        v = None
    x = max(used) + 1000
    used.add(x)
    return x


def write_into(ws, t):
    h, c0 = 1, 1
    if t['title']:
        ws.cell(1, c0, t['title']).font = Font(bold=True, size=14)
        h = 3
        if t['unit']:
            ws.cell(2, c0, '単位：円')
            h = 4
    for j, (name, _role) in enumerate(t['cols'], c0):
        ws.cell(h, j, name).font = Font(bold=True)
    r = h
    for i, row in enumerate(t['rows']):
        if i in t['blank_rows']:
            r += 1
        r += 1
        for j, (name, role) in enumerate(t['cols'], c0):
            cell = ws.cell(r, j, row.get(name))
            if role == 'num':
                cell.number_format = '#,##0'
            elif role == 'rate':
                cell.number_format = '0.0%'
    last = r
    total_at = note_at = None
    if t['total']:
        total_at = last + 1
        icol = next(j for j, (_n, role) in enumerate(t['cols'], c0) if role in ('item', 'group', 'text'))
        ws.cell(total_at, icol, '合計').font = Font(bold=True)
        for j, (name, role) in enumerate(t['cols'], c0):
            if role == 'num':
                ws.cell(total_at, j, sum(r0.get(name) or 0 for r0 in t['rows'])).number_format = '#,##0'
        r = total_at
    if t['note']:
        note_at = r + 1
        ws.cell(note_at, c0, '※金額は税込み')
    return {'h': h, 'last': last, 'ncols': len(t['cols']), 'c0': c0, 'total_at': total_at, 'note_at': note_at}


def write_before(path, t):
    wb = Workbook()
    ws = wb.active
    ws.title = t['sheet']
    lay = write_into(ws, t)
    wb.save(path)
    return lay


def col_of(t, role, nth=0):
    """役の列番号（1 から）。nth＝同じ役が複数あるときの何番目か。"""
    hit = [j for j, (_n, r) in enumerate(t['cols'], 1) if r == role]
    return hit[nth]


def name_of(t, role, nth=0):
    hit = [n for n, r in t['cols'] if r == role]
    return hit[nth]


if __name__ == '__main__':
    for s in (1, 2, 3):
        t = design(s, 6, nums=2, rate=True, title=True, total=True)
        print(s, t['sheet'], [c for c, _r in t['cols']], t['rows'][0])
