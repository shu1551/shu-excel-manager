# -*- coding: utf-8 -*-
"""選んでいる列で、値で上書きされたり形の違う式になったりしたセルを、上下と同じ式に戻す（2026-09-18 第二期）。

列の決まり: 選んでいる 1 列＝式の列。<名前>_選択.txt に書く。表は集計表（shukei）に式の列（差額・率・税込）を足したもの。
正解の決まり（依頼文に全部書く）:
  - 見出しの行＝選んでいるセルの行。本文＝見出しの次の行から最後の本文の行まで（空の行と「合計」「計」「総計」の行は触らない）。
  - 本文の式を R1C1 形式で数え、いちばん多い式を「その列の式」とする。本文で、その式でないセル（値・形の違う式）に、その式を入れる（FormulaR1C1）。
  - 空のセル・ほかの列は触らない。式の無い列なら何もしない。
組: 本番 差額に値 2 つ／A 40 行・5 つ／B 表題と単位／C 合計の行の SUM は触らない／D 途中の空行／E 税込（×1.1）の列／F 撃った後にもう一度／
  G 行のずれた式（=C5-B4）／H 途切れなし
"""
import os
import random
import shutil
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import shukei   # noqa: E402

REQUEST = ('選んでいる列で、式が途切れているセルを、上下と同じ式に戻してください（上書きしてよい）。'
           '見出しの行は必ず「選んでいるセルの行」＝Selection.Cells(1,1).Row をそのまま使う（自分で見出しの行を探さない）。'
           '本文は見出しの次の行から最後の本文の行まで。行がまるごと空の行と、行のどこかに「合計」「計」「総計」とある行は触らない。'
           '本文の式を R1C1 形式（FormulaR1C1）で数え、いちばん多い式をその列の式とする。本文で、その式でないセル（値で上書きされたセル・形の違う式）に、'
           'その式を FormulaR1C1 で入れる。空のセルとほかの列は触らない。式が 1 つも無い列なら何もしない。')
PHRASES = ['選んだ列の式の途切れを直して', '値で上書きされた所を式に戻して', '計算式が崩れている所を直して',
           '上下と同じ式に戻して', '式が抜けているセルを埋め直して', '選んでいる列の数式をそろえて']


def add_formula_col(path, t, lay, kind='diff', breaks=2, shift=False):
    wb = load_workbook(path)
    ws = wb.worksheets[0]
    h, last, c0 = lay['h'], lay['last'], lay['c0']
    a = get_column_letter(c0 + shukei.col_of(t, 'num', 0) - 1)
    b = get_column_letter(c0 + shukei.col_of(t, 'num', min(1, sum(1 for _n, r in t['cols'] if r == 'num') - 1)) - 1)
    fc = c0 + lay['ncols']
    ws.cell(h, fc, '差額' if kind == 'diff' else '税込')
    rng = random.Random(t['seed'] + 21)
    body = [r for r in range(h + 1, last + 1) if ws.cell(r, c0 + shukei.col_of(t, 'num', 0) - 1).value is not None]
    broken = set(rng.sample(body, min(breaks, max(0, len(body) - 2)))) if breaks else set()
    for r in body:
        f = f"={b}{r}-{a}{r}" if kind == 'diff' else f"=ROUND({a}{r}*1.1,0)"
        if r in broken:
            ws.cell(r, fc, rng.randrange(1, 900) * 1000)
        else:
            ws.cell(r, fc, f)
        ws.cell(r, fc).number_format = '#,##0'
    if shift and len(body) > 4:
        r = body[3]
        ws.cell(r, fc, f"={b}{r}-{a}{r - 1}")              # 1 行ずれた式（形の違う式）
    if t['total'] and lay.get('total_at'):
        L = get_column_letter(fc)
        ws.cell(lay['total_at'], fc, f"=SUM({L}{h + 1}:{L}{last})").number_format = '#,##0'
    wb.save(path)
    return fc, body


def build_truth(xl, before, truth, t, lay, fc, body):
    from collections import Counter
    wb = xl.Workbooks.Open(before)
    try:
        ws = wb.Worksheets(t['sheet'])
        fs = [str(ws.Cells(r, fc).FormulaR1C1) for r in body if ws.Cells(r, fc).HasFormula]
        if fs:
            main = Counter(fs).most_common(1)[0][0]
            for r in body:
                c = ws.Cells(r, fc)
                if c.Value is None:
                    continue
                if not c.HasFormula or str(c.FormulaR1C1) != main:
                    c.FormulaR1C1 = main
        wb.SaveAs(truth, FileFormat=51)
    finally:
        wb.Close(SaveChanges=False)


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t, kind, breaks, shift in jobs:
            before, truth = os.path.join(out, f"{stem}_前.xlsx"), os.path.join(out, f"{stem}_正解.xlsx")
            lay = shukei.write_before(before, t)
            fc, body = add_formula_col(before, t, lay, kind, breaks, shift)
            build_truth(xl, before, truth, t, lay, fc, body)
            sel = f"{get_column_letter(fc)}{lay['h']}"
            open(os.path.join(out, f"{stem}_選択.txt"), 'w', encoding='utf-8').write(sel + '\n')
            print('  ', stem, t['sheet'], [n for n, _r in t['cols']], '選ぶ列', sel, '途切れ', breaks, 'ずれ' if shift else '')
    finally:
        xl = None
        vf._quit(pid)


def d(seed, n, **kw):
    kw.setdefault('nums', 2)
    return shukei.design(seed, n, **kw)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', d(s(1), rng.randint(6, 14)), 'diff', 2, False),
          ('未見2_表題と合計', d(s(2), 9, title=True, unit=True, total=True), 'diff', 3, True),
          ('未見3_税込', d(s(3), 8, nums=1), 'tax', 2, False),
          ('未見4_空行', d(s(4), 10, blanks=2), 'diff', 1, True)], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('途切れ_本番', d(4101, 9), 'diff', 2, False),
          ('途切れ_試験A', d(4102, 40), 'diff', 5, False),
          ('途切れ_試験B', d(4103, 8, title=True, unit=True), 'diff', 2, False),
          ('途切れ_試験C', d(4104, 9, total=True), 'diff', 2, False),
          ('途切れ_試験D', d(4105, 10, blanks=2), 'diff', 2, False),
          ('途切れ_試験E', d(4106, 8, nums=1), 'tax', 2, False),
          ('途切れ_試験G', d(4107, 9), 'diff', 0, True),
          ('途切れ_試験H', d(4108, 7), 'diff', 0, False)], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '途切れ_本番_正解.xlsx'), os.path.join(HERE, f'途切れ_試験F_{x}.xlsx'))
    shutil.copyfile(os.path.join(HERE, '途切れ_本番_選択.txt'), os.path.join(HERE, '途切れ_試験F_選択.txt'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
