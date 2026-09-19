# -*- coding: utf-8 -*-
"""パワークエリで月別のクロス集計表を作る（どの表でも動く形・2026-09-18 作り直し。前の版は PQ課別月別）。

列の決まり: 選んでいる 3 列（1 つ目＝行、2 つ目＝日付、3 つ目＝値）。選んでいる列は <名前>_選択.txt に書く。
作る物の名前は元のテーブル名から（「（元のテーブル名）_クロス」）。見出しの語・シート名・テーブル名は表ごとに違う。
正解の決まり:
  - 行＝1 つ目の列（昇順）、列＝2 つ目（日付）の月（年度の順に 4月から3月・明細にある月だけ・見出しは「4月」の形）、
    値＝3 つ目の列の合計。明細に無い月の列は作らない。
  - 新しいシート（名前は 元のテーブル名 と _クロス をつないだ名前）の A1 に同じ名前のテーブルとして読み込む。クエリ名も同じ。
  - 同じ名前のクエリやシートが既にあれば消して作り直す。
組: 本番 60 行／A 600 行／B 1 か月だけ／C 年度をまたぐ（1〜3 月あり）／D 前のクエリとシート／E 撃った後にもう一度／
  F 列の並びが違う／G 途中に空行・番号の列
"""
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
import detail   # noqa: E402
import importlib.util   # noqa: E402

_spec = importlib.util.spec_from_file_location('mp_tate', os.path.join(HERE, '..', '縦持ち', 'make_pairs.py'))
T = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(T)

REQUEST = ('パワークエリで、選んでいる 3 列（1 つ目を行、2 つ目＝日付の列の月を列、3 つ目を値）から'
           '月別のクロス集計表を作ってください。行は 1 つ目の列（昇順）、列は月（年度の順に 4月から3月・'
           '明細にある月だけ・見出しは「4月」の形）、値は 3 つ目の列の合計。'
           '新しいシート（名前は 元のテーブル名 と _クロス をつないだ名前）の A1 に同じ名前のテーブルとして読み込む。'
           'クエリ名も同じ。同じ名前のクエリやシートが既にあれば消して作り直す')
PHRASES = ['パワークエリで月別のクロス集計を', 'クエリで行と月のクロス表を作って', '月を列にしたクロス集計をクエリで',
           'パワークエリでクロス集計表を', 'クエリで月別の表にして', '選んだ列でクロス集計のクエリを']


def design(seed, n, **kw):
    t = detail.design(seed, n, **kw)
    v = t['vocab']
    t['table'] = v.table('meisai')
    # M の列参照 [名前] は全角の括弧で壊れる（2026-09-18: 「金額（円）」で Expression.Error）＝括弧なしの見出しにする
    ren = {name: name.split('（')[0] for name, _r in t['cols'] if '（' in name}
    if ren:
        t['cols'] = [(ren.get(name, name), role) for name, role in t['cols']]
        t['rows'] = [{ren.get(k, k): v2 for k, v2 in row.items()} for row in t['rows']]
    return t


def write_before(path, t):
    wb = Workbook()
    ws = wb.active
    ws.title = t['sheet']
    lay = detail.write_into(ws, t)
    tb = Table(displayName=t['table'],
               ref=f"A{lay['h']}:{get_column_letter(lay['ncols'])}{lay['last']}")
    tb.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(tb)
    wb.save(path)
    return lay


def m_formula(t):
    item = detail.name_of(t, 'item')
    date = detail.name_of(t, 'date')
    num = detail.name_of(t, 'num')
    return f'''let
    Source = Excel.CurrentWorkbook(){{[Name="{t['table']}"]}}[Content],
    Typed = Table.TransformColumnTypes(Source, {{{{"{date}", type date}}, {{"{item}", type text}}, {{"{num}", Int64.Type}}}}),
    AddM = Table.AddColumn(Typed, "月", each Date.Month([{date}]), Int64.Type),
    Grouped = Table.Group(AddM, {{"{item}", "月"}}, {{{{"{num}", each List.Sum([{num}]), Int64.Type}}}}),
    Order = List.Select({{4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3}}, each List.Contains(Grouped[月], _)),
    Named = Table.TransformColumns(Grouped, {{{{"月", each Text.From(_) & "月", type text}}}}),
    Pivoted = Table.Pivot(Named, List.Transform(Order, each Text.From(_) & "月"), "月", "{num}", List.Sum),
    Sorted = Table.Sort(Pivoted, {{{{"{item}", Order.Ascending}}}})
in
    Sorted'''


def make(jobs, out=HERE):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t, kw in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            lay = write_before(before, t)
            name = t['table'] + '_クロス'
            sel = ",".join(f"{get_column_letter(detail.col_of(t, role))}{lay['h']}"
                           for role in ('item', 'date', 'num'))
            if kw.get('stale'):
                wb = xl.Workbooks.Open(before)
                try:
                    T.load_query(wb, name, m_formula(t).replace('Order.Ascending', 'Order.Descending'), name)
                    wb.Worksheets(t['sheet']).Activate()
                    wb.Save()
                finally:
                    wb.Close(SaveChanges=False)
            wb = xl.Workbooks.Open(before)
            try:
                T.load_query(wb, name, m_formula(t), name)
                wb.Worksheets(t['sheet']).Activate()
                wb.SaveAs(truth, FileFormat=51)
            finally:
                wb.Close(SaveChanges=False)
            open(os.path.join(out, f"{stem}_選択.txt"), 'w', encoding='utf-8').write(sel + '\n')
            print('  ', stem, t['sheet'], t['table'], name, '選ぶ列', sel)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', design(s(1), rng.randint(30, 150)), {}),
          ('未見2_1か月', design(s(2), 25, gap_months=(4, 5, 6, 7, 8, 9, 10, 11, 1, 2, 3)), {}),
          ('未見3_並び違いと前のクエリ', design(s(3), 50, swap=True), {'stale': True}),
          ('未見4_空行と番号', design(s(4), 40, blanks=2, no_col=True), {})], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('クロス_本番', design(3, 60), {}),
          ('クロス_試験A', design(13, 600, n_items=8), {}),
          ('クロス_試験B', design(23, 25, gap_months=(5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3)), {}),
          ('クロス_試験C', design(33, 50, gap_months=(6, 7, 8)), {}),
          ('クロス_試験D', design(43, 40), {'stale': True}),
          ('クロス_試験F', design(53, 45, swap=True), {}),
          ('クロス_試験G', design(63, 35, blanks=2, no_col=True), {})])
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, 'クロス_本番_正解.xlsx'), os.path.join(HERE, f'クロス_試験E_{x}.xlsx'))
    shutil.copyfile(os.path.join(HERE, 'クロス_本番_選択.txt'), os.path.join(HERE, 'クロス_試験E_選択.txt'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
