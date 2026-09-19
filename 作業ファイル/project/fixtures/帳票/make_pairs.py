# -*- coding: utf-8 -*-
"""帳票（結合セルの受付簿）を 1 行 1 件の一覧に直す（鍛える回路の題材・2026-09-17。印刷用の様式で受け付けた申請を集計できる形にする仕事）。

仕事: シート「受付簿」の帳票の右に 1 列空けて、1 行 1 件の一覧を作る。元の帳票は触らない。
帳票の作り:
  - 見出しは 2 段。「受付番号」「申請額」「受付日」は 2 段の結合、「申請者」は上段で、下段が「団体名」「代表者」。
  - 1 件は 2 行か 3 行。1 行目に 受付番号・団体名・代表者・申請額・受付日（受付番号・申請額・受付日は件の行数ぶん縦に結合）。
    2 行目から「所在地」「電話」のラベルと値（ラベルは団体名の列・値は代表者の列）。
正解の決まり:
  - 一覧の見出しは 1 行。2 段の見出しは下段の名前（団体名・代表者）、下段が無い所は上段の名前、件の中のラベル（所在地・電話）も見出し。
    並びは帳票に出る順（件の 1 行目の左から右、次の行）。
  - 一覧の左上＝帳票の見出しの上段の行・帳票の右端の列の 2 つ右（1 列空ける）。
  - 値は帳票のまま（受付番号の先頭の 0 を落とさない・日付は日付）。値の無い欄（電話が空）は空のまま。
  - 入れない: 受付番号の空いた枠（印刷用の空の枠）・※で始まる注記・改ページで繰り返した見出し。
  - 一覧の場所に前の一覧が残っていたら消してから書く（2 回撃っても同じ）。
組（依頼文は同じ）: 本番 8 件／A 30 件／B 1 件 3 行（電話・空の電話あり）／C 表題の行なし／D 空の枠と※注記／
  E 前の一覧（行が多い）が残っている／F 列の並びが違う／G 改ページで見出しを繰り返す／H 受付番号が先頭 0 の文字／I 撃った後にもう一度
  py make_pairs.py            → *_前.xlsx と *_正解.xlsx
  py make_pairs.py --unseen N → 未見_種N
"""
import datetime
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

HERE = os.path.dirname(os.path.abspath(__file__))
SHEET = '受付簿'
DANTAI = ['上町自治会', '下町子ども会', '本町商店会', '緑ヶ丘老人クラブ', '川口地区防災会', '中央公園愛護会', '北山婦人会',
          '南台スポーツ少年団', '旭町青年会', '若葉保存会', '東野農業者組合', '桜通り町内会', '西浜漁業者の会', 'みどり子育てサークル',
          '青葉文化協会', '栄町交通安全協会', '大沢集落営農組合', '山手ボランティア会', '駅前にぎわい会', '森の学校運営委員会']
SEI = ['佐藤', '鈴木', '高橋', '田中', '伊藤', '渡辺', '山本', '中村', '小林', '加藤', '吉田', '山田', '佐々木', '斎藤', '松本']
MEI = ['一郎', '花子', '誠', '恵', '学', '悟', '由美', '健太', '真理', '修', '陽子', '大輔']
MACHI = ['本町', '上町', '下町', '旭町', '栄町', '緑町', '青葉台', '桜木', '川原', '新田']
THIN = Side(style='thin')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build(seed, n, phone=False, zero_id=False, blank_phone=False):
    rng = random.Random(seed)
    recs = []
    day = datetime.date(2026, 4, 1) + datetime.timedelta(days=rng.randrange(0, 20))
    for i in range(n):
        no = f"{i + 1:04d}" if zero_id else f"R8-{i + 1:03d}"
        dantai = rng.choice(DANTAI) if n > len(DANTAI) else DANTAI[(seed + i * 7) % len(DANTAI)]
        rep = rng.choice(SEI) + ' ' + rng.choice(MEI)
        amt = rng.randrange(5, 300) * 1000
        day = day + datetime.timedelta(days=rng.randrange(0, 3))
        addr = f"架空市{rng.choice(MACHI)}{rng.randrange(1, 9)}-{rng.randrange(1, 30)}"
        rec = {'受付番号': no, '団体名': dantai, '代表者': rep, '申請額': amt,
               '受付日': datetime.datetime(day.year, day.month, day.day), '所在地': addr}
        if phone:
            rec['電話'] = '' if (blank_phone and rng.random() < 0.3) else f"0186-{rng.randrange(10, 99)}-{rng.randrange(1000, 9999)}"
        recs.append(rec)
    return recs


def _head(ws, r, layout, first_col=1):
    """見出し 2 段を r・r+1 行に書く（layout＝上段の並び。'申請者' は 2 列ぶん）。"""
    c = first_col
    for name in layout:
        if name == '申請者':
            ws.cell(r, c, '申請者')
            ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 1)
            ws.cell(r + 1, c, '団体名')
            ws.cell(r + 1, c + 1, '代表者')
            for cc in (c, c + 1):
                for rr in (r, r + 1):
                    ws.cell(rr, cc).border = BOX
                    ws.cell(rr, cc).font = Font(bold=True)
                    ws.cell(rr, cc).alignment = Alignment(horizontal='center')
            c += 2
        else:
            ws.cell(r, c, name)
            ws.merge_cells(start_row=r, start_column=c, end_row=r + 1, end_column=c)
            for rr in (r, r + 1):
                ws.cell(rr, c).border = BOX
                ws.cell(rr, c).font = Font(bold=True)
                ws.cell(r, c).alignment = Alignment(horizontal='center', vertical='center')
            c += 1
    return c - 1                                           # 帳票の右端の列


def _columns(layout):
    """上段の並び → {項目: 列の位置（1 始まり）}、ラベルの列・値の列。"""
    pos, c = {}, 1
    for name in layout:
        if name == '申請者':
            pos['団体名'], pos['代表者'] = c, c + 1
            c += 2
        else:
            pos[name] = c
            c += 1
    return pos, pos['団体名'], pos['代表者']


def list_fields(layout, labels):
    pos, _l, _v = _columns(layout)
    top = sorted(pos, key=lambda k: pos[k])
    return top + list(labels)


def write_pair(stem, recs, out=HERE, layout=('受付番号', '申請者', '申請額', '受付日'), title=True, empty_frames=0,
               note=False, stale=0, page=0):
    labels = ['所在地'] + (['電話'] if recs and '電話' in recs[0] else [])
    k = 1 + len(labels)
    pos, lab_c, val_c = _columns(layout)
    fields = list_fields(layout, labels)
    for truth in (False, True):
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET
        h = 1
        last_col = len(layout) + 1
        if title:
            ws.cell(1, 1, '令和8年度 地域づくり活動補助金 申請受付簿').font = Font(bold=True, size=14)
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
            h = 2
        _head(ws, h, layout)
        r = h + 2
        frames = [dict(x) for x in recs] + [None] * empty_frames
        for i, rec in enumerate(frames):
            if page and i and i % page == 0:
                _head(ws, r, layout)                       # 改ページで見出しを繰り返す
                r += 2
            for name, c in pos.items():
                v = rec.get(name, '') if rec else ''
                if v != '':
                    ws.cell(r, c, v)
                if name in ('団体名', '代表者'):
                    continue
                ws.merge_cells(start_row=r, start_column=c, end_row=r + k - 1, end_column=c)
            for j, lab in enumerate(labels, 1):
                ws.cell(r + j, lab_c, lab)
                v = rec.get(lab, '') if rec else ''
                if v != '':
                    ws.cell(r + j, val_c, v)
            for rr in range(r, r + k):
                for cc in range(1, last_col + 1):
                    ws.cell(rr, cc).border = BOX
            r += k
        if note:
            ws.cell(r, 1, '※ 受付番号は受付順に付けています。申請額は千円未満を切り捨てています。')
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last_col)
        lc = last_col + 2
        if truth:
            for j, f in enumerate(fields):
                ws.cell(h, lc + j, f)
            for i, rec in enumerate(recs, 1):
                for j, f in enumerate(fields):
                    v = rec.get(f, '')
                    if v != '':
                        ws.cell(h + i, lc + j, v)
        elif stale:
            for j, f in enumerate(fields):
                ws.cell(h, lc + j, f)
            for i in range(1, stale + 1):
                ws.cell(h + i, lc, f"旧-{i:03d}")
                ws.cell(h + i, lc + 1, f"旧団体{i}")
                ws.cell(h + i, lc + 3, 1000 * i)
        for c in range(1, last_col + 1):
            ws.column_dimensions[chr(64 + c)].width = 16
        wb.save(os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx"))


SWAP = ('受付日', '受付番号', '申請額', '申請者')


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    write_pair('未見1_標準', build(s(1), rng.randint(3, 15)), out=out)
    write_pair('未見2_3行と表題なし', build(s(2), rng.randint(4, 12), phone=True, blank_phone=True), title=False, out=out)
    write_pair('未見3_空の枠と注記', build(s(3), rng.randint(2, 9)), empty_frames=rng.randint(2, 6), note=True, out=out)
    write_pair('未見4_前の一覧', build(s(4), rng.randint(2, 6)), stale=rng.randint(12, 30), out=out)
    write_pair('未見5_並び違い', build(s(5), rng.randint(3, 10), phone=True), layout=SWAP, out=out)
    write_pair('未見6_改ページ', build(s(6), rng.randint(15, 40)), page=rng.choice((5, 8, 10)), note=True, out=out)
    write_pair('未見7_全部入り', build(s(7), 25, phone=True, zero_id=True, blank_phone=True), layout=SWAP, title=False,
               empty_frames=3, note=True, stale=40, page=10, out=out)
    write_pair('未見8_大量', build(s(8), 300, zero_id=True), page=20, out=out)
    return out


if __name__ == '__main__':
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        sys.exit(0)
    write_pair('帳票_本番', build(5, 8))
    write_pair('帳票_試験A', build(15, 30))
    write_pair('帳票_試験B', build(25, 6, phone=True, blank_phone=True))
    write_pair('帳票_試験C', build(35, 5), title=False)
    write_pair('帳票_試験D', build(45, 4), empty_frames=5, note=True)
    write_pair('帳票_試験E', build(55, 3), stale=20)
    write_pair('帳票_試験F', build(65, 7, phone=True), layout=SWAP)
    write_pair('帳票_試験G', build(75, 23), page=10)
    write_pair('帳票_試験H', build(85, 9, zero_id=True))
    shutil.copyfile(os.path.join(HERE, '帳票_本番_正解.xlsx'), os.path.join(HERE, '帳票_試験I_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, '帳票_本番_正解.xlsx'), os.path.join(HERE, '帳票_試験I_正解.xlsx'))
    print('ok')
