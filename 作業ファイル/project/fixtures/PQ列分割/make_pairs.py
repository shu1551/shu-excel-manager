# -*- coding: utf-8 -*-
"""パワークエリで、選んでいる列を区切り文字で複数の列に分ける（G パワークエリ 57・2026-09-18 第二期）。

列の決まり: 選んでいる 1 列（テーブルの列の見出しのセル）。<名前>_選択.txt に書く。
正解の決まり（依頼文に全部書く）:
  - 元＝アクティブなシートのテーブル（選んでいるセルのテーブル）。区切り文字＝その列の値に出てくる「・」「／」「/」「、」「-」「 」（半角の空白）のうち、
    いちばん多くの値に出てくるもの（同じならこの順で先）。
  - 分けた列の名前＝元の見出し＋ 1・2・…（いちばん多く分かれた数まで）。元の列の位置に並べる。ほかの列はそのまま。型は文字。
  - クエリ名・シート名・テーブル名＝元のテーブル名＋「_分割」。新しいシートの A1 に読み込む。同じ名前のクエリ・シートがあれば消して作り直す。
組: 本番 所属「課・係」／A 氏名「姓 名」／B 住所「市/町」／C 3 つに分かれる値あり／D 区切りの無い値あり／E 列の並び違い／F 撃った後にもう一度／
  G 番号の列と空欄／H 前のクエリとシート
"""
import importlib.util
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..')))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
_spec = importlib.util.spec_from_file_location('tate', os.path.join(HERE, '..', '縦持ち', 'make_pairs.py'))
tate = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(tate)

DELIMS = ['・', '／', '/', '、', '-', ' ']
REQUEST = ('パワークエリで、選んでいる列を区切り文字で複数の列に分けてください。'
           '元＝アクティブなシートの、選んでいるセルがあるテーブル（Selection.Cells(1,1).ListObject）。分ける列＝選んでいるセルの列（見出しの語はそのまま使う）。'
           '区切り文字＝その列の値に出てくる「・」「／」「/」「、」「-」「 」（半角の空白）のうち、いちばん多くの値に出てくるもの（同じ数ならこの順で先のもの）。'
           '分けた列の名前＝元の見出しの後ろに 1・2・3…（いちばん多く分かれた値の数まで）。元の列の位置に並べ、ほかの列はそのまま。分けた列の型は文字。'
           'クエリの名前・新しいシートの名前・読み込むテーブルの名前は、元のテーブル名の後ろに「_分割」。新しいシートの A1 に読み込む。'
           '同じ名前のクエリやシートが既にあれば消して作り直す。M の式は Table.SplitColumn と Splitter.SplitTextByDelimiter で書く。')
PHRASES = ['パワークエリで列を区切り文字で分割して', 'クエリで選んだ列を分けて', '課と係をパワークエリで別の列に',
           '選んでいる列を区切りで分割するクエリを', 'パワークエリで氏名を姓と名に分けて', '区切り文字で列を分けるクエリを作って']
KAKARI = ['庶務係', '企画係', '経理係', '管理係', '指導係']
SEI = ['佐藤', '鈴木', '高橋', '田中', '伊藤', '小松', '工藤', '三浦']
MEI = ['一郎', '花子', '健', '美咲', '大輔', '由美']
CITY = ['秋田市', '能代市', '横手市', '大館市', '男鹿市']
TOWN = ['山王', '中通', '旭北', '桜', '緑町']


def build(seed, n, kind='ka', **kw):
    rng = random.Random(seed)
    from vocab import Vocab
    v = Vocab(seed)
    kas = v.items(4, rng)
    vals = []
    for i in range(n):
        if kind == 'ka':
            x = f"{rng.choice(kas)}・{rng.choice(KAKARI)}"
        elif kind == 'name':
            x = f"{rng.choice(SEI)} {rng.choice(MEI)}"
        else:
            x = f"{rng.choice(CITY)}/{rng.choice(TOWN)}"
        vals.append(x)
    t = {'sheet': v.sheet('meisai'), 'table': rng.choice(['T一覧', 'T名簿', 'T明細', 'Tデータ']) + str(seed % 7),
         'split': {'ka': '所属', 'name': '氏名', 'addr': '住所'}[kind], 'no': v.word('no'), 'amount': v.word('amount').replace('（円）', ''),
         'vals': vals, 'swap': False, 'no_col': True, 'stale': False, 'blanks': False}
    t.update(kw)
    if t.get('three'):
        d = '・' if kind == 'ka' else (' ' if kind == 'name' else '/')
        t['vals'][1] = t['vals'][1] + d + '第2班'
    if t.get('plain'):
        t['vals'][2] = t['vals'][2].replace('・', '').replace(' ', '').replace('/', '')
    t['amounts'] = [rng.randrange(10, 900) * 100 for _ in range(n)]
    return t


def delim_of(vals):
    counts = [(sum(1 for x in vals if x and d in x), -i, d) for i, d in enumerate(DELIMS)]
    return max(counts)[2]


def write_before(path, t):
    wb = Workbook()
    ws = wb.active
    ws.title = t['sheet']
    cols = ([t['no']] if t['no_col'] else []) + ([t['amount'], t['split']] if t['swap'] else [t['split'], t['amount']])
    for j, c in enumerate(cols, 1):
        ws.cell(1, j, c).font = Font(bold=True)
    for i, (x, a) in enumerate(zip(t['vals'], t['amounts']), 2):
        row = {t['no']: 100 + i, t['split']: (None if t['blanks'] and i == 4 else x), t['amount']: a}
        for j, c in enumerate(cols, 1):
            if row.get(c) is not None:
                ws.cell(i, j, row[c])
    tb = Table(displayName=t['table'], ref=f"A1:{get_column_letter(len(cols))}{len(t['vals']) + 1}")
    tb.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(tb)
    wb.save(path)
    return cols


def m_formula(t, cols):
    vals = [x for x in t['vals'] if x]
    if t['blanks']:
        vals = [x for i, x in enumerate(t['vals'], 2) if i != 4]
    d = delim_of(vals)
    n = max(len(x.split(d)) for x in vals)
    names = ", ".join(f'"{t["split"]}{k}"' for k in range(1, n + 1))
    types = ", ".join(f'{{"{t["split"]}{k}", type text}}' for k in range(1, n + 1))
    return f'''let
    Source = Excel.CurrentWorkbook(){{[Name="{t['table']}"]}}[Content],
    Split = Table.SplitColumn(Source, "{t['split']}", Splitter.SplitTextByDelimiter("{d}", QuoteStyle.None), {{{names}}}),
    Typed = Table.TransformColumnTypes(Split, {{{types}}})
in
    Typed'''


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, t in jobs:
            before, truth = os.path.join(out, f"{stem}_前.xlsx"), os.path.join(out, f"{stem}_正解.xlsx")
            cols = write_before(before, t)
            name = t['table'] + '_分割'
            if t['stale']:
                wb = xl.Workbooks.Open(before)
                try:
                    tate.load_query(wb, name, f'let Source = Excel.CurrentWorkbook(){{[Name="{t["table"]}"]}}[Content] in Source', name)
                    wb.Worksheets(1).Activate()
                    wb.Save()
                finally:
                    wb.Close(SaveChanges=False)
            wb = xl.Workbooks.Open(before)
            try:
                tate.load_query(wb, name, m_formula(t, cols), name)
                wb.Worksheets(1).Activate()
                wb.SaveAs(truth, FileFormat=51)
            finally:
                wb.Close(SaveChanges=False)
            sel = f"{get_column_letter(cols.index(t['split']) + 1)}1"
            open(os.path.join(out, f"{stem}_選択.txt"), 'w', encoding='utf-8').write(sel + '\n')
            print('  ', stem, t['sheet'], t['table'], cols, '選ぶ列', sel, '区切り', repr(delim_of([x for x in t['vals'] if x])))
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    o = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(o, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_課係', build(s(1), rng.randint(5, 15), 'ka')),
          ('未見2_氏名と3つ', build(s(2), 8, 'name', three=True)),
          ('未見3_住所と並び違い', build(s(3), 7, 'addr', swap=True, plain=True)),
          ('未見4_空欄と前のクエリ', build(s(4), 6, 'ka', blanks=True, stale=True))], o)
    return o


def main():
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        return
    make([('分割Q_本番', build(3701, 8, 'ka')),
          ('分割Q_試験A', build(3702, 30, 'name')),
          ('分割Q_試験B', build(3703, 8, 'addr')),
          ('分割Q_試験C', build(3704, 8, 'ka', three=True)),
          ('分割Q_試験D', build(3705, 8, 'name', plain=True)),
          ('分割Q_試験E', build(3706, 7, 'ka', swap=True)),
          ('分割Q_試験G', build(3707, 7, 'addr', blanks=True)),
          ('分割Q_試験H', build(3708, 6, 'ka', stale=True))], HERE)
    for x in ('前', '正解'):
        shutil.copyfile(os.path.join(HERE, '分割Q_本番_正解.xlsx'), os.path.join(HERE, f'分割Q_試験F_{x}.xlsx'))
    shutil.copyfile(os.path.join(HERE, '分割Q_本番_選択.txt'), os.path.join(HERE, '分割Q_試験F_選択.txt'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(PHRASES) + '\n')
    print('ok')


if __name__ == '__main__':
    main()
