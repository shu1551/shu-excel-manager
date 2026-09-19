# -*- coding: utf-8 -*-
"""月ごとの支出明細をパワークエリで追加して課別に集計する（鍛える回路の題材・パワークエリ 1 本目・2026-09-17 深夜）。

仕事: 名前が「T支出」で始まるテーブル（月ごとのシートにある支出明細）を全部追加して 1 つにし、課名ごとの支出額の合計を
      新しいシート「PQ課別集計」の A1 にテーブル「Q課別集計」として読み込む。クエリ名も Q課別集計。
正解の決まり（依頼文に全部書く）:
  - 列は 課名・支出額（合計）、課名の昇順。
  - 同じ名前のクエリやシートが既にあれば消して作り直す（2 回撃ってもクエリは 1 つ・読み込み先の表も 1 つ）。
  - 追加するのは名前が T支出 で始まるテーブルだけ（ほかのテーブルは入れない）。
正解のブックは、直す前のブックを自分の Excel（非表示）で開き、クエリを作って読み込み、保存する（人の Excel に触らない）。
組: 本番 2 か月／A 3 か月・大量／B 列の並びが違う月がある／C 前のクエリとシートが残っている（中身が古い）／
  D T支出 で始まらないテーブルも同じブックにある（入れない）／E 撃った後にもう一度
"""
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.abspath(os.path.join(HERE, '..', '..', 'python_scripts')))
KA = ['総務課', '財政課', '企画課', '税務課', '福祉課', '建設課', '農林課', '教育総務課', '観光課', '環境課']
KAMOKU = ['需用費', '役務費', '委託料', '旅費', '備品購入費']
QNAME = 'Q課別集計'
OUT_SHEET = 'PQ課別集計'
M = '''let
    Source = Excel.CurrentWorkbook(),
    Rows = Table.SelectRows(Source, each Text.StartsWith([Name], "T支出")),
    Combined = Table.Combine(Rows[Content]),
    Typed = Table.TransformColumnTypes(Combined, {{"課名", type text}, {"支出額", Int64.Type}}),
    Grouped = Table.Group(Typed, {"課名"}, {{"支出額", each List.Sum([支出額]), Int64.Type}}),
    Sorted = Table.Sort(Grouped, {{"課名", Order.Ascending}})
in
    Sorted'''
REQUEST = ('パワークエリで、名前が T支出 で始まるテーブルを全部追加して1つにし、課名ごとの支出額の合計（列は課名・支出額、課名の昇順）を、'
           '新しいシート「PQ課別集計」の A1 にテーブル「Q課別集計」として読み込んでください。クエリ名は Q課別集計。'
           '同じ名前のクエリやシートが既にあれば消して作り直す。T支出 で始まらないテーブルは入れない')


def build(seed, months, n):
    rng = random.Random(seed)
    kas = rng.sample(KA, rng.randint(4, 8))
    return {m: [{'課名': rng.choice(kas), '科目': rng.choice(KAMOKU), '支出額': rng.randrange(1, 900) * 100} for _ in range(n)]
            for m in months}


def write_before(path, data, swap_month=None, other_table=False):
    wb = Workbook()
    wb.remove(wb.active)
    for m, rows in data.items():
        ws = wb.create_sheet(f"{m}月")
        cols = ['支出額', '課名', '科目'] if m == swap_month else ['課名', '科目', '支出額']
        for j, c in enumerate(cols, 1):
            ws.cell(1, j, c).font = Font(bold=True)
        for i, r in enumerate(rows, 2):
            for j, c in enumerate(cols, 1):
                ws.cell(i, j, r[c])
        t = Table(displayName=f"T支出{m}月", ref=f"A1:C{len(rows) + 1}")
        t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(t)
    if other_table:
        ws = wb.create_sheet('予算')
        ws.cell(1, 1, '課名')
        ws.cell(1, 2, '支出額')
        for i, k in enumerate(KA[:4], 2):
            ws.cell(i, 1, k)
            ws.cell(i, 2, 999999)
        t = Table(displayName='T予算', ref='A1:B5')
        t.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
        ws.add_table(t)
    wb.save(path)


def load_query(wb, name=QNAME, formula=M, sheet=OUT_SHEET):
    """クエリを作り（同じ名前があれば消す）、シート sheet の A1 に表として読み込む（COM）。"""
    xl = wb.Application
    for i in range(int(wb.Queries.Count), 0, -1):
        if str(wb.Queries.Item(i).Name) == name:
            wb.Queries.Item(i).Delete()
    for sh in list(wb.Worksheets):
        if str(sh.Name) == sheet:
            xl.DisplayAlerts = False
            sh.Delete()
    wb.Queries.Add(name, formula)
    ws = wb.Worksheets.Add(After=wb.Worksheets(wb.Worksheets.Count))
    ws.Name = sheet
    conn = f"OLEDB;Provider=Microsoft.Mashup.OleDb.1;Data Source=$Workbook$;Location={name};Extended Properties=\"\""
    lo = ws.ListObjects.Add(SourceType=0, Source=conn, Destination=ws.Range("$A$1"))
    lo.QueryTable.CommandType = 2
    lo.QueryTable.CommandText = f"SELECT * FROM [{name}]"
    lo.QueryTable.Refresh(False)
    lo.Name = name
    return lo


def make(jobs, out):
    import vbam_forge as vf
    xl, pid = vf._start_own()
    try:
        xl.DisplayAlerts = False
        for stem, data, kw in jobs:
            before = os.path.join(out, f"{stem}_前.xlsx")
            truth = os.path.join(out, f"{stem}_正解.xlsx")
            stale = kw.pop('stale', False)
            write_before(before, data, **kw)
            if stale:                                        # 前のクエリ（課名の降順・古い）とシートを残しておく
                wb = xl.Workbooks.Open(before)
                try:
                    load_query(wb, formula=M.replace('Order.Ascending', 'Order.Descending'))
                    wb.Worksheets(1).Activate()
                    wb.Save()
                finally:
                    wb.Close(SaveChanges=False)
            wb = xl.Workbooks.Open(before)
            try:
                load_query(wb)
                wb.Worksheets(1).Activate()
                wb.SaveAs(truth, FileFormat=51)
            finally:
                wb.Close(SaveChanges=False)
            print('  ', stem)
    finally:
        xl = None
        vf._quit(pid)


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    make([('未見1_標準', build(s(1), [4, 5], rng.randint(10, 60)), {}),
          ('未見2_4か月と並び違い', build(s(2), [4, 5, 6, 7], 40), {'swap_month': 6}),
          ('未見3_前のクエリとほかの表', build(s(3), [4, 5], 30), {'stale': True, 'other_table': True})], out)
    return out


if __name__ == '__main__':
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        sys.exit(0)
    make([('PQ_本番', build(3, [4, 5], 20), {}),
          ('PQ_試験A', build(13, [4, 5, 6], 800), {}),
          ('PQ_試験B', build(23, [4, 5], 25), {'swap_month': 5}),
          ('PQ_試験C', build(33, [4, 5], 15), {'stale': True}),
          ('PQ_試験D', build(43, [4, 5], 18), {'other_table': True})], HERE)
    shutil.copyfile(os.path.join(HERE, 'PQ_本番_正解.xlsx'), os.path.join(HERE, 'PQ_試験E_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, 'PQ_本番_正解.xlsx'), os.path.join(HERE, 'PQ_試験E_正解.xlsx'))
    open(os.path.join(HERE, 'request.txt'), 'w', encoding='utf-8').write(REQUEST + '\n')
    lines = ['パワークエリで月ごとの支出を追加して課別に集計', '月別の支出明細をパワークエリでまとめて課ごとの合計を',
             'T支出のテーブルをパワークエリで結合して課別集計', 'パワークエリで課別集計のクエリを作って',
             '各月の明細を追加して課名で集計するクエリをお願い', 'パワークエリで PQ課別集計 シートに読み込んで']
    open(os.path.join(HERE, 'phrases.txt'), 'w', encoding='utf-8').write('\n'.join(lines) + '\n')
    print('ok')
