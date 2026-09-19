# -*- coding: utf-8 -*-
"""支出明細から課別・月別の月次集計表を作る（鍛える回路の題材・2026-09-17。毎月の執行状況の報告で作る表）。

仕事: シート「支出明細」の明細の右に 1 列空けて、課×月（年度の 4 月〜3 月）の支出額の集計表を作る。明細は触らない。
正解の決まり:
  - 集計表の見出しは 1 行: 課名・4月・5月・…・12月・1月・2月・3月・合計（14 列。支出の無い月も列を出して 0）。
  - 集計表の左上＝明細の見出しの行・明細の右端の列の 2 つ右（1 列空ける）。
  - 課は明細に出てくる順（重複なし）。課名の前後の空白（半角・全角）は落として同じ課にする。
  - 各月＝その課・その月の支出額の合計（無ければ 0）。右端に行の合計、最後の行に「合計」（月ごとの合計と総合計）。
  - 支出日が文字（2026/04/15・2026年4月15日）でも日付として読む。支出額がカンマ・全角・前後の空白つきの文字でも数で読む。
  - 入れない: 明細の空行・明細の下の合計行（課名か摘要が「合計」「計」の行）。
  - 集計表の場所に前の集計が残っていたら消してから書く（2 回撃っても同じ）。
組（依頼文は同じ）: 本番 60 行・5 課／A 2,000 行／B 日付と金額が文字／C 表題の行つき／D 列の並びが違う／
  E 前の集計（課が多い）が残っている／F 空行と合計行／G 支出の無い月・課名の空白の揺れ／H 行が日付順でない／I 撃った後にもう一度
  py make_pairs.py            → *_前.xlsx と *_正解.xlsx
  py make_pairs.py --unseen N → 未見_種N
"""
import datetime
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
SHEET = '支出明細'
KA = ['総務課', '財政課', '企画課', '税務課', '福祉課', '建設課', '農林課', '教育総務課', '観光課', '環境課', '会計課', '議会事務局']
KAMOKU = ['需用費', '役務費', '委託料', '使用料及び賃借料', '工事請負費', '備品購入費', '負担金補助及び交付金', '旅費']
TEKIYO = ['コピー用紙', '郵送料', '清掃業務', '複合機リース', '修繕工事', 'パソコン購入', '協議会負担金', '出張旅費', '電気料', '広告掲載']
MONTHS = [4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3]
HEAD = ['支出日', '課名', '科目', '摘要', '支出額']
_ZEN = str.maketrans('0123456789,', '０１２３４５６７８９，')


def build(seed, n, n_ka=5, fy=2026, gap_months=(), sort=True):
    rng = random.Random(seed)
    kas = rng.sample(KA, n_ka)
    months = [m for m in MONTHS if m not in gap_months]
    rows = []
    for _ in range(n):
        m = rng.choice(months)
        y = fy if m >= 4 else fy + 1
        d = datetime.date(y, m, rng.randint(1, 28))
        rows.append([d, rng.choice(kas), rng.choice(KAMOKU), rng.choice(TEKIYO), rng.randrange(1, 800) * 100])
    if sort:
        rows.sort(key=lambda x: x[0])
    return rows


def summary(rows):
    order, tot = [], {}
    for d, ka, _k, _t, amt in rows:
        key = ka.strip().strip('　').strip()
        if key not in tot:
            order.append(key)
            tot[key] = {m: 0 for m in MONTHS}
        tot[key][d.month] += amt
    return order, tot


def write_pair(stem, rows, out=HERE, cols=HEAD, title=False, text_date=False, text_amount=False, blank_rows=0,
               total_row=False, stale=0, ka_space=False):
    rng = random.Random(stem)
    order, tot = summary(rows)
    # 空行の位置は前と正解で同じにする（ループの中で引くと前と正解で空行がずれる＝2026-09-17 夜の題材の誤り）
    blanks = set(rng.sample(range(1, max(2, len(rows) - 1)), min(blank_rows, max(0, len(rows) - 2)))) if blank_rows else set()
    for truth in (False, True):
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET
        h = 1
        if title:
            ws.cell(1, 1, '令和8年度 支出明細（一般会計）').font = Font(bold=True, size=14)
            ws.cell(2, 1, '単位：円')
            h = 4
        for j, v in enumerate(cols, 1):
            ws.cell(h, j, v).font = Font(bold=True)
        r = h + 1
        rng2 = random.Random(stem + 'v')
        for i, (d, ka, kamoku, tekiyo, amt) in enumerate(rows):
            if i in blanks:
                r += 1                                        # 空行 1 行（表の続き）
            vals = {'支出日': datetime.datetime(d.year, d.month, d.day), '課名': ka, '科目': kamoku, '摘要': tekiyo, '支出額': amt}
            if ka_space and rng2.random() < 0.3:
                vals['課名'] = ka + rng2.choice((' ', '　', '  '))
            if text_date and rng2.random() < 0.5:
                vals['支出日'] = f"{d.year}/{d.month:02d}/{d.day:02d}" if rng2.random() < 0.5 else f"{d.year}年{d.month}月{d.day}日"
            if text_amount and rng2.random() < 0.5:
                vals['支出額'] = f"{amt:,}".translate(_ZEN) if rng2.random() < 0.5 else f" {amt:,} "
            for j, c in enumerate(cols, 1):
                ws.cell(r, j, vals[c])
            r += 1
        if total_row:
            ws.cell(r, cols.index('摘要') + 1, '合計')
            ws.cell(r, cols.index('支出額') + 1, sum(x[4] for x in rows))
        lc = len(cols) + 2
        if truth:
            ws.cell(h, lc, '課名')
            for j, m in enumerate(MONTHS, 1):
                ws.cell(h, lc + j, f"{m}月")
            ws.cell(h, lc + 13, '合計')
            for i, ka in enumerate(order, 1):
                ws.cell(h + i, lc, ka)
                for j, m in enumerate(MONTHS, 1):
                    ws.cell(h + i, lc + j, tot[ka][m])
                ws.cell(h + i, lc + 13, sum(tot[ka].values()))
            rr = h + len(order) + 1
            ws.cell(rr, lc, '合計')
            for j, m in enumerate(MONTHS, 1):
                ws.cell(rr, lc + j, sum(tot[ka][m] for ka in order))
            ws.cell(rr, lc + 13, sum(sum(tot[ka].values()) for ka in order))
        elif stale:
            ws.cell(h, lc, '課名')
            for j, m in enumerate(MONTHS, 1):
                ws.cell(h, lc + j, f"{m}月")
            ws.cell(h, lc + 13, '合計')
            for i in range(1, stale + 1):
                ws.cell(h + i, lc, KA[i % len(KA)])
                for j in range(1, 14):
                    ws.cell(h + i, lc + j, 1000 * i)
            ws.cell(h + stale + 1, lc, '合計')
        wb.save(os.path.join(out, f"{stem}_{'正解' if truth else '前'}.xlsx"))


SWAP = ['課名', '支出額', '支出日', '摘要', '科目']


def unseen(seed):
    out = os.path.join(HERE, f"未見_種{seed}")
    os.makedirs(out, exist_ok=True)
    rng = random.Random(seed)
    s = lambda k: seed * 10 + k      # noqa: E731
    write_pair('未見1_標準', build(s(1), rng.randint(20, 120), n_ka=rng.randint(2, 8)), out=out)
    write_pair('未見2_文字の日付と金額', build(s(2), 80), text_date=True, text_amount=True, out=out)
    write_pair('未見3_表題と並び違い', build(s(3), 50, n_ka=4), cols=SWAP, title=True, out=out)
    write_pair('未見4_前の集計', build(s(4), 40, n_ka=3), stale=rng.randint(6, 11), out=out)
    write_pair('未見5_空行と合計行', build(s(5), 60), blank_rows=3, total_row=True, out=out)
    write_pair('未見6_月の抜けと空白', build(s(6), 30, n_ka=6, gap_months=(8, 1, 2)), ka_space=True, out=out)
    write_pair('未見7_全部入り', build(s(7), 150, n_ka=7, gap_months=(5,), sort=False), cols=SWAP, title=True, text_date=True,
               text_amount=True, blank_rows=2, total_row=True, stale=10, ka_space=True, out=out)
    write_pair('未見8_大量', build(s(8), 3000, n_ka=12), out=out)
    return out


if __name__ == '__main__':
    if len(sys.argv) > 2 and sys.argv[1] == '--unseen':
        print(unseen(int(sys.argv[2])))
        sys.exit(0)
    write_pair('月次_本番', build(5, 60))
    write_pair('月次_試験A', build(15, 2000, n_ka=10))
    write_pair('月次_試験B', build(25, 70), text_date=True, text_amount=True)
    write_pair('月次_試験C', build(35, 40, n_ka=3), title=True)
    write_pair('月次_試験D', build(45, 50), cols=SWAP)
    write_pair('月次_試験E', build(55, 30, n_ka=3), stale=9)
    write_pair('月次_試験F', build(65, 45), blank_rows=2, total_row=True)
    write_pair('月次_試験G', build(75, 25, n_ka=6, gap_months=(6, 11, 3)), ka_space=True)
    write_pair('月次_試験H', build(85, 60, sort=False))
    shutil.copyfile(os.path.join(HERE, '月次_本番_正解.xlsx'), os.path.join(HERE, '月次_試験I_前.xlsx'))
    shutil.copyfile(os.path.join(HERE, '月次_本番_正解.xlsx'), os.path.join(HERE, '月次_試験I_正解.xlsx'))
    print('ok')
