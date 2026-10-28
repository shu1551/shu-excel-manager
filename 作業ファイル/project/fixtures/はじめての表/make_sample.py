# -*- coding: utf-8 -*-
"""はじめての 10 分の試しの表を作る（README「最初の 1 回（試しの表で）」・2026-10-28）。

汚れた売上の明細を 1 枚、xlsx で書き出す。openpyxl だけで作る（Excel は要らない）。
仕込む汚れ: 空白の揺れ（後ろの空白・全角の空白・空白の有無）／重複行／文字の数字／空行／合計の行（SUM の式）。
作った後に、仕込んだ汚れを番地で出す。seiri の返事と見比べるため。

使い方: py make_sample.py [出力先.xlsx]（既定はこのフォルダの はじめての表.xlsx）
"""
import datetime
import os
import sys

from openpyxl import Workbook

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'はじめての表.xlsx')
SHEET = '売上'
HEAD = ['日付', '担当', '商品', '数量', '金額']

D = datetime.date
# 1 行目が見出し、2 行目から明細。None の行は空行。汚れは (番地の列, 何の汚れ) で印を付ける
ROWS = [
    (D(2026, 10, 1), '佐藤 一郎', 'コピー用紙', 5, 2500),
    (D(2026, 10, 1), '鈴木 花子 ', 'ボールペン', 12, 1440),       # B3 後ろの空白
    (D(2026, 10, 2), '田中 次郎', 'クリアファイル', '20', 1600),    # D4 文字の数字
    (D(2026, 10, 3), '佐藤　一郎', 'ノート', 8, 960),           # B5 全角の空白
    (D(2026, 10, 3), '田中次郎', 'コピー用紙', 3, '1500'),          # B6 空白なし・E6 文字の数字
    (D(2026, 10, 6), '鈴木 花子', 'クリアファイル', 10, 800),
    None,                                                            # 8 行目 空行
    (D(2026, 10, 7), '田中 次郎', 'ボールペン', 6, 720),
    (D(2026, 10, 1), '佐藤 一郎', 'コピー用紙', 5, 2500),          # 10 行目 2 行目と同じ（重複行）
]
DIRT = [
    ('B3', '空白の揺れ（後ろに半角の空白）', '「鈴木 花子 」→「鈴木 花子」'),
    ('B5', '空白の揺れ（姓と名の間が全角の空白）', '「佐藤　一郎」→「佐藤 一郎」'),
    ('B6', '空白の揺れ（姓と名の間に空白が無い）', '「田中次郎」→「田中 次郎」'),
    ('D4', '文字の数字', '文字の「20」→ 数の 20'),
    ('E6', '文字の数字', '文字の「1500」→ 数の 1,500'),
    ('行8', '空行', '行ごと消して詰める'),
    ('行10', '重複行（行2 と全部の列が同じ）', '行ごと消す'),
]


def _num(v):
    return float(v) if isinstance(v, str) else v


def totals():
    """合計の行の値: 開いた直後（SUM は文字の数字を数えない）と、直った後（文字の数字を数に・重複行を除く）。"""
    body = [r for r in ROWS if r is not None]
    before = [sum(r[k] for r in body if not isinstance(r[k], str)) for k in (3, 4)]
    seen, kept = set(), []
    for r in body:
        key = tuple(str(v).replace('　', ' ').strip() for v in r)
        if key not in seen:
            seen.add(key)
            kept.append(r)
    after = [sum(_num(r[k]) for r in kept) for k in (3, 4)]
    return before, after


def make(path):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(HEAD)
    for r in ROWS:
        ws.append(list(r) if r else [])
    last = 1 + len(ROWS)
    total = last + 1
    ws.cell(total, 1, '合計')
    ws.cell(total, 4, f'=SUM(D2:D{last})')
    ws.cell(total, 5, f'=SUM(E2:E{last})')
    for row in ws.iter_rows(min_row=2, max_row=last, max_col=1):
        for c in row:
            if c.value is not None:
                c.number_format = 'yyyy/m/d'
    wb.calculation.fullCalcOnLoad = True        # 式の値は Excel が開いたときに計算する
    wb.save(path)
    return total


def main(argv):
    path = os.path.abspath(argv[0]) if argv else OUT
    total = make(path)
    before, after = totals()
    try:
        shown = os.path.relpath(path)
    except ValueError:                           # 別のドライブに書いたとき
        shown = path
    print(f'作りました: {shown}（シート「{SHEET}」・見出し 1 行目・明細 2〜{total - 1} 行目・合計 {total} 行目）')
    print('仕込んだ汚れ:')
    for addr, what, fix in DIRT:
        print(f'  {addr:<5} {what}  … {fix}')
    print(f'  行{total}  合計の行（D{total}・E{total} は SUM の式）… 残す')
    print(f'合計の行の値: 開いた直後 数量 {before[0]:,.0f}・金額 {before[1]:,.0f}（SUM は文字の数字を数えない）'
          f' → 直った後 数量 {after[0]:,.0f}・金額 {after[1]:,.0f}（文字の数字を数え、重複行を除く）')


if __name__ == '__main__':
    main(sys.argv[1:])
